"""The LIVE «Verifix to'g'irlash» (lab) — `/staff-live`, the /staff API over Verifix.

From 2026-10-04 (the operator: «structure this page just like Verifix edit …
the same rule applies for everything as that page … the only difference should
be the source») every door here is the twin of a /api/staff door — same path
under `/api/staff-live`, same request and response shapes, same rights — so the
page renders /staff's own components with a different API base. What differs:

* **the source** — a unit's day is the stored live Verifix read laid out by
  `services/live_staff` (a job reads Verifix every minute), never `attendance`;
* **a worker is a Verifix `employee_id`**, so every picker sends ids, not names
  (`worker_names` on the deletion door carries them too — it is the row key);
* **nothing is applied** — an approved document is read on every request, so
  an approve, an un-post or a delete writes the document and nothing else;
* **who may be exchanged or re-titled**: only a worker who CAME (has a clock-in),
  whether still inside or already gone (the operator, 2026-10-04).

Lab for now and built to replace /staff: the documents live in `live_documents`
/ `live_document_history` / `live_deletions` / `live_placements` /
`live_day_closes`, and nothing else on the platform reads them. Page key
`staff-live`, admin-only until the operator opens it on the Access tab; inside,
each role has exactly /staff's rights (own unit for a supervisor, shift ∩ plant
for a shift-manager). Its notifications (bell + Telegram) start once the page is
opened to anybody but admins, and reach only people who can open it.
"""
import logging
import re
from collections import defaultdict
from datetime import date, datetime, timedelta, timezone
from typing import List, Optional
from uuid import uuid4

from fastapi import APIRouter, Depends, HTTPException, Request
from pydantic import BaseModel, Field
from sqlalchemy import func
from sqlalchemy.orm import Session
from sqlalchemy.orm.attributes import flag_modified

from app import identity
from app.capabilities import (
    CAP_ATTENDANCE_DELETE, CAP_DAY_REOPEN, CAP_DOCUMENTS_APPROVE, CAP_REQUESTS_APPROVE,
    cap_scope, has_cap, page_scope_is_all, profile_unit_ids, scope_is_all,
)
from app.database import get_db
from app.models import (
    Admin, ExchangeTask, LiveAllLeftNotice, LiveDayClose, LiveDeletion, LiveDocument,
    LiveDocumentHistory, LivePlacement, LiveVerifixRead, Manager,
)
from app.notify_ctx import notifications_suppressed
from app.permissions import get_page_access, require_page
from app.services import action_log, live_staff, shift_scope, verifix_live
from app.xlsx_delivery import deliver_xlsx

PAGE = live_staff.PAGE
router = APIRouter(prefix="/api/staff-live", tags=["staff-live"])
log = logging.getLogger(__name__)
_page = require_page(PAGE)

_HHMM = re.compile(r"^\s*([01]?\d|2[0-3])\s*[:.\-]\s*([0-5]\d)\s*$")
# /staff's bound: a draft dated longer ago cannot be posted (staff.STALE_APPROVE_DAYS).
STALE_APPROVE_DAYS = 14
DOC_TYPES = ("people_exchange", "role_change")


# ── who is asking, and what they may see ─────────────────────────────────────

def _tg(caller: dict) -> int:
    return int(caller["sub"])


def _who(caller: dict) -> str:
    return caller.get("full_name") or caller.get("username") or str(caller.get("sub") or "")


def _sees_all(db: Session, caller: dict) -> bool:
    """A personal page grant at "all" widens a supervisor to every unit (the
    /staff rule, `staff._staff_sees_all`); admins always pass."""
    return page_scope_is_all(db, caller, PAGE)


def _visible(db: Session, caller: dict) -> Optional[list[int]]:
    """Units the caller may read. None = every unit (admin / grant at "all")."""
    role, rid = caller.get("role"), caller.get("role_id")
    if role == "admin" or _sees_all(db, caller):
        return None
    if role == "supervisor":
        return [rid] if rid else []
    if role == "shift-manager":
        return shift_scope.unit_ids(db, rid)
    return []


def _read_unit(db: Session, caller: dict, manager_id: Optional[int]) -> int:
    """The unit a read targets — a supervisor is pinned to their own unless a
    grant at "all" widens them (`staff._staff_target_manager`)."""
    mid = manager_id
    if caller.get("role") == "supervisor" and not (manager_id and _sees_all(db, caller)):
        mid = caller.get("role_id")
    if not mid:
        raise HTTPException(status_code=400, detail="manager_id required")
    vis = _visible(db, caller)
    if vis is not None and mid not in vis:
        raise HTTPException(status_code=403, detail="Not allowed for this manager")
    return mid


def _caller_units(db: Session, caller: dict) -> Optional[list[int]]:
    return profile_unit_ids(db, identity.viewer_profile_key(db, caller))


def _cap_covers(db: Session, caller: dict, capability: str, manager_id: Optional[int]) -> bool:
    """An admin, or a grant that reaches the unit (`staff._cap_covers_unit`)."""
    if caller.get("role") == "admin":
        return True
    scope = cap_scope(db, caller, capability)
    if scope is None:
        return False
    if scope == "all":
        return True
    units = _caller_units(db, caller)
    return units is None or manager_id in units


def _day_of(raw: Optional[str]) -> Optional[date]:
    if not raw:
        return None
    try:
        return date.fromisoformat(str(raw)[:10])
    except ValueError:
        raise HTTPException(status_code=400, detail="Invalid date format")


def _hhmm(raw: Optional[str]) -> Optional[str]:
    m = _HHMM.match(str(raw or ""))
    return f"{int(m.group(1)):02d}:{m.group(2)}" if m else None


def _closed(db: Session, mid: int, d: date) -> Optional[LiveDayClose]:
    return db.query(LiveDayClose).filter(LiveDayClose.manager_id == mid,
                                         LiveDayClose.day == d).first()


def _assert_open(db: Session, mid: int, d: date) -> None:
    if _closed(db, mid, d):
        raise HTTPException(status_code=409,
                            detail="Day is closed — changes can no longer be submitted for this date")


def _units(db: Session) -> dict:
    return {m.id: m for m in db.query(Manager).all()}


def _unit_name(db: Session, mid: Optional[int]) -> Optional[str]:
    m = db.query(Manager).filter(Manager.id == mid).first() if mid else None
    return m.name if m else None


# ── the unit's day ───────────────────────────────────────────────────────────

class _NoDay(HTTPException):
    """The day could not be built (no Verifix login, the read failed …)."""

    def __init__(self, rd: dict):
        code = rd.get("error")
        status = 409 if code in ("not_configured", "no_cells", "no_unit") else 424
        super().__init__(status_code=status, detail={
            "code": code, "message": rd.get("message") or code})


def _build(db: Session, mid: int, d: Optional[date], force: bool = False):
    """(read, ctx, unit day) — or raise `_NoDay`."""
    rd = verifix_live.day_read(db, mid, d, force=force)
    if rd.get("error"):
        raise _NoDay(rd)
    ctx = live_staff.load(db, rd["day"], rd["directory"], rd["store"], rd["now"])
    return rd, ctx, live_staff.unit_day(ctx, mid)


def _named(ud: dict) -> dict:
    """employee id → the worker's primary row on this unit's day."""
    return {r["employee_id"]: r for r in ud["workers"] if r.get("split_of") is None}


# ── notifications ────────────────────────────────────────────────────────────
# Bell + Telegram DM (the operator: «as soon as we open this page for
# supervisors»): nothing is sent while the page is admin-only, and once it is
# opened, only people who can open it are told.

def _notify(db: Session, nkey: str, params: dict, *, units: tuple = (), admins: bool = True,
            supervisors: tuple = (), profiles: tuple = (), actor: Optional[int] = None,
            subject=None, ntype: str = "info") -> None:
    """One bell row per addressed PROFILE + a DM per holder (the
    `_notify_all_parties` shape): admins, the shift-managers answerable for
    `units`, the supervisors of `supervisors`, and any `profiles`; the actor's
    own profile is skipped. Never raises — a notice must not undo a decision."""
    if notifications_suppressed():
        return
    try:
        if not live_staff.lab_open(db):
            return
        from app.routers.staff import _get_user_lang, _mk_notif, _mk_notif_tg, _notify as bell
        from app.services import notification_center as notif_center
        keys: set = set(profiles)
        if admins:
            keys |= {identity.profile_key("admin", a.profile_id)
                     for a in db.query(Admin).all() if a.profile_id}
        for u in units:
            keys |= {identity.profile_key("shift-manager", rid)
                     for rid in shift_scope.role_ids_for_unit(db, u)}
        keys |= {identity.profile_key("supervisor", s) for s in supervisors if s}
        keys.discard(None)
        access = get_page_access(db)
        actor_keys = set()
        if actor:
            from app.models import TelegramUserRole
            for r in db.query(TelegramUserRole).filter(TelegramUserRole.telegram_id == actor,
                                                       TelegramUserRole.status == "approved").all():
                actor_keys.add(r.profile_key or identity.role_row_profile_key(db, r))
        dmed: set = set()
        for key in sorted(k for k in keys if live_staff.can_open(db, k, access)):
            if key in actor_keys:
                continue
            holders = identity.profile_holders(db, key)
            if not holders:
                bell(db, None, type=ntype, nkey=nkey, params=params, profile=key, subject=subject)
                continue
            bell(db, holders[0], type=ntype, dm=False, nkey=nkey, params=params,
                 profile=key, subject=subject)
            if notif_center.telegram_muted(db, key, nkey):
                continue
            from app.telegram_bot import send_tg_notification
            for tid in holders:
                if tid == actor or tid in dmed:
                    continue
                dmed.add(tid)
                lang = _get_user_lang(db, tid)
                title, body = _mk_notif(nkey, params, lang)
                try:
                    send_tg_notification(tid, title, body, html=_mk_notif_tg(nkey, params, lang))
                except Exception:  # noqa: BLE001
                    pass
        db.commit()
    except Exception:  # noqa: BLE001 — a notice that fails must not undo the decision
        db.rollback()
        log.exception("staff-live: notification %s failed", nkey)


def _creator_profile(db: Session, doc: LiveDocument) -> Optional[str]:
    if doc.created_by_role == "supervisor":
        return identity.profile_key("supervisor", doc.manager_id)
    from app.models import TelegramUserRole
    r = (db.query(TelegramUserRole)
         .filter(TelegramUserRole.telegram_id == doc.created_by_telegram_id,
                 TelegramUserRole.role == doc.created_by_role,
                 TelegramUserRole.status == "approved").first())
    return (r.profile_key or identity.role_row_profile_key(db, r)) if r else None


def _doc_params(doc: LiveDocument) -> dict:
    pl = doc.payload or {}
    target = (pl.get("target_manager_name") if pl.get("target_type") == "supervisor"
              else pl.get("task_name")) or "—"
    return {"actor_name": doc.created_by_name or "", "count": len(pl.get("employees") or []),
            "target": target,
            "target_kind": "supervisor" if pl.get("target_type") == "supervisor" else "task",
            "new_role": pl.get("new_role") or "", "date": doc.day,
            # A line whose one value is blank is dropped — a whole-day move.
            "time": ((pl.get("transfer_time") or "")
                     + (f"–{pl['return_time']}" if pl.get("return_time") else ""))}


def _notify_doc(db: Session, doc: LiveDocument, event: str, actor: int) -> None:
    pl = doc.payload or {}
    exch = doc.doc_type == "people_exchange"
    nkey = {
        ("people_exchange", "created"): "live_exchange_created",
        ("people_exchange", "approved"): "live_exchange_approved",
        ("people_exchange", "cancelled"): "live_exchange_cancelled",
        ("role_change", "created"): "live_role_change_new",
        ("role_change", "approved"): "live_role_change_approved",
        ("role_change", "cancelled"): "live_role_change_cancelled",
    }.get((doc.doc_type, event))
    if not nkey:
        return
    target = pl.get("target_manager_id") if exch and pl.get("target_type") == "supervisor" else None
    # A draft waits on its approvers; a decision is news to the units too.
    sups = (doc.manager_id, target) if event != "created" else ((target,) if target else ())
    _notify(db, nkey, _doc_params(doc), units=(doc.manager_id,), supervisors=sups,
            actor=actor, subject=("live_doc", doc.id))


# ── meta ─────────────────────────────────────────────────────────────────────

@router.get("/supervisors")
def list_supervisors(db: Session = Depends(get_db), caller: dict = Depends(_page)):
    """The units this page covers — those owning a cell counted in the загрузка
    — that the caller may read (`staff.list_supervisors`'s shape)."""
    cells, units = verifix_live._registry(db)
    with_cells = {c["manager_id"] for c in cells.values() if c["manager_id"] in units}
    vis = _visible(db, caller)
    out = [{"manager_id": mid, "full_name": u["name"], "shift": u["shift"]}
           for mid, u in units.items() if mid in with_cells and (vis is None or mid in vis)]
    return sorted(out, key=lambda u: (u["shift"] or 9, u["full_name"] or ""))


@router.get("/field-options")
def field_options(db: Session = Depends(get_db), _: dict = Depends(_page)):
    """Job titles — Verifix's own list (the directory) — and the subset a role
    change may name (`staff.is_assignable_target_role`)."""
    from app.routers.staff import is_assignable_target_role
    jobs: list = []
    row = db.query(LiveVerifixRead).filter(LiveVerifixRead.key == "dir").first()
    if row and row.data:
        jobs = sorted({j for j in (row.data.get("jobs") or {}).values() if j})
    return {"job_titles": jobs, "schedules": [],
            "assignable_job_titles": [j for j in jobs if is_assignable_target_role(j)]}


@router.get("/exchange-targets")
def exchange_targets(attend_date: str, manager_id: Optional[int] = None,
                     db: Session = Depends(get_db), caller: dict = Depends(_page)):
    """Units a worker may move INTO on the day: every unit this page covers,
    except the sender and any unit whose day is already closed."""
    if caller.get("role") not in ("admin", "supervisor"):
        raise HTTPException(status_code=403, detail="Admin or supervisor only")
    d = _day_of(attend_date)
    sender = caller.get("role_id") if caller.get("role") == "supervisor" else manager_id
    cells, units = verifix_live._registry(db)
    with_cells = {c["manager_id"] for c in cells.values() if c["manager_id"] in units}
    closed = {m for (m,) in db.query(LiveDayClose.manager_id).filter(LiveDayClose.day == d)}
    out = [{"manager_id": mid, "full_name": u["name"], "shift": u["shift"]}
           for mid, u in units.items() if mid in with_cells and mid != sender and mid not in closed]
    return sorted(out, key=lambda u: (u["shift"] or 9, u["full_name"] or ""))


@router.get("/tasks")
def list_tasks(db: Session = Depends(get_db), _: dict = Depends(_page)):
    """The shared task list a worker may be sent to (/staff's — one register)."""
    names = [t.name for t in db.query(ExchangeTask).filter(ExchangeTask.active.is_(True))
             .order_by(func.lower(ExchangeTask.name)).all()]
    return {"tasks": names}


class TaskDeleteBody(BaseModel):
    name: str


@router.post("/tasks/delete")
def delete_task(body: TaskDeleteBody, db: Session = Depends(get_db), caller: dict = Depends(_page)):
    """Admin-only, the same soft removal as /staff's (it is one shared list)."""
    if caller.get("role") != "admin":
        raise HTTPException(status_code=403, detail="Admin only")
    name = (body.name or "").strip()
    t = db.query(ExchangeTask).filter(ExchangeTask.name == name).first()
    if not t or not t.active:
        raise HTTPException(status_code=404, detail="Task not found")
    t.active = False
    db.commit()
    action_log.enrich(target_kind="task", target_id=t.id, target_name=name,
                      details=[("task", name)], changes=[("enabled", True, False)])
    return {"ok": True}


# ── attendance (the Workers tab) ─────────────────────────────────────────────

@router.get("/attendance")
def get_attendance(attend_date: Optional[str] = None, manager_id: Optional[int] = None,
                   force: bool = False, db: Session = Depends(get_db),
                   caller: dict = Depends(_page)):
    """The unit's live day in /staff's `GET /attendance` shape, plus `live` —
    what only a live source can say (the read's age, the counts, the close)."""
    mid = _read_unit(db, caller, manager_id)
    rd = verifix_live.day_read(db, mid, _day_of(attend_date), force=force)
    if rd.get("error"):
        return {"manager_id": mid, "manager_name": _unit_name(db, mid), "error": rd["error"],
                "message": rd.get("message"), "workers": [], "cells": [], "extra_hours": 0}
    day = rd["day"]
    ctx = live_staff.load(db, day, rd["directory"], rd["store"], rd["now"])
    ud = live_staff.unit_day(ctx, mid)
    close_rec = _closed(db, mid, day)
    notice = (db.query(LiveAllLeftNotice)
              .filter(LiveAllLeftNotice.manager_id == mid, LiveAllLeftNotice.day == day).first())
    pending = live_staff.pending_count(db, mid, day)
    unplaced = live_staff.unplaced(ud["workers"])
    is_admin = caller.get("role") == "admin"
    workers = ud["workers"] if is_admin else [{k: v for k, v in r.items() if k != "raw"}
                                              for r in ud["workers"]]
    formula = ctx.formula
    return {
        "manager_id": mid, "manager_name": rd["unit"]["name"], "date": day.isoformat(),
        "workers": workers,
        "cells": live_staff.cells_catalog(db, ctx, mid, ud["workers"]),
        "extra_hours": ud["extra_hours"], "extras": ud["extras"],
        "live": {
            "day": day.isoformat(), "today": rd["today"].isoformat(),
            "is_today": day == rd["today"], "shift": rd["unit"]["shift"],
            "now": live_staff._iso(rd["now"]),
            "pulled_at": live_staff._iso(rd["covered"]),
            "directory_at": live_staff._iso(rd["dir_at"]),
            "read_error": rd["read_error"],
            "auto": {"on": day == rd["today"], "every": verifix_live.HOT_S},
            "counts": ud["counts"],
            "close": live_staff.close_state(ud, day, close_rec, notice, pending),
            "needs_cell": len(unplaced),
            "formula": ({"names": formula["names"], "share": formula["share"]} if formula else None),
            "rules": {"missing_after": verifix_live.MISSING_AFTER_MIN,
                      "late_grace": verifix_live.LATE_GRACE_MIN,
                      "early_grace": verifix_live.EARLY_GRACE_MIN,
                      "min_moved_hours": live_staff.MIN_MOVED_HOURS},
            # What the read held — admins only, the page's Diagnostics line.
            "diag": live_staff.diag(ctx, mid, ud["workers"]) if is_admin else None,
        },
    }


# ── documents ────────────────────────────────────────────────────────────────

def _scope_docs(q, caller: dict, db: Session):
    """`staff._scope_documents`: own unit + exchanges addressed TO it for a
    supervisor, shift ∩ plant for a shift-manager, everything for an admin or
    a documents grant at "all"."""
    if scope_is_all(db, caller, CAP_DOCUMENTS_APPROVE):
        return q
    role, rid = caller.get("role"), caller.get("role_id")
    if role == "supervisor":
        if not rid:
            return q.filter(LiveDocument.created_by_telegram_id == _tg(caller))
        incoming = (LiveDocument.doc_type == "people_exchange") & (
            LiveDocument.payload["target_manager_id"].astext == str(rid))
        return q.filter((LiveDocument.manager_id == rid) | incoming)
    if role == "shift-manager":
        return q.filter(LiveDocument.manager_id.in_(shift_scope.unit_ids(db, rid) or [-1]))
    if role == "admin":
        return q
    return q.filter(LiveDocument.id < 0)


def _native_can_approve(doc: LiveDocument, caller: dict, db: Session) -> bool:
    """`staff._native_can_approve_doc`: role change → admin or shift-manager;
    exchange → to a supervisor: admin or the RECEIVING supervisor; to a task:
    admin or a shift-manager answerable for the sending unit."""
    role = caller.get("role")
    if role == "admin":
        return True
    if doc.doc_type != "people_exchange":
        return role == "shift-manager"
    pl = doc.payload or {}
    if pl.get("target_type") == "supervisor":
        return role == "supervisor" and caller.get("role_id") == pl.get("target_manager_id")
    return role == "shift-manager" and shift_scope.covers(db, caller.get("role_id"), doc.manager_id)


def _granted_over(doc: LiveDocument, caller: dict, db: Session) -> bool:
    scope = cap_scope(db, caller, CAP_DOCUMENTS_APPROVE)
    if scope is None:
        return False
    if scope == "all":
        return True
    units = _caller_units(db, caller)
    if units is None or doc.manager_id in units:
        return True
    pl = doc.payload or {}
    return doc.doc_type == "people_exchange" and pl.get("target_manager_id") in units


def _can_approve(doc: LiveDocument, caller: dict, db: Session) -> bool:
    return _native_can_approve(doc, caller, db) or _granted_over(doc, caller, db)


def _is_creator(doc: LiveDocument, caller: dict) -> bool:
    if caller.get("role") == "supervisor" and caller.get("role_id") == doc.manager_id:
        return True
    return doc.created_by_telegram_id == _tg(caller)


def _may_reject(doc: LiveDocument, caller: dict, db: Session) -> bool:
    return (_can_approve(doc, caller, db) or caller.get("role") in ("admin", "shift-manager")
            or _is_creator(doc, caller))


def _history(db: Session, doc: LiveDocument, action: str, caller: dict, detail: Optional[dict] = None):
    db.add(LiveDocumentHistory(document_id=doc.id, action=action, actor_telegram_id=_tg(caller),
                               actor_name=_who(caller), detail=detail))


def _serialize(doc: LiveDocument, unit_name: Optional[str] = None, detailed: bool = False) -> dict:
    pl = doc.payload or {}
    emps = pl.get("employees") or []
    out = {
        "id": doc.id, "doc_type": doc.doc_type, "doc_type_label": doc.doc_type,
        "manager_id": doc.manager_id, "supervisor_name": doc.supervisor_name or unit_name,
        "date": doc.day.isoformat() if doc.day else None,
        "status": doc.status, "approved": doc.status == "approved",
        "new_role": pl.get("new_role"),
        "target_type": pl.get("target_type"), "target_manager_id": pl.get("target_manager_id"),
        "target_manager_name": pl.get("target_manager_name"), "task_name": pl.get("task_name"),
        "transfer_time": pl.get("transfer_time"), "return_time": pl.get("return_time"),
        "employee_count": len(emps),
        "created_by_telegram_id": doc.created_by_telegram_id,
        "created_by_name": doc.created_by_name,
        "approved_by_name": doc.approved_by_name,
        "created_at": doc.created_at.isoformat() if doc.created_at else None,
        "approved_at": doc.approved_at.isoformat() if doc.approved_at else None,
    }
    if detailed:
        out["employees"] = emps
        out["payload"] = pl
    return out


def _log_doc(doc: LiveDocument, changes: Optional[list] = None) -> dict:
    pl = doc.payload or {}
    names = [e.get("worker_name") or "" for e in pl.get("employees") or []]
    details = [("doc_type", doc.doc_type), ("unit", doc.supervisor_name or "—"),
               ("date", str(doc.day))]
    if pl.get("target_type"):
        details.append(("target", pl.get("target_manager_name") or pl.get("task_name") or "—"))
    if pl.get("transfer_time"):
        details.append(("time", pl["transfer_time"] + (f"–{pl['return_time']}" if pl.get("return_time") else "")))
    if pl.get("new_role"):
        details.append(("role", pl["new_role"]))
    if names:
        details.append(("workers", f"{len(names)}: " + ", ".join(names[:10])
                        + (f" +{len(names) - 10}" if len(names) > 10 else "")))
    return {"target_kind": "live_document", "target_id": doc.id,
            "target_name": names[0] if len(names) == 1 else None,
            "unit_id": doc.manager_id, "unit_name": doc.supervisor_name, "day": doc.day,
            "details": details, "changes": list(changes or [])}


def _scope_deletions(caller: dict, db: Session) -> list:
    q = db.query(LiveDeletion)
    if not scope_is_all(db, caller, CAP_REQUESTS_APPROVE):
        role, rid = caller.get("role"), caller.get("role_id")
        if role == "supervisor":
            q = q.filter(LiveDeletion.manager_id == rid) if rid else \
                q.filter(LiveDeletion.supervisor_telegram_id == _tg(caller))
        elif role == "shift-manager":
            q = q.filter(LiveDeletion.manager_id.in_(shift_scope.unit_ids(db, rid) or [-1]))
        elif role != "admin":
            return []
    return q.order_by(LiveDeletion.day.desc(), LiveDeletion.id).all()


@router.get("/documents")
def list_documents(db: Session = Depends(get_db), caller: dict = Depends(_page)):
    """Documents + deletion batches, newest first (`staff.list_documents`)."""
    names = {m.id: m.name for m in db.query(Manager).all()}
    docs = [_serialize(d, names.get(d.manager_id)) for d in
            _scope_docs(db.query(LiveDocument), caller, db).order_by(LiveDocument.created_at.desc()).all()]
    batches: dict = defaultdict(list)
    for r in _scope_deletions(caller, db):
        batches[r.batch_id or f"solo-{r.id}"].append(r)
    items = []
    for key, reqs in batches.items():
        reqs.sort(key=lambda r: r.id)
        first = reqs[0]
        has_pending = any(r.status == "pending" for r in reqs)
        applied = any(r.status in ("approved", "undone") for r in reqs)
        status = "pending" if has_pending else ("approved" if applied else "rejected")
        created = min((r.created_at for r in reqs if r.created_at), default=None)
        items.append({
            "id": first.id, "batch_id": first.batch_id, "_source": "deletion",
            "doc_type": "deletion", "doc_type_label": "Deletion request",
            "manager_id": first.manager_id, "manager_name": names.get(first.manager_id, "—"),
            "supervisor_name": names.get(first.manager_id) or first.supervisor_name or "—",
            "supervisor_telegram_id": first.supervisor_telegram_id,
            "date": first.day.isoformat(), "status": status, "approved": status == "approved",
            "new_role": None, "employee_count": len(reqs),
            "created_by_name": first.supervisor_name,
            "approved_by_name": next((r.processed_by_name for r in reqs if r.processed_by_name), None)
            if status != "pending" else None,
            "created_at": created.isoformat() if created else first.day.isoformat(),
            "workers": [{"id": r.id, "worker_name": r.worker_name, "employee_id": r.employee_id,
                         "status": r.status, "approved_by_name": r.processed_by_name,
                         "original": r.original or {}} for r in reqs],
        })
    combined = docs + items
    combined.sort(key=lambda x: x.get("created_at") or x.get("date") or "", reverse=True)
    return combined


@router.get("/documents/pending-count")
def documents_pending_count(db: Session = Depends(get_db), caller: dict = Depends(_page)):
    n = _scope_docs(db.query(LiveDocument), caller, db).filter(LiveDocument.status == "draft").count()
    batches = {r.batch_id or f"solo-{r.id}" for r in _scope_deletions(caller, db) if r.status == "pending"}
    return {"count": n + len(batches)}


def _doc_or_404(db: Session, caller: dict, doc_id: int) -> LiveDocument:
    doc = _scope_docs(db.query(LiveDocument), caller, db).filter(LiveDocument.id == doc_id).first()
    if not doc:
        raise HTTPException(status_code=404, detail="Document not found")
    return doc


@router.get("/documents/{doc_id}")
def get_document(doc_id: int, db: Session = Depends(get_db), caller: dict = Depends(_page)):
    doc = _doc_or_404(db, caller, doc_id)
    return _serialize(doc, _unit_name(db, doc.manager_id), detailed=True)


@router.get("/documents/{doc_id}/history")
def document_history(doc_id: int, db: Session = Depends(get_db), caller: dict = Depends(_page)):
    _doc_or_404(db, caller, doc_id)
    rows = (db.query(LiveDocumentHistory).filter(LiveDocumentHistory.document_id == doc_id)
            .order_by(LiveDocumentHistory.created_at.asc()).all())
    return [{"id": h.id, "action": h.action, "actor_name": h.actor_name, "detail": h.detail,
             "created_at": h.created_at.isoformat() if h.created_at else None} for h in rows]


class DocCreateBody(BaseModel):
    doc_type: str = "role_change"
    attend_date: str
    manager_id: Optional[int] = None
    employees: List[str]                      # Verifix employee ids
    new_role: Optional[str] = None
    target_type: Optional[str] = None         # "supervisor" | "task"
    target_manager_id: Optional[int] = None
    task_name: Optional[str] = None
    transfer_time: Optional[str] = None       # "HH:MM"
    return_time: Optional[str] = None         # "HH:MM"


def _resolve_unit(caller: dict, db: Session, manager_id: Optional[int]) -> tuple[int, Optional[str]]:
    """The unit a document is filed for: a supervisor's own, else `manager_id`."""
    if caller.get("role") == "supervisor":
        mid = caller.get("role_id")
        if not mid:
            raise HTTPException(status_code=400, detail="Supervisor has no linked manager")
    else:
        if not manager_id:
            raise HTTPException(status_code=400, detail="manager_id required")
        mid = manager_id
    return mid, _unit_name(db, mid)


def _picked(ud: dict, ids: List[str], *, came: bool) -> list:
    """The picked workers' rows on the unit's day; with `came`, every one of
    them must have a clock-in (the operator: only those can be exchanged or
    re-titled)."""
    rows = _named(ud)
    out = []
    for eid in dict.fromkeys(str(i) for i in ids):
        r = rows.get(eid)
        if not r:
            continue
        if came and not r.get("clock_in"):
            raise HTTPException(status_code=400, detail=f"{r['worker_name']}: ishga kelmagan — "
                                                        f"faqat kelgan xodimni ko'chirish yoki lavozimini o'zgartirish mumkin")
        out.append(r)
    if not out:
        raise HTTPException(status_code=400, detail="None of the selected workers are on this unit's day")
    return out


def _movers(ctx, ud: dict, mid: int, ids: List[str], T: Optional[datetime]) -> list:
    """The picked workers a people-exchange from THIS unit may move — each must
    have come, and must stand in this unit at the move's time (the operator,
    2026-10-04: the unit where the worker IS files the next move — from its
    named rows or from the nameless hours it carries). A whole-day move (no
    time) is for a worker nobody has moved yet."""
    rows = _named(ud)
    extras = {x["employee_id"]: x for x in ud["extras"]}
    out = []
    for eid in dict.fromkeys(str(i) for i in ids):
        r = rows.get(eid) or extras.get(eid)
        if not r:
            continue
        w = live_staff.person(ctx, eid)
        if w.p["in"] is None:
            raise HTTPException(status_code=400, detail=(
                f"{w.name}: ishga kelmagan — faqat kelgan xodimni ko'chirish mumkin"))
        if T is None and w.moved:
            raise HTTPException(status_code=400, detail=(
                f"{w.name}: bugun allaqachon ko'chirilgan — ko'chirish vaqtini ko'rsating"))
        at = live_staff.where_at(ctx, eid, T)
        if at is None or at.unit != mid:
            when = live_staff._hm(T) if T else "kun boshida"
            raise HTTPException(status_code=400, detail=(
                f"{w.name}: {when} «{live_staff.stint_name(ctx, at)}»da — uni o'sha yerdan "
                f"ko'chirish hujjatini o'sha brigadir tuzadi"))
        out.append({"employee_id": eid, "worker_name": w.name, "job_title": w.role,
                    "verifix_code": r.get("verifix_code"),
                    "clock_in": live_staff._hm(at.start), "in_at": live_staff._iso(at.start),
                    "clock_out": r.get("clock_out")})
    if not out:
        raise HTTPException(status_code=400, detail="None of the selected workers are on this unit's day")
    return out


def _ctx_for(db: Session, doc: LiveDocument, reads: Optional[dict] = None):
    """The doc's day as stored (cached per unit-day inside one request), or None
    when there is no read to check against."""
    reads = reads if reads is not None else {}
    key = (doc.manager_id, doc.day)
    rd = reads.get(key) or verifix_live.day_read(db, doc.manager_id, doc.day)
    if rd.get("error"):
        return None
    reads[key] = rd
    db.flush()
    return live_staff.load(db, rd["day"], rd["directory"], rd["store"], rd["now"])


def _doc_list(ids) -> str:
    return ", ".join("№" + str(i) for i in sorted(ids))


def _still_here(db: Session, doc: LiveDocument, reads: Optional[dict] = None) -> None:
    """At approval: the move still holds — every worker stands in the sender's
    unit at the move's time (a document approved since it was filed may have
    moved them on: 409 `not_here`), and approving it breaks no move already
    approved (one that sends the worker on from the unit this one takes them
    out of: 409 `breaks`)."""
    if doc.doc_type != "people_exchange":
        return
    ctx = _ctx_for(db, doc, reads)
    if ctx is None:
        return                      # nothing to check against — creation already did
    eids = live_staff._doc_eids(doc)
    before = live_staff.broken_moves(ctx, eids)
    after = live_staff.broken_moves(ctx, eids, add=doc)
    pl = doc.payload or {}
    if doc.id in after:
        T = live_staff._at(ctx, (ctx.units.get(doc.manager_id) or {}).get("shift"), pl.get("transfer_time"))
        for e in pl.get("employees") or []:
            at = live_staff.where_at(ctx, str(e.get("employee_id")), T)
            if at is not None and at.unit != doc.manager_id:
                raise HTTPException(status_code=409, detail={
                    "code": "not_here", "worker": e.get("worker_name"),
                    "message": (f"{e.get('worker_name')}: {pl.get('transfer_time') or 'kun boshida'} "
                                f"«{live_staff.stint_name(ctx, at)}»da — bu hujjat tuzilgandan keyin "
                                f"boshqa hujjat uni ko'chirgan")})
    newly = after - before - {doc.id}
    if newly:
        raise HTTPException(status_code=409, detail={
            "code": "breaks", "ids": sorted(newly),
            "message": (f"Bu ko'chirish {_doc_list(newly)} hujjatni buzadi: o'sha hujjat xodimni "
                        f"keyinroq shu brigadirdan ko'chiradi. Vaqtni yoki qaytish vaqtini tekshiring.")})


def _dependents(db: Session, doc: LiveDocument) -> set:
    """Approved moves that hold only because THIS one does — the receiver moved
    the worker on. Un-posting this one would leave them moving a worker who was
    never there, so the un-post is refused until they are un-posted first."""
    if doc.doc_type != "people_exchange" or doc.status != "approved":
        return set()
    ctx = _ctx_for(db, doc)
    if ctx is None:
        return set()
    eids = live_staff._doc_eids(doc)
    return (live_staff.broken_moves(ctx, eids, drop=(doc.id,))
            - live_staff.broken_moves(ctx, eids))


def _refuse_if_depended(db: Session, doc: LiveDocument) -> None:
    deps = _dependents(db, doc)
    if deps:
        raise HTTPException(status_code=409, detail={
            "code": "depended_on", "ids": sorted(deps),
            "message": ("Bu ko'chirishdan keyin xodimni qabul qilgan brigadir uni yana ko'chirgan "
                        f"(hujjat {_doc_list(deps)}). Avval o'shani bekor qiling.")})


def _resolve_target(db: Session, sender: int, d: date, ttype: Optional[str],
                    tgt: Optional[int], task: Optional[str]):
    if ttype not in ("supervisor", "task"):
        raise HTTPException(status_code=400, detail="target_type must be 'supervisor' or 'task'")
    if ttype == "supervisor":
        if not tgt:
            raise HTTPException(status_code=400, detail="target_manager_id is required")
        if tgt == sender:
            raise HTTPException(status_code=400, detail="Cannot exchange workers to the same unit")
        m = db.query(Manager).filter(Manager.id == tgt).first()
        if not m or m.archived:
            raise HTTPException(status_code=404, detail="Target supervisor not found")
        if tgt not in {c["manager_id"] for c in verifix_live._registry(db)[0].values()}:
            raise HTTPException(status_code=400, detail="The target unit has no cell on this page")
        _assert_open(db, tgt, d)
        return ttype, m.id, m.name, None
    name = (task or "").strip()
    if not name:
        raise HTTPException(status_code=400, detail="task_name is required")
    return ttype, None, None, name


def _check_times(ctx, unit_shift: Optional[int], rows: list, ttime: Optional[str],
                 rtime: Optional[str]) -> None:
    """A transfer time must fall between the earliest arrival among the picked
    workers and now (or their latest departure); a return after it, by now."""
    if not ttime:
        return
    T = live_staff._at(ctx, unit_shift, ttime)
    ins = [datetime.fromisoformat(r["in_at"]) for r in rows if r.get("in_at")]
    if T is None:
        raise HTTPException(status_code=400, detail="time must be HH:MM")
    if T > ctx.now + timedelta(minutes=1):
        raise HTTPException(status_code=400, detail=f"«{ttime}» hali kelmagan vaqt")
    if ins and T < min(ins):
        raise HTTPException(status_code=400, detail=f"«{ttime}» — tanlangan xodimlar hali kelmagan vaqt")
    # A worker who is OUT (gone, or out on a break) cannot be moved from a
    # time after they went out — /staff clamps such a time silently and the
    # document changes nothing. Once they are back, file it then.
    for r in rows:
        p = live_staff.person(ctx, r["employee_id"]).p
        if p["status"] in ("left", "break") and p["out"] is not None and T > p["out"]:
            raise HTTPException(status_code=400, detail=(
                f"{r['worker_name']}: {live_staff._hm(p['out'])} da chiqib ketgan — "
                f"«{ttime}» dan ko'chirib bo'lmaydi"))
    if rtime:
        R = live_staff._at(ctx, unit_shift, rtime)
        if R is None or R <= T:
            raise HTTPException(status_code=400, detail="Return time must be after the transfer time")
        if R > ctx.now + timedelta(minutes=1):
            raise HTTPException(status_code=400, detail=f"«{rtime}» hali kelmagan vaqt")


def _exchange_payload(db: Session, caller: dict, body, d: date, mid: int, *, ctx, ud) -> dict:
    ttype, tgt_id, tgt_name, task_name = _resolve_target(
        db, mid, d, body.target_type, body.target_manager_id, body.task_name)
    ttime = _hhmm(body.transfer_time) if body.transfer_time else None
    if body.transfer_time and not ttime:
        raise HTTPException(status_code=400, detail="time must be HH:MM")
    rtime = _hhmm(body.return_time) if (ttime and body.return_time) else None
    shift = (ctx.units.get(mid) or {}).get("shift")
    T = live_staff._at(ctx, shift, ttime) if ttime else None
    if ttime and T is None:
        raise HTTPException(status_code=400, detail="time must be HH:MM")
    rows = _movers(ctx, ud, mid, body.employees, T)
    _check_times(ctx, shift, rows, ttime, rtime)
    # Said at filing, not only at approval: a move that would take the worker
    # away from the unit an approved later move sends them on from.
    hypo = LiveDocument(id=-1, doc_type="people_exchange", manager_id=mid, day=d, payload={
        "target_type": ttype, "target_manager_id": tgt_id, "task_name": task_name,
        "transfer_time": ttime, "return_time": rtime,
        "employees": [{"employee_id": r["employee_id"]} for r in rows]})
    eids = {r["employee_id"] for r in rows}
    newly = live_staff.broken_moves(ctx, eids, add=hypo) - live_staff.broken_moves(ctx, eids) - {-1}
    if newly:
        raise HTTPException(status_code=400, detail=(
            f"Bu ko'chirish {_doc_list(newly)} hujjatni buzadi: o'sha hujjat xodimni keyinroq "
            f"shu brigadirdan ko'chiradi. Vaqtni yoki qaytish vaqtini tekshiring."))
    if ttype == "task":
        from app.routers.staff import _ensure_exchange_task
        _ensure_exchange_task(db, task_name, caller)
    return {
        "target_type": ttype, "target_manager_id": tgt_id, "target_manager_name": tgt_name,
        "task_name": task_name, "transfer_time": ttime, "return_time": rtime,
        "employees": [{"employee_id": r["employee_id"], "worker_name": r["worker_name"],
                       "old_manager_id": mid, "old_role": r.get("job_title") or "",
                       "old_verifix_code": r.get("verifix_code"),
                       "clock_in": r.get("clock_in"), "clock_out": r.get("clock_out")}
                      for r in rows],
    }


def _role_payload(body, ud) -> dict:
    from app.routers.staff import is_assignable_target_role
    if not body.new_role:
        raise HTTPException(status_code=400, detail="new_role is required")
    if not is_assignable_target_role(body.new_role):
        raise HTTPException(status_code=400, detail="This role can only be set from verifix files and cannot be chosen as a role-change target")
    rows = _picked(ud, body.employees, came=True)
    return {"new_role": body.new_role,
            "employees": [{"employee_id": r["employee_id"], "worker_name": r["worker_name"],
                           "old_role": r.get("job_title") or ""} for r in rows]}


def _approve(doc: LiveDocument, caller: dict, db: Session, reads: Optional[dict] = None) -> None:
    if doc.status == "approved":
        return
    if doc.status == "rejected":
        raise HTTPException(status_code=409, detail="Rejected documents cannot be posted")
    _still_here(db, doc, reads)
    if doc.day:
        age = (date.today() - doc.day).days
        if age > STALE_APPROVE_DAYS:
            raise HTTPException(status_code=409, detail={
                "code": "doc_too_old", "date": doc.day.isoformat(), "age_days": age,
                "max_age_days": STALE_APPROVE_DAYS,
                "message": f"Hujjat sanasi {age} kun oldin ({doc.day}). {STALE_APPROVE_DAYS} kundan "
                           f"eski hujjatni tasdiqlab bo'lmaydi."})
    doc.status = "approved"
    doc.approved_by_telegram_id = _tg(caller)
    doc.approved_by_name = _who(caller)
    doc.approved_at = datetime.now(timezone.utc)
    _history(db, doc, "approved", caller)


def _cancel(doc: LiveDocument, caller: dict, db: Session) -> None:
    if doc.status != "approved":
        return
    doc.status = "draft"
    doc.approved_by_telegram_id = doc.approved_by_name = doc.approved_at = None
    _history(db, doc, "cancelled", caller)


def _reject(doc: LiveDocument, caller: dict, db: Session) -> None:
    if doc.status != "draft":
        raise HTTPException(status_code=409, detail="Only draft documents can be rejected")
    doc.status = "rejected"
    doc.approved_by_telegram_id = _tg(caller)
    doc.approved_by_name = _who(caller)
    doc.approved_at = datetime.now(timezone.utc)
    _history(db, doc, "rejected", caller)


def _tell_rejected(db: Session, doc: LiveDocument, caller: dict) -> None:
    if doc.created_by_telegram_id == _tg(caller):
        return
    key = _creator_profile(db, doc)
    if key:
        _notify(db, "live_document_rejected",
                {"actor_name": _who(caller), "doc_type": doc.doc_type, "date": doc.day},
                admins=False, profiles=(key,), actor=_tg(caller), subject=("live_doc", doc.id),
                ntype="error")


@router.post("/documents", status_code=201)
def create_document(body: DocCreateBody, db: Session = Depends(get_db), caller: dict = Depends(_page)):
    if caller.get("role") not in ("admin", "supervisor"):
        raise HTTPException(status_code=403, detail="Not allowed to create documents")
    if body.doc_type not in DOC_TYPES:
        raise HTTPException(status_code=400, detail="Unsupported document type")
    if not body.employees:
        raise HTTPException(status_code=400, detail="Select at least one employee")
    mid, mgr_name = _resolve_unit(caller, db, body.manager_id)
    _assert_open(db, mid, _day_of(body.attend_date))
    rd, ctx, ud = _build(db, mid, _day_of(body.attend_date))
    d = rd["day"]
    payload = (_exchange_payload(db, caller, body, d, mid, ctx=ctx, ud=ud)
               if body.doc_type == "people_exchange" else _role_payload(body, ud))
    doc = LiveDocument(doc_type=body.doc_type, manager_id=mid, supervisor_name=mgr_name, day=d,
                       payload=payload, status="draft", created_by_telegram_id=_tg(caller),
                       created_by_name=_who(caller), created_by_role=caller.get("role"))
    db.add(doc)
    db.flush()
    _history(db, doc, "created", caller, {"employee_count": len(payload["employees"])})
    # An admin's own filing IS the approval (/staff's rule) — nobody above them
    # to ask; the parties are told it is done, not asked to decide it.
    if caller.get("role") == "admin":
        _approve(doc, caller, db)
        db.commit()
        action_log.enrich(**_log_doc(doc, [("status", None, "approved")]))
        _notify_doc(db, doc, "approved", _tg(caller))
        return {"id": doc.id, "status": doc.status}
    db.commit()
    action_log.enrich(**_log_doc(doc, [("status", None, "draft")]))
    _notify_doc(db, doc, "created", _tg(caller))
    return {"id": doc.id, "status": doc.status}


class DocUpdateBody(BaseModel):
    employees: List[str]
    new_role: Optional[str] = None
    target_type: Optional[str] = None
    target_manager_id: Optional[int] = None
    task_name: Optional[str] = None
    transfer_time: Optional[str] = None
    return_time: Optional[str] = None


@router.put("/documents/{doc_id}")
def update_document(doc_id: int, body: DocUpdateBody, db: Session = Depends(get_db),
                    caller: dict = Depends(_page)):
    doc = _doc_or_404(db, caller, doc_id)
    if doc.status != "draft":
        raise HTTPException(status_code=409, detail="Only draft (Нет) documents can be edited")
    if caller.get("role") not in ("admin", "shift-manager") and not _is_creator(doc, caller):
        raise HTTPException(status_code=403, detail="Not allowed to edit this document")
    if not body.employees:
        raise HTTPException(status_code=400, detail="Select at least one employee")
    _assert_open(db, doc.manager_id, doc.day)
    rd, ctx, ud = _build(db, doc.manager_id, doc.day)
    if doc.doc_type == "people_exchange":
        prev = doc.payload or {}
        merged = DocCreateBody(
            doc_type=doc.doc_type, attend_date=doc.day.isoformat(), employees=body.employees,
            target_type=body.target_type or prev.get("target_type"),
            target_manager_id=(body.target_manager_id if body.target_manager_id is not None
                               else prev.get("target_manager_id")),
            task_name=body.task_name if body.task_name is not None else prev.get("task_name"),
            transfer_time=(body.transfer_time if body.transfer_time is not None
                           else prev.get("transfer_time")),
            return_time=body.return_time if body.return_time is not None else prev.get("return_time"))
        doc.payload = _exchange_payload(db, caller, merged, doc.day, doc.manager_id, ctx=ctx, ud=ud)
    else:
        doc.payload = _role_payload(body, ud)
    flag_modified(doc, "payload")
    _history(db, doc, "edited", caller, {"employee_count": len(doc.payload["employees"])})
    db.commit()
    action_log.enrich(**_log_doc(doc))
    return {"ok": True}


@router.post("/documents/{doc_id}/approve")
def approve_document(doc_id: int, db: Session = Depends(get_db), caller: dict = Depends(_page)):
    doc = _doc_or_404(db, caller, doc_id)
    if not _can_approve(doc, caller, db):
        raise HTTPException(status_code=403, detail="Not authorised to post this document")
    if doc.status == "approved":
        return {"ok": True, "status": doc.status}
    _approve(doc, caller, db)
    db.commit()
    action_log.enrich(**_log_doc(doc, [("status", "draft", "approved")]))
    _notify_doc(db, doc, "approved", _tg(caller))
    return {"ok": True, "status": doc.status}


@router.post("/documents/{doc_id}/reject")
def reject_document(doc_id: int, db: Session = Depends(get_db), caller: dict = Depends(_page)):
    doc = _doc_or_404(db, caller, doc_id)
    if not _may_reject(doc, caller, db):
        raise HTTPException(status_code=403, detail="Not authorised to reject this document")
    _reject(doc, caller, db)
    db.commit()
    action_log.enrich(**_log_doc(doc, [("status", "draft", "rejected")]))
    _tell_rejected(db, doc, caller)
    return {"ok": True, "status": doc.status}


@router.post("/documents/{doc_id}/cancel")
def cancel_document(doc_id: int, db: Session = Depends(get_db), caller: dict = Depends(_page)):
    doc = _doc_or_404(db, caller, doc_id)
    if not _can_approve(doc, caller, db):
        raise HTTPException(status_code=403, detail="Not authorised to un-post this document")
    if doc.status != "approved":
        return {"ok": True, "status": doc.status}
    _refuse_if_depended(db, doc)
    _cancel(doc, caller, db)
    db.commit()
    action_log.enrich(**_log_doc(doc, [("status", "approved", "draft")]))
    _notify_doc(db, doc, "cancelled", _tg(caller))
    return {"ok": True, "status": doc.status}


@router.post("/documents/{doc_id}/delete")
def delete_document(doc_id: int, db: Session = Depends(get_db), caller: dict = Depends(_page)):
    doc = _doc_or_404(db, caller, doc_id)
    if doc.status == "approved":
        if not _can_approve(doc, caller, db):
            raise HTTPException(status_code=403, detail="Approved documents can only be deleted by an approver")
        _refuse_if_depended(db, doc)
    elif doc.status == "draft":
        # Deleting a pending draft IS its rejection — the record stays (/staff).
        if not _may_reject(doc, caller, db):
            raise HTTPException(status_code=403, detail="Not allowed to reject this document")
        _reject(doc, caller, db)
        db.commit()
        action_log.enrich(action="lab.live_document_rejected", **_log_doc(doc, [("status", "draft", "rejected")]))
        _tell_rejected(db, doc, caller)
        return {"ok": True, "status": doc.status}
    elif caller.get("role") not in ("admin", "shift-manager") and not _is_creator(doc, caller):
        raise HTTPException(status_code=403, detail="Not allowed to delete this document")
    log_fields = _log_doc(doc, [("status", doc.status, None)])
    was_approved = doc.status == "approved"
    snapshot = LiveDocument(id=doc.id, doc_type=doc.doc_type, manager_id=doc.manager_id,
                            day=doc.day, payload=dict(doc.payload or {}),
                            created_by_name=doc.created_by_name, supervisor_name=doc.supervisor_name)
    db.query(LiveDocumentHistory).filter(LiveDocumentHistory.document_id == doc.id).delete()
    db.delete(doc)
    db.commit()
    action_log.enrich(**log_fields)
    if was_approved:
        _notify_doc(db, snapshot, "cancelled", _tg(caller))
    return {"ok": True}


class DocBulkBody(BaseModel):
    ids: List[int]
    action: str                       # approve | reject | cancel | delete


@router.post("/documents/bulk")
def bulk_documents(body: DocBulkBody, db: Session = Depends(get_db), caller: dict = Depends(_page)):
    docs = _scope_docs(db.query(LiveDocument), caller, db).filter(LiveDocument.id.in_(body.ids)).all()
    done, refused, notify, rejected = 0, [], [], []
    reads: dict = {}
    if body.action in ("cancel", "delete"):
        # Newest approval first: a move that depends on another was approved
        # after it, so un-posting in this order never trips over a dependent.
        docs.sort(key=lambda x: x.approved_at or datetime.min.replace(tzinfo=timezone.utc),
                  reverse=True)

    def _depended(doc) -> bool:
        db.flush()
        if _dependents(db, doc):
            refused.append({"doc_id": doc.id, "date": doc.day.isoformat() if doc.day else None,
                            "reason": "depended_on"})
            return True
        return False

    for doc in docs:
        if body.action == "approve":
            if doc.status != "draft" or not _can_approve(doc, caller, db):
                continue
            try:
                _approve(doc, caller, db, reads)
            except HTTPException as exc:
                refused.append({"doc_id": doc.id, "date": doc.day.isoformat() if doc.day else None,
                                "reason": exc.detail.get("code") if isinstance(exc.detail, dict) else "refused"})
                continue
            notify.append((doc, "approved"))
        elif body.action == "reject":
            if doc.status != "draft" or not _may_reject(doc, caller, db):
                continue
            _reject(doc, caller, db)
            rejected.append(doc)
        elif body.action == "cancel":
            if doc.status != "approved" or not _can_approve(doc, caller, db):
                continue
            if _depended(doc):
                continue
            _cancel(doc, caller, db)
            notify.append((doc, "cancelled"))
        elif body.action == "delete":
            if doc.status == "approved":
                if not _can_approve(doc, caller, db) or _depended(doc):
                    continue
                db.query(LiveDocumentHistory).filter(LiveDocumentHistory.document_id == doc.id).delete()
                db.delete(doc)
            elif doc.status == "draft":
                if not _may_reject(doc, caller, db):
                    continue
                _reject(doc, caller, db)
                rejected.append(doc)
            else:
                if caller.get("role") not in ("admin", "shift-manager") and not _is_creator(doc, caller):
                    continue
                db.query(LiveDocumentHistory).filter(LiveDocumentHistory.document_id == doc.id).delete()
                db.delete(doc)
        else:
            raise HTTPException(status_code=400, detail="Unknown action")
        done += 1
    db.commit()
    for doc, event in notify:
        _notify_doc(db, doc, event, _tg(caller))
    for doc in rejected:
        _tell_rejected(db, doc, caller)
    action_log.enrich(target_kind="batch", target_id=",".join(str(i) for i in body.ids[:20]),
                      details=[("mode", body.action), ("count", done),
                               ("skipped", max(len(body.ids) - done, 0))])
    return {"ok": True, "affected": done, "refused": refused}


# ── deletion requests ────────────────────────────────────────────────────────

class BulkDeleteBody(BaseModel):
    manager_id: Optional[int] = None
    attend_date: str
    # The page's row key — on this page a Verifix employee id. Named after
    # /staff's field so the shared deletion dialog sends one shape to both.
    worker_names: List[str] = Field(default_factory=list)
    employee_ids: List[str] = Field(default_factory=list)
    replace_batch_id: Optional[str] = None


def _original(r: dict) -> dict:
    return {"job_title": r.get("job_title") or "", "schedule": r.get("schedule") or "",
            "hours_worked": r.get("hours_worked"), "clock_in": r.get("clock_in"),
            "clock_out": r.get("clock_out"), "verifix_code": r.get("verifix_code")}


@router.post("/attendance/bulk-delete")
def bulk_delete(body: BulkDeleteBody, db: Session = Depends(get_db), caller: dict = Depends(_page)):
    """A brigadir files a batch of deletion requests; an admin (or a
    `staff.attendance.delete` grantee) deletes at once (`staff.bulk_delete_attendance`)."""
    role = caller.get("role")
    granted = has_cap(db, caller, CAP_ATTENDANCE_DELETE)
    if role not in ("admin", "supervisor") and not granted:
        raise HTTPException(status_code=403, detail="Not allowed")
    ids = list(dict.fromkeys([str(x) for x in (body.employee_ids or body.worker_names) if str(x).strip()]))
    if not ids:
        raise HTTPException(status_code=400, detail="No workers specified")
    direct = role == "admin" or granted
    if direct:
        mid = body.manager_id or caller.get("role_id")
        if not mid:
            raise HTTPException(status_code=400, detail="manager_id required for admin")
        if not _cap_covers(db, caller, CAP_ATTENDANCE_DELETE, mid):
            raise HTTPException(status_code=403, detail="Admin only")
    else:
        mid = caller.get("role_id")
        if not mid:
            raise HTTPException(status_code=400, detail="Supervisor has no linked manager")
    d = _day_of(body.attend_date)
    _assert_open(db, mid, d)
    rd, ctx, ud = _build(db, mid, d)
    d = rd["day"]
    rows = _named(ud)
    batch = str(uuid4())
    now = datetime.now(timezone.utc)
    sup = _unit_name(db, mid)
    created: list[str] = []
    if not direct and body.replace_batch_id:
        for r in db.query(LiveDeletion).filter(LiveDeletion.batch_id == body.replace_batch_id,
                                               LiveDeletion.status == "pending").all():
            r.status = "rejected"
        db.flush()
    pending = {r.employee_id for r in db.query(LiveDeletion).filter(
        LiveDeletion.manager_id == mid, LiveDeletion.day == d, LiveDeletion.status == "pending")}
    for eid in ids:
        r = rows.get(eid)
        if not r:
            continue
        if not direct and eid in pending:
            continue
        db.add(LiveDeletion(
            batch_id=batch, manager_id=mid, day=d, employee_id=eid, worker_name=r["worker_name"],
            original=_original(r), status="approved" if direct else "pending",
            supervisor_telegram_id=_tg(caller) if not direct else None,
            supervisor_name=_who(caller) if not direct else sup,
            initiated_by="admin" if direct else None,
            processed_by_telegram_id=_tg(caller) if direct else None,
            processed_by_name=_who(caller) if direct else None,
            processed_at=now if direct else None))
        created.append(r["worker_name"])
    db.commit()
    if created:
        action_log.enrich(
            action=None if direct else "lab.live_delete_requested",
            target_kind="batch", target_id=batch, unit_id=mid, unit_name=sup, day=d,
            details=[("date", d.isoformat()), ("count", len(created)),
                     ("workers", ", ".join(created[:10]) + (f" +{len(created) - 10}" if len(created) > 10 else "")),
                     ("status", "approved" if direct else "pending")])
        params = {"supervisor_name": _who(caller), "count": len(created), "date": d,
                  "admin_name": _who(caller), "worker_name": ", ".join(created[:3])
                  + (f" +{len(created) - 3}" if len(created) > 3 else "")}
        if direct:
            _notify(db, "live_record_deleted", params, units=(mid,), supervisors=(mid,),
                    actor=_tg(caller), subject=("live_day", f"{mid}:{d.isoformat()}"))
        else:
            _notify(db, "live_delete_request", params, units=(mid,), actor=_tg(caller),
                    subject=("live_batch", batch))
    return {"ok": True, "affected": len(created)}


class BatchBody(BaseModel):
    ids: Optional[List[int]] = None


def _batch_filter(token: str):
    if token.startswith("solo-"):
        try:
            return LiveDeletion.id == int(token[5:])
        except ValueError:
            return LiveDeletion.id == -1
    return LiveDeletion.batch_id == token


def _decide_deletions(reqs: list, action: str, caller: dict, db: Session) -> None:
    if caller.get("role") not in ("admin", "shift-manager") and not has_cap(db, caller, CAP_REQUESTS_APPROVE):
        raise HTTPException(status_code=403, detail="Not authorised")
    for r in reqs:
        if caller.get("role") == "shift-manager" and not shift_scope.covers(db, caller.get("role_id"), r.manager_id):
            raise HTTPException(status_code=403, detail="Not responsible for this unit")
        if (caller.get("role") not in ("admin", "shift-manager")
                and not scope_is_all(db, caller, CAP_REQUESTS_APPROVE)):
            units = _caller_units(db, caller)
            if units is not None and r.manager_id not in units:
                raise HTTPException(status_code=403, detail="Not authorised")
        if r.status == "pending" and _closed(db, r.manager_id, r.day) and action == "approved":
            pass   # /staff decides requests on a closed day too — that is what «confirmed» waits on
    now = datetime.now(timezone.utc)
    for r in reqs:
        r.status = action
        r.processed_by_telegram_id = _tg(caller)
        r.processed_by_name = _who(caller)
        r.processed_at = now


def _tell_decided(db: Session, reqs: list, action: str, caller: dict) -> None:
    by_unit: dict = defaultdict(list)
    for r in reqs:
        by_unit[(r.manager_id, r.day, r.batch_id or f"solo-{r.id}")].append(r.worker_name or "")
    for (mid, d, token), names in by_unit.items():
        _notify(db, "live_delete_approved" if action == "approved" else "live_delete_rejected",
                {"processor_name": _who(caller), "count": len(names), "date": d,
                 "worker_name": ", ".join(names[:3]) + (f" +{len(names) - 3}" if len(names) > 3 else "")},
                units=(mid,), supervisors=(mid,), actor=_tg(caller),
                subject=("live_batch", token), ntype="success" if action == "approved" else "warning")


@router.post("/requests/batch/{batch_id}/approve")
def approve_batch(batch_id: str, body: BatchBody, db: Session = Depends(get_db),
                  caller: dict = Depends(_page)):
    q = db.query(LiveDeletion).filter(_batch_filter(batch_id), LiveDeletion.status == "pending")
    if body.ids:
        q = q.filter(LiveDeletion.id.in_(body.ids))
    reqs = q.all()
    if not reqs:
        raise HTTPException(status_code=404, detail="No pending requests found in batch")
    _decide_deletions(reqs, "approved", caller, db)
    db.commit()
    action_log.enrich(target_kind="batch", target_id=batch_id,
                      details=[("count", len(reqs))], changes=[("status", "pending", "approved")])
    _tell_decided(db, reqs, "approved", caller)
    return {"ok": True, "approved": len(reqs)}


@router.post("/requests/batch/{batch_id}/reject")
def reject_batch(batch_id: str, db: Session = Depends(get_db), caller: dict = Depends(_page)):
    reqs = db.query(LiveDeletion).filter(_batch_filter(batch_id), LiveDeletion.status == "pending").all()
    if not reqs:
        raise HTTPException(status_code=404, detail="No pending requests found in batch")
    _decide_deletions(reqs, "rejected", caller, db)
    db.commit()
    action_log.enrich(target_kind="batch", target_id=batch_id,
                      details=[("count", len(reqs))], changes=[("status", "pending", "rejected")])
    _tell_decided(db, reqs, "rejected", caller)
    return {"ok": True, "rejected": len(reqs)}


@router.post("/requests/batch/{batch_id}/withdraw")
def withdraw_batch(batch_id: str, db: Session = Depends(get_db), caller: dict = Depends(_page)):
    role, rid = caller.get("role"), caller.get("role_id")
    if role not in ("admin", "supervisor"):
        raise HTTPException(status_code=403, detail="Not authorised")
    q = db.query(LiveDeletion).filter(_batch_filter(batch_id), LiveDeletion.status == "pending")
    if role == "supervisor":
        q = q.filter(LiveDeletion.manager_id == rid) if rid else \
            q.filter(LiveDeletion.supervisor_telegram_id == _tg(caller))
    reqs = q.all()
    if not reqs:
        raise HTTPException(status_code=404, detail="No pending requests found in batch")
    for r in reqs:
        r.status = "rejected"
    db.commit()
    action_log.enrich(target_kind="batch", target_id=batch_id,
                      details=[("count", len(reqs))], changes=[("status", "pending", "withdrawn")])
    return {"ok": True, "withdrawn": len(reqs)}


def _one_request(db: Session, req_id: int) -> LiveDeletion:
    r = db.query(LiveDeletion).filter(LiveDeletion.id == req_id).first()
    if not r:
        raise HTTPException(status_code=404, detail="Request not found")
    return r


@router.post("/requests/{req_id}/withdraw")
def withdraw_request(req_id: int, db: Session = Depends(get_db), caller: dict = Depends(_page)):
    role, rid = caller.get("role"), caller.get("role_id")
    if role not in ("admin", "supervisor"):
        raise HTTPException(status_code=403, detail="Not authorised")
    r = _one_request(db, req_id)
    if r.status != "pending":
        raise HTTPException(status_code=409, detail="Can only withdraw pending requests")
    if role == "supervisor" and not ((rid and r.manager_id == rid) or r.supervisor_telegram_id == _tg(caller)):
        raise HTTPException(status_code=403, detail="Not your request")
    r.status = "rejected"
    db.commit()
    action_log.enrich(target_kind="request", target_id=req_id, target_name=r.worker_name,
                      unit_id=r.manager_id, day=r.day, changes=[("status", "pending", "withdrawn")])
    return {"ok": True}


@router.post("/requests/{req_id}/approve")
def approve_request(req_id: int, db: Session = Depends(get_db), caller: dict = Depends(_page)):
    r = _one_request(db, req_id)
    if r.status != "pending":
        raise HTTPException(status_code=409, detail="Request already processed")
    _decide_deletions([r], "approved", caller, db)
    db.commit()
    action_log.enrich(target_kind="request", target_id=req_id, target_name=r.worker_name,
                      unit_id=r.manager_id, day=r.day, changes=[("status", "pending", "approved")])
    _tell_decided(db, [r], "approved", caller)
    return {"ok": True, "status": "approved"}


@router.post("/requests/{req_id}/reject")
def reject_request(req_id: int, db: Session = Depends(get_db), caller: dict = Depends(_page)):
    r = _one_request(db, req_id)
    if r.status != "pending":
        raise HTTPException(status_code=409, detail="Request already processed")
    _decide_deletions([r], "rejected", caller, db)
    db.commit()
    action_log.enrich(target_kind="request", target_id=req_id, target_name=r.worker_name,
                      unit_id=r.manager_id, day=r.day, changes=[("status", "pending", "rejected")])
    _tell_decided(db, [r], "rejected", caller)
    return {"ok": True, "status": "rejected"}


@router.post("/requests/{req_id}/undo")
def undo_request(req_id: int, db: Session = Depends(get_db), caller: dict = Depends(_page)):
    """Put a deleted worker back on the unit's day (admin / shift-manager)."""
    role = caller.get("role")
    if role not in ("admin", "shift-manager"):
        raise HTTPException(status_code=403, detail="Not authorised")
    r = _one_request(db, req_id)
    if r.status != "approved":
        raise HTTPException(status_code=409, detail="Can only undo approved requests")
    if role == "shift-manager" and not shift_scope.covers(db, caller.get("role_id"), r.manager_id):
        raise HTTPException(status_code=403, detail="Not responsible for this unit")
    r.status = "undone"
    db.commit()
    action_log.enrich(target_kind="request", target_id=req_id, target_name=r.worker_name,
                      unit_id=r.manager_id, day=r.day, changes=[("status", "approved", "undone")])
    _notify(db, "live_request_undone", {"worker_name": r.worker_name or "", "date": r.day,
                                        "undoer": _who(caller)},
            units=(r.manager_id,), supervisors=(r.manager_id,), actor=_tg(caller),
            subject=("live_day", f"{r.manager_id}:{r.day.isoformat()}"), ntype="warning")
    return {"ok": True}


# ── the day close (the Tasdiqlash tab) ───────────────────────────────────────

@router.get("/approvals/day")
def approval_day(attend_date: Optional[str] = None, manager_id: Optional[int] = None,
                 db: Session = Depends(get_db), caller: dict = Depends(_page)):
    """The day-close state of one unit-day (`staff.approval_day`)."""
    mid = _read_unit(db, caller, manager_id)
    d = _day_of(attend_date)
    closure = _closed(db, mid, d) if d else None
    pending = live_staff.pending_count(db, mid, d) if d else 0
    unplaced: list = []
    if closure is None:
        try:
            _, _, ud = _build(db, mid, d)
            unplaced = live_staff.unplaced(ud["workers"])
        except _NoDay:
            unplaced = []
    state = "open" if closure is None else ("closed" if pending else "confirmed")
    return {
        "manager_id": mid, "date": attend_date, "state": state, "closed": closure is not None,
        "closed_by": closure.closed_by_name if closure else None,
        "closed_at": closure.closed_at.isoformat() if closure and closure.closed_at else None,
        "pending_requests": pending,
        "can_reopen": _cap_covers(db, caller, CAP_DAY_REOPEN, mid),
        "needs_cell": len(unplaced),
        "needs_cell_names": [r["worker_name"] for r in unplaced[:20]],
    }


@router.get("/approvals/calendar")
def approvals_calendar(year: int, month: int, manager_id: Optional[int] = None,
                       db: Session = Depends(get_db), caller: dict = Depends(_page)):
    """Per-day close state for a month (`staff.approvals_calendar`): a day with
    a stored Verifix read is «open» until closed by hand; closed with nothing
    pending is «confirmed»."""
    mid = _read_unit(db, caller, manager_id)
    start = date(year, month, 1)
    end = date(year + (month == 12), (month % 12) + 1, 1)
    read_days = set()
    for (key,) in db.query(LiveVerifixRead.key).filter(
            LiveVerifixRead.key >= f"day:{start.isoformat()}", LiveVerifixRead.key < f"day:{end.isoformat()}"):
        try:
            read_days.add(date.fromisoformat(key[4:]))
        except ValueError:
            continue
    closes = {c.day: c for c in db.query(LiveDayClose).filter(
        LiveDayClose.manager_id == mid, LiveDayClose.day >= start, LiveDayClose.day < end)}
    pending_days = {d for (d,) in db.query(LiveDocument.day).filter(
        LiveDocument.manager_id == mid, LiveDocument.status == "draft",
        LiveDocument.day >= start, LiveDocument.day < end)} | {d for (d,) in db.query(LiveDeletion.day).filter(
            LiveDeletion.manager_id == mid, LiveDeletion.status == "pending",
            LiveDeletion.day >= start, LiveDeletion.day < end)}
    days = {}
    for d in sorted(read_days | set(closes)):
        if d in closes:
            c = closes[d]
            days[d.isoformat()] = {"status": "closed" if d in pending_days else "confirmed",
                                   "closed_by": c.closed_by_name,
                                   "closed_at": c.closed_at.isoformat() if c.closed_at else None}
        else:
            days[d.isoformat()] = {"status": "open"}
    return {"manager_id": mid, "manager_name": _unit_name(db, mid), "year": year, "month": month,
            "days": days}


class ApprovalBody(BaseModel):
    manager_id: Optional[int] = None
    date: str


@router.post("/daily/close")
def close_day(body: ApprovalBody, db: Session = Depends(get_db), caller: dict = Depends(_page)):
    """A supervisor closes their own day, an admin anybody's (`staff.close_day`) —
    refused while a counted worker has no cell."""
    role = caller.get("role")
    if role == "supervisor":
        mid = caller.get("role_id")
        if not mid:
            raise HTTPException(status_code=400, detail="Supervisor has no assigned manager")
    elif role == "admin":
        mid = body.manager_id
        if not mid:
            raise HTTPException(status_code=400, detail="manager_id required")
    else:
        raise HTTPException(status_code=403, detail="Only supervisors or admins can close a day")
    d = _day_of(body.date)
    unit = db.query(Manager).filter(Manager.id == mid).first()
    if d > verifix_live.shift_day(db, unit.shift if unit else None):
        raise HTTPException(status_code=400, detail="Cannot close a future date")
    if _closed(db, mid, d):
        raise HTTPException(status_code=409, detail="Day is already closed")
    _, _, ud = _build(db, mid, d)
    left = live_staff.unplaced(ud["workers"])
    if left:
        names = [r["worker_name"] for r in left[:5]]
        more = f" +{len(left) - len(names)}" if len(left) > len(names) else ""
        action_log.enrich(target_kind="day", target_id=f"{mid}:{d}", unit_id=mid, day=d,
                          details=[("date", str(d)), ("blocked", "cells_missing"), ("workers", len(left))])
        raise HTTPException(status_code=409, detail=(
            f"{len(left)} ta xodim yacheykaga biriktirilmagan: {', '.join(names)}{more}. "
            f"«Yacheykalar» bo'limida ularni joylashtiring."))
    # The live day still moves while anybody is inside, out on a break or due:
    # the close waits for them (a closed day must not keep changing under it).
    on = live_staff.busy(ud)
    if on:
        names = [r["worker_name"] for r in on[:5]]
        more = f" +{len(on) - len(names)}" if len(on) > len(names) else ""
        action_log.enrich(target_kind="day", target_id=f"{mid}:{d}", unit_id=mid, day=d,
                          details=[("date", str(d)), ("blocked", "still_on_shift"), ("workers", len(on))])
        raise HTTPException(status_code=409, detail=(
            f"{len(on)} ta xodim hali smenada (ichkarida, tanaffusda yoki hali kelmagan): "
            f"{', '.join(names)}{more}. Hamma chiqib ketgach kunni yoping."))
    db.add(LiveDayClose(manager_id=mid, day=d, closed_by_name=_who(caller)))
    db.commit()
    pending = live_staff.pending_count(db, mid, d)
    action_log.enrich(target_kind="day", target_id=f"{mid}:{d}", unit_id=mid,
                      unit_name=unit.name if unit else None, day=d,
                      details=[("date", str(d))] + ([("pending", pending)] if pending else []),
                      changes=[("status", "open", "closed")])
    _notify(db, "live_day_closed", {"closer_name": _who(caller), "date": d}, units=(mid,),
            supervisors=(mid,) if role == "admin" else (), actor=_tg(caller),
            subject=("live_day", f"{mid}:{d.isoformat()}"))
    return {"ok": True, "state": "closed" if pending else "confirmed", "manager_id": mid,
            "date": body.date, "pending_requests": pending}


@router.post("/approvals/reopen")
def reopen_day(body: ApprovalBody, db: Session = Depends(get_db), caller: dict = Depends(_page)):
    if not body.manager_id:
        raise HTTPException(status_code=400, detail="manager_id required")
    if not _cap_covers(db, caller, CAP_DAY_REOPEN, body.manager_id):
        raise HTTPException(status_code=403, detail="Only an admin can re-open a closed day")
    d = _day_of(body.date)
    row = _closed(db, body.manager_id, d)
    if row:
        closer = row.closed_by_name
        db.delete(row)
        db.commit()
        action_log.enrich(target_kind="day", target_id=f"{body.manager_id}:{d}", unit_id=body.manager_id,
                          day=d, details=[("date", str(d)), ("closed_by", closer or "—")],
                          changes=[("status", "closed", "open")])
        _notify(db, "live_day_reopened", {"reopener_name": _who(caller), "date": d},
                units=(body.manager_id,), supervisors=(body.manager_id,), actor=_tg(caller),
                subject=("live_day", f"{body.manager_id}:{d.isoformat()}"), ntype="warning")
    return {"ok": True, "state": "open", "manager_id": body.manager_id, "date": body.date}


# ── cell placement (the Yacheykalar tab) ─────────────────────────────────────

def _can_place(caller: dict, mid: int) -> bool:
    return caller.get("role") == "admin" or (caller.get("role") == "supervisor"
                                              and caller.get("role_id") == mid)


@router.get("/cell-placement")
def cell_placement(attend_date: Optional[str] = None, manager_id: Optional[int] = None,
                   db: Session = Depends(get_db), caller: dict = Depends(_page)):
    """The unit's day grouped by cell, the cell-less first (`staff.cell_placement`)."""
    from app.services.kpi_calculator import is_direct_role
    mid = _read_unit(db, caller, manager_id)
    rd, ctx, ud = _build(db, mid, _day_of(attend_date))
    d = rd["day"]
    closed = _closed(db, mid, d)

    def w_json(r: dict) -> dict:
        counted = is_direct_role(r.get("job_title"), r.get("hours_worked"), False)
        return {"id": r["id"], "worker_name": r["worker_name"], "job_title": r.get("job_title") or "",
                "hours": r.get("hours_worked") or 0.0, "clock_in_out": r.get("clock_in_out") or "",
                "counted": counted, "hc_weight": r.get("hc_weight"), "split_of": r.get("split_of"),
                "split_at": (r.get("clock_in") if r.get("split_of") is not None
                             else (r.get("clock_out") if r.get("hc_weight") is not None else None)),
                "from_unit": (r.get("moved") or {}).get("unit") if (r.get("moved") or {}).get("dir") == "in" else None}

    def hc(r: dict) -> float:
        if not is_direct_role(r.get("job_title"), r.get("hours_worked"), False):
            return 0.0
        return 1.0 if r.get("hc_weight") is None else float(r["hc_weight"])

    by_code: dict = defaultdict(list)
    unplaced = []
    for r in ud["workers"]:
        if r.get("verifix_code"):
            by_code[r["verifix_code"]].append(r)
        elif r.get("clock_in"):
            unplaced.append(w_json(r))
    catalog = live_staff.cells_catalog(db, ctx, mid, ud["workers"])
    cells_out = []
    for c in catalog:
        rows = by_code.pop(c["verifix_code"], [])
        if not rows and not c.get("in_load"):
            continue
        cells_out.append({"verifix_code": c["verifix_code"], "cell_id": c["cell_id"],
                          "leader_name": c["leader_name"], "workers": [w_json(r) for r in rows],
                          "counted": round(sum(hc(r) for r in rows), 2), "total": len(rows),
                          "hours": round(sum(float(r.get("hours_worked") or 0) for r in rows), 1)})
    for code, rows in by_code.items():
        cells_out.append({"verifix_code": code, "cell_id": None, "leader_name": None,
                          "workers": [w_json(r) for r in rows],
                          "counted": round(sum(hc(r) for r in rows), 2), "total": len(rows),
                          "hours": round(sum(float(r.get("hours_worked") or 0) for r in rows), 1)})
    cells_out.sort(key=lambda c: c["verifix_code"])
    return {
        "manager_id": mid, "manager_name": rd["unit"]["name"], "shift": rd["unit"]["shift"],
        "date": d.isoformat(), "day_closed": closed is not None, "cells_required": True,
        "can_edit": closed is None and _can_place(caller, mid),
        "unplaced": unplaced, "cells": cells_out,
        "totals": {"cells": len(cells_out),
                   "workers": sum(1 for r in ud["workers"] if r.get("split_of") is None),
                   "counted": sum(c["counted"] for c in cells_out),
                   "hours": round(sum(c["hours"] for c in cells_out), 1),
                   "unplaced": len(unplaced)},
    }


class PlacementMove(BaseModel):
    attendance_id: int
    verifix_code: str


class PlacementSplit(BaseModel):
    attendance_id: int
    verifix_code: str
    second_code: str
    split_at: str


class PlacementBody(BaseModel):
    manager_id: Optional[int] = None
    date: str
    moves: List[PlacementMove] = []
    splits: List[PlacementSplit] = []
    unsplits: List[int] = []


@router.put("/cell-placement")
def save_cell_placement(body: PlacementBody, db: Session = Depends(get_db),
                        caller: dict = Depends(_page)):
    """Write a batch of placements, splits and un-splits (`staff.save_cell_placement`)."""
    mid = _read_unit(db, caller, body.manager_id)
    if not _can_place(caller, mid):
        raise HTTPException(status_code=403, detail="Only this unit's supervisor or an admin may place workers")
    d = _day_of(body.date)
    if _closed(db, mid, d):
        raise HTTPException(status_code=409, detail="Day is already closed")
    rd, ctx, ud = _build(db, mid, d)
    d = rd["day"]
    by_id = {r["id"]: r for r in ud["workers"]}
    valid = {c["code"] for c in ctx.cells.values() if c["manager_id"] == mid} | \
            {r["verifix_code"] for r in ud["workers"] if r.get("verifix_code")}

    def row(att_id: int) -> dict:
        r = by_id.get(att_id)
        if not r:
            raise HTTPException(status_code=404, detail="Worker row not found for this unit and date")
        return r

    def code_ok(code: str) -> str:
        code = (code or "").strip()
        if code not in valid:
            raise HTTPException(status_code=400, detail=f"«{code}» — bu brigadaning yacheykasi emas")
        return code

    def placement(eid: str) -> LivePlacement:
        p = (db.query(LivePlacement).filter(LivePlacement.manager_id == mid, LivePlacement.day == d,
                                            LivePlacement.employee_id == eid).first())
        if p is None:
            p = LivePlacement(manager_id=mid, day=d, employee_id=eid, verifix_code="")
            db.add(p)
        return p

    moved = split_n = unsplit_n = 0
    for sec_id in dict.fromkeys(body.unsplits):
        r = row(sec_id)
        if r.get("split_of") is None:
            raise HTTPException(status_code=400, detail="That row is not the second half of a split")
        p = placement(r["employee_id"])
        p.second_code = p.split_at = None
        p.updated_by = _who(caller)
        unsplit_n += 1
    for m in body.moves:
        r = row(m.attendance_id)
        p = placement(r["employee_id"])
        p.verifix_code = code_ok(m.verifix_code)
        p.second_code = p.split_at = None
        p.updated_by = _who(caller)
        moved += 1
    for sp in body.splits:
        r = row(sp.attendance_id)
        if r.get("split_of") is not None or r.get("hc_weight") is not None:
            raise HTTPException(status_code=400, detail="That worker is already split")
        first, second = code_ok(sp.verifix_code), code_ok(sp.second_code)
        if first == second:
            raise HTTPException(status_code=400, detail="Ikkala yacheyka bir xil bo'lishi mumkin emas")
        at = _hhmm(sp.split_at)
        t = live_staff._at(ctx, (ctx.units.get(mid) or {}).get("shift"), at) if at else None
        start = datetime.fromisoformat(r["in_at"]) if r.get("in_at") else None
        stop = datetime.fromisoformat(r["out_at"]) if r.get("out_at") else ctx.now
        if not (t and start and start < t < stop):
            raise HTTPException(status_code=400,
                                detail=f"{r['worker_name']}: «{sp.split_at}» bu xodimning ish vaqtidan tashqarida")
        p = placement(r["employee_id"])
        p.verifix_code, p.second_code, p.split_at = first, second, at
        p.updated_by = _who(caller)
        split_n += 1
    db.commit()
    _, ctx2, ud2 = _build(db, mid, d)
    left = live_staff.unplaced(ud2["workers"])
    action_log.enrich(target_kind="day", target_id=f"{mid}:{d}", unit_id=mid, day=d,
                      details=[("date", str(d)), ("moved", moved), ("split", split_n),
                               ("unsplit", unsplit_n), ("still_unplaced", len(left))])
    return {"ok": True, "moved": moved, "split": split_n, "unsplit": unsplit_n, "unplaced": len(left)}


# ── Excel ────────────────────────────────────────────────────────────────────

class ExportRow(BaseModel):
    worker_name: Optional[str] = None
    job_title: Optional[str] = None
    cell: Optional[str] = None
    schedule: Optional[str] = None
    clock_in: Optional[str] = None
    clock_out: Optional[str] = None
    status: Optional[str] = None
    hours_worked: Optional[float] = None
    early_arrival_min: Optional[float] = None
    effective_hours: Optional[float] = None


class ExportBody(BaseModel):
    manager_id: int
    attend_date: str
    rows: List[ExportRow] = Field(..., max_length=5000)
    extra_hours: Optional[float] = None


@router.post("/attendance/export")
def export_attendance(request: Request, body: ExportBody, db: Session = Depends(get_db),
                      caller: dict = Depends(_page)):
    """/staff's workbook, with the clock split in two and the live status — the
    rows on screen, filtered and in the viewer's alphabet."""
    import openpyxl
    from io import BytesIO
    from openpyxl.styles import Alignment, Font, PatternFill
    _read_unit(db, caller, body.manager_id)
    manager = _unit_name(db, body.manager_id) or f"#{body.manager_id}"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Attendance"
    has_cell = any((r.cell or "").strip() for r in body.rows)
    headers = (["Date", "Manager", "Worker", "Lavozim"] + (["Yacheyka"] if has_cell else [])
               + ["Jadval", "Keldi", "Ketdi", "Holat", "Soat", "Early Arrival (min)", "Eff. Hours"])
    ws.append(headers)
    fill = PatternFill(fill_type="solid", fgColor="1C4ED8")
    for i in range(1, len(headers) + 1):
        c = ws.cell(1, i)
        c.font, c.fill = Font(bold=True, color="FFFFFF", size=10), fill
        c.alignment = Alignment(horizontal="center", vertical="center")
    even = PatternFill(fill_type="solid", fgColor="F1F5F9")
    for n, r in enumerate(body.rows, 2):
        row = [body.attend_date, manager, r.worker_name or "", r.job_title or ""]
        if has_cell:
            row.append(r.cell or "")
        row += [r.schedule or "", r.clock_in or "", r.clock_out or "", r.status or "",
                r.hours_worked, r.early_arrival_min, r.effective_hours]
        ws.append(row)
        if n % 2 == 0:
            for i in range(1, len(headers) + 1):
                ws.cell(n, i).fill = even
    if body.extra_hours:
        ws.append([])
        ws.append(["", "", "Qo'shimcha soat", "", *([""] if has_cell else []), "", "", "", "",
                   round(body.extra_hours, 2)])
    widths = [13, 30, 42, 26] + ([12] if has_cell else []) + [18, 9, 9, 16, 8, 18, 12]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[ws.cell(1, i).column_letter].width = w
    buf = BytesIO()
    wb.save(buf)
    fname = f"verifix_live_{body.attend_date}.xlsx"
    resp = deliver_xlsx(request, caller, fname, buf.getvalue(),
                        f"📊 Jonli davomat — {body.attend_date} • {manager} • {len(body.rows)}",
                        chat_id=_tg(caller))
    action_log.enrich(target_kind="day", target_id=f"{body.manager_id}:{body.attend_date}",
                      unit_id=body.manager_id, unit_name=manager,
                      details=[("date", body.attend_date), ("rows", len(body.rows)), ("file", fname)])
    return resp
