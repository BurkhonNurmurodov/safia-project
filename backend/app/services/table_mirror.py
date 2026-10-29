"""A table as it stands on the screen, laid out as a workbook.

(Since 4.339 nothing in the current bundle calls this: the «Pozitsiyalar»
«Excel» is the «ABC форма» again, with the table's extra columns inserted —
routers/production.export_positions. Kept for a tab still open on 4.338.)

The «Zagruzka fayli» positions table (and any other picker-equipped table that
wants it) is exported as EXACTLY what the reader sees: the columns the
ColumnsPicker shows, in the order it shows them; the rows the search, the
filters and the sort leave, in that order; and every cell printed the way the
screen prints it (the same decimals, the same words, the same «—»). The page
builds the spec off the very arrays it renders from, so a column added to the
table reaches the file without a change here — this module computes nothing,
it only lays the spec out in the house report look (gold title and header,
hairline grid, zebra rows, numbers as real numbers under number formats).

    title, sheet, lines[]          the banner, the sheet name, the scope lines
    cols:  [{label, align}]        align = "left" | "center" | "right"
    rows:  [[cell]]                one list per table row, one cell per column

    cell:  {t}                     text as printed («—» for nothing)
           {v, d}                  a number shown with up to `d` decimals
           {v, k: "pct"}           a fraction shown as a whole percent
           fg                      the ink the screen gives it ("#rrggbb")
           b                       bold
           m                       muted — a placeholder, never a value
"""
from __future__ import annotations

from io import BytesIO
from typing import Any, Optional

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.properties import PageSetupProperties

GOLD = "C8973F"      # `--brand` — the banner and the header row
SUBBAND = "F5F2EC"   # the scope strip under the banner
ZEBRA = "FAF9F7"     # even rows
RULE = "E5E7EB"      # hairline grid
INK = "1F2937"
SOFT = "6B7280"
MUTED = "9CA3AF"     # a placeholder («—») — absent data, never a value

MAX_COLS = 60
MAX_ROWS = 20_000
MAX_CELLS = 600_000
MAX_TEXT = 500
MAX_LINES = 8
ALIGNS = {"left", "center", "right"}


def _hex(v: Any) -> Optional[str]:
    s = str(v or "").strip().lstrip("#")
    if len(s) == 3:
        s = "".join(ch * 2 for ch in s)
    return s.upper() if len(s) == 6 and all(c in "0123456789abcdefABCDEF" for c in s) else None


def _num(v: Any) -> Optional[float]:
    if isinstance(v, bool) or v is None:
        return None
    try:
        f = float(v)
    except (TypeError, ValueError):
        return None
    return f if f == f and f not in (float("inf"), float("-inf")) else None


def validate(spec: dict) -> Optional[str]:
    """None when the spec can be laid out; else why not, in one sentence."""
    cols = spec.get("cols") or []
    rows = spec.get("rows") or []
    if not isinstance(cols, list) or not cols:
        return "The table has no columns."
    if len(cols) > MAX_COLS:
        return f"At most {MAX_COLS} columns can be exported."
    if not isinstance(rows, list) or len(rows) > MAX_ROWS:
        return f"At most {MAX_ROWS} rows can be exported."
    if len(cols) * len(rows) > MAX_CELLS:
        return "The table is too large to export."
    for r in rows:
        if not isinstance(r, list) or len(r) != len(cols):
            return "Every row must carry one cell per column."
    return None


def _number_format(v: float, d: int) -> str:
    """The screen's «up to d decimals» (toLocaleString maximumFractionDigits):
    12 prints as 12, 12.5 as 12,5 — never 12,0. Decided per cell, from the value
    as the screen rounds it, so the file reads digit for digit like the table."""
    d = max(0, min(int(d), 6))
    shown = f"{round(v, d):.{d}f}".rstrip("0").rstrip(".") if d else ""
    places = len(shown.split(".")[1]) if "." in shown else 0
    return "#,##0" + ("." + "0" * places if places else "")


def build(spec: dict) -> BytesIO:
    cols: list[dict] = spec["cols"]
    rows: list[list[dict]] = spec["rows"]
    title = str(spec.get("title") or "").strip()[:MAX_TEXT]
    lines = [str(x).strip()[:MAX_TEXT] for x in (spec.get("lines") or [])[:MAX_LINES] if str(x).strip()]
    sheet = (str(spec.get("sheet") or title or "Sheet").strip() or "Sheet")
    for ch in '[]:*?/\\':
        sheet = sheet.replace(ch, " ")
    n = len(cols)
    last = get_column_letter(n)

    hair = Side(style="thin", color=RULE)
    grid = Border(left=hair, right=hair, top=hair, bottom=hair)
    zebra = PatternFill("solid", fgColor=ZEBRA)

    wb = Workbook()
    ws = wb.active
    ws.title = sheet[:31]

    def band(row: int, text: str, *, fill: str, font: Font, height: int) -> None:
        if n > 1:
            ws.merge_cells(f"A{row}:{last}{row}")
        c = ws.cell(row, 1, text)
        c.fill = PatternFill("solid", fgColor=fill)
        c.font = font
        c.alignment = Alignment(horizontal="left", vertical="center", indent=1)
        ws.row_dimensions[row].height = height

    y = 1
    band(y, title, fill=GOLD, font=Font(bold=True, color="FFFFFF", size=14), height=30)
    for line in lines:
        y += 1
        band(y, line, fill=SUBBAND, font=Font(size=9, color=SOFT), height=18)
    y += 1
    ws.row_dimensions[y].height = 6          # breathing room, not an empty row
    head = y + 1

    widths = [0] * n
    for i, col in enumerate(cols, 1):
        label = str(col.get("label") or "")[:MAX_TEXT]
        c = ws.cell(head, i, label)
        c.fill = PatternFill("solid", fgColor=GOLD)
        c.font = Font(bold=True, color="FFFFFF", size=10)
        c.border = grid
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        widths[i - 1] = max(widths[i - 1], min(len(label), 18))
    ws.row_dimensions[head].height = 30

    for ri, row in enumerate(rows, 1):
        yy = head + ri
        for ci, (col, cell) in enumerate(zip(cols, row), 1):
            cell = cell if isinstance(cell, dict) else {"t": "" if cell is None else str(cell)}
            align = col.get("align") if col.get("align") in ALIGNS else "left"
            v = _num(cell.get("v"))
            c = ws.cell(yy, ci)
            if v is not None and cell.get("k") == "pct":
                c.value = v
                c.number_format = "0%"
                text_len = len(f"{v * 100:.0f}%")
            elif v is not None:
                c.value = v
                fmt = _number_format(v, cell.get("d") if isinstance(cell.get("d"), int) else 2)
                c.number_format = fmt
                text_len = len(f"{v:,.2f}")
            else:
                text = str(cell.get("t") if cell.get("t") is not None else "")[:MAX_TEXT]
                c.value = text
                text_len = len(text)
            ink = MUTED if cell.get("m") else (_hex(cell.get("fg")) or INK)
            c.font = Font(size=10, color=ink, bold=bool(cell.get("b")))
            c.alignment = Alignment(horizontal=align, vertical="center")
            c.border = grid
            if ri % 2 == 0:
                c.fill = zebra
            widths[ci - 1] = max(widths[ci - 1], text_len)

    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = max(7, min(w + 3, 60))
    if rows:
        ws.auto_filter.ref = f"A{head}:{last}{head + len(rows)}"
    ws.freeze_panes = ws.cell(head + 1, 1)
    ws.print_title_rows = f"{head}:{head}"
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.tabColor = GOLD
    ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.print_options.horizontalCentered = True

    out = BytesIO()
    wb.save(out)
    out.seek(0)
    return out
