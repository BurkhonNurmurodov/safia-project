"""One-off: the whole Трудоемкость register, as a workbook in the operator's DM.

The operator asked on 2026-10-14 for every catalog line's trudoyomkost as one
table: supervisor · SAP code · work centre («Команда») · process name ·
trudoyomkost. Production is not readable from the machine that wrote this, so
the table is DMed from the server. v1 printed minutes; the operator then asked
for SECONDS, the unit the catalog stores (`labor_time`), so v2 prints it as
stored. v3 (2026-10-19) adds the line's «Faza» (the SAP operation step).

What it lists (the operator's answers, 2026-10-14):
  * the CURRENT catalog — `pp_products`, which is the catalog from the latest
    dated boundary on (`pp_catalog`); frozen past versions are not listed;
  * ACTIVE lines of units that are NOT archived — what the загрузка counts;
  * the product SAP code AND the work centre, because one product is often on
    several lines (other work centres, other operations) and the code alone
    would print rows that look identical;
  * trudoyomkost per PIECE in SECONDS = `labor_time`, as stored. A line with no
    Трудоемкость is a blank cell, never 0, and is counted in the caption;
  * «Faza» by the «Позиции» table's own rule (`production._build_dashboard`):
    the фаза pinned on the line wins, else the operations the plant-wide фаза
    file lists for the line's (SKU, work centre), joined « / ». The page reads
    ONE day's file; a register has no day, so the newest file within
    `FAZA_LOOKBACK_DAYS` that names the pair answers, and its date is printed
    beside it. A line no file names is blank and counted.

It READS and writes nothing but its flag (written by `startup._send_report_once`).

Temporary: delete this module, `startup.report_trudoyomkost_register` /
`_trudoyomkost_register_job` and the call in BOTH entrypoints once it has been
sent.
"""
from __future__ import annotations

import time
from datetime import datetime, timedelta
from io import BytesIO
from zoneinfo import ZoneInfo

import requests
from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Font, PatternFill
from sqlalchemy.orm import Session

from app.config import settings
from app.models import Manager, PPProduct, PPUpload

TZ = ZoneInfo("Asia/Tashkent")
_API = "https://api.telegram.org"
SEND_RETRIES = 3

FAZA_LOOKBACK_DAYS = 90

HEAD = ["Supervisor", "SAP code", "Work centre (Команда)", "Faza (Опер.)",
        "Process name", "Trudoyomkost, sec", "Faza from"]
WIDTHS = [30, 16, 22, 14, 60, 18, 22]
SEC_COL = 6


def _txt(v) -> str:
    return ILLEGAL_CHARACTERS_RE.sub("", str(v or "").strip())


def _seconds(labor) -> float | None:
    """Seconds per piece, as stored. None stays None (blank, not 0)."""
    if labor is None:
        return None
    try:
        return round(float(labor), 4)
    except (TypeError, ValueError):
        return None


def _file_ops(db: Session, wanted: set, today) -> dict:
    """{(SKU, work centre): (ops joined « / », file date)} for the pairs in
    `wanted`, newest plant-wide фаза file first — the page's own fold of one
    file ([order, op, wc, sku, …], distinct ops sorted), applied file by file
    until every pair is answered or the look-back runs out."""
    found: dict = {}
    if not wanted:
        return found
    floor = today - timedelta(days=FAZA_LOOKBACK_DAYS)
    ups = (db.query(PPUpload.id, PPUpload.date)
           .filter(PPUpload.manager_id.is_(None), PPUpload.file_type == "faza",
                   PPUpload.date <= today, PPUpload.date >= floor)
           .order_by(PPUpload.date.desc()).all())
    for uid, day in ups:
        rows = db.query(PPUpload.rows).filter(PPUpload.id == uid).scalar() or []
        ops: dict = {}
        for fr in rows:
            if not isinstance(fr, list) or len(fr) < 4:
                continue
            op = str(fr[1] or "").strip()
            wc = str(fr[2] or "").strip()
            sku = str(fr[3] or "").strip()
            if not op or not sku or sku == "—":
                continue
            key = (sku, wc)
            if key not in wanted or key in found:
                continue
            got = ops.setdefault(key, [])
            if op not in got:
                got.append(op)
        for key, got in ops.items():
            found[key] = (" / ".join(sorted(got)), day)
        if len(found) >= len(wanted):
            break
    return found


def collect(db: Session) -> dict:
    at = datetime.now(TZ)
    units = {m.id: m for m in db.query(Manager).filter(Manager.archived.is_(False)).all()}
    lines = (db.query(PPProduct)
             .filter(PPProduct.active.is_(True),
                     PPProduct.manager_id.in_(list(units) or [-1]))
             .all())
    wanted = {(_txt(p.sap_code), _txt(p.work_center)) for p in lines
              if not (p.op or "").strip() and _txt(p.sap_code)}
    from_file = _file_ops(db, wanted, at.date())
    rows = []
    for p in lines:
        m = units[p.manager_id]
        sap, wc = _txt(p.sap_code), _txt(p.work_center)
        pinned = _txt(p.op)
        if pinned:
            faza, src = pinned, "pinned on the line"
        elif (sap, wc) in from_file:
            faza, day = from_file[(sap, wc)]
            src = f"SAP file {day.strftime('%d.%m.%Y')}"
        else:
            faza, src = "", ""
        rows.append({
            "sup": _txt(m.name),
            "sap": sap,
            "wc": wc,
            "faza": _txt(faza),
            "faza_src": src,
            "name": _txt(p.name),
            "sec": _seconds(p.labor_time),
            "_order": (p.sort_order or 0, p.id),
        })
    rows.sort(key=lambda r: (r["sup"].lower(), r["wc"], r["_order"]))
    return {
        "rows": rows,
        "units": len({r["sup"] for r in rows}),
        "no_labor": sum(1 for r in rows if r["sec"] is None),
        "no_code": sum(1 for r in rows if not r["sap"]),
        "faza_pinned": sum(1 for r in rows if r["faza_src"].startswith("pinned")),
        "faza_file": sum(1 for r in rows if r["faza_src"].startswith("SAP")),
        "no_faza": sum(1 for r in rows if not r["faza"]),
        "at": at,
    }


def build_workbook(rep: dict) -> BytesIO:
    wb = Workbook()
    ws = wb.active
    ws.title = "Trudoyomkost"
    ws.append(HEAD)
    bold = Font(bold=True)
    fill = PatternFill("solid", fgColor="F2E6CF")
    for c in ws[1]:
        c.font = bold
        c.fill = fill
        c.alignment = Alignment(vertical="center", wrap_text=True)
    for r in rep["rows"]:
        ws.append([r["sup"], r["sap"] or None, r["wc"], r["faza"] or None,
                   r["name"], r["sec"], r["faza_src"] or None])
    for row in ws.iter_rows(min_row=2, min_col=SEC_COL, max_col=SEC_COL):
        for c in row:
            c.number_format = "General"
    for row in ws.iter_rows(min_row=2, min_col=4, max_col=4):
        for c in row:
            c.number_format = "@"
    for i, w in enumerate(WIDTHS, start=1):
        ws.column_dimensions[chr(64 + i)].width = w
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:G{max(1, len(rep['rows']) + 1)}"
    out = BytesIO()
    wb.save(out)
    out.seek(0)
    return out


def caption(rep: dict) -> str:
    lines = [
        f"Trudoyomkost register · {rep['at'].strftime('%d.%m.%Y %H:%M')}",
        f"{len(rep['rows'])} catalog lines · {rep['units']} supervisors",
        "Current catalog, active lines of active units. "
        "Trudoyomkost = seconds per piece (Трудоемкость as stored).",
        f"Faza: pinned on the line {rep['faza_pinned']} · from the SAP фаза file "
        f"{rep['faza_file']} (newest file in {FAZA_LOOKBACK_DAYS} days naming the "
        f"SKU + Команда) · none {rep['no_faza']}.",
    ]
    if rep["no_labor"]:
        lines.append(f"{rep['no_labor']} line(s) have no trudoyomkost — blank cells.")
    if rep["no_code"]:
        lines.append(f"{rep['no_code']} line(s) have no SAP code.")
    return "\n".join(lines)[:1000]


def _post(method: str, data: dict, files: dict | None = None) -> None:
    last = ""
    for attempt in range(1, SEND_RETRIES + 1):
        wait = 0
        try:
            r = requests.post(f"{_API}/bot{settings.telegram_bot_token}/{method}",
                              data=data, files=files, timeout=180)
            body = r.json()
            if body.get("ok"):
                return
            last = body.get("description") or f"HTTP {r.status_code}"
            wait = int(((body.get("parameters") or {}).get("retry_after")) or 0)
        except Exception as exc:
            last = type(exc).__name__
        if attempt < SEND_RETRIES:
            time.sleep(max(wait, 3 * attempt))
    raise RuntimeError(f"{method} failed: {last}"[:300])


def send(db: Session, chat_id: int, *_window) -> int:
    if not settings.telegram_bot_token:
        raise RuntimeError("telegram bot token not configured")
    rep = collect(db)
    book = build_workbook(rep).getvalue()
    db.rollback()          # hold no transaction across the Telegram call
    stamp = rep["at"].strftime("%Y-%m-%d")
    _post("sendDocument", {"chat_id": chat_id, "caption": caption(rep)},
          files={"document": (f"trudoyomkost-register-{stamp}.xlsx", book,
                              "application/vnd.openxmlformats-officedocument."
                              "spreadsheetml.sheet")})
    return 1
