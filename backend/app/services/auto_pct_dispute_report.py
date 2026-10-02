"""Objections to task #9 («plan %»), checked against WHEN the ФАКТ was typed — DMed once.

The operator, 2026-10-02: more and more leaders object to a failed #9
(`leader_auto`, `plan_pct` — «Bajarish %» of any one of their work centres at
or above the target by the check hour) with the same words — «I entered 50%, it
was not accepted, but it was entered on time». The check stores what the page
said AT the hour (`leader_auto_checks.facts`), but nothing on the objection says
when the leader actually typed their ФАКТ. This report answers that, objection
by objection, so each can be ruled on facts rather than on the leader's word.

For every objection to an automatic #9 failure (`leader_dispute.auto_entry`,
any status but cancelled) it collects:

* what the check SAW at its hour — the ledger's own facts;
* every ФАКТ save on the leader's work centres for that production date, out of
  the action register (`production.override_set`), with who and when, plus
  every SAP фаза upload of the date that reached the unit;
* the «Bajarish %» REBUILT at the check hour from those saves (`_fact_lb`, a
  floor: only what a stored row untouched since, or a logged save, proves was
  there), the same figure as of today, and the FIRST moment the target was
  reached — so «on time» and «typed after the hour» are told apart;
* ФАКТ the same leader saved on the dates either side, on the same work
  centres — the «typed it on the wrong day's page» case.

Each objection gets ONE verdict: `on_time` (the rebuilt % at the hour meets the
target — the check got it wrong and the objection should be approved), `late`
(the target was reached, but only after the hour), `never` (below the target
even today), `unexplained` (meets the target today with no logged save or
upload after the hour — something wrote outside the logged doors), `no_plan`.

It READS and writes nothing but its own flag. The unit-date replay is COPIED
from `auto_check_restore` so deleting either one-shot cannot break the other.

Temporary: delete this module, `startup.report_auto_pct_disputes`,
`startup._auto_pct_disputes_job` and the call in BOTH entrypoints once the files
have landed — a call left behind imports a deleted module at boot, and a failed
boot rolls the deploy back.
"""
from __future__ import annotations

import io
import json
import logging
import time
from collections import Counter
from datetime import date as _date, datetime, timedelta, timezone

import requests
from sqlalchemy.orm import Session

from app.config import settings
from app.models import (
    ActionLog, Cell, LeaderAiDispute, LeaderAutoCheck, Manager, PPDaily,
    PPLineDaily, PPManagerSetting, RoleProfile,
)
from app.services import cell_lookup, latin_code, leader_auto, leader_dispute

log = logging.getLogger(__name__)

TZ = leader_auto.TASHKENT
CHECK = "plan_pct"
_API = "https://api.telegram.org"
SEND_RETRIES = 3

VERDICTS = ("on_time", "late", "never", "unexplained", "no_plan", "error")
VERDICT_TEXT = {
    "on_time": "met the target BEFORE the hour — the check was wrong",
    "late": "the target was reached only AFTER the hour",
    "never": "below the target even today",
    "unexplained": "meets the target today, but no logged save explains when",
    "no_plan": "no plan on the leader's work centres",
    "error": "could not be rebuilt",
}


# ── small helpers ────────────────────────────────────────────────────────────

def _aware(v: datetime | None) -> datetime | None:
    if v is None:
        return None
    return v if v.tzinfo else v.replace(tzinfo=timezone.utc)


def _hm(v: datetime | None) -> str:
    v = _aware(v)
    return v.astimezone(TZ).strftime("%d.%m %H:%M") if v else "—"


def _d(iso: str) -> _date:
    return datetime.strptime(str(iso)[:10], "%Y-%m-%d").date()


def _norm(code) -> str:
    return cell_lookup.norm_code(latin_code.latin_code(str(code or "")) if code else code) or ""


def _num(v):
    try:
        return float(v) if v is not None else None
    except (TypeError, ValueError):
        return None


def _details(r: ActionLog) -> dict:
    out = {}
    for pair in (r.details or []):
        try:
            out[str(pair[0])] = pair[1]
        except Exception:
            continue
    return out


def _actor(r: ActionLog) -> str:
    return f"{(r.actor_name or '?').strip()} ({r.actor_role or r.source or '?'})"


# ── the unit-date replay (ФАКТ half, copied from auto_check_restore._Unit) ──

class _Unit:
    def __init__(self, db: Session, mid: int, date: str):
        self.db, self.mid, self.date, self.d = db, mid, date, _d(date)
        st = db.query(PPManagerSetting).filter_by(manager_id=mid).first()
        self.auto_fill = True if st is None else bool(st.auto_fill)
        self._facts = self._uploads = self._sapwc = self._pp = self._pl = None

    def facts(self) -> list[dict]:
        if self._facts is None:
            out = []
            for r in (self.db.query(ActionLog)
                      .filter(ActionLog.unit_id == self.mid, ActionLog.day == self.d,
                              ActionLog.outcome == "done",
                              ActionLog.action == "production.override_set")
                      .order_by(ActionLog.created_at, ActionLog.id).all()):
                det = _details(r)
                for ch in (r.changes or []):
                    try:
                        field, old, new = ch[0], ch[1], ch[2]
                    except Exception:
                        continue
                    if field == "plan":
                        continue
                    out.append({"at": _aware(r.created_at), "wc": _norm(det.get("work_center")),
                                "sap": str(det.get("sap_code") or ""),
                                "line": det.get("line") or None,
                                "old": _num(old), "new": _num(new), "who": _actor(r),
                                "profile": r.actor_profile_key})
            self._facts = out
        return self._facts

    def uploads(self) -> list[dict]:
        if self._uploads is None:
            out = []
            for r in (self.db.query(ActionLog)
                      .filter(ActionLog.action == "production.phase_uploaded",
                              ActionLog.day == self.d, ActionLog.outcome == "done")
                      .order_by(ActionLog.created_at).all()):
                det = _details(r)
                out.append({"at": _aware(r.created_at), "mode": det.get("mode"),
                            "type": det.get("type") or "", "note": det.get("note"),
                            "who": _actor(r)})
            self._uploads = out
        return self._uploads

    def sap_wcs(self) -> set:
        if self._sapwc is None:
            self._sapwc = {_norm(w) for (w,) in self.db.query(PPDaily.work_center).filter(
                PPDaily.manager_id == self.mid, PPDaily.date == self.d,
                PPDaily.plan_qty > 0).all()}
        return self._sapwc

    def reaching_faza(self) -> list[dict]:
        got_sap = bool(self.sap_wcs())
        return [x for x in self.uploads() if "faza" in x["type"]
                and (self.auto_fill if x["note"] == "auto-fill" else got_sap)]

    def facts_at(self, t: datetime) -> tuple[dict, dict]:
        """The hand-typed ФАКТ standing at `t`: `(per line, per group)`."""
        lines: dict = {}
        groups: dict = {}
        evs = [("typed", e["at"], e) for e in self.facts()]
        evs += [("upload", x["at"], x) for x in self.reaching_faza()
                if x["mode"] in ("both", "actual")]
        for kind, at, e in sorted(evs, key=lambda z: z[1]):
            if at > t:
                break
            if kind == "upload":
                lines.clear()
                groups.clear()
                continue
            g = (e["sap"], e["wc"])
            if e["line"]:
                k = g + (str(e["line"]),)
                if e["new"] is None:
                    lines.pop(k, None)
                else:
                    lines[k] = e["new"]
            else:
                for k in [k for k in lines if k[:2] == g]:
                    lines.pop(k)
                if e["new"] is None:
                    groups.pop(g, None)
                else:
                    groups[g] = e["new"]
        return lines, groups

    def pp_rows(self) -> dict:
        if self._pp is None:
            self._pp = {(str(r.sap_code or ""), _norm(r.work_center)): r
                        for r in self.db.query(PPDaily).filter(
                            PPDaily.manager_id == self.mid, PPDaily.date == self.d).all()}
        return self._pp

    def line_rows(self) -> dict:
        if self._pl is None:
            self._pl = {(str(r.qty_key or ""), _norm(r.work_center), str(r.line_key or "")): r
                        for r in self.db.query(PPLineDaily).filter(
                            PPLineDaily.manager_id == self.mid,
                            PPLineDaily.date == self.d).all()}
        return self._pl

    def fact_lb(self, rows: list, t: datetime) -> float:
        """ФАКТ minutes that PROVABLY stood on these positions at `t` — a floor."""
        lines, groups = self.facts_at(t)
        pp, pl = self.pp_rows(), self.line_rows()
        total = 0.0
        for r in rows:
            labor = _num(r.get("labor_time"))
            if not labor:
                continue
            g = (str(r.get("qty_key") or ""), _norm(r.get("work_center")))
            k = g + (str(r.get("line_key") or ""),)
            val = None
            lrow = pl.get(k)
            if lrow is not None and _aware(lrow.updated_at) and _aware(lrow.updated_at) <= t:
                val = _num(lrow.actual_override)
            elif k in lines:
                val = lines[k]
            if val is None:
                grow = pp.get(g)
                if grow is not None and _aware(grow.updated_at) and _aware(grow.updated_at) <= t:
                    if grow.actual_override is not None:
                        val = _num(grow.actual_override)
                    elif r.get("sap_filled"):
                        val = _num(grow.actual_qty)
                elif g in groups:
                    val = groups[g]
            if val:
                total += labor * val / 60.0
        return total


# ── one objection ────────────────────────────────────────────────────────────

def _code_page(db: Session, mid: int, date: str, code: str) -> tuple[list, float, float]:
    from app.routers.production import _build_dashboard
    dash = _build_dashboard(db, mid, _d(date), wc_scope={code}, payload=None)
    t = dash.get("totals") or {}
    return (dash.get("rows") or [], float(t.get("total_plan_labor") or 0),
            float(t.get("total_actual_labor") or 0))


def _other_days(db: Session, mid: int, date: str, codes: set, prof_key: str) -> list[dict]:
    """ФАКТ this leader saved on the neighbouring dates, on their work centres."""
    d = _d(date)
    out = []
    for r in (db.query(ActionLog)
              .filter(ActionLog.unit_id == mid, ActionLog.outcome == "done",
                      ActionLog.action == "production.override_set",
                      ActionLog.day.in_([d - timedelta(days=1), d + timedelta(days=1)]),
                      ActionLog.actor_profile_key == prof_key)
              .order_by(ActionLog.created_at).all()):
        det = _details(r)
        if _norm(det.get("work_center")) not in codes:
            continue
        for ch in (r.changes or []):
            if ch and ch[0] != "plan":
                out.append({"at": _aware(r.created_at), "day": str(r.day),
                            "wc": _norm(det.get("work_center")),
                            "sap": str(det.get("sap_code") or ""), "new": _num(ch[2])})
    return out


def _one(db: Session, d: LeaderAiDispute, row: LeaderAutoCheck, now: datetime,
         timeline: list) -> dict:
    prof = db.query(RoleProfile).filter_by(id=row.leader_id).first()
    m = db.query(Manager).filter_by(id=row.manager_id).first()
    facts = row.facts or {}
    target = float(facts.get("target") or 30)
    due = _aware(row.due_at) or _aware(row.checked_at)
    cell = db.query(Cell).filter_by(id=row.cell_id).first() if row.cell_id else None
    rec = {"dispute_id": d.id, "date": row.date, "shift": getattr(m, "shift", None),
           "unit": getattr(m, "name", None), "leader": getattr(prof, "name", None) or d.leader_name,
           "status": d.status, "objection": (d.reason or "").strip(),
           "filed": _hm(d.requested_at), "hour": _hm(due), "code_then": row.code,
           "target": target, "then_pct": facts.get("best_pct", facts.get("pct")),
           "then_fact": facts.get("fact_min"), "then_plan": facts.get("plan_min")}
    if prof is None or m is None or due is None:
        rec.update(verdict="error", note="leader, unit or check hour missing")
        return rec
    ctx = leader_auto._Ctx(db, prof, m, m.shift, row.date, cell, due, now)
    mid = ctx.unit_id
    codes = sorted({c for c, _g in ctx.pairs if c})
    rec["cells"] = ", ".join(leader_auto.cell_label(c) for c in ctx.cells) or "—"
    if not codes:
        rec.update(verdict="no_plan", note="no SAP code on the leader's cells")
        return rec
    u = _Unit(db, mid, row.date)
    pages = {}
    for c in codes:
        rows, plan, fact = _code_page(db, mid, row.date, c)
        if plan > 0:
            pages[c] = (rows, plan, fact)
    if not pages:
        rec.update(verdict="no_plan", note="no plan minutes on the leader's work centres today")
        return rec

    def pct_at(t):
        best = (None, -1.0)
        for c, (rows, plan, _f) in pages.items():
            p = u.fact_lb(rows, t) / plan * 100
            if p > best[1]:
                best = (c, p)
        return best

    hour_code, hour_pct = pct_at(due)
    now_code, now_pct = max(((c, f / p * 100) for c, (_r, p, f) in pages.items()),
                            key=lambda z: z[1])
    mine = [e for e in u.facts() if e["wc"] in pages]
    ups = u.reaching_faza()
    reached = None
    for t in sorted({e["at"] for e in mine} | {x["at"] for x in ups}):
        if pct_at(t)[1] >= target:
            reached = t
            break
    before = [e for e in mine if e["at"] <= due]
    after = [e for e in mine if e["at"] > due]
    plan_now = sum(p for _r, p, _f in pages.values())
    then_plan = _num(facts.get("plan_min"))

    if hour_pct >= target:
        verdict = "on_time"
    elif now_pct < target:
        verdict = "never"
    elif reached is not None and reached > due:
        verdict = "late"
    else:
        verdict = "unexplained"
    notes = []
    if then_plan is not None and abs(plan_now - then_plan) > max(1.0, 0.005 * then_plan):
        notes.append(f"plan changed since the check ({then_plan:g} → {plan_now:.1f} min)")
    if not mine:
        notes.append("no ФАКТ was typed by hand on this date (SAP only)")
    other = _other_days(db, mid, row.date, set(pages), f"leader:{prof.id}")
    if other:
        notes.append("ФАКТ typed by this leader on another date: " + "; ".join(
            f"{o['day']} {o['wc']} {o['sap']}={o['new']:g} at {_hm(o['at'])}"
            for o in other[:4] if o["new"] is not None))
    rec.update(
        verdict=verdict, hour_pct=round(hour_pct, 1), hour_cell=hour_code,
        now_pct=round(now_pct, 1), now_cell=now_code, reached=_hm(reached),
        first_fact=_hm(mine[0]["at"]) if mine else "—",
        last_before=_hm(before[-1]["at"]) if before else "—",
        saves_before=len(before), saves_after=len(after),
        uploads=", ".join(f"{_hm(x['at'])} ({x['mode']})" for x in ups) or "—",
        note="; ".join(notes))
    for e in mine:
        timeline.append({"dispute_id": d.id, "leader": rec["leader"], "date": row.date,
                         "at": _hm(e["at"]), "when": "before" if e["at"] <= due else "AFTER",
                         "kind": "ФАКТ typed", "wc": e["wc"], "sap": e["sap"],
                         "line": e["line"], "old": e["old"], "new": e["new"], "who": e["who"],
                         "pct": round(pct_at(e["at"])[1], 1)})
    for x in ups:
        timeline.append({"dispute_id": d.id, "leader": rec["leader"], "date": row.date,
                         "at": _hm(x["at"]), "when": "before" if x["at"] <= due else "AFTER",
                         "kind": f"SAP upload ({x['mode']})", "who": x["who"],
                         "pct": round(pct_at(x["at"])[1], 1)})
    return rec


def collect(db: Session, now: datetime | None = None) -> dict:
    now = now or datetime.now(timezone.utc)
    items, timeline, errors = [], [], []
    disputes = (db.query(LeaderAiDispute)
                .filter(LeaderAiDispute.status != "cancelled",
                        LeaderAiDispute.date >= leader_auto.AUTO_FROM)
                .order_by(LeaderAiDispute.date, LeaderAiDispute.id).all())
    for d in disputes:
        try:
            e = leader_dispute.auto_entry(db, d)
            if e is None:
                continue
            row = db.query(LeaderAutoCheck).filter_by(entry_id=e.id).first()
            if row is None or row.check != CHECK:
                continue
            items.append(_one(db, d, row, now, timeline))
        except Exception as exc:
            db.rollback()
            errors.append(f"objection {d.id}: {type(exc).__name__}: {exc}")
            log.exception("auto-pct dispute report: objection %s", d.id)
    timeline.sort(key=lambda r: (r["dispute_id"], r["at"]))
    return {"built_at": _hm(now), "items": items, "timeline": timeline, "errors": errors}


# ── output ───────────────────────────────────────────────────────────────────

def summary_text(rep: dict) -> str:
    items = rep["items"]
    c = Counter(i["verdict"] for i in items)
    lines = [f"Task #9 objections, checked against when the ФАКТ was typed "
             f"({rep['built_at']})", f"{len(items)} objection(s) to an automatic #9 failure.", ""]
    for v in VERDICTS:
        if c.get(v):
            lines.append(f"• {c[v]} — {VERDICT_TEXT[v]}")
    open_on_time = [i for i in items if i["verdict"] == "on_time"
                    and i["status"] in leader_dispute.OPEN_STATES]
    if open_on_time:
        lines += ["", "Still open and on time (approve these):"]
        for i in open_on_time[:15]:
            lines.append(f"  {i['date']} {i['leader']} — {i.get('hour_pct')}% at "
                         f"{i['hour']} (check saw {i['then_pct']}%)")
    if rep["errors"]:
        lines += ["", f"{len(rep['errors'])} could not be rebuilt — see «Errors»."]
    lines += ["", "«% at the hour» is a FLOOR: only ФАКТ a logged save or an "
                  "untouched stored row proves was there counts."]
    return "\n".join(lines)


def build_workbook(rep: dict) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font
    wb = Workbook()
    sheets = (
        ("Objections", rep["items"], [
            ("Objection", "dispute_id", 9), ("Date", "date", 11), ("Shift", "shift", 6),
            ("Brigadir", "unit", 24), ("Leader", "leader", 26), ("Cells", "cells", 12),
            ("Status", "status", 10), ("Check hour", "hour", 12),
            ("Target %", "target", 8), ("Check saw %", "then_pct", 10),
            ("% at hour (rebuilt)", "hour_pct", 12), ("Work centre", "hour_cell", 10),
            ("% today", "now_pct", 9), ("Target reached at", "reached", 14),
            ("First ФАКТ save", "first_fact", 14), ("Last save before hour", "last_before", 14),
            ("Saves before", "saves_before", 8), ("Saves after", "saves_after", 8),
            ("SAP uploads", "uploads", 26), ("Verdict", "verdict", 12),
            ("Objection text", "objection", 50), ("Notes", "note", 60)]),
        ("Timeline", rep["timeline"], [
            ("Objection", "dispute_id", 9), ("Leader", "leader", 26), ("Date", "date", 11),
            ("At", "at", 12), ("vs hour", "when", 8), ("What", "kind", 18),
            ("Work centre", "wc", 10), ("SAP", "sap", 14), ("Line", "line", 22),
            ("Old", "old", 8), ("New", "new", 8), ("Who", "who", 30),
            ("% after this", "pct", 10)]),
    )
    for i, (title, rows, cols) in enumerate(sheets):
        ws = wb.active if i == 0 else wb.create_sheet()
        ws.title = title
        ws.append([c[0] for c in cols])
        for r in rows:
            ws.append([r.get(k) for _t, k, _w in cols])
        for j, (_t, _k, width) in enumerate(cols, start=1):
            ws.column_dimensions[ws.cell(1, j).column_letter].width = width
        for cl in ws[1]:
            cl.font = Font(bold=True)
        for row in ws.iter_rows(min_row=2):
            for cl in row:
                cl.alignment = Alignment(wrap_text=True, vertical="top")
        ws.freeze_panes = "A2"
    ws = wb.create_sheet("Verdicts")
    ws.append(["Verdict", "Meaning"])
    for v in VERDICTS:
        ws.append([v, VERDICT_TEXT[v]])
    if rep["errors"]:
        ws = wb.create_sheet("Errors")
        ws.append(["Error"])
        for e in rep["errors"]:
            ws.append([e])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _clean(text: str) -> str:
    token = settings.telegram_bot_token or ""
    return text.replace(token, "***") if token else text


def _post(method: str, data: dict, files: dict | None = None) -> None:
    last = ""
    for attempt in range(1, SEND_RETRIES + 1):
        wait = 0
        try:
            r = requests.post(f"{_API}/bot{settings.telegram_bot_token}/{method}",
                              data=data, files=files, timeout=120)
            body = r.json()
            if body.get("ok"):
                return
            last = body.get("description") or f"HTTP {r.status_code}"
            wait = int(((body.get("parameters") or {}).get("retry_after")) or 0)
        except Exception as exc:
            last = f"{type(exc).__name__}: {exc}"
        if attempt < SEND_RETRIES:
            time.sleep(max(wait, 3 * attempt))
    raise RuntimeError(_clean(last or f"{method} failed")[:300])


def send(db: Session, chat_id: int, *_window) -> int:
    if not settings.telegram_bot_token:
        raise RuntimeError("telegram bot token not configured")
    rep = collect(db)
    db.rollback()
    _post("sendMessage", {"chat_id": chat_id, "text": summary_text(rep)[:4000]})
    _post("sendDocument", {"chat_id": chat_id,
                           "caption": "Task #9 objections — % at the check hour, rebuilt "
                                      "from when the ФАКТ was typed."},
          files={"document": ("task9-objections.xlsx", build_workbook(rep),
                              "application/vnd.openxmlformats-officedocument."
                              "spreadsheetml.sheet")})
    _post("sendDocument", {"chat_id": chat_id, "caption": "The same, as JSON."},
          files={"document": ("task9-objections.json",
                              json.dumps(rep, ensure_ascii=False, indent=1, default=str)
                              .encode("utf-8"), "application/json")})
    return 3
