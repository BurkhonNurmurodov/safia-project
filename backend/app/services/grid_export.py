"""One /zagruzka table, as a workbook that looks like the table.

The six tables on «Zagruzka foizi» (three comparison tables, three heatmaps)
compute their figures in the browser — `utils/formulas.js` is THE definition
of every one of them, and a Python twin of it would be a second spelling that
drifts. So the table DESCRIBES ITSELF instead: the grid exactly as it is drawn
(rows in screen order, the summary column in its current AVG / MIN / MAX mode,
the P·A or D view, each cell's band colour) arrives as a spec, and this module
only lays it out. Nothing here computes a figure.

The spec (built by `utils/gridExport.js` on the client):

    title, lines[], sheet
    cols:   [{w, sep}]           one per column; `sep` = a thick LEFT edge
                                 (a date group, the summary column)
    head:   [[cell]]             the gold header rows
    body:   [[cell]]             one row per brigadir
    foot:   [[cell]]             after a blank spacer row (the unit's AVG row)
    legend: [{title, items: [{color, label, hatch, icon}]}]

    cell:   {t}                  text
            {v, k}               a whole percent; k = "pct" | "diff"
            bg / fg              the band fill and its contrast ink
            hatch                a pending day (the screen's diagonal stripes)
            dim                  the P sub-label, a softer white on gold
            span                 columns it covers (default 1)

Every data cell is FILLED with its band colour, the way the screen paints it —
never a coloured number on white — and the value is a real number under a
percent format, so the file can be sorted and summed.
"""
from __future__ import annotations

import re
from io import BytesIO
from typing import Any, Optional

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

FONT = "Segoe UI"
BRAND = "C8973F"        # the header gold — `--brand`
HEAD_INK = "FFFFFF"     # white label on gold — `--on-brand`
HEAD_DIM = "F4E9D6"     # rgba(255,255,255,.75) on gold, flattened
INK = "111827"          # `--text-1` (light theme)
NAME_INK = "374151"
SOFT = "5C6370"         # `--text-3`
FAINT = "9CA3AF"        # `--text-4`
DASH = "C9CDD3"         # the «—» of an empty cell (25% ink)
LINE = "E5E7EB"         # `--border`
EDGE = "C4C9D0"         # `--border-md`, the 2px separators
HATCH = "D9DCE1"
INNER = "F5F1EA"        # `--bg-inner`

PCT = '0"%"'
DIFF = '+0"%";-0"%";0"%"'

# A spec is the caller's own description of a table they are looking at, and
# it goes back to them alone — but it is still bounded, so a malformed request
# can never ask this process for a workbook the size of the plant's history.
MAX_COLS = 1200
MAX_ROWS = 2000
MAX_CELLS = 400_000
MAX_TEXT = 300

_HEX = re.compile(r"^#?([0-9a-fA-F]{6}|[0-9a-fA-F]{3})$")
_SHEET_BAD = re.compile(r"[\[\]:*?/\\]")


def _hex(value: Any) -> Optional[str]:
    """A CSS colour as openpyxl's RRGGBB, or None for anything else
    («transparent», a CSS variable, garbage)."""
    m = _HEX.match(str(value or "").strip())
    if not m:
        return None
    h = m.group(1)
    if len(h) == 3:
        h = "".join(ch * 2 for ch in h)
    return h.upper()


def _text(value: Any) -> str:
    return str(value if value is not None else "")[:MAX_TEXT]


def _side(color: str, style: str = "thin") -> Side:
    return Side(style=style, color=color)


THIN = _side(LINE)
THICK = _side(EDGE, "medium")


def _fill(color: str) -> PatternFill:
    return PatternFill("solid", fgColor=color)


HATCH_FILL = PatternFill("lightUp", fgColor=HATCH, bgColor="FFFFFF")


def _sheet_name(name: str) -> str:
    return (_SHEET_BAD.sub(" ", name or "").strip() or "Jadval")[:31]


def validate(spec: dict) -> Optional[str]:
    """Why this spec cannot be laid out, or None."""
    cols = spec.get("cols") or []
    if not cols or len(cols) > MAX_COLS:
        return "bad_cols"
    rows = (spec.get("head") or []) + (spec.get("body") or []) + (spec.get("foot") or [])
    if len(rows) > MAX_ROWS:
        return "too_many_rows"
    if sum(len(r or []) for r in rows) > MAX_CELLS:
        return "too_many_cells"
    return None


def _widths(ws, cols: list[dict]) -> None:
    for i, c in enumerate(cols, start=1):
        try:
            w = float(c.get("w") or 8)
        except (TypeError, ValueError):
            w = 8.0
        ws.column_dimensions[get_column_letter(i)].width = max(3.0, min(w, 80.0))


def _border(col: int, n_cols: int, seps: set[int], top_thick: bool,
            span: int = 1) -> Border:
    """The screen's rules: thin everywhere, a thick edge after the name column,
    before every date group and the summary column, and over the footer."""
    left = THICK if col in seps else THIN
    right_col = col + span          # the column after this cell
    right = THICK if right_col in seps else THIN
    return Border(left=left, right=right,
                  top=THICK if top_thick else THIN, bottom=THIN)


def _place(ws, row: int, col: int, span: int, border: Border, fill=None) -> None:
    """Border (and fill) every cell of a horizontal span, then merge it, so the
    edges wrap the merged range rather than its first cell only."""
    for c in range(col, col + span):
        cell = ws.cell(row, c)
        cell.border = border
        if fill is not None:
            cell.fill = fill
    if span > 1:
        ws.merge_cells(start_row=row, start_column=col, end_row=row, end_column=col + span - 1)


def _head_row(ws, row: int, cells: list[dict], n_cols: int, seps: set[int]) -> None:
    ws.row_dimensions[row].height = 21
    col = 1
    for cell in cells:
        span = max(1, int(cell.get("span") or 1))
        span = min(span, n_cols - col + 1)
        if span <= 0:
            break
        _place(ws, row, col, span, _border(col, n_cols, seps, False, span), _fill(BRAND))
        top = ws.cell(row, col)
        # Upper case, as the screen's header (`text-transform: uppercase`).
        top.value = _text(cell.get("t")).upper() or None
        top.font = Font(name=FONT, size=9, bold=True,
                        color=HEAD_DIM if cell.get("dim") else HEAD_INK)
        top.alignment = Alignment(horizontal="left" if col == 1 else "center",
                                  vertical="center", indent=1 if col == 1 else 0)
        col += span


def _data_row(ws, row: int, cells: list[dict], n_cols: int, seps: set[int], *,
              top_thick: bool = False, label_style: str = "name", height: float = 24) -> None:
    ws.row_dimensions[row].height = height
    col = 1
    for cell in cells:
        span = max(1, int(cell.get("span") or 1))
        span = min(span, n_cols - col + 1)
        if span <= 0:
            break
        border = _border(col, n_cols, seps, top_thick, span)
        top = ws.cell(row, col)
        if col == 1:
            # The row's name — the column the eye returns to, never coloured.
            _place(ws, row, col, span, border, _fill(INNER) if label_style == "unit" else None)
            top.value = _text(cell.get("t"))
            if label_style == "stat":
                top.font = Font(name=FONT, size=8, bold=True, color=SOFT)
            else:
                top.font = Font(name=FONT, size=10, bold=bool(cell.get("b")) or label_style == "unit",
                                color=INK if label_style == "unit" else NAME_INK)
            top.alignment = Alignment(horizontal="left", vertical="center", indent=1)
            col += span
            continue

        if cell.get("hatch"):
            # A day whose figure cannot show yet — the screen's striped cell.
            _place(ws, row, col, span, border, HATCH_FILL)
            top.value = _text(cell.get("t")) or None
            top.font = Font(name=FONT, size=9, color=FAINT)
            top.alignment = Alignment(horizontal="center", vertical="center")
            col += span
            continue

        v = cell.get("v")
        bg = _hex(cell.get("bg"))
        if isinstance(v, (int, float)) and not isinstance(v, bool):
            _place(ws, row, col, span, border, _fill(bg) if bg else None)
            top.value = v
            top.number_format = DIFF if cell.get("k") == "diff" else PCT
            top.font = Font(name=FONT, size=9, bold=True, color=_hex(cell.get("fg")) or INK)
        else:
            # No figure on this day: the screen's faint dash on a blank cell.
            _place(ws, row, col, span, border)
            top.value = _text(cell.get("t")) or "—"
            top.font = Font(name=FONT, size=9, color=DASH)
        top.alignment = Alignment(horizontal="center", vertical="center")
        col += span


def _legend(ws, row: int, groups: list[dict], n_cols: int) -> int:
    """The colour key, one band per line: a filled swatch and what it means."""
    label_to = max(2, min(n_cols, 12))
    for group in groups[:10]:
        items = (group.get("items") or [])[:16]
        if not items:
            continue
        head = ws.cell(row, 1, _text(group.get("title")))
        head.font = Font(name=FONT, size=9, bold=True, color=SOFT)
        head.alignment = Alignment(horizontal="left", vertical="center", indent=1)
        for i, item in enumerate(items):
            r = row + i
            ws.row_dimensions[r].height = 18
            sw = ws.cell(r, 2)
            sw.border = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
            color = _hex(item.get("color"))
            if item.get("hatch"):
                sw.fill = HATCH_FILL
            elif color:
                sw.fill = _fill(color)
            if item.get("icon"):
                sw.value = _text(item.get("icon"))
                sw.alignment = Alignment(horizontal="center", vertical="center")
                sw.font = Font(name=FONT, size=9, color=FAINT)
            lab = ws.cell(r, 3, _text(item.get("label")))
            lab.font = Font(name=FONT, size=9, color=NAME_INK)
            lab.alignment = Alignment(horizontal="left", vertical="center", indent=1)
            if label_to > 3:
                ws.merge_cells(start_row=r, start_column=3, end_row=r, end_column=label_to)
        row += len(items) + 1
    return row


def build_grid_workbook(spec: dict) -> BytesIO:
    cols = spec.get("cols") or []
    n_cols = len(cols)
    # A thick LEFT edge before each flagged column, and before column 2 — the
    # name column's own separator.
    seps = {i for i, c in enumerate(cols, start=1) if c.get("sep") and i > 1}
    if n_cols >= 2:
        seps.add(2)                 # the name column's 2px right edge

    wb = Workbook()
    ws = wb.active
    ws.title = _sheet_name(_text(spec.get("sheet")) or _text(spec.get("title")))
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.tabColor = BRAND
    _widths(ws, cols)

    row = 1
    title = ws.cell(row, 1, _text(spec.get("title")))
    title.font = Font(name=FONT, size=14, bold=True, color=INK)
    title.alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[row].height = 24
    row += 1
    for line in (spec.get("lines") or [])[:12]:
        if not line:
            continue
        c = ws.cell(row, 1, _text(line))
        c.font = Font(name=FONT, size=9, color=SOFT)
        c.alignment = Alignment(horizontal="left", vertical="center")
        row += 1
    row += 1

    head = spec.get("head") or []
    head_top = row
    for cells in head:
        _head_row(ws, row, cells or [], n_cols, seps)
        row += 1
    # The name header stands over every header row when the rows under it
    # leave its column blank (the comparison tables' P | A row), as on screen.
    if len(head) > 1 and all(not _text(((r or [{}])[0] or {}).get("t")) for r in head[1:]):
        ws.merge_cells(start_row=head_top, start_column=1, end_row=row - 1, end_column=1)

    body = spec.get("body") or []
    for cells in body:
        _data_row(ws, row, cells or [], n_cols, seps)
        row += 1

    foot = spec.get("foot") or []
    if foot:
        # One blank row of air: the footer answers a different question from
        # the rows above it, exactly as the screen sets it apart.
        row += 1
        for i, cells in enumerate(foot):
            kind = "stat"
            first = (cells or [{}])[0] if cells else {}
            if first.get("unit"):
                kind = "unit"
            _data_row(ws, row, cells or [], n_cols, seps, top_thick=(i == 0),
                      label_style=kind, height=22 if kind == "stat" else 24)
            row += 1

    legend = spec.get("legend") or []
    if legend:
        row += 1
        _legend(ws, row, legend, n_cols)

    # The header rows and the name column stay on screen while scrolling.
    if head:
        ws.freeze_panes = ws.cell(head_top + len(head), 2)
        ws.print_title_rows = f"{head_top}:{head_top + len(head) - 1}"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True

    bio = BytesIO()
    wb.save(bio)
    bio.seek(0)
    return bio
