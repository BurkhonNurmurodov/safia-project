"""One-off: can the platform read the ORIGINAL on-time sheet, and does its own
arithmetic equal the sheet's?

The operator (2026-10-21) pointed «O'z vaqtida yetkazish» at the logistics
team's original workbook (`ORIGINAL`). The page was built and checked against a
COPY whose raw tabs had been removed, so nobody has yet seen whether the
original carries «заявки» / «поставка» / «заход» under the headers the reader
looks for, nor whether the page's numbers equal the original's «свод» on the
same days. Only production holds the Google service account, so the answer is
read there and DMed to the operator:

  1. which of the three sheets (the original, the one set on «Ma'lumotlar»,
     the old copy) the service account can open, and the account's address;
  2. every tab of the original — size, the first rows, the formulas on top of
     the summary tabs;
  3. the reader's own pass over the raw tabs (`ontime_sync.read_*`, nothing
     stored): which header fed which field, rows, days, what was skipped, the
     wave windows;
  4. the page's arithmetic (`ontime.evaluate`) on that data against the sheet's
     «свод» (every material × wave × dated day: ordered, counted, on-time) and
     «свод по группам» (group × wave means, the group's daily mean).

It READS and writes nothing but its flag (written by `startup._send_report_once`).
It does not change which sheet the page reads.

Temporary: delete this module, `startup.report_ontime_sheet` /
`_ontime_sheet_job` and the call in BOTH entrypoints once it has been sent.
"""
from __future__ import annotations

import time
from collections import defaultdict
from datetime import date, datetime, timedelta
from io import BytesIO
from typing import Any, Optional
from zoneinfo import ZoneInfo

import requests
from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Font
from sqlalchemy.orm import Session

from app.config import settings
from app.services import ontime, ontime_sync

ORIGINAL = "1pDPWN09n0aqW7-yQwTzeh_kMmHpmAnj0sdebgBYh-2I"
COPY = "1E9WQvPQS_OOQZg8sW00zSw1idvWrdQKcTuvmtxwbt-I"
TZ = ZoneInfo("Asia/Tashkent")
_API = "https://api.telegram.org"
SEND_RETRIES = 3

SAMPLE_ROWS = 6           # the first rows of every tab, as shown in the sheet
SAMPLE_COLS = 60
FORMULA_ROWS = 8          # formulas read on top of every non-raw tab
FORMULA_COLS = 40
MAX_TABS = 20
GRID_CELLS = 600_000      # the most a summary tab is read whole
DIFF_ROWS = 3000          # «свод» rows that differ, listed in the workbook
RECENT_DAYS = 14          # our group × day means also for the newest order days
RAW = {ontime_sync.TAB_ORDERS, ontime_sync.TAB_RECEIPTS}
SVOD = "свод"
SVOD_GROUPS = "свод по группам"
EPS = 1e-6


# ── small helpers ─────────────────────────────────────────────────────────────

def _col(i: int) -> str:
    return ontime_sync._col(i)


def _num(v) -> Optional[float]:
    return float(v) if isinstance(v, (int, float)) and not isinstance(v, bool) else None


def _hms(s: Optional[int]) -> str:
    if s is None:
        return "—"
    return f"{s // 3600:02d}:{s % 3600 // 60:02d}:{s % 60:02d}"


def _pct(v) -> str:
    return "—" if v is None else f"{100 * v:.2f}%"


def _g(v, nd: int = 0) -> str:
    if v is None:
        return "—"
    s = f"{v:,.{nd}f}".replace(",", " ")
    return s


def _d(v: Optional[date]) -> str:
    return v.strftime("%d.%m.%Y") if v else "—"


def _err(exc: Exception) -> dict:
    if isinstance(exc, ontime_sync.SheetError):
        return {"code": exc.code, "detail": exc.detail[:300]}
    return {"code": type(exc).__name__, "detail": str(exc)[:300]}


def _norm(s: Any) -> str:
    return ontime_sync._norm(s)


# ── 1. who can open what ──────────────────────────────────────────────────────

def _open(client, sid: str) -> tuple[Any, dict]:
    t0 = time.monotonic()
    try:
        sh = ontime_sync._call(client.open_by_key, sid)
        meta = ontime_sync._call(sh.fetch_sheet_metadata)
        props = meta.get("properties") or {}
        return sh, {"ok": True, "title": props.get("title"), "locale": props.get("locale"),
                    "tz": props.get("timeZone"), "tabs": len(meta.get("sheets") or []),
                    "meta": meta, "ms": int(1000 * (time.monotonic() - t0))}
    except Exception as exc:  # noqa: BLE001 — every refusal is an answer here
        from gspread.exceptions import APIError
        out = {"ok": False, **_err(exc), "ms": int(1000 * (time.monotonic() - t0))}
        if isinstance(exc, APIError):
            out["status"] = getattr(getattr(exc, "response", None), "status_code", None)
        return None, out


# ── 2. every tab ──────────────────────────────────────────────────────────────

def _tabs(meta: dict) -> list[dict]:
    out = []
    for s in meta.get("sheets") or []:
        p = s.get("properties") or {}
        g = p.get("gridProperties") or {}
        out.append({"title": p.get("title"), "gid": p.get("sheetId"), "index": p.get("index"),
                    "rows": g.get("rowCount"), "cols": g.get("columnCount"),
                    "frozen": g.get("frozenRowCount"), "hidden": bool(p.get("hidden"))})
    return out


def _sample(ws, tab: dict) -> dict:
    ncols = max(1, min(tab["cols"] or 1, SAMPLE_COLS))
    out = {"rows": [], "formulas": [], "error": None}
    try:
        out["rows"] = ontime_sync._call(
            ws.get, f"A1:{_col(ncols - 1)}{SAMPLE_ROWS}") or []
    except Exception as exc:  # noqa: BLE001
        out["error"] = _err(exc)
    if _norm(tab["title"]) in {_norm(t) for t in RAW}:
        return out
    try:
        fc = max(1, min(tab["cols"] or 1, FORMULA_COLS))
        vals = ontime_sync._call(ws.get, f"A1:{_col(fc - 1)}{FORMULA_ROWS}",
                                 value_render_option="FORMULA") or []
        for ri, r in enumerate(vals, start=1):
            for ci, v in enumerate(r):
                if isinstance(v, str) and v.startswith("="):
                    out["formulas"].append((f"{_col(ci)}{ri}", v[:1500]))
    except Exception as exc:  # noqa: BLE001
        out["formula_error"] = _err(exc)
    time.sleep(ontime_sync.PAUSE_S)
    return out


def _grid(ws, tab: dict) -> list[list]:
    """A summary tab read whole (unformatted, dates as serials), bounded."""
    rows, cols = tab["rows"] or 1, tab["cols"] or 1
    if rows * cols > GRID_CELLS:
        rows = max(1, GRID_CELLS // cols)
    return ontime_sync._call(ws.get, f"A1:{_col(cols - 1)}{rows}",
                             value_render_option="UNFORMATTED_VALUE",
                             date_time_render_option="SERIAL_NUMBER") or []


def _headers(row1: list, wanted: dict) -> dict:
    """`ontime_sync._headers`'s rule over a row already read: field → column."""
    folded = [_norm(h) for h in row1]
    out = {}
    for field, names in wanted.items():
        for i, h in enumerate(folded):
            if h and any(h == n or h.startswith(n) for n in names):
                out[field] = (_col(i), str(row1[i]))
                break
    return out


# ── 3. the reader's own pass ─────────────────────────────────────────────────

def _read_raw(sh, rep: dict) -> Optional[dict]:
    notes: dict = {}
    data: dict = {}
    stage = "заход"
    try:
        t0 = time.monotonic()
        ws = ontime_sync._find_tab(sh, ontime_sync.TAB_WAVES)
        data["windows"] = ontime_sync.read_windows(ws, notes)
        rep["timing"]["windows_s"] = round(time.monotonic() - t0, 1)
        stage = ontime_sync.TAB_ORDERS
        t0 = time.monotonic()
        demand, mats, order_rows, first, last = ontime_sync.read_orders(
            ontime_sync._find_tab(sh, ontime_sync.TAB_ORDERS), notes)
        rep["timing"]["orders_s"] = round(time.monotonic() - t0, 1)
        stage = ontime_sync.TAB_RECEIPTS
        t0 = time.monotonic()
        rec, receipt_rows = ontime_sync.read_receipts(
            ontime_sync._find_tab(sh, ontime_sync.TAB_RECEIPTS), notes, mats)
        rep["timing"]["receipts_s"] = round(time.monotonic() - t0, 1)
    except Exception as exc:  # noqa: BLE001
        rep["raw_error"] = {"stage": stage, **_err(exc)}
        rep["raw_notes"] = notes
        return None
    data.update(demand=demand, mats=mats, rec=rec)
    rep["raw"] = {"order_rows": order_rows, "receipt_rows": receipt_rows,
                  "first": first, "last": last, "demand_keys": len(demand),
                  "receipt_keys": len(rec), "materials": len(mats),
                  "windows": len(data["windows"])}
    rep["raw_notes"] = notes
    return data


def _windows_by_grp(windows: list[dict]) -> dict:
    out: dict = {}
    for w in sorted(windows, key=lambda x: x["pos"]):
        g = out.setdefault(ontime._fold(w["grp"]), {"name": w["grp"], "waves": {}, "pos": w["pos"]})
        g["waves"][w["wave"]] = (w["start_s"], w["end_s"])
    return out


class _Model:
    """The page's evaluation (`ontime._rows`), over the data read just now."""

    def __init__(self, data: dict, days: set):
        self.wins = _windows_by_grp(data["windows"])
        self.mats = data["mats"]
        need = set(days) | {d + timedelta(days=1) for d in days}
        raw: dict = defaultdict(list)
        for (d, s, m), q in data["rec"].items():
            if d in need:
                raw[(d, m)].append((s, q))
        self.posts = {k: ontime._Posts(v) for k, v in raw.items()}
        self.demand: dict = {}
        self.by_day: dict = defaultdict(set)
        for (d, m, w), a in data["demand"].items():
            if d not in days:
                continue
            self.by_day[d].add(m)
            e = self.demand.setdefault((d, m), {"grp": a[2] or "", "waves": {}})
            e["waves"][w] = e["waves"].get(w, 0.0) + a[0]
            if not e["grp"] and a[2]:
                e["grp"] = a[2]
        self._ev: dict = {}

    def grp_of(self, mat: str, day: date) -> str:
        e = self.demand.get((day, mat))
        return (e["grp"] if e else "") or self.mats.get(mat, ("", ""))[1]

    def ev(self, day: date, mat: str) -> tuple[str, dict, dict]:
        """(group, its windows, {wave: row}) — what the page computes for one
        material on one order day."""
        k = (day, mat)
        if k in self._ev:
            return self._ev[k]
        grp = self.grp_of(mat, day)
        g = self.wins.get(ontime._fold(grp))
        wins = g["waves"] if g else {}
        e = self.demand.get(k)
        rows = ontime.evaluate(e["waves"] if e else {}, self.posts.get(k),
                               self.posts.get((day + timedelta(days=1), mat)), wins) if wins else {}
        out = (g["name"] if g else grp, wins, rows)
        self._ev[k] = out
        return out

    def day_rows(self, day: date):
        """Every row the page counts on that day: (material, group, wave, row)."""
        for m in sorted(self.by_day.get(day, ())):
            grp, wins, rows = self.ev(day, m)
            for w, r in rows.items():
                if r["D"] <= 0 and r["P"] <= 0:
                    continue
                yield m, grp, w, r


# ── 4. against the sheet's own summaries ─────────────────────────────────────

def _dated_cols(row1: list, start: int) -> list[tuple[int, date]]:
    """The first column of every day block in row 1. A block repeats its date
    across its columns, so only a column whose left neighbour carries another
    value opens one."""
    out = []
    prev = None
    for i, v in enumerate(row1):
        d = ontime_sync.parse_day(v) if i >= start and v not in (None, "") else None
        if d is not None and d != prev:
            out.append((i, d))
        prev = d
    return out


def _at(r: list, i: int):
    return r[i] if i < len(r) else None


def _compare_svod(grid: list[list], model: _Model, dated: list[tuple[int, date]]) -> dict:
    """«свод»: A material · B name · C group · D wave · E start · F end, then a
    block per day from G — D ordered · H counted · I on time · (row 3) the day's
    weighted mean."""
    rows = []
    for ri, r in enumerate(grid):
        if ri < 3:
            continue
        mat = ontime_sync._text(_at(r, 0))
        w = ontime_sync.parse_wave(_at(r, 3))
        if mat and w is not None:
            rows.append((ri + 1, r, mat, w))
    per_day, diffs = [], []
    win_diff, grp_diff = [], []
    for rownum, r, mat, w in rows:
        s_grp = ontime_sync._text(_at(r, 2))
        o_grp = model.mats.get(mat, ("", ""))[1]
        if o_grp and ontime._fold(o_grp) != ontime._fold(s_grp):
            grp_diff.append({"row": rownum, "material": mat, "sheet_grp": s_grp, "our_grp": o_grp})
        g = model.wins.get(ontime._fold(s_grp))
        o = (g["waves"] if g else {}).get(w)
        s_s, s_e = ontime_sync.parse_secs(_at(r, 4)), ontime_sync.parse_secs(_at(r, 5))
        if o is None or s_s != o[0] or s_e != o[1]:
            win_diff.append({"row": rownum, "material": mat, "sheet_grp": s_grp, "wave": w,
                             "sheet": (_hms(s_s), _hms(s_e)),
                             "ours": (_hms(o[0]), _hms(o[1])) if o else ("—", "—")})
    for col, day in dated:
        n = same = 0
        bad = {"D": 0, "H": 0, "c": 0}
        sw_n = sw_d = ow_n = ow_d = 0.0
        for rownum, r, mat, w in rows:
            grp, wins, ev = model.ev(day, mat)
            mine = ev.get(w)
            sD, sH, sI = _num(_at(r, col)), _num(_at(r, col + 1)), _num(_at(r, col + 2))
            if not sD and not sH and sI is None and not (mine and (mine["D"] > 0 or mine["P"] > 0)):
                continue                      # nothing on either side for this row today
            n += 1
            oD = mine["D"] if mine else 0.0
            oH = mine["H"] if mine else 0.0
            oc = mine["c"] if mine else None
            what = []
            if abs((sD or 0.0) - oD) > EPS:
                what.append("D")
            if abs((sH or 0.0) - oH) > EPS:
                what.append("H")
            if (sI is None) != (oc is None) or (oc is not None and abs(oc - sI) > 1e-9):
                what.append("c")
            for k in what:
                bad[k] += 1
            if not what:
                same += 1
            elif len(diffs) < DIFF_ROWS:
                diffs.append({"day": day, "row": rownum, "material": mat,
                              "name": ontime_sync._text(_at(r, 1)), "grp": grp, "wave": w,
                              "what": "+".join(what), "sD": sD, "sH": sH, "sI": sI,
                              "sI_raw": _at(r, col + 2), "oD": oD,
                              "oP": mine["P"] if mine else None, "oH": oH, "oc": oc})
            if sD and sI is not None:
                sw_n += sI * sD
                sw_d += sD
            if oD > 0 and oc is not None:
                ow_n += oc * oD
                ow_d += oD
        pg_n = pg_d = 0.0
        for _m, _g, _w, rr in model.day_rows(day):
            pg_d += rr["D"]
            pg_n += rr["covered"]
        per_day.append({"day": day, "col": _col(col), "rows": n, "same": same, **bad,
                        "row3": _num(_at(grid[2], col + 3)) if len(grid) > 2 else None,
                        "sheet_w": sw_n / sw_d if sw_d else None,
                        "ours_svod": ow_n / ow_d if ow_d else None,
                        "ours_page": pg_n / pg_d if pg_d else None,
                        "ours_rows": sum(1 for _ in model.day_rows(day))})
    return {"rows": len(rows), "header_days": len(dated),
            "days": [d for d in per_day if d["rows"]], "diffs": diffs,
            "win_diff": win_diff, "grp_diff": grp_diff}


def _compare_groups(grid: list[list], model: _Model, dated: list[tuple[int, date]],
                    svod_mats: dict) -> dict:
    """«свод по группам»: A group · B wave, a block per day — +2 the group-wave
    on-time mean, +3 the group's daily mean (the mean of its waves)."""
    rows = []
    for ri, r in enumerate(grid):
        grp = ontime_sync._text(_at(r, 0))
        w = ontime_sync.parse_wave(_at(r, 1))
        if grp and w is not None and ri >= 1:
            rows.append((ri + 1, r, grp, w))
    out, n_same, n = [], 0, 0
    for col, day in dated:
        page_cs: dict = defaultdict(list)
        svod_cs: dict = defaultdict(list)
        for m, g, w, rr in model.day_rows(day):
            if rr["c"] is not None:
                page_cs[(ontime._fold(g), w)].append(rr["c"])
                if (m, w) in svod_mats:
                    svod_cs[(ontime._fold(g), w)].append(rr["c"])
        means = {k: sum(v) / len(v) for k, v in page_cs.items()}
        grp_waves: dict = defaultdict(list)
        for (g, _w), v in means.items():
            grp_waves[g].append(v)
        for rownum, r, grp, w in rows:
            k = (ontime._fold(grp), w)
            s_gw, s_g = _num(_at(r, col + 2)), _num(_at(r, col + 3))
            o_gw = means.get(k)
            o_svod = (sum(svod_cs[k]) / len(svod_cs[k])) if svod_cs.get(k) else None
            gv = grp_waves.get(ontime._fold(grp))
            o_g = sum(gv) / len(gv) if gv else None
            if s_gw is None and o_gw is None:
                continue                      # neither side has a figure for this cell
            ok_gw = (s_gw is None) == (o_gw is None) and (o_gw is None or abs(o_gw - s_gw) <= 1e-9)
            ok_g = (s_g is None) == (o_g is None) and (o_g is None or abs(o_g - s_g) <= 1e-9)
            n += 1
            n_same += 1 if ok_gw and ok_g else 0
            out.append({"day": day, "row": rownum, "grp": grp, "wave": w,
                        "s_gw": s_gw, "s_gw_raw": _at(r, col + 2), "o_gw": o_gw, "o_svod": o_svod,
                        "n": len(page_cs.get(k, ())), "n_svod": len(svod_cs.get(k, ())),
                        "s_g": s_g, "o_g": o_g, "ok": ok_gw and ok_g})
    return {"rows": len(rows), "cells": n, "same": n_same, "list": out}


def _our_group_days(model: _Model, days: list[date]) -> list[dict]:
    out = []
    for day in days:
        cs: dict = defaultdict(list)
        for _m, g, w, rr in model.day_rows(day):
            if rr["c"] is not None:
                cs[(g, w)].append(rr["c"])
        by_g: dict = defaultdict(dict)
        for (g, w), v in cs.items():
            by_g[g][w] = sum(v) / len(v)
        for g, ws in sorted(by_g.items()):
            out.append({"day": day, "grp": g, "waves": ws, "mean": sum(ws.values()) / len(ws)})
    return out


# ── the report ───────────────────────────────────────────────────────────────

def collect(db: Session) -> dict:
    from app.services.sheets_reader import _reset_client, get_client, get_service_account_email
    rep: dict = {"at": datetime.now(TZ), "sa": get_service_account_email(),
                 "creds_file": settings.google_credentials_file, "timing": {}}
    rep["configured"] = ontime_sync.sheet_id(db)
    rep["state"] = ontime_sync.read_state(db)
    db.rollback()            # nothing more is read from the database; give the connection back

    t_all = time.monotonic()
    try:
        client = get_client()
    except Exception as exc:  # noqa: BLE001
        rep["client_error"] = _err(exc)
        return rep
    rep["checks"] = []
    sh_orig = None
    for label, sid in (("original", ORIGINAL), ("configured", rep["configured"]), ("copy", COPY)):
        if not sid or (label != "original" and sid == ORIGINAL) or \
                (label == "copy" and sid == rep["configured"]):
            continue
        sh, res = _open(client, sid)
        res.update(label=label, id=sid)
        rep["checks"].append(res)
        if label == "original":
            sh_orig = sh
        if sh is None:
            _reset_client()
            client = get_client()
    if sh_orig is None:
        rep["timing"]["total_s"] = round(time.monotonic() - t_all, 1)
        return rep

    meta = rep["checks"][0]["meta"]
    tabs = _tabs(meta)
    rep["tabs"] = tabs
    rep["samples"] = {}
    ws_by_id = {ws.id: ws for ws in ontime_sync._call(sh_orig.worksheets)}
    for t in tabs[:MAX_TABS]:
        if t["gid"] in ws_by_id:
            rep["samples"][t["title"]] = _sample(ws_by_id[t["gid"]], t)
    rep["timing"]["samples_s"] = round(time.monotonic() - t_all, 1)

    for name, wanted in ((ontime_sync.TAB_ORDERS, ontime_sync.ORDER_COLS),
                         (ontime_sync.TAB_RECEIPTS, ontime_sync.RECEIPT_COLS)):
        tab = next((t for t in tabs if _norm(t["title"]) == _norm(name)), None)
        s = rep["samples"].get(tab["title"]) if tab else None
        row1 = (s["rows"][0] if s and s["rows"] else [])
        rep.setdefault("columns", {})[name] = {
            "tab": tab["title"] if tab else None, "row1": row1, "found": _headers(row1, wanted),
            "needed": [f for f in wanted if f not in _headers(row1, wanted)]}

    data = _read_raw(sh_orig, rep)

    grids: dict = {}
    for name in (SVOD, SVOD_GROUPS):
        tab = next((t for t in tabs if _norm(t["title"]) == _norm(name)), None)
        if tab is None:
            continue
        try:
            grids[name] = _grid(ws_by_id[tab["gid"]], tab)
        except Exception as exc:  # noqa: BLE001
            rep.setdefault("grid_errors", {})[name] = _err(exc)
        time.sleep(ontime_sync.PAUSE_S)
    rep["grid_sizes"] = {k: (len(v), max((len(r) for r in v), default=0)) for k, v in grids.items()}

    if data is not None:
        demand = data["demand"]
        per_day: dict = defaultdict(lambda: {"lines": 0, "qty": 0.0, "mats": set(), "rec_rows": 0,
                                             "rec_qty": 0.0, "waves": defaultdict(float)})
        for (d, m, w), a in demand.items():
            x = per_day[d]
            x["qty"] += a[0]
            x["lines"] += a[1]
            x["mats"].add(m)
            x["waves"][w] += a[0]
        for (d, _s, _m), q in data["rec"].items():
            per_day[d]["rec_rows"] += 1
            per_day[d]["rec_qty"] += q
        rep["per_day"] = {d: {**v, "mats": len(v["mats"]), "waves": dict(v["waves"])}
                          for d, v in sorted(per_day.items())}
        wins = _windows_by_grp(data["windows"])
        rep["windows"] = data["windows"]
        rep["overlap"] = ontime.overlapping(wins)
        nw: dict = defaultdict(lambda: [0.0, set()])
        for (d, m, w), a in demand.items():
            grp = a[2] or data["mats"].get(m, ("", ""))[1]
            g = wins.get(ontime._fold(grp))
            if not g or w not in g["waves"]:
                nw[grp or "—"][0] += a[0]
                nw[grp or "—"][1].add(m)
        rep["no_window"] = sorted(((g, q, len(ms)) for g, (q, ms) in nw.items()), key=lambda x: -x[1])

        svod_dated = _dated_cols(grids[SVOD][0], 6) if grids.get(SVOD) else []
        grp_dated = _dated_cols(grids[SVOD_GROUPS][0], 2) if grids.get(SVOD_GROUPS) else []
        order_days = sorted({d for (d, _m, _w) in demand})
        recent = order_days[-RECENT_DAYS:]
        days = {d for _c, d in svod_dated} | {d for _c, d in grp_dated} | set(recent)
        model = _Model(data, days)
        if grids.get(SVOD):
            rep["svod"] = _compare_svod(grids[SVOD], model, svod_dated)
            svod_mats = {(ontime_sync._text(_at(r, 0)), ontime_sync.parse_wave(_at(r, 3)))
                         for r in grids[SVOD][3:]}
        else:
            svod_mats = set()
        if grids.get(SVOD_GROUPS):
            rep["groups"] = _compare_groups(grids[SVOD_GROUPS], model, grp_dated, svod_mats)
        rep["our_group_days"] = _our_group_days(model, sorted(days))
        rep["svod_head"] = [list(r) for r in (grids.get(SVOD) or [])[:3]]
    rep["grids"] = {k: v for k, v in grids.items() if k == SVOD_GROUPS}
    rep["timing"]["total_s"] = round(time.monotonic() - t_all, 1)
    rep["timing"]["peak_rss_mb"] = _peak_mb()
    return rep


def _peak_mb() -> Optional[int]:
    try:
        import resource
        return int(resource.getrusage(resource.RUSAGE_SELF).ru_maxrss / 1024)
    except Exception:  # noqa: BLE001
        return None


def text(rep: dict) -> str:
    L = [f"On-time sheet — production's own check · {rep['at']:%d.%m.%Y %H:%M}",
         f"Service account: {rep.get('sa') or '— (could not be read)'}"]
    if rep.get("client_error"):
        e = rep["client_error"]
        L.append(f"The Google client could not start: {e['code']} {e['detail']}")
        return "\n".join(L)
    st = rep.get("state") or {}
    last = st.get("last") or {}
    L.append(f"Set on «Ma'lumotlar» (ontime): {rep.get('configured') or 'nothing'}"
             + (f" · last read {last.get('finished_at') or last.get('started_at')}, "
                f"{'ok' if last.get('ok') else 'failed: ' + str(last.get('error'))}" if last else ""))
    L.append("")
    L.append("Can the service account open:")
    for c in rep.get("checks", []):
        if c["ok"]:
            L.append(f" • {c['label']} {c['id']}: YES — «{c['title']}», {c['tabs']} tabs, "
                     f"locale {c['locale']}, time zone {c['tz']}")
        else:
            L.append(f" • {c['label']} {c['id']}: NO — {c.get('status') or ''} {c['code']} "
                     f"{c['detail'][:160]}")
    orig = (rep.get("checks") or [{}])[0]
    if not orig.get("ok"):
        L.append("")
        L.append(f"The original cannot be read. Share it (Viewer) with {rep.get('sa')} — the "
                 "address above — and ask for this report again.")
        return "\n".join(L)

    L.append("")
    L.append("Tabs of the original:")
    for t in rep.get("tabs", []):
        L.append(f" • {t['title']} — {_g(t['rows'])} × {_g(t['cols'])}"
                 + (" (hidden)" if t["hidden"] else ""))
    L.append("")
    L.append("Raw tabs, as the page's reader sees them:")
    for name, c in (rep.get("columns") or {}).items():
        if not c["tab"]:
            L.append(f" • «{name}»: NO SUCH TAB")
            continue
        found = ", ".join(f"{f}={col} «{h}»" for f, (col, h) in c["found"].items())
        L.append(f" • «{c['tab']}»: {found}")
        if c["needed"]:
            L.append(f"   not found: {', '.join(c['needed'])}")
    if rep.get("raw_error"):
        e = rep["raw_error"]
        L.append(f" ✗ the read stopped at «{e['stage']}»: {e['code']} {e['detail']}")
    raw = rep.get("raw")
    if raw:
        L.append(f" • orders: {_g(raw['order_rows'])} rows, days {_d(raw['first'])} … {_d(raw['last'])}"
                 f" · receipts: {_g(raw['receipt_rows'])} rows · {raw['materials']} materials · "
                 f"{raw['windows']} wave windows")
    notes = rep.get("raw_notes") or {}
    for k in ("orders_skipped", "receipts_skipped", "windows_duplicate"):
        if notes.get(k):
            L.append(f"   {k}: {notes[k]}")
    if rep.get("overlap"):
        L.append(f"   windows 1 and 2 overlap: {', '.join(rep['overlap'])}")
    if rep.get("no_window"):
        L.append("   ordered with no window (counted nowhere): "
                 + "; ".join(f"{g} {_g(q)} ({n} materials)" for g, q, n in rep["no_window"][:8]))
    t = rep.get("timing", {})
    L.append(f"   read times: orders {t.get('orders_s')} s, receipts {t.get('receipts_s')} s"
             f" · process peak memory {t.get('peak_rss_mb') or _peak_mb()} MB")

    sv = rep.get("svod")
    if sv is not None:
        L.append("")
        L.append(f"Page arithmetic vs the sheet's «свод» ({sv['rows']} material × wave rows, "
                 f"{sv['header_days']} days in its header, {len(sv['days'])} with figures):")
        if not sv["days"]:
            L.append(" • no day with figures")
        for d in sv["days"][-20:]:
            L.append(f" • {_d(d['day'])} ({d['col']}): {d['same']}/{d['rows']} equal"
                     + (f" · differ D {d['D']}, counted {d['H']}, on-time {d['c']}"
                        if d["same"] != d["rows"] else "")
                     + f" · weighted: sheet row 3 {_pct(d['row3'])}, sheet Σ {_pct(d['sheet_w'])}, "
                       f"ours on the same rows {_pct(d['ours_svod'])}, page {_pct(d['ours_page'])}")
        if len(sv["days"]) > 20:
            L.append(f"   (+{len(sv['days']) - 20} earlier days in the workbook)")
        if sv["win_diff"]:
            L.append(f" • {len(sv['win_diff'])} rows carry another window than «заход» gives their group")
        if sv["grp_diff"]:
            L.append(f" • {len(sv['grp_diff'])} rows name another group than «заявки» does")
    gp = rep.get("groups")
    if gp is not None:
        L.append("")
        L.append(f"«свод по группам»: {gp['same']}/{gp['cells']} group × wave × day cells equal "
                 f"({gp['rows']} rows)")
    for k, v in (rep.get("grid_errors") or {}).items():
        L.append(f" ✗ «{k}» could not be read: {v['code']} {v['detail']}")
    L.append("")
    L.append(f"Took {t.get('total_s')} s. Every tab's first rows, the formulas on the summary "
             "tabs, the windows, the day totals, every differing row and our group × day means "
             "are in the attached workbook.")
    return "\n".join(L)


# ── the workbook ──────────────────────────────────────────────────────────────

def _cell(v):
    if isinstance(v, str):
        return ILLEGAL_CHARACTERS_RE.sub("", v)[:5000]
    if isinstance(v, datetime):
        return v.astimezone(TZ).replace(tzinfo=None) if v.tzinfo else v
    if isinstance(v, (list, dict, tuple, set)):
        return ILLEGAL_CHARACTERS_RE.sub("", str(v))[:2000]
    return v


def _sheet(wb: Workbook, title: str, head: list, rows: list) -> None:
    ws = wb.create_sheet(title[:31])
    if head:
        ws.append([_cell(h) for h in head])
        for c in ws[1]:
            c.font = Font(bold=True)
        ws.freeze_panes = "A2"
    for r in rows:
        ws.append([_cell(v) for v in r])
    for i in range(1, min(ws.max_column, 80) + 1):
        letter = ws.cell(1, i).column_letter
        width = max([len(str(_cell(r[i - 1]))) for r in rows[:200]
                     if i - 1 < len(r) and r[i - 1] is not None] + [len(str(head[i - 1])) if head and i - 1 < len(head) else 6])
        ws.column_dimensions[letter].width = min(max(width + 2, 8), 50)


def build_workbook(rep: dict) -> BytesIO:
    wb = Workbook()
    wb.remove(wb.active)
    _sheet(wb, "Umumiy", ["", ""], [[ln] for ln in text(rep).split("\n")])
    _sheet(wb, "Tekshiruv", ["Qaysi", "ID", "Ochildimi", "Nomi", "Til", "Vaqt zonasi", "Varaqlar",
                             "Xato", "Tafsilot", "ms"],
           [[c["label"], c["id"], "ha" if c["ok"] else "yo'q", c.get("title"), c.get("locale"),
             c.get("tz"), c.get("tabs"), c.get("code"), c.get("detail"), c.get("ms")]
            for c in rep.get("checks", [])])
    if rep.get("tabs"):
        _sheet(wb, "Varaqlar", ["Varaq", "gid", "Tartib", "Qatorlar", "Ustunlar", "Qotirilgan",
                                "Yashirin", "Formulalar (yuqorida)"],
               [[t["title"], t["gid"], t["index"], t["rows"], t["cols"], t["frozen"],
                 "ha" if t["hidden"] else "", len((rep["samples"].get(t["title"]) or {}).get("formulas", []))]
                for t in rep["tabs"]])
        for title, s in (rep.get("samples") or {}).items():
            rows = [list(r) for r in s.get("rows") or []]
            if s.get("error"):
                rows.append([f"error: {s['error']}"])
            _sheet(wb, f"N {title}"[:31], [], rows)
        fl = []
        for title, s in (rep.get("samples") or {}).items():
            for cell, f in s.get("formulas", []):
                fl.append([title, cell, f])
        _sheet(wb, "Formulalar", ["Varaq", "Katak", "Formula"], fl)
    cols = rep.get("columns") or {}
    if cols:
        cr = []
        for name, c in cols.items():
            for f, (col, h) in c["found"].items():
                cr.append([name, f, col, h])
            for f in c["needed"]:
                cr.append([name, f, "TOPILMADI", ""])
        _sheet(wb, "Ustunlar", ["Varaq", "Maydon", "Ustun", "Sarlavha"], cr)
    if rep.get("windows"):
        ov = set(rep.get("overlap") or [])
        _sheet(wb, "Oynalar", ["Guruh", "Zaxod", "Boshlanish", "Tugash", "Qator (F:I)", "1 va 2 ustma-ust"],
               [[w["grp"], w["wave"], _hms(w["start_s"]), _hms(w["end_s"]), w["pos"],
                 "ha" if w["grp"] in ov else ""] for w in rep["windows"]])
    if rep.get("no_window"):
        _sheet(wb, "Oynasiz", ["Guruh", "Buyurtma miqdori", "Materiallar"],
               [list(x) for x in rep["no_window"]])
    if rep.get("per_day"):
        _sheet(wb, "Kunlar", ["Sana", "Buyurtma qatorlari", "Buyurtma miqdori", "Materiallar",
                              "1-zaxod", "2-zaxod", "3-zaxod", "Kirim qatorlari (sek × material)",
                              "Kirim miqdori"],
               [[d, v["lines"], v["qty"], v["mats"], v["waves"].get(1), v["waves"].get(2),
                 v["waves"].get(3), v["rec_rows"], v["rec_qty"]] for d, v in rep["per_day"].items()])
    sv = rep.get("svod")
    if sv is not None:
        _sheet(wb, "Svod kunlar", ["Sana", "Ustun", "Qatorlar", "Teng", "D farq", "H farq",
                                   "On time farq", "Varaq 3-qator", "Varaq Σ (D×I)",
                                   "Biz (shu qatorlar)", "Sahifa (hammasi)", "Sahifa qatorlari"],
               [[d["day"], d["col"], d["rows"], d["same"], d["D"], d["H"], d["c"], d["row3"],
                 d["sheet_w"], d["ours_svod"], d["ours_page"], d["ours_rows"]] for d in sv["days"]])
        _sheet(wb, "Svod farqlar", ["Sana", "Qator", "Material", "Nomi", "Guruh", "Zaxod", "Farq",
                                    "Varaq D", "Varaq H", "Varaq on time", "Varaq on time (xom)",
                                    "Biz D", "Biz P", "Biz H", "Biz on time"],
               [[x["day"], x["row"], x["material"], x["name"], x["grp"], x["wave"], x["what"],
                 x["sD"], x["sH"], x["sI"], x["sI_raw"], x["oD"], x["oP"], x["oH"], x["oc"]]
                for x in sv["diffs"]])
        if sv["win_diff"]:
            _sheet(wb, "Svod oynalar", ["Qator", "Material", "Guruh", "Zaxod",
                                        "Svod oynasi", "«заход» oynasi"],
                   [[x["row"], x["material"], x["sheet_grp"], x["wave"],
                     " – ".join(x["sheet"]), " – ".join(x["ours"])] for x in sv["win_diff"]])
        if sv["grp_diff"]:
            _sheet(wb, "Svod guruhlar", ["Qator", "Material", "Svod guruhi", "«заявки» guruhi"],
                   [[x["row"], x["material"], x["sheet_grp"], x["our_grp"]] for x in sv["grp_diff"]])
    if rep.get("svod_head"):
        _sheet(wb, "Svod sarlavha", [], rep["svod_head"])
    gp = rep.get("groups")
    if gp is not None:
        _sheet(wb, "Guruhlar solishtirish", ["Sana", "Qator", "Guruh", "Zaxod", "Varaq on time",
                                             "Varaq (xom)", "Sahifa", "Biz (svod qatorlari)",
                                             "Mahsulotlar", "Svod mahsulotlari", "Varaq guruh o'rtacha",
                                             "Biz guruh o'rtacha", "Teng"],
               [[x["day"], x["row"], x["grp"], x["wave"], x["s_gw"], x["s_gw_raw"], x["o_gw"],
                 x["o_svod"], x["n"], x["n_svod"], x["s_g"], x["o_g"], "ha" if x["ok"] else "YO'Q"]
                for x in gp["list"]])
    if rep.get("our_group_days"):
        _sheet(wb, "Biz guruh × kun", ["Sana", "Guruh", "1-zaxod", "2-zaxod", "3-zaxod", "O'rtacha"],
               [[x["day"], x["grp"], x["waves"].get(1), x["waves"].get(2), x["waves"].get(3), x["mean"]]
                for x in rep["our_group_days"]])
    for name, g in (rep.get("grids") or {}).items():
        _sheet(wb, f"Varaq {name}"[:31], [], [list(r) for r in g])
    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


# ── delivery ──────────────────────────────────────────────────────────────────

def _post(method: str, data: dict, files: dict | None = None) -> None:
    last = ""
    for attempt in range(1, SEND_RETRIES + 1):
        wait = 0
        try:
            r = requests.post(f"{_API}/bot{settings.telegram_bot_token}/{method}",
                              data=data, files=files, timeout=180)
            body = r.json()
            if body.get("ok"):
                return
            last = body.get("description") or f"HTTP {r.status_code}"
            wait = int(((body.get("parameters") or {}).get("retry_after")) or 0)
        except Exception as exc:
            last = type(exc).__name__
        if attempt < SEND_RETRIES:
            time.sleep(max(wait, 3 * attempt))
    raise RuntimeError(f"{method} failed: {last}"[:300])


def send(db: Session, chat_id: int, *_window) -> int:
    if not settings.telegram_bot_token:
        raise RuntimeError("telegram bot token not configured")
    rep = collect(db)
    body = text(rep)
    book = build_workbook(rep).getvalue()
    chunks, cur = [], ""
    for line in body.split("\n"):
        if len(cur) + len(line) + 1 > 3900:
            chunks.append(cur)
            cur = ""
        cur += line + "\n"
    chunks.append(cur)
    for c in chunks:
        _post("sendMessage", {"chat_id": chat_id, "text": c})
    stamp = datetime.now(TZ).strftime("%d-%m-%Y")
    _post("sendDocument", {"chat_id": chat_id,
                           "caption": "On-time sheet — the original, read by production"},
          files={"document": (f"ontime-original-sheet-{stamp}.xlsx", book,
                              "application/vnd.openxmlformats-officedocument."
                              "spreadsheetml.sheet")})
    return len(chunks) + 1
