"""TEMPORARY boot report (2026-10-26): where «Odam soni» on «Zagruzka fayli»
and the загрузка's own headcount disagree — and what that costs.

The operator, 25.10.2026, looking at Niyozov Izzatilla's 24.10: «Файл загрузки»
showed ЛЮДЕЙ (Σ) **45** (21 + 12 + 12) and «Ежедневно» showed ОФИЦ. ЧИСЛ. **66**
for the same unit-day. Both claim to read the TYPED «Bugungi fakt» pins. They do
not read the same set of them:

  • the PAGE counts a pin only where its `work_center` string matches, EXACTLY,
    a code the day's catalog puts on the panel — the active work-centre register
    ∪ the work centres active catalog lines name (`pp_calc` builds `wc_codes`
    from those two, and looks a pin up with `ov_all.get(code)`);
  • the LOAD counts EVERY `pp_work_center_daily` row of that (unit, date) with
    `people IS NOT NULL` — `zagruzka_source.typed_pins`, keyed on the raw string
    with no `norm_code` and no join to the register at all.

So a pin survives its work centre being deactivated, deleted from the register,
dropped from the catalog, or spelled another way (a Cyrillic twin — «A code is
LATIN» — a stray space, another case), and from then on it is counted in the
загрузка while being invisible, and uneditable, on the page it was typed on.
`cell_pins` already folds two spellings onto one work centre through
`wc_group.wc_key`; `unit_people`, the one the загрузка divides by, never got it.

This report answers the two things only production can: WHICH pin rows differ,
and over how many unit-days since `zagruzka_source.ZAGRUZKA_FROM` — with the
загрузка each affected day reads now and would read off the page's own Σ, so the
blast radius of a fix is a number before it is a decision.

READ-ONLY. It writes nothing but its flag (`startup._send_report_once`) and
derives every figure from the functions the two pages already read, so it can
be stale about a pin typed after it ran — never wrong about the moment it was
taken. DMed to `UNPRICED_DM_CHAT` as text plus a workbook.

Temporary: delete this module, `startup.report_pin_mismatch` /
`_pin_mismatch_job` and the call in BOTH entrypoints once it has landed.
Changing what it reports needs a NEW flag key.
"""
from __future__ import annotations

import time
from collections import defaultdict
from datetime import date, datetime, timedelta, timezone
from io import BytesIO
from typing import Optional
from zoneinfo import ZoneInfo

import requests
from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Font
from sqlalchemy.orm import Session

from app.config import settings
from app.models import Manager, PPWorkCenter
from app.services import pp_catalog, zagruzka_source
from app.services.cell_lookup import norm_code
from app.services.latin_code import latin_code

TZ = ZoneInfo("Asia/Tashkent")
_API = "https://api.telegram.org"
SEND_RETRIES = 3

# The загрузка divides by a FLAT 480 (`kpi_calculator`), deliberately not the
# production page's `pp_shift_min` — so the «what would it read» column below
# is the figure /zagruzka and /daily actually show, not the page's own.
SHIFT_MIN = 480.0

MAX_ROWS = 4000      # workbook detail cap — the text names what it cut
TOP_N = 20           # unit-days named in the DM text


# ── why a pin is not on the page ────────────────────────────────────────────
# Narrowest first, so the reason names the FIX rather than the symptom.
R_SPELLING = "spelling"       # same code, another case/whitespace
R_TWIN     = "cyrillic_twin"  # Cyrillic lookalike of a code the page shows
R_INACTIVE = "inactive"       # the register has the row, switched off / not in the day's catalog
R_UNKNOWN  = "unknown_code"   # no register row at all on this unit

REASON_RU = {
    R_SPELLING: "другое написание того же кода",
    R_TWIN:     "кириллический двойник кода",
    R_INACTIVE: "участок есть в реестре, но не на странице этого дня",
    R_UNKNOWN:  "такого участка у бригадира нет",
}


def _classify(raw: str, page_codes: set[str], register: dict[str, bool]) -> tuple[str, Optional[str]]:
    """(reason, the page code it should have landed on) for a pin the page
    does not show. `register` is {code: active} over ALL of the unit's rows."""
    n = norm_code(raw)
    by_norm = {norm_code(c): c for c in page_codes}
    if n in by_norm:
        return R_SPELLING, by_norm[n]
    lat = norm_code(latin_code(raw))
    if lat and lat != n and lat in by_norm:
        return R_TWIN, by_norm[lat]
    if raw in register or n in {norm_code(c) for c in register}:
        return R_INACTIVE, None
    return R_UNKNOWN, None


def collect(db: Session) -> dict:
    today = datetime.now(TZ).date()
    lo, hi = zagruzka_source.ZAGRUZKA_FROM, today
    units = (db.query(Manager)
             .filter(Manager.archived.is_(False))
             .order_by(Manager.name).all())
    by_id = {m.id: m for m in units}
    ids = [m.id for m in units]
    rep: dict = {
        "from": lo, "to": hi, "units": len(units),
        "days": [], "pins": [], "per_unit": {}, "truncated": 0,
        "unit_days_with_pins": 0,
    }
    if not ids:
        return rep

    # Every typed pin of the window, and the same pins folded the way BOTH the
    # page and `unit_people` fold them (group pins win over a whole-centre one).
    pins = zagruzka_source.typed_pins(db, ids, lo, hi)
    folded = zagruzka_source.fold_pins(pins)          # {(mid, day, wc): people}
    if not folded:
        return rep

    # The unit's whole work-centre register, active or not — what tells «the
    # row exists and is switched off» from «this unit never had that code».
    register: dict[int, dict[str, bool]] = defaultdict(dict)
    for r in db.query(PPWorkCenter.manager_id, PPWorkCenter.code, PPWorkCenter.active).filter(
            PPWorkCenter.manager_id.in_(ids)):
        if r.code:
            register[int(r.manager_id)][r.code] = bool(r.active)

    # The catalog AS EACH DAY HAD IT, in runs of days (`pp_catalog.spans`), so
    # the page's own code set is rebuilt per day without one read per day.
    spans = pp_catalog.spans(db, ids, lo, hi)

    per_day: dict[tuple[int, str], dict] = defaultdict(dict)
    for (mid, day, wc, grp), n in pins.items():
        per_day[(mid, day)][(wc, grp)] = n
    # Indexed per unit-day ONCE. This runs on the deferred boot thread, and a
    # scan of every folded pin per unit-day (≈ days × units × pins) is Python
    # holding the GIL for no reason — the cost `perf_watch` keeps reporting.
    fold_by_day: dict[tuple[int, str], dict[str, float]] = defaultdict(dict)
    day_total: dict[tuple[int, str], float] = defaultdict(float)
    for (mid, day, wc), n in folded.items():
        fold_by_day[(mid, day)][wc] = n
        day_total[(mid, day)] += n
    rep["unit_days_with_pins"] = len(day_total)

    bad: list[dict] = []
    for mid in ids:
        for cat, d1, d2 in (spans.get(mid) or []):
            # Exactly the two sources `pp_calc` builds its panel from.
            page_codes = ({w.code for w in cat.active_work_centers if w.code}
                          | {p.work_center for p in cat.active_lines if p.work_center})
            d = max(d1, lo)
            while d <= min(d2, hi):
                iso = d.isoformat()
                key = (mid, iso)
                d += timedelta(days=1)
                if key not in day_total:
                    continue
                load_n = day_total[key]
                page_n = sum(n for wc, n in fold_by_day[key].items()
                             if wc in page_codes)
                if abs(load_n - page_n) < 1e-9:
                    continue
                off = []
                for (wc, grp), n in sorted(per_day[key].items()):
                    if wc in page_codes:
                        continue
                    reason, onto = _classify(wc, page_codes, register.get(mid, {}))
                    off.append({"wc": wc, "group": grp, "people": n,
                                "reason": reason, "onto": onto})
                bad.append({"unit_id": mid, "day": iso, "page_n": page_n,
                            "load_n": load_n, "diff": load_n - page_n, "off": off})

    if not bad:
        rep["days"] = []
        return rep

    # Minutes only for the affected units and window — the загрузка's own
    # numerator, so «now» and «if fixed» are the figures the pages print.
    a_ids = sorted({b["unit_id"] for b in bad})
    a_lo = date.fromisoformat(min(b["day"] for b in bad))
    a_hi = date.fromisoformat(max(b["day"] for b in bad))
    labor = zagruzka_source.unit_labor(db, a_ids, a_lo, a_hi)

    for b in bad:
        m = by_id[b["unit_id"]]
        plan = (labor.get((b["unit_id"], b["day"])) or (0.0, 0.0))[0]
        b["unit"] = m.name
        b["shift"] = m.shift
        b["on"] = bool(getattr(m, "zagruzka_on", True))
        b["plan_min"] = plan
        b["load_now"] = (plan / (SHIFT_MIN * b["load_n"])) if (plan > 0 and b["load_n"] > 0) else None
        b["load_fixed"] = (plan / (SHIFT_MIN * b["page_n"])) if (plan > 0 and b["page_n"] > 0) else None
    bad.sort(key=lambda b: (-abs(b["diff"]), b["unit"], b["day"]))

    acc: dict[int, dict] = {}
    for b in bad:
        a = acc.setdefault(b["unit_id"], {
            "unit": b["unit"], "shift": b["shift"], "on": b["on"],
            "days": 0, "max_diff": 0.0, "codes": set(), "reasons": set(),
            "first": b["day"], "last": b["day"]})
        a["days"] += 1
        a["max_diff"] = max(a["max_diff"], abs(b["diff"]))
        a["first"] = min(a["first"], b["day"])
        a["last"] = max(a["last"], b["day"])
        for o in b["off"]:
            a["codes"].add(o["wc"])
            a["reasons"].add(o["reason"])

    rep["days"] = bad[:MAX_ROWS]
    rep["truncated"] = max(0, len(bad) - MAX_ROWS)
    rep["all_days"] = len(bad)
    rep["per_unit"] = acc
    rep["pins"] = [dict(o, unit=b["unit"], day=b["day"]) for b in bad for o in b["off"]]
    return rep


def text(rep: dict) -> str:
    L: list[str] = []
    A = L.append
    A("«Odam soni»: страница и загрузка считают разные пины")
    A(f"Период: {rep['from'].strftime('%d.%m.%Y')} — {rep['to'].strftime('%d.%m.%Y')} "
      f"(с ZAGRUZKA_FROM), бригадиров: {rep['units']}")
    A("")
    if not rep.get("days"):
        A("Расхождений НЕТ: на каждом бригадиро-дне «Файл загрузки» и «Ежедневно»")
        A("считают один и тот же набор пинов. Ничего чинить не нужно.")
        A(f"Проверено пар (бригадир × день) с введёнными людьми: {rep['unit_days_with_pins']}.")
        return "\n".join(L)

    n_days = rep.get("all_days", len(rep["days"]))
    on_days = [b for b in rep["days"] if b["on"]]
    A(f"РАСХОЖДЕНИЯ: {n_days} пар (бригадир × день) из {rep['unit_days_with_pins']} "
      f"с введёнными людьми; бригадиров затронуто: {len(rep['per_unit'])}.")
    A(f"Из них в загрузке учитываются: {len(on_days)} "
      f"(остальные — «Zagruzka hisoblanmaydi», на цифры не влияют).")
    A("")
    A("Что это значит: загрузка делит на число, которого нет на странице,")
    A("где это число вводят — значит бригадир не может его ни увидеть, ни исправить.")
    A("")
    A(f"Худшие {min(TOP_N, len(rep['days']))} по разнице людей:")
    for b in rep["days"][:TOP_N]:
        d = date.fromisoformat(b["day"]).strftime("%d.%m")
        now = f"{b['load_now']*100:.0f}%" if b["load_now"] is not None else "—"
        fix = f"{b['load_fixed']*100:.0f}%" if b["load_fixed"] is not None else "—"
        mark = "" if b["on"] else " (не в загрузке)"
        A(f"• {b['unit']} · {d}{mark}: страница {b['page_n']:g} / загрузка "
          f"{b['load_n']:g} (+{b['diff']:g}) · {now} → {fix}")
        for o in b["off"][:4]:
            g = f"·{o['group']}" if o["group"] else ""
            onto = f" → должен был попасть в «{o['onto']}»" if o["onto"] else ""
            A(f"    «{o['wc']}»{g} = {o['people']:g} — {REASON_RU[o['reason']]}{onto}")
    if rep["truncated"]:
        A(f"… и ещё {rep['truncated']} — все в файле.")
    A("")
    A("По бригадирам:")
    for a in sorted(rep["per_unit"].values(), key=lambda x: (-x["days"], x["unit"])):
        mark = "" if a["on"] else " (не в загрузке)"
        A(f"• {a['unit']} (смена {a['shift']}){mark}: {a['days']} дн., "
          f"макс. +{a['max_diff']:g} чел., коды: {', '.join(sorted(a['codes'])[:6])}")
    A("")
    A("«→» — что показала бы загрузка, если делить на сумму со страницы.")
    A("Ничего не изменено: это только отчёт.")
    return "\n".join(L)


def _cell(v):
    if isinstance(v, str):
        return ILLEGAL_CHARACTERS_RE.sub("", v)
    if isinstance(v, datetime):
        v = v if v.tzinfo else v.replace(tzinfo=timezone.utc)
        return v.astimezone(TZ).replace(tzinfo=None)
    return v


def _sheet(wb: Workbook, title: str, head: list[str], rows: list[list]) -> None:
    ws = wb.create_sheet(title)
    ws.append(head)
    for c in ws[1]:
        c.font = Font(bold=True)
    for r in rows:
        ws.append([_cell(v) for v in r])
    ws.freeze_panes = "A2"
    for i, h in enumerate(head, start=1):
        width = max([len(str(h))] + [len(str(r[i - 1])) for r in rows[:300]
                                     if i - 1 < len(r) and r[i - 1] is not None])
        ws.column_dimensions[ws.cell(1, i).column_letter].width = min(max(width + 2, 9), 60)


def build_workbook(rep: dict) -> BytesIO:
    wb = Workbook()
    wb.remove(wb.active)
    _sheet(wb, "Kunlar", [
        "Brigadir", "Smena", "Zagruzkada", "Sana", "Sahifa (Σ)", "Zagruzka (Σ)",
        "Farq", "Plan, daq", "Zagruzka hozir", "Sahifa bo'yicha"],
        [[b["unit"], b["shift"], "ha" if b["on"] else "yo'q",
          date.fromisoformat(b["day"]), b["page_n"], b["load_n"], b["diff"],
          round(b["plan_min"], 1),
          (round(b["load_now"], 4) if b["load_now"] is not None else None),
          (round(b["load_fixed"], 4) if b["load_fixed"] is not None else None)]
         for b in rep.get("days", [])])
    _sheet(wb, "Pinlar", [
        "Brigadir", "Sana", "Ish markazi (yozilgan)", "Normal shakl", "Guruh",
        "Odam", "Sabab", "Qaysi kodga tegishli"],
        [[p["unit"], date.fromisoformat(p["day"]), p["wc"], norm_code(p["wc"]),
          p["group"], p["people"], REASON_RU[p["reason"]], p["onto"]]
         for p in rep.get("pins", [])[:MAX_ROWS]])
    _sheet(wb, "Brigadirlar", [
        "Brigadir", "Smena", "Zagruzkada", "Kunlar", "Maks. farq",
        "Birinchi", "Oxirgi", "Kodlar", "Sabablar"],
        [[a["unit"], a["shift"], "ha" if a["on"] else "yo'q", a["days"],
          a["max_diff"], a["first"], a["last"],
          ", ".join(sorted(a["codes"])), ", ".join(sorted(a["reasons"]))]
         for a in sorted(rep.get("per_unit", {}).values(),
                         key=lambda x: (-x["days"], x["unit"]))])
    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


def _post(method: str, data: dict, files: dict | None = None) -> None:
    last = ""
    for attempt in range(1, SEND_RETRIES + 1):
        wait = 0
        try:
            r = requests.post(f"{_API}/bot{settings.telegram_bot_token}/{method}",
                              data=data, files=files, timeout=180)
            body = r.json()
            if body.get("ok"):
                return
            last = body.get("description") or f"HTTP {r.status_code}"
            wait = int(((body.get("parameters") or {}).get("retry_after")) or 0)
        except Exception as exc:
            last = type(exc).__name__
        if attempt < SEND_RETRIES:
            time.sleep(max(wait, 3 * attempt))
    raise RuntimeError(f"{method} failed: {last}"[:300])


def send(db: Session, chat_id: int, *_window) -> int:
    if not settings.telegram_bot_token:
        raise RuntimeError("telegram bot token not configured")
    rep = collect(db)
    body = text(rep)
    book = build_workbook(rep).getvalue() if rep.get("days") else None
    db.rollback()
    chunks, cur = [], ""
    for line in body.split("\n"):
        if len(cur) + len(line) + 1 > 3900:
            chunks.append(cur)
            cur = ""
        cur += line + "\n"
    chunks.append(cur)
    for c in chunks:
        _post("sendMessage", {"chat_id": chat_id, "text": c})
    if book is None:
        return len(chunks)
    stamp = datetime.now(TZ).strftime("%d.%m.%Y")
    _post("sendDocument",
          {"chat_id": chat_id,
           "caption": f"«Odam soni» — страница против загрузки, {stamp}"},
          files={"document": (f"odam-soni-farq-{stamp}.xlsx", book,
                              "application/vnd.openxmlformats-officedocument."
                              "spreadsheetml.sheet")})
    return len(chunks) + 1
