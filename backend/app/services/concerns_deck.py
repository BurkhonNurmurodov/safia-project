"""The weekly Concerns («Xavotirlar») report as a PowerPoint deck.

`/concerns` → «Haftalik hisobot», the twin of the weekly Ojidaniya deck and
built on the same look, the same primitives and the same no-overlap rule
(`services/ojidaniya_deck.py`, `services/deck_text.py`). What it adds is the
concern register's own questions: what came in, what got closed, what is still
open and past its deadline, which departments and which brigadirs carry it,
and what went up the chain.

The rules it is built on, each of them the operator's ruling (2026-09-23):

**A fixed weekly report, never a view of the page.** The window is
`services/report_week` (last Wednesday back to the Wednesday before, both
included), the scope is both plants, both shifts and every unit, and the
page's filters change nothing. The button therefore never fires without a
confirm that writes the scope out. Admin only.

**Old concerns stay on it until they are closed.** A concern filed three weeks
ago and still open is the thing a weekly report most needs to show, so the
rows are this week's filings PLUS every older concern still open at the end of
the week. «Yangi» counts filings (by `entry_date`, the page's own period
field); «Yopildi» counts closings inside the window whenever the concern was
filed (by `completion_date`, the day the page's flow chart closes a row on).

**State is taken at the END of the window**, never at the moment the button is
pressed. «Ochiq» is a concern not closed by the end of the last day; «muddati
o'tgan» is the page's own rule as the page reads it the morning after the week
(deadline day on or before the last day). So two admins pressing it on two
different days get the same file, and last week's figures are computed the same
way this week's are, which is what makes the comparison mean anything.

**Worker-filed concerns count everywhere**, rankings included — the operator's
call for this report, overriding for it the 4 Sep reading that the worker
register «cannot be used to judge leaders». A worker's NAME is never printed and
never reaches the model: the router does not even hand it over.

**Brigadirs are ranked by overdue concerns**, then by open ones, then by
filings. Every other figure is a column beside it.

**Concern texts and closing notes are quoted, never rewritten.** They are moved
to Latin script where they were typed in Cyrillic and otherwise reproduced as
written. `services/concerns_narrative.py` writes the commentary around them.

**Every figure is computed HERE from the rows the router hands in.** This
module never queries, exactly as `ojidaniya_deck` is handed the page's output.
Always Uzbek Latin, whoever presses it.
"""
from __future__ import annotations

import io
import logging
from collections import Counter, defaultdict
from datetime import date, timedelta

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION, XL_LEGEND_POSITION
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

from app.services import deck_text
from app.services import ojidaniya_deck as _od
# ONE look and ONE set of drawing primitives for every generated deck — the
# colours, the chrome, and above all `_text`, which is where the no-overflow
# rule lives. Never copy them into this file: two copies of `_text` is how one
# deck starts drawing text over text again.
from app.services.ojidaniya_deck import (
    BODY_FONT, BROWN, CREAM, CW, DARK, FAINT, GOLD, GOLD2, GREEN, GREENBG, H,
    HEAD_FONT, INK, INK2, INNER, LINE, M, MUTED, NO_AI, ONDARK, RED, REDBG, W, WHITE,
    _ai, _ai_list, _blank, _card, _chip, _chrome, _rect, _sentence, _text,
    day_full, day_label, num, pct, period_words, short_name, signed_pct,
)
from app.translit import transliterate, transliterate_text

log = logging.getLogger(__name__)


# ── what things are called, in the deck's one language ─────────────────────
# Departments in `utils/concernCategories.js` order, labelled with the
# `concerns.category.*` uz bundle and coloured with CATEGORY_COLOR — keep all
# three in step with the frontend, which is where they are maintained.
CATS: list[tuple[str, str, str]] = [
    ("ars",          "ARS",              "EF4444"),
    ("inventory",    "Inventar",         "22C55E"),
    ("warehouse",    "Ombor",            "3B82F6"),
    ("fridge",       "Muzlatkich",       "EAB308"),
    ("procurement",  "Xarid",            "F97316"),
    ("logistics",    "Logistika",        "A855F7"),
    ("it",           "IT",               "14B8A6"),
    ("washing",      "Yuvish",           "EC4899"),
    ("plan",         "Reja",             "6366F1"),
    ("hr",           "HR",               "84CC16"),
    ("technologist", "Texnolog",         "06B6D4"),
    ("raw_material", "Keles (xomashyo)", "D946EF"),
    ("security",     "Qo'riqlash",       "0EA5E9"),
    ("kitchen",      "Oshxona",          "B45309"),
    ("other",        "Boshqa bo'lim",    "94A3B8"),
]
_CAT = {k: (label, RGBColor.from_string(hx)) for k, label, hx in CATS}
_CAT_ORDER = {k: i for i, (k, _l, _h) in enumerate(CATS)}
SLATE = RGBColor(0x94, 0xA3, 0xB8)

# The escalation chain, bottom → top, with the `concerns.level.*` uz labels and
# the identity hues `LEVEL_COLOR` gives each step on the page.
LEVELS = ["leader", "supervisor", "shift-manager", "top-manager"]
LEVEL_IDX = {lv: i for i, lv in enumerate(LEVELS)}
LEVEL_LABEL = {"leader": "Lider", "supervisor": "Brigadir",
               "shift-manager": "Smena menejeri", "top-manager": "Top-menejment"}
LEVEL_COLOUR = {"leader": RGBColor(0x94, 0xA3, 0xB8), "supervisor": RGBColor(0xEF, 0x44, 0x44),
                "shift-manager": RGBColor(0x22, 0xC5, 0x5E), "top-manager": RGBColor(0x3B, 0x82, 0xF6)}

SOFT_RED = RGBColor(0xF8, 0x71, 0x71)     # a red that reads on the dark cover
TILE = RGBColor(0x3A, 0x28, 0x1C)         # a panel on the dark cover
CLOSED_C = RGBColor(0x22, 0xC5, 0x5E)     # the page's «done» green

# A concern this short says nothing a reader can act on. Counted for the
# appendix rather than dropped, the rule the Ojidaniya deck keeps for notes.
THIN_CHARS = 12


def cat_meta(key: str | None) -> tuple[str, RGBColor]:
    """(label, colour) for a stored department key."""
    if key in _CAT:
        return _CAT[key]
    return ((key or "—"), SLATE)


def cat_key(raw) -> str:
    """Whatever the model wrote for a department → the stored key.

    It is asked for the key («warehouse») and answers with whatever reads
    naturally — the key, the label («Ombor»), or both («warehouse (Ombor)») —
    and the spellings arrive in the same response. A lookup that took one form
    only would silently drop a root-cause line, so both are accepted.
    """
    t = str(raw or "").strip()
    if not t:
        return ""
    for sep in ("—", "–", " - ", ":", "·", "("):
        if sep in t:
            t = t.split(sep, 1)[0]
    t = t.strip().casefold()
    for k, label, _hx in CATS:
        if t in (k, label.casefold(), k.replace("_", " ")):
            return k
    return t


# ── formatting ───────────────────────────────────────────────────────────────
def days1(v) -> str:
    """Days with one decimal and a comma, «—» where there is no answer."""
    if v is None:
        return "—"
    return f"{v:.1f}".replace(".", ",")


def latin_text(s: str | None) -> str:
    """A concern or a closing note in Latin script, otherwise exactly as typed —
    the SCRIPT changes and nothing else: no spelling, no wording."""
    return " ".join(transliterate_text((s or "").strip(), "uz").split())


def latin_name(s: str | None) -> str:
    return " ".join(transliterate((s or "").strip(), "uz").split())


def _mean(xs: list[float]) -> float | None:
    return (sum(xs) / len(xs)) if xs else None


def _share(part: int, whole: int) -> float:
    return (part / whole * 100.0) if whole else 0.0


def _delta_pct(cur: float, prev: float) -> float | None:
    return ((cur - prev) / prev * 100.0) if prev else None


def _vs(cur: int, prev: int) -> str:
    """«o'tgan hafta 412 (+12%)» — last week beside this one, never a verdict:
    more concerns filed can as easily mean the floor started speaking up."""
    d = _delta_pct(cur, prev)
    return f"o'tgan hafta {num(prev)}" + (f" ({signed_pct(d)})" if d is not None else "")


# ── the lifecycle, the page's own rules at a fixed instant ──────────────────
def _close_day(r: dict) -> date | None:
    """The day a concern left the open pool — the page's flow-chart rule: its
    completion date, or the filing date where none was set or it predates it."""
    if r["status"] != "done":
        return None
    c, e = r["completion"], r["entry"]
    return e if (c is None or c < e) else c


def _open_at(r: dict, end: date) -> bool:
    """Filed by the end of `end` and not closed by then."""
    if r["entry"] > end:
        return False
    cd = _close_day(r)
    return cd is None or cd > end


def _deadline(r: dict) -> date | None:
    dd = r["deadline_days"]
    if dd is None:
        return None
    # The page's rule (routers/concerns._due): counted from the day the
    # receiver took the concern into work, or — for a deadline its creator
    # typed on filing, before 2026-09-23 — from the filing day.
    start = r.get("deadline_from") or r["entry"]
    try:
        return start + timedelta(days=int(dd))
    except OverflowError:
        # A deadline typed as «999999» days is a deadline that never comes —
        # it has one, so it is not «no deadline», and it can never be overdue.
        return date.max if int(dd) > 0 else date.min


def _overdue_at(r: dict, end: date) -> bool:
    """Open at the end of `end` and past its deadline — the page's rule
    (`due < today`) read on the morning after `end`."""
    dl = _deadline(r)
    return dl is not None and dl <= end and _open_at(r, end)


# ── what the deck is made of ─────────────────────────────────────────────────
def collect(*, rows: list[dict], moves: list[dict], units: list[dict],
            win: tuple[date, date], prev_win: tuple[date, date],
            plants: list[str]) -> dict:
    """Everything the twelve slides need, from the rows the router loaded.

    `rows` carry one concern each — its dates, status, deadline, department,
    cell, unit, text, closing note, and where it sat at the end of the window.
    `moves` are the escalation steps of both windows; `units` every unit the
    register can name. Nothing here queries.
    """
    a, b = win
    pa, pb = prev_win
    ref = b + timedelta(days=1)           # the morning after the week
    days = [a + timedelta(days=i) for i in range((b - a).days + 1)]

    rs: list[dict] = []
    for raw in rows:
        r = dict(raw)
        cd = _close_day(r)
        dl = _deadline(r)
        r["close_day"] = cd
        r["deadline"] = dl
        r["filed"] = a <= r["entry"] <= b
        r["closed"] = cd is not None and a <= cd <= b
        r["open_end"] = _open_at(r, b)
        r["overdue_end"] = _overdue_at(r, b)
        r["p_filed"] = pa <= r["entry"] <= pb
        r["p_closed"] = cd is not None and pa <= cd <= pb
        r["p_open_end"] = _open_at(r, pb)
        r["p_overdue_end"] = _overdue_at(r, pb)
        r["age"] = (ref - r["entry"]).days
        r["over_days"] = (ref - dl).days if r["overdue_end"] and dl else 0
        r["to_close"] = (cd - r["entry"]).days if cd else None
        r["text_l"] = latin_text(r["text"])
        r["resolution_l"] = latin_text(r.get("resolution"))
        r["cyrillic"] = r["text_l"] != " ".join((r["text"] or "").split())
        rs.append(r)

    def n(key: str, pool=None) -> int:
        return sum(1 for r in (rs if pool is None else pool) if r[key])

    filed = [r for r in rs if r["filed"]]
    closed = [r for r in rs if r["closed"]]
    open_end = [r for r in rs if r["open_end"]]
    overdue_end = [r for r in rs if r["overdue_end"]]
    p_closed = [r for r in rs if r["p_closed"]]

    avg_close = _mean([r["to_close"] for r in closed])
    p_avg_close = _mean([r["to_close"] for r in p_closed])

    # ── departments ─────────────────────────────────────────────────────────
    by_cat: dict[str, list[dict]] = defaultdict(list)
    for r in rs:
        by_cat[r["category"] or ""].append(r)
    cats = []
    for k, g in by_cat.items():
        f = [r for r in g if r["filed"]]
        if not f and not n("open_end", g) and not n("p_filed", g):
            continue
        label, colour = cat_meta(k)
        unit_count = Counter(r["unit"] for r in f)
        cats.append({
            "key": k, "label": label, "colour": colour,
            "filed": len(f), "prev_filed": n("p_filed", g),
            "closed": n("closed", g),
            "filed_closed": sum(1 for r in f if r["close_day"] and r["close_day"] <= b),
            "open": n("open_end", g), "overdue": n("overdue_end", g),
            "share": _share(len(f), len(filed)),
            "avg_close": _mean([r["to_close"] for r in g if r["closed"]]),
            "top_units": unit_count.most_common(3),
            # What the top-3 slide may quote: this week's filings and every
            # concern of the department still open at the end of it.
            "pool": [r for r in g if r["filed"] or r["open_end"]],
        })
    cats.sort(key=lambda c: (-c["filed"], -c["open"], _CAT_ORDER.get(c["key"], 99)))

    # ── units ───────────────────────────────────────────────────────────────
    unit_meta = {u["id"]: u for u in units}
    by_unit: dict = defaultdict(list)
    for r in rs:
        by_unit[r["unit_id"]].append(r)
    # A unit is tagged with its plant only when it is NOT in the plant most
    # units belong to — twenty rows reading «Uchtepa» say nothing, the one
    # reading «Keles» is the fact.
    plant_count = Counter(u.get("plant") for u in units
                          if u.get("plant") and not u.get("archived"))
    main_plant = plant_count.most_common(1)[0][0] if plant_count else None
    unit_rows = []
    for uid, g in by_unit.items():
        f = n("filed", g)
        if not (f or n("closed", g) or n("open_end", g)):
            continue
        meta = unit_meta.get(uid) or {}
        name = latin_name(meta.get("name") or (g[0].get("unit") if g else "")) or "Biriktirilmagan"
        unit_rows.append({
            "id": uid, "name": name, "shift": meta.get("shift"),
            "plant": meta.get("plant") if meta.get("plant") != main_plant else None,
            "filed": f, "closed": n("closed", g), "open": n("open_end", g),
            "overdue": n("overdue_end", g),
            "prev_filed": n("p_filed", g), "prev_overdue": n("p_overdue_end", g),
            "avg_close": _mean([r["to_close"] for r in g if r["closed"]]),
        })
    unit_rows.sort(key=lambda u: (-u["overdue"], -u["open"], -u["filed"], u["name"]))
    filed_units = {r["unit_id"] for r in filed}
    silent = sorted(latin_name(u["name"]) for u in units
                    if not u.get("archived") and u["id"] not in filed_units)

    # ── the week, day by day ────────────────────────────────────────────────
    f_by_day = Counter(r["entry"] for r in filed)
    c_by_day = Counter(r["close_day"] for r in closed)
    daily = [{"date": d, "label": day_label(d), "full": day_full(d),
              "filed": f_by_day.get(d, 0), "closed": c_by_day.get(d, 0)} for d in days]
    peak = max(daily, key=lambda x: (x["filed"], -x["date"].toordinal())) if daily else None
    peak_cats = (Counter(r["category"] for r in filed if r["entry"] == peak["date"]).most_common(3)
                 if peak and peak["filed"] else [])

    # ── what is still open, oldest first ────────────────────────────────────
    oldest = sorted(open_end, key=lambda r: (-r["age"], r["no"]))

    # ── the chain ───────────────────────────────────────────────────────────
    def _dir(m: dict) -> int:
        return LEVEL_IDX.get(m["to_level"], 0) - LEVEL_IDX.get(m["from_level"], 0)

    moves = [dict(m, reason_l=latin_text(m.get("reason"))) for m in moves]
    cur_moves = [m for m in moves if a <= m["day"] <= b]
    prev_moves = [m for m in moves if pa <= m["day"] <= pb]
    ups = [m for m in cur_moves if _dir(m) > 0]
    downs = [m for m in cur_moves if _dir(m) <= 0]
    open_by_level = Counter(r["level_end"] for r in open_end)

    # ── what the model is shown, and what it may cite ──────────────────────
    brief = filed + [r for r in oldest if not r["filed"]]

    def delta(cur: int, prev: int) -> float | None:
        return _delta_pct(cur, prev)

    return {
        "plants": " · ".join(plants) if plants else "—",
        "window": win, "prev_window": prev_win, "ref": ref,
        "period": period_words(win), "prev_period": period_words(prev_win),
        "days": days, "daily": daily, "peak": peak, "peak_cats": peak_cats,
        "filed": len(filed), "prev_filed": n("p_filed"),
        "filed_delta": delta(len(filed), n("p_filed")),
        "closed": len(closed), "prev_closed": len(p_closed),
        "open_end": len(open_end), "prev_open_end": n("p_open_end"),
        "overdue_end": len(overdue_end), "prev_overdue_end": n("p_overdue_end"),
        "backlog_older": sum(1 for r in open_end if r["entry"] < a),
        "filed_closed": sum(1 for r in filed if r["close_day"] and r["close_day"] <= b),
        "avg_close": avg_close, "prev_avg_close": p_avg_close,
        "categories": cats,
        "units": unit_rows,
        "units_filing": len(filed_units),
        "silent_units": silent,
        "oldest": oldest,
        "moves": cur_moves, "ups": ups, "downs": downs,
        "prev_ups": sum(1 for m in prev_moves if _dir(m) > 0),
        "prev_downs": sum(1 for m in prev_moves if _dir(m) <= 0),
        "open_by_level": {lv: open_by_level.get(lv, 0) for lv in LEVELS},
        "brief": brief,
        "brief_nos": {r["no"] for r in brief},
        "by_no": {r["no"]: r for r in rs},
        "quality": {
            "cyrillic": sum(1 for r in brief if r["cyrillic"]),
            "thin": sum(1 for r in filed if len(r["text_l"]) < THIN_CHARS),
            "no_deadline_open": sum(1 for r in open_end if r["deadline"] is None),
            "worker_filed": sum(1 for r in filed if r["worker"]),
            "worker_open": sum(1 for r in open_end if r["worker"]),
        },
    }


# ── small drawing helpers of this deck's own ─────────────────────────────────
def _page(prs, eyebrow: str, title: str, page: int, footer: str):
    s = _blank(prs)
    _rect(s, 0, 0, W, H, fill=CREAM)
    _chrome(s, eyebrow, title, page, footer)
    return s


def _cols(x0: float, specs: list[tuple], gap: float) -> list[tuple]:
    """[(label, width, align)] → [(label, x, width, align)], left to right."""
    out, x = [], x0
    for label, w, al in specs:
        out.append((label, x, w, al))
        x += w + gap
    return out


def _header(slide, y: float, cols: list[tuple], h: float = 0.32):
    _rect(slide, M, y, CW, h, fill=INNER, line=LINE, radius=0.06)
    for label, x, w, al in cols:
        if label:
            _text(slide, x, y + (h - 0.2) / 2, w, 0.2, label, size=8, color=MUTED,
                  bold=True, align=al, caps=True)


def _dot(slide, x: float, y: float, colour, d: float = 0.13):
    _rect(slide, x, y, d, d, fill=colour, radius=0.03)


def _title(p: dict) -> str:
    """A prose title in sentence case — unless it is a label the register
    spells itself («ARS», «IT»), which a fallback marks `raw`."""
    t = p.get("title", "") or ""
    return t if p.get("raw") else _sentence(t)


def _state(r: dict) -> str:
    """Where a concern stood at the end of the week, in words."""
    if r["closed"] or (r["close_day"] and not r["open_end"]):
        return f"{r['to_close']} kunda yopildi"
    if r["overdue_end"]:
        return f"ochiq {r['age']} kun · muddati {r['over_days']} kun o'tgan"
    return f"ochiq {r['age']} kun"


def _meta(r: dict) -> str:
    who = " · ishchi yozgan" if r.get("worker") else ""
    return f"№{r['no']} · {r['cell'] or '—'} · {day_label(r['entry'])} · {_state(r)}{who}"


# A quoted concern is the one string on this deck whose length nobody here
# controls. It gets up to Q_LINES lines, its meta line is placed at whatever
# height the text actually took, and a stack of them is drawn against a budget
# — the lesson the Ojidaniya deck's event notes taught on 2026-09-04.
Q_SIZE, Q_META, Q_LINES, Q_GAP = 9, 8, 2, 0.09


def _quote_block(r: dict, w: float, *, resolution: bool):
    tw = w - 0.16
    text, _ = deck_text.fit(f"«{r['text_l']}»", tw, Q_SIZE, BODY_FONT,
                            max_lines=Q_LINES)
    meta = deck_text.fit(_meta(r), tw, Q_META, BODY_FONT, max_lines=1)[0]
    res: list[str] = []
    if resolution and r["resolution_l"] and r["close_day"] and not r["open_end"]:
        res = deck_text.fit(f"Yechim: «{r['resolution_l']}»", tw, Q_META,
                            BODY_FONT, max_lines=1)[0]
    h = (deck_text.block_h_in(len(text), Q_SIZE) + deck_text.block_h_in(1, Q_META)
         + deck_text.block_h_in(len(res), Q_META) + Q_GAP)
    return text, meta, res, h


def _quotes(slide, x: float, y: float, w: float, rows: list[dict], *, bottom: float,
            colour_fn, resolution: bool = True) -> tuple[float, int]:
    """A stack of quoted concerns against a vertical BUDGET: laid out while
    they fit, stopped before the next one would cross `bottom`. Returns where
    it ended and how many it drew."""
    shown = 0
    for r in rows:
        text, meta, res, h = _quote_block(r, w, resolution=resolution)
        if y + h - Q_GAP > bottom:
            break
        th = deck_text.block_h_in(len(text), Q_SIZE)
        mh = deck_text.block_h_in(1, Q_META)
        _rect(slide, x, y, 0.035, h - Q_GAP, fill=colour_fn(r) or GOLD)
        _text(slide, x + 0.16, y, w - 0.16, th, "\n".join(text), size=Q_SIZE,
              color=INK2, max_lines=len(text))
        _text(slide, x + 0.16, y + th, w - 0.16, mh, "\n".join(meta), size=Q_META, color=FAINT)
        if res:
            rh = deck_text.block_h_in(len(res), Q_META)
            _text(slide, x + 0.16, y + th + mh, w - 0.16, rh, "\n".join(res),
                  size=Q_META, color=GREEN, max_lines=len(res))
        y += h
        shown += 1
    return y, shown


def _two_series_chart(slide, x, y, w, h, categories, first, second, names, colours):
    """Two series side by side — «yangi» beside «yopildi» for each day. A native
    chart, so the reader can click into it and it stays sharp at any zoom."""
    data = CategoryChartData()
    data.categories = categories
    data.add_series(names[0], first)
    data.add_series(names[1], second)
    frame = slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED,
                                   Inches(x), Inches(y), Inches(w), Inches(h), data)
    chart = frame.chart
    chart.has_title = False
    chart.has_legend = True
    chart.legend.position = XL_LEGEND_POSITION.TOP
    chart.legend.include_in_layout = False
    chart.legend.font.size = Pt(9)
    chart.legend.font.name = BODY_FONT
    chart.legend.font.color.rgb = INK2

    plot = chart.plots[0]
    plot.gap_width = 55
    plot.overlap = -8
    plot.has_data_labels = True
    labels = plot.data_labels
    labels.number_format = '#,##0'
    labels.number_format_is_linked = False
    labels.position = XL_LABEL_POSITION.OUTSIDE_END
    labels.font.size = Pt(8.5)
    labels.font.name = BODY_FONT
    labels.font.color.rgb = MUTED
    for series, colour in zip(plot.series, colours):
        series.format.fill.solid()
        series.format.fill.fore_color.rgb = colour
        series.format.line.fill.background()

    cat_ax, val_ax = chart.category_axis, chart.value_axis
    cat_ax.has_major_gridlines = False
    cat_ax.tick_labels.font.size = Pt(9)
    cat_ax.tick_labels.font.name = BODY_FONT
    cat_ax.tick_labels.font.color.rgb = INK2
    cat_ax.format.line.color.rgb = LINE
    val_ax.has_major_gridlines = False
    val_ax.visible = False
    return chart


# ── slide 1 · cover ──────────────────────────────────────────────────────────
def _cover(prs, d: dict, narr):
    s = _blank(prs)
    _rect(s, 0, 0, W, H, fill=DARK)
    _rect(s, 0, 0, W, 0.09, fill=GOLD)

    _text(s, M, 1.05, CW, 0.25, "HAFTALIK OPERATSION TAHLIL · XAVOTIRLAR",
          size=10, color=GOLD2, bold=True)
    _text(s, M, 1.42, 8.6, 1.5, "Xavotirlar\nhaftalik tahlili",
          size=40, color=WHITE, font=HEAD_FONT, bold=True, line=1.05)
    _text(s, M, 3.05, 7.6, 0.3,
          f"{num(d['filed'])} ta yangi xavotir  ·  {d['units_filing']} brigadir  ·  "
          f"{d['period']}", size=13, color=ONDARK)
    _text(s, M, 3.42, 7.6, 0.28, d["plants"], size=11, color=GOLD2, bold=True)

    _text(s, M, 4.05, 6.6, 1.0, f"{num(d['filed'])} ta xavotir",
          size=48, color=GOLD2, font=HEAD_FONT, bold=True)
    move = (f"o'tgan haftadan {signed_pct(d['filed_delta'])}"
            if d["filed_delta"] is not None else "taqqoslash uchun ma'lumot yo'q")
    _text(s, M, 5.05, 7.4, 0.5, f"shu hafta kiritildi · {move}", size=11.5, color=ONDARK)

    # Where the week ended, as four tiles: what was closed, what is still
    # open, what is past its deadline, and what went up the chain.
    tiles = [
        (num(d["closed"]), "yopildi", WHITE),
        (num(d["open_end"]), "hafta oxirida ochiq", WHITE),
        (num(d["overdue_end"]), "muddati o'tgan", SOFT_RED),
        (num(len(d["ups"])), "yuqoriga ko'tarildi", WHITE),
    ]
    tw, th, gap = 2.12, 0.95, 0.16
    for i, (big, label, colour) in enumerate(tiles):
        x = 8.35 + (i % 2) * (tw + gap)
        y = 4.05 + (i // 2) * (th + 0.14)
        _rect(s, x, y, tw, th, fill=TILE, radius=0.1)
        _text(s, x + 0.18, y + 0.1, tw - 0.36, 0.44, big, size=22, color=colour,
              font=HEAD_FONT, bold=True)
        _text(s, x + 0.18, y + 0.58, tw - 0.36, 0.24, label, size=9, color=ONDARK)

    _text(s, M, H - 0.5, CW, 0.24, "Safia · Ishlab chiqarish boshqaruvi paneli",
          size=9, color=RGBColor(0x7A, 0x6A, 0x5C))


# ── slide 2 · executive summary ──────────────────────────────────────────────
def _summary(prs, d: dict, narr, page: int, footer: str):
    s = _page(prs, "Qisqacha", "Asosiy xulosalar", page, footer)

    open_share = _share(d["overdue_end"], d["open_end"])
    kpis = [
        (num(d["filed"]), "yangi xavotir", _vs(d["filed"], d["prev_filed"])),
        (num(d["closed"]), "yopildi",
         f"o'rtacha {days1(d['avg_close'])} kunda · {_vs(d['closed'], d['prev_closed'])}"),
        (num(d["open_end"]), "hafta oxirida ochiq",
         f"{num(d['backlog_older'])} tasi oldingi haftalardan qolgan"),
        (num(d["overdue_end"]), "muddati o'tgan",
         f"ochiqlarning {pct(open_share)} · o'tgan hafta {num(d['prev_overdue_end'])}"),
    ]
    cw = (CW - 3 * 0.16) / 4
    for i, (big, label, sub) in enumerate(kpis):
        x = M + i * (cw + 0.16)
        _card(s, x, 1.42, cw, 1.24)
        _text(s, x + 0.18, 1.58, cw - 0.36, 0.42, big, size=21,
              color=RED if i == 3 and d["overdue_end"] else INK, font=HEAD_FONT, bold=True)
        _text(s, x + 0.18, 2.03, cw - 0.36, 0.24, label, size=10, color=INK2)
        _text(s, x + 0.18, 2.28, cw - 0.36, 0.3, sub, size=8.5, color=FAINT)

    points = [p for p in _ai_list(narr, "summary_points", 3) if isinstance(p, dict)]
    if not points:
        # Without the model the slide still says something true: the three
        # departments that carried the week, stated plainly.
        points = [{"title": f"№{i + 1} — {c['label']}: {num(c['filed'])} ta yangi xavotir",
                   "raw": True,
                   "body": (f"{num(c['filed_closed'])} tasi hafta oxirigacha yopildi; "
                            f"hafta oxirida {num(c['open'])} tasi ochiq, "
                            f"{num(c['overdue'])} tasi muddati o'tgan.")}
                  for i, c in enumerate(d["categories"][:3])]

    y = 2.92
    for p in points:
        h = 0.92
        _card(s, M, y, CW, h)
        _rect(s, M, y, 0.055, h, fill=GOLD)
        _text(s, M + 0.28, y + 0.16, CW - 0.56, 0.28, _title(p),
              size=12.5, color=INK, font=HEAD_FONT, bold=True)
        _text(s, M + 0.28, y + 0.47, CW - 0.56, 0.4, p.get("body", ""),
              size=10, color=MUTED, line=1.18)
        y += h + 0.14

    head = _ai(narr, "summary_headline")
    _text(s, M, y + 0.04, CW, 0.46, head or NO_AI,
          size=10.5, color=INK2 if head else FAINT, bold=bool(head), line=1.2)


# ── slide 3 · what kept coming back ─────────────────────────────────────────
def _themes(d: dict, narr) -> list[dict]:
    """The model's recurring problems, each COUNTED here from the concern
    numbers it cited — never from a number it wrote. A cited number the week
    does not hold is dropped, and a problem left citing nothing is dropped."""
    out = []
    for t in _ai_list(narr, "themes", 6):
        if not isinstance(t, dict):
            continue
        nos = []
        for v in t.get("nums") or []:
            try:
                no = int(v)
            except (TypeError, ValueError):
                continue
            if no in d["brief_nos"] and no not in nos:
                nos.append(no)
        if not nos or not (t.get("title") or "").strip():
            continue
        out.append({"title": t.get("title", ""), "body": t.get("body", ""),
                    "cat": cat_key(t.get("cat")), "nos": nos})
    out.sort(key=lambda t: -len(t["nos"]))
    return out[:5]


def _recurring(prs, d: dict, narr, page: int, footer: str):
    themes = _themes(d, narr)
    from_ai = bool(themes)
    title = ("Hafta davomida nima qayta-qayta ko'tarildi?" if from_ai
             else "Eng ko'p xavotir kelgan bo'limlar")
    s = _page(prs, "Takrorlanuvchi muammolar", title, page, footer)

    if not from_ai:
        # No model: the five departments that took the most concerns, each
        # with its most recent one quoted — true, and still worth the slide.
        themes = []
        for c in d["categories"][:5]:
            recent = sorted((r for r in c["pool"] if r["filed"]),
                            key=lambda r: (r["entry"], r["no"]), reverse=True)
            if not recent:
                continue
            themes.append({"title": c["label"],
                           "body": (f"{num(c['filed_closed'])} tasi yopildi, "
                                    f"{num(c['open'])} tasi hafta oxirida ochiq."),
                           "cat": c["key"], "nos": [r["no"] for r in recent],
                           "count": c["filed"], "label_title": True})

    y = 1.36
    rh, gap = 0.98, 0.1
    rw = CW - 3.02                      # the text column stops 0.1" short of the count
    for i, t in enumerate(themes):
        # A department the model named in words of its own is not printed as
        # one: a label here is a claim the register makes, not the model.
        label, colour = cat_meta(t["cat"]) if t["cat"] in _CAT else ("", GOLD)
        _card(s, M, y, CW, rh)
        _rect(s, M, y, 0.055, rh, fill=colour)
        _chip(s, M + 0.22, y + 0.14, 0.34, 0.34, str(i + 1), fill=INK, color=GOLD2, size=11)
        # A department's own label («ARS», «IT») is printed as the register
        # spells it; only the model's prose gets the sentence-case pass.
        _text(s, M + 0.72, y + 0.12, rw, 0.26,
              t["title"] if t.get("label_title") else _sentence(t["title"]), size=12,
              color=INK, font=HEAD_FONT, bold=True)
        _text(s, M + 0.72, y + 0.39, rw, 0.34, t["body"], size=9, color=MUTED)
        ex = d["by_no"].get(t["nos"][0])
        if ex:
            _text(s, M + 0.72, y + 0.74, rw, 0.2,
                  f"«{ex['text_l']}» — №{ex['no']} · {ex['cell'] or '—'}",
                  size=8.5, color=FAINT)
        count = t.get("count", len(t["nos"]))
        _text(s, M + CW - 2.2, y + 0.13, 2.0, 0.42, f"{num(count)} ta", size=20,
              color=INK, font=HEAD_FONT, bold=True, align=PP_ALIGN.RIGHT)
        if label:
            _text(s, M + CW - 2.2, y + 0.58, 2.0, 0.22, label, size=9, color=colour,
                  bold=True, align=PP_ALIGN.RIGHT)
        y += rh + gap

    note = ("Soni — AI shu muammoga bog'lagan xavotirlar, ular raqami bo'yicha sanaldi. "
            "Misol matnlari o'zgartirilmagan." if from_ai
            else NO_AI + " Soni — bo'limga shu hafta kelgan xavotirlar.")
    _text(s, M, min(y + 0.02, 6.72), CW, 0.22, note, size=8.5, color=FAINT)


# ── slide 4 · departments ────────────────────────────────────────────────────
def _departments(prs, d: dict, narr, page: int, footer: str):
    s = _page(prs, "Bo'limlar kesimida", "Qaysi bo'limga qancha xavotir tushdi?", page, footer)

    cols = _cols(M + 0.2, [
        ("Bo'lim", 2.75, PP_ALIGN.LEFT), ("Yangi", 0.9, PP_ALIGN.RIGHT),
        ("O'tgan hafta", 1.12, PP_ALIGN.RIGHT), ("O'zgarish", 1.0, PP_ALIGN.RIGHT),
        ("Yopildi", 1.5, PP_ALIGN.RIGHT), ("Ochiq", 0.9, PP_ALIGN.RIGHT),
        ("Muddati o'tgan", 1.3, PP_ALIGN.RIGHT), ("", 1.6, PP_ALIGN.LEFT),
    ], 0.12)
    _header(s, 1.36, cols)

    rows = d["categories"][:15]
    if not rows:
        _text(s, M, 3.2, CW, 0.4, "Bu hafta xavotir kiritilmagan.", size=12,
              color=FAINT, align=PP_ALIGN.CENTER)
        return
    top = max(c["filed"] for c in rows) or 1
    y0, bottom = 1.76, 6.18
    rh = min(0.34, (bottom - y0) / len(rows))
    size = 9.5 if rh >= 0.28 else 9
    ty = (rh - deck_text.block_h_in(1, size)) / 2
    for i, c in enumerate(rows):
        y = y0 + i * rh
        (_, x0, w0, _a), (_, x1, w1, a1), (_, x2, w2, a2), (_, x3, w3, a3), \
            (_, x4, w4, a4), (_, x5, w5, a5), (_, x6, w6, a6), (_, x7, w7, _a7) = cols
        _dot(s, x0, y + (rh - 0.13) / 2, c["colour"])
        _text(s, x0 + 0.22, y + ty, w0 - 0.22, rh - ty, c["label"], size=size, color=INK,
              bold=True, max_lines=1)
        _text(s, x1, y + ty, w1, rh - ty, num(c["filed"]), size=size, color=INK,
              bold=True, align=a1, max_lines=1)
        _text(s, x2, y + ty, w2, rh - ty, num(c["prev_filed"]), size=size, color=MUTED,
              align=a2, max_lines=1)
        diff = c["filed"] - c["prev_filed"]
        arrow = "▲" if diff > 0 else ("▼" if diff < 0 else "—")
        _text(s, x3, y + ty, w3, rh - ty, f"{arrow} {num(abs(diff))}" if diff else "—",
              size=size, color=MUTED, align=a3, max_lines=1)
        closed_pct = _share(c["filed_closed"], c["filed"]) if c["filed"] else None
        _text(s, x4, y + ty, w4, rh - ty,
              f"{num(c['filed_closed'])} · {pct(closed_pct)}" if closed_pct is not None else "—",
              size=size, color=GREEN if c["filed_closed"] else MUTED, align=a4, max_lines=1)
        _text(s, x5, y + ty, w5, rh - ty, num(c["open"]), size=size, color=INK2,
              align=a5, max_lines=1)
        _text(s, x6, y + ty, w6, rh - ty, num(c["overdue"]) if c["overdue"] else "—",
              size=size, color=RED if c["overdue"] else MUTED, bold=bool(c["overdue"]),
              align=a6, max_lines=1)
        if c["filed"]:
            _rect(s, x7, y + (rh - 0.12) / 2, max(0.04, w7 * c["filed"] / top), 0.12,
                  fill=c["colour"], radius=0.03)
        _rect(s, M + 0.2, y + rh - 0.006, CW - 0.4, 0.006, fill=LINE)

    note = _ai(narr, "categories_note")
    _rect(s, M, 6.22, CW, 0.46, fill=INNER, line=LINE, radius=0.08)
    _text(s, M + 0.22, 6.27, CW - 0.44, 0.36, note or NO_AI, size=9.5,
          color=INK2 if note else FAINT, line=1.1)
    _text(s, M, 6.72, CW, 0.2,
          "Yangi — shu hafta kiritilganlar. Yopildi — ulardan hafta oxirigacha yopilgani. "
          "Ochiq va muddati o'tgan — hafta oxiridagi holat, eski xavotirlar bilan birga.",
          size=8, color=FAINT)


# ── slide 5 · the three departments, opened up ───────────────────────────────
def _pick_quotes(c: dict) -> list[dict]:
    """What a department card quotes: its oldest overdue problems first (at most
    two), then the ones that took longest to close, with how they were closed,
    then whatever else came in this week."""
    pool = c["pool"]
    stuck = sorted((r for r in pool if r["open_end"]),
                   key=lambda r: (not r["overdue_end"], -r["age"], r["no"]))[:2]
    done = sorted((r for r in pool if r["closed"] and r["resolution_l"]),
                  key=lambda r: (-(r["to_close"] or 0), r["no"]))
    seen = {r["no"] for r in stuck}
    rest = sorted((r for r in pool if r["no"] not in seen and not r["closed"]),
                  key=lambda r: (r["entry"], r["no"]), reverse=True)
    out, used = [], set()
    for r in stuck + done + rest:
        if r["no"] not in used:
            out.append(r)
            used.add(r["no"])
    return out


def _top_departments(prs, d: dict, narr, page: int, footer: str):
    s = _page(prs, "So'ralgan kesim", "Top-3 bo'lim: aynan nima haqida xavotir?", page, footer)

    roots = {cat_key(r.get("cat")): r.get("root")
             for r in _ai_list(narr, "category_roots", 6) if isinstance(r, dict)}
    three = [c for c in d["categories"] if c["filed"]][:3]
    if not three:
        _text(s, M, 3.2, CW, 0.4, "Bu hafta xavotir kiritilmagan.", size=12,
              color=FAINT, align=PP_ALIGN.CENTER)
        return
    cw = (CW - 2 * 0.18) / 3
    for i, c in enumerate(three):
        x = M + i * (cw + 0.18)
        _card(s, x, 1.32, cw, 4.86)
        _rect(s, x, 1.32, cw, 0.055, fill=c["colour"])
        _dot(s, x + 0.18, 1.56, c["colour"], d=0.18)
        _text(s, x + 0.46, 1.5, cw - 0.64, 0.28, c["label"], size=12.5, color=INK,
              font=HEAD_FONT, bold=True)
        _text(s, x + 0.18, 1.82, cw - 0.36, 0.22,
              f"{num(c['filed'])} yangi · {num(c['filed_closed'])} yopildi · "
              f"{num(c['open'])} ochiq", size=9.5, color=BROWN, bold=True)
        who = " · ".join(f"{short_name(latin_name(u), 1.2, 8.5)} {k}" for u, k in c["top_units"] if u)
        _text(s, x + 0.18, 2.08, cw - 0.36, 0.22, f"eng ko'p: {who}" if who else "—",
              size=8.5, color=FAINT)

        _quotes(s, x + 0.18, 2.4, cw - 0.36, _pick_quotes(c), bottom=5.36,
                colour_fn=lambda r: RED if r["overdue_end"] else (
                    CLOSED_C if not r["open_end"] else GOLD))

        root = roots.get(c["key"])
        _rect(s, x + 0.18, 5.42, cw - 0.36, 0.72, fill=INNER, radius=0.06)
        _text(s, x + 0.32, 5.5, cw - 0.64, 0.58,
              ("Ildiz: " + root) if root else
              (f"Muddati o'tgan: {num(c['overdue'])} · o'rtacha yopilish "
               f"{days1(c['avg_close'])} kun"),
              size=8.5, color=INK2 if root else FAINT, line=1.18)

    total = sum(c["filed"] for c in three)
    _text(s, M, 6.35, CW, 0.4,
          f"Uchtasi birgalikda {num(total)} ta xavotir — shu hafta kelganlarning "
          f"{pct(_share(total, d['filed']))} qismi. Qizil chiziq — muddati o'tgan, "
          f"yashil — yopilgan.", size=9.5, color=MUTED)


# ── slide 6 · every unit, ranked by what is past its deadline ────────────────
def _units(prs, d: dict, narr, page: int, footer: str):
    s = _page(prs, "Brigadirlar kesimida", "Muddati o'tgan xavotirlar kimda ko'p?", page, footer)

    cols = _cols(M + 0.2, [
        ("#", 0.36, PP_ALIGN.LEFT), ("Brigadir", 3.3, PP_ALIGN.LEFT),
        ("Yangi", 0.85, PP_ALIGN.RIGHT), ("Yopildi", 0.9, PP_ALIGN.RIGHT),
        ("Ochiq", 0.85, PP_ALIGN.RIGHT), ("Muddati o'tgan", 1.3, PP_ALIGN.RIGHT),
        ("O'rt. yopilish", 1.3, PP_ALIGN.RIGHT), ("", 1.9, PP_ALIGN.LEFT),
    ], 0.12)
    _header(s, 1.34, cols)

    rows = d["units"]
    shown = rows[:22]
    y0, bottom = 1.72, 6.36
    if not shown:
        _text(s, M, 3.2, CW, 0.4, "Bu hafta xavotir kiritilmagan.", size=12,
              color=FAINT, align=PP_ALIGN.CENTER)
        return
    rh = min(0.3, (bottom - y0) / len(shown))
    size = 9.5 if rh >= 0.28 else (9 if rh >= 0.24 else 8.5)
    ty = max(0.0, (rh - deck_text.block_h_in(1, size)) / 2)
    top = max((u["overdue"] for u in shown), default=0) or 1
    for i, u in enumerate(shown):
        y = y0 + i * rh
        (_, xn, wn, an), (_, x0, w0, _a0), (_, x1, w1, a1), (_, x2, w2, a2), \
            (_, x3, w3, a3), (_, x4, w4, a4), (_, x5, w5, a5), (_, x6, w6, _a6) = cols
        if i % 2:
            _rect(s, M + 0.1, y, CW - 0.2, rh, fill=INNER)
        _text(s, xn, y + ty, wn, rh - ty, str(i + 1), size=size, color=FAINT,
              align=an, max_lines=1)
        tag = " · ".join(t for t in (f"{u['shift']}-smena" if u["shift"] else "",
                                     u["plant"] or "") if t)
        name = short_name(u["name"], w0 - 0.2 - (deck_text.width_pt(f" ({tag})", size) / 72
                                                if tag else 0), size, bold=True)
        _text(s, x0, y + ty, w0, rh - ty, f"{name} ({tag})" if tag else name, size=size,
              color=INK, bold=True, max_lines=1)
        _text(s, x1, y + ty, w1, rh - ty, num(u["filed"]), size=size, color=INK2,
              align=a1, max_lines=1)
        _text(s, x2, y + ty, w2, rh - ty, num(u["closed"]), size=size, color=INK2,
              align=a2, max_lines=1)
        _text(s, x3, y + ty, w3, rh - ty, num(u["open"]), size=size, color=INK2,
              align=a3, max_lines=1)
        _text(s, x4, y + ty, w4, rh - ty, num(u["overdue"]) if u["overdue"] else "—",
              size=size, color=RED if u["overdue"] else MUTED, bold=bool(u["overdue"]),
              align=a4, max_lines=1)
        _text(s, x5, y + ty, w5, rh - ty,
              f"{days1(u['avg_close'])} kun" if u["avg_close"] is not None else "—",
              size=size, color=MUTED, align=a5, max_lines=1)
        if u["overdue"]:
            _rect(s, x6, y + (rh - 0.1) / 2, max(0.04, w6 * u["overdue"] / top), 0.1,
                  fill=RED, radius=0.03)

    note = _ai(narr, "units_note")
    more = len(rows) - len(shown)
    line = note or NO_AI
    if more > 0:
        line = f"Yana {more} brigadir ro'yxatga sig'madi. " + line
    _text(s, M, 6.42, CW, 0.28, line, size=9, color=INK2 if note else FAINT)
    _text(s, M, 6.72, CW, 0.2,
          "Tartib: hafta oxirida muddati o'tgan xavotirlar, keyin ochiqlar, keyin yangilar "
          "soni bo'yicha. Ishchilar yozgan xavotirlar ham hisobga olingan.",
          size=8, color=FAINT)


# ── slide 7 · the oldest open concerns ───────────────────────────────────────
def _oldest(prs, d: dict, narr, page: int, footer: str):
    s = _page(prs, "Hal qilinmaganlar", "Eng uzoq ochiq turgan xavotirlar", page, footer)
    b = d["window"][1]
    _text(s, M, 1.3, CW, 0.22,
          f"{day_label(b)} kun oxiriga holat · jami ochiq {num(d['open_end'])}, "
          f"shundan {num(d['overdue_end'])} tasi muddati o'tgan",
          size=9.5, color=FAINT)

    cols = _cols(M + 0.15, [
        ("№", 0.62, PP_ALIGN.LEFT), ("Yacheyka", 0.8, PP_ALIGN.LEFT),
        ("Bo'lim", 1.3, PP_ALIGN.LEFT), ("Xavotir", 5.05, PP_ALIGN.LEFT),
        ("Ochiq", 0.8, PP_ALIGN.RIGHT), ("Muddat", 1.15, PP_ALIGN.RIGHT),
        ("Kimda", 1.6, PP_ALIGN.LEFT),
    ], 0.1)
    _header(s, 1.6, cols, h=0.3)

    rows = d["oldest"][:10]
    if not rows:
        _text(s, M, 3.4, CW, 0.4, "Hafta oxirida ochiq xavotir qolmagan.", size=12,
              color=GREEN, align=PP_ALIGN.CENTER)
        return
    y, rh = 1.96, 0.46
    for i, r in enumerate(rows):
        (_, xa, wa, _1), (_, xb, wb, _2), (_, xc, wc, _3), (_, xd, wd, _4), \
            (_, xe, we, ae), (_, xf, wf, af), (_, xg, wg, _7) = cols
        if i % 2:
            _rect(s, M + 0.05, y, CW - 0.1, rh, fill=INNER)
        label, colour = cat_meta(r["category"])
        _text(s, xa, y + 0.07, wa, 0.2, f"№{r['no']}", size=9, color=INK, bold=True)
        _text(s, xb, y + 0.07, wb, 0.2, r["cell"] or "—", size=9, color=INK2)
        _dot(s, xc, y + 0.1, colour, d=0.11)
        _text(s, xc + 0.18, y + 0.07, wc - 0.18, 0.2, label, size=8.5, color=INK2)
        _text(s, xd, y + 0.06, wd, 0.34, f"«{r['text_l']}»", size=9,
              color=INK2)
        _text(s, xe, y + 0.07, we, 0.2, f"{r['age']} kun", size=9, color=INK,
              bold=True, align=ae)
        _text(s, xf, y + 0.07, wf, 0.2,
              f"{r['over_days']} kun o'tdi" if r["overdue_end"] else
              ("belgilanmagan" if r["deadline"] is None else "muddatida"),
              size=8.5, color=RED if r["overdue_end"] else MUTED,
              bold=bool(r["overdue_end"]), align=af)
        lv = LEVEL_LABEL.get(r["level_end"], r["level_end"] or "—")
        _text(s, xg, y + 0.04, wg, 0.18, lv, size=8, color=LEVEL_COLOUR.get(r["level_end"], MUTED),
              bold=True)
        _text(s, xg, y + 0.22, wg, 0.18,
              short_name(latin_name(r.get("holder_end")), wg, 8) or "—", size=8, color=MUTED)
        y += rh

    note = _ai(narr, "oldest_note")
    _text(s, M, 6.64, CW, 0.28, note or NO_AI, size=9, color=INK2 if note else FAINT)


# ── slide 8 · the chain ──────────────────────────────────────────────────────
def _move_block(m: dict, w: float):
    tw = w - 0.16
    fl = LEVEL_LABEL.get(m["from_level"], m["from_level"])
    tl = LEVEL_LABEL.get(m["to_level"], m["to_level"])
    head = deck_text.fit(f"№{m['no']} · {m['cell'] or '—'} · {day_label(m['day'])} "
                         f"{m['at']} · {fl} → {tl}", tw, 9, BODY_FONT, True, 1)[0]
    who = " → ".join(v for v in (latin_name(m.get("from_name")),
                                 latin_name(m.get("target_name"))) if v)
    names = deck_text.fit(who or "—", tw, 8.5, BODY_FONT, False, 1)[0]
    reason = deck_text.fit(f"Sabab: «{m['reason_l']}»", tw, 9,
                           BODY_FONT, False, 2)[0]
    h = (deck_text.block_h_in(1, 9) + deck_text.block_h_in(1, 8.5)
         + deck_text.block_h_in(len(reason), 9) + Q_GAP)
    return head, names, reason, h


def _chain(prs, d: dict, narr, page: int, footer: str):
    s = _page(prs, "Zanjir bo'ylab", "Nima yuqoriga ko'tarildi va nega?", page, footer)

    # left: the counts, and where the open pool sat when the week ended
    tiles = [(len(d["ups"]), "yuqoriga ko'tarildi", f"o'tgan hafta {num(d['prev_ups'])}"),
             (len(d["downs"]), "pastga qaytarildi", f"o'tgan hafta {num(d['prev_downs'])}")]
    tw = 1.72
    for i, (big, label, sub) in enumerate(tiles):
        x = M + i * (tw + 0.16)
        _card(s, x, 1.36, tw, 1.0)
        _text(s, x + 0.16, 1.44, tw - 0.32, 0.44, num(big), size=24, color=INK,
              font=HEAD_FONT, bold=True)
        _text(s, x + 0.16, 1.92, tw - 0.32, 0.2, label, size=9, color=INK2)
        _text(s, x + 0.16, 2.12, tw - 0.32, 0.2, sub, size=8, color=FAINT)

    pw = 2 * tw + 0.16
    _card(s, M, 2.52, pw, 2.2)
    _text(s, M + 0.18, 2.62, pw - 0.36, 0.22, "Ochiqlar hafta oxirida qayerda",
          size=9, color=GOLD, bold=True, caps=True)
    tot = sum(d["open_by_level"].values()) or 1
    y = 2.96
    for lv in LEVELS:
        k = d["open_by_level"].get(lv, 0)
        _dot(s, M + 0.18, y + 0.05, LEVEL_COLOUR[lv], d=0.12)
        _text(s, M + 0.4, y, pw - 1.2, 0.22, LEVEL_LABEL[lv], size=9.5, color=INK2)
        _text(s, M + pw - 0.78, y, 0.6, 0.22, num(k), size=9.5, color=INK, bold=True,
              align=PP_ALIGN.RIGHT)
        _rect(s, M + 0.4, y + 0.26, max(0.02, (pw - 0.6) * k / tot), 0.06,
              fill=LEVEL_COLOUR[lv], radius=0.02)
        y += 0.42

    note = _ai(narr, "moves_note")
    _rect(s, M, 4.86, pw, 1.9, fill=INNER, line=LINE, radius=0.08)
    _text(s, M + 0.18, 4.98, pw - 0.36, 1.66, note or NO_AI, size=9.5,
          color=INK2 if note else FAINT, line=1.2)

    # right: the moves themselves, the ones UP first — those are the problems a
    # level could not solve — each with the reason whoever moved it gave
    x = M + pw + 0.3
    w = CW - pw - 0.3
    moves = sorted(d["ups"], key=lambda m: (m["day"], m["at"])) + \
        sorted(d["downs"], key=lambda m: (m["day"], m["at"]))
    if not moves:
        _text(s, x, 3.3, w, 0.4, "Bu hafta hech bir xavotir zanjir bo'ylab ko'chirilmagan.",
              size=11, color=FAINT, align=PP_ALIGN.CENTER)
        return
    y, bottom, shown = 1.36, 6.52, 0
    for m in moves:
        head, names, reason, h = _move_block(m, w)
        if y + h - Q_GAP > bottom:
            break
        up = LEVEL_IDX.get(m["to_level"], 0) > LEVEL_IDX.get(m["from_level"], 0)
        _rect(s, x, y, 0.035, h - Q_GAP,
              fill=LEVEL_COLOUR.get(m["to_level"], GOLD) if up else SLATE)
        hh = deck_text.block_h_in(1, 9)
        nh = deck_text.block_h_in(1, 8.5)
        _text(s, x + 0.16, y, w - 0.16, hh, "\n".join(head), size=9, color=INK, bold=True)
        _text(s, x + 0.16, y + hh, w - 0.16, nh, "\n".join(names), size=8.5, color=MUTED)
        _text(s, x + 0.16, y + hh + nh, w - 0.16, deck_text.block_h_in(len(reason), 9),
              "\n".join(reason), size=9, color=INK2, max_lines=len(reason))
        y += h
        shown += 1
    if shown < len(moves):
        _text(s, x, 6.6, w, 0.22, f"yana {len(moves) - shown} ta ko'chirish ro'yxatga sig'madi",
              size=8.5, color=FAINT)


# ── slide 9 · the week, day by day ───────────────────────────────────────────
def _daily(prs, d: dict, narr, page: int, footer: str):
    peak = d["peak"]
    title = (f"Eng ko'p xavotir — {day_full(peak['date'])}" if peak and peak["filed"]
             else "Hafta dinamikasi")
    s = _page(prs, "Hafta dinamikasi", title, page, footer)

    _two_series_chart(s, M - 0.1, 1.35, 8.6, 4.35,
                      [x["label"] for x in d["daily"]],
                      [x["filed"] for x in d["daily"]],
                      [x["closed"] for x in d["daily"]],
                      ("Yangi", "Yopildi"), (GOLD, CLOSED_C))

    x = M + 8.75
    cw = CW - 8.75
    if peak and peak["filed"]:
        _card(s, x, 1.35, cw, 1.5, fill=INK)
        _text(s, x + 0.22, 1.5, cw - 0.44, 0.5, f"{num(peak['filed'])} ta",
              size=26, color=GOLD2, font=HEAD_FONT, bold=True)
        _text(s, x + 0.22, 2.02, cw - 0.44, 0.7, f"yangi xavotir — {day_full(peak['date'])}",
              size=10, color=ONDARK, line=1.18)
        _text(s, x, 3.0, cw, 0.22, "O'sha kuni eng ko'p", size=9, color=GOLD, bold=True, caps=True)
        y = 3.3
        for key, k in d["peak_cats"]:
            label, colour = cat_meta(key)
            _dot(s, x, y + 0.05, colour, d=0.12)
            _text(s, x + 0.22, y, cw - 0.9, 0.22, label, size=9.5, color=INK2)
            _text(s, x + cw - 0.6, y, 0.6, 0.22, num(k), size=9.5, color=INK, bold=True,
                  align=PP_ALIGN.RIGHT)
            y += 0.32
    best = max(d["daily"], key=lambda z: z["closed"]) if d["daily"] else None
    if best and best["closed"]:
        _card(s, x, 4.5, cw, 1.2)
        _text(s, x + 0.2, 4.62, cw - 0.4, 0.4, f"{num(best['closed'])} ta", size=20,
              color=GREEN, font=HEAD_FONT, bold=True)
        _text(s, x + 0.2, 5.08, cw - 0.4, 0.5, f"eng ko'p yopilgan kun — {day_full(best['date'])}",
              size=9, color=MUTED, line=1.15)

    note = _ai(narr, "daily_note")
    _rect(s, M, 5.84, 8.5, 0.74, fill=INNER, line=LINE, radius=0.08)
    _text(s, M + 0.22, 5.94, 8.1, 0.56, note or NO_AI, size=9.5,
          color=INK2 if note else FAINT, line=1.2)
    n_days = len(d["daily"]) or 1
    _text(s, M, 6.66, 8.5, 0.24,
          f"Kuniga o'rtacha {days1(d['filed'] / n_days)} ta yangi, "
          f"{days1(d['closed'] / n_days)} ta yopildi.", size=8.5, color=FAINT)


# ── slide 10 · what to do ────────────────────────────────────────────────────
def _actions(prs, d: dict, narr, page: int, footer: str):
    s = _page(prs, "Tavsiyalar", "Nima qilish kerak?", page, footer)

    actions = [a for a in _ai_list(narr, "actions", 4) if isinstance(a, dict)]
    if not actions:
        actions = [{"cat": c["key"],
                    "text": (f"{c['label']} — hafta oxirida {num(c['open'])} ta ochiq, "
                             f"{num(c['overdue'])} tasi muddati o'tgan. Muddati o'tganlarni "
                             f"mas'ullar bilan birma-bir ko'rib chiqish.")}
                   for c in sorted(d["categories"], key=lambda c: (-c["overdue"], -c["open"]))[:3]]

    by_key = {c["key"]: c for c in d["categories"]}
    y = 1.36
    for i, a in enumerate(actions):
        key = cat_key(a.get("cat"))
        c = by_key.get(key)
        colour = c["colour"] if c else GOLD
        h = 1.06
        _card(s, M, y, CW, h)
        _rect(s, M, y, 0.055, h, fill=colour)
        _chip(s, M + 0.24, y + 0.2, 0.34, 0.3, str(i + 1), fill=colour, color=WHITE, size=10.5)
        if c:
            _text(s, M + 0.72, y + 0.18, 3.3, 0.24, c["label"], size=10.5, color=INK, bold=True)
            _text(s, M + 0.72, y + 0.42, 3.3, 0.22,
                  f"{num(c['open'])} ochiq · {num(c['overdue'])} muddati o'tgan",
                  size=9, color=BROWN, bold=True)
        _text(s, M + 4.25, y + 0.2, CW - 4.5, 0.72, a.get("text", ""), size=10,
              color=INK2, line=1.22)
        y += h + 0.13

    covered = sum(by_key[k]["open"] for k in {cat_key(a.get("cat")) for a in actions}
                  if k in by_key)
    if covered and d["open_end"]:
        _rect(s, M, y + 0.06, CW, 0.62, fill=INK, radius=0.08)
        _text(s, M + 0.24, y + 0.2, CW - 0.48, 0.36,
              f"Bu choralar {num(covered)} ta ochiq xavotirni — hafta oxiridagi "
              f"ochiqlarning {pct(_share(covered, d['open_end']))} qismini qamrab oladi.",
              size=10.5, color=GOLD2, bold=True)


# ── slide 11 · the conclusion ────────────────────────────────────────────────
def _conclusion(prs, d: dict, narr, page: int, footer: str):
    s = _blank(prs)
    _rect(s, 0, 0, W, H, fill=DARK)
    _rect(s, 0, 0, W, 0.09, fill=GOLD)
    _chrome(s, "Xulosa", "", page, footer, dark=True)

    head = _ai(narr, "conclusion_headline")
    _text(s, M, 1.1, CW - 1.0, 1.42,
          head or (f"{num(d['filed'])} ta yangi xavotir, {num(d['closed'])} tasi yopildi — "
                   f"hafta oxirida {num(d['open_end'])} ta ochiq."),
          size=26, color=WHITE, font=HEAD_FONT, bold=True, line=1.12)

    tcat = [c for c in d["categories"] if c["filed"]][:3]
    worst = d["units"][0] if d["units"] and d["units"][0]["overdue"] else None
    line = (f"Top-3 bo'lim ({' · '.join(c['label'] for c in tcat)}) — yangi xavotirlarning "
            f"{pct(sum(c['share'] for c in tcat))}.") if tcat else ""
    if worst:
        line += (f" Muddati o'tganlari eng ko'p: {short_name(worst['name'], 4.0, 11)} "
                 f"({num(worst['overdue'])}).")
    _text(s, M, 2.62, CW, 0.3, line, size=11, color=ONDARK)

    points = [p for p in _ai_list(narr, "conclusion_points", 3) if isinstance(p, dict)]
    if not points:
        points = [{"title": c["label"], "raw": True,
                   "body": f"{num(c['filed'])} ta yangi, hafta oxirida {num(c['open'])} ta ochiq."}
                  for c in tcat]
    cw = (CW - 2 * 0.2) / 3
    for i, p in enumerate(points):
        x = M + i * (cw + 0.2)
        _rect(s, x, 3.0, cw, 2.5, fill=TILE, radius=0.12)
        _chip(s, x + 0.22, 3.2, 0.34, 0.34, str(i + 1), fill=GOLD2, color=DARK, size=12)
        _text(s, x + 0.22, 3.72, cw - 0.44, 0.5, _title(p), size=13,
              color=WHITE, font=HEAD_FONT, bold=True, line=1.1)
        _text(s, x + 0.22, 4.3, cw - 0.44, 1.05, p.get("body", ""), size=9.5,
              color=ONDARK, line=1.22)

    nxt = d["window"][1] + timedelta(days=7)
    _text(s, M, 5.72, CW, 0.3, f"Keyingi hisobot: {day_label(nxt)} chorshanba yakunida.",
          size=10, color=GOLD2)


# ── slide 12 · how it was counted ────────────────────────────────────────────
def _appendix(prs, d: dict, narr, page: int, footer: str):
    s = _page(prs, "Ilova", "Metodika va ma'lumot sifati", page, footer)
    q = d["quality"]
    b = d["window"][1]
    cw = (CW - 0.2) / 2

    def bullets(x, y, w, items, mark_fn, step):
        for t in items:
            _rect(s, x, y + 0.06, 0.055, 0.2, fill=mark_fn(t))
            _text(s, x + 0.18, y, w - 0.18, step - 0.04, t, size=9, color=MUTED, line=1.2)
            y += step

    _card(s, M, 1.32, cw, 2.64)
    _text(s, M + 0.24, 1.46, cw - 0.48, 0.26, "Raqamlar qanday olindi",
          size=11.5, color=INK, font=HEAD_FONT, bold=True)
    bullets(M + 0.24, 1.8, cw - 0.48, [
        f"Manba: «Xavotirlar» reyestri — ikkala zavod ({d['plants']}), ikkala smena, "
        f"barcha brigadirlar. Sahifadagi filtrlar hisobga olinmaydi.",
        f"Yangi — kiritilgan sanasi {d['period']} oralig'ida bo'lganlar. Yopildi — "
        f"shu oraliqda yopilganlar, qachon kiritilganidan qat'i nazar.",
        f"Ochiq va muddati o'tgan — {day_label(b)} kun oxiriga holat. Muddati o'tgan — "
        f"belgilangan muddati shu kungacha tugagan ochiq xavotir.",
        "Yopilish vaqti — kiritilgan kundan yopilgan kungacha, kunlarda.",
    ], lambda t: GOLD, 0.52)

    _card(s, M + cw + 0.2, 1.32, cw, 2.64)
    _text(s, M + cw + 0.44, 1.46, cw - 0.48, 0.26, "Kimlar kirdi, kimlar yo'q",
          size=11.5, color=INK, font=HEAD_FONT, bold=True)
    right = [
        f"Ishchilar yozgan {num(q['worker_filed'])} ta yangi xavotir hamma hisobda, "
        f"reytingda ham bor. Ishchilarning ismlari faylga chiqarilmagan.",
        f"{num(q['no_deadline_open'])} ta ochiq xavotirga muddat belgilanmagan — ular "
        f"«muddati o'tgan» bo'la olmaydi.",
    ]
    if d["silent_units"]:
        right.append(f"Bu hafta xavotir kiritmagan brigadirlar ({len(d['silent_units'])}): "
                     + " · ".join(d["silent_units"]))
    bullets(M + cw + 0.44, 1.8, cw - 0.48, right,
            lambda t: RED if t.startswith("Bu hafta") else GOLD, 0.68)

    _card(s, M, 4.12, CW, 1.58)
    _text(s, M + 0.24, 4.26, CW - 0.48, 0.26, "Sifat bo'yicha kuzatuvlar",
          size=11.5, color=INK, font=HEAD_FONT, bold=True)
    facts = [
        (f"{num(q['cyrillic'])} ta xavotir kirill alifbosida yozilgan — bu faylda lotinga "
         f"o'girildi, matn o'zgartirilmadi.", MUTED),
        (f"{num(q['thin'])} ta yangi xavotir matni juda qisqa ({THIN_CHARS} belgidan kam) — "
         f"nima haqida ekani tushunarsiz.", RED if q["thin"] else GREEN),
        (f"Taqqoslash davri: {d['prev_period']}. Ikkala davr ham 8 kunlik va bir kunni "
         f"baham ko'radi.", MUTED),
    ]
    y = 4.6
    for t, tone in facts:
        _rect(s, M + 0.24, y + 0.05, 0.055, 0.2, fill=tone)
        _text(s, M + 0.42, y, CW - 0.7, 0.3, t, size=9, color=MUTED, line=1.2)
        y += 0.34

    ai_line = ("Sahifalardagi izoh matnlari sun'iy intellekt tomonidan yozilgan; raqamlar, "
               "jadvallar va grafiklar to'g'ridan-to'g'ri ma'lumotlar bazasidan olingan. "
               "Xavotir matnlari va yechimlar qayta yozilmagan — faqat qo'shtirnoq ichida "
               "keltirilgan." if narr else
               "Bu faylda AI izohlari yo'q — barcha matn ma'lumotlar bazasidan olingan "
               "raqamlardan iborat.")
    _rect(s, M, 5.88, CW, 0.7, fill=INNER, line=LINE, radius=0.08)
    _text(s, M + 0.24, 6.0, CW - 0.48, 0.5, ai_line, size=8.5, color=FAINT, line=1.2)


# ── the deck ─────────────────────────────────────────────────────────────────
_SLIDES = [
    _summary, _recurring, _departments, _top_departments, _units, _oldest,
    _chain, _daily, _actions, _conclusion, _appendix,
]


def build(d: dict, narrative: dict | None = None) -> bytes:
    """The finished .pptx. `narrative` may be None — see `concerns_narrative`."""
    _od._TRIMS.clear()
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(W), Inches(H)

    footer = f"Xavotirlar haftalik tahlili · {d['period']} · {d['plants']}"
    _cover(prs, d, narrative)
    for i, fn in enumerate(_SLIDES, start=2):
        fn(prs, d, narrative, i, footer)

    if _od._TRIMS:
        log.info("CONCERNS DECK built with %d trimmed text block(s)", len(_od._TRIMS))
    # The no-overlap rule's second half, checked on the real week — logged,
    # never raised, the Ojidaniya deck's call: a report with a cosmetic flaw
    # beats no report.
    for problem in deck_text.check_layout(prs, W, H):
        log.warning("CONCERNS DECK layout: %s", problem)

    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


def filename(d: dict) -> str:
    a, b = d["window"]
    return f"Xavotirlar_haftalik_tahlili_{a.strftime('%d-%m')}_{b.strftime('%d-%m-%Y')}.pptx"
