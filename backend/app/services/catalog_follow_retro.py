"""TEMPORARY one-shot (2026-10-26): the catalog follows its cells — for the
cells that ALREADY moved (docs/plan-catalog-follows-cells.md §3 WP6, rulings 3
and 4, decision Q3: applied in the first release, strict guards).

The live follow (`services/catalog_follow`) moves a cell's positions when its
brigadir change lands. Every cell that changed brigadir from 6 October up to
this release left its positions behind in the old unit; the operator's rulings
are that each follows from its OWN move day (ruling 3) and both brigadirs'
figures are recomputed from that day, closed days included (ruling 4). This
pass finds those cells, works out each one's move day from the record, and
hands each to `catalog_follow.follow_group(kind="retro", strict=True)` in a
transaction of its own. Its flag reading `done` is what switches the live sweep
on (`catalog_follow.enabled`).

**Who is considered** (the snapshot is taken once, at the start):

* every cell with an APPLIED brigadir change (`cell_link_events`, `value ≠
  prev`, from `FLOOR`) that has no move record yet;
* every active coded cell whose code has, in ANOTHER unit's current catalog,
  ORPHAN lines — lines no cell of that unit can own: the unit has no living
  cell at the code, or the line's letter is carried by none of them. These are
  the cells that moved before link events existed (12 October).

**When it moved is read off evidence, per day over [FLOOR−7, today]**: the
dated links from `EVENTS_FROM` (12 Oct, when they began — authoritative from
then on); before that the closed live day's roster (`LiveProjection`, the unit
whose copy filed the cell's Verifix code, when exactly one did) and the
attendance majority (Σ `hc_weight` of named rows with hours at the code, a unit
holding more than ⅔). The action register only CORROBORATES (it names units by
spelling; never resolved): a logged unit change near the move day is shown
beside it. A cell that stood in its new unit since before `FLOOR` follows from
`FLOOR` («moved before 6 Oct — followed from 6 Oct») only when attendance older
than the window places it in the old unit.

**Conservative by construction**: every case the evidence cannot settle is
NAMED and skipped, never guessed — `chain` (more than one hop, a round trip),
`from_unknown`, `day_unknown`, `letter_ambiguous` (a landed change included
whose old unit keeps, at the code, positions of a letter no departing or
staying cell carries — another cell's, one that left before link events:
`orphans_of_another_cell`, and that cell is asked about too), `shared_lines`,
`y_has_lines`, `window_cells_moved` (another cell at the code moved into or out
of either unit within the window — its link events from `FLOOR`, its evidence
hops from `FLOOR−7` and from its placement before that; such moves are followed
together or not at all), `stayer_unproven` (a cell the dated links stand at the code in either
unit on h−1 and h with NO evidence of standing there by `h` — cells carry no
creation date, so a cell opened after the move would otherwise count as having
stayed and shape whose the positions are and which letters are free),
`letters_after`, `archived`, `no_code`, `conservation` (the engine's own
strict check). A skipped brigadir change gets a FINAL `named` move record, so
the live sweep never takes it up: a person decides. A `busy` one is recorded
for the sweep to retry — and the sweep, before it retries it, runs these very
guards again (`recheck`): what changed in the meantime is named, never followed
past. What brought nothing to move (`nothing`, `no_code`) is SETTLED: counted
on a line of its own, never in the hand-resolution list or the named count.

**`volume`** — more than `MAX_CELLS` cells, or more than `MAX_ROWS` per-day rows
the follows would actually re-key (`_volume`: each cell's OLD unit's quantity
rows at its own positions' pairs and the unit's pins at the code, from its move
day): NOTHING is applied and nothing is recorded per cell. The flag reads
`volume:<cells>/<rows>` (the sweep stays off: production as before this
release), the pass runs again at every boot, a deploy raising the cap completes
it, and the operator is DMed once per distinct count.

The flag (`dry:<n>`, `done:<moved>/<named>` or `volume:<cells>/<rows>`) and the
RECORD are written in one transaction after the cells; an applying pass that
finished writes ONE Jurnal row (`org.catalog_moved`, `kind` retro — the live
sweep's shape). A pass stopped half-way (a deploy retires the
worker; a failure at its very end) is run again at the next boot: what it
already settled is asked about no second time, and is reported — marked as an
earlier attempt's — beside what the new attempt did. The DM — text + a
workbook — is built from the stored RECORD only, sent once per mode behind its
own flag (`<FLAG>_dm_<mode>`, `startup._send_report_once`): a failed DM is
re-sent at the next boot, the pass is never re-run for it.

Delete this module, `startup.schedule_catalog_follow_retro` /
`_catalog_follow_retro_job` and both entrypoint calls once the flag reads `done`,
the DM has landed AND no retroactive move record waits for a retry (`kind`
retro, `state` busy / failed) — the sweep rechecks those through `recheck`, and
without this module names them `recheck_unavailable` instead of following them.
`catalog_follow.RETRO_FLAG` (the switch) STAYS.
"""
from __future__ import annotations

import json
import logging
import time as _time
from collections import Counter, defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta, timezone
from io import BytesIO
from typing import Optional
from zoneinfo import ZoneInfo

from sqlalchemy import func, text
from sqlalchemy.orm import Session

from app.models import (ActionLog, AppSetting, Attendance, Cell, CellLinkEvent, LiveProjection,
                        Manager, PPCatalogMove, PPDaily, PPLineDaily, PPProduct, PPWorkCenterDaily)
from app.services import catalog_follow as cf
from app.services import catalog_ownership as own
from app.services import pp_catalog
from app.services.cell_lookup import norm_code
from app.services.pp_calc import _field, daily_key

log = logging.getLogger(__name__)
TZ = ZoneInfo("Asia/Tashkent")

FLAG = cf.RETRO_FLAG                        # "catalog_follow_retro_2026_10_26_v1" — the sweep's switch
RECORD = "catalog_follow_retro_2026_10_26"
RETRO_APPLY = True             # decided 2026-10-26: applied in the first release (strict guards)
RETRO_LOCK = 8_140_573_112_026_002
MAX_CELLS, MAX_ROWS = 40, 5000
TRIES = 3                                   # a pass that keeps failing stops; the sweep stays off

FLOOR = cf.FOLLOW_FROM                      # 6 Oct — nothing before it moves (ruling 3)
EVENTS_FROM = date(2026, 10, 12)            # the dated cell links began: authoritative from here
LOOK_BACK = 7                               # evidence from FLOOR−7
PRIOR_DAYS = 90                             # how far back «it stood in X before» is looked for
GAP_MAX = 3                                 # days with no evidence between the last X and first Y day
MAJORITY = 2 / 3

_API = "https://api.telegram.org"
SEND_RETRIES = 3

# The words of the DM and the workbook (Uzbek Latin, the operator's).
REASONS = {
    "chain": "bir necha ko'chish (zanjir yoki qaytish)",
    "from_unknown": "qaysi brigadirdan kelgani aniq emas",
    "day_unknown": "ko'chish kuni aniq emas",
    "letter_ambiguous": "yacheyka harfi aniq emas",
    "shared_lines": "boshqa yacheykalar bilan umumiy pozitsiyalar bor (harf qo'yish kerak)",
    "y_has_lines": "yangi brigadirda shu ish markazida egasiz pozitsiyalar bor",
    "window_cells_moved": "shu ish markazida boshqa yacheyka ham ko'chgan yoki arxivlangan",
    "stayer_unproven": ("shu ish markazida qolgan yacheyka ko'chish kuni u yerda turgani "
                        "dalil bilan tasdiqlanmagan (keyinroq ochilgan bo'lishi mumkin)"),
    "recheck_unavailable": "qayta urinishda tekshiruv o'tkazib bo'lmadi",
    "letters_after": "yacheyka harfi keyinroq o'zgargan",
    "archived": "arxivda",
    "no_code": "SAP kodi yo'q",
    "conservation": "daqiqalar yig'indisi saqlanmadi",
    "volume": "hajm juda katta — hech narsa qo'llanmadi",
    "nothing": "ko'chadigan pozitsiya yo'q",
    "busy": "band — avtomatik qayta uriniladi",
    "failed": "xato — avtomatik qayta uriniladi",
    "shared_left": "umumiy pozitsiyalar qoldi",
    "shared_left_split": "umumiy pozitsiyalar qoldi (ikki brigadirga bo'lingan)",
    "shared_left_archived": "arxivdagi umumiy pozitsiyalar eski brigadirda qoldi",
    "no_free_letter": "bo'sh harf qolmagan",
    "y_register_break": "yangi brigadirda harfsiz yacheyka bor",
    "pin_conflict": "odam soni ikki joyda kiritilgan",
    "code_changed": "SAP kodi o'zgargan",
    "to_no_unit": "brigadirsiz qoldi",
    "from_no_unit": "avval brigadirsiz edi",
    "owned_nothing": "yacheykaga tegishli pozitsiya yo'q",
    "cell_gone": "yacheyka o'chirilgan",
}
SOURCES = {"event": "hodisa", "column": "ustun (12.10 dan)", "roster": "jonli kun ro'yxati",
           "attendance": "davomat", "floor": "6.10 dan (avvalroq ko'chgan)",
           "earlier": "oldingi urinishda"}
# A `why` worth saying beside its kind's words in the DM.
WHYS = {"orphans_of_another_cell": "eski brigadirda boshqa yacheykaning pozitsiyalari ham qolgan"}

# What leaves nothing for a person: the cell brought nothing to move (state
# `nothing` — a change from no unit, a cell that owned nothing there) or has no
# SAP code to carry positions by. Counted on a line of its own, never in the
# hand-resolution list or the flag's named count — a batch of new cells given
# their first brigadir would otherwise crowd out every move a person must decide.
SETTLED = frozenset({"nothing", "no_code", "from_no_unit", "owned_nothing", "cell_gone"})
# What a person must decide, listed first: the engine's strict stops and this
# pass's own guards.
DECIDE = cf.STRICT_STOP | {
    "chain", "from_unknown", "day_unknown", "letter_ambiguous", "shared_lines", "y_has_lines",
    "window_cells_moved", "stayer_unproven", "letters_after", "archived", "conservation",
    "volume", "recheck_unavailable"}


# ── small helpers ───────────────────────────────────────────────────────────

def _days(a: date, b: date) -> list[date]:
    return [a + timedelta(days=i) for i in range((b - a).days + 1)] if a <= b else []


def _iso(v):
    return v.isoformat() if hasattr(v, "isoformat") else v


def _vkey(code) -> str:
    """A Verifix code as the roster, attendance and cells all meet on it —
    normalised, leading zeros dropped (`cell_lookup`'s zero-padded rule)."""
    n = norm_code(code)
    return n.lstrip("0") or n


def _setting(db: Session, key: str) -> Optional[AppSetting]:
    return db.query(AppSetting).filter(AppSetting.key == key).first()


def _put(db: Session, key: str, value: str) -> None:
    row = _setting(db, key)
    if row is None:
        db.add(AppSetting(key=key, value=value))
    else:
        row.value = value


def _flag(db: Session) -> Optional[str]:
    """The flag's value — "" when absent, None when it cannot be read (a box
    before the table: nothing to do)."""
    try:
        v = db.query(AppSetting.value).filter(AppSetting.key == FLAG).scalar()
    except Exception:  # noqa: BLE001
        db.rollback()
        return None
    return v or ""


def _tries(v: str) -> int:
    try:
        return int(v.split(":")[1])
    except (IndexError, ValueError):
        return TRIES


def pending(db: Session) -> bool:
    """Should the pass run: the flag is absent, it reads a rehearsal (`dry`)
    and the moves are to be applied now, an earlier pass failed and has tries
    left, or an earlier pass found too much at once (`volume:`) — asked again
    at every boot, so a later deploy raising `MAX_CELLS` / `MAX_ROWS` completes
    it."""
    v = _flag(db)
    if v is None:
        return False
    if not v:
        return True
    if v.startswith(("dry", "volume:")):
        return RETRO_APPLY
    if v.startswith("failed:"):
        return _tries(v) < TRIES
    return False


def dm_mode(db: Session) -> Optional[str]:
    """Which DM is owed: `done`, `dry`, `failed` once the tries are spent, or
    `volume_<cells>_<rows>` — one per distinct count, so a pass that trips the
    cap again on the same numbers at the next boot says nothing new."""
    v = _flag(db) or ""
    if v.startswith("done"):
        return "done"
    if v.startswith("dry"):
        return "dry"
    if v.startswith("failed:") and _tries(v) >= TRIES:
        return "failed"
    if v.startswith("volume:"):
        return "volume_" + "_".join(p for p in v[len("volume:"):].split("/") if p)
    return None


def dm_flag(mode: str) -> str:
    return f"{FLAG}_dm_{mode}"


# ── the lock ────────────────────────────────────────────────────────────────
# A SESSION lock on a connection of its own, not a transaction lock: every cell
# is its own transaction, and a transaction lock would be gone at the first
# commit — the second app copy of a worker swap could then start a pass beside
# this one.

def _claim():
    from app.database import SessionLocal
    eng = SessionLocal.kw.get("bind")
    if eng is None or eng.dialect.name != "postgresql":
        return True
    try:
        conn = eng.connect()
    except Exception:  # noqa: BLE001
        log.exception("catalog retro: no connection for the lock")
        return None
    try:
        got = bool(conn.execute(text("SELECT pg_try_advisory_lock(:k)"), {"k": RETRO_LOCK}).scalar())
        conn.commit()
    except Exception:  # noqa: BLE001
        log.exception("catalog retro: could not ask for the lock")
        got = False
    if got:
        return conn
    conn.close()
    return None


def _release(conn) -> None:
    if conn is True or conn is None:
        return
    try:
        conn.execute(text("SELECT pg_advisory_unlock(:k)"), {"k": RETRO_LOCK})
        conn.commit()
    except Exception:  # noqa: BLE001
        pass
    finally:
        try:
            conn.close()
        except Exception:  # noqa: BLE001
            pass


# ── the evidence ────────────────────────────────────────────────────────────

def _rosters(db: Session, a: date, b: date, keys: set) -> dict:
    """{(day, Verifix key): {unit}} — the units whose closed live-day copy
    filed the cell's code that day."""
    out: dict = defaultdict(set)
    for mid, day, roster in db.query(LiveProjection.manager_id, LiveProjection.day,
                                     LiveProjection.roster).filter(
            LiveProjection.day >= a, LiveProjection.day <= b).all():
        for item in roster or []:
            if isinstance(item, (list, tuple)) and len(item) >= 2:
                k = _vkey(item[1])
                if k in keys:
                    out[(day, k)].add(int(mid))
    return out


def _attendance(db: Session, a: date, b: date, keys: set) -> dict:
    """{(day, Verifix key): {unit: Σ people}} — named rows with hours at the
    code, a split worker by their share."""
    w = func.coalesce(Attendance.hc_weight, 1.0)
    out: dict = defaultdict(lambda: defaultdict(float))
    for d, code, mid, s in (db.query(Attendance.date, Attendance.verifix_code,
                                     Attendance.manager_id, func.sum(w))
                            .filter(Attendance.date >= a, Attendance.date <= b,
                                    Attendance.verifix_code.isnot(None),
                                    Attendance.worker_name.isnot(None),
                                    Attendance.worker_name != "",
                                    Attendance.hours_worked > 0)
                            .group_by(Attendance.date, Attendance.verifix_code,
                                      Attendance.manager_id).all()):
        k = _vkey(code)
        if k in keys and mid is not None:
            out[(d, k)][int(mid)] += float(s or 0)
    return out


def _majority(weights: dict) -> Optional[tuple[int, float, bool]]:
    """(the biggest unit, its share, whether it holds more than ⅔)."""
    tot = sum(weights.values())
    if tot <= 0:
        return None
    u, w = max(weights.items(), key=lambda kv: (kv[1], -kv[0]))
    share = w / tot
    return u, round(share, 3), share > MAJORITY


def _log_days(db: Session, cell_ids: set) -> tuple[dict, dict]:
    """({cell id: [days]} on which the register logged a change of the cell's
    brigadir — shown beside the evidence, never decided by (the register names
    units by spelling); {cell id: [moments]} at which it logged a change of the
    cell's LETTER by hand — what makes a move's letter unknowable when its
    event kept none of its own (`_decide_event`))."""
    units: dict = defaultdict(list)
    letters: dict = defaultdict(list)
    if not cell_ids:
        return units, letters
    try:
        rows = db.query(ActionLog.target_id, ActionLog.created_at, ActionLog.changes).filter(
            ActionLog.target_kind == "cell", ActionLog.outcome == "done",
            ActionLog.target_id.in_([str(i) for i in cell_ids])).all()
    except Exception:  # noqa: BLE001 — the register is a corroboration, never a must
        db.rollback()
        return units, letters
    for tid, at, changes in rows:
        if not at or not isinstance(changes, list):
            continue
        keys = {ch[0] for ch in changes if isinstance(ch, (list, tuple)) and ch}
        try:
            cid = int(tid)
        except (TypeError, ValueError):
            continue
        at = _aware(at)
        if keys & {"unit", "manager_id", "brigadir"}:
            units[cid].append(at.astimezone(TZ).date())
        if "wc_group" in keys:
            letters[cid].append(at)
    return units, letters


def _aware(v: datetime) -> datetime:
    """A timestamp the driver handed back naive is UTC (the columns are
    timestamptz) — never compared with an aware one as it stands."""
    return v.replace(tzinfo=timezone.utc) if v.tzinfo is None else v


@dataclass
class _Ev:
    """One cell's evidence over the window, day by day."""
    days: list = field(default_factory=list)     # [{d, ev, ro, at, unit, src, conflict}]
    prior: Optional[tuple] = None                 # (day, unit) — the last placement before the window
    log: list = field(default_factory=list)


def _timeline(cell: Cell, window: list[date], ro: dict, at: dict, prior_days: list[date]) -> _Ev:
    from app.services import cell_links
    k = _vkey(cell.verifix_code)
    has_events = bool([e for e in cell_links.events().get("manager_id", {}).get(cell.id) or []
                       if e.value != e.prev])
    ev = _Ev()
    for d in window:
        rs = sorted(ro.get((d, k)) or ())
        maj = _majority(at.get((d, k)) or {})
        e = {"d": d, "ev": None, "ro": rs, "at": [maj[0], maj[1]] if maj else None,
             "unit": None, "src": None, "conflict": False}
        if d >= EVENTS_FROM:
            e["ev"] = cell_links.value_on(cell, "manager_id", d)
            e["unit"], e["src"] = e["ev"], ("event" if has_events else "column")
        else:
            r1 = rs[0] if len(rs) == 1 else None
            a1 = maj[0] if maj and maj[2] else None
            if r1 is not None and a1 is not None and r1 != a1:
                e["conflict"] = True
            elif r1 is not None:
                e["unit"], e["src"] = r1, "roster"
            elif a1 is not None:
                e["unit"], e["src"] = a1, "attendance"
        ev.days.append(e)
    for d in reversed(prior_days):
        maj = _majority(at.get((d, k)) or {})
        if maj and maj[2]:
            ev.prior = (d, maj[0])
            break
    return ev


def _evidence(db: Session, interest: dict, log_ids: set, today: date) -> tuple[dict, dict, dict]:
    """The evidence of every cell in `interest` ({id: Cell}), read once:

    * {cell id: its `_Ev` — day by day over [FLOOR−7, today], the attendance
      placement before that window, the days the register logged a brigadir
      change of it};
    * {cell id: the moments its letter was changed by hand} for `log_ids`;
    * {(Verifix key, unit): the FIRST day a closed live-day roster or a named
      attendance row with hours put the code in the unit} — over
      [FLOOR−7−PRIOR_DAYS, today], what `_stood` asks."""
    a = FLOOR - timedelta(days=LOOK_BACK)
    p = a - timedelta(days=PRIOR_DAYS)
    window = _days(a, today)
    prior_days = _days(p, a - timedelta(days=1))
    keys = {_vkey(o.verifix_code) for o in interest.values() if o.verifix_code}
    ro = _rosters(db, a, today, keys)
    at = _attendance(db, p, today, keys)
    logs, relettered = _log_days(db, log_ids)
    tls: dict = {}
    for oid, o in interest.items():
        tls[oid] = _timeline(o, window, ro, at, prior_days)
        tls[oid].log = sorted(logs.get(oid) or [])
    first: dict = {}

    def seen(k, u, d) -> None:
        if (k, u) not in first or d < first[(k, u)]:
            first[(k, u)] = d

    for (d, k), us in ro.items():
        for u in us:
            seen(k, u, d)
    for (d, k), ws in at.items():
        for u, w in ws.items():
            if w > 0:
                seen(k, u, d)
    return tls, relettered, first


def _stood(o: Cell, unit: int, h: date, first: dict) -> bool:
    """Is there EVIDENCE that cell `o` stood in `unit` on or before `h` — a
    closed live day's roster or a named attendance row at its code filed under
    the unit (`first`, from `_evidence`), or a brigadir change of its own (a
    link event) that put it there by `h`?

    The dated links alone prove nothing for a cell no event moved: they answer
    its CURRENT unit for every day, and a cell carries no creation date — one
    opened after `h` reads as having stood there all along."""
    from app.services import cell_links
    d = first.get((_vkey(o.verifix_code), int(unit)))
    if d is not None and d <= h:
        return True
    return any(e.value != e.prev and e.value == unit and e.eff <= h
               for e in cell_links.events().get("manager_id", {}).get(o.id) or [])


def _runs(days: list) -> list[dict]:
    """The evidence as runs of one unit, oldest first; a day with no answer
    joins no run and breaks none."""
    runs: list = []
    for e in days:
        u = e["unit"]
        if u is None:
            continue
        if runs and runs[-1]["unit"] == u:
            runs[-1]["last"] = e["d"]
        else:
            runs.append({"unit": u, "first": e["d"], "last": e["d"], "src": e["src"]})
    return runs


# ── the candidates ──────────────────────────────────────────────────────────

@dataclass
class _Cand:
    cell: Cell
    code: str
    event: Optional[CellLinkEvent] = None
    src: Optional[int] = None
    dst: Optional[int] = None
    h: Optional[date] = None
    letter_src: Optional[str] = None
    source: Optional[str] = None
    orphan_x: list = field(default_factory=list)
    named: Optional[dict] = None
    corroborated: bool = False
    seen_x: bool = False                 # the evidence ever placed it in the old unit
    quiet: bool = False                  # a bystander at the code: not reported
    owned: int = 0                       # lines of the old unit the cell would take (the guards)
    pairs: set = field(default_factory=set)   # their (qty key, Команда) — what `_volume` counts
    rows: int = 0                        # the per-day rows the follow would re-key (`_volume`)

    def name(self, kind: str, **params) -> None:
        if self.named is None:
            self.named = {"kind": kind, **{k: v for k, v in params.items() if v is not None}}


def _orphans(db: Session, by_code: dict, today: date, units: Optional[list] = None) -> dict:
    """{(unit, code): orphan info} — every unit's current ACTIVE lines at a
    code where cells stand, that no living cell of the unit can own. `units`
    narrows it to those units (the sweep's `recheck` asks about one)."""
    from app.services import cell_archive
    if units is None:
        units = [m for (m,) in db.query(PPProduct.manager_id).filter(
            PPProduct.active.is_(True)).distinct().all()]
    units = [u for u in units if u is not None]
    out: dict = {}
    for x, cat in pp_catalog._current_many(db, units).items():
        at: dict = defaultdict(list)
        for p in cat.active_lines:
            k = norm_code(p.work_center)
            if k and k in by_code:
                at[k].append(p)
        for code, lines in at.items():
            here = [c for c in by_code[code] if c.manager_id == x]
            living = [c for c in here if cell_archive.alive(c, today)]
            dead = [c for c in here if not cell_archive.alive(c, today)]
            carried = {c.wc_group for c in living if c.wc_group}
            orph = [p for p in lines if p.wc_group and p.wc_group not in carried] if living \
                else list(lines)
            if orph:
                out[(x, code)] = {"lines": orph, "has_cells": bool(living),
                                  "letters": sorted({p.wc_group for p in orph if p.wc_group}),
                                  "all_letters": sorted({p.wc_group for p in lines if p.wc_group}),
                                  "dead": [c.verifix_code for c in dead]}
    return out


def _letter_of(ev: CellLinkEvent, cell: Cell) -> Optional[str]:
    """The letter the cell left its old unit with — the sweep's own reading."""
    vfx = ev.vfx if isinstance(ev.vfx, dict) else {}
    if "letter_before" in vfx:
        return vfx["letter_before"] or None
    if "wc_group_was" in vfx:
        return vfx["wc_group_was"] or None
    return cell.wc_group or None


def _left_letters(evs: list, cell_of: dict) -> dict:
    """{(old unit, code): {the letters the cells whose landed change took them
    out of that unit carried}} — what a landed change can account for of the
    positions it left behind."""
    out: dict = defaultdict(set)
    for e in evs:
        cell = cell_of.get(e.cell_id)
        if cell is not None and e.prev is not None:
            out[(e.prev, norm_code(cell.sap_code))].add(_letter_of(e, cell))
    return out


def _foreign(info: Optional[dict], letters: set) -> list:
    """The letters of an old unit's orphaned positions at a code (`_orphans`'
    info) that no cell whose landed change left the unit there carried — a
    letter no stayer carries either (those are no orphans). Such a position is
    ANOTHER cell's: one that left before link events (12 Oct), so it has no
    change of its own. Handed to a change that did land it would go with the
    wrong cell; a lone departer takes every line nobody claims."""
    if not info:
        return []
    return sorted({g for p in info["lines"] if (g := _field(p, "wc_group"))}
                  - {g for g in letters if g})


def _foreign_name(c: _Cand, letters: Optional[list]) -> None:
    """A landed change whose old unit keeps, at the code, orphaned positions of
    a letter it never carried (`_foreign`): whose they are the change cannot
    say — named, never handed to it."""
    if c.named is None and letters and c.src is not None and c.dst is not None:
        c.name("letter_ambiguous", why="orphans_of_another_cell", letters=letters,
               letter=c.letter_src)


def _decide_event(c: _Cand, tl: _Ev, relettered: list) -> None:
    """A landed brigadir change: the move day is the event's own; the evidence
    must show nothing but the old unit before it and the new one from it.

    The letter the cell left with is the event's own record of it. An event
    that landed before the follow existed kept one only when the move CHANGED
    the letter (`wc_group_was`); otherwise the letter is read off the column —
    which a person may have re-lettered by hand in the new unit since, and the
    old unit's positions would then be handed by the wrong letter. A logged
    hand change after the landing makes it unknowable: named."""
    ev = c.event
    x, y, h = ev.prev, ev.value, ev.effective_from
    c.src, c.dst, c.h, c.source = x, y, h, "event"
    c.letter_src = _letter_of(ev, c.cell)
    if x is None or y is None:
        return                                 # the engine records it (from/to no unit)
    before = {r["unit"] for r in _runs([e for e in tl.days if e["d"] < h])}
    after = {r["unit"] for r in _runs([e for e in tl.days if e["d"] >= h])}
    if before - {x} or after - {y} or c.cell.manager_id != y:
        c.name("chain", before=sorted(before), after=sorted(after), now=c.cell.manager_id)
        return
    vfx = ev.vfx if isinstance(ev.vfx, dict) else {}
    landed = ev.applied_at or ev.effective_at
    if ("letter_before" not in vfx and "wc_group_was" not in vfx and landed is not None
            and any(at >= _aware(landed) for at in relettered)):
        c.name("letter_ambiguous", why="letter_changed_after_move",
               letter_now=c.letter_src)


def _decide_orphan(c: _Cand, tl: _Ev, orphans: dict) -> None:
    """A cell with no landing event whose positions sit orphaned in another
    unit: one hop X → Y, its day read off the evidence."""
    xs = sorted(c.orphan_x)
    y = c.cell.manager_id
    c.dst = y
    if len(xs) != 1:
        c.name("from_unknown", units=xs, why="orphans_in_several_units")
        return
    x = c.src = xs[0]
    info = orphans[(x, c.code)]
    runs = _runs(tl.days)
    seq = [r["unit"] for r in runs]
    c.seen_x = x in seq or (tl.prior is not None and tl.prior[1] == x)
    if not runs or y not in seq:
        c.name("day_unknown", why="no_evidence", seen=seq or None)
        return
    if set(seq) - {x, y}:
        c.name("from_unknown" if x not in seq else "chain", seen=seq)
        return
    if seq == [y]:
        first = runs[0]["first"]
        if first >= FLOOR or runs[0]["src"] not in ("roster", "attendance"):
            c.name("day_unknown", first_y=_iso(first), why="no_old_unit_evidence")
            return
        if tl.prior is None or tl.prior[1] != x:
            c.name("from_unknown", prior=[_iso(tl.prior[0]), tl.prior[1]] if tl.prior else None,
                   why="never_seen_in_old_unit")
            return
        c.h, c.source = FLOOR, "floor"
    elif seq == [x, y]:
        lx, fy = runs[0]["last"], runs[1]["first"]
        gap = (fy - lx).days - 1
        if gap > GAP_MAX:
            c.name("day_unknown", last_x=_iso(lx), first_y=_iso(fy), gap=gap)
            return
        if fy < FLOOR:
            c.h, c.source = FLOOR, "floor"
        else:
            c.h, c.source = fy, runs[1]["src"]
    else:
        c.name("chain", seen=seq)
        return
    # The letter it left X with: the one orphan letter while X keeps cells at
    # the code; none when X has no cell there (it owned everything) — unless
    # X's lines carry several letters, a sign two cells left.
    if info["dead"]:
        c.name("archived", x_cells=info["dead"], why="old_unit_has_archived_cell")
    elif info["has_cells"]:
        if len(info["letters"]) != 1:
            c.name("letter_ambiguous", letters=info["letters"])
        else:
            c.letter_src = info["letters"][0]
    elif len(info["all_letters"]) >= 2:
        c.name("letter_ambiguous", letters=info["all_letters"])
    if tl.log and c.h is not None:
        c.corroborated = any(c.h - timedelta(days=3) <= d <= c.h + timedelta(days=1) for d in tl.log)


def _changes(o: Cell, tl: Optional[_Ev], today: date) -> list[tuple]:
    """(day, from, to, what) — every change of a cell's unit or existence the
    pass could follow or be misled by: its link events from FLOOR to today
    (pending ones due by today included), its archive flips from FLOOR, and the
    hops its evidence shows over the WHOLE window — from FLOOR−7, and from the
    placement before it (`prior`) into the window's first run.

    The evidence hops reach back past FLOOR because a cell that left a unit
    shortly before 6 Oct is followed by this very pass from FLOOR
    (`_decide_orphan`), and its positions may still sit in the unit it left:
    a cell at the code moving out of that unit later must not be followed past
    it, one at a time."""
    from app.services import cell_archive, cell_links
    out = []
    for e in cell_links.events().get("manager_id", {}).get(o.id) or []:
        if e.value != e.prev and FLOOR <= e.eff <= today:
            out.append((e.eff, e.prev, e.value, "event"))
    if tl is not None:
        runs = _runs(tl.days)
        if runs and tl.prior is not None and tl.prior[1] != runs[0]["unit"] \
                and runs[0]["first"] <= today:
            out.append((runs[0]["first"], tl.prior[1], runs[0]["unit"], "evidence"))
        for r0, r1 in zip(runs, runs[1:]):
            if r1["first"] <= today:
                out.append((r1["first"], r0["unit"], r1["unit"], "evidence"))
    prev = cell_archive.alive(o, FLOOR - timedelta(days=1))
    for d in _days(FLOOR, today):
        a = cell_archive.alive(o, d)
        if a != prev:
            u = cell_links.value_on(o, "manager_id", d)
            out.append((d, u, u, "archive"))
        prev = a
    return out


# ── the pass ────────────────────────────────────────────────────────────────

def run(db: Session, *, today: Optional[date] = None, now: Optional[datetime] = None) -> dict:
    """One pass, under the retro lock; the flag is read again under it. Each
    cell is its own transaction, then the flag and the RECORD in one. Returns
    the record (or {"state": "skipped" | "locked"})."""
    if not pending(db):
        return {"state": "skipped"}
    conn = _claim()
    if conn is None:
        return {"state": "locked"}
    try:
        db.expire_all()
        if not pending(db):
            return {"state": "skipped"}
        prior = _flag(db) or ""
        try:
            return _run(db, today or _today(), now or datetime.now(timezone.utc))
        except Exception as exc:  # noqa: BLE001 — tried again at the next boot, then stops
            db.rollback()
            log.exception("catalog retro failed")
            n = (_tries(prior) if prior.startswith("failed:") else 0) + 1
            _put(db, FLAG, f"failed:{n}")
            _put(db, RECORD, json.dumps({"result": "failed", "attempt": n,
                                         "at": datetime.now(timezone.utc).isoformat(),
                                         "error": f"{type(exc).__name__}: {exc}"[:800]}))
            db.commit()
            return {"state": "failed", "attempt": n}
    finally:
        _release(conn)


def _today() -> date:
    from app.services import cell_archive
    return cell_archive.today()


def _run(db: Session, today: date, now: datetime) -> dict:
    from app.services import cell_archive, cell_links
    cell_links.forget()
    cell_archive.forget()
    apply = RETRO_APPLY

    cells = db.query(Cell).order_by(Cell.verifix_code, Cell.id).all()
    by_id = {c.id: c for c in cells}
    by_code: dict = defaultdict(list)
    for c in cells:
        k = norm_code(c.sap_code)
        if k:
            by_code[k].append(c)
    units = {m.id: m.name for m in db.query(Manager.id, Manager.name).all()}

    # What an earlier attempt of this pass already settled (one stopped by a
    # deploy, or failed at its very end): those cells are not asked about again
    # — their records stand — and the RECORD reports them, or a re-run that
    # finds nothing left would tell the operator that nothing moved.
    earlier = [i for (i,) in db.query(PPCatalogMove.id).filter(PPCatalogMove.kind == "retro")]

    # ── the snapshot: landed changes with no record, orphaned positions ──
    E = CellLinkEvent
    evs = [e for e in db.query(E).filter(
        E.field == "manager_id", E.applied_at.isnot(None), E.effective_from >= FLOOR,
        E.value.is_distinct_from(E.prev)).order_by(E.cell_id, E.effective_from, E.id).all()
        if not cf._mark(e.vfx)]
    rowed = {r.event_id for r in db.query(PPCatalogMove.event_id).filter(
        PPCatalogMove.event_id.in_([e.id for e in evs]))} if evs else set()
    by_cell_ev: dict = defaultdict(list)
    for e in evs:
        by_cell_ev[e.cell_id].append(e)
    cands: list[_Cand] = []
    chain_rows: list[tuple] = []                 # (event, cell, named) — a cell's several changes
    for cid, es in by_cell_ev.items():
        open_ = [e for e in es if e.id not in rowed]
        if not open_ or cid not in by_id:
            continue
        cell = by_id[cid]
        if len(es) >= 2:
            hops = [[_iso(e.effective_from), e.prev, e.value] for e in es]
            for e in open_:
                chain_rows.append((e, cell, {"kind": "chain", "events": hops}))
            continue
        cands.append(_Cand(cell=cell, code=norm_code(cell.sap_code), event=open_[0]))
    ev_cells = set(by_cell_ev)

    # Event-less: a record already waiting for a retry is the sweep's; one an
    # earlier attempt settled (moved, or named by the engine) stands — asked
    # again it would only write a second record of the same cell.
    taken = {cid for (cid,) in db.query(PPCatalogMove.cell_id).filter(
        PPCatalogMove.event_id.is_(None), PPCatalogMove.kind == "retro").all()}
    orphans = _orphans(db, by_code, today)
    # Positions a landed change left behind are that change's to follow (or
    # to name): no other cell at the code is asked about them — unless some of
    # them carry a letter none of the cells that left with a change carried
    # (`_foreign`). Those are another cell's, one that left before link events;
    # the change is named for them (below), and the cells at the code are
    # asked about the unit's orphans as if no change had landed, so the cell
    # they belong to is named too.
    foreign = {k: _foreign(orphans.get(k), ls) for k, ls in _left_letters(evs, by_id).items()}
    explained = {k for k, f in foreign.items() if not f}
    claim: dict = defaultdict(list)
    for (x, code) in sorted(orphans):
        if (x, code) in explained:
            continue
        for c in by_code[code]:
            if (c.manager_id is not None and c.manager_id != x and c.id not in ev_cells
                    and c.id not in taken and cell_archive.alive(c, today)):
                claim[c.id].append(x)
    for cid, xs in claim.items():
        cell = by_id[cid]
        cands.append(_Cand(cell=cell, code=norm_code(cell.sap_code), orphan_x=xs))

    # ── the evidence, for every cell at a code a candidate stands at ──
    codes = {c.code for c in cands if c.code}
    interest = {o.id: o for k in codes for o in by_code[k]} | {c.cell.id: c.cell for c in cands}
    tls, relettered, first = _evidence(db, interest, {c.cell.id for c in cands}, today)

    # ── decide, then the guards ──
    for c in cands:
        cell = c.cell
        if not c.code:
            c.name("no_code")
            if c.event is not None:
                c.src, c.dst, c.h = c.event.prev, c.event.value, c.event.effective_from
            continue
        if c.event is not None:
            _decide_event(c, tls[cell.id], relettered.get(cell.id) or [])
            _foreign_name(c, foreign.get((c.src, c.code)))
        else:
            _decide_orphan(c, tls[cell.id], orphans)
        if c.named is None and c.h is not None and not cell_archive.alive(cell, c.h):
            c.name("archived", day=_iso(c.h))
        elif c.named is None and not cell_archive.alive(cell, today):
            c.name("archived", day=_iso(today))

    # Several event-less cells claiming one unit's orphans: whose is whose?
    # Only a cell the evidence ever placed in that unit is a claimant — every
    # other cell at the code (a cell that has always stood in the new unit, or
    # one of another unit or shift sharing the work centre) is asked too, and
    # named `from_unknown` for itself; counted here it would hold back the one
    # cell that really came from there.
    claimers: dict = defaultdict(list)
    for c in cands:
        if c.event is None and c.src is not None and c.seen_x:
            claimers[(c.src, c.code)].append(c)
    for _k, cs in claimers.items():
        if len(cs) > 1:
            for c in cs:
                c.name("letter_ambiguous", cells=sorted(o.cell.verifix_code or "" for o in cs),
                       why="several_cells_claim_the_old_unit")
    # The others at that code the evidence never placed there are no news once
    # a cell that did came from there: they are left out of the report (said,
    # they would read as cells whose move is in doubt — they never moved).
    for c in cands:
        if (c.event is None and not c.seen_x and (c.src, c.code) in claimers
                and (c.named or {}).get("kind") == "from_unknown"):
            c.quiet = True

    for c in cands:
        if c.named is not None or c.h is None or c.src is None or c.dst is None:
            continue
        _guards(db, c, by_code, tls, today, first)

    ready = [c for c in cands if c.named is None and c.h is not None]
    # A change from or to no unit, and one whose cell owns nothing in the old
    # unit, moves nothing (the engine records why) — they are no part of the
    # volume, or a batch of new cells given a brigadir would hold every real
    # move back.
    real = [c for c in ready if c.src is not None and c.dst is not None and c.owned]
    volume = _volume(db, real)
    if len(real) > MAX_CELLS or volume > MAX_ROWS:
        if apply:
            # Too much at once: nothing at all is applied and nothing is
            # recorded per cell — not a final `named` row, not the `done` that
            # switches the sweep on. The flag reads `volume:` and the pass runs
            # again at every boot, so a deploy raising the cap completes it;
            # until then production behaves as before this release.
            return _stop_volume(db, today, now, cands, real, volume, units, earlier)
        for c in real:
            c.name("volume", cells=len(real), rows=volume, max_cells=MAX_CELLS, max_rows=MAX_ROWS)
        ready = [c for c in ready if c.named is None]

    # ── apply, one cell at a time ──
    moved, skipped, busy = [], [], []
    for c in sorted(ready, key=lambda c: (c.h, c.cell.verifix_code or "", c.cell.id)):
        out = _follow(db, c, apply, now)
        {"moved": moved, "busy": busy}.get(out["state"], skipped).append(out)
    for c in cands:
        if c.named is not None and not c.quiet:
            skipped.append(_entry(c, state="named", kind=c.named["kind"],
                                  params={k: v for k, v in c.named.items() if k != "kind"}))
    for e, cell, nm in chain_rows:
        skipped.append({"cell": cell.id, "vcode": cell.verifix_code, "code": norm_code(cell.sap_code),
                        "from": e.prev, "to": e.value, "h": _iso(e.effective_from), "event": e.id,
                        "state": "named", "kind": "chain", "params": {"events": nm["events"]}})

    # ── the final rows, the flag and the record, in one transaction ──
    db.expire_all()
    for out in _earlier(db, earlier, by_id):
        {"moved": moved, "busy": busy}.get(out["state"], skipped).append(out)
    if apply:
        for c in cands:
            if c.named is not None and c.event is not None:
                _named_row(db, c.event, c.cell, c.named, c.src, c.dst, c.h, c.letter_src, now)
        for e, cell, nm in chain_rows:
            _named_row(db, e, cell, nm, e.prev, e.value, e.effective_from, _letter_of(e, cell), now)
    # What brought nothing to move is counted apart, never among what a
    # person must decide — and what they must decide is listed first.
    skipped, settled = _split(skipped)
    rec = _record(db, today, now, apply, cands, moved, skipped, busy, tls, units,
                  chained=len(chain_rows), settled=settled)
    _put(db, FLAG, f"done:{len(moved)}/{len(skipped)}" if apply else f"dry:{len(moved)}")
    _put(db, RECORD, json.dumps(rec, default=str, ensure_ascii=False))
    db.commit()
    from app.services import cell_register
    cell_register.after_commit()
    cf._forget_caches()
    if apply:
        _journal(rec)
    return rec


def _stop_volume(db: Session, today: date, now: datetime, cands: list, real: list, volume: int,
                 units: dict, earlier: list) -> dict:
    """The volume guard tripped on an applying pass: the flag
    `volume:<cells>/<rows>` and the RECORD (what would have moved, each cell's
    rows), in one transaction — nothing else is written."""
    rec = {"result": "volume", "mode": "apply", "at": now.isoformat(), "today": _iso(today),
           "floor": _iso(FLOOR), "cells": len(real), "rows": volume,
           "max_cells": MAX_CELLS, "max_rows": MAX_ROWS, "earlier": len(earlier),
           "real": [_entry(c, lines=c.owned, rows=c.rows)
                    for c in sorted(real, key=lambda c: (-c.rows, c.cell.verifix_code or ""))],
           "units": {str(k): v for k, v in units.items()},
           "counts": {"candidates": len(cands), "moved": 0, "skipped": 0, "busy": 0,
                      "volume_cells": len(real), "volume_rows": volume}}
    _put(db, FLAG, f"volume:{len(real)}/{volume}")
    _put(db, RECORD, json.dumps(rec, default=str, ensure_ascii=False))
    db.commit()
    log.warning("catalog retro: %d cells / %d rows over the cap (%d / %d) — nothing applied",
                len(real), volume, MAX_CELLS, MAX_ROWS)
    return rec


def _settled(e: dict) -> bool:
    """An outcome that leaves nothing for a person (`SETTLED`): the engine
    wrote `nothing`, or the pass's own kind is one of them (`no_code`)."""
    return e.get("state") == "nothing" or e.get("kind") in SETTLED


def _settled_kind(e: dict) -> str:
    """Why a settled entry brought nothing — the engine's own word for it where
    it named one («avval brigadirsiz edi»), else `nothing`."""
    k = e.get("kind") or "nothing"
    if k != "nothing":
        return k
    return next((n for n in (e.get("params") or {}).get("named") or [] if n in SETTLED), k)


def _split(entries: list) -> tuple[list, list]:
    """(what a person must decide — `DECIDE` kinds first, each group in its
    order —, what is settled)."""
    decide = [e for e in entries if not _settled(e)]
    decide.sort(key=lambda e: 0 if e.get("kind") in DECIDE else 1)
    return decide, [e for e in entries if _settled(e)]


def _journal(rec: dict) -> None:
    """One Jurnal row for the pass that applied (`catalog_follow._report`'s
    shape, marked `kind: retro`): the tally of outcomes, then at most 40 rows
    in all — «code from→to h: state · lines · named». Never raises."""
    entries = [*(rec.get("moved") or []), *(rec.get("skipped") or []),
               *(rec.get("settled") or []), *(rec.get("busy") or [])]
    if not entries:
        return
    tally = Counter(e.get("state") for e in entries)
    details = [("kind", "retro")] + [(f"catalog_{k}", v) for k, v in sorted(tally.items())]
    lines = []
    for e in entries:
        named = list(e.get("named") or []) if e.get("state") == "moved" else \
            sorted({e.get("kind"), *((e.get("params") or {}).get("named") or [])} - {None, "nothing"})
        lines.append(f"{e.get('vcode') or e.get('cell')} {e.get('from')}→{e.get('to')} {e.get('h')}: "
                     f"{e.get('state')}"
                     + (f" · {e['lines']}" if e.get("lines") else "")
                     + (f" · {', '.join(named)}" if named else "")
                     + (" · earlier" if e.get("earlier") else ""))
    details += [("catalog_cell", x) for x in lines[:max(0, 40 - len(details))]]
    try:
        from app.services import action_log
        action_log.record_system("org", "org.catalog_moved", details=details)
    except Exception:  # noqa: BLE001
        log.exception("catalog retro: action log failed")


def _guards(db: Session, c: _Cand, by_code: dict, tls: dict, today: date, first: dict) -> None:
    """The guards a decided hop must pass before the engine is handed it —
    `run` asks them of every candidate, `recheck` again of a transfer the
    sweep is about to retry. `first` is `_evidence`'s third answer."""
    from app.services import cell_links
    cell, code, x, y, h = c.cell, c.code, c.src, c.dst, c.h
    before = h - timedelta(days=1)
    at_code = {o.id: o for o in by_code.get(code) or []}

    def unproven(unit: int, standing: dict) -> bool:
        # stayer_unproven: the dated links stand a cell at the code in the unit
        # on h−1 and h, and nothing shows it was there by h. Counted as a
        # stayer it would decide whose the positions are and which letters
        # are free; a cell opened since would be counted all the same.
        for oid in standing:
            o = at_code.get(oid) or db.get(Cell, oid)
            if o is None or not _stood(o, unit, h, first):
                c.name("stayer_unproven", other=getattr(o, "verifix_code", None) or oid,
                       unit=unit, day=_iso(h))
                return True
        return False
    # letters_after: a later move already dated this cell's letter.
    later = [eff for eff, _w in cell_links.events().get(cell_links.LETTERS, {}).get(cell.id) or []
             if eff > h]
    if later:
        c.name("letters_after", day=_iso(min(later)))
        return
    # window_cells_moved: another cell at the code moved into or out of either
    # unit (or was archived or restored there) from FLOOR on. Followed one at a
    # time such moves take each other's positions, so neither is followed.
    for o in by_code.get(code) or []:
        if o.id == cell.id:
            continue
        for d, frm, to, what in _changes(o, tls.get(o.id), today):
            if frm in (x, y) or to in (x, y):
                c.name("window_cells_moved", other=o.verifix_code, day=_iso(d), what=what)
                return
    # shared_lines / y_has_lines, on the old unit's catalog at h and today.
    stay_x = own.cells_standing(db, x, code, before, h, exclude={cell.id})
    cats_x = (pp_catalog.at(db, x, h), pp_catalog.current(db, x))
    at_x = [own.lines_at(cat.lines, code) for cat in cats_x]
    rows_x = any(norm_code(w.code) == code for cat in cats_x for w in cat.work_centers)
    if not any(at_x) and not rows_x:
        return                  # X holds nothing at the code: whoever stays there, the
                                # cell brings nothing (the engine records `nothing`)
    if unproven(x, stay_x):
        return
    A = dict(stay_x)
    A[cell.id] = c.letter_src
    mine: list = []
    for lines in at_x:
        hd = own.handed(lines, A, cell.id)
        # Only an ACTIVE shared line stops the move (§1.9: archived lines
        # never block — the engine leaves them in the old unit and says so).
        shared = [p for p in hd.shared if getattr(p, "active", True) is not False]
        if shared and stay_x:
            c.name("shared_lines", count=len(shared),
                   lines=[p.name for p in shared[:10]])
            return
        mine += hd.mine
    c.owned = len({id(p) for p in mine})
    c.pairs = {(daily_key(_field(p, "sap_code"), _field(p, "name")), _field(p, "work_center") or "")
               for p in mine}
    if not mine and not rows_x:
        return                  # it brings nothing: the engine records `nothing`
    # The new unit's stayers decide the arrival's letter, the sibling rule and
    # a join: each must be shown to have been there.
    stay_y = own.cells_standing(db, y, code, before, h, exclude={cell.id})
    if unproven(y, stay_y):
        return
    if not mine:
        return
    # y_has_lines: the new unit's positions at the code that no cell staying
    # there claims — the arrival would take a share of them from `h` (all of
    # them, arriving alone). Unclaimed is `wc_group.share`'s reading: no
    # letter, or a letter no staying cell carries, while two or more cells stay
    # (a lone stayer takes every such line under its own letter, §1.6's
    # sibling rule). A line the moved positions JOIN is no stranger — nor is
    # one that is ALREADY the arrival's own: a cell no link event moved (one
    # moved before 12 Oct) stands in the new unit on h−1 already, and the
    # lines `wc_group.share` hands it alone there (its own letter, or every
    # line while nobody else stays) were typed for it since it came.
    if len(stay_y) == 1:
        return
    carried = {g for g in stay_y.values() if g}
    idents = {own.join_ident(p) for p in mine}
    already = c.event is None and cell_links.value_on(cell, "manager_id", before) == y
    if already and stay_y:
        # Beside other cells it keeps its letter only where the engine lets it
        # (`catalog_follow.build_plan`'s `own_dst`): no other cell row of the
        # unit at the code and no stayer carries it. Otherwise it is re-lettered
        # and the lines under its old letter are no longer its own.
        mine_y = cf._own_letter(cell, y, before)
        already = bool(mine_y) and mine_y not in (cf._row_letters(db, y, code, {cell.id}) | carried)
    A_y = {**stay_y, cell.id: cell_links.letter_on(cell, before)} if already else None
    for cat in (pp_catalog.at(db, y, h), pp_catalog.current(db, y)):
        at_y = own.lines_at(cat.active_lines, code)
        its = {id(p) for p in own.handed(at_y, A_y, cell.id).mine} if already else set()
        extra = [p for p in at_y
                 if not (stay_y and (getattr(p, "wc_group", None) or None) in carried)
                 and own.join_ident(p) not in idents and id(p) not in its]
        if extra:
            c.name("y_has_lines", count=len(extra), lines=[p.name for p in extra[:10]],
                   staying=len(stay_y))
            return


def recheck(db: Session, row: PPCatalogMove, *, today: Optional[date] = None) -> Optional[dict]:
    """The one-shot's OWN guards, run again for a retroactive transfer the live
    sweep is about to retry — a record this pass left `busy` (another writer
    held a unit) or `failed` (`catalog_follow._recheck_retro` calls it). The
    engine's strict stops alone would not see what changed since: another
    cell moving at the code, a position added to the new unit, a cell opened
    beside it, the evidence of the move day…

    The hop is read again exactly as `run` read it — off the landing event, or
    (a cell no event moved) off the old unit's orphaned positions and the
    evidence — and must still be the hop the record names: the same day, the
    same units, the same letter. Then the archive and `_guards`. Returns the
    named condition ({"kind", …params}) that stops it, or None when it may be
    followed (a cell gone, or a change from or to no unit, are the engine's to
    record). Reads only."""
    from app.services import cell_archive
    today = today or _today()
    cell = db.get(Cell, row.cell_id)
    if cell is None:
        return None
    code = norm_code(cell.sap_code)
    if not code:
        return {"kind": "no_code"}
    if row.code and code != row.code:
        return {"kind": "code_changed", "was": row.code, "now": code}
    at_code = [o for o in db.query(Cell).order_by(Cell.verifix_code, Cell.id).all()
               if norm_code(o.sap_code) == code]
    by_code = {code: at_code}
    ev = None
    if row.event_id is not None:
        ev = db.get(CellLinkEvent, row.event_id)
        if ev is None:
            return {"kind": "from_unknown", "why": "event_gone"}
    c = _Cand(cell=cell, code=code, event=ev)
    interest = {o.id: o for o in at_code} | {cell.id: cell}
    tls, relettered, first = _evidence(db, interest, {cell.id}, today)
    # `run`'s snapshot, for this one code: every unit's orphans at it and the
    # landed changes of the cells standing at it.
    orphans = _orphans(db, by_code, today)
    evs = [e for e in db.query(CellLinkEvent).filter(
        CellLinkEvent.field == "manager_id", CellLinkEvent.applied_at.isnot(None),
        CellLinkEvent.effective_from >= FLOOR,
        CellLinkEvent.value.is_distinct_from(CellLinkEvent.prev),
        CellLinkEvent.cell_id.in_([o.id for o in at_code])).all() if not cf._mark(e.vfx)]
    foreign = {k: _foreign(orphans.get(k), ls)
               for k, ls in _left_letters(evs, {o.id: o for o in at_code}).items()}
    if ev is not None:
        _decide_event(c, tls[cell.id], relettered.get(cell.id) or [])
        _foreign_name(c, foreign.get((c.src, code)))
    else:
        x = row.from_unit
        # Less the positions a landed change left behind (that change's to
        # follow — `run`'s `explained`); the cell claims those of every unit
        # but its own. Orphans in a second unit make the old unit unknowable,
        # as they did for `run`.
        evented = {e.cell_id for e in evs}
        explained = {k for k, f in foreign.items() if not f}
        xs = sorted(u for (u, k) in orphans
                    if k == code and (u, k) not in explained and u != cell.manager_id)
        if x not in xs:
            c.src, c.dst = x, cell.manager_id
            c.name("from_unknown", why="orphans_gone")
        else:
            c.orphan_x = xs
            _decide_orphan(c, tls[cell.id], orphans)
            if c.named is None:
                # Another cell the evidence places in the old unit claims the
                # same orphans: whose is whose? (`run`'s claimers, for one.)
                rowed = {cid for (cid,) in db.query(PPCatalogMove.cell_id).filter(
                    PPCatalogMove.event_id.is_(None), PPCatalogMove.kind == "retro").all()}
                for o in at_code:
                    if (o.id == cell.id or o.id in evented or o.id in rowed
                            or o.manager_id in (None, x) or not cell_archive.alive(o, today)):
                        continue
                    oc = _Cand(cell=o, code=code, orphan_x=[x])
                    _decide_orphan(oc, tls[o.id], orphans)
                    if oc.seen_x:
                        c.name("letter_ambiguous", why="several_cells_claim_the_old_unit",
                               cells=sorted([cell.verifix_code or "", o.verifix_code or ""]))
                        break
    if c.named is None and (c.h, c.src, c.dst) != (row.home_from, row.from_unit, row.to_unit):
        c.name("day_unknown", why="evidence_changed",
               was=[_iso(row.home_from), row.from_unit, row.to_unit],
               now=[_iso(c.h), c.src, c.dst])
    if c.named is None and (c.letter_src or None) != (row.letter_from or None):
        c.name("letter_ambiguous", why="letter_changed_since",
               was=row.letter_from, now=c.letter_src)
    if c.named is None and c.h is not None and not cell_archive.alive(cell, c.h):
        c.name("archived", day=_iso(c.h))
    elif c.named is None and not cell_archive.alive(cell, today):
        c.name("archived", day=_iso(today))
    if c.named is None and c.h is not None and c.src is not None and c.dst is not None:
        _guards(db, c, by_code, tls, today, first)
    return c.named


def _volume(db: Session, ready: list) -> int:
    """How many per-day rows the follows would re-key — what each one actually
    touches: its OLD unit's `pp_line_daily` and `pp_daily` rows at the (qty key,
    Команда) pairs of the lines the guards found to be the cell's own
    (`_Cand.pairs`), and that unit's typed pins at the code, from the move day
    on. The new unit's own positions and the stayers' are never re-keyed, so
    they are no part of it. Each candidate's share is left on `_Cand.rows`."""
    if not ready:
        return 0
    for c in ready:
        c.rows = 0
    by_src: dict = defaultdict(list)
    for c in ready:
        by_src[c.src].append(c)
    units = sorted(by_src)
    lo = min(c.h for c in ready)
    for model, key in ((PPLineDaily, PPLineDaily.qty_key), (PPDaily, PPDaily.sap_code)):
        for mid, qk, wc, d, k in (db.query(model.manager_id, key, model.work_center, model.date,
                                           func.count())
                                  .filter(model.manager_id.in_(units), model.date >= lo)
                                  .group_by(model.manager_id, key, model.work_center, model.date)
                                  .all()):
            for c in by_src.get(int(mid)) or []:
                if d >= c.h and (qk, wc or "") in c.pairs:
                    c.rows += int(k)
    W = PPWorkCenterDaily
    for mid, wc, d, k in (db.query(W.manager_id, W.work_center, W.date, func.count())
                          .filter(W.manager_id.in_(units), W.date >= lo)
                          .group_by(W.manager_id, W.work_center, W.date).all()):
        for c in by_src.get(int(mid)) or []:
            if d >= c.h and norm_code(wc) == c.code:
                c.rows += int(k)
    return sum(c.rows for c in ready)


def _entry(c: _Cand, **kw) -> dict:
    out = {"cell": c.cell.id, "vcode": c.cell.verifix_code, "code": c.code, "from": c.src,
           "to": c.dst, "h": _iso(c.h), "event": c.event.id if c.event is not None else None,
           "source": c.source, "letter_from": c.letter_src, "corroborated": c.corroborated}
    out.update(kw)
    return out


def _follow(db: Session, c: _Cand, apply: bool, now: datetime) -> dict:
    """One cell through the engine, in its own transaction."""
    from app.services import cell_archive, cell_links
    t = cf.Transfer(cell_id=c.cell.id, event_id=c.event.id if c.event is not None else None,
                    src=c.src, dst=c.dst, home_from=c.h, letter_src=c.letter_src)
    try:
        res = cf.follow_group(db, [t], kind="retro", strict=True, dry=not apply, now=now)
    except Exception as exc:  # noqa: BLE001 — the sweep retries it
        db.rollback()
        log.exception("catalog retro: follow of cell %s failed", c.cell.id)
        if apply:
            cf._record_retry([t], "failed", repr(exc), now, kind="retro")
        return _entry(c, state="busy", kind="failed", params={"error": repr(exc)[:300]})
    finally:
        cell_links.forget()
        cell_archive.forget()
    if res.state == "busy":
        db.rollback()
        if apply:
            cf._record_retry([t], "busy", None, now, kind="retro")
        return _entry(c, state="busy", kind="busy")
    row = next((r for r in res.rows if r.cell_id == c.cell.id), None)
    named = list((res.report or {}).get("named") or [])
    kinds = sorted({n.get("kind") for n in named if n.get("kind")})
    if apply:
        try:
            db.commit()
        except Exception as exc:  # noqa: BLE001 — one cell's commit never stops the pass
            db.rollback()
            log.exception("catalog retro: commit of cell %s failed", c.cell.id)
            cf._record_retry([t], "failed", repr(exc), now, kind="retro")
            return _entry(c, state="busy", kind="failed", params={"error": repr(exc)[:300]})
        finally:
            cell_links.forget()
            cell_archive.forget()
        # Readers of the units the cell moved between see the move at once,
        # not at the end of a pass of many cells.
        from app.services import cell_register
        cell_register.after_commit()
        cf._forget_caches()
    else:
        db.rollback()
    cell_links.forget()
    cell_archive.forget()
    db.expire_all()
    lines = list(getattr(row, "line_ids", None) or []) if row is not None else []
    wcs = list(getattr(row, "wc_ids", None) or []) if row is not None else []
    if res.state == "named":
        stop = next((k for k in kinds if k in cf.STRICT_STOP), kinds[0] if kinds else "named")
        return _entry(c, state="named", kind=stop, params={"named": kinds})
    if not (lines or wcs):
        return _entry(c, state="nothing", kind="nothing", params={"named": kinds})
    entries = list(getattr(row, "lines", None) or [])
    return _entry(c, state="moved", S=_iso(res.S), lines=len(lines),
                  joined=sum(1 for e in entries if e.get("joined_into")),
                  wc=len(wcs), wc_dropped=len(getattr(row, "wc_dropped", None) or []),
                  letter_to=getattr(row, "letter_to", None),
                  letters=dict(getattr(row, "letters", None) or {}),
                  rekey=dict(getattr(row, "rekey", None) or {}), named=kinds,
                  impact=res.impact or {})


def _earlier(db: Session, ids: list, by_id: dict) -> list[dict]:
    """The records an earlier attempt of this pass wrote, as the RECORD's own
    entries (marked `earlier`) — read as they stand now, a retry the sweep may
    yet make included."""
    if not ids:
        return []
    out = []
    for r in db.query(PPCatalogMove).filter(PPCatalogMove.id.in_(ids)).order_by(PPCatalogMove.id):
        cell = by_id.get(r.cell_id)
        kinds = sorted({n.get("kind") for n in r.named or [] if isinstance(n, dict) and n.get("kind")})
        base = {"cell": r.cell_id, "vcode": getattr(cell, "verifix_code", None) or str(r.cell_id),
                "code": r.code, "from": r.from_unit, "to": r.to_unit, "h": _iso(r.home_from),
                "event": r.event_id, "source": "earlier", "letter_from": r.letter_from,
                "corroborated": False, "earlier": True}
        if r.state == "moved":
            out.append({**base, "state": "moved", "S": _iso(r.storage_from),
                        "lines": len(r.line_ids or []),
                        "joined": sum(1 for e in r.lines or [] if (e or {}).get("joined_into")),
                        "wc": len(r.wc_ids or []), "wc_dropped": len(r.wc_dropped or []),
                        "letter_to": r.letter_to, "letters": dict(r.letters or {}),
                        "rekey": dict(r.rekey or {}), "named": kinds,
                        "impact": dict(r.impact or {})})
        elif r.state in cf.RETRY:
            out.append({**base, "state": "busy", "kind": r.state,
                        "params": {"error": (r.error or "")[:300]} if r.error else {}})
        else:
            stop = next((k for k in kinds if k in cf.STRICT_STOP), kinds[0] if kinds else r.state)
            out.append({**base, "state": r.state if r.state == "nothing" else "named",
                        "kind": "nothing" if r.state == "nothing" else stop,
                        "params": {"named": kinds}})
    return out


def _named_row(db: Session, ev: CellLinkEvent, cell: Cell, named: dict, src, dst, h, letter,
               now: datetime) -> None:
    """A brigadir change the pass left alone: a FINAL record, so the live
    sweep never takes it up — a person decides (the DM names it)."""
    if db.query(PPCatalogMove.id).filter(PPCatalogMove.event_id == ev.id).first() is not None:
        return
    t = cf.Transfer(cell_id=cell.id, event_id=ev.id, src=src, dst=dst,
                    home_from=h or ev.effective_from, letter_src=letter)
    r = cf._fill_row(PPCatalogMove(event_id=ev.id, cell_id=cell.id), t, kind="retro",
                     state="named", code=norm_code(cell.sap_code), by_name=ev.by_name,
                     named=[named])
    db.add(r)
    ev.vfx = {**(ev.vfx or {}), "catalog_follow": {
        "state": "named", "storage_from": None, "lines": 0, "letter_to": None,
        "at": now.isoformat()}}


def _record(db: Session, today: date, now: datetime, apply: bool, cands: list, moved: list,
            skipped: list, busy: list, tls: dict, units: dict, chained: int = 0,
            settled: Optional[list] = None) -> dict:
    """Everything the DM and the workbook print — they read only this.
    `skipped` is what a person must decide, `settled` what brought nothing to
    move (`_split`)."""
    settled = settled or []
    impact = []
    for m in moved:
        for u, days in sorted((m.pop("impact", None) or {}).items(), key=lambda kv: str(kv[0])):
            for d, v in sorted(days.items()):
                impact.append({"unit": int(u), "day": d, "cell": m["vcode"],
                               "plan": v.get("plan"), "actual": v.get("actual"),
                               "people": v.get("people")})
    evidence = {}
    for c in cands:
        tl = tls.get(c.cell.id)
        if tl is None or c.quiet:
            continue
        evidence[str(c.cell.id)] = {
            "vcode": c.cell.verifix_code,
            "days": [[_iso(e["d"]), e["ev"], e["ro"], e["at"], e["unit"], e["src"], e["conflict"]]
                     for e in tl.days if e["ev"] is not None or e["ro"] or e["at"] or e["conflict"]],
            "prior": [_iso(tl.prior[0]), tl.prior[1]] if tl.prior else None,
            "log": [_iso(d) for d in tl.log]}
    # Every cell moved with no landing event has no link history — whatever
    # day it moved on (one read off the column from 12 Oct included): a
    # per-cell reader still shows it in its new unit before `h`.
    pre12 = sorted(m["vcode"] or str(m["cell"]) for m in moved if m.get("event") is None)
    return {"result": "done" if apply else "dry", "mode": "apply" if apply else "dry",
            "at": now.isoformat(), "today": _iso(today), "floor": _iso(FLOOR),
            "events_from": _iso(EVENTS_FROM), "moved": moved, "skipped": skipped,
            "settled": settled, "busy": busy,
            "impact": impact, "evidence": evidence, "pre12": pre12,
            "units": {str(k): v for k, v in units.items()},
            "counts": {"candidates": len(cands) + chained,
                       "moved": len(moved), "skipped": len(skipped), "settled": len(settled),
                       "busy": len(busy),
                       "reasons": dict(Counter(e["kind"] for e in skipped)),
                       "settled_reasons": dict(Counter(_settled_kind(e) for e in settled))}}


# ── the DM ──────────────────────────────────────────────────────────────────

def _d(v) -> str:
    """«2026-10-08» → «08.10»."""
    try:
        d = date.fromisoformat(str(v)[:10])
    except ValueError:
        return str(v or "—")
    return d.strftime("%d.%m")


def _uname(rec: dict, u) -> str:
    if u is None:
        return "—"
    return (rec.get("units") or {}).get(str(u)) or f"#{u}"


def _volume_text(rec: dict) -> str:
    """The DM of a pass the volume guard stopped: nothing applied, the counts
    against the cap, the cells, and what a person does about it."""
    real = rec.get("real") or []
    L = ["📦 Pozitsiyalarni yacheykalar bilan ko'chirish (bir martalik): HECH NARSA QO'LLANMADI — "
         "hajm chegaradan katta.",
         f"Ko'chadigan yacheykalar: {rec.get('cells')} (chegara {rec.get('max_cells')}) · "
         f"qayta yoziladigan kunlik qatorlar: {rec.get('rows')} (chegara {rec.get('max_rows')})."]
    if real:
        L.append("")
        for m in real[:30]:
            L.append(f"• {m.get('vcode')}: {_uname(rec, m.get('from'))} → {_uname(rec, m.get('to'))}, "
                     f"{_d(m.get('h'))} dan — {m.get('lines', 0)} pozitsiya, {m.get('rows', 0)} qator")
        if len(real) > 30:
            L.append(f"… yana {len(real) - 30}")
    if rec.get("earlier"):
        L.append("")
        L.append(f"Oldingi urinishda yozilgani o'z joyida qoladi: {rec['earlier']} ta yozuv.")
    L.append("")
    L.append("Chegarani (MAX_CELLS / MAX_ROWS, catalog_follow_retro.py) kattalashtirib qayta "
             "joylash kerak — o'tish har yuklanishda qayta ishga tushadi va chegara yetarli "
             "bo'lganda o'zi bajariladi.")
    L.append("Avtomatik kuzatuv o'chiq qoldi: ko'chgan yacheykalarning pozitsiyalari hozircha "
             "eski brigadirda turibdi, avvalgidek.")
    return "\n".join(L)


def _reason(e: dict) -> str:
    """An entry's reason in the DM's words — its kind's, and the `why` beside
    it where that says more (`WHYS`)."""
    p = e.get("params") or {}
    out = REASONS.get(e.get("kind"), e.get("kind"))
    if p.get("why") in WHYS:
        out += f" — {WHYS[p['why']]}"
        if p.get("letters"):
            out += f" ({', '.join(p['letters'])})"
    return out


def text_of(rec: dict) -> str:
    """The DM's text, out of the RECORD alone."""
    if rec.get("result") == "failed":
        return ("📦 Pozitsiyalarni yacheykalar bilan ko'chirish (bir martalik) "
                f"{rec.get('attempt')} marta muvaffaqiyatsiz bo'ldi: {rec.get('error')}\n"
                "Avtomatik kuzatuv o'chiq qoldi — dasturchiga murojaat qiling.")
    if rec.get("result") == "volume":
        return _volume_text(rec)
    dry = rec.get("result") == "dry"
    moved, busy = rec.get("moved") or [], rec.get("busy") or []
    skipped, settled = _split(rec.get("skipped") or [])
    settled = (rec.get("settled") or []) + settled
    if not (moved or skipped or settled or busy):
        return ("📦 Pozitsiyalar yacheykalar bilan: 6-oktabrdan beri ko'chiriladigan narsa "
                "topilmadi." + ("" if dry else " Avtomatik kuzatuv yoqildi."))
    done_word = "Rejada" if dry else "Ko'chirildi"
    L = [("📦 SINOV (hech narsa yozilmadi): pozitsiyalar" if dry else "📦 Pozitsiyalar")
         + " yacheykalar bilan ko'chdi — bir martalik o'tish, 6-oktabrdan",
         f"{done_word}: {len(moved)} yacheyka · Qoldirildi: {len(skipped)} · "
         f"Qayta urinadi: {len(busy)}"
         + (f" · Ko'chadigan narsa yo'q: {len(settled)}" if settled else "")]
    if moved:
        L.append("")
        L.append("✅ " + ("Ko'chiriladi:" if dry else "Ko'chirildi:"))
        for m in moved[:30]:
            extra = f", {m['joined']} birlashtirildi" if m.get("joined") else ""
            s = f" (saqlash {_d(m['S'])} dan)" if m.get("S") and m.get("S") != m.get("h") else ""
            L.append(f"• {m.get('vcode')}: {_uname(rec, m.get('from'))} → {_uname(rec, m.get('to'))}, "
                     f"{_d(m.get('h'))} dan{s} — {m.get('lines', 0)} pozitsiya{extra} "
                     f"[{SOURCES.get(m.get('source'), m.get('source') or '—')}]")
        if len(moved) > 30:
            L.append(f"… yana {len(moved) - 30}")
    if skipped:
        L.append("")
        L.append("⚠️ Qoldirildi (qo'lda hal qiling):")
        counts = Counter(e["kind"] for e in skipped)
        for kind, n in sorted(counts.items(), key=lambda kv: (kv[0] not in DECIDE, -kv[1], kv[0])):
            L.append(f"• {REASONS.get(kind, kind)}: {n}")
        for e in skipped[:25]:
            L.append(f"  – {e.get('vcode')} ({_uname(rec, e.get('from'))} → {_uname(rec, e.get('to'))}): "
                     f"{_reason(e)}")
        if len(skipped) > 25:
            L.append(f"  … yana {len(skipped) - 25}")
    if busy:
        L.append("")
        L.append(f"⏳ Band edi, avtomatik qayta uriniladi: {', '.join(str(b.get('vcode')) for b in busy[:20])}")
    if settled:
        L.append("")
        why = Counter(_settled_kind(e) for e in settled).most_common()
        L.append(f"✔️ Hal bo'ldi, ko'chadigan narsa yo'q: {len(settled)} — "
                 + ", ".join(f"{REASONS.get(k, k)}: {n}" for k, n in why))
    if rec.get("pre12"):
        L.append("")
        L.append("ℹ️ Hodisasiz ko'chgan (havola tarixi yo'q): " + ", ".join(rec["pre12"][:30])
                 + " — bu yacheykalarning ko'chish kunidan oldingi o'z ko'rsatkichlari "
                   "yangi brigadirda ko'rinishda qoladi.")
    L.append("")
    L.append("Sinov: avtomatik kuzatuv hali o'chiq." if dry
             else "Avtomatik kuzatuv yoqildi: endi har bir yangi ko'chishda pozitsiyalar o'zi ko'chadi.")
    L.append("Batafsil — ilova qilingan Excel faylda.")
    return "\n".join(L)


def workbook(rec: dict) -> bytes:
    """«Ko'chirildi» (or «Rejada»), «Qoldirildi», «Dalil», «Ta'sir» — out of
    the RECORD alone."""
    from openpyxl import Workbook
    from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    head = Font(bold=True, color="FFFFFF")
    fill = PatternFill("solid", fgColor="C8973F")

    def clean(v):
        if isinstance(v, (dict, list)):
            v = json.dumps(v, ensure_ascii=False, default=str)
        return ILLEGAL_CHARACTERS_RE.sub("", v) if isinstance(v, str) else v

    def sheet(ws, title, header, data, widths):
        ws.title = title
        ws.append(header)
        for c in ws[1]:
            c.font, c.fill = head, fill
            c.alignment = Alignment(vertical="center", wrap_text=True)
        for row in data:
            ws.append([clean(v) for v in row])
        for i, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = w
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions

    dry = rec.get("result") == "dry"
    un = lambda u: _uname(rec, u)  # noqa: E731
    moved = [[m.get("vcode"), m.get("code"), un(m.get("from")), un(m.get("to")), m.get("h"), m.get("S"),
              SOURCES.get(m.get("source"), m.get("source")), m.get("lines"), m.get("joined"),
              f"{m.get('wc', 0)} ko'chdi, {m.get('wc_dropped', 0)} olib tashlandi",
              f"{m.get('letter_from') or '—'} → {m.get('letter_to') or '—'}",
              m.get("letters"), m.get("rekey"), ", ".join(m.get("named") or []),
              "ha" if m.get("corroborated") else ""]
             for m in rec.get("moved") or []]
    sheet(wb.active, "Rejada" if dry else "Ko'chirildi",
          ["Yacheyka", "SAP kod", "Qayerdan", "Qayerga", "Ko'chish kuni", "Saqlash kuni",
           "Kun manbasi", "Pozitsiyalar", "Birlashtirilgan", "Ish markazi qatori", "Harf",
           "Harflar tarixi", "Qayta kalitlash", "Eslatmalar", "Jurnalda tasdiq"],
          moved, [10, 12, 26, 26, 13, 13, 22, 12, 14, 24, 10, 24, 30, 30, 14])
    # Every entry the pass left unmoved — the settled ones last, by the word
    # the engine had for them.
    skipped = [[e.get("vcode"), e.get("code"), un(e.get("from")), un(e.get("to")), e.get("h"),
                _reason({**e, "kind": k}), k, e.get("params"),
                e.get("event"), SOURCES.get(e.get("source"), e.get("source") or "")]
               for e in (rec.get("skipped") or []) + (rec.get("busy") or [])
               + (rec.get("settled") or [])
               for k in [_settled_kind(e) if _settled(e) else e.get("kind")]]
    sheet(wb.create_sheet(), "Qoldirildi",
          ["Yacheyka", "SAP kod", "Qayerdan", "Qayerga", "Kun", "Sabab", "Kod", "Tafsilot",
           "Hodisa", "Kun manbasi"], skipped, [10, 12, 26, 26, 12, 44, 18, 60, 10, 22])
    ev_rows = []
    for _cid, ev in sorted((rec.get("evidence") or {}).items(), key=lambda kv: kv[1].get("vcode") or ""):
        for d, e, rs, a, u, src, conflict in ev.get("days") or []:
            ev_rows.append([ev.get("vcode"), d, un(e) if e is not None else "",
                            ", ".join(un(x) for x in rs or []),
                            f"{un(a[0])} ({round(a[1] * 100)}%)" if a else "",
                            un(u) if u is not None else ("ziddiyat" if conflict else ""),
                            SOURCES.get(src, src or "")])
        if ev.get("prior"):
            ev_rows.append([ev.get("vcode"), ev["prior"][0], "", "", un(ev["prior"][1]),
                            un(ev["prior"][1]), "oyna oldidan oxirgi davomat"])
        for d in ev.get("log") or []:
            ev_rows.append([ev.get("vcode"), d, "", "", "", "", "jurnal: brigadir o'zgartirilgan"])
    sheet(wb.create_sheet(), "Dalil",
          ["Yacheyka", "Kun", "Hodisa (havola)", "Jonli kun ro'yxati", "Davomat (ko'pchilik)",
           "Xulosa", "Manba"], ev_rows, [10, 12, 26, 30, 30, 26, 30])
    imp_rows = []
    for i in rec.get("impact") or []:
        plan, act, ppl = i.get("plan") or [None, None], i.get("actual") or [None, None], \
            i.get("people") or [None, None]
        imp_rows.append([un(i.get("unit")), i.get("day"), i.get("cell"), plan[0], plan[1],
                         act[0], act[1], ppl[0], ppl[1],
                         "yacheyka odamlari kiritilmagan" if ppl[1] is None else ""])
    sheet(wb.create_sheet(), "Ta'sir",
          ["Brigadir", "Kun", "Yacheyka", "Reja daq. (oldin)", "Reja daq. (keyin)",
           "Fakt daq. (oldin)", "Fakt daq. (keyin)", "Odamlar (oldin)", "Odamlar (keyin)", "Belgi"],
          imp_rows, [26, 12, 10, 16, 16, 16, 16, 14, 14, 30])
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _post(method: str, data: dict, files: Optional[dict] = None) -> None:
    import requests
    from app.config import settings
    last = ""
    for attempt in range(1, SEND_RETRIES + 1):
        wait = 0
        try:
            r = requests.post(f"{_API}/bot{settings.telegram_bot_token}/{method}",
                              data=data, files=files, timeout=180)
            body = r.json()
            if body.get("ok"):
                return
            last = body.get("description") or f"HTTP {r.status_code}"
            wait = int(((body.get("parameters") or {}).get("retry_after")) or 0)
        except Exception as exc:  # noqa: BLE001
            last = type(exc).__name__
        if attempt < SEND_RETRIES:
            _time.sleep(max(wait, 3 * attempt))
    raise RuntimeError(f"{method} failed: {last}"[:300])


def sender(db: Session, chat_id: int, *_window) -> int:
    """`startup._send_report_once`'s sender: the stored RECORD as text (+ the
    workbook when anything was considered). Never runs the pass. Returns the
    messages sent."""
    row = _setting(db, RECORD)
    if row is None or not row.value:
        raise RuntimeError("catalog retro: no record to send")
    rec = json.loads(row.value)
    body = text_of(rec)
    print("[startup] catalog follow retro:\n" + body)
    n = 0
    if rec.get("moved") or rec.get("skipped") or rec.get("busy") or rec.get("settled"):
        stamp = str(rec.get("at") or "")[:16].replace(":", "-").replace("T", "_")
        _post("sendDocument",
              {"chat_id": chat_id,
               "caption": ("Pozitsiyalar yacheykalar bilan — "
                           + ("sinov (hech narsa yozilmadi)" if rec.get("result") == "dry"
                              else "bir martalik ko'chirish") + f", {_d(rec.get('today'))}")},
              files={"document": (f"catalog-follow-retro-{stamp}.xlsx", workbook(rec),
                                  "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})
        n += 1
    chunks, cur = [], ""
    for line in body.split("\n"):
        if len(cur) + len(line) + 1 > 3900:
            chunks.append(cur)
            cur = ""
        cur += line + "\n"
    chunks.append(cur)
    for ch in chunks:
        _post("sendMessage", {"chat_id": chat_id, "text": ch})
        n += 1
    return n


def send(chat_id: int) -> int:
    """Text + workbook from the stored RECORD (never re-runs the pass)."""
    from app.database import SessionLocal
    with SessionLocal() as db:
        return sender(db, chat_id)
