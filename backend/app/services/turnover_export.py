"""«Kadrlar qo'nimsizligi» as a workbook — the page's month, five tabs.

A pure formatter over `turnover.month_payload` (+ the working lists of
`turnover.people_of`): it never re-derives a figure, so the file and the page
can never state two answers. The WORDS come from the client in the viewer's
language (`labels`), with Uzbek as the floor so a tab open on an older bundle
still gets a readable file. The house report style is `quality_export`'s.

    Umumiy          banner · scope · KPI cards · the formula · the score bands
                    · what counts toward nobody
    Liderlar        one row per leader — the KPI itself, in the screen's order
    Yacheykalar     one row per cell
    Ketganlar       every person who left, with where they were counted
    Ishlayotganlar  every person counted as working at the month's end
"""
from __future__ import annotations

from datetime import date, datetime
from io import BytesIO
from typing import Any, Optional

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font
from openpyxl.utils import get_column_letter

from app.services.quality_export import (
    BAND, BOX, BRAND, FONT, GREEN, INK, INK_FAINT, INK_SOFT, LEFT, NUM, PANEL, RED, AMBER, SLATE, TINT,
    _banner, _block, _fill, _head_row, _kpi_cards, _meta_strip, _section, _sheet, _side,
)
from app.translit import transliterate, transliterate_text

LABELS_UZ = {
    "sheetSummary": "Umumiy", "sheetLeaders": "Liderlar", "sheetCells": "Yacheykalar",
    "sheetLeft": "Ketganlar", "sheetWorking": "Ishlayotganlar",
    "left": "Ketganlar", "working": "Oy oxirida ishlayotganlar", "rate": "Qo'nimsizlik (yillik)",
    "leaders": "Baholangan liderlar", "formula": "Qanday hisoblanadi",
    "bands": "Ball shkalasi", "bandRate": "Qo'nimsizlik", "bandScore": "Ball", "bandPoints": "KPI bali",
    "nobody": "Hech kimga hisoblanmaganlar",
    "colN": "№", "colLeader": "Lider", "colBrigadir": "Brigadir", "colCells": "Yacheykalar",
    "colCell": "Yacheyka", "colShift": "Smena", "colWorking": "Ishlayotganlar", "colLeft": "Ketganlar",
    "colRate": "Qo'nimsizlik, %", "colScore": "Ball", "colPoints": "KPI bali",
    "colName": "F.I.Sh.", "colJob": "Lavozim", "colHired": "Ishga kirgan", "colLeftOn": "Ketgan sana",
    "colDays": "Ishlagan kuni", "colReason": "Sabab", "colNote": "Izoh", "colCounted": "Kimga hisoblandi",
    "countedLeader": "Liderga", "countedNoLeader": "Lidersiz yacheyka", "countedNoCell": "Yacheyka /cells'da yo'q",
    "noLeader": "Lider yo'q", "noRate": "—",
    "formulaLines": [],
    "nobodyNoLeader": "Lidersiz yacheykalardan ketganlar",
    "nobodyNoCell": "/cells'da yo'q kodlar",
    "nobodyOther": "Yacheyka bo'lmagan bo'limlardan ketganlar (KPI'ga kirmaydi)",
}


HAIR_B = Border(bottom=_side())


def _l(p: dict, key: str) -> Any:
    return (p.get("labels") or {}).get(key) or LABELS_UZ.get(key) or key


def _name(p: dict, v: Optional[str]) -> str:
    return transliterate(v, p.get("lang") or "uz") if v else "—"


def _text(p: dict, v: Optional[str]) -> str:
    return transliterate_text(v, p.get("lang") or "uz") if v else "—"


def _d(v: Optional[str]):
    if not v:
        return None
    try:
        return date.fromisoformat(v[:10])
    except ValueError:
        return v


def _score_color(score: Optional[int]) -> str:
    if score is None:
        return SLATE
    return GREEN if score >= 4 else AMBER if score == 3 else RED


def _row(ws, r: int, c1: int, values: list, *, zebra: int, fmts: Optional[dict] = None,
         tints: Optional[dict] = None, centers: tuple = ()) -> None:
    ws.row_dimensions[r].height = 16
    bg = _fill(PANEL if zebra % 2 == 0 else BAND)
    for j, v in enumerate(values):
        tone = (tints or {}).get(j)
        fill = _fill(TINT.get(tone, BAND)) if tone else bg
        font = Font(name=FONT, size=9.5, bold=bool(tone), color=tone or INK)
        align = Alignment(horizontal="center", vertical="center") if (j in centers or tone) else LEFT
        if isinstance(v, (int, float)) and not isinstance(v, bool):
            align = Alignment(horizontal="right", vertical="center", indent=1)
        _block(ws, r, c1 + j, r, c1 + j, "—" if v is None or v == "" else v, fill=fill, border=BOX,
               align=align, font=font, fmt=(fmts or {}).get(j))


def _table(ws, head_row: int, c1: int, labels: list[str], rows: list[tuple]) -> int:
    _head_row(ws, head_row, c1, labels, height=28)
    r = head_row + 1
    first = r
    for i, (vals, kw) in enumerate(rows):
        _row(ws, r, c1, vals, zebra=i, **kw)
        r += 1
    c2 = c1 + len(labels) - 1
    ws.auto_filter.ref = f"{get_column_letter(c1)}{head_row}:{get_column_letter(c2)}{max(first, r - 1)}"
    ws.freeze_panes = ws.cell(first, c1 + 1)
    ws.print_title_rows = f"{head_row}:{head_row}"
    return r


def _summary(wb: Workbook, p: dict) -> None:
    pay = p["payload"]
    t = pay.get("totals") or {}
    ws = _sheet(wb, _l(p, "sheetSummary"), {c: 13.0 for c in range(2, 14)})
    C1, C2 = 2, 13
    row = _banner(ws, 2, C1, C2, p.get("title") or "", p.get("subtitle") or "")
    row = _meta_strip(ws, row, C1, C2, p.get("meta") or [])
    by = t.get("by_score") or {}
    kpis = [
        {"value": t.get("left", 0), "label": _l(p, "left"), "color": RED,
         "hint": p.get("hints", {}).get("left", "")},
        {"value": t.get("working", 0), "label": _l(p, "working"), "color": BRAND,
         "hint": p.get("hints", {}).get("working", "")},
        {"value": t.get("rate"), "label": _l(p, "rate"), "color": AMBER, "fmt": '0.0"%"',
         "hint": p.get("hints", {}).get("rate", "")},
        {"value": t.get("scored", 0), "label": _l(p, "leaders"), "color": GREEN,
         "hint": " · ".join(f"{s}: {by.get(str(s), 0)}" for s in (5, 4, 3, 2, 1))},
    ]
    row = _kpi_cards(ws, row, C1, kpis)

    row = _section(ws, row, C1, C2, _l(p, "formula"))
    for line in _l(p, "formulaLines") or []:
        ws.row_dimensions[row].height = 30
        _block(ws, row, C1, row, C2, line, font=Font(name=FONT, size=10, color=INK),
               align=Alignment(horizontal="left", vertical="center", wrap_text=True, indent=1))
        row += 1
    row += 1

    row = _section(ws, row, C1, C2, _l(p, "bands"))
    _head_row(ws, row, C1, [_l(p, "bandRate"), _l(p, "bandScore"), _l(p, "bandPoints")], span=3,
              first_span=3, height=22)
    row += 1
    rule = pay.get("rule") or {}
    prev = None
    bands = sorted(rule.get("bands") or [], key=lambda b: float(b["max"]))
    steps = [(b["max"], b["score"]) for b in bands] + [(None, rule.get("worst", 1))]
    for mx, sc in steps:
        if mx is None:
            rng = f"> {prev:g}%"
        elif prev is None:
            rng = f"≤ {float(mx):g}%"
        else:
            rng = f"{float(prev) + 0.1:g} – {float(mx):g}%"
        tone = _score_color(int(sc))
        for k, (val, fmt) in enumerate(((rng, None), (int(sc), NUM),
                                         (round(int(sc) * float(rule.get("weight", 0.2)), 2), "0.00"))):
            c = C1 + k * 3
            _block(ws, row, c, row, c + 2, val, fill=_fill(TINT.get(tone, BAND) if k == 1 else PANEL), border=BOX,
                   font=Font(name=FONT, size=10, bold=k == 1, color=tone if k == 1 else INK),
                   align=Alignment(horizontal="center", vertical="center"), fmt=fmt)
        if mx is not None:
            prev = float(mx)
        row += 1
    row += 1

    nobody = []
    if t.get("left_no_leader"):
        nobody.append((_l(p, "nobodyNoLeader"), t["left_no_leader"]))
    for u in pay.get("unregistered") or []:
        nobody.append((f"{_l(p, 'nobodyNoCell')}: {u.get('code') or '—'} {_text(p, u.get('unit'))}",
                       f"{u.get('left', 0)} / {u.get('working', 0)}"))
    if t.get("other_left"):
        nobody.append((_l(p, "nobodyOther"), t["other_left"]))
    if nobody:
        row = _section(ws, row, C1, C2, _l(p, "nobody"))
        for text, n in nobody:
            ws.row_dimensions[row].height = 18
            _block(ws, row, C1, row, C2 - 2, text, border=HAIR_B, font=Font(name=FONT, size=9.5, color=INK_SOFT))
            _block(ws, row, C2 - 1, row, C2, n, border=HAIR_B, font=Font(name=FONT, size=10, bold=True, color=INK),
                   align=Alignment(horizontal="right", vertical="center", indent=1))
            row += 1



def _leaders(wb: Workbook, p: dict) -> None:
    pay = p["payload"]
    rows = list(pay.get("leaders") or [])
    order = p.get("leader_order") or []
    if order:
        pos = {lid: i for i, lid in enumerate(order)}
        rows.sort(key=lambda x: pos.get(x["id"], len(pos)))
    ws = _sheet(wb, _l(p, "sheetLeaders"), {2: 6, 3: 34, 4: 28, 5: 18, 6: 14, 7: 12, 8: 14, 9: 9, 10: 11},
                landscape=True)
    C1 = 2
    labels = [_l(p, k) for k in ("colN", "colLeader", "colBrigadir", "colCells", "colWorking", "colLeft",
                                 "colRate", "colScore", "colPoints")]
    row = _banner(ws, 2, C1, C1 + len(labels) - 1, p.get("title") or "", p.get("subtitle") or "")
    out = []
    for i, x in enumerate(rows, 1):
        brig = ", ".join(_name(p, b.get("name")) for b in x.get("brigadirs") or []) or "—"
        cells = ", ".join(c.get("code") or "" for c in x.get("cells") or [])
        tone = _score_color(x.get("score")) if x.get("score") is not None else None
        out.append(([i, _name(p, x.get("name")), brig, cells, x.get("working", 0), x.get("left", 0),
                     x.get("rate"), x.get("score"), x.get("points")],
                    {"fmts": {6: '0.0"%"', 8: "0.00"}, "tints": {7: tone} if tone else {}, "centers": (0, 3)}))
    _table(ws, row, C1, labels, out)


def _cells(wb: Workbook, p: dict) -> None:
    pay = p["payload"]
    ws = _sheet(wb, _l(p, "sheetCells"), {2: 12, 3: 32, 4: 28, 5: 8, 6: 14, 7: 12, 8: 14}, landscape=True)
    C1 = 2
    labels = [_l(p, k) for k in ("colCell", "colLeader", "colBrigadir", "colShift", "colWorking", "colLeft",
                                 "colRate")]
    row = _banner(ws, 2, C1, C1 + len(labels) - 1, p.get("title") or "", p.get("subtitle") or "")
    out = []
    for c in pay.get("cells") or []:
        out.append(([c.get("code"), _name(p, c.get("leader")) if c.get("leader") else _l(p, "noLeader"),
                     _name(p, c.get("brigadir")), c.get("shift"), c.get("working", 0), c.get("left", 0),
                     c.get("rate")],
                    {"fmts": {6: '0.0"%"'}, "centers": (0, 3)}))
    _table(ws, row, C1, labels, out)


def _left(wb: Workbook, p: dict) -> None:
    pay = p["payload"]
    ws = _sheet(wb, _l(p, "sheetLeft"), {2: 34, 3: 10, 4: 28, 5: 24, 6: 22, 7: 12, 8: 12, 9: 10, 10: 30, 11: 30,
                                        12: 22}, landscape=True)
    C1 = 2
    labels = [_l(p, k) for k in ("colName", "colCell", "colLeader", "colBrigadir", "colJob", "colHired",
                                 "colLeftOn", "colDays", "colReason", "colNote", "colCounted")]
    row = _banner(ws, 2, C1, C1 + len(labels) - 1, p.get("title") or "", p.get("subtitle") or "")
    counted = {"leader": _l(p, "countedLeader"), "no_leader": _l(p, "countedNoLeader"),
               "no_cell": _l(p, "countedNoCell")}
    out = []
    for x in pay.get("leavers") or []:
        out.append(([_name(p, x.get("name")), x.get("code"), _name(p, x.get("leader")) if x.get("leader") else None,
                     _name(p, x.get("brigadir")) if x.get("brigadir") else None, _text(p, x.get("job")),
                     _d(x.get("hired")), _d(x.get("left")), x.get("days"), _text(p, x.get("reason")),
                     _text(p, x.get("note")), counted.get(x.get("counted"), x.get("counted"))],
                    {"fmts": {5: "DD.MM.YYYY", 6: "DD.MM.YYYY"}, "centers": (1, 5, 6),
                     "tints": {} if x.get("counted") == "leader" else {10: SLATE}}))
    _table(ws, row, C1, labels, out)


def _working(wb: Workbook, p: dict) -> None:
    pay = p["payload"]
    names = p.get("working") or {}
    ws = _sheet(wb, _l(p, "sheetWorking"), {2: 12, 3: 30, 4: 36, 5: 26, 6: 13}, landscape=False)
    C1 = 2
    labels = [_l(p, k) for k in ("colCell", "colLeader", "colName", "colJob", "colHired")]
    row = _banner(ws, 2, C1, C1 + len(labels) - 1, p.get("title") or "", p.get("subtitle") or "")
    out = []
    for c in pay.get("cells") or []:
        for person in names.get(c.get("key")) or []:
            out.append(([c.get("code"), _name(p, c.get("leader")) if c.get("leader") else _l(p, "noLeader"),
                         _name(p, person.get("name")), _text(p, person.get("job")), _d(person.get("hired"))],
                        {"fmts": {4: "DD.MM.YYYY"}, "centers": (0, 4)}))
    _table(ws, row, C1, labels, out)


def build_turnover_workbook(p: dict) -> BytesIO:
    wb = Workbook()
    wb.remove(wb.active)
    _summary(wb, p)
    _leaders(wb, p)
    _cells(wb, p)
    _left(wb, p)
    _working(wb, p)
    wb.properties.title = p.get("title") or "Turnover"
    wb.properties.creator = "Safia IMS"
    wb.properties.created = datetime.now()
    bio = BytesIO()
    wb.save(bio)
    bio.seek(0)
    return bio
