"""Every objection waiting on an ADMIN ruling — DMed once, with everything a
ruling needs.

The operator, 2026-10-20, looking at «Lider nazorati» → «Norozliklar» →
«Adminlarda 116»: «Send me these 116 objections as a report on startup file.
Include every single detail of that task that you need to check if the
objection is true or not.» Production is not readable from a cloud session, so
the report is a boot job like every other one-off in `startup.py`, and it lands
in the operator's chat.

WHICH rows: every objection whose `status` is «admin» — passed up by the
brigadir (or filed straight to the admin stage by a brigadir or an admin) and
not ruled yet — whatever its checklist day. The screen the operator looked at
was narrowed to 1–19 Oct (`PAGE_FROM` / `PAGE_TO`); a row outside it is
included and marked, so nothing waiting on an admin is missed.

WHAT each carries — every figure off the platform's own functions, never a
second spelling:

  * the card the queue and the appeal page draw (`routers.leaders._dispute_items`):
    who, unit, cell, shift, day, task, which judge refused it (the AI · an
    automatic check · an admin's «not done»), the AI's flags, its reason in four
    languages, the clocks it read, the date verdict and the window;
  * the task as resolved for that leader, day and cell: name, what to
    photograph, the instruction the leader reads, the AI criteria
    (`leader_ai.criteria_for`, the one-process rule included), photos required,
    weight, proof kind, deadline, the date rule — plus the AI prompt rebuilt
    from today's config (`ai_prompt.txt`) and every criteria edit made after the
    verdict;
  * the leader's answer: done or not, their comment, when it was saved and
    submitted, the day's state and score;
  * every proof photo — the archive channel's own bytes — with the server's
    capture time for in-app camera shots, and which of them the AI saw;
  * the example photos the AI was shown in front of the proofs;
  * the whole appeal chat, every entry in order, with every attached file;
  * for an automatic check: what it read at its hour (the ledger), when the
    leader acted (`auto_check_timing`, #8 / #9), the «Ish grafigi» lists at the
    check and now (#11), the unit's people / plan / fact saves (#1).

FILES, to the operator's chat:

  1. a message — how many, and what follows;
  2. `objections-admin-<date>.xlsx` — one row per objection, to skim;
  3. ZIP parts, at most `PART_BYTES` each. EVERY part carries README.md,
     objections.md (every objection in full, in one file), objections.json (the
     same as data) and the workbook. Each objection's folder — objection.md,
     ai_prompt.txt, its proof photos and its chat's files — is whole inside ONE
     part (unless it alone is bigger than a part); the example photos
     (`examples/`) ride in the first part. Unzip every part into ONE folder;
  4. a closing message — what went out, and anything that could not be fetched.

It READS, and writes nothing but its flag (`startup._send_report_once`) and its
progress row. A restart RESUMES: an objection already shipped in a sent part is
not sent again, and the next part is numbered on.

Temporary: delete this module, `startup.report_admin_objections` /
`_admin_objections_job` and the call in BOTH entrypoints once the files have
landed — a call left behind imports a deleted module at boot, and a failed boot
rolls the deploy back.
"""
from __future__ import annotations

import json
import logging
import os
import re
import shutil
import tempfile
import threading
import time
import zipfile
import zlib
from concurrent.futures import ThreadPoolExecutor
from datetime import date as _date, datetime, timedelta
from io import BytesIO
from zoneinfo import ZoneInfo

import requests
from sqlalchemy.orm import Session

from app.config import settings
from app.database import SessionLocal
from app.models import (
    ActionLog, AppSetting, Cell, Factory, LeaderAiDispute, LeaderAiReview,
    LeaderAutoCheck, LeaderTaskDay, LeaderTaskDef, LeaderTaskEntry,
    LeaderTaskExample, LeaderTaskLeaderSetting, LeaderTaskMedia,
    LeaderTaskPhoto, LeaderTaskSetting, Manager, RoleProfile,
)

log = logging.getLogger(__name__)

TZ = ZoneInfo("Asia/Tashkent")              # the plant's wall clock

# The period selected on the operator's screen when they asked. Rows outside it
# still wait on an admin and are included — marked, never dropped.
PAGE_FROM = "2026-10-01"
PAGE_TO = "2026-10-19"

PROGRESS_KEY = "admin_objections_report_2026_10_20_v1:progress"

PART_BYTES = 45 * 1024 * 1024               # sendDocument refuses above 50 MB
MAX_PARTS = 30                               # a ceiling against a runaway, said if hit
PART_PAUSE_S = 3
SEND_RETRIES = 3
FETCH_WORKERS = 6
FETCH_RETRIES = 4
FETCH_BACKOFF_S = 2
CHAT_FILE_MAX = 20 * 1024 * 1024            # the most the bot API hands back
TIMELINE_MAX = 80                           # unit saves listed for a #1 objection
TEXT_MAX = 4000                             # free text is cut, never dropped
ERRORS_NAMED = 15

_API = "https://api.telegram.org"
_tls = threading.local()

_WEEKDAYS = ("Monday", "Tuesday", "Wednesday", "Thursday", "Friday",
             "Saturday", "Sunday")


# ── the words ────────────────────────────────────────────────────────────────
# English, because the operator asked in English and means to read it with
# Claude; every text a PERSON wrote (the leader, the brigadir, the chat) and
# every reason the AI wrote stays exactly as it was typed.

KIND_TEXT = {
    "ai": "the AI (it refused the proof photos)",
    "auto": "an automatic check (no photo — the platform read its own data at a fixed hour)",
    "admin": "an admin's hand ruling «not done»",
}

FLAG_TEXT = {
    "unreadable": "unreadable — the AI could not read the photo",
    "off_topic": "off topic — the photo is not about this task",
    "not_proven": "not proven — the photo does not show the work done by the criteria",
    "no_date": "no date — no readable date/time on the photo (required by the date rule)",
    "date_mismatch": "date mismatch — the photo's date/time is outside the allowed window",
}

MODE_TEXT = {
    "strict": "strict — the photo's DAY and HOUR must both fall inside the window",
    "date_only": "date only — the DAY must match (the hour is not judged); a date shown inside the app/document counts",
    "time_only": "time only — the HOUR must be inside the window (the day is not judged)",
    "not_checked": "not checked — neither the day nor the hour is judged",
}

PROOF_TEXT = {
    "camera": "in-app camera — the time on each photo is the SERVER's, the leader cannot change it",
    "screenshot": "photos/screenshots sent to the bot — the AI reads the date/time off the picture",
    "auto": "automatic — the platform answers it, no photo",
}

CHAT_KIND = {
    "filed": "FILED the objection", "message": "message",
    "sup_rejected": "brigadir REFUSED", "uplifted": "brigadir PASSED IT UP to the admins",
    "approved": "admin APPROVED", "rejected": "admin REFUSED",
    "undone": "a ruling was UNDONE (reopened)", "invited": "people were invited",
}

CHECK_TEXT = {
    "plan_staffing": "#1 — a plan on the leader's cell AND that cell's people («Bugungi fakt») typed by the hour",
    "concerns": "#8 — at least one concern written (by the leader, or by a worker of their cell) between 00:00 and the hour",
    "plan_pct": "#9 — «Bajarish %» of one of the leader's work centres at or above the pass mark by the hour",
    "staff_list": "#11 — every worker on the cell's «Ish grafigi» list marked «Keladi» / «Kelmaydi» by the end of the shift",
}

TIMING_TEXT = {
    "on_time": "ON TIME — the data was on the page at the hour (the check should have passed)",
    "late": "LATE — it arrived after the hour",
    "never": "NEVER — not reached even now",
    "unexplained": "unexplained — the replay cannot place when the mark was crossed",
    "no_plan": "no plan was on the page for the leader's work centres",
}

SAVE_TEXT = {
    "production.staffing_set": "«Odamlar soni» saved (people / staffing)",
    "production.override_set": "ПЛАН/ФАКТ typed on a position",
    "production.wc_override_set": "work-centre figure typed",
}

CODE_TEXT = {
    "ok": "passed",
    "no_plan": "no plan was on the page for the leader's cells at the hour",
    "no_staffing": "the people («Bugungi fakt») were not typed at the hour",
    "no_concern": "no concern had been written by the hour",
    "under_target": "the plan completion was below the pass mark at the hour",
    "no_sap_code": "the leader's cell has no SAP code (a register fault)",
    "started_late": "the checklist was opened after the check hour (legacy rule)",
    "not_checked": "the check could not run — a platform fault",
    "no_data": "the check could not read its data",
    "not_filled": "not every worker on «Ish grafigi» was marked",
}


# ── plain values ─────────────────────────────────────────────────────────────

def _aware(v):
    if v is None:
        return None
    if isinstance(v, str):
        try:
            v = datetime.fromisoformat(v)
        except ValueError:
            return None
    if isinstance(v, datetime):
        # Every timestamptz column comes back aware. The naive ones are the
        # Google-Form submit times, which are the plant's wall clock.
        return v if v.tzinfo else v.replace(tzinfo=TZ)
    return None


def _iso(v) -> str | None:
    """An instant as ISO on the plant's clock; a non-instant unchanged."""
    a = _aware(v)
    if a is None:
        return v if isinstance(v, str) else None
    return a.astimezone(TZ).isoformat(timespec="seconds")


def _t(v) -> str:
    """«dd.mm.yyyy HH:MM» on the plant's clock, «—» for nothing."""
    a = _aware(v)
    if a is None:
        return str(v) if v else "—"
    return a.astimezone(TZ).strftime("%d.%m.%Y %H:%M")


def _ts(v) -> str:
    a = _aware(v)
    if a is None:
        return str(v) if v else "—"
    return a.astimezone(TZ).strftime("%d.%m.%Y %H:%M:%S")


def _dmy(day: str | None) -> str:
    try:
        d = _date.fromisoformat(str(day)[:10])
    except (TypeError, ValueError):
        return str(day or "—")
    return f"{d:%d.%m.%Y} ({_WEEKDAYS[d.weekday()]})"


def _cut(v):
    return v[:TEXT_MAX] if isinstance(v, str) else v


def _why(exc: Exception) -> str:
    return _clean(f"{type(exc).__name__}: {exc}")[:300]


def _clean(text: str) -> str:
    """Telegram's URLs carry the bot token; never let it into a file or chat."""
    token = settings.telegram_bot_token or ""
    return text.replace(token, "***") if token else text


_SLUG_DROP = re.compile(r"[^\w\-.]", re.UNICODE)


def _slug(v, fallback: str = "-", n: int = 40) -> str:
    s = _SLUG_DROP.sub("", str(v or "").strip().replace(" ", "_"))
    return s[:n] or fallback


def _plain(v):
    if isinstance(v, datetime):
        return _iso(v)
    if isinstance(v, _date):
        return v.isoformat()
    if isinstance(v, (list, tuple)):
        return [_plain(x) for x in v]
    if isinstance(v, dict):
        return {str(k): _plain(x) for k, x in v.items()}
    if hasattr(v, "_asdict"):
        return _plain(v._asdict())
    try:
        json.dumps(v)
        return v
    except TypeError:
        return str(v)


def _mode(checked, dayed, timed) -> str:
    if not checked or not (dayed or timed):
        return "not_checked"
    if dayed and timed:
        return "strict"
    return "date_only" if dayed else "time_only"


def _ext(data: bytes | None, mime: str | None = None, name: str | None = None) -> str:
    """The file's real type off its first bytes; its name, then its declared
    type, when the bytes do not say."""
    b = data or b""
    if b[:3] == b"\xff\xd8\xff":
        return "jpg"
    if b[:8] == b"\x89PNG\r\n\x1a\n":
        return "png"
    if b[:4] == b"RIFF" and b[8:12] == b"WEBP":
        return "webp"
    if b[:6] in (b"GIF87a", b"GIF89a"):
        return "gif"
    if b[4:12] in (b"ftypheic", b"ftypheix", b"ftypmif1", b"ftypmsf1"):
        return "heic"
    if b[:4] == b"%PDF":
        return "pdf"
    if name and "." in name:
        e = _slug(name.rsplit(".", 1)[1].lower(), "", 8)
        if e:
            return e
    m = (mime or "").lower()
    for k, e in (("jpeg", "jpg"), ("png", "png"), ("webp", "webp"), ("gif", "gif"),
                 ("pdf", "pdf"), ("mp4", "mp4"), ("quicktime", "mov")):
        if k in m:
            return e
    return "bin"


def _sentinel(raw: str | None) -> str | None:
    """A stored reason in words: the two platform sentinels expanded, anything
    a person typed returned as typed."""
    s = str(raw or "")
    if s.startswith("__auto__|"):
        parts = s.split("|")
        hh = parts[1] if len(parts) > 1 else "?"
        code = parts[2] if len(parts) > 2 else "?"
        return f"[automatic check at {hh}: {CODE_TEXT.get(code, code)}]"
    if s.startswith("__missed__|"):
        return f"[not submitted before {s.split('|', 1)[1]} — closed by the deadline]"
    return s or None


# ── the collection ───────────────────────────────────────────────────────────

def collect(db: Session) -> dict:
    """{"report": …, "downloads": […], "examples": […]}, plain values only.

    One objection that cannot be read costs that objection's own detail and
    nothing else: each is assembled under its own try, and what failed is
    written into it (`errors`)."""
    from app.routers.leaders import _dispute_items
    from app.services import leader_dispute

    now = datetime.now(TZ)
    rows = (db.query(LeaderAiDispute)
            .filter(LeaderAiDispute.status == leader_dispute.ADMIN)
            .order_by(LeaderAiDispute.date, LeaderAiDispute.id).all())
    ids = [d.id for d in rows]

    # The card the page draws, for every row at once; one by one if the batch
    # trips on a single row.
    payload = {"role": "admin", "sub": "0", "full_name": "report"}
    cards: dict[int, dict] = {}
    try:
        for it in _dispute_items(db, payload, rows):
            cards[int(it["id"])] = _plain(it)
    except Exception:
        db.rollback()
        log.warning("admin objections report: batch card failed", exc_info=True)
        for i in ids:
            d = db.query(LeaderAiDispute).filter_by(id=i).first()
            if d is None:
                continue
            try:
                for it in _dispute_items(db, payload, [d]):
                    cards[int(it["id"])] = _plain(it)
            except Exception as exc:
                db.rollback()
                cards[i] = {"error": _why(exc)}

    ctx = _Lookups(db)
    objections: list[dict] = []
    downloads: list[dict] = []
    examples: dict[int, dict] = {}
    for i in ids:
        d = db.query(LeaderAiDispute).filter_by(id=i).first()
        if d is None or d.status != leader_dispute.ADMIN:
            continue                        # ruled while this ran
        try:
            o, dl = _one(db, d, cards.get(i) or {}, ctx, examples)
        except Exception as exc:
            db.rollback()
            log.exception("admin objections report: objection %s", i)
            o, dl = _bare(db, i, cards.get(i) or {}, exc, ctx), []
        objections.append(o)
        downloads.extend(dl)

    objections.sort(key=lambda o: (o["date"] or "", o["who"].get("unit") or "",
                                   o["who"].get("leader") or "",
                                   o["task"].get("id") or 0, o["id"]))
    for n, o in enumerate(objections, 1):
        o["seq"] = n

    def _count(key):
        out: dict = {}
        for o in objections:
            k = key(o)
            out[k] = out.get(k, 0) + 1
        return dict(sorted(out.items(), key=lambda kv: (-kv[1], str(kv[0]))))

    report = {
        "generated_at": now.isoformat(timespec="seconds"),
        "what": "Every objection to a rejected checklist task that is waiting on an "
                "ADMIN ruling (status «admin»), whatever its checklist day.",
        "page_period": {"from": PAGE_FROM, "to": PAGE_TO,
                        "note": "the period selected on the operator's screen — rows "
                                "outside it are marked in_page_period=false"},
        "counts": {
            "objections": len(objections),
            "in_page_period": sum(1 for o in objections if o["in_page_period"]),
            "outside_page_period": sum(1 for o in objections if not o["in_page_period"]),
            "by_judge": _count(lambda o: o["kind"]),
            "by_task": _count(lambda o: f"{o['task'].get('id')} {o['task'].get('name') or ''}".strip()),
            "by_flag": _flag_counts(objections),
            "by_unit": _count(lambda o: o["who"].get("unit") or "—"),
            "by_day": dict(sorted(_count(lambda o: o["date"] or "—").items())),
        },
        "objections": objections,
    }
    return {"report": report, "downloads": downloads,
            "examples": [examples[k] for k in sorted(examples)]}


def _flag_counts(objections: list[dict]) -> dict:
    out: dict = {}
    for o in objections:
        for f in ((o.get("ai") or {}).get("flags") or []):
            out[f] = out.get(f, 0) + 1
    return dict(sorted(out.items(), key=lambda kv: -kv[1]))


class _Lookups:
    """Names read once, as plain values — a rollback inside one objection must
    not expire what every other one is printed with."""

    def __init__(self, db: Session):
        self.units = {r.id: {"name": r.name, "shift": r.shift, "factory_id": r.factory_id}
                      for r in db.query(Manager.id, Manager.name, Manager.shift,
                                        Manager.factory_id).all()}
        self.factories = {r.id: (r.name_uz or r.code)
                          for r in db.query(Factory.id, Factory.name_uz, Factory.code).all()}
        self.profiles = {r.id: {"name": r.name, "manager_id": r.manager_id}
                         for r in db.query(RoleProfile.id, RoleProfile.name,
                                           RoleProfile.manager_id).all()}
        self.cells = {r.id: {"code": r.verifix_code, "one_process": getattr(r, "one_process", None)}
                      for r in db.query(Cell).all()}
        self.defs = {r.id: {"uz": r.name_uz, "ru": r.name_ru, "en": r.name_en,
                            "weight": int(r.default_weight or 0),
                            "auto_check": getattr(r, "auto_check", None)}
                     for r in db.query(LeaderTaskDef).all()}


def _bare(db: Session, i: int, card: dict, exc: Exception, ctx: _Lookups) -> dict:
    """What can still be said about an objection whose detail could not be read."""
    d = db.query(LeaderAiDispute).filter_by(id=i).first()
    from app.services import leader_appeal_chat as chat
    task_id = d.task_id if d else card.get("taskId")
    date = d.date if d else card.get("date")
    o = {
        "id": i, "date": date, "seq": 0,
        "in_page_period": bool(date and PAGE_FROM <= str(date)[:10] <= PAGE_TO),
        "kind": card.get("kind") or "ai",
        "folder": f"{date or 'no-date'}_d{i}_t{int(task_id or 0):02d}_{_slug(d.leader_name if d else card.get('leader'))}",
        "links": {"chat": chat.chat_url(chat.DISPUTE, i), "report": None},
        "who": {"leader": (d.leader_name if d else card.get("leader")),
                "unit": card.get("supervisor") or ctx.units.get(d.manager_id if d else None, {}).get("name")},
        "task": {"id": task_id, "name": ctx.defs.get(task_id or 0, {}).get("uz")},
        "objection": {"reason": _cut(d.reason) if d else card.get("reason")},
        "errors": [f"the detail of this objection could not be read: {_why(exc)}"],
        "files": [],
    }
    return o


def _one(db: Session, d: LeaderAiDispute, card: dict, ctx: _Lookups,
         examples: dict) -> tuple[dict, list[dict]]:
    """One objection, whole — and what has to be downloaded for it."""
    from app.services import (
        auto_check_timing, leader_ai, leader_appeal_chat as chat, leader_auto,
        leader_dispute, leader_reports, leader_shift, leader_tasks, one_process,
    )

    errors: list[str] = []

    def guard(what, fn, default=None):
        try:
            return fn()
        except Exception as exc:
            db.rollback()
            errors.append(f"{what}: {_why(exc)}")
            return default

    # The ORM row's values, copied before anything can roll the session back.
    p = {
        "id": d.id, "ref": d.ref, "review_id": d.review_id, "date": d.date,
        "task_id": d.task_id, "leader_id": d.leader_id, "leader_name": d.leader_name,
        "manager_id": d.manager_id, "status": d.status, "reason": d.reason,
        "requested_by_profile": d.requested_by_profile,
        "requested_by_name": d.requested_by_name, "requested_at": d.requested_at,
        "sup_action": d.sup_action, "sup_note": d.sup_note,
        "sup_case": leader_dispute.sup_case(d), "sup_by": d.sup_by_name,
        "sup_at": d.sup_at, "admin_ruling": d.admin_ruling,
    }
    kind = card.get("kind") or guard("kind", lambda: leader_dispute.kind_of(db, d), "ai")
    date = str(p["date"] or "")[:10]
    task_id = int(p["task_id"])

    rev = guard("verdict row", lambda: (
        (db.query(LeaderAiReview).filter_by(id=p["review_id"]).first() if p["review_id"] else None)
        or db.query(LeaderAiReview).filter_by(ref=p["ref"]).first()))
    review = None
    if rev is not None:
        review = {
            "id": rev.id, "ref": rev.ref, "source": rev.source, "date": rev.date,
            "task_id": rev.task_id, "leader_id": rev.leader_id,
            "manager_id": rev.manager_id, "shift": rev.shift, "status": rev.status,
            "flags": list(rev.flags or []), "image_date": rev.image_date,
            "clocks": _plain(list(rev.clocks or [])),
            "reason": {l: getattr(rev, f"reason_{l}") for l in leader_ai.LANGS},
            "photos_sent": rev.photos, "model": rev.model, "attempts": rev.attempts,
            "error": rev.error, "created_at": _iso(rev.created_at),
            "reviewed_at": _iso(rev.reviewed_at), "resolution": rev.resolution,
            "resolved_by": rev.resolved_by, "resolved_at": _iso(rev.resolved_at),
            "resolution_note": rev.resolution_note,
        }

    # Whose rule: the verdict row's own stamps where there is one (the reviewer
    # judged by them), the objection's otherwise.
    mid = (review or {}).get("manager_id") or p["manager_id"]
    lid = (review or {}).get("leader_id") or p["leader_id"]
    unit = ctx.units.get(mid) or ctx.units.get(p["manager_id"]) or {}
    shift = card.get("shift") or (review or {}).get("shift") or unit.get("shift")

    uid = card.get("uid") or guard("report id", lambda: leader_reports.uid_of_ref(db, p["ref"]))
    rep = guard("day report", lambda: leader_reports.day_report(db, uid, allow_open=True)) if uid else None
    rtask = None
    if rep:
        rtask = next((t for t in rep.get("tasks") or []
                      if int(t.get("id") or 0) == task_id), None)

    # ── the bot entry, its day, its photos ───────────────────────────────────
    entry = day = None
    media: list[dict] = []
    roll: list[dict] = []
    if str(p["ref"]).startswith("bot:") and str(p["ref"])[4:].isdigit():
        eid = int(str(p["ref"])[4:])
        e = guard("entry", lambda: db.query(LeaderTaskEntry).filter_by(id=eid).first())
        if e is not None:
            entry = {"id": e.id, "day_id": e.day_id, "done": bool(e.done),
                     "reason": e.reason, "saved_at": e.saved_at, "closed_at": e.closed_at}
            dd = guard("day", lambda: db.query(LeaderTaskDay).filter_by(id=entry["day_id"]).first())
            if dd is not None:
                day = {"id": dd.id, "date": dd.date, "cell_id": dd.cell_id,
                       "manager_id": dd.manager_id, "leader_id": dd.leader_id,
                       "closed_at": dd.closed_at,
                       "completion": float(dd.completion) if dd.completion is not None else None,
                       "finished_at": dd.finished_at, "reopened": dd.reopened}
            media = guard("photos", lambda: [
                {"id": m.id, "file_id": m.file_id, "pos": m.pos}
                for m in db.query(LeaderTaskMedia).filter_by(entry_id=eid)
                .order_by(LeaderTaskMedia.pos, LeaderTaskMedia.id).all()], []) or []
            roll = guard("camera roll", lambda: [
                {"slot": s.slot, "file_id": s.file_id, "captured_at": s.captured_at,
                 "received_at": s.received_at, "stamp": s.stamp, "late": bool(s.late),
                 "deferred": bool(s.deferred), "skew_s": s.skew_s}
                for s in db.query(LeaderTaskPhoto)
                .filter_by(day_id=entry["day_id"], task_id=task_id)
                .order_by(LeaderTaskPhoto.slot, LeaderTaskPhoto.id).all()], []) or []

    cell_id = (day or {}).get("cell_id") or (rep or {}).get("cellId")
    cell = ctx.cells.get(cell_id) if cell_id else None

    # ── the task as configured for this leader, day and cell ─────────────────
    prof = guard("leader profile", lambda: db.query(RoleProfile).filter_by(id=lid).first()) if lid else None
    cfg = {}
    if prof is not None:
        cfg = guard("task config", lambda: (leader_tasks.effective_leader_config(
            db, prof, shift, day=date) or {}).get(task_id) or {}, {}) or {}
    day_cell = (guard("day cell", lambda: one_process.day_cell(db, p["ref"]))
                if task_id == one_process.TASK else None)
    criteria = guard("criteria", lambda: leader_ai.criteria_for(
        db, task_id, mid, lid, date=date, cell_id=day_cell), "") or ""
    level = guard("criteria level", lambda: _criteria_level(db, task_id, mid, lid, criteria))
    label = guard("task name", lambda: leader_ai.task_label(db, task_id, mid, lid, date=date))
    note = guard("task note", lambda: leader_ai.task_note(db, task_id), "") or ""
    rule = guard("date rule", lambda: leader_ai.date_rule_for(db, task_id, mid, lid, shift, date=date))
    one_proc = None
    if task_id == one_process.TASK:
        one_proc = guard("one process", lambda: one_process.applies(
            db, task_id=task_id, leader_id=lid, day=date, cell_id=day_cell))

    names = card.get("taskName") if isinstance(card.get("taskName"), dict) else {}
    tdef = ctx.defs.get(task_id) or {}
    task = {
        "id": task_id,
        "name": (names or {}).get("uz") or tdef.get("uz") or label,
        "name_ru": (names or {}).get("ru") or tdef.get("ru"),
        "name_en": (names or {}).get("en") or tdef.get("en"),
        "label_ai_saw": label,
        "what_to_photograph": note or None,
        "leader_instruction": cfg.get("description") or None,
        "ai_criteria": criteria or None,
        "ai_criteria_level": level,
        "one_process_cell": one_proc,
        "photos_required": cfg.get("min_media"),
        "weight": cfg.get("weight") if cfg.get("weight") is not None else tdef.get("weight"),
        "proof_kind": ("auto" if cfg.get("auto_check") or kind == "auto"
                       else cfg.get("proof_kind")),
        "auto_check": cfg.get("auto_check") or tdef.get("auto_check"),
        "deadline": cfg.get("deadline"),
        "enabled": cfg.get("enabled"),
    }
    if rule is not None and task["proof_kind"] != "auto":
        verdict_card = card.get("verdict") or {}
        task["date_rule"] = {
            "mode": _mode(rule.checked, rule.dayed, rule.timed),
            "mode_text": MODE_TEXT[_mode(rule.checked, rule.dayed, rule.timed)],
            "window": list(rule.win),
            "days_after_report_allowed": rule.plus,
            "expected": verdict_card.get("expected") or guard("expected", lambda: leader_ai.expected_text(
                date, shift, rule.win, check=rule.checked, days=rule.dayed,
                times=rule.timed, plus=rule.plus)),
        }

    # ── what the AI was shown, and the prompt it is given (rebuilt today) ────
    ex_paths: list[str] = []
    if kind != "auto":
        def _examples():
            ex_mid = mid
            if lid and leader_shift.moved(db, lid) and leader_shift.shifted(
                    db, lid, date, leader_shift.unit_shift(db, mid)):
                ex_mid = None
            ex_ids = leader_ai.example_ids_map(db, manager_id=ex_mid, leader_id=lid).get(task_id, [])
            out = []
            for ex in (db.query(LeaderTaskExample).filter(LeaderTaskExample.id.in_(ex_ids))
                       .order_by(LeaderTaskExample.id).all() if ex_ids else []):
                if ex.id not in examples:
                    lvl = leader_ai.example_level(ex)
                    raw = bytes(ex.data or b"")
                    examples[ex.id] = {
                        "id": ex.id, "task_id": ex.task_id, "level": lvl,
                        "manager_id": ex.manager_id, "leader_id": ex.leader_id,
                        "path": f"examples/ex{ex.id}_t{ex.task_id:02d}_{lvl}.{_ext(raw, ex.mime)}",
                        "data": raw,
                    }
                out.append(examples[ex.id]["path"])
            return out
        ex_paths = guard("examples", _examples, []) or []

    sheet_urls = _sheet_urls(rtask)
    if not media and not sheet_urls and rev is not None and review and review.get("source") == "sheet":
        # The day report could not be read — take the photos off the form row
        # the AI itself reviewed (`leader_ai._sheet_row`).
        def _form_urls():
            row = leader_ai._sheet_row(db, rev)
            tk = next((t for t in (row.tasks or []) if int(t.get("id") or 0) == task_id),
                      None) if row is not None else None
            return leader_ai._sheet_photos(tk) if tk else []
        sheet_urls = guard("form photos", _form_urls, []) or []
    n_photos = len(media) if media else len(sheet_urls)
    prompt = None
    if kind != "auto" and task["proof_kind"] != "auto" and rule is not None:
        prompt = guard("prompt", lambda: leader_ai._prompt(
            task=label or task["name"] or f"#{task_id}", note=note, criteria=criteria,
            n_images=max(1, min(n_photos, leader_ai.MAX_IMAGES)),
            omitted=max(0, n_photos - leader_ai.MAX_IMAGES),
            n_examples=len(ex_paths),
            screen_dates=bool(rule.checked and rule.dayed and not rule.timed)))

    edits = []
    if review and review.get("reviewed_at"):
        edits = guard("criteria edits", lambda: _criteria_edits(
            db, task_id, mid, _aware(review["reviewed_at"])), []) or []

    # ── the photos ───────────────────────────────────────────────────────────
    folder = (f"{date or 'no-date'}_d{p['id']}_t{task_id:02d}_"
              f"{_slug(card.get('leader') or p['leader_name'])}")
    downloads: list[dict] = []
    files: list[dict] = []
    cams = list(roll) if len(roll) == len(media) else []
    for n, m in enumerate(media, 1):
        key = f"m{m['id']}"
        cam = cams[n - 1] if cams else None
        files.append({
            "key": key, "role": "proof", "n": n, "base": f"proof_{n}",
            "sent_to_ai": n <= leader_ai.MAX_IMAGES,
            "camera": ({"captured_at": _iso(cam["captured_at"]),
                        "received_at": _iso(cam["received_at"]),
                        "stamp": cam["stamp"], "late": cam["late"],
                        "deferred": cam["deferred"], "skew_s": cam["skew_s"]}
                       if cam else None),
        })
        downloads.append({"key": key, "obj": p["id"], "src": "tg", "file_id": m["file_id"]})
    if not media:
        for n, u in enumerate(sheet_urls, 1):
            key = f"u{p['id']}_{n}"
            files.append({"key": key, "role": "proof", "n": n, "base": f"proof_{n}",
                          "sent_to_ai": n <= leader_ai.MAX_IMAGES, "url": u,
                          "camera": None})
            downloads.append({"key": key, "obj": p["id"], "src": "url", "url": u})
    extra_roll = []
    if roll and not cams:
        # The camera roll does not line up with the photos the entry holds —
        # list it as it stands rather than guess which shot is which.
        extra_roll = [{"slot": r["slot"], "captured_at": _iso(r["captured_at"]),
                       "received_at": _iso(r["received_at"]), "late": r["late"],
                       "deferred": r["deferred"]} for r in roll]

    # ── the chat ─────────────────────────────────────────────────────────────
    msgs_out: list[dict] = []
    members_out: list[dict] = []

    def _chat():
        msgs = chat.messages(db, chat.DISPUTE, p["id"])
        att = chat.files_of(db, [m.id for m in msgs])
        for m in msgs:
            fl = []
            for k, f in enumerate(att.get(m.id, []), 1):
                key = f"c{f.id}"
                stem, _, _e = (f.name or "file").rpartition(".")
                base = f"chat_{m.id}_{k}_{_slug(stem or f.name, 'file', 30)}"
                fe = {"key": key, "role": "chat", "base": base, "name": f.name,
                      "mime": f.mime, "size": f.size, "message_id": m.id}
                if (f.size or 0) > CHAT_FILE_MAX:
                    fe["skipped"] = "larger than 20 MB"
                else:
                    downloads.append({"key": key, "obj": p["id"], "src": "tg",
                                      "file_id": f.file_id, "name": f.name, "mime": f.mime})
                files.append(fe)
                fl.append(key)
            msgs_out.append({
                "id": m.id, "kind": m.kind, "text": _cut(m.text),
                "author": m.author_name, "author_role": m.author_role,
                "author_profile": m.author_profile,
                "at": _iso(m.created_at), "edited_at": _iso(m.edited_at), "files": fl,
            })
        for mem in chat.members(db, chat.DISPUTE, p["id"]):
            members_out.append({"name": mem.name, "profile": mem.profile_key,
                                "invited_by": mem.invited_by_name, "at": _iso(mem.created_at)})
    guard("chat", _chat)

    # ── an automatic check ───────────────────────────────────────────────────
    auto = None
    if kind == "auto" or (card.get("auto") or {}).get("code"):
        auto = guard("automatic check", lambda: _auto_detail(
            db, d, p, entry, uid, task_id, mid, card, leader_auto, auto_check_timing))

    # ── the admin ruling under objection ─────────────────────────────────────
    ruling = None
    if p["admin_ruling"]:
        r = p["admin_ruling"] if isinstance(p["admin_ruling"], dict) else {}
        ruling = {"by": r.get("by"), "at": _iso(r.get("at")),
                  "override_now": {"done": (rtask or {}).get("admin_done"),
                                   "by": (rtask or {}).get("admin_by"),
                                   "at": _iso((rtask or {}).get("admin_at"))}}

    verdict_card = card.get("verdict") or {}
    ai = None
    if review is not None:
        ai = {**review,
              "flags_text": [FLAG_TEXT.get(f, f) for f in review["flags"]],
              "date_reason": verdict_card.get("dateReason"),
              "expected": verdict_card.get("expected"),
              "rejects_task": bool((rtask or {}).get("ai_rejected"))}

    leader_name = card.get("leader") or p["leader_name"]
    o = {
        "id": p["id"], "seq": 0, "status": p["status"], "kind": kind,
        "kind_text": KIND_TEXT.get(kind, kind), "date": date,
        "in_page_period": bool(date and PAGE_FROM <= date <= PAGE_TO),
        "folder": folder,
        "links": {"chat": chat.chat_url(chat.DISPUTE, p["id"]),
                  "report": leader_reports.report_url(uid) if uid else None},
        "who": {
            "leader": leader_name, "leader_profile_id": lid,
            "unit": card.get("supervisor") or unit.get("name"), "unit_id": mid,
            "shift": shift,
            "factory": ctx.factories.get(unit.get("factory_id")),
            "cell": (cell or {}).get("code") or (rep or {}).get("cell"),
            "cell_id": cell_id,
        },
        "day": {
            "report_uid": uid, "source": (rep or {}).get("source"),
            "open": (rep or {}).get("open"),
            "closed_at": _iso((day or {}).get("closed_at")) or _iso((rep or {}).get("submittedAt")),
            "score": (rep or {}).get("score"), "raw_score": (rep or {}).get("rawScore"),
            "counts": (rep or {}).get("counts"),
            "excluded": _plain((rep or {}).get("excluded")),
            "voided": (rep or {}).get("voided"),
        },
        "task": task,
        "answer": {
            "answered": (rtask or {}).get("answered") if rtask else (entry is not None),
            "done": entry["done"] if entry else (rtask or {}).get("done"),
            "leader_comment": _sentinel(entry["reason"] if entry else (rtask or {}).get("reason")),
            "comment_by_platform": str((entry or {}).get("reason") or (rtask or {}).get("reason") or "")
            .startswith(("__auto__|", "__missed__|")),
            "saved_at": _iso(entry["saved_at"]) if entry else None,
            "submitted_at": _iso(entry["closed_at"]) if entry else _iso((rep or {}).get("submittedAt")),
            "rejected_by_now": (rtask or {}).get("rejected_by"),
            "camera_roll_unaligned": extra_roll or None,
        },
        "ai": ai,
        "examples_shown_to_ai": ex_paths,
        "criteria_edits_after_verdict": edits,
        "ai_prompt_rebuilt_today": prompt,
        "auto": auto,
        "admin_ruling": ruling,
        "objection": {
            "filed_by": p["requested_by_name"],
            "filed_by_role": (str(p["requested_by_profile"] or "").split(":")[0] or None),
            "filed_at": _iso(p["requested_at"]),
            "reason": _cut(p["reason"]),
            "brigadir": ({"action": p["sup_action"], "case": _cut(p["sup_case"]),
                          "by": p["sup_by"], "at": _iso(p["sup_at"])}
                         if p["sup_action"] else None),
            "waiting_since": _iso(p["sup_at"] or p["requested_at"]),
        },
        "chat": msgs_out,
        "chat_members_invited": members_out,
        "files": files,
        "errors": errors,
    }
    return o, downloads


def _sheet_urls(task: dict | None) -> list[str]:
    return [u.strip() for u in str((task or {}).get("photo") or "").split(",") if "http" in u]


def _criteria_level(db: Session, task_id: int, mid, lid, criteria: str) -> str | None:
    """Which level of the chain the criteria text comes from — informational."""
    from app.services import one_process
    if not criteria:
        return None
    own = (db.query(LeaderTaskLeaderSetting).filter_by(leader_id=lid, task_id=task_id).first()
           if lid else None)
    if own is not None and (own.criteria or "").strip() == criteria:
        return "leader"
    if criteria == one_process.CRITERIA:
        return "one_process"
    sup = (db.query(LeaderTaskSetting).filter_by(manager_id=mid, task_id=task_id).first()
           if mid else None)
    if sup is not None and (sup.criteria or "").strip() == criteria:
        return "unit"
    return "global"


def _criteria_edits(db: Session, task_id: int, mid, after) -> list[dict]:
    """Every edit of this task's AI criteria AFTER the verdict that could reach
    this unit — the criteria printed are today's, and this is how to tell
    whether the AI saw a different text."""
    if after is None:
        return []
    out = []
    for a in (db.query(ActionLog)
              .filter(ActionLog.action == "ltask.criteria_set",
                      ActionLog.outcome == "done",
                      ActionLog.target_id == str(task_id),
                      ActionLog.created_at > after)
              .order_by(ActionLog.created_at, ActionLog.id).limit(50).all()):
        if a.unit_id not in (None, mid):
            continue
        out.append({"at": _iso(a.created_at), "by": a.actor_name,
                    "unit": a.unit_name, "details": _plain(a.details)})
    return out


def _auto_detail(db, d, p, entry, uid, task_id, mid, card, leader_auto, timing_mod) -> dict:
    """What the automatic check read at its hour, and the evidence of when the
    leader acted."""
    a = card.get("auto") or {}
    row = None
    if entry:
        row = (db.query(LeaderAutoCheck).filter_by(entry_id=entry["id"])
               .order_by(LeaderAutoCheck.id.desc()).first())
    out = {
        "reason": _sentinel(a.get("reason")), "hour": a.get("time"),
        "code": a.get("code"), "code_text": CODE_TEXT.get(a.get("code"), a.get("code")),
        "facts": _plain(a.get("facts")),
    }
    if row is not None:
        out.update({
            "check": row.check, "check_text": CHECK_TEXT.get(row.check, row.check),
            "asked_at": _iso(row.due_at), "warned_at": _iso(row.warned_at),
            "checked_at": _iso(row.checked_at), "outcome": row.outcome,
            "code": row.code or out["code"],
            "code_text": CODE_TEXT.get(row.code, row.code) if row.code else out["code_text"],
            "facts": _plain(row.facts),
        })
    facts = out.get("facts") if isinstance(out.get("facts"), dict) else {}
    out["result_lines_en"] = leader_auto.result_lines(facts, "en") if facts else []
    out["result_lines_uz"] = leader_auto.result_lines(facts, "uz") if facts else []
    try:
        out["timing"] = _plain(timing_mod.for_dispute(db, d))
    except Exception as exc:
        db.rollback()
        out["timing_error"] = _why(exc)
    check = (row.check if row is not None else None)
    if check == "staff_list" and uid:
        try:
            out["staff_list"] = _plain(leader_auto.staff_list_snapshot(db, uid, task_id))
        except Exception as exc:
            db.rollback()
            out["staff_list_error"] = _why(exc)
    if check in ("plan_staffing", "plan_pct") and mid:
        try:
            out["unit_saves"] = _unit_saves(db, mid, p["date"])
        except Exception as exc:
            db.rollback()
            out["unit_saves_error"] = _why(exc)
    return out


def _unit_saves(db: Session, mid: int, day: str) -> list[dict]:
    """Every «Odamlar soni» save and ПЛАН/ФАКТ edit of this unit for the day
    (and the next, where a night's work may sit) — who, when, refused or not."""
    try:
        d0 = _date.fromisoformat(str(day)[:10])
    except ValueError:
        return []
    rows = (db.query(ActionLog)
            .filter(ActionLog.unit_id == mid,
                    ActionLog.day.in_([d0, d0 + timedelta(days=1)]),
                    ActionLog.action.in_(("production.staffing_set",
                                          "production.wc_override_set",
                                          "production.override_set")))
            .order_by(ActionLog.created_at, ActionLog.id).limit(TIMELINE_MAX).all())
    return [{"at": _iso(r.created_at), "day": r.day.isoformat() if r.day else None,
             "action": r.action, "outcome": r.outcome, "by": r.actor_name,
             "by_role": r.actor_role, "target": r.target_name,
             "details": _plain(r.details), "changes": _plain(r.changes)} for r in rows]


# ── the words a reader reads ─────────────────────────────────────────────────

def _pairs(v) -> str:
    """An action-register `details` / `changes` list as one short line."""
    if isinstance(v, list):
        out = []
        for x in v[:12]:
            if isinstance(x, (list, tuple)):
                out.append("=".join(str(y) for y in x))
            else:
                out.append(str(x))
        return "; ".join(out)[:300]
    return json.dumps(v, ensure_ascii=False)[:300]


def _quote(text) -> list[str]:
    s = str(text or "").strip()
    if not s:
        return ["> —"]
    return ["> " + ln if ln.strip() else ">" for ln in s.splitlines()]


def _file_line(f: dict, names: dict, failed: dict) -> str:
    path = names.get(f["key"])
    if f.get("skipped"):
        return f"`{f.get('name')}` — not included ({f['skipped']})"
    if not path:
        why = failed.get(f["key"]) or "not fetched"
        return f"`{f['base']}` — COULD NOT BE FETCHED ({why})"
    return f"`{path}`"


def card_md(o: dict, total: int, names: dict, failed: dict) -> str:
    """One objection, in full, as markdown — the same text in objections.md and
    in the objection's own folder. Paths are from the root of the unzipped
    folder."""
    L: list[str] = []
    w, t, ans = o.get("who") or {}, o.get("task") or {}, o.get("answer") or {}
    L.append(f"## {o.get('seq')}/{total} · objection #{o['id']} · {w.get('leader') or '—'} · "
             f"task №{t.get('id')} «{t.get('name') or '—'}» · {_dmy(o.get('date'))}")
    L.append("")
    L.append(f"- **Open the chat to rule:** {o['links']['chat']}")
    if o["links"].get("report"):
        L.append(f"- Day report: {o['links']['report']}")
    L.append(f"- Folder: `{o.get('folder')}/`")
    if not o.get("in_page_period"):
        L.append(f"- ⚠ The checklist day is OUTSIDE {_dmy(PAGE_FROM)[:10]}–{_dmy(PAGE_TO)[:10]}, "
                 f"the period on your screen — it waits on an admin all the same.")
    ob = o.get("objection") or {}
    L.append(f"- Waiting on an admin since: {_t(ob.get('waiting_since'))}")
    flags = (o.get("ai") or {}).get("flags") or []
    L.append(f"- Refused by: {o.get('kind_text') or o.get('kind')}"
             + (f" — flags: {', '.join(flags)}" if flags else ""))
    if t.get("weight") is not None:
        L.append(f"- At stake: the task's weight, {t['weight']} points (approving gives them back)")
    for e in o.get("errors") or []:
        L.append(f"- ⚠ Not fully read: {e}")
    L.append("")

    L.append("### Who and when")
    L.append(f"- Checklist day: {_dmy(o.get('date'))}, shift {w.get('shift') or '—'}")
    L.append(f"- Leader: {w.get('leader') or '—'} (profile {w.get('leader_profile_id') or '—'})")
    L.append(f"- Brigadir (unit): {w.get('unit') or '—'} (unit {w.get('unit_id') or '—'})"
             + (f" · plant {w['factory']}" if w.get("factory") else ""))
    if w.get("cell"):
        L.append(f"- Cell: {w['cell']}")
    dy = o.get("day") or {}
    if dy:
        state = "still OPEN (being filed)" if dy.get("open") else f"closed {_t(dy.get('closed_at'))}"
        score = (f" · day score now {dy.get('score')}%" if dy.get("score") is not None else "")
        L.append(f"- Day: {state}{score} · source: {dy.get('source') or '—'}")
        if dy.get("excluded"):
            L.append(f"- Day is EXCLUDED from the results: {dy['excluded']}")
        if dy.get("voided"):
            L.append("- Day is VOIDED by the filing-window rule")
    L.append("")

    L.append("### The task")
    L.append(f"- Task №{t.get('id')}: «{t.get('name') or '—'}»"
             + (f" (ru: «{t['name_ru']}»)" if t.get("name_ru") else ""))
    if t.get("what_to_photograph"):
        L.append(f"- What to photograph (the task's note): {t['what_to_photograph']}")
    L.append("- Instruction the LEADER reads (never shown to the AI):")
    L += ["  " + q for q in _quote(t.get("leader_instruction"))]
    if t.get("auto_check") or o.get("kind") == "auto":
        L.append("- This task is answered by an AUTOMATIC CHECK, not by photos.")
    else:
        lvl = t.get("ai_criteria_level")
        if t.get("ai_criteria"):
            L.append("- AI criteria — the text the AI judged the photos against "
                     f"(resolved today for this leader, day and cell; level: {lvl or '—'}):")
            L += ["  " + q for q in _quote(t.get("ai_criteria"))]
        else:
            L.append("- AI criteria: NONE written at any level — the AI judged the photos against "
                     "the task's name and note only (and the date rule).")
        if t.get("one_process_cell") is not None:
            L.append(f"- One-process cell (task 3's lenient text): {'yes' if t['one_process_cell'] else 'no'}")
    bits = []
    if t.get("photos_required") is not None and t.get("proof_kind") != "auto":
        bits.append(f"photos required: {t['photos_required']}")
    if t.get("weight") is not None:
        bits.append(f"weight: {t['weight']}")
    if t.get("deadline"):
        bits.append(f"deadline: {t['deadline']}")
    if bits:
        L.append("- " + " · ".join(bits))
    if t.get("proof_kind"):
        L.append(f"- Proof: {PROOF_TEXT.get(t['proof_kind'], t['proof_kind'])}")
    dr = t.get("date_rule")
    if dr and o.get("kind") != "auto":
        L.append(f"- Date rule: {dr.get('mode_text')}")
        L.append(f"  - accepted: {dr.get('expected') or '—'} (window {'–'.join(dr.get('window') or [])}"
                 f", days after the report allowed: {dr.get('days_after_report_allowed')})")
    edits = o.get("criteria_edits_after_verdict") or []
    if edits:
        L.append("- ⚠ The AI criteria were EDITED after the verdict (the AI may have seen an older text):")
        for e in edits:
            L.append(f"  - {_t(e.get('at'))} by {e.get('by') or '—'}"
                     + (f" ({e['unit']})" if e.get("unit") else "")
                     + f": {json.dumps(e.get('details'), ensure_ascii=False)[:300]}")
    L.append("")

    L.append("### The leader's answer")
    done = ans.get("done")
    by_platform = bool(ans.get("comment_by_platform"))
    who_ans = "Recorded by the platform" if by_platform else "Answer"
    L.append(f"- {who_ans}: {'Yes — done' if done else ('No — not done' if done is False else '—')}"
             f" · saved {_t(ans.get('saved_at'))} · submitted {_t(ans.get('submitted_at'))}")
    if ans.get("leader_comment"):
        L.append("- Reason the platform stored on it:" if by_platform
                 else "- The leader's own comment on the answer:")
        L += ["  " + q for q in _quote(ans["leader_comment"])]
    proofs = [f for f in o.get("files") or [] if f.get("role") == "proof"]
    if proofs:
        L.append(f"- Proof photos ({len(proofs)}; the AI sees at most 4, in this order):")
        for f in proofs:
            extra = []
            cam = f.get("camera")
            if cam:
                extra.append(f"in-app camera, taken {_ts(cam.get('captured_at'))} (server clock)")
                if cam.get("late"):
                    extra.append("OUTSIDE the photo window")
                if cam.get("deferred"):
                    extra.append("uploaded later (offline)")
            if not f.get("sent_to_ai"):
                extra.append("NOT sent to the AI (beyond the first 4)")
            L.append(f"  {f.get('n')}. {_file_line(f, names, failed)}"
                     + (f" — {'; '.join(extra)}" if extra else ""))
    elif o.get("kind") != "auto":
        L.append("- Proof photos: none stored")
    if ans.get("camera_roll_unaligned"):
        L.append("- In-app camera roll (does not line up one-to-one with the photos above):")
        for r in ans["camera_roll_unaligned"]:
            L.append(f"  - slot {r.get('slot')}: taken {_ts(r.get('captured_at'))}"
                     + (" · OUTSIDE the window" if r.get("late") else ""))
    L.append("")

    ai = o.get("ai")
    if ai:
        L.append("### The AI verdict")
        L.append(f"- Status: {ai.get('status')} · reviewed {_t(ai.get('reviewed_at'))}"
                 f" · model {ai.get('model') or '—'} · photos it saw: {ai.get('photos_sent')}")
        if ai.get("flags"):
            L.append("- Flags:")
            for ft in ai.get("flags_text") or []:
                L.append(f"  - {ft}")
        else:
            L.append("- Flags: none (the AI passed it)")
        for lang, title in (("uz", "uz"), ("en", "en"), ("ru", "ru")):
            txt = (ai.get("reason") or {}).get(lang)
            if txt:
                L.append(f"- AI's reason ({title}):")
                L += ["  " + q for q in _quote(txt)]
        if ai.get("image_date"):
            L.append(f"- Date/time the AI read on the photos: {ai['image_date']}")
        dr_txt = ai.get("date_reason")
        if isinstance(dr_txt, dict):
            for lang in ("en", "uz"):
                if dr_txt.get(lang):
                    L.append(f"- Date verdict ({lang}): {dr_txt[lang]}")
        elif dr_txt:
            L.append(f"- Date verdict: {dr_txt}")
        if ai.get("expected"):
            L.append(f"- Measured against: {ai['expected']}")
        if ai.get("error"):
            L.append(f"- AI error: {ai['error']}")
        if ai.get("resolution"):
            L.append(f"- A ruling already on the verdict: {ai['resolution']} by {ai.get('resolved_by') or '—'}"
                     f" at {_t(ai.get('resolved_at'))}"
                     + (f" — {ai['resolution_note']}" if ai.get("resolution_note") else ""))
        ex = o.get("examples_shown_to_ai") or []
        L.append("- Example photos the AI was shown in FRONT of the proofs: "
                 + (", ".join(f"`{x}`" for x in ex) if ex else "none"))
        if o.get("ai_prompt_rebuilt_today"):
            L.append(f"- The AI's full prompt (rebuilt from today's settings): `{o.get('folder')}/ai_prompt.txt`")
        L.append("")

    au = o.get("auto")
    if au:
        L.append("### The automatic check")
        if au.get("check_text"):
            L.append(f"- Check: {au['check_text']}")
        L.append(f"- Asked at {_t(au.get('asked_at'))} · read at {_t(au.get('checked_at'))}"
                 f" · leader warned at {_t(au.get('warned_at'))}")
        L.append(f"- Verdict: {au.get('outcome') or '—'} — {au.get('code_text') or au.get('code') or '—'}")
        lines = au.get("result_lines_en") or []
        if lines:
            L.append("- What it read at the hour:")
            L += [f"  - {x}" for x in lines]
        tm = au.get("timing")
        if isinstance(tm, dict) and tm:
            L.append("- When the leader acted (rebuilt from the action register and the concern times):")
            L.append(f"  - timing verdict: {TIMING_TEXT.get(tm.get('verdict'), tm.get('verdict') or '—')}")
            if tm.get("hour_pct") is not None:
                L.append(f"  - completion at the hour: {tm['hour_pct']}% · now: {tm.get('now_pct')}%"
                         + (f" · pass mark {tm.get('target')}%" if tm.get("target") is not None else ""))
            if tm.get("found") is not None:
                L.append(f"  - concerns that counted (written before the hour): {tm['found']}")
            if tm.get("reached"):
                L.append(f"  - reached at {_t(tm['reached'])}"
                         + (f" — {tm['late_min']} min after the hour" if tm.get("late_min") else ""))
            for ev in (tm.get("events") or [])[:40]:
                after = " · AFTER the hour" if ev.get("after") else ""
                if ev.get("kind") == "concern":
                    L.append(f"  - {_ts(ev.get('at'))} · concern №{ev.get('seq')} by "
                             f"{'a worker' if ev.get('worker') else 'the leader'}"
                             f" · cell {ev.get('cell') or '—'}{after}: {ev.get('position') or ''}")
                elif ev.get("kind") == "sap":
                    L.append(f"  - {_ts(ev.get('at'))} · SAP file upload → {ev.get('pct')}%{after}")
                else:
                    L.append(f"  - {_ts(ev.get('at'))} · ФАКТ typed by {ev.get('who') or '—'}"
                             f" ({ev.get('role') or '—'}) · {ev.get('code') or ''}"
                             f" {ev.get('position') or ''} = {ev.get('value')}"
                             f" → {ev.get('pct')}%{after}")
            if tm.get("other_days"):
                L.append(f"  - activity on other days: "
                         f"{json.dumps(tm['other_days'], ensure_ascii=False)[:400]}")
        sl = au.get("staff_list")
        if isinstance(sl, dict) and sl.get("lists"):
            L.append(f"- «Ish grafigi» at the check (cutoff {_t(sl.get('cutoff'))}) and now:")
            for li in sl["lists"]:
                L.append(f"  - cell {li.get('cell')} — list of {li.get('day')} ({li.get('kind')}):"
                         f" at the check {json.dumps(li.get('check'), ensure_ascii=False)},"
                         f" now {json.dumps(li.get('now'), ensure_ascii=False)}")
                for r in (li.get("rows") or [])[:80]:
                    c, n = r.get("check") or {}, r.get("now") or {}
                    L.append(f"    - {r.get('name')}: at check {c.get('mark') or '—'}"
                             f" · now {n.get('mark') or '—'}"
                             + (f" (set {n.get('at')} by {n.get('by')})" if n.get("at") else "")
                             + (" · marked AFTER the check" if r.get("late") else ""))
        saves = au.get("unit_saves") or []
        if saves:
            L.append(f"- The unit's people / plan / fact saves for the day and the next ({len(saves)}):")
            for sv in saves:
                L.append(f"  - {_ts(sv.get('at'))} · {SAVE_TEXT.get(sv.get('action'), sv.get('action'))}"
                         f" for {sv.get('day') or '—'} · {sv.get('outcome')}"
                         f" · {sv.get('by') or '—'} ({sv.get('by_role') or '—'})"
                         + (f" · {sv['target']}" if sv.get("target") else "")
                         + (f" · {_pairs(sv.get('details'))}" if sv.get("details") else "")
                         + (f" · changes {_pairs(sv.get('changes'))}" if sv.get("changes") else ""))
        for k in ("timing_error", "staff_list_error", "unit_saves_error"):
            if au.get(k):
                L.append(f"- ⚠ {k}: {au[k]}")
        L.append("")

    ru = o.get("admin_ruling")
    if ru:
        L.append("### The admin's ruling under objection")
        L.append(f"- An admin ({ru.get('by') or '—'}) ruled the task NOT done by hand at {_t(ru.get('at'))}")
        on = ru.get("override_now") or {}
        if on.get("done") is False:
            now_txt = f"still NOT done (set by {on.get('by') or '—'} at {_t(on.get('at'))})"
        elif on.get("done") is True:
            now_txt = f"changed to DONE since (by {on.get('by') or '—'} at {_t(on.get('at'))})"
        else:
            now_txt = "lifted since — no hand ruling stands on the task now"
        L.append(f"- The hand ruling now: {now_txt}")
        L.append("")

    L.append("### The objection")
    L.append(f"- Filed by {ob.get('filed_by') or '—'} ({ob.get('filed_by_role') or '—'}) at {_t(ob.get('filed_at'))}:")
    L += ["  " + q for q in _quote(ob.get("reason"))]
    br = ob.get("brigadir")
    if br:
        what = "passed it up to the admins" if br.get("action") == "uplifted" else br.get("action")
        L.append(f"- The brigadir ({br.get('by') or '—'}) {what} at {_t(br.get('at'))}"
                 + (":" if br.get("case") else " (their note repeats the filing above)"))
        if br.get("case"):
            L += ["  " + q for q in _quote(br["case"])]
    L.append("")

    msgs = o.get("chat") or []
    L.append(f"### The whole chat ({len(msgs)} entries)")
    fmap = {f["key"]: f for f in o.get("files") or []}
    for m in msgs:
        head = (f"- {_ts(m.get('at'))} · {m.get('author') or '—'} ({m.get('author_role') or '—'})"
                f" · {CHAT_KIND.get(m.get('kind'), m.get('kind'))}")
        L.append(head + (":" if (m.get("text") or "").strip() else ""))
        if (m.get("text") or "").strip():
            L += ["  " + q for q in _quote(m["text"])]
        for k in m.get("files") or []:
            f = fmap.get(k)
            if f:
                L.append(f"  - file «{f.get('name')}»: {_file_line(f, names, failed)}")
    mem = o.get("chat_members_invited") or []
    if mem:
        L.append("- People an admin brought into the chat: "
                 + ", ".join(f"{x.get('name')} (by {x.get('invited_by') or '—'})" for x in mem))
    L.append("")
    return "\n".join(L)


README = """# Objections waiting on an admin — how to read this

Unzip EVERY part into ONE folder. Every part carries the same README.md,
objections.md, objections.json and objections.xlsx; each objection's folder is
whole inside one part, and `examples/` rides in the first part.

* `objections.md` — every objection in full, one after another: who, which
  task, what the task asks (the leader's instruction AND the AI's criteria),
  the date rule, the leader's answer, every proof photo (path), the AI verdict
  (flags + its own reason in uz / en / ru + the date it read), the automatic
  check's evidence where the task is automatic, the admin ruling where an
  admin refused it by hand, the objection (the leader's text + the brigadir's
  case) and the whole appeal chat.
* `<date>_d<id>_t<task>_<leader>/` — one folder per objection:
  `objection.md` (the same card), `ai_prompt.txt` (the prompt the AI is given,
  rebuilt from today's settings), `proof_N.*` (the proof photos, in the order
  the leader filed them — the AI sees the first 4), `chat_*` (files attached in
  the appeal chat).
* `examples/` — the example photos the AI is shown in FRONT of the proofs
  («this is what a correct proof looks like»). The card names which ones.
* `objections.json` — all of it as data; `objections.xlsx` — one row each.

Worth knowing before judging:

* The AI criteria, the instruction and the prompt are resolved from TODAY's
  settings. An edit made after the verdict is listed under the card's
  «criteria were EDITED after the verdict» line.
* Flags: unreadable · off_topic · not_proven are the AI's judgement of the
  picture; no_date · date_mismatch are arithmetic on the clocks it read
  against the task's date rule (for an in-app camera photo the clock is the
  SERVER's capture time, which the leader cannot change).
* A date rule is «strict» (day and hour inside the window), «date only» (the
  day must match), «time only» (the hour must be inside the window) or «not
  checked».
* An automatic check (#1 #8 #9 #11) has no photo: its card shows what the
  check read at its hour and the evidence of when the leader acted.
* To rule: open the chat link on the card. Approving gives the task its full
  weight back; refusing needs a reason, which is sent to the leader.
"""


def build_workbook(report: dict, names: dict, failed: dict) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = "Objections"
    cols = [
        ("#", 5), ("Objection", 10), ("Checklist day", 12), ("Shift", 6),
        ("In 1–19 Oct", 9), ("Brigadir (unit)", 22), ("Leader", 28), ("Cell", 8),
        ("Task", 34), ("Refused by", 14), ("AI flags", 22), ("AI reason (uz)", 60),
        ("AI reason (en)", 60), ("Accepted date/time", 30), ("Date/time the AI read", 26),
        ("Leader's objection", 60), ("Brigadir's case", 50), ("Chat entries", 8),
        ("Last chat entry", 50), ("Photos", 7), ("Weight", 7), ("Waiting since", 16),
        ("Automatic check verdict", 40), ("Folder", 44), ("Chat link", 22),
        ("Day report", 22),
    ]
    head_fill = PatternFill("solid", fgColor="C8973F")
    for c, (title, width) in enumerate(cols, 1):
        cell = ws.cell(row=1, column=c, value=title)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = head_fill
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(c)].width = width
    wrap = Alignment(vertical="top", wrap_text=True)
    for r, o in enumerate(report["objections"], 2):
        w, t = o.get("who") or {}, o.get("task") or {}
        ai, au, ob = o.get("ai") or {}, o.get("auto") or {}, o.get("objection") or {}
        msgs = o.get("chat") or []
        last = msgs[-1] if msgs else {}
        photos = sum(1 for f in o.get("files") or [] if f.get("role") == "proof")
        vals = [
            o.get("seq"), o["id"], o.get("date"), w.get("shift"),
            "yes" if o.get("in_page_period") else "no", w.get("unit"), w.get("leader"),
            w.get("cell"), f"№{t.get('id')} {t.get('name') or ''}",
            {"ai": "AI", "auto": "automatic check", "admin": "admin ruling"}.get(o.get("kind"), o.get("kind")),
            ", ".join(ai.get("flags") or []), (ai.get("reason") or {}).get("uz"),
            (ai.get("reason") or {}).get("en"),
            ai.get("expected") or (t.get("date_rule") or {}).get("expected"),
            ai.get("image_date"), ob.get("reason"), (ob.get("brigadir") or {}).get("case"),
            len(msgs),
            (f"{_t(last.get('at'))} {last.get('author') or ''}: {last.get('text') or ''}" if last else None),
            photos, t.get("weight"), _t(ob.get("waiting_since")),
            (f"{au.get('outcome') or ''} — {au.get('code_text') or ''}" if au else None),
            o.get("folder"), o["links"]["chat"], o["links"].get("report"),
        ]
        for c, v in enumerate(vals, 1):
            if isinstance(v, str):
                v = v[:32000]
            cell = ws.cell(row=r, column=c, value=v)
            cell.alignment = wrap
        for c in (25, 26):
            cell = ws.cell(row=r, column=c)
            if cell.value:
                cell.hyperlink = cell.value
                cell.font = Font(color="1F5FBF", underline="single")
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(cols))}{max(1, len(report['objections']) + 1)}"

    s2 = wb.create_sheet("Counts")
    s2.column_dimensions["A"].width = 46
    s2.column_dimensions["B"].width = 10
    row = 1
    c = report["counts"]
    for title, data in (("Total", {"objections": c["objections"],
                                   "about 1–19 Oct": c["in_page_period"],
                                   "about other days": c["outside_page_period"]}),
                        ("By judge", c["by_judge"]), ("By AI flag", c["by_flag"]),
                        ("By task", c["by_task"]), ("By brigadir", c["by_unit"]),
                        ("By checklist day", c["by_day"])):
        s2.cell(row=row, column=1, value=title).font = Font(bold=True)
        row += 1
        for k, v in data.items():
            s2.cell(row=row, column=1, value=str(k))
            s2.cell(row=row, column=2, value=v)
            row += 1
        row += 1
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ── fetching ─────────────────────────────────────────────────────────────────

class _Gone(Exception):
    """A file Telegram will never serve — retrying cannot help."""


def _session() -> requests.Session:
    s = getattr(_tls, "s", None)
    if s is None:
        s = _tls.s = requests.Session()
    return s


def _fetch_tg(file_id: str) -> bytes:
    """One archived file's bytes, with Telegram's flood answer obeyed."""
    token = settings.telegram_bot_token
    last = "fetch failed"
    for attempt in range(1, FETCH_RETRIES + 1):
        wait = FETCH_BACKOFF_S * attempt
        try:
            s = _session()
            r = s.get(f"{_API}/bot{token}/getFile", params={"file_id": file_id}, timeout=60)
            body = r.json()
            if body.get("ok"):
                path = (body.get("result") or {}).get("file_path") or ""
                if not path:
                    raise _Gone("file no longer on Telegram")
                res = s.get(f"{_API}/file/bot{token}/{path}", timeout=180)
                if res.status_code == 200 and res.content:
                    return res.content
                last = f"download HTTP {res.status_code}"
            else:
                last = body.get("description") or f"getFile HTTP {r.status_code}"
                if body.get("error_code") == 400:
                    raise _Gone(last)
                wait = max(wait, int((body.get("parameters") or {}).get("retry_after") or 0))
        except _Gone:
            raise
        except Exception as exc:
            last = type(exc).__name__
        if attempt < FETCH_RETRIES:
            time.sleep(wait)
    raise RuntimeError(last)


def _fetch(item: dict) -> tuple[bytes | None, str | None]:
    try:
        if item["src"] == "tg":
            return _fetch_tg(item["file_id"]), None
        if item["src"] == "url":
            from app.services import leader_ai
            data, _mime = leader_ai.fetch_sheet_image(item["url"])
            return data, None
        return None, "unknown source"
    except Exception as exc:
        return None, _why(exc)


# ── delivery ─────────────────────────────────────────────────────────────────

def _post(method: str, data: dict, files: dict | None = None, timeout: int = 60) -> dict:
    last = ""
    for attempt in range(1, SEND_RETRIES + 1):
        wait = 0
        try:
            if files:
                for f in files.values():
                    fh = f[1]
                    if hasattr(fh, "seek"):
                        fh.seek(0)
            r = requests.post(f"{_API}/bot{settings.telegram_bot_token}/{method}",
                              data=data, files=files, timeout=timeout)
            body = r.json()
            if body.get("ok"):
                return body
            last = body.get("description") or f"HTTP {r.status_code}"
            wait = int(((body.get("parameters") or {}).get("retry_after")) or 0)
        except Exception as exc:
            last = _clean(f"{type(exc).__name__}: {exc}")
        if attempt < SEND_RETRIES:
            time.sleep(max(wait, PART_PAUSE_S * attempt))
    raise RuntimeError(_clean(f"{method} failed: {last}")[:300])


def _say(chat_id: int, text: str) -> None:
    _post("sendMessage", {"chat_id": chat_id, "text": text[:4000],
                          "disable_web_page_preview": "true"})


def _send_doc(chat_id: int, name: str, path: str, caption: str, mime: str) -> None:
    with open(path, "rb") as fh:
        _post("sendDocument", {"chat_id": chat_id, "caption": caption[:1000]},
              files={"document": (name, fh, mime)}, timeout=900)


def _load_progress() -> dict:
    s = SessionLocal()
    try:
        row = s.query(AppSetting).filter_by(key=PROGRESS_KEY).first()
        return json.loads(row.value) if row and row.value else {}
    except Exception:
        log.warning("admin objections report: progress unreadable", exc_info=True)
        return {}
    finally:
        s.close()


def _save_progress(prog: dict) -> None:
    s = SessionLocal()
    try:
        value = json.dumps(prog, separators=(",", ":"), ensure_ascii=False)
        row = s.query(AppSetting).filter_by(key=PROGRESS_KEY).first()
        if row:
            row.value = value
        else:
            s.add(AppSetting(key=PROGRESS_KEY, value=value))
        s.commit()
    except Exception:
        s.rollback()
        log.warning("admin objections report: progress not saved", exc_info=True)
    finally:
        s.close()


def send(db: Session, chat_id: int, *_window) -> int:
    """Collect, fetch, pack and DM. Returns how many files went out in this
    run. Raises on a delivery that failed outright, so `_send_report_once`
    retries it on the next boot (and resumes from the progress row)."""
    if not settings.telegram_bot_token:
        raise RuntimeError("telegram bot token not configured")

    prog = _load_progress()
    if prog.get("closed"):
        return 0
    got = collect(db)
    db.rollback()                         # plain values from here on
    report, downloads, examples = got["report"], got["downloads"], got["examples"]
    objs = report["objections"]
    c = report["counts"]
    stamp = datetime.now(TZ).strftime("%Y-%m-%d")
    total = len(objs)

    shipped = set(int(x) for x in (prog.get("shipped") or []))
    names: dict = dict(prog.get("names") or {})
    failed: dict = dict(prog.get("failed") or {})
    sent = 0

    if not prog.get("started"):
        j = c["by_judge"]
        _say(chat_id, (
            f"Objections waiting on an ADMIN ruling: {total}"
            f" ({c['in_page_period']} about 1–19 Oct, the period on your screen;"
            f" {c['outside_page_period']} about other days).\n"
            f"Refused by: the AI {j.get('ai', 0)} · an automatic check {j.get('auto', 0)}"
            f" · an admin's ruling {j.get('admin', 0)}.\n\n"
            f"Next: objections.xlsx (one row each, to skim), then ZIP parts with every "
            f"photo, example and chat file. Unzip ALL parts into ONE folder and open "
            f"README.md / objections.md. A closing message comes last."))
        prog["started"] = datetime.now(TZ).isoformat(timespec="seconds")
        _save_progress(prog)

    tmp = tempfile.mkdtemp(prefix="admin-objections-")
    try:
        # ── fetch everything not shipped yet, to disk ───────────────────────
        todo = [it for it in downloads if it["obj"] not in shipped]
        on_disk: dict[str, str] = {}
        base_of = {f["key"]: (o["folder"], f) for o in objs for f in o.get("files") or []}
        with ThreadPoolExecutor(max_workers=FETCH_WORKERS) as pool:
            for it, (data, err) in zip(todo, pool.map(_fetch, todo)):
                folder, f = base_of.get(it["key"], ("misc", {"base": it["key"]}))
                if err is not None or not data:
                    failed[it["key"]] = err or "empty file"
                    continue
                failed.pop(it["key"], None)
                rel = f"{folder}/{f['base']}.{_ext(data, it.get('mime'), it.get('name'))}"
                path = os.path.join(tmp, "files", it["key"])
                os.makedirs(os.path.dirname(path), exist_ok=True)
                with open(path, "wb") as fh:
                    fh.write(data)
                on_disk[it["key"]] = path
                names[it["key"]] = rel

        # ── the shared files every part carries ─────────────────────────────
        cards = {o["id"]: card_md(o, total, names, failed) for o in objs}
        for o in objs:
            for f in o.get("files") or []:
                if f["key"] in names:
                    f["path"] = names[f["key"]]
                elif f["key"] in failed:
                    f["error"] = failed[f["key"]]
        md = (f"# Objections waiting on an admin — {total}\n\n"
              f"Generated {report['generated_at']} (Tashkent). {c['in_page_period']} are about "
              f"1–19 Oct (the period on your screen), {c['outside_page_period']} about other days. "
              f"Read README.md first.\n\n"
              + "\n\n---\n\n".join(cards[o["id"]] for o in objs) + "\n")
        js = json.dumps({**report, "examples": [{k: v for k, v in e.items() if k != "data"}
                                                 for e in examples]},
                        ensure_ascii=False, indent=1, default=str)
        xlsx = build_workbook(report, names, failed)
        shared = {"README.md": README.encode(), "objections.md": md.encode(),
                  "objections.json": js.encode(), "objections.xlsx": xlsx}
        shared_cost = sum(len(zlib.compress(v)) for v in shared.values()) + 4096

        if not prog.get("xlsx"):
            xp = os.path.join(tmp, "objections.xlsx")
            with open(xp, "wb") as fh:
                fh.write(xlsx)
            _send_doc(chat_id, f"objections-admin-{stamp}.xlsx", xp,
                      f"Objections waiting on an admin: {total} — one row each.",
                      "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
            sent += 1
            prog["xlsx"] = True
            _save_progress(prog)

        # ── pack: one objection's folder whole inside one part ──────────────
        part_no = int(prog.get("parts") or 0)
        zf = None
        zpath = None
        size = 0
        in_part: list[int] = []
        capped = False

        def _open():
            nonlocal zf, zpath, size, part_no
            part_no += 1
            zpath = os.path.join(tmp, f"part-{part_no}.zip")
            zf = zipfile.ZipFile(zpath, "w")
            for name, data in shared.items():
                zf.writestr(name, data, compress_type=zipfile.ZIP_DEFLATED)
            size = shared_cost
            if not prog.get("examples"):
                for ex in examples:
                    zf.writestr(ex["path"], ex["data"], compress_type=zipfile.ZIP_STORED)
                    size += len(ex["data"]) + 512

        def _ship(last: bool = False):
            nonlocal zf, in_part, sent
            zf.close()
            _send_doc(chat_id, f"objections-admin-{stamp}-{part_no:02d}.zip", zpath,
                      f"Objections waiting on an admin — part {part_no}"
                      f"{' (last)' if last else ''}. Unzip every part into ONE folder; "
                      f"start with README.md.", "application/zip")
            sent += 1
            os.remove(zpath)
            shipped.update(in_part)
            prog.update({"parts": part_no, "shipped": sorted(shipped), "names": names,
                         "failed": failed, "examples": True})
            _save_progress(prog)
            zf, in_part = None, []
            time.sleep(PART_PAUSE_S)

        pending = [o for o in objs if o["id"] not in shipped]
        for i, o in enumerate(pending):
            own = [(names[f["key"]], on_disk[f["key"]]) for f in o.get("files") or []
                   if f["key"] in on_disk]
            texts = {f"{o['folder']}/objection.md": cards[o["id"]].encode()}
            if o.get("ai_prompt_rebuilt_today"):
                texts[f"{o['folder']}/ai_prompt.txt"] = o["ai_prompt_rebuilt_today"].encode()
            need = sum(os.path.getsize(pth) + 512 for _r, pth in own) + 8192
            if zf is not None and size + need > PART_BYTES and in_part:
                _ship()
            if zf is None:
                if part_no >= MAX_PARTS:
                    capped = True
                    break
                _open()
            for name, data in texts.items():
                zf.writestr(name, data, compress_type=zipfile.ZIP_DEFLATED)
            for rel, pth in own:
                fsz = os.path.getsize(pth)
                # One objection bigger than a whole part spills into the next.
                if size + fsz + 512 > PART_BYTES and size > shared_cost:
                    _ship()
                    if part_no >= MAX_PARTS:
                        capped = True
                        break
                    _open()
                zf.write(pth, rel, compress_type=zipfile.ZIP_STORED)
                size += fsz + 512
            if capped:
                break
            in_part.append(o["id"])
            size += 8192
        if zf is not None:
            _ship(last=True)
        if part_no == 0:                  # nothing waited on an admin
            _open()
            _ship(last=True)

        lost = {k: v for k, v in failed.items()}
        label = {f["key"]: f"{o['folder']}/{f['base']}" for o in objs for f in o.get("files") or []}
        lines = [
            f"Objections waiting on an admin — sent: {total} objections in {part_no} ZIP part(s).",
            f"Photos and files fetched: {len(names)}; could not be fetched: {len(lost)}.",
            "Unzip ALL parts into ONE folder, then read README.md and objections.md.",
        ]
        if capped:
            lines.append(f"⚠ NOT COMPLETE — stopped after {MAX_PARTS} parts.")
        for k, v in list(lost.items())[:ERRORS_NAMED]:
            lines.append(f"  · {label.get(k, k)}: {v}")
        if len(lost) > ERRORS_NAMED:
            lines.append(f"  · …and {len(lost) - ERRORS_NAMED} more (listed in objections.json).")
        _say(chat_id, "\n".join(lines))
        sent += 1
        _save_progress({"closed": True, "parts": part_no, "started": prog.get("started"),
                        "objections": total, "fetched": len(names), "failed": len(lost)})
    finally:
        shutil.rmtree(tmp, ignore_errors=True)
    return sent
