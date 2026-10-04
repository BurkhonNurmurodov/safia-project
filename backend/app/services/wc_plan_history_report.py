"""One-off: did work centre A1437 EVER carry a plan on Suvonov Elshod OF's page?

The operator asked on 2026-10-04, looking at the «Jamoalar» card of A1437 on
«Zagruzka fayli» (O.soni —, Shtat 0, 0 min). Production is not readable from the
machine that wrote this, so the answer is DMed from the server.

«Plan» has three meanings on that page, and the report answers each one
separately, because they can disagree:

  1. SAP QUANTITIES — `pp_daily` rows of this unit at this work centre (what the
     фаза file wrote, or what somebody typed over it).
  2. MINUTES — `zagruzka_source.wc_labor`, the page's own trudoyomkost: quantity ×
     Трудоемкость of the unit's catalog lines AT this work centre. No catalog line
     there means 0 minutes, whatever quantities are stored.
  3. CATALOG — whether the unit's catalog (today's, or any frozen version of it)
     ever held a line at this work centre.

Plus the context that explains a mismatch: the unit's work-centre register row
(an upload reaches every work centre the register names), the typed people pins,
what the plant-wide фаза file itself held at this code, which OTHER units carry
the code (catalog, register, cells), and the Jurnal rows that touched it.

It READS and writes nothing but its flag (written by `startup._send_report_once`).

Temporary: delete this module, `startup.report_wc_plan_history` /
`_wc_plan_history_job` and the call in BOTH entrypoints once it has been sent.
"""
from __future__ import annotations

import time
from collections import defaultdict
from datetime import date, datetime, timezone
from io import BytesIO
from zoneinfo import ZoneInfo

import requests
from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Font
from sqlalchemy import Text, cast, func, or_, text as sql_text
from sqlalchemy.orm import Session

from app.config import settings
from app.models import (
    ActionLog, Cell, Manager, PPCatalogVersion, PPDaily, PPLineDaily,
    PPManagerSetting, PPProduct, PPWorkCenter, PPWorkCenterDaily, RoleProfile,
)
from app.services import zagruzka_source
from app.services.cell_lookup import norm_code

UNIT_NAME = "Suvonov Elshod OF"
WORK_CENTER = "A1437"
TZ = ZoneInfo("Asia/Tashkent")
_API = "https://api.telegram.org"
SEND_RETRIES = 3
_JURNAL_MAX = 400


def _f(v) -> float:
    try:
        return float(v or 0)
    except (TypeError, ValueError):
        try:
            return float(str(v).replace(" ", "").replace(",", "."))
        except (TypeError, ValueError):
            return 0.0


def _g(v) -> str:
    """A quantity as it reads: 240 not 240.0, «—» for nothing."""
    if v is None:
        return "—"
    v = float(v)
    return f"{v:,.0f}".replace(",", " ") if v == int(v) else f"{v:,.1f}".replace(",", " ")


def _dt(v) -> str:
    if v is None:
        return "—"
    v = v if v.tzinfo else v.replace(tzinfo=timezone.utc)
    return v.astimezone(TZ).strftime("%d.%m.%Y %H:%M")


def _d(v) -> str:
    if v is None:
        return "—"
    if isinstance(v, str):
        v = date.fromisoformat(v[:10])
    return v.strftime("%d.%m.%Y")


def _unit(db: Session, name: str) -> Manager:
    """By NAME, never by id — ids differ between a checkout and production."""
    m = db.query(Manager).filter(Manager.name == name).first()
    if m is None:
        m = (db.query(Manager).filter(func.lower(Manager.name) == name.lower())
             .first())
    if m is None:
        known = ", ".join(sorted(x.name for x in db.query(Manager).all()
                                 if x.name and "suvonov" in x.name.lower()))
        raise RuntimeError(f"unit «{name}» not found (Suvonov units: {known or 'none'})")
    return m


def _names(db: Session) -> dict[int, str]:
    return {m.id: m.name for m in db.query(Manager).all()}


def _auto_fill(db: Session, mids) -> dict[int, bool]:
    rows = {r.manager_id: bool(r.auto_fill)
            for r in db.query(PPManagerSetting)
            .filter(PPManagerSetting.manager_id.in_(list(mids) or [-1])).all()}
    return {m: rows.get(m, True) for m in mids}     # absent row = ON


def _raw_file(db: Session, wc: str) -> dict[str, dict]:
    """What the PLANT-WIDE фаза file itself held at this code, per stored date.

    Read in SQL (`jsonb_array_elements`), so the thousands of rows of every
    stored file never come into Python. The column is found by its header
    («Команда»), falling back to the stored layout."""
    out: dict[str, dict] = {}
    q = sql_text("""
        SELECT u.date, u.columns, r
        FROM pp_uploads u, jsonb_array_elements(u.rows) r
        WHERE u.manager_id IS NULL AND u.file_type = 'faza'
          AND upper(replace(coalesce(
                r ->> (CASE WHEN u.columns ? 'Команда'
                            THEN (SELECT i - 1 FROM jsonb_array_elements_text(u.columns)
                                  WITH ORDINALITY AS c(h, i) WHERE h = 'Команда' LIMIT 1)
                            ELSE 2 END)::int, ''), ' ', '')) = :wc
    """)
    for day, cols, row in db.execute(q, {"wc": wc}):
        cols = list(cols or [])

        def at(header, fallback):
            i = cols.index(header) if header in cols else fallback
            return row[i] if isinstance(row, list) and i < len(row) else None

        k = day.isoformat()
        g = out.setdefault(k, {"ops": 0, "plan": 0.0, "orders": set(), "skus": set(),
                               "names": set()})
        g["ops"] += 1
        g["plan"] += _f(at("План", 5))
        if at("Заказ", 0):
            g["orders"].add(str(at("Заказ", 0)))
        if at("SKU", 3):
            g["skus"].add(str(at("SKU", 3)))
        if at("Наименование", 4):
            g["names"].add(str(at("Наименование", 4)).strip())
    return out


def collect(db: Session, unit_name: str = UNIT_NAME, wc: str = WORK_CENTER) -> dict:
    """Everything the message and the workbook say, as plain data."""
    wc = norm_code(wc)
    m = _unit(db, unit_name)
    names = _names(db)
    today = datetime.now(TZ).date()

    def here(col):
        return func.upper(func.replace(col, " ", "")) == wc

    # ── 1. the SAP quantities stored for this unit at this work centre ───────
    daily = (db.query(PPDaily)
             .filter(PPDaily.manager_id == m.id, here(PPDaily.work_center))
             .order_by(PPDaily.date, PPDaily.sap_code).all())
    sap_rows = [{
        "date": r.date, "sku": r.sap_code, "wc": r.work_center,
        "plan_qty": _f(r.plan_qty), "fact_qty": _f(r.actual_qty),
        "plan_ovr": None if r.plan_override is None else _f(r.plan_override),
        "fact_ovr": None if r.actual_override is None else _f(r.actual_override),
        "updated": r.updated_at,
    } for r in daily]
    for s in sap_rows:
        s["plan"] = s["plan_ovr"] if s["plan_ovr"] is not None else s["plan_qty"]
        s["fact"] = s["fact_ovr"] if s["fact_ovr"] is not None else s["fact_qty"]

    line_rows = [{
        "date": r.date, "key": r.qty_key, "line": r.line_key,
        "plan_ovr": None if r.plan_override is None else _f(r.plan_override),
        "fact_ovr": None if r.actual_override is None else _f(r.actual_override),
        "updated": r.updated_at,
    } for r in (db.query(PPLineDaily)
                .filter(PPLineDaily.manager_id == m.id, here(PPLineDaily.work_center))
                .order_by(PPLineDaily.date).all())]

    # ── 2. the page's own minutes ─────────────────────────────────────────────
    first_any = db.query(func.min(PPDaily.date)).filter(PPDaily.manager_id == m.id).scalar()
    lo = min([d for d in (first_any, sap_rows[0]["date"] if sap_rows else None) if d]
             or [today])
    minutes: dict[str, tuple[float, float]] = {}
    minutes_err = None
    try:
        for (_mid, dk, w), (pm, am) in zagruzka_source.wc_labor(db, [m.id], lo, today).items():
            if norm_code(w) == wc:
                prev = minutes.get(dk, (0.0, 0.0))
                minutes[dk] = (prev[0] + pm, prev[1] + am)
    except Exception as exc:                       # the rest of the report stands
        db.rollback()
        minutes_err = f"{type(exc).__name__}: {exc}"[:200]

    # ── 3. the catalog: today's, and every frozen version ────────────────────
    cur_lines = [{
        "from": "current", "sku": p.sap_code, "name": p.name, "wc": p.work_center,
        "op": p.op, "labor": None if p.labor_time is None else _f(p.labor_time),
        "active": bool(p.active), "auto": bool(p.auto_fill), "grp": p.wc_group,
        "created": p.created_at,
    } for p in (db.query(PPProduct)
                .filter(PPProduct.manager_id == m.id, here(PPProduct.work_center))
                .order_by(PPProduct.sort_order, PPProduct.id).all())]
    versions = []
    for v in (db.query(PPCatalogVersion).filter(PPCatalogVersion.manager_id == m.id)
              .order_by(PPCatalogVersion.valid_to).all()):
        lines = [d for d in (v.lines or []) if norm_code(d.get("work_center")) == wc]
        wcs = [d for d in (v.work_centers or []) if norm_code(d.get("code")) == wc]
        versions.append({"from": v.valid_from, "to": v.valid_to, "reason": v.reason,
                         "created": v.created_at, "n_lines": len(v.lines or []),
                         "lines": lines, "wc": wcs[0] if wcs else None})

    wc_row = (db.query(PPWorkCenter)
              .filter(PPWorkCenter.manager_id == m.id, here(PPWorkCenter.code)).first())
    register = None if wc_row is None else {
        "code": wc_row.code, "shtatka": wc_row.shtatka,
        "capacity": None if wc_row.capacity is None else _f(wc_row.capacity),
        "active": bool(wc_row.active)}

    # ── 4. typed people pins ──────────────────────────────────────────────────
    pins = [{"date": p.date, "grp": p.wc_group, "people": p.people,
             "shtatka": p.shtatka, "updated": p.updated_at}
            for p in (db.query(PPWorkCenterDaily)
                      .filter(PPWorkCenterDaily.manager_id == m.id,
                              here(PPWorkCenterDaily.work_center))
                      .order_by(PPWorkCenterDaily.date).all())]

    # ── 5. the plant-wide file itself ─────────────────────────────────────────
    raw_err = None
    try:
        raw = _raw_file(db, wc)
    except Exception as exc:
        db.rollback()
        raw, raw_err = {}, f"{type(exc).__name__}: {exc}"[:200]

    # ── 6. who else carries the code ──────────────────────────────────────────
    others: dict[int, dict] = defaultdict(lambda: {
        "lines": 0, "lines_active": 0, "register": None, "cells": [],
        "sap_days": 0, "sap_plan": 0.0, "first": None, "last": None})
    for p in db.query(PPProduct).filter(here(PPProduct.work_center)).all():
        o = others[p.manager_id]
        o["lines"] += 1
        o["lines_active"] += 1 if p.active else 0
    for w in db.query(PPWorkCenter).filter(here(PPWorkCenter.code)).all():
        others[w.manager_id]["register"] = f"shtat {w.shtatka}" + ("" if w.active else " (off)")
    leaders = {p.id: p.name for p in db.query(RoleProfile).all()}
    cells_all = db.query(Cell).filter(here(Cell.sap_code)).order_by(Cell.verifix_code).all()
    for c in cells_all:
        if c.manager_id is None:
            continue
        others[c.manager_id]["cells"].append(
            c.verifix_code + (f"·{c.wc_group}" if c.wc_group else "")
            + (f" ({leaders.get(c.leader_id)})" if c.leader_id else ""))
    for mid, n, plan, a, b in (
            db.query(PPDaily.manager_id, func.count(func.distinct(PPDaily.date)),
                     func.sum(func.coalesce(PPDaily.plan_override, PPDaily.plan_qty)),
                     func.min(PPDaily.date), func.max(PPDaily.date))
            .filter(here(PPDaily.work_center)).group_by(PPDaily.manager_id).all()):
        o = others[mid]
        o["sap_days"], o["sap_plan"], o["first"], o["last"] = n, _f(plan), a, b
    af = _auto_fill(db, others.keys() | {m.id})
    units = [{"id": mid, "name": names.get(mid, f"#{mid}"), "auto_fill": af.get(mid, True),
              **o} for mid, o in sorted(others.items(), key=lambda kv: kv[0] != m.id)]
    unplaced = [c.verifix_code for c in cells_all if c.manager_id is None]

    # ── 7. the Jurnal ─────────────────────────────────────────────────────────
    pat = f"%{wc}%"
    hit = or_(cast(ActionLog.details, Text).ilike(pat),
              cast(ActionLog.changes, Text).ilike(pat),
              ActionLog.target_id.ilike(pat), ActionLog.target_name.ilike(pat))
    jur = (db.query(ActionLog)
           .filter(hit, or_(ActionLog.unit_id == m.id, ActionLog.unit_id.is_(None),
                            ActionLog.action.ilike("production.%"),
                            ActionLog.action.ilike("%cell%")))
           .order_by(ActionLog.created_at.desc()).limit(_JURNAL_MAX).all())
    jurnal = [{"at": r.created_at, "action": r.action, "outcome": r.outcome,
               "who": r.actor_name, "role": r.actor_role,
               "unit": r.unit_name or names.get(r.unit_id), "day": r.day, "target": r.target_name or r.target_id,
               "changes": "; ".join(
                   f"{c[0]}: {c[1] if c[1] is not None else '—'} → "
                   f"{c[2] if c[2] is not None else '—'}"
                   for c in (r.changes or []) if isinstance(c, list) and len(c) >= 3)[:500]}
              for r in reversed(jur)]

    # ── per-day fold ──────────────────────────────────────────────────────────
    days: dict[str, dict] = {}

    def day(k):
        return days.setdefault(k, {"skus": 0, "plan": 0.0, "fact": 0.0, "typed": 0,
                                   "pmin": None, "amin": None, "people": None,
                                   "raw_ops": 0, "raw_plan": None})
    for s in sap_rows:
        g = day(s["date"].isoformat())
        g["skus"] += 1
        g["plan"] += s["plan"]
        g["fact"] += s["fact"]
        g["typed"] += (s["plan_ovr"] is not None) + (s["fact_ovr"] is not None)
    for k, (pm, am) in minutes.items():
        g = day(k)
        g["pmin"], g["amin"] = pm, am
    for p in pins:
        g = day(p["date"].isoformat())
        if p["people"] is not None:
            g["people"] = (g["people"] or 0) + p["people"]
    for k, r in raw.items():
        g = day(k)
        g["raw_ops"], g["raw_plan"] = r["ops"], r["plan"]

    plan_days = sorted(k for k, s in days.items() if s["plan"] > 0)
    min_days = sorted(k for k, s in days.items() if (s["pmin"] or 0) > 0)
    raw_days = sorted(k for k, s in days.items() if (s["raw_plan"] or 0) > 0)
    ever_line = bool(cur_lines) or any(v["lines"] for v in versions)

    return {
        "unit": m.name, "unit_id": m.id, "shift": m.shift, "wc": wc,
        "auto_fill": af.get(m.id, True), "today": today,
        "register": register, "sap_rows": sap_rows, "line_rows": line_rows,
        "minutes_err": minutes_err, "raw_err": raw_err,
        "cur_lines": cur_lines, "versions": versions, "ever_line": ever_line,
        "pins": pins, "raw": raw, "units": units, "unplaced": unplaced,
        "jurnal": jurnal, "days": dict(sorted(days.items())),
        "plan_days": plan_days, "min_days": min_days, "raw_days": raw_days,
        "range_from": lo,
    }


def _span(keys: list[str]) -> str:
    if not keys:
        return "never"
    return f"{len(keys)} days, {_d(keys[0])} → {_d(keys[-1])}"


def text(rep: dict) -> str:
    wc, unit = rep["wc"], rep["unit"]
    sap = rep["sap_rows"]
    L = [f"{wc} on «{unit}» (unit #{rep['unit_id']}, shift {rep['shift'] or '—'})",
         f"Read {datetime.now(TZ).strftime('%d.%m.%Y %H:%M')}; stored days from "
         f"{_d(rep['range_from'])}.", ""]

    L.append("ANSWER")
    L.append(f"1. SAP plan quantities stored for this unit at {wc}: {_span(rep['plan_days'])}"
             + (f" · Σ plan {_g(sum(s['plan'] for s in sap))}, Σ fact "
                f"{_g(sum(s['fact'] for s in sap))} pcs" if sap else ""))
    if rep["minutes_err"]:
        L.append(f"2. Plan MINUTES (the card's «min»): could not be computed — {rep['minutes_err']}")
    else:
        tot = sum((d["pmin"] or 0) for d in rep["days"].values())
        L.append(f"2. Plan MINUTES (the card's «min»): {_span(rep['min_days'])}"
                 + (f" · Σ {_g(tot)} min" if rep["min_days"] else ""))
    nver = sum(1 for v in rep["versions"] if v["lines"])
    L.append(f"3. Catalog line at {wc}: "
             + (f"YES — {len(rep['cur_lines'])} in today's catalog, in {nver} of "
                f"{len(rep['versions'])} frozen versions" if rep["ever_line"]
                else f"NEVER — not in today's catalog, not in any of "
                     f"{len(rep['versions'])} frozen versions"))
    typed = [p for p in rep["pins"] if p["people"] is not None]
    L.append(f"4. People typed («Bugungi fakt») at {wc}: "
             + (f"{len(typed)} days, last {_d(typed[-1]['date'])}" if typed else "never"))
    if rep["raw_err"]:
        L.append(f"5. Plant-wide SAP file: could not be read — {rep['raw_err']}")
    else:
        L.append(f"5. Plant-wide SAP file had operations at {wc}: {_span(rep['raw_days'])}")
    L.append("")

    reg = rep["register"]
    L.append("WHY THE ROWS ARE THERE")
    L.append(f"• Unit's work-centre register: "
             + (f"{reg['code']} · shtat {reg['shtatka']}"
                + (f" · capacity {_g(reg['capacity'])}" if reg["capacity"] else "")
                + ("" if reg["active"] else " · OFF")
                if reg else f"{wc} is NOT in it"))
    L.append(f"• SAP auto-fill for this unit: {'ON' if rep['auto_fill'] else 'OFF (manual)'}")
    L.append("  An upload writes every фаза operation at a work centre the register "
             "names, for SKUs the unit's catalog carries anywhere — so quantities can "
             "land at a work centre that has no catalog line, and then read 0 min.")
    L.append("")

    L.append(f"WHO ELSE CARRIES {wc}")
    for u in rep["units"]:
        mark = " ◀ this unit" if u["id"] == rep["unit_id"] else ""
        bits = [f"catalog {u['lines_active']}/{u['lines']} lines" if u["lines"] else "no catalog line",
                f"register {u['register']}" if u["register"] else "not in register",
                ("cells " + ", ".join(u["cells"])) if u["cells"] else "no cell",
                (f"SAP rows {u['sap_days']} days ({_d(u['first'])} → {_d(u['last'])}), "
                 f"Σ plan {_g(u['sap_plan'])}") if u["sap_days"] else "no SAP rows",
                "auto-fill " + ("ON" if u["auto_fill"] else "OFF")]
        L.append(f"• {u['name']}{mark}: " + " · ".join(bits))
    if rep["unplaced"]:
        L.append(f"• cells with no unit: {', '.join(rep['unplaced'])}")
    L.append("")

    L.append("LAST 14 DAYS WITH ANY TRACE (plan pcs · fact pcs · plan min · people · file ops)")
    rows = [(k, v) for k, v in rep["days"].items()
            if v["skus"] or v["pmin"] or v["people"] is not None or v["raw_ops"]]
    for k, v in rows[-14:]:
        L.append(f"{_d(k)}: {_g(v['plan'] if v['skus'] else None)} · "
                 f"{_g(v['fact'] if v['skus'] else None)} · {_g(v['pmin'])} · "
                 f"{_g(v['people'])} · {v['raw_ops'] or '—'}"
                 + (" · typed" if v["typed"] else ""))
    if not rows:
        L.append("none")
    L.append("")
    L.append(f"Jurnal rows naming {wc}: {len(rep['jurnal'])}"
             + (f" (newest {_JURNAL_MAX})" if len(rep["jurnal"]) >= _JURNAL_MAX else "")
             + ". Everything per day, per SKU, the catalog versions and the Jurnal "
               "are in the attached workbook.")
    return "\n".join(L)


def _cell(v):
    if isinstance(v, str):
        return ILLEGAL_CHARACTERS_RE.sub("", v)
    if isinstance(v, datetime):
        v = v if v.tzinfo else v.replace(tzinfo=timezone.utc)
        return v.astimezone(TZ).replace(tzinfo=None)
    return v


def _sheet(wb: Workbook, title: str, head: list[str], rows: list[list]) -> None:
    ws = wb.create_sheet(title)
    ws.append(head)
    for c in ws[1]:
        c.font = Font(bold=True)
    for r in rows:
        ws.append([_cell(v) for v in r])
    ws.freeze_panes = "A2"
    for i, h in enumerate(head, start=1):
        width = max([len(str(h))] + [len(str(r[i - 1])) for r in rows[:300]
                                     if i - 1 < len(r) and r[i - 1] is not None])
        ws.column_dimensions[ws.cell(1, i).column_letter].width = min(max(width + 2, 9), 60)


def build_workbook(rep: dict) -> BytesIO:
    wb = Workbook()
    wb.remove(wb.active)
    _sheet(wb, "Kunlik", [
        "Sana", "SAP qatorlar (SKU)", "Plan, dona", "Fakt, dona", "Qo'lda kiritilgan",
        "Plan, daq", "Fakt, daq", "Odam soni (kiritilgan)", "Fayldagi operatsiyalar",
        "Fayldagi plan, dona"],
        [[date.fromisoformat(k), v["skus"] or None,
          v["plan"] if v["skus"] else None, v["fact"] if v["skus"] else None,
          v["typed"] or None, v["pmin"], v["amin"], v["people"],
          v["raw_ops"] or None, v["raw_plan"]] for k, v in rep["days"].items()])
    _sheet(wb, "SAP qatorlar", [
        "Sana", "SKU", "Ish markazi", "Plan (fayl)", "Fakt (fayl)", "Plan (qo'lda)",
        "Fakt (qo'lda)", "Plan (amalda)", "Fakt (amalda)", "O'zgargan"],
        [[s["date"], s["sku"], s["wc"], s["plan_qty"], s["fact_qty"], s["plan_ovr"],
          s["fact_ovr"], s["plan"], s["fact"], s["updated"]] for s in rep["sap_rows"]])
    _sheet(wb, "Qator bo'yicha", [
        "Sana", "SKU / kalit", "Qator", "Plan (qo'lda)", "Fakt (qo'lda)", "O'zgargan"],
        [[r["date"], r["key"], r["line"], r["plan_ovr"], r["fact_ovr"], r["updated"]]
         for r in rep["line_rows"]])
    cat = [["bugungi", "", s["sku"], s["name"], s["wc"], s["op"], s["labor"],
            "ha" if s["active"] else "yo'q", "ha" if s["auto"] else "yo'q", s["grp"],
            s["created"]] for s in rep["cur_lines"]]
    for v in rep["versions"]:
        for d in v["lines"]:
            cat.append([_d(v["from"]) if v["from"] else "boshidan", _d(v["to"]),
                        d.get("sap_code"), d.get("name"), d.get("work_center"), d.get("op"),
                        d.get("labor_time"), "ha" if d.get("active", True) else "yo'q",
                        "ha" if d.get("auto_fill", True) else "yo'q", d.get("wc_group"), None])
    _sheet(wb, "Katalog", [
        "Versiya dan", "Versiya gacha", "SKU", "Nomi", "Ish markazi", "Opr.",
        "Trudoyomkost, s", "Faol", "SAP avto", "Guruh", "Yaratilgan"], cat)
    _sheet(wb, "Katalog versiyalari", [
        "Dan", "Gacha", "Sabab", "Yaratilgan", "Jami qatorlar",
        f"{rep['wc']} qatorlari", f"{rep['wc']} reestrda"],
        [[v["from"], v["to"], v["reason"], v["created"], v["n_lines"], len(v["lines"]),
          (f"shtat {v['wc'].get('shtatka')}" if v["wc"] else "yo'q")]
         for v in rep["versions"]])
    _sheet(wb, "Odam soni", ["Sana", "Guruh", "Odam", "Shtat", "O'zgargan"],
           [[p["date"], p["grp"], p["people"], p["shtatka"], p["updated"]]
            for p in rep["pins"]])
    _sheet(wb, "SAP fayl", ["Sana", "Operatsiyalar", "Plan, dona", "Buyurtmalar", "SKU",
                            "Nomlar"],
           [[date.fromisoformat(k), r["ops"], r["plan"], len(r["orders"]),
             ", ".join(sorted(r["skus"])), " | ".join(sorted(r["names"]))[:500]]
            for k, r in sorted(rep["raw"].items())])
    _sheet(wb, "Boshqa bo'limlar", [
        "Bo'lim", "Katalog qatorlari (faol/jami)", "Reestr", "Yacheykalar",
        "SAP kunlar", "Σ plan", "Birinchi", "Oxirgi", "SAP avto"],
        [[u["name"] + (" ◀" if u["id"] == rep["unit_id"] else ""),
          f"{u['lines_active']}/{u['lines']}", u["register"], ", ".join(u["cells"]),
          u["sap_days"], u["sap_plan"], u["first"], u["last"],
          "ha" if u["auto_fill"] else "yo'q"] for u in rep["units"]])
    _sheet(wb, "Jurnal", ["Vaqt", "Amal", "Natija", "Kim", "Rol", "Bo'lim", "Kun",
                          "Nishon", "O'zgarishlar"],
           [[j["at"], j["action"], j["outcome"], j["who"], j["role"], j["unit"], j["day"],
             j["target"], j["changes"]] for j in rep["jurnal"]])
    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


def _post(method: str, data: dict, files: dict | None = None) -> None:
    last = ""
    for attempt in range(1, SEND_RETRIES + 1):
        wait = 0
        try:
            r = requests.post(f"{_API}/bot{settings.telegram_bot_token}/{method}",
                              data=data, files=files, timeout=180)
            body = r.json()
            if body.get("ok"):
                return
            last = body.get("description") or f"HTTP {r.status_code}"
            wait = int(((body.get("parameters") or {}).get("retry_after")) or 0)
        except Exception as exc:
            last = type(exc).__name__
        if attempt < SEND_RETRIES:
            time.sleep(max(wait, 3 * attempt))
    raise RuntimeError(f"{method} failed: {last}"[:300])


def send(db: Session, chat_id: int, *_window) -> int:
    if not settings.telegram_bot_token:
        raise RuntimeError("telegram bot token not configured")
    rep = collect(db)
    body = text(rep)
    book = build_workbook(rep).getvalue()
    db.rollback()
    chunks, cur = [], ""
    for line in body.split("\n"):
        if len(cur) + len(line) + 1 > 3900:
            chunks.append(cur)
            cur = ""
        cur += line + "\n"
    chunks.append(cur)
    for c in chunks:
        _post("sendMessage", {"chat_id": chat_id, "text": c})
    stamp = datetime.now(TZ).strftime("%d.%m.%Y")
    _post("sendDocument", {"chat_id": chat_id,
                           "caption": f"{rep['wc']} · {rep['unit']} — full history"},
          files={"document": (f"{rep['wc']}-{rep['unit'].replace(' ', '-')}-{stamp}.xlsx",
                              book,
                              "application/vnd.openxmlformats-officedocument."
                              "spreadsheetml.sheet")})
    return len(chunks) + 1
