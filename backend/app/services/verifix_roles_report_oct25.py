"""TEMPORARY boot report (2026-10-25): every shift-manager, brigadir and leader
AS VERIFIX LISTS THEM, with the shift each works.

The operator (2026-10-25): «report me the names, shifts of all shift-managers,
supervisors (brigadirs), and leaders from the verifix (not from our platform)».
Nothing on this platform is read but the Verifix login: the people, their jobs,
their org units and schedules come from Verifix's API, fresh, and so does the
shift — from the last `DAYS_BACK` days of «Отчёт по посещениям» (when each
person actually came in), else the planned start of their schedule.

- **Who**: a WORKING employee (status W, or U — taken on, not started yet)
  whose Verifix job is a leader's (`verifix_org.is_leader_job`), a brigadir's
  (`verifix_org.is_supervisor_job`), or a shift-manager's (`_is_shift_manager`:
  «смен» + начальник / менеджер / мастер / руководитель / старший …).
  Assistants, deputies and «и.о.» are left out and listed on their own sheet,
  so a job title the rules missed is visible.
- **Shift**: 1 (day) = arrivals between 04:00 and 15:00, 2 (night) = the rest,
  over the days the person came; one side needs ≥ 75 % or the person reads
  «both». Nobody came → the planned start of the days in the report → the
  schedule's name («08:00-20:00»).

READS and writes nothing but its flag (`startup.VERIFIX_ROLES_FLAG`). DMed to
`UNPRICED_DM_CHAT` as text in chunks plus a workbook. Delete this module,
`startup.report_verifix_roles` and both entrypoint calls once it has landed.
"""
from __future__ import annotations

import re
import time
from collections import Counter, defaultdict
from datetime import datetime, timedelta
from io import BytesIO
from typing import Optional
from zoneinfo import ZoneInfo

import requests
from sqlalchemy.orm import Session

from app.config import settings
from app.services import verifix, verifix_explore, verifix_live, verifix_org
from app.translit import transliterate

TZ = ZoneInfo("Asia/Tashkent")
_API = "https://api.telegram.org"
SEND_RETRIES = 3
BUDGET_S = 300.0
DAYS_BACK = 14
DAY_FROM_H, DAY_TO_H = 4, 15      # an arrival in [04:00, 15:00) is the day shift
MAJORITY = 0.75
STATUSES = ("W", "U")

ROLES = ("shift_manager", "brigadir", "leader")
ROLE_LABEL = {"shift_manager": "Shift-managers", "brigadir": "Brigadirs (supervisors)",
              "leader": "Leaders"}
ROLE_ONE = {"shift_manager": "Shift-manager", "brigadir": "Brigadir", "leader": "Leader"}

_SHIFT_RE = re.compile(r"смен|smena|shift", re.I)
_HEAD_RE = re.compile(r"начальн|менедж|мастер|руковод|старш|бошлиғ|бошлиг|boshlig|"
                      r"manager|head|master|supervisor", re.I)
_NOT_RE = re.compile(r"помощ|ассист|замест|\bзам\b|\bи\.?\s*о\b|yordamchi|o['’ʻ‘`]?rinbosar|"
                     r"assistant|deputy|acting", re.I)
# A job that names one of the roles at all — the «similar, left out» sheet.
_ROLE_WORD_RE = re.compile(r"лидер|lider|leader|бригадир|brigadir|смен|smena|shift", re.I)


def _is_shift_manager(job: str) -> bool:
    j = job or ""
    return bool(_SHIFT_RE.search(j) and _HEAD_RE.search(j)) and not _NOT_RE.search(j)


def role_of(job: str) -> Optional[str]:
    if verifix_org.is_leader_job(job):
        return "leader"
    if verifix_org.is_supervisor_job(job):
        return "brigadir"
    if _is_shift_manager(job):
        return "shift_manager"
    return None


def _shift_of_hour(h: int) -> int:
    return 1 if DAY_FROM_H <= h < DAY_TO_H else 2


def _verdict(n1: int, n2: int) -> Optional[str]:
    tot = n1 + n2
    if not tot:
        return None
    if n1 / tot >= MAJORITY:
        return "1"
    if n2 / tot >= MAJORITY:
        return "2"
    return "both"


SHIFT_LABEL = {"1": "Shift 1 (day)", "2": "Shift 2 (night)", "both": "Both shifts",
               None: "Shift unknown"}


# ── the read ──────────────────────────────────────────────────────────────────

def _timesheet(c: dict, ids: list[int], first, last, deadline: float) -> dict[str, dict]:
    """Per person: arrivals by shift, planned starts by shift, last arrival,
    the schedule name the report prints."""
    out: dict[str, dict] = {}
    if not ids:
        return out
    with verifix.client(c) as cl:
        for i in range(0, len(ids), 200):
            body = {"period_begin_date": verifix._dmy(first), "period_end_date": verifix._dmy(last),
                    "division_ids": [], "employee_ids": ids[i:i + 200]}
            for page in verifix.each_page(cl, "core/timesheet$export", body,
                                          limit=verifix.LIMIT_TIMESHEET, deadline=deadline):
                for r in page or []:
                    eid = str(r.get("employee_id") or "")
                    if not eid:
                        continue
                    rec = out.setdefault(eid, {"came": Counter(), "plan": Counter(), "last": None,
                                               "sched": "", "days": 0})
                    rec["sched"] = rec["sched"] or str(r.get("schedule_name") or "").strip()
                    for d in r.get("days") or []:
                        rec["days"] += 1
                        tin = verifix_live._dt(d.get("input_time"))
                        if tin:
                            rec["came"][_shift_of_hour(tin.hour)] += 1
                            if rec["last"] is None or tin > rec["last"]:
                                rec["last"] = tin
                        beg = verifix_live._dt(d.get("begin_time"))
                        if beg and str(d.get("day_kind") or "W") == "W":
                            rec["plan"][_shift_of_hour(beg.hour)] += 1
    return out


def collect(db: Session) -> dict:
    c = verifix_explore.config(db)
    db.rollback()   # no transaction held across the Verifix reads
    deadline = time.monotonic() + BUDGET_S
    emps, _ = verifix_explore._employees(c, deadline, force=True)
    divs, _ = verifix_explore._divisions(c, deadline, force=True)
    jobs, _ = verifix_explore._jobs(c, deadline, force=True)
    scheds, _ = verifix_explore._schedules(c, deadline, force=True)

    def jname(e):
        return (jobs.get(e.get("job") or "") or {}).get("name") or ""

    working = [e for e in emps.values() if (e.get("status") or "W") in STATUSES]
    job_count: Counter = Counter(jname(e) for e in working)
    picked, similar = [], defaultdict(list)
    for e in working:
        j = jname(e)
        r = role_of(j)
        if r:
            picked.append((r, e))
        elif _ROLE_WORD_RE.search(j):
            similar[j].append(e)

    today = datetime.now(TZ).date()
    first = today - timedelta(days=DAYS_BACK)
    ids = sorted({int(e["id"]) for _, e in picked if str(e["id"]).isdigit()})
    ts_err = None
    try:
        ts = _timesheet(c, ids, first, today, deadline)
    except verifix.VerifixError as exc:
        ts, ts_err = {}, f"{exc.code}: {exc}"

    def unit_of(e):
        u = divs.get(e.get("unit") or "") or divs.get(e.get("div") or "") or {}
        parent = divs.get(u.get("parent") or "") or {}
        dept = divs.get(e.get("div") or "") or {}
        return u, parent, dept

    rows = []
    for r, e in picked:
        u, parent, dept = unit_of(e)
        sname = (scheds.get(e.get("sched") or "") or {}).get("name") or ""
        t = ts.get(e["id"]) or {}
        came, plan = t.get("came") or Counter(), t.get("plan") or Counter()
        shift, basis = _verdict(came[1], came[2]), "arrivals"
        if shift is None:
            shift, basis = _verdict(plan[1], plan[2]), "planned days"
        if shift is None:
            b, _e = verifix_live._schedule_window(sname or t.get("sched") or "", today)
            if b is not None:
                shift, basis = str(_shift_of_hour(b.hour)), "schedule name"
            else:
                basis = "—"
        rows.append({
            "role": r, "shift": shift, "basis": basis,
            "name": e.get("name") or f"#{e['id']}", "latin": transliterate(e.get("name") or ""),
            "job": jname(e), "status": e.get("status") or "W",
            "unit": u.get("name") or "", "code": u.get("raw_code") or "",
            "workshop": parent.get("name") or "", "dept": dept.get("name") or "",
            "schedule": sname or t.get("sched") or "",
            "came1": came[1], "came2": came[2], "plan1": plan[1], "plan2": plan[2],
            "last": t.get("last"), "hired": e.get("hired") or "", "id": e["id"],
        })
    order = {"1": 0, "2": 1, "both": 2, None: 3}
    rows.sort(key=lambda x: (ROLES.index(x["role"]), order.get(x["shift"], 3), x["name"].lower()))
    return {
        "at": datetime.now(TZ), "first": first, "last": today, "ts_err": ts_err,
        "read": len(emps), "working": len(working), "rows": rows,
        "jobs": sorted(job_count.items(), key=lambda kv: (-kv[1], kv[0])),
        "similar": sorted(((j, ps) for j, ps in similar.items()), key=lambda kv: -len(kv[1])),
    }


# ── the text ──────────────────────────────────────────────────────────────────

def text(rep: dict) -> str:
    rows = rep["rows"]
    L = [f"VERIFIX — shift-managers, brigadirs and leaders · read {rep['at']:%d.%m.%Y %H:%M}",
         f"{rep['read']} employees read, {rep['working']} working. Shift = when each person came in, "
         f"{rep['first']:%d.%m}–{rep['last']:%d.%m} (day = arrival 04:00–15:00); "
         f"nobody came → planned start → schedule name.",
         "Names and units exactly as Verifix spells them. Full detail in the Excel file."]
    if rep["ts_err"]:
        L.append(f"⚠ The attendance report could not be read ({rep['ts_err']}) — shifts come "
                 f"from schedule names only.")
    L.append("")
    for r in ROLES:
        mine = [x for x in rows if x["role"] == r]
        by = Counter(x["shift"] for x in mine)
        L.append(f"=== {ROLE_LABEL[r].upper()} — {len(mine)} "
                 f"(shift 1: {by['1']} · shift 2: {by['2']} · both: {by['both']} · unknown: {by[None]}) ===")
        if not mine:
            L.append("  none found")
        for s in ("1", "2", "both", None):
            grp = [x for x in mine if x["shift"] == s]
            if not grp:
                continue
            L.append(f"— {SHIFT_LABEL[s]}: {len(grp)}")
            for x in grp:
                where = " ".join(p for p in (x["code"], x["unit"]) if p) or "—"
                tail = "" if s in ("1", "2") else f" · came d{x['came1']}/n{x['came2']}"
                extra = " · not started yet" if x["status"] == "U" else ""
                L.append(f"  · {x['name']} — {x['job']} · {where}{tail}{extra}")
        L.append("")
    if rep["similar"]:
        L.append("=== SIMILAR JOBS LEFT OUT (assistants, deputies, other «смена» jobs) ===")
        for j, ps in rep["similar"][:25]:
            L.append(f"  · {j}: {len(ps)}")
    return "\n".join(L)


# ── the workbook ──────────────────────────────────────────────────────────────

def workbook(rep: dict) -> bytes:
    from openpyxl import Workbook
    from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    head = Font(bold=True, color="FFFFFF")
    fill = PatternFill("solid", fgColor="C8973F")

    def sheet(ws, title, header, data, widths):
        ws.title = title
        ws.append(header)
        for c in ws[1]:
            c.font, c.fill = head, fill
            c.alignment = Alignment(vertical="center", wrap_text=True)
        for row in data:
            ws.append([ILLEGAL_CHARACTERS_RE.sub("", v) if isinstance(v, str) else v for v in row])
        for i, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = w
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions

    header = ["Role", "Shift", "Shift basis", "Name (Verifix)", "Name (Latin)", "Job (Verifix)",
              "Status", "Cell / org unit code", "Org unit", "Workshop (parent)", "Department",
              "Schedule", f"Came, day ({DAYS_BACK} d)", f"Came, night ({DAYS_BACK} d)",
              "Planned day starts", "Planned night starts", "Last arrival", "Hired", "Verifix id"]
    data = [[ROLE_ONE[x["role"]], SHIFT_LABEL[x["shift"]], x["basis"], x["name"], x["latin"], x["job"],
             "working" if x["status"] == "W" else "not started", x["code"], x["unit"], x["workshop"],
             x["dept"], x["schedule"], x["came1"], x["came2"], x["plan1"], x["plan2"],
             x["last"].strftime("%d.%m.%Y %H:%M") if x["last"] else "", x["hired"], x["id"]]
            for x in rep["rows"]]
    sheet(wb.active, "All three roles", header, data,
          [14, 16, 14, 34, 34, 30, 11, 12, 32, 28, 28, 26, 10, 10, 10, 10, 16, 12, 12])

    for r in ROLES:
        mine = [d for d, x in zip(data, rep["rows"]) if x["role"] == r]
        sheet(wb.create_sheet(), ROLE_LABEL[r][:31], header, mine,
              [14, 16, 14, 34, 34, 30, 11, 12, 32, 28, 28, 26, 10, 10, 10, 10, 16, 12, 12])

    counted = {x["job"]: ROLE_ONE[x["role"]] for x in rep["rows"]}
    sheet(wb.create_sheet(), "All jobs", ["Job (Verifix)", "Working people", "Counted as"],
          [[j or "—", n, counted.get(j, "")] for j, n in rep["jobs"]], [44, 14, 16])
    sheet(wb.create_sheet(), "Similar, left out", ["Job (Verifix)", "People", "Names"],
          [[j, len(ps), "; ".join(sorted(p.get("name") or "" for p in ps))] for j, ps in rep["similar"]],
          [44, 10, 120])
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ── send ──────────────────────────────────────────────────────────────────────

def _post(method: str, data: dict, files: dict | None = None) -> None:
    last = ""
    for attempt in range(1, SEND_RETRIES + 1):
        wait = 0
        try:
            r = requests.post(f"{_API}/bot{settings.telegram_bot_token}/{method}",
                              data=data, files=files, timeout=180)
            body = r.json()
            if body.get("ok"):
                return
            last = body.get("description") or f"HTTP {r.status_code}"
            wait = int(((body.get("parameters") or {}).get("retry_after")) or 0)
        except Exception as exc:  # noqa: BLE001
            last = type(exc).__name__
        if attempt < SEND_RETRIES:
            time.sleep(max(wait, 3 * attempt))
    raise RuntimeError(f"{method} failed: {last}"[:300])


def send(db: Session, chat_id: int, *_window) -> int:
    """The workbook first, then the text in chunks. Returns the messages sent."""
    if not settings.telegram_bot_token:
        raise RuntimeError("telegram bot token not configured")
    rep = collect(db)
    body = text(rep)
    print("[startup] Verifix roles report:\n" + body)
    stamp = rep["at"].strftime("%Y-%m-%d_%H-%M")
    n = Counter(x["role"] for x in rep["rows"])
    _post("sendDocument",
          {"chat_id": chat_id,
           "caption": (f"Verifix — {n['shift_manager']} shift-managers, {n['brigadir']} brigadirs, "
                       f"{n['leader']} leaders, with their shifts. Read straight from Verifix "
                       f"{rep['at']:%d.%m.%Y %H:%M}.")},
          files={"document": (f"verifix-roles-{stamp}.xlsx", workbook(rep),
                              "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})
    chunks, cur = [], ""
    for line in body.split("\n"):
        if len(cur) + len(line) + 1 > 3900:
            chunks.append(cur)
            cur = ""
        cur += line + "\n"
    chunks.append(cur)
    for ch in chunks:
        _post("sendMessage", {"chat_id": chat_id, "text": ch})
    return len(chunks) + 1
