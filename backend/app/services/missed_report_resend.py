"""Leader day reports since 16 Sep — which ones the leaders never got, listed
once, and sent on the operator's tap.

⚠ TEMPORARY (2026-09-23). Leaders reported that their day report never came
and that they could not object to the AI's verdicts: the DM is the only door a
leader has onto `/leaders/report/<uid>`, where «Norozilik bildirish» lives.
The cause was `leader_reports.send_for_uid` parking a report tried while its
day was still OPEN and never coming back to it (fixed the same day — see the
comment there and `leader_ai.sweep_unreported`). This module confirms it on
production data and repairs the backlog, in the two steps the operator chose:

1. `send` — a boot one-shot (`startup.report_missed_reports_resend`) that DMs
   the operator, for every bot leader-day from `DATE_FROM` to today: was the
   report received, and if not, why — plus an Excel naming which leaders missed
   which dates, and TWO buttons. The list behind the buttons is STORED
   (`LIST_KEY`) exactly as it was sent.
2. `apply` — a tap (`telegram_bot._mrr_callback`): «all» sends every listed
   report, «rej» only the days carrying an AI-rejected task (the ones an
   objection can be filed on). The sending runs as a background job at one DM a
   second and DMs the result when it is done.

The resent DM is the NORMAL day report, one per day, each with its «open the
report» button (the operator's choice: unchanged, no «sent late» line) — sent
through `leader_reports.send_for_uid` itself, so every rule a report obeys
(exclusions, cutoffs, rehearsal days, the ledger) is applied again at the moment
of sending, and a report sent since the list was built is never sent twice.
Leaders only: from `leader_unit_report.DIGEST_FROM` a brigadir's copy is a row
of the unit digest, which never depended on the leader's park.

Days from `leader_ai.OPEN_PARK_RETRY_FROM` (23 Sep) on are NOT on the list: the
fixed sweep sends those itself. They are still counted in the report, so the
confirmation covers today as well.

Temporary: delete this module, `startup.report_missed_reports_resend`,
`startup._missed_reports_resend_job`, the call in BOTH entrypoints and the
`mrr:` callback in `telegram_bot.py` once the operator has tapped (or declined).
A call left behind imports a deleted module at boot, and a failed boot rolls the
deploy back.
"""
from __future__ import annotations

import hashlib
import io
import json
import logging
import time
from collections import Counter, defaultdict
from datetime import datetime, timedelta, timezone

import requests
from sqlalchemy.orm import Session

from app.config import settings
from app.identity import profile_holders
from app.models import (
    ActionLog, AppSetting, Cell, LeaderAiDispute, LeaderAiReview, LeaderDayReport,
    LeaderTaskDay, LeaderTaskEntry, Manager, RoleProfile, TelegramUser,
)
from app.services import leader_ai, leader_cutoffs, leader_exclusions, leader_tasks

log = logging.getLogger(__name__)

TZ = timezone(timedelta(hours=5))                  # the plant's wall clock
DATE_FROM = "2026-09-16"                           # the operator's window: the last 7 days
RESEND_BEFORE = leader_ai.OPEN_PARK_RETRY_FROM     # from here the fixed sweep sends them
LIST_KEY = "missed_report_resend_list_2026_09_23"
CALLBACK = "mrr"
PARK_OPEN = "report no longer exists"              # what send_for_uid logged for an open day
LATE_H = 3                  # a report sent this long after its day closed counts as late
PACE_S = 1.1                # between two DMs — Telegram allows one a second per chat
RUN_STALE_MIN = 10          # a run silent this long (no heartbeat) died with its process
_BEAT_EVERY = 20            # items between two heartbeats
_FAIL_SLACK = timedelta(minutes=10)
_CHUNK = 800

_API = "https://api.telegram.org"
SEND_RETRIES = 3

# ── the verdict on one leader-day's report ───────────────────────────────────
RECEIVED = "received"
RECEIVED_LATE = f"received late ({LATE_H}h+ after the day closed)"
NO_ACCOUNT = "sent, but the leader has no Telegram account (bell only)"
BLOCKED = "sent, but Telegram refused it (bot blocked / never started)"
MISSED_PARKED = "NOT SENT: tried while the day was open, never retried"
MISSED_UNSENT = "NOT SENT: finished, never sent"
WAITING = "waiting for the AI review"
NO_PHOTO = "no report by design: no photo task was reviewed"
ON_PURPOSE = "not reported on purpose (excluded, cut off, …)"
NOT_AUTO = "outside the automatic regime"
OPEN = "day still open"
MISSED = (MISSED_PARKED, MISSED_UNSENT)
UNREACHED = (NO_ACCOUNT, BLOCKED)
ORDER = (RECEIVED, RECEIVED_LATE, MISSED_PARKED, MISSED_UNSENT, NO_ACCOUNT,
         BLOCKED, WAITING, NO_PHOTO, ON_PURPOSE, NOT_AUTO, OPEN)

TG_OK = "ok"
TG_NONE = "no account"
TG_BLOCKED = "bot blocked / never started"


def _aware(v):
    if v is None:
        return None
    return v if v.tzinfo else v.replace(tzinfo=timezone.utc)


def _t(v) -> str | None:
    v = _aware(v)
    return v.astimezone(TZ).strftime("%Y-%m-%d %H:%M") if v else None


def _dm(date: str) -> str:
    return f"{date[8:10]}.{date[5:7]}"


def _chunks(seq):
    seq = list(seq)
    for i in range(0, len(seq), _CHUNK):
        yield seq[i:i + _CHUNK]


def _closer(reason: str | None) -> str:
    r = reason or ""
    if r.startswith(leader_tasks.AUTO_PREFIX):
        return "automatic check"
    if r.startswith(leader_tasks.MISSED_PREFIX):
        return "deadline (nothing filed)"
    return "leader / deadline submit"


def _tg_state(tids: list[int], fails: dict[int, object], at) -> str:
    """Could the bot reach this leader at `at` (None = now)?

    `telegram_users.dm_failed_at` marks the FIRST failure of the current
    failing streak and is cleared by any success, so an account failing now
    whose streak began before `at` was failing at `at` too."""
    if not tids:
        return TG_NONE
    at = _aware(at)
    for tid in tids:
        f = _aware(fails.get(tid))
        if f is None or (at is not None and f > at + _FAIL_SLACK):
            return TG_OK
    return TG_BLOCKED


# ── step 1: what happened to every report ────────────────────────────────────

def collect(db: Session, now: datetime | None = None) -> dict:
    from app.services import leader_reports

    now = now or datetime.now(timezone.utc)
    today = now.astimezone(TZ).date().isoformat()

    D, E, R = LeaderTaskDay, LeaderTaskEntry, LeaderAiReview
    days = (db.query(D.id, D.leader_id, D.manager_id, D.date, D.cell_id,
                     D.closed_at, D.completion)
            .filter(D.date >= DATE_FROM, D.date <= today)
            .order_by(D.date, D.manager_id, D.leader_id, D.id).all())
    day_ids = [d.id for d in days]

    entries: dict[int, list] = defaultdict(list)
    for ch in _chunks(day_ids):
        for e in (db.query(E.id, E.day_id, E.task_id, E.reason, E.closed_at)
                  .filter(E.day_id.in_(ch)).all()):
            entries[e.day_id].append(e)
    ref_day = {f"bot:{e.id}": e.day_id for es in entries.values() for e in es}

    reviews: dict[int, list] = defaultdict(list)
    disputes: dict[int, int] = defaultdict(int)
    for ch in _chunks(ref_day):
        for r in (db.query(R.ref, R.task_id, R.status, R.attempts, R.resolution)
                  .filter(R.ref.in_(ch)).all()):
            reviews[ref_day[r.ref]].append(r)
        for d in (db.query(LeaderAiDispute.ref).filter(LeaderAiDispute.ref.in_(ch)).all()):
            disputes[ref_day[d.ref]] += 1

    keys = {f"bot:{i}": i for i in day_ids}
    uids = {f"bot-{i}": i for i in day_ids}
    ledger: dict[int, object] = {}
    for ch in _chunks(keys):
        for L in (db.query(LeaderDayReport.report_key, LeaderDayReport.sends,
                           LeaderDayReport.first_sent_at, LeaderDayReport.last_sent_at)
                  .filter(LeaderDayReport.report_key.in_(ch)).all()):
            ledger[keys[L.report_key]] = L
    parks: dict[int, tuple] = {}
    sends: dict[int, object] = {}
    for ch in _chunks(uids):
        for a in (db.query(ActionLog.action, ActionLog.target_id, ActionLog.reason,
                           ActionLog.created_at)
                  .filter(ActionLog.action.in_(("report.parked", "report.sent")),
                          ActionLog.target_id.in_(ch))
                  .order_by(ActionLog.created_at).all()):
            did = uids[a.target_id]
            if a.action == "report.parked":
                parks.setdefault(did, (a.reason, a.created_at))
            else:
                sends.setdefault(did, a.created_at)       # the first real send

    lids = {d.leader_id for d in days}
    names = {p.id: p.name for p in db.query(RoleProfile.id, RoleProfile.name)
             .filter(RoleProfile.id.in_(lids or [0])).all()}
    units = {m.id: (m.name, m.shift) for m in
             db.query(Manager.id, Manager.name, Manager.shift).all()}
    cells = {c.id: c.verifix_code for c in db.query(Cell.id, Cell.verifix_code).all()}
    holders: dict[int, list[int]] = {}
    for lid in lids:
        try:
            holders[lid] = profile_holders(db, f"leader:{lid}")
        except Exception:
            db.rollback()
            holders[lid] = []
    tids = {t for ts in holders.values() for t in ts}
    fails = {u.telegram_id: u.dm_failed_at for u in
             db.query(TelegramUser.telegram_id, TelegramUser.dm_failed_at)
             .filter(TelegramUser.telegram_id.in_(tids or [0])).all()}

    rows = []
    for d in days:
        unit, shift = units.get(d.manager_id, (f"#{d.manager_id}", None))
        name = names.get(d.leader_id, f"#{d.leader_id}")
        es = entries.get(d.id, [])
        rs = reviews.get(d.id, [])
        closed = [e for e in es if e.closed_at is not None]
        last = max(closed, key=lambda e: _aware(e.closed_at)) if closed else None
        unfinished = [r for r in rs if r.status in ("pending", "error")
                      and (r.attempts or 0) < leader_ai.MAX_ATTEMPTS]
        flagged = sorted({r.task_id for r in rs if r.status == "flagged"
                          and r.resolution != "approved"})
        L = ledger.get(d.id)
        park_reason, park_at = parks.get(d.id, (None, None))
        sent = L is not None and (L.sends or 0) > 0
        sent_at = sends.get(d.id) or (L.last_sent_at if sent else None)
        closed_at = _aware(d.closed_at)
        auto = leader_ai.in_auto_regime(d.date, shift)
        tg_now = _tg_state(holders.get(d.leader_id, []), fails, None)
        out = False

        if d.closed_at is None:
            verdict = OPEN
        elif not auto:
            verdict = NOT_AUTO
        elif sent:
            tg = _tg_state(holders.get(d.leader_id, []), fails, sent_at)
            if tg == TG_NONE:
                verdict = NO_ACCOUNT
            elif tg == TG_BLOCKED:
                verdict = BLOCKED
            elif sent_at and closed_at and _aware(sent_at) - closed_at > timedelta(hours=LATE_H):
                verdict = RECEIVED_LATE
            else:
                verdict = RECEIVED
        else:
            # The same two questions `send_for_uid` asks before it reports.
            out = (leader_exclusions.excluded(db, d.leader_id, d.date)
                   or leader_cutoffs.active(db, d.leader_id, name, d.date) is not None)
            if out or (park_reason and park_reason != PARK_OPEN):
                verdict = ON_PURPOSE
            elif unfinished:
                verdict = WAITING
            elif not rs:
                verdict = NO_PHOTO
            elif L is not None:
                verdict = MISSED_PARKED
            else:
                verdict = MISSED_UNSENT

        rows.append({
            "date": d.date, "shift": shift, "unit": unit, "manager_id": d.manager_id,
            "leader": name, "leader_id": d.leader_id,
            "cell": cells.get(d.cell_id) if d.cell_id else None,
            "day_id": d.id, "uid": f"bot-{d.id}", "key": f"bot:{d.id}",
            "verdict": verdict,
            "day_closed_at": _t(d.closed_at),
            "raw_score": round(float(d.completion)) if d.completion is not None else None,
            "tasks_closed": len(closed),
            "last_task": last.task_id if last else None,
            "last_task_closed_by": _closer(last.reason) if last else None,
            "ai_reviews": len(rs), "ai_waiting": len(unfinished),
            "ai_rejected_tasks": ", ".join(f"#{t}" for t in flagged),
            "ai_rejected": len(flagged),
            "objections": disputes.get(d.id, 0),
            "report_sent_at": _t(sent_at) if sent else None,
            "hours_close_to_report": (round((_aware(sent_at) - closed_at).total_seconds() / 3600, 1)
                                      if sent and sent_at and closed_at else None),
            "park_reason": park_reason, "parked_at": _t(park_at),
            "telegram_accounts": len(holders.get(d.leader_id, [])),
            "telegram_now": tg_now,
            "fixed_automatically": d.date >= RESEND_BEFORE,
        })

    # The list behind the buttons: the days nobody was told about, before the
    # fix's own reach. Each is read through the day report itself, so the
    # score and the rejected tasks it carries are the ones the DM will carry.
    items, unbuildable = [], []
    for r in rows:
        if r["verdict"] not in MISSED or r["date"] >= RESEND_BEFORE:
            continue
        try:
            rep = leader_reports.day_report(db, r["uid"])
        except Exception as exc:          # noqa: BLE001 - one day never sinks the list
            db.rollback()
            log.exception("missed-report resend: day report for %s", r["uid"])
            rep, why = None, f"{type(exc).__name__}: {str(exc)[:120]}"
        else:
            why = "no report could be built"
        if rep is None:
            unbuildable.append({**r, "why": why})
            continue
        rejected = [t["id"] for t in rep["tasks"] if t["ai_rejected"]]
        r["score"] = rep["score"]
        r["ai_rejected_tasks"] = ", ".join(f"#{t}" for t in rejected)
        r["ai_rejected"] = len(rejected)
        items.append({
            "uid": r["uid"], "key": r["key"], "date": r["date"], "shift": r["shift"],
            "unit": r["unit"], "leader": r["leader"], "leader_id": r["leader_id"],
            "cell": r["cell"], "score": rep["score"],
            "rejected_tasks": r["ai_rejected_tasks"], "has_rejected": bool(rejected),
            "objections": r["objections"], "why": r["verdict"],
            "telegram": r["telegram_now"],
        })

    return {"generated_at": _t(now), "date_from": DATE_FROM, "date_to": today,
            "resend_before": RESEND_BEFORE, "rows": rows, "items": items,
            "unbuildable": unbuildable, "summary": _summary(rows, items)}


def _summary(rows: list[dict], items: list[dict]) -> dict:
    closed = [r for r in rows if r["verdict"] not in (OPEN, NOT_AUTO)]
    by = Counter(r["verdict"] for r in rows)
    missed = [r for r in rows if r["verdict"] in MISSED]
    lost = [r for r in rows if r["verdict"] in MISSED + UNREACHED]
    per_date = []
    for (dt, sh) in sorted({(r["date"], r["shift"]) for r in rows},
                           key=lambda k: (k[0], k[1] or 0)):
        sub = [r for r in rows if r["date"] == dt and r["shift"] == sh]
        c = Counter(r["verdict"] for r in sub)
        per_date.append({
            "date": dt, "shift": sh, "days": len(sub),
            "closed": sum(1 for r in sub if r["verdict"] not in (OPEN, NOT_AUTO)),
            "received": c[RECEIVED] + c[RECEIVED_LATE], "late": c[RECEIVED_LATE],
            "missed": c[MISSED_PARKED] + c[MISSED_UNSENT],
            "unreached": c[NO_ACCOUNT] + c[BLOCKED],
            "waiting": c[WAITING], "no_photo": c[NO_PHOTO],
            "on_purpose": c[ON_PURPOSE], "open": c[OPEN],
        })
    leaders: dict[int, dict] = {}
    for r in rows:
        if r["verdict"] in (OPEN, NOT_AUTO):
            continue
        p = leaders.setdefault(r["leader_id"], {
            "leader": r["leader"], "unit": r["unit"], "shift": r["shift"],
            "closed": 0, "received": 0, "not_received": 0, "dates": [],
            "with_rejected": 0, "telegram": r["telegram_now"],
            "accounts": r["telegram_accounts"], "to_send": 0})
        p["closed"] += 1
        if r["verdict"] in (RECEIVED, RECEIVED_LATE):
            p["received"] += 1
        if r["verdict"] in MISSED + UNREACHED:
            p["not_received"] += 1
            p["dates"].append(_dm(r["date"]) + (f" {r['cell']}" if r["cell"] else ""))
            if r["ai_rejected"]:
                p["with_rejected"] += 1
    for it in items:
        if it["leader_id"] in leaders:
            leaders[it["leader_id"]]["to_send"] += 1
    return {
        "by_verdict": {v: by[v] for v in ORDER if by[v]},
        "closed": len(closed),
        "leaders_closed": len({r["leader_id"] for r in closed}),
        "leaders_missed": len({r["leader_id"] for r in lost}),
        "missed_rejected_days": sum(1 for r in missed if r["ai_rejected"]),
        "missed_rejected_tasks": sum(r["ai_rejected"] for r in missed),
        "missed_objections": sum(r["objections"] for r in missed),
        "missed_last_closed_by": dict(Counter(
            r["last_task_closed_by"] or "—" for r in missed).most_common()),
        "fixed_now": sum(1 for r in missed if r["fixed_automatically"]),
        "per_date": per_date,
        "leaders": sorted(leaders.values(),
                          key=lambda p: (-p["not_received"], p["unit"], p["leader"])),
        "to_send": len(items),
        "to_send_rejected": sum(1 for i in items if i["has_rejected"]),
        "to_send_leaders": len({i["leader_id"] for i in items}),
        "to_send_leaders_rejected": len({i["leader_id"] for i in items if i["has_rejected"]}),
        "to_send_unreachable_leaders": len({i["leader_id"] for i in items
                                            if i["telegram"] != TG_OK}),
    }


# ── the workbook ─────────────────────────────────────────────────────────────

ALL_COLS = [
    ("date", 11), ("shift", 6), ("unit", 24), ("leader", 26), ("cell", 7),
    ("verdict", 46), ("day_closed_at", 16), ("raw_score", 8), ("tasks_closed", 8),
    ("last_task", 7), ("last_task_closed_by", 22), ("ai_reviews", 8),
    ("ai_waiting", 8), ("ai_rejected", 8), ("ai_rejected_tasks", 16),
    ("objections", 8), ("report_sent_at", 16), ("hours_close_to_report", 9),
    ("park_reason", 26), ("parked_at", 16), ("telegram_accounts", 9),
    ("telegram_now", 22), ("fixed_automatically", 9), ("uid", 12),
    ("leader_id", 8), ("manager_id", 8),
]


def build_workbook(rep: dict) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    bold = Font(bold=True)
    head_fill = PatternFill("solid", fgColor="E8E3D6")
    miss_fill = PatternFill("solid", fgColor="FBD5D5")
    s = rep["summary"]

    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    ws.append([f"Leader day reports {_dm(rep['date_from'])} → {_dm(rep['date_to'])}"
               f" — did the leaders get them? (generated {rep['generated_at']} Tashkent)"])
    ws["A1"].font = Font(bold=True, size=13)
    ws.append([f"Closed leader-days: {s['closed']} ({s['leaders_closed']} leaders). "
               f"Days from {_dm(rep['resend_before'])} on are sent by the fix itself; "
               f"the days before it are on the «To send» sheet."])
    ws.append([])
    ws.append(["Verdict", "Leader-days"])
    for c in ws[ws.max_row]:
        c.font = bold
    for k, v in s["by_verdict"].items():
        ws.append([k, v])
    ws.append([])
    ws.append(["What closed the day last, on the days never sent", ""])
    ws[ws.max_row][0].font = bold
    for k, v in s["missed_last_closed_by"].items():
        ws.append([k, v])
    ws.append([])
    for label, val in (
            ("Leaders who did not get at least one report", s["leaders_missed"]),
            ("Never-sent days carrying AI-rejected tasks", s["missed_rejected_days"]),
            ("AI-rejected tasks on never-sent days", s["missed_rejected_tasks"]),
            ("Objections filed on never-sent days", s["missed_objections"]),
            ("On the «To send» list", s["to_send"]),
            ("…of them with AI-rejected tasks", s["to_send_rejected"])):
        ws.append([label, val])
    ws.column_dimensions["A"].width = 60
    ws.column_dimensions["B"].width = 14

    def table(name, header, data, widths, fill_if=None):
        t = wb.create_sheet(name)
        t.append(header)
        for c in t[1]:
            c.font = bold
            c.fill = head_fill
            c.alignment = Alignment(wrap_text=True, vertical="top")
        for vals in data:
            t.append(vals)
            if fill_if and fill_if(vals):
                for c in t[t.max_row]:
                    c.fill = miss_fill
        for i, w in enumerate(widths, start=1):
            t.column_dimensions[t.cell(row=1, column=i).column_letter].width = w
        t.freeze_panes = "A2"
        if t.max_row > 1:
            t.auto_filter.ref = t.dimensions
        return t

    pd = ["date", "shift", "days", "closed", "received", "late", "missed",
          "unreached", "waiting", "no_photo", "on_purpose", "open"]
    table("By date", ["date", "shift", "leader-days", "closed", "received",
                      "…of them late", "NOT sent", "sent, not reachable",
                      "waiting for AI", "no photo reviewed", "on purpose",
                      "still open"],
          [[r[c] for c in pd] for r in s["per_date"]],
          [11, 6, 10, 8, 9, 9, 9, 10, 9, 10, 9, 8], lambda v: v[6] or v[7])

    table("Leaders", ["leader", "brigadir", "shift", "closed days", "received",
                      "NOT received", "dates not received", "…with AI-rejected tasks",
                      "Telegram accounts", "Telegram now", "on the send list"],
          [[p["leader"], p["unit"], p["shift"], p["closed"], p["received"],
            p["not_received"], ", ".join(p["dates"]), p["with_rejected"],
            p["accounts"], p["telegram"], p["to_send"]] for p in s["leaders"]],
          [28, 26, 6, 8, 9, 9, 40, 10, 9, 22, 9], lambda v: v[5])

    table("To send", ["date", "shift", "brigadir", "leader", "cell", "verified score",
                      "AI-rejected tasks", "objections filed", "why not sent",
                      "Telegram now", "uid"],
          [[i["date"], i["shift"], i["unit"], i["leader"], i["cell"], i["score"],
            i["rejected_tasks"], i["objections"], i["why"], i["telegram"], i["uid"]]
           for i in rep["items"]],
          [11, 6, 26, 28, 7, 9, 16, 9, 46, 22, 12])

    keys = [c for c, _ in ALL_COLS]
    table("All days", keys, [[r.get(k) for k in keys] for r in rep["rows"]],
          [w for _, w in ALL_COLS], lambda v: v[5] in MISSED + UNREACHED)

    if rep["unbuildable"]:
        table("Cannot be sent", ["date", "leader", "brigadir", "uid", "why"],
              [[r["date"], r["leader"], r["unit"], r["uid"], r["why"]]
               for r in rep["unbuildable"]], [11, 28, 26, 12, 60])

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ── delivery ─────────────────────────────────────────────────────────────────

def _clean(text: str) -> str:
    token = settings.telegram_bot_token or ""
    return text.replace(token, "***") if token else text


def _post(method: str, data: dict, files: dict | None = None, timeout: int = 120) -> dict:
    """One Bot API call, retried; the token (which rides in the URL) never
    reaches an exception text or the journal."""
    last = ""
    for attempt in range(1, SEND_RETRIES + 1):
        wait = 0
        try:
            r = requests.post(f"{_API}/bot{settings.telegram_bot_token}/{method}",
                              data=data, files=files, timeout=timeout)
            body = r.json()
            if body.get("ok"):
                return body.get("result") or {}
            last = body.get("description") or f"HTTP {r.status_code}"
            wait = int(((body.get("parameters") or {}).get("retry_after")) or 0)
        except Exception as exc:
            last = f"{type(exc).__name__}: {exc}"
        if attempt < SEND_RETRIES:
            time.sleep(max(wait, 3 * attempt))
    raise RuntimeError(_clean(last or f"{method} failed")[:300])


def summary_text(rep: dict) -> str:
    s = rep["summary"]
    by = s["by_verdict"]
    received = by.get(RECEIVED, 0) + by.get(RECEIVED_LATE, 0)
    not_received = sum(by.get(v, 0) for v in MISSED + UNREACHED)
    lines = [
        f"Leader day reports {_dm(rep['date_from'])} → {_dm(rep['date_to'])} — "
        f"did the leaders get them?",
        "",
        f"Closed leader-days: {s['closed']} ({s['leaders_closed']} leaders)",
        f"✅ Received: {received}"
        + (f" (of them late, {LATE_H}h+ after the day closed: {by[RECEIVED_LATE]})"
           if by.get(RECEIVED_LATE) else ""),
        f"❌ NOT received: {not_received}",
        f"   • tried while the day was still open, never retried: {by.get(MISSED_PARKED, 0)}",
        f"   • finished, never sent: {by.get(MISSED_UNSENT, 0)}",
        f"   • sent, but the leader has no Telegram account: {by.get(NO_ACCOUNT, 0)}",
        f"   • sent, but Telegram refused it (bot blocked / never started): "
        f"{by.get(BLOCKED, 0)}",
        f"⏳ Waiting for the AI review: {by.get(WAITING, 0)}",
        f"➖ No report by design (no photo task reviewed): {by.get(NO_PHOTO, 0)}",
        f"➖ Not reported on purpose (excluded, cut off, …): {by.get(ON_PURPOSE, 0)}",
        f"Day still open: {by.get(OPEN, 0)}",
        "",
        f"Leaders who did not get at least one report: {s['leaders_missed']} "
        f"of {s['leaders_closed']}.",
        f"AI-rejected tasks on the never-sent days: {s['missed_rejected_tasks']}, "
        f"on {s['missed_rejected_days']} days — objections filed on them: "
        f"{s['missed_objections']}.",
    ]
    if s["missed_last_closed_by"]:
        lines.append("What closed those days last: " + ", ".join(
            f"{k} {v}" for k, v in s["missed_last_closed_by"].items()) + ".")
    lines += ["", "Per date — closed · NOT received:"]
    per: dict[str, list[str]] = defaultdict(list)
    for p in s["per_date"]:
        if p["closed"]:
            per[p["date"]].append(f"S{p['shift'] or '?'} {p['closed']}·"
                                  f"{p['missed'] + p['unreached']}")
    for dt in sorted(per):
        lines.append(f"{_dm(dt)}  " + "   ".join(per[dt]))
    lines += [
        "",
        f"From {_dm(rep['resend_before'])} on the fix sends them by itself "
        f"({s['fixed_now']} held back today are going out now).",
        "Excel: «Leaders» names who missed which dates; «To send» is the list "
        "behind the buttons.",
    ]
    return "\n".join(lines)


def _fname(rep: dict) -> str:
    return f"missed-day-reports-{rep['date_from']}-to-{rep['date_to']}"


def list_id(items: list[dict]) -> str:
    blob = json.dumps([i["uid"] for i in items], sort_keys=True)
    return hashlib.sha1(blob.encode()).hexdigest()[:10]


def _card_text(s: dict, rep: dict) -> str:
    last = _dm((datetime.fromisoformat(rep["resend_before"]) - timedelta(days=1))
               .date().isoformat())
    text = (f"Send the missed day reports for {_dm(rep['date_from'])}–{last} to the "
            f"leaders: the normal day report, one DM per day, each with its "
            f"«Open the report» button — where they can object. Leaders only; "
            f"the brigadirs already have their daily digest.\n\n"
            f"To send: {s['to_send']} report(s) to {s['to_send_leaders']} leader(s) — "
            f"{s['to_send_rejected']} of them carry AI-rejected tasks "
            f"({s['to_send_leaders_rejected']} leaders).")
    if s["to_send_unreachable_leaders"]:
        text += (f"\n⚠ {s['to_send_unreachable_leaders']} of these leaders have no "
                 f"working Telegram (no account / bot blocked): their reports land "
                 f"in the app's bell only — sheet «To send», column «Telegram now».")
    return text


def _keyboard(lid: str, s: dict) -> str:
    rows = [[{"text": f"📨 Send all {s['to_send']}",
              "callback_data": f"{CALLBACK}:all:{lid}"}]]
    if 0 < s["to_send_rejected"] < s["to_send"]:
        rows.append([{"text": f"📨 Only the {s['to_send_rejected']} with rejected tasks",
                      "callback_data": f"{CALLBACK}:rej:{lid}"}])
    return json.dumps({"inline_keyboard": rows})


def _after_keyboard(lid: str, mode: str, rest: int, failed: int) -> str:
    """What the card offers once a run is over: a retry when something failed,
    the clean days when only the rejected ones went, else nothing."""
    if failed:
        rows = [[{"text": "🔁 Retry the failed ones",
                  "callback_data": f"{CALLBACK}:{mode}:{lid}"}]]
    elif mode == "rej" and rest:
        rows = [[{"text": f"📨 Send the other {rest} (no rejected tasks)",
                  "callback_data": f"{CALLBACK}:all:{lid}"}]]
    else:
        rows = []
    return json.dumps({"inline_keyboard": rows})


def _save_run(db: Session, idx: int, **fields) -> dict:
    """Merge `fields` into run `idx` of the stored list and commit."""
    row = db.query(AppSetting).filter_by(key=LIST_KEY).first()
    data = json.loads(row.value)
    data["runs"][idx].update(fields)
    row.value = json.dumps(data, ensure_ascii=False)
    db.commit()
    return data


def send(db: Session, chat_id: int, *_window) -> int:
    """Step 1: build the list, STORE it, DM the report with its buttons.
    Returns how many messages went out (the `_send_report_once` contract)."""
    if not settings.telegram_bot_token:
        raise RuntimeError("telegram bot token not configured")
    rep = collect(db)
    db.rollback()
    items = rep["items"]
    lid = list_id(items)
    row = db.query(AppSetting).filter_by(key=LIST_KEY).first()
    stored = json.loads(row.value) if row and row.value else {}
    if stored.get("runs"):
        raise RuntimeError("the resend list was already used — not rebuilt")
    value = json.dumps({"list_id": lid, "built_at": rep["generated_at"],
                        "items": items, "runs": []}, ensure_ascii=False)
    if row is None:
        db.add(AppSetting(key=LIST_KEY, value=value))
    else:
        row.value = value
    db.commit()

    _post("sendMessage", {"chat_id": chat_id, "text": summary_text(rep)[:4000]})
    _post("sendDocument", {"chat_id": chat_id,
                           "caption": "Leader day reports — received or not, per "
                                      "leader-day; who missed which dates."},
          files={"document": (f"{_fname(rep)}.xlsx", build_workbook(rep),
                              "application/vnd.openxmlformats-officedocument."
                              "spreadsheetml.sheet")}, timeout=300)
    blob = json.dumps(rep, ensure_ascii=False, indent=1, default=str).encode()
    _post("sendDocument", {"chat_id": chat_id, "caption": "Same report, full data."},
          files={"document": (f"{_fname(rep)}.json", blob,
                              "application/json")}, timeout=300)
    s = rep["summary"]
    if not items:
        _post("sendMessage", {"chat_id": chat_id,
                              "text": "Nothing to send: every closed day before "
                                      f"{_dm(rep['resend_before'])} was reported."})
        return 4
    _post("sendMessage", {"chat_id": chat_id, "text": _card_text(s, rep),
                          "reply_markup": _keyboard(lid, s)})
    return 4


# ── step 2: the tap ──────────────────────────────────────────────────────────

def apply(db: Session, *, actor_tid: int, actor_name: str, lid: str, mode: str,
          chat_id: int | None, message_id: int | None) -> dict:
    """Start sending the STORED list — never a recomputation. The sending
    itself runs as a background job (`_run`): a few hundred DMs at one a second
    would hold the bot's update loop for minutes."""
    from app.scheduler import schedule_at

    if mode not in ("all", "rej"):
        return {"status": "bad"}
    row = db.query(AppSetting).filter_by(key=LIST_KEY).with_for_update().first()
    if row is None or not row.value:
        return {"status": "missing"}
    data = json.loads(row.value)
    if data.get("list_id") != lid:
        return {"status": "stale"}
    runs = data.setdefault("runs", [])
    now = datetime.now(timezone.utc)
    if runs and runs[-1].get("status") == "sending":
        beat = datetime.fromisoformat(runs[-1].get("beat") or runs[-1]["started"])
        if now - beat < timedelta(minutes=RUN_STALE_MIN):
            return {"status": "busy"}
        runs[-1]["status"] = "interrupted"   # died with its process; sending is idempotent
    items = [i for i in data.get("items") or [] if mode == "all" or i["has_rejected"]]
    runs.append({"mode": mode, "by": actor_name, "by_tid": actor_tid,
                 "started": now.isoformat(), "status": "sending", "planned": len(items)})
    row.value = json.dumps(data, ensure_ascii=False)
    db.commit()
    idx = len(runs) - 1
    if not schedule_at(f"missed-report-resend-{idx}", now + timedelta(seconds=2),
                       _run, args=(lid, idx, chat_id, message_id)):
        return {"status": "error"}
    return {"status": "started", "planned": len(items), "mode": mode}


def _run(lid: str, idx: int, chat_id: int | None, message_id: int | None) -> None:
    """Send every report of one run, one DM a second, then say what happened.

    `send_for_uid` decides each day again at the moment of sending — excluded
    since, cut off since, sent since (the ledger), all of it — so this sends
    exactly the reports that are still owed and nothing twice. That is also
    what makes a run safe to repeat after a restart killed it half-way."""
    from app.database import SessionLocal
    from app.services import action_log, leader_reports

    db = SessionLocal()
    try:
        row = db.query(AppSetting).filter_by(key=LIST_KEY).first()
        data = json.loads(row.value)
        run = data["runs"][idx]
        mode = run["mode"]
        items = [i for i in data.get("items") or [] if mode == "all" or i["has_rejected"]]
        waiting = leader_ai.unfinished_reports(db, {i["key"] for i in items})
        sent = already = waiting_n = skipped = failed = 0
        leaders: set[int] = set()
        for n, it in enumerate(items, start=1):
            try:
                if it["key"] in waiting:
                    waiting_n += 1
                elif (db.query(LeaderDayReport.id)
                      .filter(LeaderDayReport.report_key == it["key"],
                              LeaderDayReport.sends > 0).first()):
                    already += 1
                elif leader_reports.send_for_uid(db, it["uid"], key=it["key"]):
                    sent += 1
                    leaders.add(it["leader_id"])
                    time.sleep(PACE_S)
                else:
                    skipped += 1
            except Exception:           # noqa: BLE001 - one report never stops the run
                db.rollback()
                failed += 1
                log.exception("missed-report resend: %s failed", it["uid"])
            if n % _BEAT_EVERY == 0:
                _save_run(db, idx, beat=datetime.now(timezone.utc).isoformat(),
                          progress=n)

        data = _save_run(db, idx, status="done",
                         finished=datetime.now(timezone.utc).isoformat(),
                         sent=sent, already=already, waiting=waiting_n,
                         skipped=skipped, failed=failed)
        action_log.record_bot(
            db, run.get("by_tid"), "leader_review", "report.missed_resent",
            actor_name=run.get("by"), target_kind="report",
            target_id=f"batch:{sent}",
            details=[("mode", mode), ("sent", sent), ("leaders", len(leaders)),
                     ("already", already or None), ("waiting", waiting_n or None),
                     ("skipped", skipped or None), ("failed", failed or None),
                     ("list", lid)])

        text = (f"✅ Missed day reports sent: {sent} to {len(leaders)} leader(s)"
                f" ({'every listed day' if mode == 'all' else 'days with rejected tasks'}).")
        if already:
            text += f"\nAlready sent since the list was built: {already}."
        if waiting_n:
            text += (f"\nStill waiting for the AI review — they go out by themselves "
                     f"when it lands: {waiting_n}.")
        if skipped:
            text += (f"\nNot sent — no longer reportable (excluded / cut off / "
                     f"rehearsal since): {skipped}.")
        if failed:
            text += f"\n⚠ Failed: {failed} — tap «Retry» below, or tell the developer."
        if chat_id:
            _post("sendMessage", {"chat_id": chat_id, "text": text})
            if message_id:
                rest = sum(1 for i in data.get("items") or [] if not i["has_rejected"])
                try:
                    _post("editMessageReplyMarkup", {
                        "chat_id": chat_id, "message_id": message_id,
                        "reply_markup": _after_keyboard(lid, mode, rest, failed)})
                except Exception:       # noqa: BLE001 - the card is cosmetic
                    log.warning("missed-report resend: could not edit the card",
                                exc_info=True)
    except Exception:
        db.rollback()
        log.exception("missed-report resend: run %s failed", idx)
        try:
            _save_run(db, idx, status="error")
        except Exception:               # noqa: BLE001
            db.rollback()
        if chat_id:
            try:
                _post("sendMessage", {"chat_id": chat_id,
                                      "text": "⚠ Sending the missed reports stopped "
                                              "with an error. Tap the button again — "
                                              "reports already sent are not sent twice."})
            except Exception:           # noqa: BLE001
                pass
    finally:
        db.close()


def result_text(res: dict) -> str:
    st = res.get("status")
    if st == "started":
        return (f"Sending {res['planned']} report(s) now, one a second — "
                f"I'll message you when it is done.")
    if st == "busy":
        return "Already sending — wait for the result message."
    if st == "stale":
        return "This button belongs to an older list — nothing was sent."
    if st == "error":
        return "Could not start sending — nothing was sent. Tap again."
    return "The list is gone — nothing was sent."
