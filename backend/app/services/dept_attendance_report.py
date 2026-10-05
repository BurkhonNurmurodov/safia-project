"""One-off: the clock-in / clock-out of «Департамент по продукту и технологиям»
for Mon 28.09 – Sun 04.10.2026, read from Verifix and DMed to the operator as
text + a workbook.

The operator asked on 2026-10-05 whether the platform can see the clock in/out
of that department's people and, if so, for the week as Excel: full name,
position, clock in, clock out, work schedule, worked hours. The session that
wrote this cannot reach Verifix or Telegram (the Verifix login and the bot token
live only on the server), so the answer is produced here.

Read exactly the way the admin «Davomat» tab reads a day (`verifix_attendance`):

* `core/timesheet$export` («Отчёт по посещениям») for the whole week, its own
  `input_time` / `output_time` as the clocks — proven equal to the Excel export's
  for 1,470 of 1,475 people (CLAUDE.md, «Davomat reads Verifix»);
* hours = «Отработано» by `verifix_attendance.hours_rule` (Явка + Свободное
  время unless the parity check found a rule of its own), and only for a day
  carrying BOTH clocks, as the export prints it;
* the work schedule = the report's own schedule name, and the day's planned
  begin / end (`begin_time` / `end_time`).

The department is found by NAME in `core/division$list`: every division whose
folded name holds «продукт» and «технолог» (a match naming «департамент» wins
over one that does not), with every division under it. Its people are every
employee — any status — whose department or org unit is in that subtree and who
was employed during the week, plus whoever the report itself files under it.

It READS and writes nothing but its flag (written by `startup._send_report_once`).
No passport, PINFL, phone or photo is read into the file.

Temporary: delete this module, `startup.report_dept_attendance` /
`_dept_attendance_job` and the call in BOTH entrypoints once it has been sent.
"""
from __future__ import annotations

import time
from collections import Counter
from datetime import date, datetime, timedelta
from io import BytesIO
from typing import Any, Optional

import requests
from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Font, PatternFill
from sqlalchemy.orm import Session

from app.config import settings
from app.services import verifix, verifix_attendance, verifix_live

DEPT_LABEL = "Департамент по продукту и технологиям"
PERIOD_FROM = date(2026, 9, 28)
PERIOD_TO = date(2026, 10, 4)
DEADLINE_S = 420.0        # a boot job, not a request: no Cloudflare 100 s here
READ_TRIES = 3            # the whole Verifix read, with a pause between
ID_CHUNK = 400
GRACE_MIN = verifix_live.LATE_GRACE_MIN

_API = "https://api.telegram.org"
SEND_RETRIES = 3

DAY_KIND = {"W": "Work day", "R": "Day off", "A": "Extra day off",
            "H": "Holiday", "N": "Non-working day"}
WEEKDAY = ("Mon", "Tue", "Wed", "Thu", "Fri", "Sat", "Sun")

_HEAD_FILL = PatternFill("solid", fgColor="C8973F")
_RED = PatternFill("solid", fgColor="FDE2E2")
_AMBER = PatternFill("solid", fgColor="FEF3C7")
_GREY = PatternFill("solid", fgColor="EEF0F3")
_GREEN = PatternFill("solid", fgColor="DCFCE7")


# ── small helpers ─────────────────────────────────────────────────────────────

def _s(v: Any) -> str:
    return str(v).strip() if v is not None else ""


def _fold(s: str) -> str:
    return (s or "").lower().replace("ё", "е")


def _is_dept(name: str) -> bool:
    f = _fold(name)
    return (("продукт" in f and "технолог" in f)
            or ("product" in f and "technolog" in f)
            or ("mahsulot" in f and "texnolog" in f))


def _days() -> list[date]:
    n = (PERIOD_TO - PERIOD_FROM).days + 1
    return [PERIOD_FROM + timedelta(days=i) for i in range(n)]


def _num(v: Any) -> float:
    try:
        return float(v or 0)
    except (TypeError, ValueError):
        return 0.0


def _clock(dt: Optional[datetime], day: date) -> Optional[str]:
    """«08:57», or «00:40 (+1)» when it falls on the next calendar day."""
    if not dt:
        return None
    s = dt.strftime("%H:%M")
    shift = (dt.date() - day).days
    return s if shift == 0 else f"{s} ({'+' if shift > 0 else ''}{shift})"


def _hm(hours: Optional[float]) -> str:
    if hours is None:
        return ""
    m = int(round(hours * 60))
    return f"{m // 60}:{m % 60:02d}"


# ── reading Verifix ───────────────────────────────────────────────────────────

def _pages(cl, path: str, body: dict, limit: int, deadline: float):
    for page in verifix.each_page(cl, path, body, limit=limit, deadline=deadline):
        for r in page or []:
            if isinstance(r, dict):
                yield r


def _merge_ts(out: dict, r: dict, src: str) -> None:
    eid = _s(r.get("employee_id"))
    if not eid:
        return
    rec = out.setdefault(eid, {"name": _s(r.get("employee_name")), "job": _s(r.get("job_name")),
                               "schedule": _s(r.get("schedule_name")),
                               "div": _s(r.get("division_id")), "days": {}, "src": set()})
    rec["src"].add(src)
    for k, v in (("name", "employee_name"), ("job", "job_name"),
                 ("schedule", "schedule_name"), ("div", "division_id")):
        if not rec[k] and _s(r.get(v)):
            rec[k] = _s(r.get(v))
    for d in r.get("days") or []:
        if not isinstance(d, dict):
            continue
        day = verifix_live._d(d.get("date"))
        if day is None or not (PERIOD_FROM <= day <= PERIOD_TO):
            continue
        # A row per (employee, schedule) can split one week; each day keeps the
        # schedule / position of the row it came in.
        prev = rec["days"].get(day)
        cand = {**d, "_schedule": _s(r.get("schedule_name")), "_job": _s(r.get("job_name"))}
        if prev is None or (not prev.get("input_time") and cand.get("input_time")):
            rec["days"][day] = cand


def _read(cfg: dict) -> dict:
    deadline = time.monotonic() + DEADLINE_S
    divs: dict[str, dict] = {}
    jobs: dict[str, str] = {}
    kinds: dict[str, str] = {}
    emps: dict[str, dict] = {}
    notes: list[str] = []
    with verifix.client(cfg) as cl:
        for d in _pages(cl, "core/division$list", {"division_ids": []},
                        verifix.LIMIT_LIST, deadline):
            did = _s(d.get("division_id"))
            if did:
                divs[did] = {"id": did, "name": _s(d.get("name")) or f"#{did}",
                             "parent": _s(d.get("parent_id")) or None,
                             "code": _s(d.get("code")) or None, "state": _s(d.get("state")) or "A"}

        matched = [x for x in divs.values() if _is_dept(x["name"])]
        dept_named = [x for x in matched if "департамент" in _fold(x["name"])]
        roots = dept_named or matched
        if not roots:
            cands = sorted({x["name"] for x in divs.values()
                            if any(w in _fold(x["name"])
                                   for w in ("департамент", "продукт", "технолог"))})
            return {"found": False, "divisions": len(divs), "candidates": cands}
        ignored = [x for x in matched if x not in roots]

        children: dict[str, list[str]] = {}
        for x in divs.values():
            if x["parent"]:
                children.setdefault(x["parent"], []).append(x["id"])
        subtree: set[str] = set()
        stack = [x["id"] for x in roots]
        while stack:
            cur = stack.pop()
            if cur in subtree:
                continue
            subtree.add(cur)
            stack.extend(children.get(cur, []))

        for j in _pages(cl, "core/job$list", {"job_ids": []}, verifix.LIMIT_LIST, deadline):
            if _s(j.get("job_id")):
                jobs[_s(j.get("job_id"))] = _s(j.get("name"))
        scheds: dict[str, str] = {}
        try:
            for x in _pages(cl, "core/schedule$list", {"schedule_ids": []},
                            verifix.LIMIT_LIST, deadline):
                if _s(x.get("schedule_id")):
                    scheds[_s(x.get("schedule_id"))] = _s(x.get("name"))
        except verifix.VerifixError as exc:
            notes.append(f"schedule names not read ({exc.code})")
        try:
            for k in _pages(cl, "core/time_kind$list", {"time_kind_ids": []},
                            verifix.LIMIT_LIST, deadline):
                if _s(k.get("time_kind_id")):
                    kinds[_s(k.get("time_kind_id"))] = _s(k.get("name"))
        except verifix.VerifixError as exc:
            notes.append(f"time kinds not read ({exc.code}) — the breakdown shows ids")

        body = {"employee_ids": [], "statuses": [], "npins": []}
        for e in _pages(cl, "core/employee$list", body, verifix.LIMIT_LIST, deadline):
            eid = _s(e.get("employee_id"))
            div, unit = _s(e.get("division_id")), _s(e.get("org_unit_id"))
            if not eid or not (div in subtree or unit in subtree):
                continue
            hired = verifix_live._d(e.get("hiring_date"))
            gone = verifix_live._d(e.get("dismissal_date"))
            if (hired and hired > PERIOD_TO) or (gone and gone < PERIOD_FROM):
                continue
            name = " ".join(_s(e.get(k)) for k in ("last_name", "first_name", "middle_name")
                            if _s(e.get(k)))
            emps[eid] = {"name": name or f"#{eid}", "job": jobs.get(_s(e.get("job_id")), ""),
                         "schedule": scheds.get(_s(e.get("schedule_id")), ""),
                         "div": div, "unit": unit, "status": _s(e.get("status")) or "W",
                         "hired": hired, "dismissed": gone}

        ts: dict[str, dict] = {}
        ids = sorted(int(x) for x in emps if x.isdigit())
        # An empty employee filter means EVERYBODY — never send one.
        for i in range(0, len(ids), ID_CHUNK):
            body = {"period_begin_date": verifix._dmy(PERIOD_FROM),
                    "period_end_date": verifix._dmy(PERIOD_TO),
                    "division_ids": [], "employee_ids": ids[i:i + ID_CHUNK]}
            for r in _pages(cl, "core/timesheet$export", body, verifix.LIMIT_TIMESHEET, deadline):
                if _s(r.get("employee_id")) in emps:
                    _merge_ts(ts, r, "people")

        # Whoever the REPORT files under the department, wherever the directory
        # places them today (a transfer since the week). The row's own division
        # decides, so a server that ignored the filter adds nobody.
        div_ids = sorted(int(x) for x in subtree if x.isdigit())
        try:
            body = {"period_begin_date": verifix._dmy(PERIOD_FROM),
                    "period_end_date": verifix._dmy(PERIOD_TO),
                    "division_ids": div_ids, "employee_ids": []}
            for r in _pages(cl, "core/timesheet$export", body, verifix.LIMIT_TIMESHEET, deadline):
                if _s(r.get("division_id")) in subtree:
                    _merge_ts(ts, r, "division")
        except verifix.VerifixError as exc:
            notes.append(f"the by-department report read stopped ({exc.code}); "
                         "people were read by name list")

    return {"found": True, "divs": divs, "roots": roots, "ignored": ignored,
            "subtree": subtree, "kinds": kinds, "emps": emps, "ts": ts, "notes": notes}


# ── shaping it ────────────────────────────────────────────────────────────────

def _day_row(day: date, d: Optional[dict], rec: dict, rule: dict, kinds: dict) -> dict:
    if d is None:
        return {"date": day, "kind": None, "schedule": rec.get("schedule") or "",
                "job": rec.get("job") or "", "plan": "", "in": None, "out": None,
                "hours": None, "late": None, "early": None, "breakdown": "",
                "status": "No report data", "tone": "grey"}
    kind = _s(d.get("day_kind")) or "W"
    schedule = d.get("_schedule") or rec.get("schedule") or ""
    begin, end = verifix_live._dt(d.get("begin_time")), verifix_live._dt(d.get("end_time"))
    if (begin is None or end is None) and kind == "W" and schedule:
        b2, e2 = verifix_live._schedule_window(schedule, day)
        begin, end = begin or b2, end or e2
    t_in, t_out = verifix_live._dt(d.get("input_time")), verifix_live._dt(d.get("output_time"))

    facts: dict[str, float] = {}
    for f in d.get("facts") or []:
        k = _s(f.get("time_kind_id"))
        v = _num(f.get("fact_value"))
        if k and v:
            facts[k] = facts.get(k, 0.0) + v

    hours = None
    if t_in and t_out:
        hours = round(sum(facts.get(k, 0.0) for k in rule["kinds"]) / rule["div"], 2)
    late = early = None
    if begin and t_in:
        delta = verifix_live._mins(verifix_live._minute(begin), verifix_live._minute(t_in))
        if delta > GRACE_MIN:
            late = round(delta)
    if end and t_out:
        delta = verifix_live._mins(verifix_live._minute(t_out), verifix_live._minute(end))
        if delta > GRACE_MIN:
            early = round(delta)

    if t_in and t_out:
        status, tone = "Worked", ("amber" if late or early else "green")
    elif t_in:
        status, tone = "No clock-out", "red"
    elif t_out:
        status, tone = "No clock-in", "red"
    else:
        other = [kinds.get(k) or f"kind {k}" for k, v in sorted(facts.items(), key=lambda x: -x[1])
                 if k not in rule["kinds"] and v > 0]
        if kind == "W":
            status, tone = ("Absent" + (f" — {', '.join(other)}" if other else "")), "red"
        else:
            status, tone = DAY_KIND.get(kind, f"Day kind {kind}"), "grey"

    breakdown = "; ".join(f"{kinds.get(k) or 'kind ' + k} {_hm(v / rule['div'])}"
                          for k, v in sorted(facts.items(), key=lambda x: -x[1]))
    plan = (f"{begin.strftime('%H:%M')}–{_clock(end, day)}" if begin and end else "")
    return {"date": day, "kind": kind, "schedule": schedule, "job": d.get("_job") or rec.get("job") or "",
            "plan": plan, "in": _clock(t_in, day), "out": _clock(t_out, day), "hours": hours,
            "late": late, "early": early, "breakdown": breakdown, "status": status, "tone": tone}


def collect(db: Session) -> dict:
    cfg = verifix.config(db, with_password=True)
    if not (cfg["login"] and cfg.get("password") and cfg["filial_id"]):
        return {"configured": False}
    rule = verifix_attendance.hours_rule(db)
    db.rollback()                      # nothing below needs the session open

    last: Optional[Exception] = None
    raw = None
    for attempt in range(1, READ_TRIES + 1):
        try:
            raw = _read(cfg)
            break
        except verifix.VerifixError as exc:
            last = exc
            if attempt < READ_TRIES:
                time.sleep(30 * attempt)
    if raw is None:
        raise RuntimeError(f"Verifix read failed: {getattr(last, 'code', '')} "
                           f"{getattr(last, 'message', last)}"[:300])
    out = {"configured": True, "rule": rule, **raw}
    if not raw["found"]:
        return out

    divs, emps, ts, kinds = raw["divs"], raw["emps"], raw["ts"], raw["kinds"]
    people = []
    for eid in set(emps) | set(ts):
        e = emps.get(eid) or {}
        rec = ts.get(eid) or {}
        unit_id = e.get("unit") if e.get("unit") in raw["subtree"] else (e.get("div") or rec.get("div"))
        sub = (divs.get(unit_id) or {}).get("name") or (divs.get(rec.get("div") or "") or {}).get("name") or ""
        status = e.get("status")
        if not e:
            emp_state = "Not in today's directory under this department (moved since)"
        elif status == "W":
            emp_state = "Working"
        elif e.get("dismissed"):
            emp_state = f"Dismissed {e['dismissed']:%d.%m.%Y}"
        else:
            emp_state = f"Status {status}"
        if not rec.get("schedule") and e.get("schedule"):
            rec = {**rec, "schedule": e["schedule"]}
        rows = [_day_row(day, (rec.get("days") or {}).get(day), rec, rule, kinds) for day in _days()]
        if eid not in ts:
            for r in rows:
                r["status"] = "No row in Verifix's report"
        scheds = Counter(r["schedule"] for r in rows if r["schedule"])
        people.append({
            "id": eid, "name": rec.get("name") or e.get("name") or f"#{eid}",
            "job": rec.get("job") or e.get("job") or "",
            "sub": sub, "state": emp_state,
            "schedule": scheds.most_common(1)[0][0] if scheds else (rec.get("schedule") or ""),
            "in_report": eid in ts, "rows": rows,
        })
    people.sort(key=lambda p: (_fold(p["sub"]), _fold(p["name"])))
    out["people"] = people
    return out


# ── the text ──────────────────────────────────────────────────────────────────

def _totals(people: list[dict]) -> dict:
    t = Counter()
    for p in people:
        for r in p["rows"]:
            if r["in"]:
                t["clocked"] += 1
            if r["in"] and r["out"]:
                t["full"] += 1
                t["hours"] += r["hours"] or 0
            if r["status"] == "No clock-out":
                t["no_out"] += 1
            if r["status"] == "No clock-in":
                t["no_in"] += 1
            if r["status"].startswith("Absent"):
                t["absent"] += 1
            if r["late"]:
                t["late"] += 1
    return t


def text(rep: dict) -> str:
    span = f"{PERIOD_FROM:%d.%m}–{PERIOD_TO:%d.%m.%Y}"
    if not rep.get("configured"):
        return (f"Verifix — «{DEPT_LABEL}», {span}\n\n"
                "Could not read it: no Verifix login is saved on the server "
                "(admin → «Verifix» card).")
    if not rep["found"]:
        lines = [f"Verifix — «{DEPT_LABEL}», {span}", "",
                 f"No Verifix division matched this name (searched all {rep['divisions']} "
                 "divisions for «продукт» + «технолог»). So no report yet.", "",
                 "Divisions that exist with «департамент», «продукт» or «технолог» in the name:"]
        lines += [f"• {n}" for n in rep["candidates"][:60]] or ["• (none)"]
        lines += ["", "Tell me which one is meant and I will send the week for it."]
        return "\n".join(lines)

    people = rep["people"]
    t = _totals(people)
    working = sum(1 for p in people if p["state"] == "Working")
    in_rep = sum(1 for p in people if p["in_report"])
    roots = ", ".join(f"«{r['name']}» (Verifix id {r['id']}"
                      + (f", code {r['code']}" if r.get("code") else "") + ")" for r in rep["roots"])
    subs = len(rep["subtree"]) - len(rep["roots"])
    rule = rep["rule"]
    lines = [
        f"Verifix — «{DEPT_LABEL}», {span} (Mon–Sun)", "",
        ("Access: YES — Verifix's attendance report («Отчёт по посещениям») returns this "
         "department's clock-in/out." if in_rep else
         "Access: the department and its people are visible, but Verifix's attendance report "
         "returned NO rows for them — the API role may not cover this department."),
        "",
        f"Department: {roots}" + (f" + {subs} unit(s) under it" if subs else ""),
        f"People: {len(people)} employed during the week ({working} working today), "
        f"{in_rep} in the report.",
        f"Days with a clock-in: {t['clocked']} · full in+out: {t['full']} "
        f"({t['hours']:.1f} h worked in total)",
        f"No clock-out: {t['no_out']} · no clock-in: {t['no_in']} · absent on a work day: "
        f"{t['absent']} · late (>{GRACE_MIN} min): {t['late']}",
        "",
        f"Worked hours = Verifix «Отработано» ({' + '.join(rule['names'])}), the same rule "
        "as the «Davomat» tab; counted only on a day with both clocks.",
        "The workbook: Summary · Daily · Week grid · Notes.",
    ]
    if rep["ignored"]:
        lines.append("Also matched the name but left out (not a «департамент»): "
                     + ", ".join(f"«{x['name']}»" for x in rep["ignored"][:10]))
    for n in rep["notes"]:
        lines.append(f"Note: {n}")
    return "\n".join(lines)


# ── the workbook ──────────────────────────────────────────────────────────────

def _cell(v):
    if isinstance(v, str):
        return ILLEGAL_CHARACTERS_RE.sub("", v)
    return v


def _sheet(wb: Workbook, title: str, head: list[str], rows: list[list],
           tones: Optional[list[Optional[PatternFill]]] = None, tone_cols=(),
           widths: Optional[dict[int, int]] = None):
    ws = wb.create_sheet(title)
    ws.append(head)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = _HEAD_FILL
        c.alignment = Alignment(vertical="center", wrap_text=True)
    for i, r in enumerate(rows):
        ws.append([_cell(v) for v in r])
        fill = tones[i] if tones else None
        if fill is not None:
            for col in tone_cols:
                ws.cell(i + 2, col).fill = fill
    if title != "Notes":
        ws.freeze_panes = "B2"
    if rows and title != "Notes":
        ws.auto_filter.ref = ws.dimensions
    for i, h in enumerate(head, start=1):
        width = max([len(str(h))] + [len(str(r[i - 1])) for r in rows[:500]
                                     if i - 1 < len(r) and r[i - 1] is not None])
        if widths and i in widths:
            width = widths[i]
        ws.column_dimensions[ws.cell(1, i).column_letter].width = min(max(width + 2, 8), 60)
    ws.row_dimensions[1].height = 30
    return ws


_TONE = {"red": _RED, "amber": _AMBER, "grey": _GREY, "green": _GREEN}


def build_workbook(rep: dict) -> BytesIO:
    people = rep["people"]
    days = _days()
    wb = Workbook()
    wb.remove(wb.active)

    summ = []
    for p in people:
        rows = p["rows"]
        full = [r for r in rows if r["in"] and r["out"]]
        total = round(sum(r["hours"] or 0 for r in full), 2)
        summ.append([
            p["name"], p["job"], p["sub"], p["schedule"], p["state"],
            sum(1 for r in rows if r["kind"] == "W"),
            sum(1 for r in rows if r["in"]), len(full), total,
            round(total / len(full), 2) if full else None,
            sum(1 for r in rows if r["late"]), sum(1 for r in rows if r["early"]),
            sum(1 for r in rows if r["status"] == "No clock-out"),
            sum(1 for r in rows if r["status"].startswith("Absent")),
        ])
    ws = _sheet(wb, "Summary", [
        "Full name", "Position", "Unit", "Work schedule", "Employment",
        "Work days planned", "Days clocked in", "Days with in + out", "Worked hours (total)",
        "Avg hours per full day", "Late days", "Left-early days", "No clock-out days",
        "Absent work days"], summ)
    for row in ws.iter_rows(min_row=2, min_col=9, max_col=10):
        for c in row:
            c.number_format = "0.00"

    daily, tones = [], []
    for p in people:
        for r in p["rows"]:
            daily.append([
                p["name"], r["job"] or p["job"], p["sub"], r["date"], WEEKDAY[r["date"].weekday()],
                r["schedule"], DAY_KIND.get(r["kind"], r["kind"] or ""), r["plan"],
                r["in"], r["out"], r["hours"], r["late"], r["early"], r["status"], r["breakdown"]])
            tones.append(_TONE.get(r["tone"]))
    ws = _sheet(wb, "Daily", [
        "Full name", "Position", "Unit", "Date", "Day", "Work schedule", "Day type",
        "Planned (start–end)", "Clock in", "Clock out", "Worked hours", "Late (min)",
        "Left early (min)", "Status", "Verifix time breakdown (h:mm)"],
        daily, tones, tone_cols=(14,), widths={4: 11, 15: 50})
    for row in ws.iter_rows(min_row=2, min_col=4, max_col=4):
        for c in row:
            c.number_format = "DD.MM.YYYY"
    for row in ws.iter_rows(min_row=2, min_col=11, max_col=11):
        for c in row:
            c.number_format = "0.00"

    grid_head = ["Full name", "Position", "Work schedule"] + [
        f"{WEEKDAY[d.weekday()]} {d:%d.%m}" for d in days] + ["Total hours"]
    grid, grid_tones = [], []
    for p in people:
        cells = []
        for r in p["rows"]:
            if r["in"] and r["out"]:
                cells.append(f"{r['in']}–{r['out']} · {r['hours']:.2f} h" if r["hours"] is not None
                             else f"{r['in']}–{r['out']}")
            elif r["in"]:
                cells.append(f"{r['in']}–? (no clock-out)")
            elif r["out"]:
                cells.append(f"?–{r['out']} (no clock-in)")
            else:
                cells.append(r["status"])
        total = round(sum(r["hours"] or 0 for r in p["rows"] if r["in"] and r["out"]), 2)
        grid.append([p["name"], p["job"], p["schedule"]] + cells + [total])
        grid_tones.append([_TONE.get(r["tone"]) for r in p["rows"]])
    ws = _sheet(wb, "Week grid", grid_head, grid, widths={i: 24 for i in range(4, 4 + len(days))})
    for i, ts in enumerate(grid_tones):
        for j, fill in enumerate(ts):
            if fill is not None:
                ws.cell(i + 2, 4 + j).fill = fill

    rule = rep["rule"]
    notes = [
        ["Department", DEPT_LABEL],
        ["Matched in Verifix", "; ".join(f"{r['name']} (id {r['id']}"
                                         + (f", code {r['code']}" if r.get("code") else "") + ")"
                                         for r in rep["roots"])],
        ["Units included", "; ".join(sorted(rep["divs"][x]["name"] for x in rep["subtree"]
                                            if x in rep["divs"]))],
        ["Period", f"{PERIOD_FROM:%d.%m.%Y} (Mon) – {PERIOD_TO:%d.%m.%Y} (Sun)"],
        ["Source", "Verifix API, core/timesheet$export («Отчёт по посещениям»), read "
                   f"{verifix_live.now_local():%d.%m.%Y %H:%M} (Tashkent)"],
        ["Who is listed", "Every employee (any status) whose department or org unit is the "
                          "department above or a unit under it and who was employed during the week, "
                          "plus anybody the report itself files under it."],
        ["Clock in / Clock out", "The report's own check-in / check-out (input_time / output_time) — "
                                 "the same values as Verifix's Excel export. «(+1)» = the next day."],
        ["Work schedule / Planned", "The schedule Verifix assigns, and that day's planned start–end."],
        ["Worked hours", f"Verifix «Отработано» = {' + '.join(rule['names'])}, only on a day with "
                         "both clocks (as the export prints it). Same rule as the «Davomat» tab."],
        ["Late / Left early", f"Minutes after the planned start / before the planned end, "
                              f"shown when over {GRACE_MIN} min."],
        ["Status colours", "Green = worked · Amber = worked, late or left early · "
                           "Red = absent or a clock missing · Grey = day off / no data."],
    ]
    for n in rep["notes"]:
        notes.append(["Note", n])
    ws = _sheet(wb, "Notes", ["Item", "Explanation"], notes, widths={1: 24, 2: 100})
    for row in ws.iter_rows(min_row=2, min_col=2, max_col=2):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    for row in ws.iter_rows(min_row=2, min_col=1, max_col=1):
        for c in row:
            c.font = Font(bold=True)

    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


# ── sending ───────────────────────────────────────────────────────────────────

def _post(method: str, data: dict, files: dict | None = None) -> None:
    last = ""
    for attempt in range(1, SEND_RETRIES + 1):
        wait = 0
        try:
            r = requests.post(f"{_API}/bot{settings.telegram_bot_token}/{method}",
                              data=data, files=files, timeout=180)
            body = r.json()
            if body.get("ok"):
                return
            last = body.get("description") or f"HTTP {r.status_code}"
            wait = int(((body.get("parameters") or {}).get("retry_after")) or 0)
        except Exception as exc:
            last = type(exc).__name__
        if attempt < SEND_RETRIES:
            time.sleep(max(wait, 3 * attempt))
    raise RuntimeError(f"{method} failed: {last}"[:300])


def _send_text(chat_id: int, body: str) -> int:
    chunks, cur = [], ""
    for line in body.split("\n"):
        if len(cur) + len(line) + 1 > 3900:
            chunks.append(cur)
            cur = ""
        cur += line + "\n"
    chunks.append(cur)
    for c in chunks:
        _post("sendMessage", {"chat_id": chat_id, "text": c})
    return len(chunks)


def send(db: Session, chat_id: int, *_window) -> int:
    if not settings.telegram_bot_token:
        raise RuntimeError("telegram bot token not configured")
    try:
        rep = collect(db)
    except Exception as exc:
        # Say so rather than leave the operator waiting; the raise lets the
        # next boot try again (`startup._send_report_once`, three attempts).
        _send_text(chat_id, f"Verifix — «{DEPT_LABEL}», {PERIOD_FROM:%d.%m}–{PERIOD_TO:%d.%m.%Y}\n\n"
                            f"The read failed this time ({exc}). It will be tried again on the "
                            "next server restart.")
        raise
    db.rollback()
    n = _send_text(chat_id, text(rep))
    if rep.get("configured") and rep.get("found") and rep.get("people"):
        book = build_workbook(rep).getvalue()
        _post("sendDocument", {"chat_id": chat_id,
                               "caption": f"{DEPT_LABEL} — attendance "
                                          f"{PERIOD_FROM:%d.%m}–{PERIOD_TO:%d.%m.%Y}"},
              files={"document": (f"product-technology-dept-attendance-"
                                  f"{PERIOD_FROM:%Y%m%d}-{PERIOD_TO:%Y%m%d}.xlsx", book,
                                  "application/vnd.openxmlformats-officedocument."
                                  "spreadsheetml.sheet")})
        n += 1
    return n
