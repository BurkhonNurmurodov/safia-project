"""Who still files worker concerns in the GOOGLE SHEETS — one month, as a workbook.

**Why this exists.** From 2026-09-06 (v4.49.0) a worker files a concern on
`/cell-concerns`, on their leader's shop-floor PC, instead of in one of the
~180 per-cell spreadsheets the «Liderlar Havotirlar» registry links. The sheets
were never switched off, and nothing on the platform says who kept writing in
them. The operator asked, on 2026-09-15, for the whole of September once, in
their own chat, as a detailed file. It REPORTS and changes nothing.

**Nothing is re-measured.**

* A sheet row is a `worker_concerns` row — the crawl `/worker-concerns` already
  reads (`services/worker_concerns.py`): only rows with concern text, dated by
  the sheet's own «Дата заполнения». A row whose date cannot be read carries
  ``date = None`` there and belongs to no month; those are COUNTED, never
  silently dropped.
* A row belongs to the REGISTRY's (brigadir, leader, cell) — the business's own
  attribution and the one the KPI page uses. The registry is read live, once,
  so a sheet nobody ever wrote in is still listed; the stamp on the stored rows
  is the fallback, and the file says when it had to fall back. The platform's
  CURRENT leader, unit, shift and plant for the cell's code ride beside the
  registry's spelling, because the registry can lag the /cells register. The
  cell is its CODE — no workshop name anywhere.
* An in-app filing is a `leader_concerns` row carrying `worker_name`, which
  only `/cell-concerns` writes, by its `entry_date` — at whatever level it has
  since been uplifted to, since an uplift does not un-file a concern.

Delivery is the caller's business — `startup.report_sheet_concerns_xlsx` fires
`send_xlsx` once behind a flag, after an incremental crawl so the file reads the
sheets as they stand. Always Uzbek Latin: a boot job has no browser in the loop
to send the words, the rule `ojidaniya_deck` already follows.
"""
from __future__ import annotations

import os
import re
from collections import Counter
from datetime import date, datetime, timedelta
from io import BytesIO
from zoneinfo import ZoneInfo

import requests
from openpyxl import Workbook
from openpyxl.utils import get_column_letter
from sqlalchemy import func
from sqlalchemy.orm import Session

from app.config import settings
from app.models import (Cell, Factory, LeaderConcern, Manager, RoleProfile,
                        WorkerConcern, WorkerConcernSyncMeta)
from app.services.ojidaniya_export import (CENTER, DATE_FMT, INDIGO, WRAP,
                                           _unp_cell, _unp_head, _xl)
from app.services.quality_export import (AMBER, BAND, BOX, BRAND_SOFT, CREAM,
                                         GREEN, INK_SOFT, NUM, ORANGE, PANEL, RED,
                                         RIGHT, SLATE, _banner, _fill, _kpi_cards,
                                         _meta_strip, _section, _sheet)
from app.services.latin_code import _CYRILLIC, _TWINS
from app.translit import transliterate

TZ = ZoneInfo("Asia/Tashkent")
# The sheets are crawled nightly and again right before this file is built; a
# last read older than this means the crawl is failing and the file says so.
SYNC_STALE_AFTER = timedelta(hours=36)

# /cell-concerns went live with v4.49.0 on 06.09.2026 at 15:40. A sheet row
# carries a DATE and no time, so the launch day counts as «from» — its morning
# rows included; the scope strip names the exact minute.
APP_FROM = date(2026, 9, 6)
APP_OPENED = "06.09.2026 15:40"
RECENT_DAYS = 7
SHEET_URL = "https://docs.google.com/spreadsheets/d/{}"

_MONTHS = ("yanvar", "fevral", "mart", "aprel", "may", "iyun", "iyul",
           "avgust", "sentabr", "oktabr", "noyabr", "dekabr")

# The sheets' own vocabulary — what the leaders type, so what they recognise.
STATUS_LABEL = {"todo": "To do", "doing": "Doing", "done": "Done",
                "deferred": "O'tqazish"}

# `worker_concerns._FAILURE_CODES`, in words.
FAIL_LABEL = {"header": "sarlavha topilmadi", "permission": "ruxsat yo'q",
              "missing": "jadval topilmadi", "quota": "Google limiti",
              "network": "tarmoq xatosi", "other": "o'qib bo'lmadi"}

ST_SHEET, ST_BOTH, ST_MOVED, ST_APP, ST_NONE = "sheet", "both", "moved", "app", "none"
STATE_LABEL = {ST_SHEET: "Faqat jadvalda", ST_BOTH: "Jadvalda va ilovada",
               ST_MOVED: "Ilovaga o'tgan", ST_APP: "Faqat ilovada",
               ST_NONE: "Hech qayerda yozilmagan"}
STATE_ORDER = {ST_SHEET: 0, ST_BOTH: 1, ST_MOVED: 2, ST_APP: 3, ST_NONE: 4}
STATE_COLOR = {ST_SHEET: RED, ST_BOTH: AMBER, ST_MOVED: GREEN, ST_APP: GREEN,
               ST_NONE: INK_SOFT}


def _dmy(d: date | None) -> str:
    return d.strftime("%d.%m.%Y") if d else "—"


def _dm(d: date) -> str:
    return d.strftime("%d.%m")


def _day(v) -> date | None:
    try:
        return date.fromisoformat(v) if v else None
    except (TypeError, ValueError):
        return None


def _code(v) -> str:
    """A cell code in the ONE spelling every join here compares: no spaces,
    upper case, leading zeros dropped — the registry writes «822» where the
    /cells register writes «0822»."""
    k = re.sub(r"\s+", "", str(v or "")).upper()
    return k.lstrip("0") or k


_TWIN_TABLE = str.maketrans(_TWINS)
_LATIN = re.compile("[A-Za-z]")


def _word(w: str) -> str:
    """One word of a name, in Latin letters. A Latin word carrying a Cyrillic
    twin («ХAKIMOV», typed on the Russian layout) is folded letter for letter —
    transliterating it would read the Х as «Kh». A Cyrillic word is
    transliterated, and keeps its capitals when it was written in capitals."""
    if not _CYRILLIC.search(w):
        return w
    if _LATIN.search(w) and all(ch in _TWINS for ch in _CYRILLIC.findall(w)):
        return w.translate(_TWIN_TABLE)
    out = transliterate(w, "uz")
    return out.upper() if w.isupper() else out


def _name(v) -> str:
    """A person's name in Latin letters with its spacing folded — the registry
    spells people in both alphabets, sometimes inside one word."""
    return " ".join(_word(w) for w in str(v).split()) if v else ""


def _who(v) -> str:
    """The identity two spellings of one name are counted under."""
    return _name(v).upper()


def read_registry() -> tuple[list[dict] | None, str]:
    """The «Liderlar» registry tab, read live off the crawl's own client and
    parser. ``(None, reason)`` when it cannot be read — the caller then lists
    the sheets the stored rows name, and the file says so."""
    if not os.path.exists(settings.google_credentials_file):
        return None, "Google kaliti bu serverda yo'q"
    try:
        from app.services import worker_concerns
        return worker_concerns.read_registry(worker_concerns._crawl_client()), ""
    except PermissionError:
        # gspread turns a 403 into a BARE PermissionError — no message at all.
        return None, "Google ruxsat bermadi (403)"
    except Exception as exc:  # a report degrades, it never dies on the registry
        return None, (str(exc) or type(exc).__name__)[:200]


def _blank(sid: str | None, cell, leader, brigadir, in_registry) -> dict:
    return {
        "sheet_id": sid, "cell": (cell or "").strip(),
        "reg_leader": _name(leader), "reg_brigadir": _name(brigadir),
        "in_registry": in_registry,
        "rows": 0, "before": 0, "after": 0, "recent": 0,
        "open": 0, "done": 0, "deferred": 0, "other": 0,
        "days": Counter(), "owners": set(), "first": None, "last": None,
        "ever": 0, "last_ever": None, "prev": 0, "undated": 0,
        "app_rows": 0, "app_leader": "", "app_brigadir": "",
    }


def collect(db: Session, date_from: date, date_to: date, *,
            registry: list[dict] | None = None, registry_error: str = "",
            today: date | None = None) -> dict:
    """Everything the file and its caption state, computed once — the workbook
    and the caption are formatters and re-derive nothing.

    ``registry`` is `read_registry`'s list, or None when it could not be read
    (``registry_error`` then says why)."""
    today = today or datetime.now(TZ).date()
    recent_from = today - timedelta(days=RECENT_DAYS - 1)
    prev_to = date_from.replace(day=1) - timedelta(days=1)
    prev_from = prev_to.replace(day=1)
    month = _MONTHS[date_from.month - 1]

    # ── the platform's own register: a code's unit, leader, shift and plant ──
    # Columns, not entities, so a column this report never reads cannot make it
    # fail on a database a migration has not reached yet.
    mgrs = {int(m.id): m for m in db.query(
        Manager.id, Manager.name, Manager.shift, Manager.factory_id).all()}
    factories = {int(f.id): (f.name_uz or f.name_ru or f.code or str(f.id))
                 for f in db.query(Factory.id, Factory.name_uz, Factory.name_ru,
                                   Factory.code).all()}
    leader_names = {int(p.id): p.name for p in db.query(RoleProfile.id, RoleProfile.name)
                    .filter(RoleProfile.role == "leader").all()}
    cells: dict[str, dict] = {}
    for c in (db.query(Cell.verifix_code, Cell.manager_id, Cell.leader_id)
              .order_by(Cell.verifix_code).all()):
        k = _code(c.verifix_code)
        if not k or k in cells:
            continue
        mid = int(c.manager_id) if c.manager_id is not None else None
        m = mgrs.get(mid)
        cells[k] = {
            "code": (c.verifix_code or "").strip(),
            "unit_id": mid if m else None,
            "unit": _name(m.name) if m else "",
            "shift": m.shift if m else None,
            "factory": (factories.get(int(m.factory_id))
                        if m and m.factory_id is not None else None),
            "leader": _name(leader_names.get(int(c.leader_id))) if c.leader_id else "",
        }

    meta = db.query(WorkerConcernSyncMeta).filter_by(id=1).first()
    failures = {f.get("sheet_id"): f for f in ((meta.failures if meta else None) or [])
                if isinstance(f, dict) and f.get("sheet_id")}

    # ── every sheet: the registry first, so its CURRENT spelling names it ────
    # The stamp on the stored rows is only as fresh as that sheet's last crawl.
    sheets: dict[str, dict] = {}

    def sheet(sid: str, cell, leader, brigadir) -> dict:
        if sid not in sheets:
            sheets[sid] = _blank(sid, cell, leader, brigadir,
                                 None if registry is None else False)
        return sheets[sid]

    for e in registry or []:
        sheet(e["sheet_id"], e.get("cell"), e.get("leader"),
              e.get("brigadir"))["in_registry"] = True

    for sid, cell, leader, brig, n, last, undated, prev in (
            db.query(WorkerConcern.sheet_id,
                     func.max(WorkerConcern.reg_cell),
                     func.max(WorkerConcern.reg_leader),
                     func.max(WorkerConcern.reg_brigadir),
                     func.count(WorkerConcern.id),
                     func.max(WorkerConcern.date),
                     func.count(WorkerConcern.id).filter(WorkerConcern.date.is_(None)),
                     func.count(WorkerConcern.id).filter(WorkerConcern.date.between(
                         prev_from.isoformat(), prev_to.isoformat())))
            .group_by(WorkerConcern.sheet_id).all()):
        s = sheet(sid, cell, leader, brig)
        s.update(ever=int(n or 0), last_ever=_day(last),
                 undated=int(undated or 0), prev=int(prev or 0))

    # ── the month's sheet rows ────────────────────────────────────────────────
    register: list[dict] = []
    for r in (db.query(WorkerConcern.sheet_id, WorkerConcern.reg_cell,
                       WorkerConcern.reg_leader, WorkerConcern.reg_brigadir,
                       WorkerConcern.row_leader, WorkerConcern.owner,
                       WorkerConcern.text, WorkerConcern.date,
                       WorkerConcern.status, WorkerConcern.status_raw)
              .filter(WorkerConcern.date >= date_from.isoformat(),
                      WorkerConcern.date <= date_to.isoformat())
              .order_by(WorkerConcern.date.desc(), WorkerConcern.reg_cell,
                        WorkerConcern.id).all()):
        d = _day(r.date)
        if d is None:
            continue
        s = sheet(r.sheet_id, r.reg_cell, r.reg_leader, r.reg_brigadir)
        st = r.status or "todo"
        s["rows"] += 1
        s["before" if d < APP_FROM else "after"] += 1
        if d >= recent_from:
            s["recent"] += 1
        s["open" if st in ("todo", "doing") else st if st in ("done", "deferred")
          else "other"] += 1
        s["days"][d] += 1
        if (r.owner or "").strip():
            s["owners"].add(_who(r.owner))
        s["first"] = d if s["first"] is None else min(s["first"], d)
        s["last"] = d if s["last"] is None else max(s["last"], d)
        register.append({"date": d, "sheet_id": r.sheet_id,
                         "row_leader": _name(r.row_leader), "owner": _name(r.owner),
                         "text": r.text or "",
                         "status": STATUS_LABEL.get(st) or (r.status_raw or st)})

    # ── the month's in-app filings, per cell ─────────────────────────────────
    app: dict[str, dict] = {}
    for c in (db.query(LeaderConcern.cell_code, LeaderConcern.worker_name,
                       LeaderConcern.entry_date, LeaderConcern.leader_name,
                       LeaderConcern.brigadir_name)
              .filter(LeaderConcern.worker_name.isnot(None),
                      LeaderConcern.entry_date >= date_from,
                      LeaderConcern.entry_date <= date_to).all()):
        k = _code(c.cell_code)
        a = app.setdefault(k, {"code": (c.cell_code or "").strip(), "rows": 0,
                               "workers": set(), "leader": _name(c.leader_name),
                               "brigadir": _name(c.brigadir_name)})
        a["rows"] += 1
        a["workers"].add(_who(c.worker_name))

    def place(s: dict, k: str) -> None:
        reg = cells.get(k) or {}
        s.update(key=k, code=reg.get("code") or s["cell"], in_cells=bool(reg),
                 leader_now=reg.get("leader") or s["app_leader"],
                 unit_id=reg.get("unit_id"), unit=reg.get("unit") or "",
                 shift=reg.get("shift"), factory=reg.get("factory"))

    listed: list[dict] = []
    for s in sheets.values():
        k = _code(s["cell"])
        a = app.get(k)
        if a:
            s.update(app_rows=a["rows"], app_leader=a["leader"],
                     app_brigadir=a["brigadir"])
        f = failures.get(s["sheet_id"])
        place(s, k)
        s.update(n_days=len(s["days"]), n_owners=len(s["owners"]),
                 failed=(FAIL_LABEL.get(f.get("code"), FAIL_LABEL["other"])
                         if f else None),
                 url=SHEET_URL.format(s["sheet_id"]))
        s["state"] = (ST_BOTH if s["rows"] and s["app_rows"] else
                      ST_SHEET if s["rows"] else
                      ST_MOVED if s["app_rows"] else ST_NONE)
        listed.append(s)

    # A cell that filed in the app and has no sheet at all is listed too — it
    # is part of «who uses what», and dropping it would shrink the app's half.
    sheet_keys = {s["key"] for s in listed}
    for k, a in app.items():
        if k in sheet_keys:
            continue
        s = _blank(None, a["code"], None, None, None)
        s.update(app_rows=a["rows"], app_leader=a["leader"],
                 app_brigadir=a["brigadir"], n_days=0, n_owners=0,
                 failed=None, url=None, state=ST_APP)
        place(s, k)
        listed.append(s)

    # ── per brigadir: the platform's unit where the code resolves, else the
    # registry's own spelling — never a guess about which unit a code means ──
    units: dict[str, dict] = {}
    for s in listed:
        if s["unit_id"] is not None:
            key, uname = f"u{s['unit_id']}", s["unit"]
        else:
            uname = s["reg_brigadir"] or s["app_brigadir"]
            key = f"r{uname.upper()}"
        u = units.get(key)
        if u is None:
            u = units[key] = {"name": uname, "shift": s["shift"],
                              "factory": s["factory"], "sheets": 0, "still": 0,
                              "after_sheets": 0, "rows": 0, "after": 0,
                              "recent": 0, "app_keys": set(), "app_rows": 0,
                              "last": None}
        s["unit_key"] = key
        if s["sheet_id"]:
            u["sheets"] += 1
        if s["rows"]:
            u["still"] += 1
            u["rows"] += s["rows"]
            u["after"] += s["after"]
            u["recent"] += s["recent"]
            u["after_sheets"] += 1 if s["after"] else 0
            u["last"] = s["last"] if u["last"] is None else max(u["last"], s["last"])
        # Two sheets can name one cell; its app filings count once.
        if s["app_rows"] and s["key"] not in u["app_keys"]:
            u["app_keys"].add(s["key"])
            u["app_rows"] += s["app_rows"]
    for u in units.values():
        u["app_cells"] = len(u.pop("app_keys"))

    still = sorted((s for s in listed if s["rows"]),
                   key=lambda s: (-s["last"].toordinal(), -s["rows"], s["code"]))
    everyone = sorted(listed, key=lambda s: (STATE_ORDER[s["state"]], -s["rows"],
                                             -s["app_rows"], s["code"]))
    unit_list = sorted(units.values(),
                       key=lambda u: (-u["rows"], -u["app_rows"], u["name"]))

    by_sid = {s["sheet_id"]: s for s in listed if s["sheet_id"]}
    for x in register:
        s = by_sid[x["sheet_id"]]
        x.update(code=s["code"], reg_leader=s["reg_leader"],
                 reg_brigadir=s["reg_brigadir"], url=s["url"])

    last_row = max((s["last"] for s in still), default=None)
    axis_to = min(date_to, max(today, last_row) if last_row else today)
    axis = [date_from + timedelta(days=i)
            for i in range((axis_to - date_from).days + 1)]

    counts = {
        "still": len(still),
        "rows": sum(s["rows"] for s in still),
        "after_sheets": sum(1 for s in still if s["after"]),
        "after_rows": sum(s["after"] for s in still),
        "recent_sheets": sum(1 for s in still if s["recent"]),
        "recent_rows": sum(s["recent"] for s in still),
        "open": sum(s["open"] for s in still),
        "done": sum(s["done"] for s in still),
        "leaders": len({_who(s["reg_leader"]) for s in still if s["reg_leader"]}),
        "no_leader": sum(1 for s in still if not s["reg_leader"]),
        "brigadirs": len({s["unit_key"] for s in still}),
        "only_sheet": sum(1 for s in still if s["state"] == ST_SHEET),
        "app_cells": len(app),
        "app_rows": sum(a["rows"] for a in app.values()),
        "both": len({s["key"] for s in still if s["app_rows"]}),
        "moved": sum(1 for s in listed if s["state"] == ST_MOVED),
        "silent": sum(1 for s in listed if s["state"] == ST_NONE),
        "sheets": sum(1 for s in listed if s["sheet_id"]),
        "undated": sum(s["undated"] for s in listed),
    }

    notes: list[str] = []
    # Freshness first: a crawl that stopped working would otherwise read as
    # «nobody writes in the sheets any more», the one wrong answer that looks
    # like good news.
    synced = meta.last_synced.astimezone(TZ) if meta and meta.last_synced else None
    stale = synced is None or datetime.now(TZ) - synced > SYNC_STALE_AFTER
    if synced is None:
        notes.append("Jadvallar hali bir marta ham to'liq o'qilmagan — fayldagi "
                     "jadval raqamlari to'liq emas.")
    elif stale:
        notes.append(f"Jadvallar oxirgi marta {synced:%d.%m.%Y %H:%M} da o'qilgan — "
                     "undan keyin yozilgan havotirlar bu faylda yo'q.")
    if meta and meta.ok is False and meta.message and not failures:
        notes.append(f"Oxirgi o'qish xato bilan to'xtadi: {meta.message[:240]}")
    if registry is None:
        why = registry_error or "sababi noma'lum"
        notes.append(f"Registr o'qilmadi ({why}) — ro'yxat saqlangan qatorlardan "
                     "tuzildi, shuning uchun hech qachon yozilmagan jadvallar unda yo'q.")
    if failures:
        bad = ", ".join(f"{f.get('cell') or '—'} ({FAIL_LABEL.get(f.get('code'), FAIL_LABEL['other'])})"
                        for f in failures.values())
        notes.append(f"{len(failures)} ta jadval oxirgi o'qishda o'qilmadi — ularning "
                     f"qatorlari avvalgi o'qishdagi holicha qoldi: {bad}.")
    removed = [s["code"] for s in still if s["in_registry"] is False]
    if removed:
        notes.append(f"{len(removed)} ta jadval registrdan olib tashlangan, lekin {month} "
                     f"qatorlari bazada qolgan (qayta o'qilmaydi): {', '.join(removed)}.")
    unknown = sorted({s["code"] for s in still if not s["in_cells"]})
    if unknown:
        notes.append(f"{len(unknown)} ta yacheyka kodi /cells registrida topilmadi — "
                     f"hozirgi lider, smena va zavod bo'sh: {', '.join(unknown)}.")
    if counts["no_leader"]:
        notes.append(f"{counts['no_leader']} ta jadvalga registrda lider yozilmagan — "
                     "«Lider (registrda)» bo'sh.")
    if counts["undated"]:
        notes.append(f"{counts['undated']} ta qatorning sanasi o'qilmadi (butun tarix "
                     "bo'yicha) — ular hech bir oyga qo'shilmadi.")

    return {
        "from": date_from, "to": date_to, "today": today,
        "recent_from": recent_from, "month": month, "year": date_from.year,
        "prev_month": _MONTHS[prev_from.month - 1],
        "synced": synced, "stale": stale,
        "still": still, "everyone": everyone, "units": unit_list,
        "register": register, "axis": axis, "counts": counts, "notes": notes,
    }


# ── the workbook ─────────────────────────────────────────────────────────────
def _title(rep: dict) -> tuple[str, str]:
    return (f"Havotirlar hali Google jadvalda — {rep['month']} {rep['year']}",
            f"{_dmy(rep['from'])} – {_dmy(rep['to'])} · «Liderlar Havotirlar» "
            "jadvallari va /cell-concerns ilovasi · barcha brigadirlar · ikkala smena")


def _scope(rep: dict) -> list[dict]:
    return [
        {"label": "Davr", "value": f"{_dmy(rep['from'])} – {_dmy(rep['to'])}"},
        {"label": "Jadvallar o'qilgan", "value":
            rep["synced"].strftime("%d.%m.%Y %H:%M") if rep["synced"] else "hali yo'q"},
        {"label": "Ilova ochilgan", "value": APP_OPENED},
        {"label": "Oxirgi 7 kun", "value":
            f"{_dmy(rep['recent_from'])} – {_dmy(rep['today'])}"},
        {"label": "Jadval qatori", "value": "matni bor · «Дата заполнения»"},
        {"label": "Ilova qatori", "value": "/cell-concerns · kiritilgan kun"},
        {"label": "Egasi", "value": "registrdagi yacheyka va lider"},
        {"label": "Bu hisobot", "value": "faqat xabar beradi"},
    ]


def _kpis(rep: dict) -> list[dict]:
    c, m, a = rep["counts"], rep["month"], _dm(APP_FROM)
    return [
        {"value": c["still"], "label": "Hali jadvalga yozayotgan yacheyka",
         "color": ORANGE, "hint": f"{c['leaders']} lider · {c['brigadirs']} brigadir"},
        {"value": c["after_sheets"], "label": f"{a} dan keyin ham jadvalda",
         "color": RED, "hint": f"{c['after_rows']} havotir · ilova ochilgandan keyin"},
        {"value": c["recent_sheets"], "label": "Oxirgi 7 kunda jadvalda",
         "color": AMBER, "hint": f"{c['recent_rows']} havotir"},
        {"value": c["rows"], "label": f"Jadvaldagi havotir ({m})", "color": SLATE,
         "hint": f"{c['open']} ochiq · {c['done']} bajarilgan"},
        {"value": c["app_cells"], "label": "Ilovada yozgan yacheyka", "color": GREEN,
         "hint": f"{c['app_rows']} havotir · {c['both']} tasi jadvalda ham"},
        {"value": c["moved"], "label": "Jadvaldan ilovaga o'tgan", "color": INDIGO,
         "hint": f"{c['silent']} jadvalda {m}da hech narsa yo'q"},
    ]


def _how(rep: dict) -> str:
    m, a = rep["month"], _dm(APP_FROM)
    return (
        f"Qanday o'qiladi. {APP_OPENED} dan ishchi havotirini /cell-concerns "
        "sahifasida, liderning kompyuterida yozadi. «Liderlar Havotirlar» Google "
        f"jadvallari o'chirilmagan — bu fayl {m}da kim hali ham jadvalga yozayotganini "
        "ko'rsatadi. Havotir = jadvaldagi matni bor qator, sanasi — o'sha qatordagi "
        "«Дата заполнения». Yacheyka, lider va brigadir — registrdagi yozuv; "
        "«Hozirgi lider», smena va zavod — platformaning /cells registridan. "
        f"«{a} dan» — ilova ochilgan kundan boshlab (qatorda vaqt yo'q, shuning uchun "
        f"{a} ning o'zi ham shu yerda). Holat: «Faqat jadvalda» — {m}da faqat "
        "jadvalga yozgan; «Jadvalda va ilovada» — ikkalasiga; «Ilovaga o'tgan» — "
        "jadvalga yozmagan, ilovada yozgan; «Hech qayerda yozilmagan» — ikkalasida "
        "ham havotir yo'q. Hisobot hech narsani o'zgartirmaydi."
    )


def _row_height(text: str, chars_per_line: int = 150) -> float:
    return 15 * max(1, -(-len(text) // chars_per_line)) + 4


def _summary_sheet(wb: Workbook, rep: dict) -> None:
    ws = _sheet(wb, "Xulosa", {2: 32, 3: 9, 4: 13, 5: 11, 6: 11, 7: 12, 8: 12,
                               9: 12, 10: 11, 11: 12, 12: 12, 13: 13})
    title, subtitle = _title(rep)
    r = _banner(ws, 2, 2, 13, title, subtitle)
    r = _meta_strip(ws, r, 2, 13, _scope(rep))
    r = _kpi_cards(ws, r, 2, _kpis(rep))

    how = _how(rep)
    _unp_cell(ws, r, 2, _xl(how), _fill(BRAND_SOFT), align=WRAP, size=10)
    ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=13)
    ws.row_dimensions[r].height = _row_height(how, 140)
    r += 2

    a = _dm(APP_FROM)
    r = _section(ws, r, 2, 13, "Brigadirlar bo'yicha",
                 f"{len(rep['units'])} brigadir · jadvaldagi havotir soni bo'yicha")
    head = r
    r = _unp_head(ws, r, 2, [("Brigadir", 32), ("Smena", 9), ("Zavod", 13),
                             ("Jadvallar", 11), ("Hali jadvalda", 11),
                             (f"{a} dan keyin ham", 12), ("Jadvaldagi havotir", 12),
                             (f"shundan {a} dan", 12), ("Oxirgi 7 kun", 11),
                             ("Ilovada: yacheyka", 12), ("Ilovada: havotir", 12),
                             ("Oxirgi jadval sanasi", 13)])
    first = r
    tot = Counter()
    last_all = None
    for i, u in enumerate(rep["units"]):
        bg = _fill(PANEL if i % 2 == 0 else BAND)
        _unp_cell(ws, r, 2, _xl(u["name"]), bg, bold=bool(u["still"]))
        _unp_cell(ws, r, 3, u["shift"], bg, align=CENTER)
        _unp_cell(ws, r, 4, _xl(u["factory"]), bg, align=CENTER)
        _unp_cell(ws, r, 5, u["sheets"] or None, bg, fmt=NUM, align=RIGHT)
        _unp_cell(ws, r, 6, u["still"] or None, bg, fmt=NUM, align=RIGHT,
                  bold=bool(u["still"]), color=ORANGE if u["still"] else INK_SOFT)
        _unp_cell(ws, r, 7, u["after_sheets"] or None, bg, fmt=NUM, align=RIGHT,
                  bold=bool(u["after_sheets"]),
                  color=RED if u["after_sheets"] else INK_SOFT)
        _unp_cell(ws, r, 8, u["rows"] or None, bg, fmt=NUM, align=RIGHT)
        _unp_cell(ws, r, 9, u["after"] or None, bg, fmt=NUM, align=RIGHT)
        _unp_cell(ws, r, 10, u["recent"] or None, bg, fmt=NUM, align=RIGHT)
        _unp_cell(ws, r, 11, u["app_cells"] or None, bg, fmt=NUM, align=RIGHT,
                  color=GREEN if u["app_cells"] else INK_SOFT)
        _unp_cell(ws, r, 12, u["app_rows"] or None, bg, fmt=NUM, align=RIGHT,
                  color=GREEN if u["app_rows"] else INK_SOFT)
        _unp_cell(ws, r, 13, u["last"], bg, fmt=DATE_FMT, align=CENTER)
        for k in ("sheets", "still", "after_sheets", "rows", "after", "recent",
                  "app_cells", "app_rows"):
            tot[k] += u[k]
        if u["last"] and (last_all is None or u["last"] > last_all):
            last_all = u["last"]
        r += 1
    if r > first:
        ws.auto_filter.ref = f"B{head}:M{r - 1}"
        bg = _fill(BRAND_SOFT)
        _unp_cell(ws, r, 2, "Jami", bg, bold=True)
        _unp_cell(ws, r, 3, None, bg)
        _unp_cell(ws, r, 4, None, bg)
        for col, k in zip(range(5, 13), ("sheets", "still", "after_sheets", "rows",
                                         "after", "recent", "app_cells", "app_rows")):
            _unp_cell(ws, r, col, tot[k] or None, bg, fmt=NUM, align=RIGHT, bold=True)
        _unp_cell(ws, r, 13, last_all, bg, fmt=DATE_FMT, align=CENTER, bold=True)
        r += 1

    r += 1
    notes = rep["notes"]
    r = _section(ws, r, 2, 13, "Diqqat",
                 f"{len(notes)} ta izoh" if notes else "izoh yo'q")
    for n in notes or ["Hamma jadval o'qildi — ogohlantirish yo'q."]:
        _unp_cell(ws, r, 2, _xl(n), _fill(PANEL), align=WRAP, size=9.5)
        ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=13)
        ws.row_dimensions[r].height = _row_height(n)
        r += 1


def _link(ws, r: int, col: int, url: str | None, bg) -> None:
    if not url:
        _unp_cell(ws, r, col, None, bg, align=CENTER)
        return
    cell = _unp_cell(ws, r, col, "ochish", bg, align=CENTER, size=9, color=INDIGO)
    cell.hyperlink = url


def _put(ws, r: int, bg):
    """A cursor along one row: each call writes the next column."""
    col = iter(range(2, 200))

    def put(value, **kw):
        return _unp_cell(ws, r, next(col), value, bg, **kw)

    return put, col


def _still_sheet(wb: Workbook, rep: dict) -> None:
    a, m = _dm(APP_FROM), rep["month"]
    before = f"{_dm(rep['from'])}–{_dm(APP_FROM - timedelta(days=1))}"
    cols = [("№", 5), ("Yacheyka", 10), ("Lider (registrda)", 28),
            ("Brigadir (registrda)", 26), ("Hozirgi lider (/cells)", 26),
            ("Smena", 7), ("Zavod", 12), ("Jami havotir", 9), (before, 10),
            (f"{a} dan", 9), ("Oxirgi 7 kun", 9), ("Kun", 7), ("Ishchi", 8),
            ("Birinchi", 11), ("Oxirgi", 11), ("Ochiq", 8), ("Bajarilgan", 10),
            ("O'tqazish", 9), (f"Ilovada ({m})", 10), ("Holat", 18), ("Jadval", 8)]
    ws = _sheet(wb, "Hali jadvalda", {}, landscape=True)
    last = 1 + len(cols)
    r = _banner(ws, 2, 2, last, _title(rep)[0],
                f"{m.capitalize()}da Google jadvalga kamida bitta havotir yozilgan "
                "har bir yacheyka — oxirgi yozilgan kun bo'yicha, yangisi tepada")
    c = rep["counts"]
    r = _section(ws, r, 2, last, "Hali jadvalda",
                 f"{c['still']} yacheyka · {c['rows']} havotir · «{a} dan» — "
                 "ilova ochilgan kundan")
    head = r
    r = _unp_head(ws, r, 2, cols)
    first = r
    for i, s in enumerate(rep["still"], 1):
        bg = _fill(PANEL if i % 2 else BAND)
        put, col = _put(ws, r, bg)
        put(i, fmt=NUM, align=CENTER, color=INK_SOFT)
        put(_xl(s["code"]), align=CENTER, bold=True)
        put(_xl(s["reg_leader"]))
        put(_xl(s["reg_brigadir"]))
        put(_xl(s["leader_now"]), color=INK_SOFT)
        put(s["shift"], align=CENTER)
        put(_xl(s["factory"]), align=CENTER)
        put(s["rows"], fmt=NUM, align=RIGHT, bold=True)
        put(s["before"] or None, fmt=NUM, align=RIGHT)
        put(s["after"] or None, fmt=NUM, align=RIGHT, bold=bool(s["after"]),
            color=RED if s["after"] else INK_SOFT)
        put(s["recent"] or None, fmt=NUM, align=RIGHT, bold=bool(s["recent"]),
            color=ORANGE if s["recent"] else INK_SOFT)
        put(s["n_days"], fmt=NUM, align=RIGHT)
        put(s["n_owners"] or None, fmt=NUM, align=RIGHT)
        put(s["first"], fmt=DATE_FMT, align=CENTER)
        put(s["last"], fmt=DATE_FMT, align=CENTER, bold=True)
        put(s["open"] or None, fmt=NUM, align=RIGHT)
        put(s["done"] or None, fmt=NUM, align=RIGHT)
        put(s["deferred"] or None, fmt=NUM, align=RIGHT)
        put(s["app_rows"] or None, fmt=NUM, align=RIGHT, bold=bool(s["app_rows"]),
            color=GREEN if s["app_rows"] else INK_SOFT)
        put(STATE_LABEL[s["state"]], align=CENTER, size=9,
            bold=s["state"] == ST_SHEET, color=STATE_COLOR[s["state"]])
        _link(ws, r, next(col), s["url"], bg)
        r += 1
    if r > first:
        ws.auto_filter.ref = f"B{head}:{get_column_letter(last)}{r - 1}"
    else:
        _unp_cell(ws, r, 2, f"{m.capitalize()}da Google jadvalga hech kim yozmagan.",
                  _fill(PANEL), bold=True, color=GREEN)
        ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=last)
    ws.freeze_panes = f"E{first}"
    ws.print_title_rows = f"{head}:{head}"


def _daily_sheet(wb: Workbook, rep: dict) -> None:
    axis = rep["axis"]
    cols = ([("Yacheyka", 10), ("Lider (registrda)", 28)]
            + [(_dm(d), 6.5) for d in axis] + [("Jami", 8)])
    ws = _sheet(wb, "Kunlik", {}, landscape=True)
    last = 1 + len(cols)
    r = _banner(ws, 2, 2, last, _title(rep)[0],
                "Har bir yacheyka Google jadvalga kuniga nechta havotir yozgani")
    r = _section(ws, r, 2, last, "Kunma-kun",
                 f"sariq sarlavha — {_dm(APP_FROM)} dan, ilova ochilgandan keyin")
    head = r
    r = _unp_head(ws, r, 2, cols)
    for j, d in enumerate(axis):
        if d >= APP_FROM:
            ws.cell(head, 4 + j).fill = _fill(CREAM)
    first = r
    totals = Counter()
    for i, s in enumerate(rep["still"]):
        bg = _fill(PANEL if i % 2 == 0 else BAND)
        _unp_cell(ws, r, 2, _xl(s["code"]), bg, align=CENTER, bold=True)
        _unp_cell(ws, r, 3, _xl(s["reg_leader"]), bg)
        for j, d in enumerate(axis):
            n = s["days"].get(d, 0)
            totals[d] += n
            if n:
                _unp_cell(ws, r, 4 + j, n, _fill(BRAND_SOFT), fmt=NUM,
                          align=CENTER, bold=True)
            else:
                # Blank, not a dash: thirty columns of «—» bury the days that
                # carry something, which are the whole point of this sheet.
                cell = ws.cell(r, 4 + j)
                cell.fill = bg
                cell.border = BOX
        _unp_cell(ws, r, 4 + len(axis), s["rows"], bg, fmt=NUM, align=RIGHT,
                  bold=True)
        r += 1
    if rep["still"]:
        bg = _fill(BRAND_SOFT)
        _unp_cell(ws, r, 2, "Jami", bg, bold=True)
        _unp_cell(ws, r, 3, f"{len(rep['still'])} yacheyka", bg, bold=True)
        for j, d in enumerate(axis):
            _unp_cell(ws, r, 4 + j, totals[d] or None, bg, fmt=NUM, align=CENTER,
                      bold=True)
        _unp_cell(ws, r, 4 + len(axis), rep["counts"]["rows"], bg, fmt=NUM,
                  align=RIGHT, bold=True)
    ws.freeze_panes = f"D{first}"
    ws.print_title_rows = f"{head}:{head}"


def _register_sheet(wb: Workbook, rep: dict) -> None:
    cols = [("Sana", 11), ("Yacheyka", 10), ("Lider (registrda)", 26),
            ("Brigadir (registrda)", 24), ("Qatordagi lider", 24),
            ("Havotir egasi", 22), ("Havotir", 70), ("Holat", 11), ("Jadval", 8)]
    ws = _sheet(wb, "Reyestr", {}, landscape=True)
    last = 1 + len(cols)
    r = _banner(ws, 2, 2, last, _title(rep)[0],
                f"{rep['month'].capitalize()}da Google jadvallarga yozilgan har bir "
                "havotir — xom dalil")
    r = _section(ws, r, 2, last, "Reyestr",
                 f"{len(rep['register'])} havotir · yangisi tepada")
    head = r
    r = _unp_head(ws, r, 2, cols)
    first = r
    for i, x in enumerate(rep["register"]):
        bg = _fill(PANEL if i % 2 == 0 else BAND)
        _unp_cell(ws, r, 2, x["date"], bg, fmt=DATE_FMT, align=CENTER)
        _unp_cell(ws, r, 3, _xl(x["code"]), bg, align=CENTER, bold=True)
        _unp_cell(ws, r, 4, _xl(x["reg_leader"]), bg)
        _unp_cell(ws, r, 5, _xl(x["reg_brigadir"]), bg)
        _unp_cell(ws, r, 6, _xl(x["row_leader"]), bg, size=9, color=INK_SOFT)
        _unp_cell(ws, r, 7, _xl(x["owner"]), bg)
        _unp_cell(ws, r, 8, _xl(x["text"][:32000]), bg, size=9)
        _unp_cell(ws, r, 9, _xl(x["status"]), bg, align=CENTER, size=9)
        _link(ws, r, 10, x["url"], bg)
        r += 1
    if r > first:
        ws.auto_filter.ref = f"B{head}:{get_column_letter(last)}{r - 1}"
    ws.freeze_panes = f"D{first}"
    ws.print_title_rows = f"{head}:{head}"


def _all_sheet(wb: Workbook, rep: dict) -> None:
    a, m = _dm(APP_FROM), rep["month"]
    cols = [("Holat", 20), ("Yacheyka", 10), ("Lider (registrda)", 26),
            ("Brigadir (registrda)", 24), ("Hozirgi lider (/cells)", 24),
            ("Smena", 7), ("Zavod", 12), (f"Jadval: {m}", 10),
            (f"shundan {a} dan", 10), (f"Ilova: {m}", 10),
            ("Oxirgi jadval sanasi", 12), (f"Jadval: {rep['prev_month']}", 11),
            ("Jadvaldagi jami havotir", 12), ("Sanasi o'qilmagan", 10),
            ("Registrda", 9), ("Oxirgi o'qish", 18), ("Jadval", 8)]
    ws = _sheet(wb, "Barcha yacheykalar", {}, landscape=True)
    last = 1 + len(cols)
    r = _banner(ws, 2, 2, last, _title(rep)[0],
                "Registrdagi har bir jadval va ilovada yozgan har bir yacheyka — "
                f"{m}dagi holati bilan")
    c = rep["counts"]
    r = _section(ws, r, 2, last, "Barcha yacheykalar",
                 f"{c['sheets']} jadval · {len(rep['everyone'])} qator · "
                 "«Jadval: …» — faqat Google jadval, «Ilova: …» — faqat /cell-concerns")
    head = r
    r = _unp_head(ws, r, 2, cols)
    first = r
    for i, s in enumerate(rep["everyone"]):
        bg = _fill(PANEL if i % 2 == 0 else BAND)
        put, col = _put(ws, r, bg)
        put(STATE_LABEL[s["state"]], size=9, bold=s["state"] == ST_SHEET,
            color=STATE_COLOR[s["state"]])
        put(_xl(s["code"]), align=CENTER, bold=True)
        put(_xl(s["reg_leader"]))
        put(_xl(s["reg_brigadir"]))
        put(_xl(s["leader_now"]), color=INK_SOFT)
        put(s["shift"], align=CENTER)
        put(_xl(s["factory"]), align=CENTER)
        put(s["rows"] or None, fmt=NUM, align=RIGHT, bold=bool(s["rows"]))
        put(s["after"] or None, fmt=NUM, align=RIGHT,
            color=RED if s["after"] else INK_SOFT)
        put(s["app_rows"] or None, fmt=NUM, align=RIGHT,
            color=GREEN if s["app_rows"] else INK_SOFT)
        put(s["last_ever"], fmt=DATE_FMT, align=CENTER)
        put(s["prev"] or None, fmt=NUM, align=RIGHT)
        put(s["ever"] or None, fmt=NUM, align=RIGHT)
        put(s["undated"] or None, fmt=NUM, align=RIGHT)
        put(None if s["in_registry"] is None else ("ha" if s["in_registry"] else "yo'q"),
            align=CENTER, size=9,
            color=RED if s["in_registry"] is False else INK_SOFT)
        put(None if not s["sheet_id"] else
            (f"o'qilmadi: {s['failed']}" if s["failed"] else "o'qildi"),
            align=CENTER, size=9, color=RED if s["failed"] else INK_SOFT)
        _link(ws, r, next(col), s["url"], bg)
        r += 1
    if r > first:
        ws.auto_filter.ref = f"B{head}:{get_column_letter(last)}{r - 1}"
    ws.freeze_panes = f"D{first}"
    ws.print_title_rows = f"{head}:{head}"


def build_workbook(rep: dict) -> BytesIO:
    """`collect`'s output as the five-sheet file. A formatter: it re-derives
    nothing, so the file and the caption can only ever state one set of
    numbers."""
    wb = Workbook()
    wb.remove(wb.active)
    _summary_sheet(wb, rep)
    _still_sheet(wb, rep)
    _daily_sheet(wb, rep)
    _register_sheet(wb, rep)
    _all_sheet(wb, rep)
    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


def caption(rep: dict) -> str:
    """Telegram caps a document caption at 1024 chars, so this is the finding
    and the figures that frame it — the file carries the names."""
    c, m, a = rep["counts"], rep["month"], _dm(APP_FROM)
    synced = rep["synced"].strftime("%d.%m %H:%M") if rep["synced"] else "—"
    lines = [f"📋 <b>Havotirlar hali Google jadvalda — {m} {rep['year']}</b>",
             f"{_dmy(rep['from'])} – {_dmy(rep['to'])} · jadvallar o'qilgan: {synced}",
             ""]
    if rep["stale"]:
        lines.insert(2, f"⚠️ <b>Jadvallar {synced} dan beri o'qilmagan</b> — undan "
                        "keyingi yozuvlar faylda yo'q")
    if c["still"]:
        lines += [
            f"🟠 <b>{c['still']}</b> yacheyka {m}da jadvalga yozgan — {c['rows']} "
            f"havotir ({c['leaders']} lider, {c['brigadirs']} brigadir)",
            f"🔴 <b>{c['after_sheets']}</b> tasi {a} dan keyin ham, ilova "
            f"ochilgandan keyin — {c['after_rows']} havotir",
            f"🟡 <b>{c['recent_sheets']}</b> tasi oxirgi 7 kunda "
            f"({_dm(rep['recent_from'])}–{_dm(rep['today'])}) — {c['recent_rows']} havotir",
            f"⛔ <b>{c['only_sheet']}</b> tasi {m}da ilovada umuman yozmagan",
        ]
    else:
        lines.append(f"✅ {m.capitalize()}da Google jadvalga hech kim yozmagan")
    lines += [
        f"🟢 <b>{c['app_cells']}</b> yacheyka ilovada yozgan — {c['app_rows']} "
        f"havotir; {c['both']} tasi jadvalda ham",
        f"🔵 <b>{c['moved']}</b> yacheyka jadvaldan ilovaga o'tgan · {c['silent']} "
        f"ta jadvalda {m}da hech narsa yo'q",
    ]
    if rep["notes"]:
        lines.append(f"⚠️ {len(rep['notes'])} ta izoh — «Xulosa» varag'ining oxirida")
    lines += ["", "<i>«Xulosa» — brigadirlar bo'yicha · «Hali jadvalda» — kim, qancha, "
                  "qachon · «Kunlik» — kunma-kun · «Reyestr» — har bir havotir · "
                  "«Barcha yacheykalar» — har bir jadvalning holati</i>"]
    text = "\n".join(lines)
    return text if len(text) <= 1024 else "\n".join(lines[:-2])


def send_xlsx(db: Session, chat_id: int, date_from: date, date_to: date) -> int:
    """Read the registry, build the workbook and DM it. Returns 1 on delivery.

    Deliberately NOT `xlsx_delivery.deliver_file`: that decides between a
    browser download and a Telegram DM from the REQUEST it was called on, and a
    boot job has no request — its one surface is the chat."""
    registry, err = read_registry()
    rep = collect(db, date_from, date_to, registry=registry, registry_error=err)
    stamp = datetime.now(TZ).strftime("%d.%m.%Y")
    name = f"havotirlar-google-jadvalda-{rep['month']}-{stamp}.xlsx"
    r = requests.post(
        f"https://api.telegram.org/bot{settings.telegram_bot_token}/sendDocument",
        data={"chat_id": chat_id, "caption": caption(rep), "parse_mode": "HTML"},
        files={"document": (name, build_workbook(rep).getvalue(),
                            "application/vnd.openxmlformats-officedocument."
                            "spreadsheetml.sheet")},
        timeout=180)
    j = r.json()
    if not j.get("ok"):
        raise RuntimeError(j.get("description") or f"HTTP {r.status_code}")
    return 1
