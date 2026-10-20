"""One-off: did Yog'mirov Feruz's 18.10 загрузка change AFTER he closed the day?

The supervisor wrote on 2026-10-20 (17:58), under a screenshot of the shift-2
comparison grid on /zagruzka: «mani 18.10 kunim bunaqa mas edi, hozir qizarib
qolgan» — his 18.10 did not look like this when he closed it, and now it is red
(A 64–66% on the rows that read 100–143% on every other day). Production is not
readable from the machine that wrote this, so the answer is DMed from the server.

A closed (unit, day) is computed from FOUR inputs at read time
(`brigadirs.build_metrics_list`), and the day-close lock (`idle_lock`) guards
only the per-day writers of /production and /idle-cell. So the report reads
every input with its own timestamps and puts each against the moment the day
was closed:

  1. ПЛАН/ФАКТ — `pp_daily` (an SAP ingest DELETES the date's rows and clears
     its overrides, and checks no lock) and `pp_line_daily` (typed per line),
     priced by the dated catalog (`pp_catalog_versions`).
  2. HEADCOUNT — the typed «Bugungi fakt» pins (`pp_work_center_daily.people`).
  3. OJIDANIYA — the APPROVED intervals of the cells that stood in the unit that
     day (`cell_ojidaniya_intervals` × `cell_link_events`). An approval, an edit
     or a back-dated cell move after the close changes the day without touching
     any row filed under it.
  4. ATTENDANCE — `attendance` rows carry no timestamp, so the live day's copy
     (`live_projections.copied_at`), its close, documents, exchanges, clock
     fixes, deletions, placements and the edit requests stand in for them.

Plus the Jurnal (`action_logs`) for the unit, the day and the unit's cells, and
the shift's grid for 16–19.10 in all three comparison tables, computed the way
`components/charts/ComparisonTable.jsx` computes them, so the screenshot's row
can be matched to a unit and to its table.

It READS and writes nothing but its flag (written by `startup._send_report_once`).

Temporary: delete this module, `startup.report_closed_day_trace` /
`_closed_day_trace_job` and the call in BOTH entrypoints once it has been sent.
"""
from __future__ import annotations

import statistics
import time
from datetime import date, datetime, time as dtime, timedelta, timezone
from io import BytesIO
from zoneinfo import ZoneInfo

import requests
from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Font
from sqlalchemy import Text, and_, cast, or_
from sqlalchemy.orm import Session

from app.config import settings
from app.models import (
    ActionLog, Attendance, AttendanceBatch, Cell, CellLinkEvent,
    CellOjidaniyaInterval, DayApproval, EditRequest, LiveClockFix, LiveDayClose,
    LiveDeletion, LiveDocument, LivePlacement, LiveProjection, LiveStaffEvent,
    Manager, PPCatalogVersion, PPDaily, PPLineDaily, PPManagerSetting, PPUpload,
    PPWorkCenterDaily,
)
from app.routers.brigadirs import build_metrics_list
from app.services import cell_links, zagruzka_source
from app.services.day_state import day_state
from app.services.idle_source import cells_with_moved
from app.services.name_map import _norm

# `name_map._norm` folds Ёгмиров / Yog'mirov / Yogmirov / Yo‘g‘mirov to this.
UNIT_KEY = "OGMIROV FERUZ"
DAY = date(2026, 10, 18)
GRID_FROM, GRID_TO = date(2026, 10, 16), date(2026, 10, 19)
CTX_FROM = date(2026, 10, 14)
TZ = ZoneInfo("Asia/Tashkent")
_API = "https://api.telegram.org"
SEND_RETRIES = 3
_JURNAL_MAX = 600

# utils/formulas.js — the page's constants, never a second spelling of them.
SHIFT_MIN = 480
PRODUCTIVE_SHARE = 0.9
VERIFIX_EFFICIENCY = 0.85
KAIZEN_BUFFER = 10


def _f(v):
    return None if v is None else float(v)


def _g(v, nd: int = 0) -> str:
    """A number as it reads: 1 240 not 1240.0, «—» for nothing."""
    if v is None:
        return "—"
    return f"{float(v):,.{nd}f}".replace(",", " ")


def _pct(v) -> str:
    return "—" if v is None else f"{round(v * 100)}%"


def _aware(v):
    if v is None:
        return None
    return v if v.tzinfo else v.replace(tzinfo=timezone.utc)


def _dt(v) -> str:
    v = _aware(v)
    return "—" if v is None else v.astimezone(TZ).strftime("%d.%m %H:%M")


def _d(v) -> str:
    if v is None:
        return "—"
    if isinstance(v, str):
        v = date.fromisoformat(v[:10]) if "-" in v else datetime.strptime(v, "%d.%m.%Y").date()
    return v.strftime("%d.%m")


def _after(v, ref) -> bool:
    v = _aware(v)
    return bool(v and ref and v > ref)


def _today() -> date:
    return datetime.now(TZ).date()


# ── the unit ──────────────────────────────────────────────────────────────────

def _unit(db: Session) -> Manager:
    """By NAME, never by id — `startup.MANAGERS` ids are a stale seed and a
    checkout's ids are not production's."""
    surname, given = UNIT_KEY.split()
    near = [m for m in db.query(Manager).all() if surname in _norm(m.name or "").split()]
    for pick in ([m for m in near if not m.archived],
                 [m for m in near if not m.archived and given in _norm(m.name or "")],
                 [m for m in near if _norm(m.name or "") == UNIT_KEY]):
        if len(pick) == 1:
            return pick[0]
    roster = [f"#{m.id} {m.name} (shift {m.shift}{', archived' if m.archived else ''})"
              for m in near]
    raise RuntimeError(f"unit «{UNIT_KEY}» did not resolve to one: {roster or 'none'}")


# ── the three comparison tables, as ComparisonTable.jsx reads them ────────────

def _simple(minutes, hc):
    den = SHIFT_MIN * PRODUCTIVE_SHARE * float(hc or 0)
    return (float(minutes) / den) if minutes and den > 0 else None


def _full90_a(m):
    """`formulas.full90ActualUtil`: the full A with Verifix hours at 0.9."""
    if m.net_util is None:
        return None
    base = m.avail_min
    if base is None and m.prod_plan:
        base = SHIFT_MIN * (m.prod_actual / m.prod_plan)
    if not base or base <= 0:
        return None
    pa, hc = float(m.prod_actual or 0), float(m.official_hc or 0)
    if m.verifix_labor is not None:
        vl = float(m.verifix_labor)
    elif m.effective_hc is not None:
        vl = pa + (m.effective_hc - hc) * base
    else:
        return None
    ehc = hc + ((vl / VERIFIX_EFFICIENCY) * PRODUCTIVE_SHARE - pa) / base
    den = ehc * (base - float(m.equip_downtime or 0) - float(m.avg_early_arrival or 0)
                 - KAIZEN_BUFFER)
    return pa / den if den > 0 else None


def _row(m) -> dict:
    hc = float(m.official_hc or 0)
    vmin = (float(m.verifix_labor) / VERIFIX_EFFICIENCY) if m.verifix_labor else None
    return {
        "date": datetime.strptime(m.date, "%d.%m.%Y").date(),
        "full_p": m.baseline_util, "full_a": m.net_util,
        "simple_p": _simple(m.prod_plan, hc), "simple_a": _simple(m.prod_actual, hc),
        "full90_a": _full90_a(m),
        "plan_min": float(m.prod_plan or 0), "fact_min": float(m.prod_actual or 0),
        "hc": hc, "verifix_min": vmin, "verifix_hc": float(m.verifix_hc or 0),
        "effective_hc": m.effective_hc, "avail_min": m.avail_min,
        "ojid": float(m.equip_downtime or 0), "early": float(m.avg_early_arrival or 0),
        "by_cat": dict(m.downtime_by_cat or {}), "basis": m.basis,
    }


def _grid(db: Session, shift) -> list[dict]:
    """Every unit of the shift, 16–19.10, all three tables — to match the
    screenshot's rows (sorted by name, as `/api/heatmap` sorts them)."""
    out: dict[str, dict] = {}
    for m in build_metrics_list(db, GRID_FROM, GRID_TO, shift, None, require_closed=False):
        u = out.setdefault(m.manager_name, {"id": m.manager_id, "name": m.manager_name,
                                            "days": {}})
        r = _row(m)
        r["state"] = day_state(db, m.manager_id, r["date"])[0]
        u["days"][r["date"]] = r
    return [out[k] for k in sorted(out)]


# ── what happened to the day ─────────────────────────────────────────────────

def _closes(db: Session, mid: int) -> dict:
    """Every trace of the day being closed or re-opened, oldest first, and the
    two moments everything else is measured against: the FIRST close (what the
    brigadir saw when he closed) and the current one."""
    jur = (db.query(ActionLog)
           .filter(ActionLog.unit_id == mid, ActionLog.day == DAY,
                   ActionLog.action.in_(("attendance.day_closed", "attendance.day_reopened")))
           .order_by(ActionLog.created_at).all())
    appr = db.query(DayApproval).filter_by(manager_id=mid, date=DAY).first()
    live = (db.query(LiveDayClose).filter_by(manager_id=mid, day=DAY)
            .order_by(LiveDayClose.closed_at).all())
    events = [(_aware(j.created_at), j.action.split(".")[-1], j.actor_name, j.source,
               j.outcome) for j in jur]
    if appr:
        events.append((_aware(appr.approved_at), "approval row", appr.approved_by_name,
                       "day_approvals", "current"))
    for lc in live:
        events.append((_aware(lc.closed_at), "live close", lc.closed_by_name,
                       "live_day_closes", ""))
    events.sort(key=lambda e: e[0] or datetime.max.replace(tzinfo=timezone.utc))
    shut = [e[0] for e in events if e[0] and e[1] in ("day_closed", "approval row", "live close")
            and e[4] in ("done", "current", "")]
    state, closure, counts = day_state(db, mid, DAY)
    return {
        "events": events, "first": min(shut) if shut else None,
        "current": _aware(appr.approved_at) if appr else None,
        "reopens": sum(1 for e in events if e[1] == "day_reopened" and e[4] == "done"),
        "state": state, "counts": counts,
    }


def _per_day(db: Session, mid: int, lo: date, hi: date) -> dict:
    """Per day: what each stored input holds and when it was last written."""
    days: dict[date, dict] = {}

    def day(d):
        return days.setdefault(d, {
            "pp_rows": 0, "pp_plan": 0.0, "pp_fact": 0.0, "pp_ovr": 0, "pp_last": None,
            "ln_rows": 0, "ln_ovr": 0, "ln_last": None,
            "pin_people": None, "pin_last": None,
            "att_rows": 0, "att_workers": set(), "att_hours": 0.0, "att_early": 0.0,
            "att_sup": 0, "att_split": 0,
            "iv_n": 0, "iv_appr": 0, "iv_last": None, "copied_at": None, "copy_rows": None,
        })

    def later(a, b):
        a, b = _aware(a), _aware(b)
        return b if a is None or (b and b > a) else a

    for r in db.query(PPDaily).filter(PPDaily.manager_id == mid,
                                      PPDaily.date >= lo, PPDaily.date <= hi):
        x = day(r.date)
        x["pp_rows"] += 1
        x["pp_plan"] += float((r.plan_override if r.plan_override is not None else r.plan_qty) or 0)
        x["pp_fact"] += float((r.actual_override if r.actual_override is not None else r.actual_qty) or 0)
        x["pp_ovr"] += (r.plan_override is not None) + (r.actual_override is not None)
        x["pp_last"] = later(x["pp_last"], r.updated_at)
    for r in db.query(PPLineDaily).filter(PPLineDaily.manager_id == mid,
                                          PPLineDaily.date >= lo, PPLineDaily.date <= hi):
        x = day(r.date)
        x["ln_rows"] += 1
        x["ln_ovr"] += (r.plan_override is not None) + (r.actual_override is not None)
        x["ln_last"] = later(x["ln_last"], r.updated_at)
    for r in db.query(PPWorkCenterDaily).filter(PPWorkCenterDaily.manager_id == mid,
                                                PPWorkCenterDaily.date >= lo,
                                                PPWorkCenterDaily.date <= hi):
        x = day(r.date)
        x["pin_last"] = later(x["pin_last"], r.updated_at)
    people = zagruzka_source.unit_people(zagruzka_source.typed_people(db, [mid], lo, hi))
    for (m, iso), n in people.items():
        day(date.fromisoformat(iso))["pin_people"] = n
    for r in db.query(Attendance.date, Attendance.worker_name, Attendance.hours_worked,
                      Attendance.early_arrival_min, Attendance.is_supervisor,
                      Attendance.split_of).filter(Attendance.manager_id == mid,
                                                  Attendance.date >= lo,
                                                  Attendance.date <= hi):
        x = day(r.date)
        x["att_rows"] += 1
        x["att_workers"].add((r.worker_name or "").strip().lower())
        x["att_hours"] += float(r.hours_worked or 0)
        x["att_early"] += float(r.early_arrival_min or 0)
        x["att_sup"] += bool(r.is_supervisor)
        x["att_split"] += bool(r.split_of)
    for r in db.query(LiveProjection.day, LiveProjection.rows, LiveProjection.copied_at).filter(
            LiveProjection.manager_id == mid, LiveProjection.day >= lo, LiveProjection.day <= hi):
        x = day(r.day)
        x["copied_at"], x["copy_rows"] = _aware(r.copied_at), r.rows
    for d in days.values():
        d["att_workers"] = len(d["att_workers"] - {""})
    return days


def _cells(db: Session, mid: int, lo: date, hi: date) -> tuple[list, dict]:
    """The cells that stood in the unit on some day of the range — today's, and
    any that has since moved away (`cell_links`) — and, per day, which of them
    stood HERE."""
    cells = cells_with_moved(db, db.query(Cell).filter(Cell.manager_id == mid).all(),
                             [mid], lo, hi)
    here: dict[date, set] = {}
    d = lo
    while d <= hi:
        on = cell_links.units_of_cells_on(cells, d)
        here[d] = {cid for cid, unit in on.items() if unit == mid}
        d += timedelta(days=1)
    return cells, here


def _intervals(db: Session, cells: list, here: dict, ref) -> list[dict]:
    if not cells:
        return []
    code = {c.id: c.verifix_code for c in cells}
    out = []
    for iv in (db.query(CellOjidaniyaInterval)
               .filter(CellOjidaniyaInterval.cell_id.in_(list(code)),
                       CellOjidaniyaInterval.date >= CTX_FROM.isoformat(),
                       CellOjidaniyaInterval.date <= _today().isoformat())
               .order_by(CellOjidaniyaInterval.date, CellOjidaniyaInterval.start)):
        d = date.fromisoformat(iv.date)
        stamps = [t for t in (iv.created_at, iv.updated_at, iv.decided_at) if t]
        out.append({
            "date": d, "cell": code.get(iv.cell_id), "here": iv.cell_id in here.get(d, set()),
            "category": iv.category, "start": iv.start, "end": iv.end,
            "stopped": bool(iv.stopped), "cleaners": iv.cleaners, "status": iv.status,
            "note": (iv.note or "")[:200], "by": iv.entered_by_profile,
            "decided_by": iv.decided_by_profile, "created": iv.created_at,
            "updated": iv.updated_at, "decided": iv.decided_at,
            "after_close": any(_after(t, ref) for t in stamps),
        })
    return out


def _links(db: Session, mid: int, cells: list) -> list[dict]:
    """Every dated brigadir change of those cells, and any naming this unit."""
    ids = [c.id for c in cells]
    q = db.query(CellLinkEvent).filter(CellLinkEvent.field == "manager_id", or_(
        CellLinkEvent.value == mid, CellLinkEvent.prev == mid,
        CellLinkEvent.cell_id.in_(ids or [-1])))
    names = {m.id: m.name for m in db.query(Manager.id, Manager.name)}
    code = {c.id: c.verifix_code for c in db.query(Cell.id, Cell.verifix_code)}
    return [{
        "cell": code.get(e.cell_id, e.cell_id), "from": names.get(e.prev, e.prev),
        "to": names.get(e.value, e.value), "effective_from": e.effective_from,
        "effective_at": e.effective_at, "applied": e.applied_at, "src": e.src,
        "by": e.by_name, "created": e.created_at,
        "reaches_day": e.effective_from is not None and e.effective_from <= DAY
                       and (e.value == mid or e.prev == mid),
    } for e in q.order_by(CellLinkEvent.created_at)]


def _jurnal(db: Session, mid: int, cells: list) -> list[dict]:
    """The unit's rows, the day's rows (any unit — an SAP upload is plant-wide
    and names no unit), and rows naming one of the unit's cells, from the
    morning of the day on."""
    since = datetime.combine(DAY, dtime(0, 0), tzinfo=TZ)
    codes = [c.verifix_code for c in cells if c.verifix_code]
    by_cell = [cast(ActionLog.details, Text).ilike(f"%{c}%") for c in codes[:40]]
    cell_ids = [str(c.id) for c in cells]
    q = (db.query(ActionLog)
         .filter(ActionLog.created_at >= since, or_(
             ActionLog.unit_id == mid,
             ActionLog.day == DAY,
             and_(ActionLog.target_kind == "cell", ActionLog.target_id.in_(cell_ids or ["-"])),
             *by_cell))
         .order_by(ActionLog.created_at).limit(_JURNAL_MAX))

    def why(j) -> str:
        """A text match on a 4-digit code can be a coincidence, so the DM lists
        only rows tied to the unit, the day or a cell; the rest go to the book."""
        if j.unit_id == mid:
            return "unit"
        if j.day == DAY:
            return "day"
        if j.target_kind == "cell" and j.target_id in cell_ids:
            return "cell"
        return "mention"

    return [{
        "why": why(j),
        "at": j.created_at, "action": j.action, "outcome": j.outcome, "source": j.source,
        "who": j.actor_name, "role": j.actor_role, "unit": j.unit_name,
        "unit_id": j.unit_id, "day": j.day, "target": j.target_name or j.target_id,
        "changes": j.changes, "details": j.details, "path": j.path,
        "version": j.app_version,
    } for j in q]


def collect(db: Session) -> dict:
    u = _unit(db)
    mid = u.id
    hi = _today()
    close = _closes(db, mid)
    ref = close["first"]

    mine = {}
    for m in build_metrics_list(db, CTX_FROM, hi, None, [mid], require_closed=False):
        r = _row(m)
        r["state"] = day_state(db, mid, r["date"])[0]
        mine[r["date"]] = r
    days = _per_day(db, mid, CTX_FROM, hi)
    cells, here = _cells(db, mid, CTX_FROM, hi)
    ivs = _intervals(db, cells, here, ref)

    auto = db.query(PPManagerSetting).filter_by(manager_id=mid).first()
    versions = [{
        "from": v.valid_from, "to": v.valid_to, "reason": v.reason, "created": v.created_at,
        "lines": len(v.lines or []),
        "covers_day": (v.valid_from is None or v.valid_from <= DAY) and DAY <= v.valid_to,
    } for v in db.query(PPCatalogVersion).filter_by(manager_id=mid)
        .order_by(PPCatalogVersion.created_at)]
    uploads = [{
        "date": r.date, "type": r.file_type, "file": r.filename, "rows": r.row_count,
        "at": r.uploaded_at, "unit": r.manager_id,
    } for r in db.query(PPUpload.date, PPUpload.file_type, PPUpload.filename,
                        PPUpload.row_count, PPUpload.uploaded_at, PPUpload.manager_id)
        .filter(PPUpload.date >= DAY - timedelta(days=2), PPUpload.date <= hi)
        .order_by(PPUpload.uploaded_at)]

    sap_rows = [{
        "sku": r.sap_code, "wc": r.work_center, "plan_qty": _f(r.plan_qty),
        "fact_qty": _f(r.actual_qty), "plan_ovr": _f(r.plan_override),
        "fact_ovr": _f(r.actual_override), "updated": r.updated_at,
    } for r in db.query(PPDaily).filter_by(manager_id=mid, date=DAY)
        .order_by(PPDaily.work_center, PPDaily.sap_code)]
    line_rows = [{
        "key": r.qty_key, "wc": r.work_center, "line": r.line_key,
        "plan_ovr": _f(r.plan_override), "fact_ovr": _f(r.actual_override),
        "cap": _f(r.container_cap), "updated": r.updated_at,
    } for r in db.query(PPLineDaily).filter_by(manager_id=mid, date=DAY)
        .order_by(PPLineDaily.work_center, PPLineDaily.qty_key)]
    pins = [{
        "wc": r.work_center, "grp": r.wc_group, "people": r.people, "shtatka": r.shtatka,
        "updated": r.updated_at,
    } for r in db.query(PPWorkCenterDaily).filter_by(manager_id=mid, date=DAY)
        .order_by(PPWorkCenterDaily.work_center, PPWorkCenterDaily.wc_group)]
    attendance = [{
        "worker": r.worker_name, "job": r.job_title, "clock": r.clock_in_out,
        "hours": _f(r.hours_worked), "early": _f(r.early_arrival_min),
        "cell": r.verifix_code, "sup": bool(r.is_supervisor), "w": r.hc_weight,
        "split": r.split_of,
    } for r in db.query(Attendance).filter_by(manager_id=mid, date=DAY)
        .order_by(Attendance.verifix_code, Attendance.worker_name)]

    live = []
    for e in db.query(LiveStaffEvent).filter(LiveStaffEvent.day == DAY, or_(
            LiveStaffEvent.from_manager_id == mid, LiveStaffEvent.to_manager_id == mid)):
        live.append({"what": f"exchange {e.kind}", "status": e.status,
                     "who": e.worker_name, "detail": f"{e.from_cell} → {e.to_cell}",
                     "by": e.created_by_name, "at": e.created_at,
                     "decided": e.decided_at, "decided_by": e.decided_by_name})
    for doc in db.query(LiveDocument).filter_by(manager_id=mid, day=DAY):
        live.append({"what": f"document {doc.doc_type}", "status": doc.status, "who": "",
                     "detail": f"#{doc.id}", "by": doc.created_by_name, "at": doc.created_at,
                     "decided": doc.approved_at or doc.updated_at,
                     "decided_by": doc.approved_by_name})
    for x in db.query(LiveDeletion).filter_by(manager_id=mid, day=DAY):
        live.append({"what": "deletion", "status": x.status, "who": x.worker_name,
                     "detail": x.initiated_by or "", "by": x.supervisor_name,
                     "at": x.created_at, "decided": x.processed_at,
                     "decided_by": x.processed_by_name})
    for x in db.query(LiveClockFix).filter_by(manager_id=mid, day=DAY):
        live.append({"what": f"clock fix {x.action}", "status": "", "who": x.worker_name,
                     "detail": x.out_time or "", "by": x.created_by, "at": x.created_at,
                     "decided": None, "decided_by": None})
    for x in db.query(LivePlacement).filter_by(manager_id=mid, day=DAY):
        live.append({"what": "placement", "status": "", "who": x.employee_id,
                     "detail": f"{x.verifix_code}" + (f" → {x.second_code} @{x.split_at}"
                                                      if x.second_code else ""),
                     "by": x.updated_by, "at": x.updated_at, "decided": None,
                     "decided_by": None})
    for x in db.query(EditRequest).filter_by(manager_id=mid, date=DAY):
        live.append({"what": "edit request", "status": x.status, "who": x.worker_name,
                     "detail": f"{x.original} → {x.changes}"[:300], "by": x.supervisor_name,
                     "at": x.created_at, "decided": x.processed_at,
                     "decided_by": x.processed_by_name})
    for x in live:
        x["after_close"] = _after(x["at"], ref) or _after(x["decided"], ref)
    live.sort(key=lambda x: _aware(x["at"]) or datetime.min.replace(tzinfo=timezone.utc))
    batch = db.query(AttendanceBatch).filter_by(date=DAY).first()

    return {
        "unit": u.name, "unit_id": mid, "shift": u.shift, "zagruzka_on": u.zagruzka_on,
        "auto_fill": True if auto is None else bool(auto.auto_fill),
        "close": close, "ref": ref, "mine": mine, "days": days,
        "cells": [{"id": c.id, "code": c.verifix_code, "now": c.manager_id,
                   "on_day": c.id in here.get(DAY, set())} for c in cells],
        "intervals": ivs, "links": _links(db, mid, cells), "versions": versions,
        "uploads": uploads, "sap_rows": sap_rows, "line_rows": line_rows, "pins": pins,
        "attendance": attendance, "live": live,
        "batch": ({"status": batch.status, "uploaded": batch.uploaded_at,
                   "by": batch.uploaded_by_name, "saved": batch.saved_at,
                   "saved_by": batch.saved_by_name} if batch else None),
        "jurnal": _jurnal(db, mid, cells), "grid": _grid(db, u.shift),
    }


# ── the answer, as text ───────────────────────────────────────────────────────

def _outliers(rep: dict) -> list[str]:
    """18.10 against the median of the unit's other days, input by input — the
    question is WHICH of the four moved, before WHEN."""
    mine = rep["mine"]
    day = mine.get(DAY)
    others = [r for d, r in mine.items() if d != DAY and r["plan_min"] and r["hc"]]
    if not day or not others:
        return ["not enough days with figures to compare"]

    def med(fn):
        vals = [fn(r) for r in others if fn(r) is not None]
        return statistics.median(vals) if vals else None

    probes = [
        ("ФАКТ ÷ ПЛАН (ВЫП)", lambda r: r["fact_min"] / r["plan_min"] if r["plan_min"] else None, 2),
        ("ПЛАН min per typed person", lambda r: r["plan_min"] / r["hc"] if r["hc"] else None, 0),
        ("ФАКТ min per typed person", lambda r: r["fact_min"] / r["hc"] if r["hc"] else None, 0),
        ("typed people", lambda r: r["hc"], 0),
        ("Verifix min per typed person",
         lambda r: r["verifix_min"] / r["hc"] if r["verifix_min"] and r["hc"] else None, 0),
        ("ojidaniya min per person", lambda r: r["ojid"], 1),
        ("early arrival min", lambda r: r["early"], 1),
    ]
    out = []
    for label, fn, nd in probes:
        v, m = fn(day), med(fn)
        if v is None or m is None:
            out.append(f"{label}: 18.10 {_g(v, nd)} · usual {_g(m, nd)}")
            continue
        ratio = (v / m) if m else None
        flag = "  ◀ OFF" if ratio is not None and (ratio < 0.75 or ratio > 1.33) else ""
        out.append(f"{label}: 18.10 {_g(v, nd)} · usual {_g(m, nd)}"
                   + (f" (×{ratio:.2f})" if ratio is not None else "") + flag)
    return out


def text(rep: dict) -> str:
    ref, close = rep["ref"], rep["close"]
    L = [f"{rep['unit'].upper()} · {DAY:%d.%m.%Y} — did the closed day change?", ""]
    L.append(f"UNIT #{rep['unit_id']} {rep['unit']} · shift {rep['shift']} · загрузка "
             f"{'on' if rep['zagruzka_on'] else 'OFF'} · SAP auto-fill "
             f"{'ON' if rep['auto_fill'] else 'OFF (manual)'}")
    L.append(f"DAY STATE NOW: {close['state']} · pending requests "
             f"{close['counts'].get('pending_requests', 0)} · draft docs "
             f"{close['counts'].get('draft_docs', 0)} · re-opened {close['reopens']}×")
    for at, what, who, src, outcome in close["events"]:
        L.append(f"  {_dt(at)} {what} — {who or '?'} ({src}{', ' + outcome if outcome else ''})")
    if not close["events"]:
        L.append("  no close recorded for this day")
    L.append(f"FIRST CLOSE (what the brigadir saw): {_dt(ref)}")
    L.append("")

    L.append("THE DAY NOW vs ITS NEIGHBOURS — To'liq P/A · Smena boshi/oxiri P/A · ×0.9 A")
    for d in sorted(rep["mine"]):
        r = rep["mine"][d]
        mark = (" ◀" if d == DAY else "") + ("" if r["state"] == "confirmed"
                                              else f" ({r['state']})")
        L.append(f"{_d(d)}{mark}: {_pct(r['full_p'])}/{_pct(r['full_a'])} · "
                 f"{_pct(r['simple_p'])}/{_pct(r['simple_a'])} · {_pct(r['full90_a'])} | "
                 f"plan {_g(r['plan_min'])} · fact {_g(r['fact_min'])} min · people "
                 f"{_g(r['hc'])} · Verifix {_g(r['verifix_min'])} min · ojid "
                 f"{_g(r['ojid'], 1)} · early {_g(r['early'], 1)}")
    L.append("Which input of 18.10 is out of line:")
    L.extend("  " + s for s in _outliers(rep))
    L.append("")

    L.append(f"WHAT WAS WRITTEN AFTER THE FIRST CLOSE ({_dt(ref)})")
    dd = rep["days"].get(DAY, {})
    found = []

    def say(label, at, extra=""):
        late = _after(at, ref)
        found.append(late)
        L.append(f"• {label}: last written {_dt(at)}"
                 + ("  ◀ AFTER the close" if late else "") + (f" · {extra}" if extra else ""))

    say("ПЛАН/ФАКТ quantities (pp_daily)", dd.get("pp_last"),
        f"{dd.get('pp_rows', 0)} rows · Σ plan {_g(dd.get('pp_plan'))} · Σ fact "
        f"{_g(dd.get('pp_fact'))} pcs · {dd.get('pp_ovr', 0)} typed overrides")
    say("per-line typed ПЛАН/ФАКТ (pp_line_daily)", dd.get("ln_last"),
        f"{dd.get('ln_rows', 0)} rows · {dd.get('ln_ovr', 0)} values")
    say("typed people «Bugungi fakt»", dd.get("pin_last"),
        f"Σ {_g(dd.get('pin_people'))} people")
    say("attendance copy (live_projections)", dd.get("copied_at"),
        f"{dd.get('copy_rows')} rows copied · now {dd.get('att_rows', 0)} rows, "
        f"{dd.get('att_workers', 0)} people, Σ {_g(dd.get('att_hours'), 1)} h")
    late_iv = [i for i in rep["intervals"] if i["date"] == DAY and i["after_close"]]
    found.append(bool(late_iv))
    n_day = sum(1 for i in rep["intervals"] if i["date"] == DAY and i["here"])
    L.append(f"• ojidaniya intervals on the unit's cells that day: {n_day}, "
             f"{len(late_iv)} created/edited/decided AFTER the close")
    for i in late_iv[:12]:
        L.append(f"    {i['cell']} {i['category']} {i['start']}–{i['end']} {i['status']}"
                 f"{'' if i['here'] else ' (cell not in unit that day)'} · created "
                 f"{_dt(i['created'])} · decided {_dt(i['decided'])} by {i['decided_by'] or '—'}")
    reach = [x for x in rep["links"] if x["reaches_day"] and _after(x["created"], ref)]
    found.append(bool(reach))
    L.append(f"• cell moves recorded after the close that reach 18.10: {len(reach)}")
    for x in reach[:8]:
        L.append(f"    {x['cell']}: {x['from']} → {x['to']} from {_d(x['effective_from'])} "
                 f"({x['src']}, {x['by'] or '—'}, recorded {_dt(x['created'])})")
    cat = [v for v in rep["versions"] if v["covers_day"]]
    late_cat = [v for v in cat if _after(v["created"], ref)]
    found.append(bool(late_cat))
    L.append(f"• catalog versions covering 18.10: {len(cat)}, frozen after the close: "
             f"{len(late_cat)}" + ("" if cat else " (the day reads TODAY's catalog)"))
    for v in late_cat[:5]:
        L.append(f"    {_d(v['from'])}–{_d(v['to'])} · {v['reason']} · {_dt(v['created'])}")
    late_live = [x for x in rep["live"] if x["after_close"]]
    found.append(bool(late_live))
    L.append(f"• live-day documents / exchanges / fixes / edit requests: {len(rep['live'])}, "
             f"{len(late_live)} filed or decided after the close")
    for x in late_live[:10]:
        L.append(f"    {x['what']} {x['status']} · {x['detail'][:60]} · at {_dt(x['at'])} · "
                 f"decided {_dt(x['decided'])} {x['decided_by'] or ''}")
    up = [x for x in rep["uploads"] if x["date"] == DAY]
    L.append("• SAP files stored for 18.10: "
             + (", ".join(f"{x['type']} {_dt(x['at'])}" + (" ◀ AFTER" if _after(x["at"], ref) else "")
                          for x in up) or "none"))
    b = rep["batch"]
    if b:
        L.append(f"• «Davomat» batch 18.10: {b['status']} · uploaded {_dt(b['uploaded'])} "
                 f"({b['by'] or '—'}) · saved {_dt(b['saved'])} ({b['saved_by'] or '—'})")
    if ref is None:
        L.append("  (no close time found, so nothing can be called «after»)")
    elif not any(found):
        L.append("  Nothing filed under this day was written after the close — look at the "
                 "inputs computed at read time (catalog, cells) and the deploys below.")
    L.append("")

    late_all = [j for j in rep["jurnal"] if _after(j["at"], ref) and j["outcome"] == "done"]
    late_j = [j for j in late_all if j["why"] != "mention"]
    L.append(f"JURNAL AFTER THE CLOSE naming this unit, this day or its cells: {len(late_j)}"
             + (f" (+{len(late_all) - len(late_j)} that only mention a cell code — workbook)"
                if len(late_all) > len(late_j) else "")
             + (f" (first {_JURNAL_MAX} read)" if len(rep["jurnal"]) >= _JURNAL_MAX else ""))
    for j in late_j[:30]:
        ch = ""
        if j["changes"]:
            ch = " · " + "; ".join(f"{c[0]} {c[1]}→{c[2]}" for c in j["changes"][:3]
                                   if isinstance(c, (list, tuple)) and len(c) >= 3)
        L.append(f"{_dt(j['at'])} {j['action']} — {j['who'] or j['source']}"
                 f" [{j['unit'] or '—'} · {_d(j['day'])}] v{j['version'] or '?'}{ch}"[:300])
    L.append("")

    L.append(f"SHIFT {rep['shift']} GRID {GRID_FROM:%d.%m}–{GRID_TO:%d.%m} — match the "
             "screenshot's rows (To'liq P/A · Smena boshi/oxiri P/A)")
    for unit in rep["grid"]:
        bits = []
        for d in sorted(unit["days"]):
            r = unit["days"][d]
            bits.append(f"{_d(d)} {_pct(r['full_p'])}/{_pct(r['full_a'])}·"
                        f"{_pct(r['simple_p'])}/{_pct(r['simple_a'])}"
                        + ("" if r["state"] == "confirmed" else f" ({r['state']})"))
        mark = " ◀" if unit["id"] == rep["unit_id"] else ""
        L.append(f"{unit['name']}{mark}: " + " | ".join(bits))
    L.append("")
    L.append("Every row behind these lines — the SAP rows, typed values, people pins, "
             "attendance, intervals, cell moves, catalog versions, uploads and the whole "
             "Jurnal — is in the attached workbook.")
    return "\n".join(L)


# ── the workbook ──────────────────────────────────────────────────────────────

def _cell(v):
    if isinstance(v, str):
        return ILLEGAL_CHARACTERS_RE.sub("", v)
    if isinstance(v, datetime):
        return _aware(v).astimezone(TZ).replace(tzinfo=None)
    if isinstance(v, (list, dict, tuple, set)):
        return ILLEGAL_CHARACTERS_RE.sub("", str(v))[:2000]
    return v


def _sheet(wb: Workbook, title: str, head: list[str], rows: list[list]) -> None:
    ws = wb.create_sheet(title)
    ws.append(head)
    for c in ws[1]:
        c.font = Font(bold=True)
    for r in rows:
        ws.append([_cell(v) for v in r])
    ws.freeze_panes = "A2"
    for i, h in enumerate(head, start=1):
        width = max([len(str(h))] + [len(str(_cell(r[i - 1]))) for r in rows[:300]
                                     if i - 1 < len(r) and r[i - 1] is not None])
        ws.column_dimensions[ws.cell(1, i).column_letter].width = min(max(width + 2, 9), 60)


def build_workbook(rep: dict) -> BytesIO:
    ref = rep["ref"]
    wb = Workbook()
    wb.remove(wb.active)
    _sheet(wb, "Kunlar", [
        "Sana", "Holat", "To'liq P", "To'liq A", "Smena P", "Smena A", "×0.9 A", "Plan, daq",
        "Fakt, daq", "Odam soni", "Verifix, daq", "Verifix odam", "Samarali odam",
        "Mavjud daq", "Ojidaniya", "Erta kelish", "SAP qatorlar", "SAP oxirgi yozuv",
        "Qator qiymatlari oxirgi", "Odam soni oxirgi", "Davomat qatorlar", "Davomat odam",
        "Davomat soat", "Nusxa vaqti"],
        [[d, r["state"], r["full_p"], r["full_a"], r["simple_p"], r["simple_a"], r["full90_a"],
          r["plan_min"], r["fact_min"], r["hc"], r["verifix_min"], r["verifix_hc"],
          r["effective_hc"], r["avail_min"], r["ojid"], r["early"],
          rep["days"].get(d, {}).get("pp_rows"), rep["days"].get(d, {}).get("pp_last"),
          rep["days"].get(d, {}).get("ln_last"), rep["days"].get(d, {}).get("pin_last"),
          rep["days"].get(d, {}).get("att_rows"), rep["days"].get(d, {}).get("att_workers"),
          rep["days"].get(d, {}).get("att_hours"), rep["days"].get(d, {}).get("copied_at")]
         for d, r in sorted(rep["mine"].items())])
    _sheet(wb, "Yopilish", ["Vaqt", "Nima", "Kim", "Manba", "Natija"],
           [list(e) for e in rep["close"]["events"]])
    _sheet(wb, "SAP qatorlar 18.10", [
        "SKU", "Ish markazi", "Plan (fayl)", "Fakt (fayl)", "Plan (qo'lda)",
        "Fakt (qo'lda)", "O'zgargan", "Yopilgandan keyin"],
        [[s["sku"], s["wc"], s["plan_qty"], s["fact_qty"], s["plan_ovr"], s["fact_ovr"],
          s["updated"], "ha" if _after(s["updated"], ref) else ""] for s in rep["sap_rows"]])
    _sheet(wb, "Qator qiymatlari 18.10", [
        "Kalit", "Ish markazi", "Qator", "Plan (qo'lda)", "Fakt (qo'lda)", "Kontiner krishka",
        "O'zgargan", "Yopilgandan keyin"],
        [[r["key"], r["wc"], r["line"], r["plan_ovr"], r["fact_ovr"], r["cap"], r["updated"],
          "ha" if _after(r["updated"], ref) else ""] for r in rep["line_rows"]])
    _sheet(wb, "Odam soni 18.10", ["Ish markazi", "Guruh", "Odam", "Shtat", "O'zgargan",
                                   "Yopilgandan keyin"],
           [[p["wc"], p["grp"], p["people"], p["shtatka"], p["updated"],
             "ha" if _after(p["updated"], ref) else ""] for p in rep["pins"]])
    _sheet(wb, "Davomat 18.10", ["Xodim", "Lavozim", "Kirish/chiqish", "Soat", "Erta, daq",
                                 "Yacheyka", "Brigadir", "Vazn", "Bo'lingan"],
           [[a["worker"], a["job"], a["clock"], a["hours"], a["early"], a["cell"],
             "ha" if a["sup"] else "", a["w"], a["split"]] for a in rep["attendance"]])
    _sheet(wb, "Jonli kun 18.10", ["Nima", "Holat", "Kim haqida", "Tafsilot", "Kiritgan",
                                   "Vaqt", "Qaror vaqti", "Qaror qilgan", "Yopilgandan keyin"],
           [[x["what"], x["status"], x["who"], x["detail"], x["by"], x["at"], x["decided"],
             x["decided_by"], "ha" if x["after_close"] else ""] for x in rep["live"]])
    _sheet(wb, "Ojidaniya", ["Sana", "Yacheyka", "Shu bo'limda", "Toifa", "Boshlanish",
                             "Tugash", "To'xtadi", "Tozalovchi", "Holat", "Izoh", "Kiritgan",
                             "Qaror qilgan", "Yaratilgan", "O'zgargan", "Qaror vaqti",
                             "Yopilgandan keyin"],
           [[i["date"], i["cell"], "ha" if i["here"] else "yo'q", i["category"], i["start"],
             i["end"], "ha" if i["stopped"] else "", i["cleaners"], i["status"], i["note"],
             i["by"], i["decided_by"], i["created"], i["updated"], i["decided"],
             "ha" if i["after_close"] and i["date"] == DAY else ""] for i in rep["intervals"]])
    _sheet(wb, "Yacheykalar", ["Yacheyka", "Hozirgi bo'lim", "18.10 shu bo'limda"],
           [[c["code"], c["now"], "ha" if c["on_day"] else "yo'q"] for c in rep["cells"]])
    _sheet(wb, "Yacheyka ko'chishlari", ["Yacheyka", "Kimdan", "Kimga", "Kuchga kiradi",
                                         "Kuchga kirish vaqti", "Qo'llangan", "Manba", "Kim",
                                         "Yozilgan", "18.10 ga yetadi"],
           [[x["cell"], x["from"], x["to"], x["effective_from"], x["effective_at"],
             x["applied"], x["src"], x["by"], x["created"],
             "ha" if x["reaches_day"] else ""] for x in rep["links"]])
    _sheet(wb, "Katalog versiyalari", ["Dan", "Gacha", "Sabab", "Yaratilgan", "Qatorlar",
                                       "18.10 ni qamraydi"],
           [[v["from"], v["to"], v["reason"], v["created"], v["lines"],
             "ha" if v["covers_day"] else ""] for v in rep["versions"]])
    _sheet(wb, "SAP fayllar", ["Sana", "Turi", "Fayl", "Qatorlar", "Yuklangan", "Bo'lim"],
           [[x["date"], x["type"], x["file"], x["rows"], x["at"], x["unit"]]
            for x in rep["uploads"]])
    grid = []
    for unit in rep["grid"]:
        for d, r in sorted(unit["days"].items()):
            grid.append([unit["name"], d, r["state"], r["full_p"], r["full_a"], r["simple_p"],
                         r["simple_a"], r["full90_a"], r["plan_min"], r["fact_min"], r["hc"],
                         r["verifix_min"], r["ojid"], r["early"]])
    _sheet(wb, f"{rep['shift']}-smena", [
        "Bo'lim", "Sana", "Holat", "To'liq P", "To'liq A", "Smena P", "Smena A", "×0.9 A",
        "Plan, daq", "Fakt, daq", "Odam soni", "Verifix, daq", "Ojidaniya", "Erta kelish"],
        grid)
    _sheet(wb, "Jurnal", ["Nega", "Vaqt", "Amal", "Natija", "Manba", "Kim", "Rol", "Bo'lim",
                          "Kun", "Nishon", "O'zgarishlar", "Tafsilotlar", "Yo'l", "Versiya",
                          "Yopilgandan keyin"],
           [[j["why"], j["at"], j["action"], j["outcome"], j["source"], j["who"], j["role"],
             j["unit"],
             j["day"], j["target"], j["changes"], j["details"], j["path"], j["version"],
             "ha" if _after(j["at"], ref) else ""] for j in rep["jurnal"]])
    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


# ── delivery ──────────────────────────────────────────────────────────────────

def _post(method: str, data: dict, files: dict | None = None) -> None:
    last = ""
    for attempt in range(1, SEND_RETRIES + 1):
        wait = 0
        try:
            r = requests.post(f"{_API}/bot{settings.telegram_bot_token}/{method}",
                              data=data, files=files, timeout=180)
            body = r.json()
            if body.get("ok"):
                return
            last = body.get("description") or f"HTTP {r.status_code}"
            wait = int(((body.get("parameters") or {}).get("retry_after")) or 0)
        except Exception as exc:
            last = type(exc).__name__
        if attempt < SEND_RETRIES:
            time.sleep(max(wait, 3 * attempt))
    raise RuntimeError(f"{method} failed: {last}"[:300])


def send(db: Session, chat_id: int, *_window) -> int:
    if not settings.telegram_bot_token:
        raise RuntimeError("telegram bot token not configured")
    rep = collect(db)
    body = text(rep)
    book = build_workbook(rep).getvalue()
    db.rollback()
    chunks, cur = [], ""
    for line in body.split("\n"):
        if len(cur) + len(line) + 1 > 3900:
            chunks.append(cur)
            cur = ""
        cur += line + "\n"
    chunks.append(cur)
    for c in chunks:
        _post("sendMessage", {"chat_id": chat_id, "text": c})
    stamp = datetime.now(TZ).strftime("%d.%m.%Y")
    _post("sendDocument", {"chat_id": chat_id,
                           "caption": f"{rep['unit']} · {DAY:%d.%m} — what changed after the close"},
          files={"document": (f"closed-day-{DAY:%d-%m}-{rep['unit'].replace(' ', '-')}-{stamp}.xlsx",
                              book,
                              "application/vnd.openxmlformats-officedocument."
                              "spreadsheetml.sheet")})
    return len(chunks) + 1
