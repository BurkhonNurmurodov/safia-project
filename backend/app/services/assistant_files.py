"""Files in a «Yordamchi» chat — what the user attaches, and what it makes.

IN: a photo, a screenshot, a PDF, an Excel or CSV sheet, a text file. Images
and PDFs go to Gemini as they are (images shrunk first, like the proof
reviewer's); a sheet is read HERE into plain rows, because Gemini reads tables
as text far better than it reads a zip of XML. A file it cannot read is refused
at the door rather than sent along to fail later.

OUT: a workbook it builds from data it gathered (``make_xlsx``), or a file an
existing export endpoint returned. In a browser the file is a download link in
the chat; in Telegram, where a web view cannot save a download, the bot DMs it
— the platform's standing rule for every export (``xlsx_delivery``).
"""
from __future__ import annotations

import base64
import csv
import io
import json
import logging
import re
from typing import Any, Optional

from sqlalchemy.orm import Session

from app.models import AssistantFile

log = logging.getLogger(__name__)

MAX_UPLOAD = 20 * 1024 * 1024
MAX_PER_MESSAGE = 5
_PDF_INLINE_MAX = 12 * 1024 * 1024
_SHEET_CHARS = 150_000

XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def sniff(name: str, data: bytes, declared: str = "") -> Optional[str]:
    """The file's real kind by its BYTES (a name can claim anything):
    image/* · application/pdf · xlsx · text/csv · text/plain — or None."""
    head = data[:16]
    if head.startswith(b"\xff\xd8\xff"):
        return "image/jpeg"
    if head.startswith(b"\x89PNG\r\n\x1a\n"):
        return "image/png"
    if head[:4] == b"RIFF" and data[8:12] == b"WEBP":
        return "image/webp"
    if head[:6] in (b"GIF87a", b"GIF89a"):
        return "image/gif"
    if head.startswith(b"%PDF-"):
        return "application/pdf"
    if head.startswith(b"PK\x03\x04"):
        if b"xl/" in data[:4096] or (name or "").lower().endswith(".xlsx"):
            return XLSX
        return None
    try:
        data[:200_000].decode("utf-8")
    except UnicodeDecodeError:
        return None
    low = (name or "").lower()
    if low.endswith((".csv", ".tsv")):
        return "text/csv"
    if low.endswith((".txt", ".md", ".json", ".log")) or (declared or "").startswith("text/"):
        return "text/plain"
    return "text/plain"


def store(db: Session, *, owner_key: str, opened_by: Optional[str], thread_id: Optional[int],
          kind: str, name: str, mime: str, data: bytes) -> AssistantFile:
    f = AssistantFile(owner_key=owner_key, opened_by=opened_by, thread_id=thread_id,
                      kind=kind, name=(name or "file")[:200], mime=mime, size=len(data),
                      data=data)
    db.add(f)
    db.flush()
    return f


def meta(f: AssistantFile) -> dict:
    return {"id": f.id, "name": f.name, "mime": f.mime, "size": f.size, "kind": f.kind}


def sheet_text(data: bytes) -> str:
    """An .xlsx as tab-separated rows, sheet by sheet, values not formulas."""
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    out: list[str] = []
    used = 0
    for ws in wb.worksheets:
        out.append(f"## Sheet «{ws.title}»")
        rows = 0
        for row in ws.iter_rows(values_only=True):
            cells = ["" if v is None else str(v).replace("\t", " ").replace("\n", " ")
                     for v in row[:60]]
            while cells and cells[-1] == "":
                cells.pop()
            if not cells:
                continue
            line = "\t".join(cells)
            used += len(line) + 1
            if used > _SHEET_CHARS:
                out.append("… (the rest of the workbook is cut off)")
                return "\n".join(out)
            out.append(line)
            rows += 1
            if rows >= 3000:
                out.append("… (more rows in this sheet are cut off)")
                break
    return "\n".join(out)


def to_parts(f: AssistantFile) -> list[dict]:
    """The Gemini parts that show the model this file."""
    label = f"[Attachment #{f.id} «{f.name}», {f.mime}, {f.size} bytes]"
    try:
        if f.mime.startswith("image/"):
            from app.services.gemini import shrink_image
            data, mime = shrink_image(bytes(f.data), f.mime)
            return [{"text": label},
                    {"inline_data": {"mime_type": mime, "data": base64.b64encode(data).decode()}}]
        if f.mime == "application/pdf":
            if f.size > _PDF_INLINE_MAX:
                return [{"text": label + " (too large to read — ask for a smaller PDF)"}]
            return [{"text": label},
                    {"inline_data": {"mime_type": "application/pdf",
                                     "data": base64.b64encode(bytes(f.data)).decode()}}]
        if f.mime == XLSX:
            return [{"text": label + "\n" + sheet_text(bytes(f.data))}]
        text = bytes(f.data).decode("utf-8", errors="replace")
        if len(text) > _SHEET_CHARS:
            text = text[:_SHEET_CHARS] + "\n… (cut off)"
        return [{"text": label + "\n" + text}]
    except Exception as exc:
        log.warning("assistant: could not read attachment %s: %s", f.id, exc)
        return [{"text": label + " (this file could not be read)"}]


# ── a workbook it builds ─────────────────────────────────────────────────────

def _safe_sheet(name: str, used: set) -> str:
    n = re.sub(r"[\[\]\*\?/\\:]", " ", (name or "Sheet").strip())[:31] or "Sheet"
    base, i = n, 2
    while n.lower() in used:
        n = f"{base[:28]} {i}"
        i += 1
    used.add(n.lower())
    return n


def _num(v: Any) -> Any:
    """Keep numbers numeric (a spreadsheet must be able to sum them)."""
    if isinstance(v, (int, float)) or v is None:
        return v
    s = str(v).strip()
    if re.fullmatch(r"-?\d+", s) and not (len(s) > 1 and s.startswith("0")):
        try:
            return int(s)
        except ValueError:
            return s
    if re.fullmatch(r"-?\d+[.,]\d+", s):
        try:
            return float(s.replace(",", "."))
        except ValueError:
            return s
    return s


def make_xlsx(title: str, sheets: list[dict]) -> bytes:
    """A plain, readable workbook: a header row in bold on a soft gold band,
    frozen, filterable, columns sized to what they hold."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    wb.remove(wb.active)
    used: set = set()
    head_fill = PatternFill("solid", fgColor="F3E7CC")
    for sh in sheets or []:
        ws = wb.create_sheet(_safe_sheet(sh.get("name") or title, used))
        cols = [str(c) for c in (sh.get("columns") or [])]
        rows = sh.get("rows") or []
        if cols:
            ws.append(cols)
            for i in range(1, len(cols) + 1):
                c = ws.cell(row=1, column=i)
                c.font = Font(bold=True)
                c.fill = head_fill
                c.alignment = Alignment(vertical="center", wrap_text=True)
        for r in rows:
            ws.append([_num(v) for v in (r if isinstance(r, (list, tuple)) else [r])])
        width = max(len(cols), max((len(r) for r in rows if isinstance(r, (list, tuple))), default=0))
        for i in range(1, width + 1):
            longest = max((len(str(ws.cell(row=j, column=i).value or ""))
                           for j in range(1, min(ws.max_row, 400) + 1)), default=8)
            ws.column_dimensions[get_column_letter(i)].width = min(max(longest + 2, 8), 50)
        if cols:
            ws.freeze_panes = "A2"
            if ws.max_row > 1:
                ws.auto_filter.ref = f"A1:{get_column_letter(len(cols))}{ws.max_row}"
    if not wb.worksheets:
        wb.create_sheet("Sheet")
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def parse_rows(raw: Any) -> list:
    """Rows the model sent as a JSON string (or already a list)."""
    if isinstance(raw, list):
        return raw
    if not raw:
        return []
    try:
        val = json.loads(raw)
    except (TypeError, ValueError):
        reader = csv.reader(io.StringIO(str(raw)))
        return [r for r in reader]
    return val if isinstance(val, list) else []
