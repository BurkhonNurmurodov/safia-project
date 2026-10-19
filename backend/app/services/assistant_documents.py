"""Documents «Yordamchi» builds — ANY kind, from one block model (2026-10-19).

The operator: «the AI assistant should be able to create any kind of
documents when it's asked even when we don't have a template for it». Until
then it could hand over an Excel workbook (a plain sheet, gone since 2026-10-19) or
a file an export endpoint already produced, and it refused the «Samaradorlik»
table because the browser computes it and nothing exports it.

Now the model composes a document as BLOCKS — heading, paragraph, bullets,
numbered, table, facts, quote, page_break — and names a format; this module
lays them out. One model, eight outputs, one look:

* **docx** (python-docx) — a report somebody edits further.
* **pdf** (fpdf2, a Unicode TTF off the server — the camera stamp's font
  resolver) — a report somebody prints or forwards as it is.
* **pptx** (python-pptx) — every level-1 heading opens a slide; a long table
  continues onto the next with «(davomi)».
* **xlsx** — every table its own sheet (the house look — title, scope line, gold header, zebra, frozen panes, autofilter; a
  «grid» block is the /zagruzka «Excel» template itself, laid out by ``grid_export.lay_out``), the text on a first
  sheet; **csv** — the tables, comma-separated, UTF-8 with a BOM so Excel
  opens Cyrillic right.
* **md**, **txt**, **html** — the same blocks as text a chat, a wiki or a
  browser takes.

The look is the platform's: the title, a muted line under it (period · scope
· the day), headings in ink, a table whose header row is brand gold with
white text, zebra rows, numbers right-aligned, and a quiet footer naming
Safia IMS and the day. Nothing decorative. A table's rows may come as a
``rows_ref`` — a `compute` result or an API answer kept for the run — so a
figure never passes through the model's hands on its way into the file.
"""
from __future__ import annotations

import csv
import html
import io
import json
import logging
import math
import re
from datetime import datetime
from functools import lru_cache
from typing import Any, Callable, Optional
from zoneinfo import ZoneInfo

log = logging.getLogger(__name__)

TZ = ZoneInfo("Asia/Tashkent")

DOCX = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
PPTX = "application/vnd.openxmlformats-officedocument.presentationml.presentation"
XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

FORMATS: dict[str, tuple[str, str]] = {        # format → (extension, mime)
    "docx": ("docx", DOCX),
    "pdf": ("pdf", "application/pdf"),
    "pptx": ("pptx", PPTX),
    "xlsx": ("xlsx", XLSX),
    "csv": ("csv", "text/csv"),
    "md": ("md", "text/markdown"),
    "txt": ("txt", "text/plain"),
    "html": ("html", "text/html"),
}
_ALIASES = {"word": "docx", "doc": "docx", "excel": "xlsx", "powerpoint": "pptx",
            "ppt": "pptx", "markdown": "md", "text": "txt", "htm": "html"}

BRAND = "C8973F"
INK = "0F1117"
MUTED = "6B7280"
ZEBRA = "F6F4EF"
RULE = "DDDAD2"
QUOTE_BG = "FAF3E5"

MAX_BLOCKS = 400
MAX_ROWS = 20_000
MAX_COLS = 80
MAX_ITEMS = 500
MAX_TEXT = 20_000

BLOCK_TYPES = ("heading", "paragraph", "bullets", "numbered", "table", "grid", "facts", "quote",
               "page_break")
MAX_GRID_COLS = 200
MAX_GRID_TEXT = 60


class DocError(Exception):
    """A document that cannot be built, in words the model can act on."""


# ── the block model ──────────────────────────────────────────────────────────

def norm_format(fmt: Any) -> str:
    f = str(fmt or "").strip().lower().lstrip(".")
    f = _ALIASES.get(f, f)
    if f not in FORMATS:
        raise DocError(f"unknown format «{fmt}» — one of: {', '.join(FORMATS)}")
    return f


def _text(v: Any, cap: int = MAX_TEXT) -> str:
    if v is None:
        return ""
    s = v if isinstance(v, str) else str(v)
    return s.strip()[:cap]


def _rows_from_ref(obj: Any, columns: list[str]) -> tuple[list[str], list[list]]:
    """A kept ref as a table: a list of lists, a list of dicts (the columns
    are the keys, or the ones asked for), or a {columns, rows} result."""
    if isinstance(obj, dict):
        if isinstance(obj.get("rows"), list):
            cols = [str(c) for c in (obj.get("columns") or columns or [])]
            inner = obj["rows"]
            if inner and isinstance(inner[0], dict):
                cols = cols or [str(k) for k in inner[0].keys()]
                return cols, [[r.get(c) for c in cols] for r in inner if isinstance(r, dict)]
            return cols, [list(r) if isinstance(r, (list, tuple)) else [r] for r in inner]
        lists = [v for v in obj.values() if isinstance(v, list)]
        if len(lists) == 1:
            return _rows_from_ref(lists[0], columns)
        raise DocError("the ref is an object with no single list in it — compute the table "
                       "first and set result = {\"columns\": [...], \"rows\": [[...]]}")
    if isinstance(obj, list):
        if not obj:
            return columns, []
        if isinstance(obj[0], dict):
            cols = columns or [str(k) for k in obj[0].keys()]
            return cols, [[r.get(c) for c in cols] for r in obj if isinstance(r, dict)]
        return columns, [list(r) if isinstance(r, (list, tuple)) else [r] for r in obj]
    raise DocError("the ref is not a table (not a list) — compute one first")


def _norm_table_rows(b: dict, i: int, resolve_ref, what: str = "table",
                     max_cols: int = MAX_COLS) -> tuple[list[str], list[list], int]:
    """A table's or a grid's columns and rows, checked and padded: rows_json or
    a list, or rows_ref resolved through the caller. → (columns, rows, how
    many rows the cap dropped)."""
    cols = [str(c) for c in (b.get("columns") or [])][:max_cols]
    rows = b.get("rows")
    if isinstance(rows, str):
        try:
            rows = json.loads(rows)
        except ValueError as exc:
            raise DocError(f"block #{i} ({what}): rows is not valid JSON: {exc}") from exc
    ref = b.get("rows_ref")
    if ref:
        if resolve_ref is None:
            raise DocError(f"block #{i} ({what}): rows_ref cannot be used here")
        found, obj = resolve_ref(str(ref))
        if not found:
            raise DocError(f"block #{i} ({what}): ref «{ref}» is not available — "
                           f"compute it again or fetch it again")
        cols, rows = _rows_from_ref(obj, cols)
        cols = cols[:max_cols]
    if rows is None:
        rows = []
    if not isinstance(rows, list):
        raise DocError(f"block #{i} ({what}): rows must be a list of rows")
    clean = []
    for r in rows[:MAX_ROWS]:
        if isinstance(r, dict):
            r = [r.get(c) for c in cols] if cols else list(r.values())
        elif not isinstance(r, (list, tuple)):
            r = [r]
        clean.append(list(r)[:max_cols])
    if not cols and clean:
        cols = [f"{j + 1}" for j in range(max(len(r) for r in clean))]
    if not cols:
        raise DocError(f"block #{i} ({what}): no columns and no rows")
    width = len(cols)
    clean = [r + [None] * (width - len(r)) if len(r) < width else r[:width] for r in clean]
    return cols, clean, max(0, len(rows) - MAX_ROWS)


def normalise(raw: Any, resolve_ref: Optional[Callable[[str], tuple[bool, Any]]] = None) -> list[dict]:
    """The model's blocks (a JSON string or a list) checked and filled in:
    bad types named, tables resolved from their refs, caps applied."""
    if isinstance(raw, str):
        try:
            raw = json.loads(raw) if raw.strip() else []
        except ValueError as exc:
            raise DocError(f"blocks_json is not valid JSON: {exc}") from exc
    if isinstance(raw, dict):
        raw = raw.get("blocks") if isinstance(raw.get("blocks"), list) else [raw]
    if not isinstance(raw, list):
        raise DocError("blocks_json must be a JSON array of blocks")
    if len(raw) > MAX_BLOCKS:
        raise DocError(f"too many blocks ({len(raw)}; at most {MAX_BLOCKS})")
    out: list[dict] = []
    for i, b in enumerate(raw, 1):
        if isinstance(b, str):
            b = {"type": "paragraph", "text": b}
        if not isinstance(b, dict):
            raise DocError(f"block #{i} is not an object")
        typ = str(b.get("type") or "paragraph").strip().lower()
        typ = {"h1": "heading", "h2": "heading", "h3": "heading", "text": "paragraph", "p": "paragraph",
               "list": "bullets", "ul": "bullets", "ol": "numbered", "steps": "numbered",
               "kv": "facts", "note": "quote", "callout": "quote", "break": "page_break",
               "pagebreak": "page_break", "heatmap": "grid", "zagruzka_grid": "grid",
               "zagruzka": "grid", "bands": "grid", "pct_grid": "grid"}.get(typ, typ)
        if typ not in BLOCK_TYPES:
            raise DocError(f"block #{i}: unknown type «{b.get('type')}» — one of: {', '.join(BLOCK_TYPES)}")
        if typ == "heading":
            lvl = b.get("level")
            if lvl is None and str(b.get("type", "")).lower() in ("h1", "h2", "h3"):
                lvl = int(str(b["type"])[1])
            try:
                lvl = int(lvl or 1)
            except (TypeError, ValueError):
                lvl = 1
            out.append({"type": "heading", "text": _text(b.get("text"), 300), "level": min(max(lvl, 1), 3)})
        elif typ in ("paragraph", "quote"):
            out.append({"type": typ, "text": _text(b.get("text"))})
        elif typ in ("bullets", "numbered"):
            items = b.get("items")
            if isinstance(items, str):
                items = [x for x in items.splitlines() if x.strip()]
            if not isinstance(items, list):
                raise DocError(f"block #{i} ({typ}): items must be a list of strings")
            out.append({"type": typ, "items": [_text(x, 2000) for x in items[:MAX_ITEMS]]})
        elif typ == "facts":
            items = b.get("items")
            if isinstance(items, dict):
                items = list(items.items())
            if not isinstance(items, list):
                raise DocError(f"block #{i} (facts): items must be a list of [label, value] pairs")
            pairs = []
            for it in items[:MAX_ITEMS]:
                if isinstance(it, dict):
                    pairs.append([_text(it.get("label") or it.get("name"), 200), _text(it.get("value"), 500)])
                elif isinstance(it, (list, tuple)) and len(it) >= 2:
                    pairs.append([_text(it[0], 200), _text(it[1], 500)])
                else:
                    raise DocError(f"block #{i} (facts): every item is a [label, value] pair")
            out.append({"type": "facts", "items": pairs})
        elif typ == "table":
            cols, clean, truncated = _norm_table_rows(b, i, resolve_ref)
            out.append({"type": "table", "columns": cols, "rows": clean,
                        "caption": _text(b.get("caption"), 200), "truncated": truncated})
        elif typ == "grid":
            out.append(_norm_grid(b, i, resolve_ref))
        else:
            out.append({"type": "page_break"})
    if not out:
        raise DocError("the document has no content — give it at least one block")
    return out


# ── shared helpers ───────────────────────────────────────────────────────────

_NUM = re.compile(r"^[-+−]?\d[\d\s\u00a0]*([.,]\d+)?\s*%?$")
_BOLD = re.compile(r"\*\*(.+?)\*\*")


def fmt_cell(v: Any, *, machine: bool = False) -> str:
    """A cell as text. A float is rounded to two decimals; thousands get the
    platform's space; None reads «—». `machine` (CSV) keeps numbers raw."""
    if v is None:
        return "" if machine else "—"
    if isinstance(v, bool):
        return "true" if machine else ("✓" if v else "—")
    if isinstance(v, int):
        return str(v) if machine else _thousands(str(v))
    if isinstance(v, float):
        if math.isnan(v) or math.isinf(v):
            return "" if machine else "—"
        if machine:
            return repr(v) if v != int(v) else str(int(v))
        s = f"{v:.2f}".rstrip("0").rstrip(".")
        return _thousands(s)
    if isinstance(v, (dict, list)):
        return json.dumps(v, ensure_ascii=False)
    return str(v)


def _thousands(s: str) -> str:
    sign = "-" if s.startswith("-") else ""
    s = s.lstrip("-")
    whole, _, frac = s.partition(".")
    if len(whole) >= 5:
        groups = []
        while whole:
            groups.insert(0, whole[-3:])
            whole = whole[:-3]
        whole = " ".join(groups)
    return sign + whole + ("." + frac if frac else "")


def is_num(s: str) -> bool:
    s = (s or "").strip()
    return bool(s) and s != "—" and bool(_NUM.match(s))


def numeric_columns(rows: list[list[str]], ncols: int) -> list[bool]:
    """A column is numeric when most of its filled cells are numbers."""
    out = []
    for j in range(ncols):
        vals = [r[j] for r in rows if j < len(r) and r[j] not in ("", "—")]
        out.append(bool(vals) and sum(1 for v in vals if is_num(v)) >= 0.6 * len(vals))
    return out


def runs(text: str) -> list[tuple[str, bool]]:
    """«a **b** c» → [(a, False), (b, True), (c, False)]."""
    out, pos = [], 0
    for m in _BOLD.finditer(text or ""):
        if m.start() > pos:
            out.append((text[pos:m.start()], False))
        out.append((m.group(1), True))
        pos = m.end()
    if pos < len(text or ""):
        out.append((text[pos:], False))
    return out or [("", False)]


def plain(text: str) -> str:
    return _BOLD.sub(r"\1", text or "")


def stamp() -> str:
    return datetime.now(TZ).strftime("%d.%m.%Y")


def footer_text() -> str:
    return f"Safia IMS · {stamp()}"


def table_text(block: dict, *, machine: bool = False) -> list[list[str]]:
    return [[fmt_cell(v, machine=machine) for v in r] for r in block["rows"]]


def _table_width(weights: list[float], per_char: float, full: float) -> float:
    """What the table's columns need at this font — never more than the page,
    so a two-column table is not stretched across it."""
    return min(float(full), sum(weights) * per_char)


def _continued(name: str) -> str:
    base = name[: -len(" (davomi)")] if name.endswith(" (davomi)") else name
    return base + " (davomi)"


def _col_weights(cols: list[str], rows: list[list[str]], sample: int = 200) -> list[float]:
    """Relative column widths from what the cells hold (longest line, capped),
    so a code column stays narrow beside a note column."""
    w = []
    for j, c in enumerate(cols):
        longest = max([len(c)] + [len(r[j]) for r in rows[:sample] if j < len(r)])
        # +3 is the cell's own padding, so a 4-digit code never wraps to «000 / 1»
        w.append(float(min(max(longest, 5), 40) + 3))
    return w


# ── docx ─────────────────────────────────────────────────────────────────────


# ── The /zagruzka grid — the platform's template for a brigadir × day table ──
# From 2026-10-19 (the operator: «If there's template for it, use that
# template» — the assistant had handed over the «Samaradorlik» month as a bare
# sheet with filter arrows) a table of whole percents per unit per day —
# загрузка, samaradorlik, plan fulfilment, the comparison tables' P / A or D,
# /workers' davomat — is never a plain sheet. It is laid out through
# `grid_export.lay_out`, the very module the «Excel» button on /zagruzka draws
# its file with: the title, the scope line, the «generated» line, the gold
# header, every cell FILLED with its band colour under a percent format, «—»
# blanks, hatched pending days, the summary column, the legend, frozen panes,
# landscape. The bands are the admin's own edges for the table named
# (`routers.settings.ZAGRUZKA_BANDS`, read by the assistant and attached here
# as `segments`), so the file is coloured exactly as the screen is. In a
# format that is not a workbook (docx, pdf, pptx, csv, md, txt) the grid is a
# table of percents (`grid_as_table`); the HTML keeps the colours.

GRID_WORDS = {
    "uz": {"generated": "Tayyorlandi", "colors": "Ranglar",
           "pending": "Ko'rsatkich hali chiqmagan kunlar",
           "pending_title": "Tasdiqlanishi kutilmoqda", "pending_hc": "Odam soni kutilmoqda",
           "pending_labor": "Trudoyomkost yo'q",
           "stat": {"avg": "O'RT", "min": "MIN", "max": "MAKS"},
           "below": "D < {v}", "range": "{a} dan {b} gacha", "above": "> {v}"},
    "uz_cyrl": {"generated": "Тайёрланди", "colors": "Ранглар",
                "pending": "Кўрсаткич ҳали чиқмаган кунлар",
                "pending_title": "Тасдиқланиши кутилмоқда", "pending_hc": "Одам сони кутилмоқда",
                "pending_labor": "Трудоёмкость йўқ",
                "stat": {"avg": "ЎРТ", "min": "МИН", "max": "МАКС"},
                "below": "D < {v}", "range": "{a} дан {b} гача", "above": "> {v}"},
    "ru": {"generated": "Сформирован", "colors": "Цвета",
           "pending": "Дни, где показатель ещё не показан",
           "pending_title": "Ожидает подтверждения", "pending_hc": "Ожидается «Одам сони»",
           "pending_labor": "Нет трудоёмкости",
           "stat": {"avg": "СРД", "min": "МИН", "max": "МАКС"},
           "below": "D < {v}", "range": "от {a} до {b}", "above": "> {v}"},
    "en": {"generated": "Generated", "colors": "Colours",
           "pending": "Days whose figure is not shown yet",
           "pending_title": "Awaiting confirmation", "pending_hc": "Waiting for headcount",
           "pending_labor": "No trudoyomkost",
           "stat": {"avg": "AVG", "min": "MIN", "max": "MAX"},
           "below": "D < {v}", "range": "{a} to {b}", "above": "> {v}"},
}
# The screen's pending icons (components/ui/PendingInfoModal.jsx) and the
# legend line each one gets.
PENDING = {"⏳": "pending_title", "👥": "pending_hc", "📄": "pending_labor"}
PENDING_REASONS = {"not_closed": "⏳", "requests": "⏳", "no_headcount": "👥", "no_labor": "📄"}

# `bands` names the /zagruzka table whose colour edges the grid wears — the
# registry's own names (`ZAGRUZKA_BANDS`), plus `att` (/workers' davomat
# scale, fixed) and `default` (the platform's default heatmap / diff edges).
# The aliases are the words a question is asked with.
BAND_NAMES = ("load", "fulfil", "eff", "full", "full90", "simple", "att", "default")
BAND_ALIASES = {
    "zagruzka": "load", "загрузка": "load", "heatmap": "load", "nagruzka": "load",
    "yuklanish": "load", "yuk": "load", "karta": "load", "map": "load",
    "fulfilment": "fulfil", "fulfillment": "fulfil", "plan": "fulfil", "reja": "fulfil",
    "bajarilish": "fulfil", "reja_bajarilishi": "fulfil", "plan_pct": "fulfil", "vyp": "fulfil",
    "efficiency": "eff", "samaradorlik": "eff", "effektivnost": "eff", "эффективность": "eff",
    "comparison": "full", "toliq": "full", "to'liq": "full", "pa": "full", "full_table": "full",
    "verifix90": "full90", "full_90": "full90", "verifix": "full90",
    "simplified": "simple", "smena": "simple", "shift": "simple", "simple_pa": "simple",
    "soddalashtirilgan": "simple",
    "attendance": "att", "davomat": "att", "kelish": "att", "workers": "att",
    "": "default", "none": "default", "pct": "default", "diff": "default", "heat": "default",
}
# /workers' «Davomat issiqlik xaritasi» scale — the twin of `ATT_SEGMENTS` in
# frontend/src/utils/statusBands.js; change one and change the other.
ATT_SEGMENTS = [
    {"from": 0, "color": "#b91c1c"}, {"from": 60, "color": "#ef4444"},
    {"from": 75, "color": "#eab308"}, {"from": 85, "color": "#84cc16"},
    {"from": 90, "color": "#22c55e"}, {"from": 95, "color": "#15803d"},
]
_HEX6 = re.compile(r"^#?([0-9a-fA-F]{6}|[0-9a-fA-F]{3})$")
_GRID_NUM = re.compile(r"^[-+−]?\d[\d\s ]*([.,]\d+)?\s*%?$")


def _grid_words(lang: Optional[str]) -> dict:
    return GRID_WORDS.get(str(lang or "").strip()) or GRID_WORDS["uz"]


def _grid_value(v: Any) -> Any:
    """One grid cell as the template reads it: a number (the figure a whole
    percent is rounded from), None (no figure — the «—»), a pending day
    ({"pending": icon}) or a short text."""
    if v is None or isinstance(v, bool):
        return None
    if isinstance(v, (int, float)):
        if isinstance(v, float) and (math.isnan(v) or math.isinf(v)):
            return None
        return float(v)
    if isinstance(v, dict):
        p = v.get("pending") or v.get("hatch") or v.get("reason")
        if p:
            key = str(p).strip().lower()
            icon = PENDING_REASONS.get(key) or (str(p)[:1] if str(p)[:1] in PENDING else "⏳")
            return {"pending": icon}
        return _grid_value(v.get("v", v.get("value")))
    s = str(v).strip()
    if not s or s in ("—", "-", "–", "n/a", "N/A", "null", "None"):
        return None
    if s[:1] in PENDING:
        return {"pending": s[:1]}
    if s.lower() in PENDING_REASONS:
        return {"pending": PENDING_REASONS[s.lower()]}
    if _GRID_NUM.match(s):
        return float(s.replace("−", "-").replace("%", "").replace(",", ".")
                     .replace(" ", "").replace(" ", ""))
    return {"text": s[:MAX_GRID_TEXT]}


def _custom_scale(items: list, i: int) -> list[dict]:
    """A scale the model wrote itself: [{from, color}], colours as hex."""
    out = []
    for it in items[:16]:
        if not isinstance(it, dict):
            raise DocError(f"block #{i} (grid): every band is {{from, color}}")
        m = _HEX6.match(str(it.get("color") or ""))
        try:
            frm = float(it.get("from"))
        except (TypeError, ValueError):
            raise DocError(f"block #{i} (grid): a band's from must be a number") from None
        if not m:
            raise DocError(f"block #{i} (grid): a band's color must be a hex colour (#22c55e)")
        h = m.group(1)
        if len(h) == 3:
            h = "".join(ch * 2 for ch in h)
        out.append({"from": frm, "color": "#" + h.lower()})
    if not out:
        raise DocError(f"block #{i} (grid): bands is an empty list")
    return _ordered(out)


def _ordered(segs: list) -> list[dict]:
    """The scale sorted by its edges, numbers as numbers — `orderedSegments`."""
    out = []
    for sg in segs or []:
        try:
            out.append({"from": float(sg.get("from")), "color": str(sg.get("color") or "#9ca3af")})
        except (TypeError, ValueError, AttributeError):
            continue
    return sorted(out, key=lambda x: x["from"])


def _norm_grid(b: dict, i: int, resolve_ref) -> dict:
    cols, rows, truncated = _norm_table_rows(b, i, resolve_ref, "grid", MAX_GRID_COLS)
    if len(cols) < 2:
        raise DocError(f"block #{i} (grid): the first column is the row's name and every other "
                       f"column a day (or a period) — give at least two columns")
    kind = str(b.get("kind") or "pct").strip().lower()
    kind = {"percent": "pct", "%": "pct", "p": "pct", "a": "pct", "d": "diff", "difference": "diff",
            "delta": "diff", "farq": "diff"}.get(kind, kind)
    if kind not in ("pct", "diff"):
        raise DocError(f"block #{i} (grid): kind is pct (a percentage) or diff (D = P − A)")
    stat = str(b.get("stat") if b.get("stat") is not None else "avg").strip().lower()
    stat = {"average": "avg", "mean": "avg", "o'rt": "avg", "": "avg", "no": "none", "false": "none",
            "off": "none", "maks": "max", "minimum": "min", "maximum": "max"}.get(stat, stat)
    if stat not in ("avg", "min", "max", "none"):
        raise DocError(f"block #{i} (grid): stat is avg, min, max or none")
    bands = b.get("bands")
    segments = None
    if isinstance(bands, list):
        segments = _custom_scale(bands, i)
        bands = "custom"
    else:
        key = str(bands or "default").strip().lower().replace("-", "_").replace(" ", "_")
        name = BAND_ALIASES.get(key, key)
        if name not in BAND_NAMES:
            raise DocError(f"block #{i} (grid): bands «{bands}» names no /zagruzka table — one of: "
                           f"load (the загрузка heatmap), fulfil (reja bajarilishi), eff "
                           f"(samaradorlik), full / full90 / simple (the comparison tables: kind "
                           f"pct for P or A, diff for D), att (davomat), default")
        bands = name
    body = [[("" if r[0] is None else str(r[0])[:200]), *[_grid_value(v) for v in r[1:]]] for r in rows]
    foot: list[list] = []
    if b.get("foot") is not None or b.get("foot_ref"):
        _c, frows, _t = _norm_table_rows({"columns": cols, "rows": b.get("foot"), "rows_ref": b.get("foot_ref")},
                                         i, resolve_ref, "grid foot", MAX_GRID_COLS)
        foot = [[("" if r[0] is None else str(r[0])[:200]), *[_grid_value(v) for v in r[1:]]] for r in frows[:12]]
    return {"type": "grid", "columns": cols, "rows": body, "foot": foot, "kind": kind, "stat": stat,
            "bands": bands, "segments": segments, "caption": _text(b.get("caption"), 200),
            "truncated": truncated}


def band_requests(blocks: list[dict]) -> set[tuple[str, str]]:
    """The (table, kind) scales the grids still need — read by the assistant
    from the admin's settings and attached with `attach_scales`."""
    return {(b["bands"], b["kind"]) for b in blocks
            if b["type"] == "grid" and b.get("segments") is None}


def attach_scales(blocks: list[dict], scales: dict) -> None:
    for b in blocks:
        if b["type"] != "grid" or b.get("segments") is not None:
            continue
        segs = scales.get((b["bands"], b["kind"]))
        if not segs:
            raise DocError(f"no colour scale for the «{b['bands']}» grid")
        b["segments"] = _ordered(segs)


def _whole(v: float) -> int:
    """JS `Math.round` — the heatmap rounds every cell before it looks."""
    return int(math.floor(v + 0.5))


def _contrast(color: str) -> str:
    """`contrastText` of utils/statusBands.js: the ink on a band colour."""
    h = color.lstrip("#")
    if len(h) == 3:
        h = "".join(ch * 2 for ch in h)
    try:
        r, g, b = int(h[0:2], 16), int(h[2:4], 16), int(h[4:6], 16)
    except ValueError:
        return "#111827"
    return "#111827" if (0.299 * r + 0.587 * g + 0.114 * b) / 255 > 0.52 else "#ffffff"


def _paint(p: float, segs: list[dict]) -> dict:
    """The band a whole value falls in — `getSegmentColor` / `segmentFill`."""
    if not segs:
        return {}
    hit = segs[0]
    for sg in segs:
        if p >= sg["from"]:
            hit = sg
        else:
            break
    return {"bg": hit["color"], "fg": _contrast(hit["color"])}


def _edge(x: float, minus: bool = False) -> str:
    s = str(int(x)) if float(x).is_integer() else f"{x:g}"
    return s.replace("-", "−") if minus else s


def legend_items(segs: list[dict], kind: str, words: dict) -> list[dict]:
    """The colour key — `segmentBands` (a percent) or the comparison table's
    `diffLegendBands` (a difference), band by band."""
    out = []
    for i, sg in enumerate(segs):
        nxt = segs[i + 1] if i + 1 < len(segs) else None
        if kind == "diff":
            if i == 0 and nxt:
                label = words["below"].replace("{v}", _edge(nxt["from"], True))
            elif not nxt:
                label = words["above"].replace("{v}", _edge(sg["from"] - 1, True))
            else:
                label = (words["range"].replace("{a}", _edge(sg["from"], True))
                         .replace("{b}", _edge(nxt["from"] - 1, True)))
        else:
            if i == 0:
                label = f"< {_edge(nxt['from'])}%" if nxt else f"≥ {_edge(sg['from'])}%"
            elif not nxt:
                label = f"≥ {_edge(sg['from'])}%"
            else:
                label = f"{_edge(sg['from'])} to {_edge(nxt['from'] - 1)}%"
        out.append({"color": sg["color"], "label": label})
    return out


def grid_spec(block: dict, *, title: str, lines: list[str], lang: Optional[str] = None) -> dict:
    """A normalised grid block as the spec `grid_export.lay_out` draws —
    the shape `utils/gridExport.js` builds for the page's own «Excel»."""
    words = _grid_words(lang)
    segs = block.get("segments") or []
    kind, stat = block["kind"], block["stat"]
    cols = block["columns"]
    spec_cols = [{"w": 30}] + [{"w": 9 if kind == "diff" else 8} for _ in cols[1:]]
    head = [{"t": cols[0]}] + [{"t": c} for c in cols[1:]]
    if stat != "none":
        spec_cols.append({"w": 9, "sep": True})
        head.append({"t": words["stat"][stat]})
    seen: dict[str, str] = {}

    def cell(v: Any) -> dict:
        if isinstance(v, dict):
            if "pending" in v:
                icon = v["pending"]
                seen.setdefault(icon, words[PENDING.get(icon, "pending_title")])
                return {"hatch": True, "t": icon}
            return {"t": v.get("text", "")}
        if v is None:
            return {"t": "—"}
        p = _whole(v)
        return {"v": p, "k": kind, **_paint(p, segs)}

    def row_cells(r: list, unit: bool = False) -> list[dict]:
        name: dict = {"t": r[0]}
        if unit:
            name["unit"] = True
        out = [name, *[cell(v) for v in r[1:]]]
        if stat != "none":
            nums = [_whole(v) for v in r[1:] if isinstance(v, (int, float))]
            if nums:
                sv = (max(nums) if stat == "max" else min(nums) if stat == "min"
                      else _whole(sum(nums) / len(nums)))
                out.append({"v": sv, "k": kind, **_paint(sv, segs)})
            else:
                out.append({"t": "—"})
        return out

    body = [row_cells(r) for r in block["rows"]]
    foot = [row_cells(r, unit=True) for r in block.get("foot") or []]
    legend = []
    if segs:
        legend.append({"title": words["colors"], "items": legend_items(segs, kind, words)})
    if seen:
        legend.append({"title": words["pending"],
                       "items": [{"hatch": True, "icon": ic, "label": lb} for ic, lb in seen.items()]})
    return {"title": title, "lines": list(lines), "sheet": title, "cols": spec_cols,
            "head": [head], "body": body, "foot": foot, "legend": legend}


def grid_as_tables(block: dict, lang: Optional[str] = None, per: int = 0,
                   plain_pending: bool = False) -> list[dict]:
    """The grid for a format that cannot fill a cell: a table of percents, the
    unit rows after the body. A page or a slide holds only so many day
    columns, so `per` cuts the days into chunks — the name column repeated,
    the chunk's days named in the caption, the summary column (over the WHOLE
    period) on the last chunk."""
    cols = block["columns"]
    days = cols[1:]
    if not per or len(days) <= per:
        return [grid_as_table(block, lang, plain_pending)]
    out = []
    for k in range(0, len(days), per):
        sub = dict(block)
        sub["columns"] = [cols[0], *days[k:k + per]]
        sub["rows"] = [[r[0], *r[1 + k:1 + k + per]] for r in block["rows"]]
        sub["foot"] = [[r[0], *r[1 + k:1 + k + per]] for r in block.get("foot") or []]
        last = k + per >= len(days)
        if not last:
            sub["stat"] = "none"
        cap = block.get("caption") or ""
        sub["caption"] = f"{cap} ({days[k]} – {days[min(k + per, len(days)) - 1]})".strip()
        t = grid_as_table(sub, lang, plain_pending)
        if last and block["stat"] != "none":
            # The summary is the whole period's, not this chunk's.
            whole = grid_as_table(block, lang, plain_pending)
            for row, full in zip(t["rows"], whole["rows"]):
                row[-1] = full[-1]
        out.append(t)
    return out


def grid_as_table(block: dict, lang: Optional[str] = None, plain_pending: bool = False) -> dict:
    """The grid for a format that cannot fill a cell: a table of percents,
    the unit rows after the body. `plain_pending` prints a pending day as «…»
    — the PDF's one font has no emoji."""
    words = _grid_words(lang)
    kind, stat = block["kind"], block["stat"]

    def txt(v: Any) -> str:
        if isinstance(v, dict):
            if "pending" in v:
                return "…" if plain_pending else v["pending"]
            return v.get("text", "")
        if v is None:
            return "—"
        p = _whole(v)
        if kind == "diff":
            return f"+{p}%" if p > 0 else (f"−{-p}%" if p < 0 else "0%")
        return f"{p}%"

    cols = list(block["columns"]) + ([words["stat"][stat]] if stat != "none" else [])
    rows = []
    for r in list(block["rows"]) + list(block.get("foot") or []):
        cells = [r[0] or ""] + [txt(v) for v in r[1:]]
        if stat != "none":
            nums = [_whole(v) for v in r[1:] if isinstance(v, (int, float))]
            cells.append(txt(float(max(nums) if stat == "max" else min(nums) if stat == "min"
                                   else _whole(sum(nums) / len(nums)))) if nums else "—")
        rows.append(cells)
    return {"type": "table", "columns": cols, "rows": rows,
            "caption": block.get("caption", ""), "truncated": block.get("truncated", 0)}


def build_docx(title: str, subtitle: str, blocks: list[dict]) -> bytes:
    from docx import Document
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt, RGBColor

    ink, muted, gold = RGBColor.from_string(INK), RGBColor.from_string(MUTED), RGBColor.from_string(BRAND)

    def shade(cell, fill: str) -> None:
        tcPr = cell._tc.get_or_add_tcPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear")
        shd.set(qn("w:color"), "auto")
        shd.set(qn("w:fill"), fill)
        tcPr.append(shd)

    def no_borders(table) -> None:
        tblPr = table._tbl.tblPr
        borders = OxmlElement("w:tblBorders")
        for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
            el = OxmlElement(f"w:{edge}")
            el.set(qn("w:val"), "nil")
            borders.append(el)
        tblPr.append(borders)

    def light_borders(table) -> None:
        tblPr = table._tbl.tblPr
        borders = OxmlElement("w:tblBorders")
        for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
            el = OxmlElement(f"w:{edge}")
            el.set(qn("w:val"), "single")
            el.set(qn("w:sz"), "4")
            el.set(qn("w:color"), RULE)
            borders.append(el)
        tblPr.append(borders)

    def write_runs(par, text: str, size: Optional[float] = None, color=None, bold_all: bool = False) -> None:
        for s, bold in runs(text):
            r = par.add_run(s)
            r.bold = bold or bold_all
            if size:
                r.font.size = Pt(size)
            if color is not None:
                r.font.color.rgb = color

    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
    sec.left_margin = sec.right_margin = Cm(2.0)
    sec.top_margin, sec.bottom_margin = Cm(1.8), Cm(1.8)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    normal.font.color.rgb = ink
    normal.paragraph_format.space_after = Pt(6)
    for name, size in (("Heading 1", 15), ("Heading 2", 12.5), ("Heading 3", 11.5)):
        st = doc.styles[name]
        st.font.name = "Calibri"
        st.font.size = Pt(size)
        st.font.bold = True
        st.font.color.rgb = ink
        st.element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
        st.paragraph_format.space_before = Pt(14 if name == "Heading 1" else 10)
        st.paragraph_format.space_after = Pt(4)
        st.paragraph_format.keep_with_next = True

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    write_runs(p, plain(title), size=22, color=ink, bold_all=True)
    if subtitle:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(4)
        write_runs(p, plain(subtitle), size=10, color=muted)
    # a thin gold rule under the title block
    rule = doc.add_paragraph()
    rule.paragraph_format.space_after = Pt(10)
    pPr = rule._p.get_or_add_pPr()
    bdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "12")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), BRAND)
    bdr.append(bottom)
    pPr.append(bdr)

    for b in blocks:
        t = b["type"]
        if t == "heading":
            doc.add_heading(plain(b["text"]), level=b["level"])
        elif t == "paragraph":
            write_runs(doc.add_paragraph(), b["text"])
        elif t in ("bullets", "numbered"):
            style = "List Bullet" if t == "bullets" else "List Number"
            for item in b["items"]:
                par = doc.add_paragraph(style=style)
                par.paragraph_format.space_after = Pt(2)
                write_runs(par, item)
        elif t == "quote":
            par = doc.add_paragraph()
            par.paragraph_format.left_indent = Cm(0.5)
            par.paragraph_format.space_before = Pt(4)
            par.paragraph_format.space_after = Pt(8)
            pPr = par._p.get_or_add_pPr()
            bdr = OxmlElement("w:pBdr")
            left = OxmlElement("w:left")
            left.set(qn("w:val"), "single")
            left.set(qn("w:sz"), "18")
            left.set(qn("w:space"), "8")
            left.set(qn("w:color"), BRAND)
            bdr.append(left)
            pPr.append(bdr)
            write_runs(par, b["text"], color=RGBColor.from_string("3F3F46"))
        elif t == "facts":
            tbl = doc.add_table(rows=0, cols=2)
            tbl.alignment = WD_TABLE_ALIGNMENT.LEFT
            no_borders(tbl)
            for label, value in b["items"]:
                cells = tbl.add_row().cells
                cells[0].paragraphs[0].paragraph_format.space_after = Pt(2)
                cells[1].paragraphs[0].paragraph_format.space_after = Pt(2)
                write_runs(cells[0].paragraphs[0], plain(label), size=10, color=muted)
                write_runs(cells[1].paragraphs[0], value, size=11, bold_all=True)
                cells[0].width, cells[1].width = Cm(5.5), Cm(11.5)
            doc.add_paragraph().paragraph_format.space_after = Pt(2)
        elif t == "table":
            cols, rows = b["columns"], table_text(b)
            nums = numeric_columns(rows, len(cols))
            if b.get("caption"):
                cap = doc.add_paragraph()
                cap.paragraph_format.space_after = Pt(3)
                cap.paragraph_format.keep_with_next = True
                write_runs(cap, plain(b["caption"]), size=10, color=muted, bold_all=True)
            tbl = doc.add_table(rows=1, cols=len(cols))
            tbl.alignment = WD_TABLE_ALIGNMENT.LEFT
            tbl.autofit = False
            light_borders(tbl)
            weights = _col_weights(cols, rows)
            tw = _table_width(weights, 0.2, 17.0)
            widths = [Cm(tw * w / sum(weights)) for w in weights]
            hdr = tbl.rows[0]
            trPr = hdr._tr.get_or_add_trPr()
            th = OxmlElement("w:tblHeader")
            th.set(qn("w:val"), "true")
            trPr.append(th)
            for j, c in enumerate(cols):
                cell = hdr.cells[j]
                shade(cell, BRAND)
                par = cell.paragraphs[0]
                par.paragraph_format.space_after = Pt(0)
                if nums[j]:
                    par.alignment = WD_ALIGN_PARAGRAPH.RIGHT
                r = par.add_run(c)
                r.bold = True
                r.font.size = Pt(9.5)
                r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
                cell.width = widths[j]
            for j, wd in enumerate(widths):
                tbl.columns[j].width = wd
            for i, row in enumerate(rows):
                cells = tbl.add_row().cells
                for j, v in enumerate(row):
                    cells[j].width = widths[j]
                    par = cells[j].paragraphs[0]
                    par.paragraph_format.space_after = Pt(0)
                    if nums[j]:
                        par.alignment = WD_ALIGN_PARAGRAPH.RIGHT
                    r = par.add_run(v)
                    r.font.size = Pt(9.5)
                    if v == "—":
                        r.font.color.rgb = muted
                    if i % 2 == 1:
                        shade(cells[j], ZEBRA)
            if b.get("truncated"):
                note = doc.add_paragraph()
                write_runs(note, f"… yana {b['truncated']} qator faylga sig'madi", size=9, color=muted)
            doc.add_paragraph().paragraph_format.space_after = Pt(2)
        elif t == "page_break":
            doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    foot = sec.footer.paragraphs[0]
    foot.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    write_runs(foot, footer_text(), size=8.5, color=muted)

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


# ── pdf ──────────────────────────────────────────────────────────────────────

@lru_cache(maxsize=2)
def _ttf(bold: bool) -> str:
    """A Unicode TTF on this server — the camera stamp's own resolver, so the
    two never disagree about which font the box has."""
    from app.services.downtime_card import _find_font
    return _find_font(bold)


def build_pdf(title: str, subtitle: str, blocks: list[dict]) -> bytes:
    from fpdf import FPDF
    from fpdf.enums import TableBordersLayout, TableCellFillMode, XPos, YPos
    from fpdf.fonts import FontFace

    try:
        regular, bold = _ttf(False), _ttf(True)
    except Exception as exc:  # the stamp's CardError, or no fonts at all
        raise DocError("a PDF needs a Unicode font on the server and none was found — "
                       "make a Word (docx) file instead") from exc

    gold = (0xC8, 0x97, 0x3F)
    ink = (0x0F, 0x11, 0x17)
    mute = (0x6B, 0x72, 0x80)
    zebra = (0xF6, 0xF4, 0xEF)
    rule = (0xDD, 0xDA, 0xD2)

    class _PDF(FPDF):
        def footer(self):
            self.set_y(-14)
            self.set_font("Body", "", 8)
            self.set_text_color(*mute)
            self.cell(0, 6, f"{footer_text()}  ·  {self.page_no()}/{{nb}}", align="R")

    pdf = _PDF(orientation="P", unit="mm", format="A4")
    pdf.alias_nb_pages()
    pdf.set_margins(18, 16, 18)
    pdf.set_auto_page_break(auto=True, margin=20)
    pdf.add_font("Body", "", regular)
    pdf.add_font("Body", "B", bold)
    pdf.set_font("Body", "", 10.5)
    pdf.add_page()
    width = pdf.w - pdf.l_margin - pdf.r_margin

    def para(text: str, size: float = 10.5, color=ink, h: float = 5.4, bold_all: bool = False,
             after: float = 2.5, indent: float = 0.0) -> None:
        pdf.set_font("Body", "B" if bold_all else "", size)
        pdf.set_text_color(*color)
        if indent:
            pdf.set_x(pdf.l_margin + indent)
        pdf.multi_cell(width - indent, h, text, markdown=not bold_all,
                       new_x=XPos.LMARGIN, new_y=YPos.NEXT)
        pdf.ln(after)

    pdf.set_font("Body", "B", 20)
    pdf.set_text_color(*ink)
    pdf.multi_cell(width, 9, plain(title), new_x=XPos.LMARGIN, new_y=YPos.NEXT)
    if subtitle:
        pdf.ln(0.5)
        pdf.set_font("Body", "", 9.5)
        pdf.set_text_color(*mute)
        pdf.multi_cell(width, 5, plain(subtitle), new_x=XPos.LMARGIN, new_y=YPos.NEXT)
    pdf.ln(2)
    pdf.set_draw_color(*gold)
    pdf.set_line_width(0.6)
    y = pdf.get_y()
    pdf.line(pdf.l_margin, y, pdf.l_margin + width, y)
    pdf.ln(6)

    for b in blocks:
        t = b["type"]
        if t == "heading":
            size = {1: 14, 2: 12, 3: 11}[b["level"]]
            if pdf.get_y() > pdf.h - 45:
                pdf.add_page()
            pdf.ln(2 if b["level"] == 1 else 1)
            para(plain(b["text"]), size=size, h=size * 0.5, bold_all=True, after=1.5)
        elif t == "paragraph":
            para(b["text"])
        elif t in ("bullets", "numbered"):
            for n, item in enumerate(b["items"], 1):
                mark = "•" if t == "bullets" else f"{n}."
                pdf.set_font("Body", "", 10.5)
                pdf.set_text_color(*ink)
                pdf.set_x(pdf.l_margin)
                pdf.cell(7, 5.4, mark)
                pdf.multi_cell(width - 7, 5.4, item, markdown=True, new_x=XPos.LMARGIN, new_y=YPos.NEXT)
                pdf.ln(0.8)
            pdf.ln(2)
        elif t == "quote":
            pdf.set_font("Body", "", 10)
            txt = plain(b["text"])
            h = pdf.multi_cell(width - 8, 5.2, txt, dry_run=True, output="HEIGHT")
            if pdf.get_y() + h + 6 > pdf.h - pdf.b_margin:
                pdf.add_page()
            y0 = pdf.get_y()
            pdf.set_fill_color(0xFA, 0xF3, 0xE5)
            pdf.rect(pdf.l_margin, y0, width, h + 5, style="F")
            pdf.set_fill_color(*gold)
            pdf.rect(pdf.l_margin, y0, 1.2, h + 5, style="F")
            pdf.set_xy(pdf.l_margin + 5, y0 + 2.5)
            pdf.set_text_color(0x3F, 0x3F, 0x46)
            pdf.multi_cell(width - 8, 5.2, txt, new_x=XPos.LMARGIN, new_y=YPos.NEXT)
            pdf.set_y(y0 + h + 5)
            pdf.set_fill_color(255, 255, 255)
            pdf.ln(4)
        elif t == "facts":
            pdf.set_draw_color(*rule)
            pdf.set_line_width(0.2)
            pdf.set_fill_color(255, 255, 255)   # the table's base style reads it
            with pdf.table(borders_layout=TableBordersLayout.NONE, first_row_as_headings=False,
                           col_widths=(32, 68), line_height=5.6, padding=(1.2, 1.5),
                           text_align=("LEFT", "LEFT"), width=width, align="LEFT",
                           markdown=True) as tbl:
                for label, value in b["items"]:
                    row = tbl.row()
                    row.cell(plain(label), style=FontFace(family="Body", size_pt=9.5, color=mute))
                    row.cell(value, style=FontFace(family="Body", emphasis="BOLD", size_pt=10.5, color=ink))
            pdf.ln(3)
        elif t == "table":
            cols, rows = b["columns"], table_text(b)
            nums = numeric_columns(rows, len(cols))
            if b.get("caption"):
                para(plain(b["caption"]), size=9.5, color=mute, h=5, bold_all=True, after=1)
            pdf.set_draw_color(*rule)
            pdf.set_line_width(0.2)
            pdf.set_font("Body", "", 9)
            pdf.set_text_color(*ink)
            pdf.set_fill_color(255, 255, 255)   # the table's base style reads it
            weights = _col_weights(cols, rows)
            aligns = tuple("RIGHT" if n else "LEFT" for n in nums)
            with pdf.table(borders_layout=TableBordersLayout.HORIZONTAL_LINES,
                           cell_fill_color=zebra, cell_fill_mode=TableCellFillMode.EVEN_ROWS,
                           col_widths=tuple(weights), line_height=5.2, padding=(1.4, 1.6),
                           text_align=aligns, width=_table_width(weights, 1.9, width),
                           align="LEFT", repeat_headings=1,
                           headings_style=FontFace(family="Body", emphasis="BOLD", size_pt=9,
                                                   color=(255, 255, 255), fill_color=gold)) as tbl:
                row = tbl.row()
                for c in cols:
                    row.cell(c)
                for r in rows:
                    row = tbl.row()
                    for v in r:
                        row.cell(v)
            if b.get("truncated"):
                para(f"… yana {b['truncated']} qator faylga sig'madi", size=8.5, color=mute, h=4.5)
            pdf.ln(3)
        elif t == "page_break":
            pdf.add_page()

    return bytes(pdf.output())


# ── pptx ─────────────────────────────────────────────────────────────────────

_SLIDE_LINES = 14          # body lines a content slide holds
_TABLE_ROWS = 11           # table rows a slide holds (plus the header)
_CHARS_PER_LINE = 92


def _lines(text: str, per: int = _CHARS_PER_LINE) -> int:
    return max(1, math.ceil(len(plain(text)) / per))


def _slides(title: str, subtitle: str, blocks: list[dict]) -> list[dict]:
    """Blocks packed onto slides by a line budget. A level-1 heading opens a
    slide and names it; a table too long for one slide continues onto the
    next with «(davomi)»; everything else flows."""
    slides: list[dict] = [{"kind": "title", "title": plain(title), "subtitle": plain(subtitle)}]
    cur: Optional[dict] = None
    used = 0

    def new(name: str) -> None:
        nonlocal cur, used
        cur = {"kind": "content", "title": name, "items": []}
        slides.append(cur)
        used = 0

    def ensure(name: str) -> None:
        if cur is None:
            new(name)

    for b in blocks:
        t = b["type"]
        if t == "heading" and b["level"] == 1:
            new(plain(b["text"]))
            continue
        if t == "page_break":
            if cur is not None and cur["items"]:
                new(_continued(cur["title"]))
            continue
        ensure(plain(title))
        if t == "heading":
            need = 1
        elif t in ("paragraph", "quote"):
            need = _lines(b["text"])
        elif t in ("bullets", "numbered"):
            need = sum(_lines(x, 88) for x in b["items"])
        elif t == "facts":
            need = len(b["items"])
        else:  # table
            need = min(len(b["rows"]), _TABLE_ROWS) + 2
        if t == "table":
            rows = b["rows"]
            start = 0
            first = True
            while True:
                if used + 3 > _SLIDE_LINES:
                    new(_continued(cur["title"]))
                room = max(1, min(_TABLE_ROWS, (_SLIDE_LINES - used - 2)))
                chunk = rows[start:start + room]
                cur["items"].append({**b, "rows": chunk, "caption": b["caption"] if first else ""})
                used += len(chunk) + 2
                start += len(chunk)
                first = False
                if start >= len(rows):
                    break
                new(_continued(cur["title"]))
            continue
        if used + need > _SLIDE_LINES and cur["items"]:
            new(_continued(cur["title"]))
        if t in ("bullets", "numbered") and need > _SLIDE_LINES:
            items = list(b["items"])
            while items:
                room, take = _SLIDE_LINES - used, []
                while items and (room - _lines(items[0], 88)) >= 0:
                    room -= _lines(items[0], 88)
                    take.append(items.pop(0))
                if not take:
                    take.append(items.pop(0))
                cur["items"].append({**b, "items": take})
                used = _SLIDE_LINES
                if items:
                    new(_continued(cur["title"]))
            continue
        cur["items"].append(b)
        used += need
    return slides


def build_pptx(title: str, subtitle: str, blocks: list[dict]) -> bytes:
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
    from pptx.util import Emu, Inches, Pt

    gold, ink, mute = RGBColor.from_string(BRAND), RGBColor.from_string(INK), RGBColor.from_string(MUTED)
    white = RGBColor(0xFF, 0xFF, 0xFF)
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    blank = prs.slide_layouts[6]
    W, H = prs.slide_width, prs.slide_height
    margin = Inches(0.7)

    def textbox(slide, left, top, width, height):
        tb = slide.shapes.add_textbox(left, top, width, height)
        tf = tb.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_right = Emu(0)
        tf.margin_top = tf.margin_bottom = Emu(0)
        return tf

    def put_runs(par, text: str, size: float, color, bold_all: bool = False) -> None:
        for s, bold in runs(text):
            r = par.add_run()
            r.text = s
            r.font.size = Pt(size)
            r.font.bold = bold or bold_all
            r.font.color.rgb = color
            r.font.name = "Calibri"

    def footer(slide, n: int) -> None:
        tf = textbox(slide, margin, H - Inches(0.45), W - 2 * margin, Inches(0.3))
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.RIGHT
        put_runs(p, f"{footer_text()}  ·  {n}", 9, mute)

    def fill(shape, color) -> None:
        shape.fill.solid()
        shape.fill.fore_color.rgb = color
        shape.line.fill.background()

    plan = _slides(title, subtitle, blocks)
    for n, s in enumerate(plan, 1):
        slide = prs.slides.add_slide(blank)
        if s["kind"] == "title":
            bg = slide.background.fill
            bg.solid()
            bg.fore_color.rgb = ink
            bar = slide.shapes.add_shape(1, margin, Inches(2.9), Inches(0.09), Inches(1.5))
            fill(bar, gold)
            tf = textbox(slide, margin + Inches(0.3), Inches(2.75), W - 2 * margin - Inches(0.3), Inches(1.8))
            tf.vertical_anchor = MSO_ANCHOR.TOP
            put_runs(tf.paragraphs[0], s["title"], 36, white, bold_all=True)
            if s["subtitle"]:
                p = tf.add_paragraph()
                p.space_before = Pt(10)
                put_runs(p, s["subtitle"], 16, RGBColor(0xB4, 0xB8, 0xC2))
            tf2 = textbox(slide, margin, H - Inches(0.45), W - 2 * margin, Inches(0.3))
            tf2.paragraphs[0].alignment = PP_ALIGN.RIGHT
            put_runs(tf2.paragraphs[0], footer_text(), 9, RGBColor(0x8A, 0x8F, 0x98))
            continue
        bar = slide.shapes.add_shape(1, margin, Inches(0.55), Inches(0.07), Inches(0.55))
        fill(bar, gold)
        tf = textbox(slide, margin + Inches(0.25), Inches(0.5), W - 2 * margin - Inches(0.25), Inches(0.7))
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        put_runs(tf.paragraphs[0], s["title"], 26, ink, bold_all=True)
        y = Inches(1.45)
        for b in s["items"]:
            t = b["type"]
            if t == "table":
                cols, rows = b["columns"], table_text(b)
                nums = numeric_columns(rows, len(cols))
                if b.get("caption"):
                    cap = textbox(slide, margin, y, W - 2 * margin, Inches(0.3))
                    put_runs(cap.paragraphs[0], b["caption"], 12, mute, bold_all=True)
                    y += Inches(0.35)
                row_h = Inches(0.34)
                weights = _col_weights(cols, rows)
                total = sum(weights) or 1.0
                tw = Emu(int(_table_width(weights, Inches(0.095), W - 2 * margin)))
                shape = slide.shapes.add_table(len(rows) + 1, len(cols), margin, y,
                                               tw, row_h * (len(rows) + 1))
                tbl = shape.table
                tblPr = tbl._tbl.tblPr
                tblPr.set("bandRow", "0")
                tblPr.set("firstRow", "0")
                for j, wgt in enumerate(weights):
                    tbl.columns[j].width = Emu(int(tw * (wgt / total)))
                for i in range(len(rows) + 1):
                    tbl.rows[i].height = row_h
                    for j in range(len(cols)):
                        cell = tbl.cell(i, j)
                        cell.margin_left = cell.margin_right = Inches(0.08)
                        cell.margin_top = cell.margin_bottom = Inches(0.03)
                        cell.vertical_anchor = MSO_ANCHOR.MIDDLE
                        text = cols[j] if i == 0 else rows[i - 1][j]
                        tfc = cell.text_frame
                        tfc.word_wrap = True
                        p = tfc.paragraphs[0]
                        p.alignment = PP_ALIGN.RIGHT if nums[j] else PP_ALIGN.LEFT
                        r = p.add_run()
                        r.text = text
                        r.font.name = "Calibri"
                        r.font.size = Pt(12 if i == 0 else 11)
                        r.font.bold = i == 0
                        r.font.color.rgb = white if i == 0 else (mute if text == "—" else ink)
                        cell.fill.solid()
                        cell.fill.fore_color.rgb = gold if i == 0 else (
                            RGBColor.from_string(ZEBRA) if i % 2 == 0 else white)
                y += row_h * (len(rows) + 1) + Inches(0.25)
                continue
            if t == "facts":
                rows_f = b["items"]
                row_h = Inches(0.36)
                shape = slide.shapes.add_table(len(rows_f), 2, margin, y, W - 2 * margin, row_h * len(rows_f))
                tbl = shape.table
                tbl._tbl.tblPr.set("bandRow", "0")
                tbl._tbl.tblPr.set("firstRow", "0")
                tbl.columns[0].width = Emu(int((W - 2 * margin) * 0.32))
                tbl.columns[1].width = Emu(int((W - 2 * margin) * 0.68))
                for i, (label, value) in enumerate(rows_f):
                    tbl.rows[i].height = row_h
                    for j, text in enumerate((plain(label), value)):
                        cell = tbl.cell(i, j)
                        cell.margin_left = Inches(0.05)
                        cell.vertical_anchor = MSO_ANCHOR.MIDDLE
                        cell.fill.background()
                        p = cell.text_frame.paragraphs[0]
                        put_runs(p, text, 12 if j == 0 else 14, mute if j == 0 else ink, bold_all=(j == 1))
                y += row_h * len(rows_f) + Inches(0.25)
                continue
            est = Inches(0.32) * (
                1 if t == "heading" else
                _lines(b["text"]) if t in ("paragraph", "quote") else
                sum(_lines(x, 88) for x in b["items"]))
            if t == "quote":
                bar = slide.shapes.add_shape(1, margin, y, Inches(0.05), est + Inches(0.1))
                fill(bar, gold)
                tfb = textbox(slide, margin + Inches(0.22), y, W - 2 * margin - Inches(0.22), est + Inches(0.1))
            else:
                tfb = textbox(slide, margin, y, W - 2 * margin, est + Inches(0.1))
            first = True
            if t == "heading":
                put_runs(tfb.paragraphs[0], plain(b["text"]), 18 if b["level"] == 2 else 15, ink, bold_all=True)
            elif t == "paragraph":
                put_runs(tfb.paragraphs[0], b["text"], 15, ink)
            elif t == "quote":
                put_runs(tfb.paragraphs[0], plain(b["text"]), 15, RGBColor(0x3F, 0x3F, 0x46))
            else:
                for i, item in enumerate(b["items"], 1):
                    p = tfb.paragraphs[0] if first else tfb.add_paragraph()
                    first = False
                    p.space_after = Pt(4)
                    mark = "•  " if t == "bullets" else f"{i}.  "
                    put_runs(p, mark + item, 15, ink)
            y += est + Inches(0.18)
        footer(slide, n)

    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


# ── xlsx / csv ───────────────────────────────────────────────────────────────

def _sheet_name(block: dict, last_heading: str, i: int) -> str:
    return plain(block.get("caption") or last_heading or f"Jadval {i}")


def _unique_sheet(name: str, used: set) -> str:
    from app.services.grid_export import _sheet_name as clean
    base = clean(name)
    out, n = base, 2
    while out.lower() in used:
        suffix = f" ({n})"
        out = base[: 31 - len(suffix)] + suffix
        n += 1
    used.add(out.lower())
    return out


def _num_format(v: float) -> str:
    if float(v).is_integer():
        return "#,##0"
    return "#,##0.0" if round(v, 1) == v else "#,##0.00"


def _sheet_head(ws, title: str, lines: list[str]) -> int:
    """The template's opening on every sheet — title, scope lines, a blank
    row — as the /zagruzka export writes it. → the first free row."""
    from openpyxl.styles import Alignment, Font
    from app.services import grid_export as ge

    ws.sheet_view.showGridLines = False
    ws.sheet_properties.tabColor = ge.BRAND
    row = 1
    c = ws.cell(row, 1, title[:300])
    c.font = Font(name=ge.FONT, size=14, bold=True, color=ge.INK)
    c.alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[row].height = 24
    row += 1
    for line in lines[:12]:
        if not line:
            continue
        c = ws.cell(row, 1, str(line)[:300])
        c.font = Font(name=ge.FONT, size=9, color=ge.SOFT)
        c.alignment = Alignment(horizontal="left", vertical="center")
        row += 1
    return row + 1


def _table_sheet(ws, title: str, lines: list[str], block: dict) -> None:
    """An ordinary table in the house look: the template's head, a gold
    header row, thin rules, zebra rows, numbers right-aligned under a
    number format, «—» blanks, frozen header + name column, an autofilter,
    landscape fit to width."""
    from openpyxl.styles import Alignment, Border, Font
    from openpyxl.utils import get_column_letter
    from app.services import grid_export as ge

    cols, rows = block["columns"], block["rows"]
    n = len(cols)
    row = _sheet_head(ws, title, lines)
    thin = Border(left=ge.THIN, right=ge.THIN, top=ge.THIN, bottom=ge.THIN)
    head_row = row
    ws.row_dimensions[row].height = 21
    for j, name in enumerate(cols, 1):
        c = ws.cell(row, j, str(name).upper()[:300])
        c.fill = ge._fill(ge.BRAND)
        c.font = Font(name=ge.FONT, size=9, bold=True, color=ge.HEAD_INK)
        c.alignment = Alignment(horizontal="left" if j == 1 else "center", vertical="center",
                                indent=1 if j == 1 else 0, wrap_text=True)
        c.border = thin
    row += 1
    text_rows = table_text(block)
    nums = numeric_columns(text_rows, n)
    zebra = ge._fill(ZEBRA)
    for k, r in enumerate(rows):
        ws.row_dimensions[row].height = 20
        for j, v in enumerate(r, 1):
            c = ws.cell(row, j)
            c.border = thin
            if k % 2:
                c.fill = zebra
            if v is None or isinstance(v, bool) or (isinstance(v, float) and (math.isnan(v) or math.isinf(v))):
                c.value = "✓" if v is True else "—"
                c.font = Font(name=ge.FONT, size=9, color=ge.INK if v is True else ge.DASH)
                c.alignment = Alignment(horizontal="center", vertical="center")
            elif isinstance(v, (int, float)):
                c.value = v
                c.number_format = _num_format(v)
                c.font = Font(name=ge.FONT, size=10, color=ge.INK)
                c.alignment = Alignment(horizontal="right", vertical="center")
            else:
                txt = json.dumps(v, ensure_ascii=False) if isinstance(v, (dict, list)) else str(v)
                c.value = txt[:MAX_TEXT]
                c.font = Font(name=ge.FONT, size=10, color=ge.NAME_INK if j == 1 else ge.INK)
                c.alignment = Alignment(horizontal="right" if nums[j - 1] and j > 1 else "left",
                                        vertical="center", indent=1 if j == 1 else 0)
        row += 1
    if block.get("truncated"):
        c = ws.cell(row, 1, f"… +{block['truncated']}")
        c.font = Font(name=ge.FONT, size=9, color=ge.SOFT)
        row += 1
    for j in range(1, n + 1):
        longest = max([len(str(cols[j - 1]))] +
                      [len(text_rows[k][j - 1]) for k in range(min(len(text_rows), 500))])
        ws.column_dimensions[get_column_letter(j)].width = min(max(longest + 3, 8), 48 if j == 1 else 28)
    ws.freeze_panes = ws.cell(head_row + 1, 2 if n >= 3 else 1)
    if rows:
        ws.auto_filter.ref = f"A{head_row}:{get_column_letter(n)}{head_row + len(rows)}"
    ws.print_title_rows = f"{head_row}:{head_row}"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True


def _notes_sheet(ws, title: str, lines: list[str], notes: list[tuple]) -> None:
    """The document's prose — headings, paragraphs, lists, facts — on a sheet
    of its own, first in the workbook."""
    from openpyxl.styles import Alignment, Font
    from app.services import grid_export as ge

    row = _sheet_head(ws, title, lines)
    ws.column_dimensions["A"].width = 34
    ws.column_dimensions["B"].width = 80
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    for note in notes:
        kind = note[0]
        if kind == "blank":
            row += 1
            continue
        if kind == "h":
            c = ws.cell(row, 1, note[1][:300])
            c.font = Font(name=ge.FONT, size=11, bold=True, color=ge.INK)
            ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=2)
        elif kind == "kv":
            a = ws.cell(row, 1, note[1][:300])
            a.font = Font(name=ge.FONT, size=10, color=ge.SOFT)
            b = ws.cell(row, 2, note[2][:MAX_TEXT])
            b.font = Font(name=ge.FONT, size=10, bold=True, color=ge.INK)
            b.alignment = Alignment(wrap_text=True, vertical="top")
        else:
            c = ws.cell(row, 1, note[1][:MAX_TEXT])
            c.font = Font(name=ge.FONT, size=10, color=ge.INK)
            c.alignment = Alignment(wrap_text=True, vertical="top")
            ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=2)
            ws.row_dimensions[row].height = max(15, 15 * (1 + len(note[1]) // 110))
        row += 1


def build_xlsx(title: str, subtitle: str, blocks: list[dict], lang: Optional[str] = None) -> bytes:
    """The workbook in the platform's own look. A «grid» block is the
    /zagruzka «Excel» template, sheet for sheet (`grid_export.lay_out`); a
    table is a house sheet; the prose goes on a first sheet. Every sheet
    opens with the title, the scope line and «Tayyorlandi: …», as the page's
    file does."""
    from openpyxl import Workbook
    from app.services import grid_export as ge

    words = _grid_words(lang)
    lines = [plain(subtitle)] if subtitle else []
    lines.append(f"{words['generated']}: {datetime.now(TZ).strftime('%d.%m.%Y, %H:%M')}")
    wb = Workbook()
    wb.remove(wb.active)
    used: set = set()
    notes: list[tuple] = []
    last_heading = ""
    n = 0
    for b in blocks:
        t = b["type"]
        if t == "grid":
            n += 1
            name = _sheet_name(b, last_heading, n)
            ws = wb.create_sheet(_unique_sheet(name, used))
            ge.lay_out(ws, grid_spec(b, title=plain(b.get("caption") or title), lines=lines, lang=lang))
        elif t == "table":
            n += 1
            name = _sheet_name(b, last_heading, n)
            ws = wb.create_sheet(_unique_sheet(name, used))
            _table_sheet(ws, plain(b.get("caption") or title), lines, b)
        elif t == "heading":
            last_heading = plain(b["text"])
            notes.append(("blank",))
            notes.append(("h", last_heading))
        elif t in ("paragraph", "quote"):
            notes.append(("p", plain(b["text"])))
        elif t in ("bullets", "numbered"):
            for i, item in enumerate(b["items"], 1):
                notes.append(("p", ("• " if t == "bullets" else f"{i}. ") + plain(item)))
        elif t == "facts":
            for label, value in b["items"]:
                notes.append(("kv", plain(label), plain(value)))
    if notes:
        ws = wb.create_sheet(_unique_sheet(plain(title)[:31] or "Hisobot", used), 0)
        _notes_sheet(ws, plain(title), lines, notes)
    if not wb.worksheets:
        raise DocError("an Excel file needs a table or some text")
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def build_csv(title: str, subtitle: str, blocks: list[dict]) -> bytes:
    tables = [b for b in blocks if b["type"] == "table"]
    if not tables:
        raise DocError("a CSV needs at least one table block — for text choose docx, pdf, md or txt")
    buf = io.StringIO()
    w = csv.writer(buf, lineterminator="\r\n")
    for i, b in enumerate(tables):
        if i:
            w.writerow([])
        if len(tables) > 1:
            w.writerow([plain(b.get("caption") or f"{plain(title)} — {i + 1}")])
        w.writerow(b["columns"])
        for r in table_text(b, machine=True):
            w.writerow(r)
    return ("\ufeff" + buf.getvalue()).encode("utf-8")


# ── md / txt / html ──────────────────────────────────────────────────────────

def _md_cell(s: str) -> str:
    return s.replace("|", "\\|").replace("\n", " ")


def build_md(title: str, subtitle: str, blocks: list[dict]) -> bytes:
    out = [f"# {title}".rstrip()]
    if subtitle:
        out.append(f"_{plain(subtitle)}_")
    out.append("")
    for b in blocks:
        t = b["type"]
        if t == "heading":
            out += ["#" * (b["level"] + 1) + " " + plain(b["text"]), ""]
        elif t == "paragraph":
            out += [b["text"], ""]
        elif t == "quote":
            out += ["> " + b["text"].replace("\n", "\n> "), ""]
        elif t == "bullets":
            out += ["- " + x for x in b["items"]] + [""]
        elif t == "numbered":
            out += [f"{i}. {x}" for i, x in enumerate(b["items"], 1)] + [""]
        elif t == "facts":
            out += ["| | |", "|---|---|"] + [f"| {_md_cell(plain(k))} | **{_md_cell(plain(v))}** |" for k, v in b["items"]] + [""]
        elif t == "table":
            cols, rows = b["columns"], table_text(b)
            nums = numeric_columns(rows, len(cols))
            if b.get("caption"):
                out.append(f"**{plain(b['caption'])}**")
                out.append("")
            out.append("| " + " | ".join(_md_cell(c) for c in cols) + " |")
            out.append("|" + "|".join("---:" if n else "---" for n in nums) + "|")
            out += ["| " + " | ".join(_md_cell(v) for v in r) + " |" for r in rows]
            if b.get("truncated"):
                out.append(f"\n_… yana {b['truncated']} qator_")
            out.append("")
        elif t == "page_break":
            out += ["---", ""]
    out.append(f"_{footer_text()}_")
    return "\n".join(out).encode("utf-8")


def _txt_table(cols: list[str], rows: list[list[str]]) -> list[str]:
    nums = numeric_columns(rows, len(cols))
    widths = [max([len(c)] + [len(r[j]) for r in rows]) for j, c in enumerate(cols)]
    widths = [min(w, 40) for w in widths]

    def line(cells: list[str]) -> str:
        parts = []
        for j, v in enumerate(cells):
            v = v if len(v) <= widths[j] else v[:widths[j] - 1] + "…"
            parts.append(v.rjust(widths[j]) if nums[j] else v.ljust(widths[j]))
        return "  ".join(parts).rstrip()
    return [line(cols), "  ".join("─" * w for w in widths)] + [line(r) for r in rows]


def build_txt(title: str, subtitle: str, blocks: list[dict]) -> bytes:
    out = [plain(title), "═" * min(len(plain(title)), 72)]
    if subtitle:
        out.append(plain(subtitle))
    out.append("")
    for b in blocks:
        t = b["type"]
        if t == "heading":
            text = plain(b["text"])
            out += ["", text.upper() if b["level"] == 1 else text, "─" * min(len(text), 72), ""]
        elif t == "paragraph":
            out += [plain(b["text"]), ""]
        elif t == "quote":
            out += ["│ " + plain(b["text"]).replace("\n", "\n│ "), ""]
        elif t == "bullets":
            out += ["• " + plain(x) for x in b["items"]] + [""]
        elif t == "numbered":
            out += [f"{i}. {plain(x)}" for i, x in enumerate(b["items"], 1)] + [""]
        elif t == "facts":
            wl = max((len(plain(k)) for k, _ in b["items"]), default=0)
            out += [f"{plain(k).ljust(wl)}  {plain(v)}" for k, v in b["items"]] + [""]
        elif t == "table":
            if b.get("caption"):
                out.append(plain(b["caption"]))
            out += _txt_table(b["columns"], table_text(b))
            if b.get("truncated"):
                out.append(f"… yana {b['truncated']} qator")
            out.append("")
        elif t == "page_break":
            out += ["", "·" * 24, ""]
    out += ["", footer_text()]
    return "\n".join(out).encode("utf-8")


_HTML_CSS = f"""
:root{{--ink:#{INK};--muted:#{MUTED};--gold:#{BRAND};--zebra:#{ZEBRA};--rule:#{RULE}}}
*{{box-sizing:border-box}}
body{{margin:0;padding:40px 24px;font:15px/1.55 -apple-system,"Segoe UI",Inter,Roboto,Arial,sans-serif;color:var(--ink);background:#fff}}
main{{max-width:860px;margin:0 auto}}
h1{{font-size:28px;line-height:1.2;margin:0 0 4px;font-weight:700;letter-spacing:-.01em}}
.sub{{color:var(--muted);font-size:13px;margin:0 0 14px}}
hr.rule{{border:0;border-top:2px solid var(--gold);margin:0 0 28px}}
h2{{font-size:20px;margin:30px 0 8px}}h3{{font-size:16px;margin:22px 0 6px}}h4{{font-size:14.5px;margin:18px 0 4px}}
p{{margin:0 0 12px}}ul,ol{{margin:0 0 14px;padding-left:22px}}li{{margin:3px 0}}
table{{border-collapse:collapse;width:100%;margin:6px 0 22px;font-size:13.5px}}
caption{{text-align:left;color:var(--muted);font-weight:600;font-size:13px;padding:0 0 6px}}
th{{background:var(--gold);color:#fff;text-align:left;font-weight:600;padding:8px 10px;white-space:nowrap}}
td{{padding:7px 10px;border-bottom:1px solid var(--rule);vertical-align:top}}
tr:nth-child(even) td{{background:var(--zebra)}}
td.n,th.n{{text-align:right;font-variant-numeric:tabular-nums;white-space:nowrap}}
td.dash{{color:var(--muted)}}
table.facts td{{border:0;padding:5px 10px 5px 0;background:none}}
table.facts td:first-child{{color:var(--muted);width:32%}}table.facts td:last-child{{font-weight:600}}
blockquote{{margin:12px 0 18px;padding:10px 14px;border-left:4px solid var(--gold);background:#FAF3E5;color:#3F3F46}}
.note{{color:var(--muted);font-size:12.5px}}
footer{{margin-top:48px;color:var(--muted);font-size:12px;text-align:right}}
.pb{{page-break-after:always;break-after:page;height:0;margin:0}}
table.grid{{width:auto;max-width:100%}}table.grid th{{text-align:center;font-size:12px;text-transform:uppercase}}
table.grid th:first-child{{text-align:left}}
table.grid td{{text-align:center;font-weight:600;border:1px solid var(--rule);padding:6px 8px;font-variant-numeric:tabular-nums;white-space:nowrap}}
table.grid td:first-child{{text-align:left;font-weight:400;color:#374151;border-right:2px solid #C4C9D0}}
table.grid tr:nth-child(even) td{{background:none}}table.grid td.dash{{color:#C9CDD3;font-weight:400}}
table.grid td.hatch{{background:repeating-linear-gradient(135deg,#fff 0 4px,#D9DCE1 4px 6px);color:#9CA3AF;font-weight:400}}
table.grid td.stat,table.grid th.stat{{border-left:2px solid #C4C9D0}}
table.grid tfoot td{{font-weight:700;background:#F5F1EA;border-top:2px solid #C4C9D0}}
ul.legend{{list-style:none;padding:0;margin:-14px 0 22px;display:flex;flex-wrap:wrap;gap:6px 18px;font-size:12.5px;color:var(--muted)}}
ul.legend .sw{{display:inline-block;width:14px;height:14px;border:1px solid var(--rule);vertical-align:-2px;margin-right:6px}}
ul.legend .sw.hatch{{background:repeating-linear-gradient(135deg,#fff 0 3px,#D9DCE1 3px 5px)}}
@media print{{body{{padding:0}}}}
"""


def _html_inline(text: str) -> str:
    return "".join(f"<strong>{html.escape(s)}</strong>" if b else html.escape(s) for s, b in runs(text))


def build_html(title: str, subtitle: str, blocks: list[dict], lang: Optional[str] = None) -> bytes:
    out = ["<!doctype html>", '<html lang="uz"><head><meta charset="utf-8">',
           '<meta name="viewport" content="width=device-width, initial-scale=1">',
           f"<title>{html.escape(plain(title))}</title>", f"<style>{_HTML_CSS}</style></head><body><main>",
           f"<h1>{_html_inline(title)}</h1>"]
    if subtitle:
        out.append(f'<p class="sub">{html.escape(plain(subtitle))}</p>')
    out.append('<hr class="rule">')
    for b in blocks:
        t = b["type"]
        if t == "heading":
            out.append(f"<h{b['level'] + 1}>{_html_inline(b['text'])}</h{b['level'] + 1}>")
        elif t == "paragraph":
            out.append(f"<p>{_html_inline(b['text'])}</p>")
        elif t == "quote":
            out.append(f"<blockquote>{_html_inline(b['text'])}</blockquote>")
        elif t in ("bullets", "numbered"):
            tag = "ul" if t == "bullets" else "ol"
            out.append(f"<{tag}>" + "".join(f"<li>{_html_inline(x)}</li>" for x in b["items"]) + f"</{tag}>")
        elif t == "facts":
            out.append('<table class="facts">' + "".join(
                f"<tr><td>{html.escape(plain(k))}</td><td>{_html_inline(v)}</td></tr>" for k, v in b["items"]) + "</table>")
        elif t == "table":
            cols, rows = b["columns"], table_text(b)
            nums = numeric_columns(rows, len(cols))
            cap = f"<caption>{html.escape(plain(b['caption']))}</caption>" if b.get("caption") else ""
            head = "".join(f'<th{" class=n" if nums[j] else ""}>{html.escape(c)}</th>' for j, c in enumerate(cols))
            body = "".join("<tr>" + "".join(
                f'<td class="{"n " if nums[j] else ""}{"dash" if v == "—" else ""}">{html.escape(v)}</td>'
                for j, v in enumerate(r)) + "</tr>" for r in rows)
            out.append(f"<table>{cap}<thead><tr>{head}</tr></thead><tbody>{body}</tbody></table>")
            if b.get("truncated"):
                out.append(f'<p class="note">… yana {b["truncated"]} qator</p>')
        elif t == "grid":
            spec = grid_spec(b, title=plain(b.get("caption") or title), lines=[], lang=lang)
            cap = f"<caption>{html.escape(plain(b['caption']))}</caption>" if b.get("caption") else ""
            n_cols = len(spec["cols"])

            def gcell(c: dict, j: int, tag: str = "td") -> str:
                cls = ["stat"] if (j == n_cols - 1 and spec["cols"][j].get("sep")) else []
                if tag == "th":
                    return f'<th class="{" ".join(cls)}">{html.escape(str(c.get("t") or ""))}</th>'
                if c.get("hatch"):
                    cls.append("hatch")
                    return f'<td class="{" ".join(cls)}">{html.escape(str(c.get("t") or "⏳"))}</td>'
                v = c.get("v")
                if isinstance(v, (int, float)):
                    val = (f"+{v}%" if v > 0 else f"−{-v}%" if v < 0 else "0%") if c.get("k") == "diff" else f"{v}%"
                    style = f' style="background:{c["bg"]};color:{c["fg"]}"' if c.get("bg") else ""
                    return f'<td class="{" ".join(cls)}"{style}>{val}</td>'
                txt = str(c.get("t") or "—")
                if txt == "—":
                    cls.append("dash")
                return f'<td class="{" ".join(cls)}">{html.escape(txt)}</td>'

            head = "".join(gcell(c, j, "th") for j, c in enumerate(spec["head"][0]))
            body = "".join("<tr>" + "".join(gcell(c, j) for j, c in enumerate(r)) + "</tr>" for r in spec["body"])
            foot = ("<tfoot>" + "".join("<tr>" + "".join(gcell(c, j) for j, c in enumerate(r)) + "</tr>"
                                        for r in spec["foot"]) + "</tfoot>") if spec["foot"] else ""
            out.append(f'<table class="grid">{cap}<thead><tr>{head}</tr></thead><tbody>{body}</tbody>{foot}</table>')
            items = []
            for group in spec["legend"]:
                for it in group["items"]:
                    sw = ('<span class="sw hatch"></span>' if it.get("hatch")
                          else f'<span class="sw" style="background:{it["color"]}"></span>')
                    items.append(f"<li>{sw}{html.escape(str(it.get('icon') or ''))} {html.escape(it['label'])}</li>")
            if items:
                out.append('<ul class="legend">' + "".join(items) + "</ul>")
            if b.get("truncated"):
                out.append(f'<p class="note">… yana {b["truncated"]} qator</p>')
        elif t == "page_break":
            out.append('<hr class="pb">')
    out.append(f"<footer>{html.escape(footer_text())}</footer></main></body></html>")
    return "\n".join(out).encode("utf-8")


# ── the one door ─────────────────────────────────────────────────────────────

_BUILDERS = {"docx": build_docx, "pdf": build_pdf, "pptx": build_pptx, "xlsx": build_xlsx,
             "csv": build_csv, "md": build_md, "txt": build_txt, "html": build_html}


def build(fmt: str, title: str, subtitle: str, blocks: list[dict],
          lang: Optional[str] = None) -> tuple[str, str, bytes]:
    """→ (extension, mime, bytes). Blocking (python-docx, fpdf2 and python-pptx
    are CPU work) — call it from a worker thread."""
    fmt = norm_format(fmt)
    ext, mime = FORMATS[fmt]
    title = _text(title, 300) or "Hujjat"
    subtitle = _text(subtitle, 400)
    if fmt not in ("xlsx", "html"):
        # A page or a slide holds only so many day columns; a sheet holds them all.
        per = {"pdf": 8, "docx": 12, "pptx": 8}.get(fmt, 0)
        flat: list[dict] = []
        for b in blocks:
            flat.extend(grid_as_tables(b, lang, per, fmt == "pdf") if b["type"] == "grid" else [b])
        blocks = flat
    try:
        if fmt in ("xlsx", "html"):
            data = _BUILDERS[fmt](title, subtitle, blocks, lang)
        else:
            data = _BUILDERS[fmt](title, subtitle, blocks)
    except DocError:
        raise
    except Exception as exc:
        log.exception("assistant document (%s) failed", fmt)
        raise DocError(f"the {fmt} file could not be built: {exc}") from exc
    return ext, mime, data


def safe_filename(name: Any, fallback: str = "Hujjat") -> str:
    s = re.sub(r"[\\/:*?\"<>|\x00-\x1f]+", " ", str(name or "")).strip().strip(".")
    s = re.sub(r"\.(docx|pdf|pptx|xlsx|csv|md|txt|html?)$", "", s, flags=re.I).strip()
    return (s or fallback)[:80]
