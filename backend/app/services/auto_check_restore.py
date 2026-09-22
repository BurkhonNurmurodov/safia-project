"""Points the «one cell is enough» rule gives back — listed once, restored on one tap.

The operator's ruling of 2026-09-22 (`leader_auto`, «ONE CELL IS ENOUGH»): a
leader who owns several cells passes #1 / #8 / #9 when ANY ONE of their cells
meets the task. Every verdict taken from `leader_auto.AUTO_FROM` until that
shipped was taken under the old reading — every cell for #1, the leader's
combined percentage for #9, and on a unit with one checklist per cell each
checklist judged on its own cell for all three — so a leader could lose a point
while one of their cells was filled on time. This module finds those verdicts
and gives the point back, in the two steps the operator chose:

1. `send` — a boot one-shot (`startup.report_auto_check_restore`) that DMs the
   operator a summary, an Excel list and ONE button. The list is STORED
   (`LIST_KEY`) exactly as it was sent.
2. `apply` — the button (`telegram_bot._acr_callback`), all or nothing: every
   row of the stored list gets its task back through the ordinary admin overlay
   (`LeaderTaskOverride`, done=True — reversible on the day report like any
   ruling), and each changed day re-sends its corrected report
   (`leader_reports.resend_if_changed`). The tap restores what the operator
   read, never a recomputation that a number typed since could have moved.

ONLY MULTI-CELL CASES (the operator's scope). A verdict qualifies only when the
leader had more than one cell and the new rule, applied to what stood AT THE
HOUR, passes. «Started late» on a unit with one checklist per leader, platform
faults and every other verdict stand.

ONLY PROVEN ON TIME (the operator's ruling). Nothing typed after the hour
counts. The evidence, strongest first:

* the check's OWN record at the hour (`leader_auto_checks`) — on a per-cell
  unit a sibling cell checklist that PASSED; for #1 which cells had their people
  typed (`facts.untyped` against `facts.cells`);
* concerns, whose `created_at` is exact — the check itself re-run at its hour;
* the action register's time-stamped saves (`action_logs`) — «Odamlar soni»,
  ПЛАН/ФАКТ typed on a position, SAP фаза uploads — replayed up to the hour.
  The replay is the one `auto_check_report` built for the 20–21 Sep report,
  COPIED here so deleting either one-shot cannot break the other;
* a stored row untouched since before the hour (`updated_at`) is its own proof.

A case the new rule might pass but none of that can place before the hour is
listed as NOT restored with the reason, and it never rides the button. A plan
that came in the SAP file is taken as the file stated it once an upload had
reached the unit before the hour — a SAP plan is not something a leader types
on time or late.

Temporary: delete this module, `startup.report_auto_check_restore`,
`startup._auto_check_restore_job`, the call in BOTH entrypoints and the `acr:`
callback in `telegram_bot.py` once the operator has pressed (or declined) the
button. A call left behind imports a deleted module at boot, and a failed boot
rolls the deploy back.
"""
from __future__ import annotations

import hashlib
import io
import json
import logging
import time
from collections import defaultdict
from datetime import date as _date, datetime, timezone

import requests
from sqlalchemy.orm import Session

from app.config import settings
from app.models import (
    ActionLog, AppSetting, Cell, LeaderAutoCheck, LeaderTaskDef, LeaderTaskEntry,
    LeaderTaskOverride, Manager, PPDaily, PPLineDaily, PPManagerSetting, RoleProfile,
)
from app.services import (
    cell_lookup, latin_code, leader_auto, leader_exclusions, leader_tasks,
    zagruzka_source,
)

log = logging.getLogger(__name__)

TZ = leader_auto.TASHKENT
LIST_KEY = "auto_check_restore_list_2026_09_22"
CALLBACK = "acr:apply"
SET_BY = "bitta yacheyka yetarli"     # what the overlay's «set by» says after the name
_API = "https://api.telegram.org"
SEND_RETRIES = 3

RESTORE, REFUSE = "restore", "refuse"
PASSED, FAILED = leader_auto.PASSED, leader_auto.FAILED


# ── small helpers ────────────────────────────────────────────────────────────

def _aware(v: datetime | None) -> datetime | None:
    if v is None:
        return None
    return v if v.tzinfo else v.replace(tzinfo=timezone.utc)


def _hm(v: datetime | None) -> str:
    """«21.09 09:47», plant wall clock."""
    v = _aware(v)
    return v.astimezone(TZ).strftime("%d.%m %H:%M") if v else "—"


def _d(iso: str) -> _date:
    return datetime.strptime(str(iso)[:10], "%Y-%m-%d").date()


def _norm(code) -> str:
    """Latin twins first (a code logged before 11 Sep may be Cyrillic), then
    `cell_lookup.norm_code` — the spelling `auto_check_report` matches on."""
    return cell_lookup.norm_code(latin_code.latin_code(str(code or "")) if code else code) or ""


def _details(r: ActionLog) -> dict:
    out = {}
    for pair in (r.details or []):
        try:
            out[str(pair[0])] = pair[1]
        except Exception:
            continue
    return out


def _people_of(v):
    """A whole-centre staffing change is logged «people/shtatka»; a group one is
    the bare number. Both come back as the people figure or None."""
    if v is None:
        return None
    if isinstance(v, (int, float)):
        return float(v)
    head = str(v).split("/", 1)[0].strip()
    if head in ("", "None", "null"):
        return None
    try:
        return float(head)
    except ValueError:
        return None


def _num(v):
    try:
        return float(v) if v is not None else None
    except (TypeError, ValueError):
        return None


# ── the unit-date replay (copied from auto_check_report._Unit) ──────────────

class _Unit:
    """One unit on one date: every logged save, replayable to any instant."""

    def __init__(self, db: Session, m: Manager, date: str):
        self.db, self.m, self.date, self.d = db, m, date, _d(date)
        st = db.query(PPManagerSetting).filter_by(manager_id=m.id).first()
        self.auto_fill = True if st is None else bool(st.auto_fill)
        self.cells = db.query(Cell).filter(Cell.manager_id == m.id).all()
        self._events = self._uploads = self._sapwc = None
        self._pp = self._pl = self._replay_ok = None

    def events(self) -> dict:
        if self._events is not None:
            return self._events
        rows = (self.db.query(ActionLog)
                .filter(ActionLog.unit_id == self.m.id, ActionLog.day == self.d,
                        ActionLog.outcome == "done",
                        ActionLog.action.in_((
                            "production.staffing_set", "production.wc_override_set",
                            "production.override_set")))
                .order_by(ActionLog.created_at, ActionLog.id).all())
        pins, plans, facts = [], [], []
        for r in rows:
            at = _aware(r.created_at)
            if r.action == "production.staffing_set":
                for ch in (r.changes or []):
                    try:
                        name, _old, new = ch[0], ch[1], ch[2]
                    except Exception:
                        continue
                    if name == "minutes":
                        continue
                    name = str(name)
                    code, grp = (name.split(" · ", 1) + [None])[:2] if " · " in name \
                        else (name, None)
                    pins.append({"at": at, "wc": _norm(code.strip()),
                                 "grp": (grp.strip() or None) if grp else None,
                                 "people": _people_of(new)})
            elif r.action == "production.wc_override_set":
                parts = str(r.target_id or "").split(":")
                if len(parts) < 3:
                    continue
                grp = parts[3] if len(parts) > 3 and parts[3] else None
                for ch in (r.changes or []):
                    try:
                        name, _old, new = ch[0], ch[1], ch[2]
                    except Exception:
                        continue
                    if name == "workers":
                        pins.append({"at": at, "wc": _norm(parts[2].strip()),
                                     "grp": grp, "people": _people_of(new)})
            else:                       # production.override_set — ПЛАН / ФАКТ
                det = _details(r)
                for ch in (r.changes or []):
                    try:
                        field, _old, new = ch[0], ch[1], ch[2]
                    except Exception:
                        continue
                    ev = {"at": at, "wc": _norm(det.get("work_center")),
                          "sap": str(det.get("sap_code") or ""),
                          "line": det.get("line") or None, "new": _num(new)}
                    (plans if field == "plan" else facts).append(ev)
        self._events = {"pins": pins, "plans": plans, "facts": facts}
        return self._events

    def uploads(self) -> list[dict]:
        if self._uploads is None:
            out = []
            for r in (self.db.query(ActionLog)
                      .filter(ActionLog.action == "production.phase_uploaded",
                              ActionLog.day == self.d, ActionLog.outcome == "done")
                      .order_by(ActionLog.created_at).all()):
                det = _details(r)
                out.append({"at": _aware(r.created_at), "mode": det.get("mode"),
                            "type": det.get("type") or "", "note": det.get("note")})
            self._uploads = out
        return self._uploads

    def sap_plan_wcs(self) -> set:
        """Work centres the stored SAP snapshot gives a plan on this date."""
        if self._sapwc is None:
            self._sapwc = {_norm(w) for (w,) in self.db.query(PPDaily.work_center).filter(
                PPDaily.manager_id == self.m.id, PPDaily.date == self.d,
                PPDaily.plan_qty > 0).all()}
        return self._sapwc

    def reaching_faza(self, before: datetime | None = None) -> list[dict]:
        """SAP фаза uploads of this date that WROTE this unit. An «auto-fill»
        upload reaches it while its switch is on; a «picked» one records no
        targets, so the stored snapshot is the only witness."""
        got_sap = bool(self.sap_plan_wcs())
        out = []
        for x in self.uploads():
            if "faza" not in x["type"]:
                continue
            reached = self.auto_fill if x["note"] == "auto-fill" else got_sap
            if reached and (before is None or x["at"] <= before):
                out.append(x)
        return out

    def typed_plan_at(self, t: datetime, codes: set) -> dict:
        """{(sap, wc, line): value} of the hand-typed ПЛАН standing at `t`."""
        state: dict = {}
        evs = [("typed", e["at"], e) for e in self.events()["plans"] if e["wc"] in codes]
        evs += [("upload", x["at"], x) for x in self.reaching_faza()
                if x["mode"] in ("both", "plan")]
        for kind, at, e in sorted(evs, key=lambda z: z[1]):
            if at > t:
                break
            if kind == "upload":
                state = {}
                continue
            key = (e["sap"], e["wc"], e["line"])
            if e["new"]:
                state[key] = e["new"]
            else:
                state.pop(key, None)
        return state

    def facts_at(self, t: datetime) -> tuple[dict, dict]:
        """The hand-typed ФАКТ standing at `t`: `(per line, per group)`.

        A line write sets that line; a group write (no line — an older tab)
        sets EVERY line of the group, which is why it forgets the lines' own
        earlier values; an upload that restates ФАКТ clears everything. The
        last is broader than the server (it spares a line kept by hand), which
        only ever makes this a LOWER bound."""
        lines: dict = {}
        groups: dict = {}
        evs = [("typed", e["at"], e) for e in self.events()["facts"]]
        evs += [("upload", x["at"], x) for x in self.reaching_faza()
                if x["mode"] in ("both", "actual")]
        for kind, at, e in sorted(evs, key=lambda z: z[1]):
            if at > t:
                break
            if kind == "upload":
                lines.clear()
                groups.clear()
                continue
            g = (e["sap"], e["wc"])
            if e["line"]:
                k = g + (str(e["line"]),)
                if e["new"] is None:
                    lines.pop(k, None)
                else:
                    lines[k] = e["new"]
            else:
                for k in [k for k in lines if k[:2] == g]:
                    lines.pop(k)
                if e["new"] is None:
                    groups.pop(g, None)
                else:
                    groups[g] = e["new"]
        return lines, groups

    def pins_at(self, t: datetime) -> dict:
        """The typed pins as the log says they stood at `t` — the
        `zagruzka_source.typed_pins` shape."""
        state: dict = {}
        for ev in self.events()["pins"]:
            if ev["at"] > t:
                continue
            key = (self.m.id, self.date, ev["wc"], ev["grp"])
            if ev["people"] is None:
                state.pop(key, None)
            else:
                state[key] = ev["people"]
        return state

    def replay_ok(self) -> bool:
        """Does the replay, run to the end, reproduce the pins stored now? When
        it does not, something wrote pins outside the logged endpoints and the
        at-the-hour picture must not be trusted."""
        if self._replay_ok is None:
            now = zagruzka_source.typed_pins(self.db, [self.m.id], self.d, self.d)

            def canon(p):
                return {(_norm(k[2]), k[3] or None): float(v) for k, v in p.items()
                        if str(k[1])[:10] == self.date}
            end = datetime(2100, 1, 1, tzinfo=timezone.utc)
            self._replay_ok = canon(self.pins_at(end)) == canon(now)
        return self._replay_ok

    def cell_value(self, pins: dict, cell: Cell):
        got = zagruzka_source.cell_pins(self.cells, pins).get((cell.id, self.date))
        return None if got is None else got[0]

    def pp_rows(self) -> dict:
        if self._pp is None:
            self._pp = {(str(r.sap_code or ""), _norm(r.work_center)): r
                        for r in self.db.query(PPDaily).filter(
                            PPDaily.manager_id == self.m.id, PPDaily.date == self.d).all()}
        return self._pp

    def line_rows(self) -> dict:
        if self._pl is None:
            self._pl = {(str(r.qty_key or ""), _norm(r.work_center), str(r.line_key or "")): r
                        for r in self.db.query(PPLineDaily).filter(
                            PPLineDaily.manager_id == self.m.id,
                            PPLineDaily.date == self.d).all()}
        return self._pl


# ── what stood at the hour ───────────────────────────────────────────────────

def _plan_at_hour(u: _Unit, rows_today: list, cell: Cell, due: datetime):
    """Did this cell carry a plan AT the hour? `(True|False|None, why)`.

    False only where there is no plan on the cell even now — the cell was never
    asked for anything that day. None is «cannot be placed before the hour»."""
    code = _norm(cell.sap_code)
    if not code:
        return False, "no SAP code"
    if not leader_auto.cell_planned(rows_today, cell):
        return False, "no plan on this cell"
    if u.typed_plan_at(due, {code}):
        return True, "plan typed before the hour"
    reach = u.reaching_faza(before=due)
    if reach and code in u.sap_plan_wcs():
        return True, f"SAP plan, uploaded {_hm(reach[0]['at'])}"
    return None, "its plan cannot be placed before the hour"


def _fact_lb(u: _Unit, rows: list, due: datetime) -> float:
    """ФАКТ minutes that PROVABLY stood on these positions at the hour.

    Per line, the page's own precedence (the line's typed value → the group's
    typed value → the SAP snapshot, the last only where the line takes the
    file), each level read either off a stored row nobody has written since the
    hour or off the replayed saves. Anything the record cannot place counts 0,
    so this is a floor, never an estimate."""
    lines, groups = u.facts_at(due)
    pp, pl = u.pp_rows(), u.line_rows()
    total = 0.0
    for r in rows:
        labor = _num(r.get("labor_time"))
        if not labor:
            continue
        g = (str(r.get("qty_key") or ""), _norm(r.get("work_center")))
        k = g + (str(r.get("line_key") or ""),)
        val = None
        lrow = pl.get(k)
        if lrow is not None and _aware(lrow.updated_at) and _aware(lrow.updated_at) <= due:
            val = _num(lrow.actual_override)
        elif k in lines:
            val = lines[k]
        if val is None:
            grow = pp.get(g)
            if grow is not None and _aware(grow.updated_at) and _aware(grow.updated_at) <= due:
                if grow.actual_override is not None:
                    val = _num(grow.actual_override)
                elif r.get("sap_filled"):
                    val = _num(grow.actual_qty)
            elif g in groups:
                val = groups[g]
        if val:
            total += labor * val / 60.0
    return total


class _Code:
    """One work centre's page for the day — the rows and the «Bajarish %» tile
    `leader_auto.code_totals` reads, uncut by group."""

    def __init__(self, db: Session, mid: int, date: str, code: str):
        from app.routers.production import _build_dashboard
        dash = _build_dashboard(db, mid, _d(date), wc_scope={code}, payload=None)
        t = dash.get("totals") or {}
        self.rows = dash.get("rows") or []
        self.plan = float(t.get("total_plan_labor") or 0)
        self.fact = float(t.get("total_actual_labor") or 0)
        self.pct = round(float(t.get("completion") or 0) * 100, 1)


# ── judging one leader-date-task ─────────────────────────────────────────────

class _Walk:
    """Caches shared by every group of one pass."""

    def __init__(self, db: Session, now: datetime):
        self.db, self.now = db, now
        self._units: dict = {}
        self._codes: dict = {}
        self._cells: dict = {}

    def unit(self, m: Manager, date: str) -> _Unit:
        key = (m.id, date)
        if key not in self._units:
            self._units[key] = _Unit(self.db, m, date)
        return self._units[key]

    def code(self, mid: int, date: str, code: str) -> _Code:
        key = (mid, date, code)
        if key not in self._codes:
            self._codes[key] = _Code(self.db, mid, date, code)
        return self._codes[key]

    def cell(self, cell_id: int | None) -> Cell | None:
        return (self.db.query(Cell).filter(Cell.id == cell_id).first()
                if cell_id else None)

    def cells(self, leader_id: int) -> list[Cell]:
        if leader_id not in self._cells:
            self._cells[leader_id] = (self.db.query(Cell)
                                      .filter(Cell.leader_id == leader_id)
                                      .order_by(Cell.verifix_code, Cell.id).all())
        return self._cells[leader_id]

    def leader_rows(self, prof, m, date, due) -> list:
        ctx = leader_auto._Ctx(self.db, prof, m, m.shift, date, None, due, self.now)
        return (ctx.dashboard.get("rows") or []) if ctx.pairs else []


def _label(c: Cell | None) -> str:
    return leader_auto.cell_label(c) if c is not None else "—"


def _judge(w: _Walk, prof, m, check: str, date: str, due: datetime,
           sibs: list, failed: list, target) -> list[tuple]:
    """`[(ledger row, RESTORE|REFUSE, text)]` for one (leader, date, task).
    Empty when the leader had one cell, or the new rule fails too."""
    db, hh = w.db, due.astimezone(TZ).strftime("%H:%M")
    if any(s.cell_id is not None for s in sibs):
        return _judge_per_cell(w, prof, m, check, date, due, hh, sibs, failed, target)

    f = failed[0]
    facts = f.facts or {}
    cells = w.cells(prof.id)
    if check == "plan_staffing":
        if f.code != "no_staffing" or int(facts.get("cells") or 0) < 2:
            return []
        untyped = {str(x) for x in (facts.get("untyped") or [])}
        labels = [_label(c) for c in cells]
        if len(cells) != int(facts.get("cells") or 0) or not untyped <= set(labels):
            return [(f, REFUSE, f"the leader's cells changed since the {hh} check "
                                f"({facts.get('cells')} then, {len(cells)} now)")]
        typed = [c for c in cells if _label(c) not in untyped]
        if not typed:
            return []
        u = w.unit(m, date)
        rows = w.leader_rows(prof, m, date, due)
        unknown = []
        for c in typed:
            ok, why = _plan_at_hour(u, rows, c, due)
            if ok:
                return [(f, RESTORE, f"{_label(c)}: people were typed at the {hh} check "
                                     f"(the check's own record) · {why}")]
            if ok is None:
                unknown.append(f"{_label(c)}: people typed at the hour, but {why}")
        return [(f, REFUSE, "; ".join(unknown))] if unknown else []

    if check == "plan_pct":
        if f.code != "under_target":
            return []
        codes = sorted({_norm(c.sap_code) for c in cells if _norm(c.sap_code)})
        if len(codes) < 2:
            return []
        return _judge_pct(w, prof, m, date, due, hh, cells, codes, f, target,
                          plan_then=_num(facts.get("plan_min")),
                          fact_then=_num(facts.get("fact_min")))
    return []                           # #8 already counted every cell


def _judge_pct(w: _Walk, prof, m, date, due, hh, cells, codes, f, target,
               plan_then, fact_then) -> list[tuple]:
    want = float(target if target is not None else 30)
    per = {c: w.code(m.id, date, c) for c in codes}
    names = defaultdict(list)
    for c in cells:
        if _norm(c.sap_code):
            names[_norm(c.sap_code)].append(_label(c))
    planned = {c: p for c, p in per.items() if p.plan > 0}
    if not planned:
        return []
    if plan_then is not None:
        plan_now = sum(p.plan for p in per.values())
        if abs(plan_now - plan_then) > max(1.0, 0.005 * plan_then):
            return [(f, REFUSE, f"the plan changed after the {hh} check "
                                f"({plan_then:g} min then, {plan_now:.1f} now) — the "
                                f"per-cell % at the hour cannot be rebuilt")]
    if fact_then is not None:
        # Every minute of ФАКТ the check saw, put on ONE work centre: if even
        # that misses the mark, no cell of theirs reached it at the hour.
        ub = max(fact_then / p.plan * 100 for p in planned.values())
        if ub < want:
            return []
    elif max(p.pct for p in planned.values()) < want:
        return []                       # not even today's figure reaches it
    u = w.unit(m, date)
    lb = {c: _fact_lb(u, p.rows, due) for c, p in planned.items()}
    if fact_then is not None and sum(lb.values()) > fact_then + max(1.0, 0.005 * fact_then):
        return [(f, REFUSE, f"the saves rebuild more ФАКТ ({sum(lb.values()):.1f} min) than "
                            f"the {hh} check saw ({fact_then:g}) — not trusted")]
    best = max(planned, key=lambda c: lb[c] / planned[c].plan)
    pct = round(lb[best] / planned[best].plan * 100, 1)
    who = ", ".join(sorted(names.get(best) or [best]))
    if pct >= want:
        return [(f, RESTORE, f"{who}: {pct}% at {hh} — {lb[best]:.1f} of "
                             f"{planned[best].plan:.1f} plan-min provably in before the hour")]
    return [(f, REFUSE, f"{who} may have reached {want:g}% at {hh}, but only {pct}% of its "
                        f"ФАКТ can be shown to have been in before the hour")]


def _judge_per_cell(w: _Walk, prof, m, check, date, due, hh, sibs, failed,
                    target) -> list[tuple]:
    """A unit with one checklist per cell. The check was taken per cell, so its
    own record answers most of it: one cell checklist that PASSED at the hour
    is one cell that met the task."""
    if len({s.cell_id for s in sibs if s.cell_id is not None}) < 2:
        return []
    measured = [s for s in sibs if s.cell_id is not None
                and s.outcome in (PASSED, FAILED) and s.code != "started_late"]
    if not measured:
        return []                       # no checklist of theirs existed at the hour
    passed = [s for s in measured if s.outcome == PASSED]
    if passed:
        s = passed[0]
        why = (f"cell {_label(w.cell(s.cell_id))} passed the {hh} check "
               f"(the check's own record)")
        return [(f, RESTORE, why) for f in failed]
    if check == "concerns":
        # Every cell of theirs, at the hour — the check itself, re-run with its
        # own clock. `created_at` cannot move, so this is exact.
        ctx = leader_auto._Ctx(w.db, prof, m, m.shift, date, None, due, w.now)
        v = leader_auto._RUNNERS["concerns"](ctx, None)
        if v.done:
            why = f"{v.facts.get('found')} concern(s) created before {hh} on the leader's cells"
            return [(f, RESTORE, why) for f in failed]
        return []
    unmeasured = [w.cell(s.cell_id) for s in sibs if s.code == "started_late"]
    unmeasured = [c for c in unmeasured if c is not None]
    if not unmeasured:
        return []                       # every cell measured at the hour; none met it
    if check == "plan_staffing":
        u = w.unit(m, date)
        if not u.replay_ok():
            return [(f, REFUSE, f"{', '.join(_label(c) for c in unmeasured)} was not checked "
                                f"at {hh} and its headcount saves cannot be replayed") for f in failed]
        pins, rows = u.pins_at(due), w.leader_rows(prof, m, date, due)
        unknown = []
        for c in unmeasured:
            v = u.cell_value(pins, c)
            if v is None:
                continue
            ok, why = _plan_at_hour(u, rows, c, due)
            if ok:
                text = f"{_label(c)}: {v:g} people typed before {hh} (action log) · {why}"
                return [(f, RESTORE, text) for f in failed]
            if ok is None:
                unknown.append(f"{_label(c)}: people typed before {hh}, but {why}")
        return [(f, REFUSE, "; ".join(unknown)) for f in failed] if unknown else []
    if check == "plan_pct":
        codes = sorted({_norm(c.sap_code) for c in unmeasured if _norm(c.sap_code)})
        if not codes:
            return []
        out = _judge_pct(w, prof, m, date, due, hh, unmeasured, codes, failed[0],
                         target, plan_then=None, fact_then=None)
        return [(f, kind, why) for f in failed for (_r, kind, why) in out]
    return []


# ── the list ─────────────────────────────────────────────────────────────────

def collect(db: Session, now: datetime | None = None) -> dict:
    """Every failed automatic verdict of a multi-cell leader, judged again under
    «one cell is enough» on what stood at the hour. Reads only."""
    from app.services.leader_bot import day_uid
    now = now or datetime.now(timezone.utc)
    w = _Walk(db, now)
    defs = {td.id: td for td in db.query(LeaderTaskDef).all()}
    rows = (db.query(LeaderAutoCheck)
            .filter(LeaderAutoCheck.date >= leader_auto.AUTO_FROM,
                    LeaderAutoCheck.check.in_(leader_auto.CHECKS))
            .order_by(LeaderAutoCheck.date, LeaderAutoCheck.leader_id,
                      LeaderAutoCheck.task_id, LeaderAutoCheck.id).all())
    groups: dict = defaultdict(list)
    for r in rows:
        groups[(r.leader_id, r.date, r.task_id)].append(r)

    items, refused, errors = [], [], []
    for (lid, date, tid), sibs in groups.items():
        failed = [s for s in sibs if s.outcome == FAILED and s.entry_id is not None]
        if not failed:
            continue
        try:
            prof = db.query(RoleProfile).filter(RoleProfile.id == lid).first()
            m = db.query(Manager).filter(Manager.id == failed[0].manager_id).first()
            if prof is None or m is None:
                continue
            if leader_exclusions.excluded(db, prof.id, date, leader_name=prof.name):
                continue                # a day that counts for nothing
            td = defs.get(tid)
            parsed = leader_auto.parse_check(td.auto_check if td else None)
            target = parsed[1] if parsed else None
            due = _aware(failed[0].due_at)
            decided = _judge(w, prof, m, failed[0].check, date, due, sibs, failed, target)
            if not decided:
                continue
            weight = (leader_tasks.effective_leader_config(db, prof, m.shift, day=date)
                      .get(tid) or {}).get("weight")
            for r, kind, why in decided:
                entry = db.query(LeaderTaskEntry).filter(
                    LeaderTaskEntry.id == r.entry_id).first()
                if entry is None or entry.done:
                    continue
                uid = day_uid(entry.day_id)
                if db.query(LeaderTaskOverride.id).filter_by(uid=uid, task_id=tid).first():
                    continue            # an admin has already ruled on it
                rec = {"row_id": r.id, "uid": uid, "task_id": tid, "date": date,
                       "leader_id": prof.id, "leader": prof.name,
                       "manager_id": m.id, "unit": m.name, "shift": m.shift,
                       "task": (td.name_uz if td else f"#{tid}") or f"#{tid}",
                       "check": r.check,
                       "cell": _label(w.cell(r.cell_id)) if r.cell_id else "—",
                       "hour": due.astimezone(TZ).strftime("%H:%M") if due else "—",
                       "code": r.code, "weight": _num(weight) or 0.0, "why": why}
                (items if kind == RESTORE else refused).append(rec)
        except Exception as exc:        # noqa: BLE001 - one group never sinks the list
            db.rollback()
            log.exception("auto-check restore: %s %s #%s", lid, date, tid)
            errors.append(f"leader {lid} {date} #{tid}: {str(exc)[:160]}")
    return {"built_at": now.isoformat(), "items": items, "refused": refused,
            "errors": errors, "groups": len(groups)}


def list_id(items: list[dict]) -> str:
    blob = json.dumps([(i["uid"], i["task_id"]) for i in items], sort_keys=True)
    return hashlib.sha1(blob.encode()).hexdigest()[:10]


def summary_text(rep: dict) -> str:
    items, refused = rep["items"], rep["refused"]
    by_task = defaultdict(int)
    for i in items:
        by_task[i["task_id"]] += 1
    pts = sum(i["weight"] for i in items)
    leaders = len({i["leader_id"] for i in items})
    lines = [
        "«One cell is enough» — points to give back "
        f"({leader_auto.AUTO_FROM[8:10]}.{leader_auto.AUTO_FROM[5:7]} → now).",
        "",
        f"Give back: {len(items)} task(s) to {leaders} leader(s), {pts:g} points in total"
        + (" — " + ", ".join(f"#{t}: {n}" for t, n in sorted(by_task.items())) if items else "")
        + ".",
        f"Not restored — the new rule might pass, but it cannot be shown on time: "
        f"{len(refused)}.",
    ]
    if rep.get("errors"):
        lines.append(f"⚠ {len(rep['errors'])} case(s) could not be computed — sheet «Errors».")
    lines += ["", "Only leaders with two or more cells; only what stood at the check hour. "
                  "The Excel lists each one with its evidence."]
    return "\n".join(lines)


def build_workbook(rep: dict) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font
    wb = Workbook()
    cols = [("Date", "date", 11), ("Shift", "shift", 6), ("Brigadir", "unit", 26),
            ("Leader", "leader", 26), ("Task", "task_id", 6), ("Task name", "task", 30),
            ("Cell", "cell", 8), ("Check hour", "hour", 10), ("Verdict then", "code", 14),
            ("Points", "weight", 7)]
    for i, (title, rows, last) in enumerate((
            ("Restore", rep["items"], "Evidence (before the hour)"),
            ("Not restored", rep["refused"], "Why not"))):
        ws = wb.active if i == 0 else wb.create_sheet()
        ws.title = title
        head = [c[0] for c in cols] + [last]
        ws.append(head)
        for rec in rows:
            ws.append([rec.get(k) for _t, k, _w in cols] + [rec.get("why")])
        for j, (_t, _k, width) in enumerate(cols, start=1):
            ws.column_dimensions[ws.cell(1, j).column_letter].width = width
        ws.column_dimensions[ws.cell(1, len(head)).column_letter].width = 90
        for c in ws[1]:
            c.font = Font(bold=True)
        for row in ws.iter_rows(min_row=2):
            row[-1].alignment = Alignment(wrap_text=True, vertical="top")
        ws.freeze_panes = "A2"
    if rep.get("errors"):
        ws = wb.create_sheet("Errors")
        ws.append(["Error"])
        for e in rep["errors"]:
            ws.append([e])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ── delivery ─────────────────────────────────────────────────────────────────

def _clean(text: str) -> str:
    token = settings.telegram_bot_token or ""
    return text.replace(token, "***") if token else text


def _post(method: str, data: dict, files: dict | None = None) -> None:
    """One Bot API call, retried; the token (which rides in the URL) never
    reaches an exception text or the journal."""
    last = ""
    for attempt in range(1, SEND_RETRIES + 1):
        wait = 0
        try:
            r = requests.post(f"{_API}/bot{settings.telegram_bot_token}/{method}",
                              data=data, files=files, timeout=120)
            body = r.json()
            if body.get("ok"):
                return
            last = body.get("description") or f"HTTP {r.status_code}"
            wait = int(((body.get("parameters") or {}).get("retry_after")) or 0)
        except Exception as exc:
            last = f"{type(exc).__name__}: {exc}"
        if attempt < SEND_RETRIES:
            time.sleep(max(wait, 3 * attempt))
    raise RuntimeError(_clean(last or f"{method} failed")[:300])


def send(db: Session, chat_id: int, *_window) -> int:
    """Step 1: build the list, STORE it, DM it with the button. Returns how many
    messages went out (the `_send_report_once` contract)."""
    if not settings.telegram_bot_token:
        raise RuntimeError("telegram bot token not configured")
    rep = collect(db)
    db.rollback()
    items = rep["items"]
    lid = list_id(items)
    row = db.query(AppSetting).filter_by(key=LIST_KEY).first()
    stored = json.loads(row.value) if row and row.value else {}
    if stored.get("applied"):
        raise RuntimeError("the restore list was already applied — not rebuilt")
    value = json.dumps({"list_id": lid, "built_at": rep["built_at"],
                        "items": items, "applied": None}, ensure_ascii=False)
    if row is None:
        db.add(AppSetting(key=LIST_KEY, value=value))
    else:
        row.value = value
    db.commit()

    _post("sendMessage", {"chat_id": chat_id, "text": summary_text(rep)[:4000]})
    _post("sendDocument", {"chat_id": chat_id,
                           "caption": "«One cell is enough» — the list, with the evidence "
                                      "for each point."},
          files={"document": ("one-cell-restore.xlsx", build_workbook(rep),
                              "application/vnd.openxmlformats-officedocument."
                              "spreadsheetml.sheet")})
    if not items:
        return 2
    pts = sum(i["weight"] for i in items)
    text = (f"Tap to give back the {len(items)} task(s) on the «Restore» sheet "
            f"({pts:g} points). Each becomes an admin ruling «done» on that task — "
            f"reversible on the day report — and each changed day re-sends its "
            f"corrected report to the leader.")
    markup = {"inline_keyboard": [[{"text": f"✅ Give back {len(items)} ({pts:g} pts)",
                                    "callback_data": f"{CALLBACK}:{lid}"}]]}
    _post("sendMessage", {"chat_id": chat_id, "text": text,
                          "reply_markup": json.dumps(markup)})
    return 3


# ── step 2: the button ───────────────────────────────────────────────────────

def apply(db: Session, *, actor_tid: int, actor_name: str, lid: str) -> dict:
    """Give back every task on the STORED list — all or nothing, once.

    A task an admin has ruled on since the list was built is left alone: their
    ruling is the newer statement about it. Nothing is written until every row
    is ready, and the corrected reports go out only after the commit."""
    from app.services import action_log, leader_reports
    row = (db.query(AppSetting).filter_by(key=LIST_KEY)
           .with_for_update().first())
    if row is None or not row.value:
        return {"status": "missing"}
    data = json.loads(row.value)
    if data.get("list_id") != lid:
        return {"status": "stale"}
    if data.get("applied"):
        return {"status": "already", **data["applied"]}
    at = datetime.now(timezone.utc)
    who = f"{actor_name} · {SET_BY}"[:160]
    restored, skipped, uids = 0, 0, set()
    for it in data.get("items") or []:
        if db.query(LeaderTaskOverride.id).filter_by(
                uid=it["uid"], task_id=it["task_id"]).first():
            skipped += 1
            continue
        db.add(LeaderTaskOverride(uid=it["uid"], task_id=it["task_id"],
                                  date=it["date"], leader=(it.get("leader") or "")[:160] or None,
                                  done=True, set_by=who, set_at=at))
        restored += 1
        uids.add(it["uid"])
    data["applied"] = {"at": at.astimezone(TZ).strftime("%d.%m %H:%M"), "by": actor_name,
                       "restored": restored, "skipped": skipped}
    row.value = json.dumps(data, ensure_ascii=False)
    db.commit()

    resent = 0
    for uid in sorted(uids):
        try:
            if leader_reports.resend_if_changed(db, uid):
                resent += 1
        except Exception:               # noqa: BLE001 - a DM never undoes the ruling
            db.rollback()
            log.exception("auto-check restore: corrected report for %s failed", uid)
    action_log.record_bot(
        db, actor_tid, "leader_review", "checklist.auto_points_restored",
        actor_name=actor_name, target_kind="task", target_id=f"batch:{restored}",
        details=[("count", restored), ("skipped", skipped or None),
                 ("reports", resent or None), ("list", lid)])
    return {"status": "done", "restored": restored, "skipped": skipped,
            "resent": resent}


def result_text(res: dict) -> str:
    st = res.get("status")
    if st == "done":
        return (f"✅ Given back: {res['restored']} task(s)"
                + (f", {res['skipped']} left alone (an admin had ruled since)"
                   if res.get("skipped") else "")
                + f". Corrected reports sent: {res.get('resent', 0)}.")
    if st == "already":
        return (f"Already given back on {res.get('at')} by {res.get('by')} — "
                f"{res.get('restored')} task(s).")
    if st == "stale":
        return "This button belongs to an older list — nothing was changed."
    return "The list is gone — nothing was changed."
