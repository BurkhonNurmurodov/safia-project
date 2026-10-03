"""EVERYTHING the Verifix API returns for one day, as a workbook — a one-time
diagnostic (the operator's request, 2026-10-02).

The live «Verifix to'g'irlash» printed every arrival early and every departure
late against the Excel the same API feeds, and the fix (v4.201.1) was made from
the two exports alone, without the raw rows behind them. This dumps those rows:
every form the API user may read, every field each one returns, for
2026-10-01 — and, beside the day's uploaded file, what the report, the marks
and the live page say about each person in it, so the next fix is made on the
data and not on a guess about it.

TEMPORARY: delete together with `startup.report_verifix_dump` and the call in
BOTH entrypoints once the file has landed. Nothing is stored and nothing is
written to Verifix. Names, ids and clocks DO leave this module — that is the
point — into the operator's own chat only.
"""
from __future__ import annotations

import json
import logging
import time
from collections import Counter, defaultdict
from datetime import date, datetime, time as dtime, timedelta
from io import BytesIO
from typing import Any, Optional

import requests
from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
from sqlalchemy.orm import Session

from app.config import settings
from app.models import AttendanceBatch, AttendanceBatchRow
from app.services import verifix, verifix_live
from app.services.verifix_parity import _file_clock, _keys

log = logging.getLogger(__name__)

DAY = date(2026, 10, 1)
# Marks are read into the next afternoon: a night shift ends the morning after.
TRACKS_END = dtime(14, 0)
# A scheduler job, not a request — minutes are fine; Cloudflare is not in the way.
BUDGET_S = 20 * 60
PAGE_CAP = 400
CELL_MAX = 32000            # an Excel cell holds 32,767 characters
SPLIT_BYTES = 45 * 1024 * 1024   # sendDocument refuses above 50 MB
MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

# Every form the read-only role opens, with the widest body each accepts.
FORMS: list[tuple[str, str, dict, int]] = [
    ("divisions", "core/division$list", {"division_ids": []}, verifix.LIMIT_LIST),
    ("employees", "core/employee$list",
     {"employee_ids": [], "statuses": [], "npins": []}, verifix.LIMIT_LIST),
    ("jobs", "core/job$list", {"job_ids": []}, verifix.LIMIT_LIST),
    ("schedules", "core/schedule$list", {"schedule_ids": []}, verifix.LIMIT_LIST),
    ("time_kinds", "core/time_kind$list", {"time_kind_ids": []}, verifix.LIMIT_LIST),
    ("locations", "core/location$list", {"location_ids": []}, verifix.LIMIT_LIST),
    ("timesheet", "core/timesheet$export",
     {"period_begin_date": verifix._dmy(DAY), "period_end_date": verifix._dmy(DAY),
      "division_ids": [], "employee_ids": []}, verifix.LIMIT_TIMESHEET),
    ("tracks", "core/track$list",
     {"begin_datetime": datetime.combine(DAY, dtime(0)).strftime("%d.%m.%Y %H:%M:%S"),
      "end_datetime": datetime.combine(DAY + timedelta(days=1), TRACKS_END)
      .strftime("%d.%m.%Y %H:%M:%S")}, verifix.LIMIT_TRACKS),
]


# ── fetching ──────────────────────────────────────────────────────────────────

def _pages(cl, path: str, body: dict, limit: int, deadline: float) -> tuple[list, dict]:
    """Every page of one form. Never raises: the status says what stopped it."""
    rows: list = []
    pages, cursor = 0, None
    status: dict[str, Any] = {"ok": True}
    try:
        while True:
            if time.monotonic() > deadline:
                status = {"ok": False, "code": "slow", "message": "time limit"}
                break
            data, nxt = verifix.call(cl, path, body, limit=limit, cursor=cursor)
            pages += 1
            rows.extend(data if isinstance(data, list) else [data])
            if not nxt or nxt == cursor:
                break
            if pages >= PAGE_CAP:
                status = {"ok": False, "code": "slow", "message": "page cap"}
                break
            cursor = nxt
    except verifix.VerifixError as exc:
        status = {"ok": False, "code": exc.code, "message": exc.message, "status": exc.status}
    except Exception as exc:                       # pragma: no cover — logged
        log.exception("verifix dump: %s failed", path)
        status = {"ok": False, "code": "error", "message": type(exc).__name__}
    status.update(pages=pages, rows=len(rows))
    return rows, status


def collect(cfg: dict, deadline: float) -> dict:
    """Every form, page by page. A fatal answer (bad login, wrong host) stops
    the rest — every further call would fail the same way."""
    out: dict[str, Any] = {"forms": {}, "meta": [], "at": datetime.now(verifix.TZ)}
    stopped: Optional[str] = None
    with verifix.client(cfg) as cl:
        for key, path, body, limit in FORMS:
            row: dict[str, Any] = {"form": key, "path": path,
                                   "body": json.dumps(body, ensure_ascii=False)}
            if stopped:
                row.update(ok=False, code="skipped", message=f"after {stopped}",
                           pages=0, rows=0, ms=0)
                out["forms"][key] = []
                out["meta"].append(row)
                continue
            t0 = time.monotonic()
            rows, st = _pages(cl, path, body, limit, deadline)
            row.update(st, ms=int((time.monotonic() - t0) * 1000))
            out["forms"][key] = rows
            out["meta"].append(row)
            if not st["ok"] and st.get("code") in verifix._FATAL:
                stopped = st["code"]
    return out


# ── flattening ────────────────────────────────────────────────────────────────

def _flat(obj: Any, prefix: str = "", out: Optional[dict] = None, *, skip: tuple = ()) -> dict:
    """A nested record as one row: `a.b` for a dict inside, a joined string for
    a list of scalars, JSON for a list of records. Nothing is dropped."""
    out = {} if out is None else out
    if not isinstance(obj, dict):
        out[prefix.rstrip(".") or "value"] = (json.dumps(obj, ensure_ascii=False)
                                              if isinstance(obj, list) else obj)
        return out
    for k, v in obj.items():
        if not prefix and k in skip:
            continue
        key = f"{prefix}{k}"
        if isinstance(v, dict):
            _flat(v, key + ".", out)
        elif isinstance(v, list):
            if all(not isinstance(x, (dict, list)) for x in v):
                out[key] = ", ".join("" if x is None else str(x) for x in v)
            else:
                out[key] = json.dumps(v, ensure_ascii=False)
        else:
            out[key] = v
    return out


def _cell(v: Any) -> Any:
    if v is None or isinstance(v, (bool, int, float)):
        return v
    if isinstance(v, (datetime, date)):
        return v
    s = v if isinstance(v, str) else json.dumps(v, ensure_ascii=False)
    s = ILLEGAL_CHARACTERS_RE.sub("", s)
    return s[:CELL_MAX]


def _sheet(wb: Workbook, title: str, rows: list[dict], first: tuple = ()) -> None:
    """One sheet per form: the columns are every key any row carries, in the
    order first seen; bold header, frozen, filtered."""
    ws = wb.create_sheet(title[:31])
    cols: list[str] = list(first)
    seen = set(cols)
    for r in rows:
        for k in r:
            if k not in seen:
                seen.add(k)
                cols.append(k)
    ws.append(cols or ["(bo'sh)"])
    for c in ws[1]:
        c.font = Font(bold=True)
    for r in rows:
        ws.append([_cell(r.get(k)) for k in cols])
    ws.freeze_panes = "A2"
    if rows and cols:
        ws.auto_filter.ref = f"A1:{get_column_letter(len(cols))}{len(rows) + 1}"
    for i, k in enumerate(cols, 1):
        w = max([len(str(k))] + [len(str(r.get(k) or "")) for r in rows[:200]])
        ws.column_dimensions[get_column_letter(i)].width = min(max(w + 2, 8), 50)


def _census(rows: list[dict]) -> list[dict]:
    """Which fields a form sends, how many rows fill each, and a sample."""
    filled: Counter[str] = Counter()
    sample: dict[str, Any] = {}
    order: list[str] = []
    for r in rows:
        for k, v in _flat(r).items():
            if k not in sample:
                order.append(k)
                sample[k] = None
            if v not in (None, "", [], {}):
                filled[k] += 1
                if sample[k] is None:
                    sample[k] = v
    return [{"field": k, "filled": filled[k], "of": len(rows), "sample": sample[k]}
            for k in order]


# ── the workbook ──────────────────────────────────────────────────────────────

def _hm(dt: Optional[datetime]) -> Optional[str]:
    return dt.strftime("%H:%M") if dt else None


def _stamp(dt: Optional[datetime]) -> Optional[str]:
    if not dt:
        return None
    return dt.strftime("%H:%M") if dt.date() == DAY else dt.strftime("%d.%m %H:%M")


def _mins(a: Optional[str], b: Optional[str], day_b: Optional[date] = None) -> Optional[int]:
    """b − a in minutes, both «HH:MM» on DAY (b may sit on the day after)."""
    if not a or not b:
        return None
    try:
        ta = datetime.combine(DAY, dtime(*(int(x) for x in a.split(":")[:2])))
        tb = datetime.combine(day_b or DAY, dtime(*(int(x) for x in b.split(":")[:2])))
    except (TypeError, ValueError):
        return None
    return int(round((tb - ta).total_seconds() / 60))


def build(db: Session, data: dict, *, tracks_apart: bool = False) -> tuple[BytesIO, Optional[BytesIO], dict]:
    """The workbook(s) and the summary the caption prints. With `tracks_apart`
    the marks go to a second file (one file must stay under Telegram's cap)."""
    forms = data["forms"]
    divs = {str(d.get("division_id") or ""): d for d in forms.get("divisions", [])}
    kinds = {str(k.get("time_kind_id") or ""): k for k in forms.get("time_kinds", [])}
    emps = {str(e.get("employee_id") or ""): e for e in forms.get("employees", [])}
    cells, units = verifix_live._registry(db)

    def ename(e: dict) -> str:
        return " ".join(str(x).strip() for x in (e.get("last_name"), e.get("first_name"),
                                                 e.get("middle_name")) if x).strip()

    def place(eid: str) -> dict:
        """Where WE would put this person: the cell their org unit's code names
        (the live page's rule) and the division beside it, for comparison."""
        e = emps.get(eid) or {}
        unit = str(e.get("org_unit_id") or "")
        div = str(e.get("division_id") or "")
        ucode = str((divs.get(unit) or {}).get("code") or "").strip()
        dcode = str((divs.get(div) or {}).get("code") or "").strip()
        cell = cells.get(verifix._code_key(ucode)) if ucode else None
        u = units.get(cell["manager_id"]) if cell else None
        return {"our_unit": (u or {}).get("name"), "our_cell": cell["code"] if cell else None,
                "org_unit_code": ucode or None, "org_unit_name": (divs.get(unit) or {}).get("name"),
                "division_code": dcode or None, "division_name": (divs.get(div) or {}).get("name")}

    # The report: one row per employee × day entry, facts as columns.
    report_rows: list[dict] = []
    report_by_eid: dict[str, dict] = {}
    for r in forms.get("timesheet", []):
        eid = str(r.get("employee_id") or "")
        base = _flat(r, skip=("days",))
        days = r.get("days") or []
        report_by_eid[eid] = {"row": r, "days": [d for d in days if verifix_live._d(d.get("date")) == DAY]}
        if not days:
            report_rows.append({**base, **place(eid), "day.date": None})
        for d in days:
            dflat = {f"day.{k}": v for k, v in _flat(d, skip=("facts",)).items()}
            facts: dict[str, float] = {}
            for f in d.get("facts") or []:
                k = str(f.get("time_kind_id") or "")
                try:
                    v = float(f.get("fact_value") or 0)
                except (TypeError, ValueError):
                    v = 0.0
                if k:
                    facts[k] = facts.get(k, 0.0) + v
            fcols = {f"fact:{k} {(kinds.get(k) or {}).get('name') or ''}".strip(): v
                     for k, v in facts.items()}
            report_rows.append({**base, **place(eid), **dflat, **fcols,
                                "day.facts_json": json.dumps(d.get("facts") or [], ensure_ascii=False)})

    # The marks: one row per track, in order.
    track_rows: list[dict] = []
    marks_by_eid: dict[str, list] = defaultdict(list)
    for t in forms.get("tracks", []):
        eid = str(t.get("employee_id") or "")
        at = verifix_live._dt(t.get("track_datetime"))
        typ = str(t.get("track_type") or "").upper()
        marks_by_eid[eid].append((at, typ, t))
        track_rows.append({"employee_id": eid, "employee_name": ename(emps.get(eid) or {}),
                           **place(eid), **_flat(t, skip=("employee_id",)), "_at": at})
    track_rows.sort(key=lambda r: (r.get("our_unit") or "~", r.get("employee_name") or "",
                                   r["_at"] or datetime.max))
    for r in track_rows:
        r.pop("_at", None)
    for v in marks_by_eid.values():
        v.sort(key=lambda m: m[0] or datetime.max)

    # The file for the day, matched to the API by NAME (the parity check's rule:
    # the full folded name, else surname + first name when exactly one answers).
    by_full: dict[str, set] = defaultdict(set)
    by_two: dict[str, set] = defaultdict(set)
    for eid, rec in report_by_eid.items():
        full, two = _keys(rec["row"].get("employee_name") or ename(emps.get(eid) or {}))
        by_full[full].add(eid)
        by_two[two].add(eid)
    dir_full: dict[str, set] = defaultdict(set)
    dir_two: dict[str, set] = defaultdict(set)
    for eid, e in emps.items():
        full, two = _keys(ename(e))
        dir_full[full].add(eid)
        dir_two[two].add(eid)

    def match(name: str) -> tuple[Optional[str], str]:
        full, two = _keys(name)
        for idx in (by_full, by_two, dir_full, dir_two):
            ids = idx.get(full if idx in (by_full, dir_full) else two) or set()
            if len(ids) == 1:
                return next(iter(ids)), "ok"
            if len(ids) > 1:
                return None, f"ambiguous ({len(ids)})"
        return None, "not found"

    now = datetime.now(verifix.TZ).replace(tzinfo=None)
    formula = verifix_live._formula(db)
    cmp_rows: list[dict] = []
    totals: Counter[str] = Counter()
    batch = db.query(AttendanceBatch).filter(AttendanceBatch.date == DAY).first()
    file_rows = (db.query(AttendanceBatchRow).filter(AttendanceBatchRow.batch_id == batch.id)
                 .order_by(AttendanceBatchRow.id).all()) if batch else []
    for fr in file_rows:
        code = (fr.verifix_code or "").strip()
        cell = cells.get(verifix._code_key(code)) if code else None
        unit = units.get(cell["manager_id"]) if cell else None
        f_in, f_out = _file_clock(fr.clock_in_out)
        eid, how = match(fr.worker_name or "")
        totals["file_rows"] += 1
        row: dict[str, Any] = {
            "file.unit": (unit or {}).get("name"), "file.shift": (unit or {}).get("shift"),
            "file.cell": code or None, "file.worker": fr.worker_name,
            "file.job": fr.job_title, "file.schedule": fr.schedule,
            "file.clock": fr.clock_in_out, "file.in": f_in, "file.out": f_out,
            "file.hours": float(fr.hours_worked) if fr.hours_worked is not None else None,
            "file.status": fr.status,
            "file.edited": bool(fr.edited), "file.manual": bool(fr.manual),
            "match": how, "employee_id": eid,
        }
        if eid:
            totals["matched"] += 1
            rec = report_by_eid.get(eid)
            e = emps.get(eid) or {}
            row.update({"api.name": (rec or {}).get("row", {}).get("employee_name") or ename(e),
                        "api.status": e.get("status"), **place(eid)})
            d = (rec or {}).get("days") or []
            d0 = d[0] if d else None
            if d0:
                r_in = verifix_live._dt(d0.get("input_time"))
                r_out = verifix_live._dt(d0.get("output_time"))
                row.update({
                    "report.input_time": d0.get("input_time"),
                    "report.output_time": d0.get("output_time"),
                    "report.begin_time": d0.get("begin_time"), "report.end_time": d0.get("end_time"),
                    "report.day_kind": d0.get("day_kind"), "report.plan_time": d0.get("plan_time"),
                    "report.facts": "; ".join(
                        f"{(kinds.get(str(f.get('time_kind_id'))) or {}).get('name') or f.get('time_kind_id')}"
                        f"={f.get('fact_value')}" for f in d0.get("facts") or []),
                    "report.days_on_row": len((rec or {}).get("row", {}).get("days") or []),
                })
            else:
                r_in = r_out = None
                row["report.input_time"] = "(no day entry)" if rec else "(no report row)"
            marks = marks_by_eid.get(eid, [])
            m_in = marks[0][0] if marks else None
            m_out = marks[-1][0] if marks else None
            row.update({
                "marks.n": len(marks), "marks.first": _stamp(m_in), "marks.last": _stamp(m_out),
                "marks.types": ", ".join(f"{k}:{v}" for k, v in Counter(m[1] or "∅" for m in marks).items()),
                "marks.all": " · ".join(f"{_stamp(m[0]) or '?'}{(' ' + m[1]) if m[1] else ''}" for m in marks),
            })
            # What the live page prints for this person NOW (v4.201.1 rule).
            try:
                p = verifix_live._person(DAY, now, {"name": row.get("api.name"), "job": "",
                                                    "schedule": ((rec or {}).get("row") or {}).get("schedule_name") or "",
                                                    "days": d}, [(m[0], m[1]) for m in marks if m[0]], formula)
                row.update({"live.in": _hm(p["in"]), "live.in_src": p["in_src"],
                            "live.out": _stamp(p["out"]), "live.out_src": p["out_src"],
                            "live.status": p["status"],
                            "live.hours": round(p["hours"], 2) if p["hours"] is not None else None})
            except Exception as exc:                 # pragma: no cover — one row
                row["live.status"] = f"error: {type(exc).__name__}"
            row.update({
                "Δ report_in − file_in": _mins(f_in, _hm(r_in)),
                "Δ first_mark − file_in": _mins(f_in, _hm(m_in), m_in.date() if m_in else None),
                "Δ report_out − file_out": _mins(f_out, _hm(r_out), r_out.date() if r_out else None),
                "Δ last_mark − file_out": _mins(f_out, _hm(m_out), m_out.date() if m_out else None),
                "Δ live_in − file_in": _mins(f_in, row.get("live.in")),
            })
            if f_in and r_in:
                totals["in_compared"] += 1
                totals["in_equal"] += _hm(r_in) == f_in
            if f_out and r_out:
                totals["out_compared"] += 1
                totals["out_equal"] += _hm(r_out) == f_out
            if f_in and m_in:
                totals["first_mark_earlier"] += (_mins(f_in, _hm(m_in), m_in.date()) or 0) < 0
            if f_out and m_out:
                totals["last_mark_later"] += (_mins(f_out, _hm(m_out), m_out.date()) or 0) > 0
        cmp_rows.append(row)
    cmp_rows.sort(key=lambda r: (r.get("file.unit") or "~", r.get("file.worker") or ""))

    summary = {"at": data["at"], "meta": data["meta"], "totals": dict(totals),
               "report_rows": len(report_rows), "track_rows": len(track_rows)}

    def workbook(with_tracks: bool, with_rest: bool) -> BytesIO:
        wb = Workbook()
        ws = wb.active
        ws.title = "Xulosa"
        ws.append(["Verifix API — hamma narsa", DAY.strftime("%d.%m.%Y"),
                   f"olindi {data['at'].strftime('%d.%m.%Y %H:%M')}"])
        ws["A1"].font = Font(bold=True, size=13)
        ws.append([])
        ws.append(["form", "path", "body", "pages", "rows", "ms", "ok", "code", "message"])
        for c in ws[3]:
            c.font = Font(bold=True)
        for m in data["meta"]:
            ws.append([m.get(k) for k in ("form", "path", "body", "pages", "rows", "ms", "ok", "code", "message")])
        ws.append([])
        ws.append(["Fayl bilan solishtirish (Solishtirish varag'i)"])
        ws[ws.max_row][0].font = Font(bold=True)
        for k, v in sorted(totals.items()):
            ws.append([k, v])
        ws.append([])
        ws.append(["Har bir forma qaysi maydonlarni yuboradi"])
        ws[ws.max_row][0].font = Font(bold=True)
        ws.append(["form", "field", "filled", "of", "sample"])
        for c in ws[ws.max_row]:
            c.font = Font(bold=True)
        for key, _p, _b, _l in FORMS:
            for c in _census(forms.get(key, [])):
                ws.append([key, c["field"], c["filled"], c["of"], _cell(c["sample"])])
        for col, w in zip("ABCDEFGHI", (18, 28, 60, 8, 8, 8, 6, 12, 40)):
            ws.column_dimensions[col].width = w
        if with_rest:
            _sheet(wb, "Solishtirish", cmp_rows)
            _sheet(wb, "Hisobot", report_rows,
                   first=("employee_id", "employee_name", "our_unit", "our_cell"))
        if with_tracks:
            _sheet(wb, "Belgilar", track_rows,
                   first=("employee_id", "employee_name", "our_unit", "our_cell"))
        if with_rest:
            _sheet(wb, "Xodimlar", [{"employee_id": str(e.get("employee_id") or ""),
                                     "name": ename(e), **place(str(e.get("employee_id") or "")),
                                     **_flat(e, skip=("employee_id",))}
                                    for e in forms.get("employees", [])])
            for key, title in (("divisions", "Bo'limlar"), ("jobs", "Lavozimlar"),
                               ("schedules", "Jadvallar"), ("time_kinds", "Vaqt turlari"),
                               ("locations", "Joylar")):
                _sheet(wb, title, [_flat(r) for r in forms.get(key, [])])
        bio = BytesIO()
        wb.save(bio)
        return bio

    if tracks_apart:
        return workbook(False, True), workbook(True, False), summary
    return workbook(True, True), None, summary


# ── delivery ──────────────────────────────────────────────────────────────────

def _caption(summary: dict) -> str:
    t = summary["totals"]
    forms = " · ".join(f"{m['form']} {m['rows']}" + ("" if m.get("ok") else f" ({m.get('code')})")
                       for m in summary["meta"])
    lines = [f"📦 Verifix API — {DAY.strftime('%d.%m.%Y')}, hamma narsa.",
             f"Formalar: {forms}.",
             f"Fayl: {t.get('file_rows', 0)} qator, {t.get('matched', 0)} tasi API'da topildi. "
             f"Hisobot kelish vaqti faylga teng: {t.get('in_equal', 0)}/{t.get('in_compared', 0)}; "
             f"ketish: {t.get('out_equal', 0)}/{t.get('out_compared', 0)}. "
             f"Birinchi belgi fayldan erta: {t.get('first_mark_earlier', 0)}; "
             f"oxirgi belgi fayldan kech: {t.get('last_mark_later', 0)}.",
             "Varaqlar: Xulosa (formalar, maydonlar) · Solishtirish (fayl × hisobot × belgilar × "
             "jonli sahifa) · Hisobot · Belgilar · Xodimlar · Bo'limlar · Lavozimlar · Jadvallar · "
             "Vaqt turlari · Joylar."]
    return "\n".join(lines)[:1000]


def _post_document(chat_id: int, name: str, blob: bytes, caption: str) -> None:
    r = requests.post(
        f"https://api.telegram.org/bot{settings.telegram_bot_token}/sendDocument",
        data={"chat_id": chat_id, "caption": caption},
        files={"document": (name, blob, MIME)},
        timeout=300)
    j = r.json()
    if not j.get("ok"):
        raise RuntimeError(j.get("description") or f"HTTP {r.status_code}")


def send(db: Session, chat_id: int, *_window) -> int:
    """Fetch, build and DM. Returns how many files went. Raises on a failure,
    so `_send_report_once` records the attempt and retries on the next boot."""
    cfg = verifix.config(db, with_password=True)
    if not (cfg["login"] and cfg.get("password") and cfg["filial_id"]):
        raise RuntimeError("Verifix is not configured")
    data = collect(cfg, time.monotonic() + BUDGET_S)
    if not any(m.get("rows") for m in data["meta"]):
        first = data["meta"][0] if data["meta"] else {}
        raise RuntimeError(f"nothing came back: {first.get('code')} {first.get('message')}")
    one, _, summary = build(db, data)
    blob = one.getvalue()
    stamp = DAY.strftime("%Y-%m-%d")
    if len(blob) <= SPLIT_BYTES:
        _post_document(chat_id, f"verifix-api-{stamp}.xlsx", blob, _caption(summary))
        return 1
    rest, tracks, summary = build(db, data, tracks_apart=True)
    _post_document(chat_id, f"verifix-api-{stamp}.xlsx", rest.getvalue(), _caption(summary))
    _post_document(chat_id, f"verifix-api-{stamp}-belgilar.xlsx", tracks.getvalue(),
                   f"📦 Belgilar (track$list) — {DAY.strftime('%d.%m.%Y')} 00:00 → "
                   f"{(DAY + timedelta(days=1)).strftime('%d.%m')} {TRACKS_END.strftime('%H:%M')}, "
                   f"{summary['track_rows']} ta.")
    return 2
