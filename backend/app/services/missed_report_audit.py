"""One-shot: which leader-days never got their day report, and why.

⚠ TEMPORARY (2026-09-22). A leader reported that after closing their last task
the final report never came, and that they could not object to the AI's
verdict. Reading the code found one way that happens:

  1. a leader closes a photo task, the AI reviews it, and the drain tries to
     send the day report (`leader_ai.report_finished`);
  2. the DAY is still open — another task, since 20 Sep usually an automatic
     check (#8 at 17:00 / 06:00), has not closed yet — so
     `routers.leaders.build_report_row` answers None and
     `leader_reports.send_for_uid` PARKS the key as «report no longer exists»;
  3. the day closes later with no AI review behind its last task, so no drain
     pass touches the key again, and `leader_ai.sweep_unreported` skips it
     because a ledger row (the park) now exists. The report is never sent, and
     the DM's button — the leader's only way to the report page and its
     «Norozilik bildirish» — never arrives.

This module collects, for every bot leader-day from `DATE_FROM` to today, what
actually happened to its report so the claim can be checked on production data
rather than on a copy: the ledger row, the park and its logged reason, when the
day closed and which task closed it last, the AI reviews and objections behind
it, and whether the leader's bell holds a report notification. It READS only;
`startup.report_missed_day_reports` DMs it once as a summary, an .xlsx and the
full .json. Delete this module and its startup job once the files have landed.
"""
from __future__ import annotations

import io
import json
import logging
import time
from collections import Counter, defaultdict
from datetime import datetime, timedelta, timezone

import requests
from sqlalchemy.orm import Session

from app.config import settings
from app.identity import profile_holders
from app.models import (
    ActionLog, Cell, LeaderAiDispute, LeaderAiReview, LeaderDayReport,
    LeaderTaskDay, LeaderTaskEntry, Manager, Notification, RoleProfile,
)
from app.services import leader_ai, leader_tasks

log = logging.getLogger(__name__)

TZ = timezone(timedelta(hours=5))           # the plant's wall clock
DATE_FROM = "2026-09-01"                    # overlaps the local 11 Sep copy
AUTO_CHECKS_FROM = "2026-09-20"             # leader_auto.AUTO_FROM
REPORT_NKEYS = ("leader_day_clean", "leader_day_flagged", "leader_day_corrected")
PARK_OPEN = "report no longer exists"       # what send_for_uid logs for an open day
_CHUNK = 800

_API = "https://api.telegram.org"
SEND_RETRIES = 3

# ── the verdict on one leader-day's report ───────────────────────────────────
SENT = "sent"
OPEN = "day still open"
WAITING = "waiting for AI review"
MISSED_PARKED = "MISSED: tried while the day was open, never retried"
MISSED_NO_REVIEW = "MISSED: no AI-reviewed task, nothing sends it"
MISSED_FINISHED = "MISSED: finished, no ledger row"
PARKED_UNKNOWN = "parked, reason not logged"
PARKED_OK = "not reported on purpose"
NOT_AUTO = "outside the automatic regime"
MISSED = (MISSED_PARKED, MISSED_NO_REVIEW, MISSED_FINISHED)


def _aware(v):
    if v is None:
        return None
    return v if v.tzinfo else v.replace(tzinfo=timezone.utc)


def _t(v) -> str | None:
    v = _aware(v)
    return v.astimezone(TZ).strftime("%Y-%m-%d %H:%M") if v else None


def _chunks(seq):
    seq = list(seq)
    for i in range(0, len(seq), _CHUNK):
        yield seq[i:i + _CHUNK]


def _closer(reason: str | None) -> str:
    r = reason or ""
    if r.startswith(leader_tasks.AUTO_PREFIX):
        return "auto check"
    if r.startswith(leader_tasks.MISSED_PREFIX):
        return "deadline (nothing filed)"
    return "leader / deadline submit"


def collect(db: Session, date_from: str = DATE_FROM, now: datetime | None = None) -> dict:
    now = now or datetime.now(timezone.utc)
    today = now.astimezone(TZ).date().isoformat()

    D, E, R = LeaderTaskDay, LeaderTaskEntry, LeaderAiReview
    days = (db.query(D.id, D.leader_id, D.manager_id, D.date, D.cell_id,
                     D.closed_at, D.completion)
            .filter(D.date >= date_from, D.date <= today)
            .order_by(D.date, D.manager_id, D.leader_id, D.id).all())
    day_ids = [d.id for d in days]

    entries: dict[int, list] = defaultdict(list)
    for ch in _chunks(day_ids):
        for e in (db.query(E.id, E.day_id, E.task_id, E.done, E.reason,
                           E.saved_at, E.closed_at)
                  .filter(E.day_id.in_(ch)).all()):
            entries[e.day_id].append(e)
    ref_day = {f"bot:{e.id}": e.day_id for es in entries.values() for e in es}

    reviews: dict[int, list] = defaultdict(list)
    disputes: dict[int, list] = defaultdict(list)
    for ch in _chunks(ref_day):
        for r in (db.query(R.ref, R.task_id, R.status, R.attempts, R.resolution,
                           R.created_at, R.reviewed_at)
                  .filter(R.ref.in_(ch)).all()):
            reviews[ref_day[r.ref]].append(r)
        for d in (db.query(LeaderAiDispute.ref, LeaderAiDispute.status,
                           LeaderAiDispute.requested_at)
                  .filter(LeaderAiDispute.ref.in_(ch)).all()):
            disputes[ref_day[d.ref]].append(d)

    keys = {f"bot:{i}": i for i in day_ids}
    uids = {f"bot-{i}": i for i in day_ids}
    ledger: dict[int, object] = {}
    for ch in _chunks(keys):
        for L in (db.query(LeaderDayReport.report_key, LeaderDayReport.sends,
                           LeaderDayReport.score_sent, LeaderDayReport.first_sent_at,
                           LeaderDayReport.last_sent_at)
                  .filter(LeaderDayReport.report_key.in_(ch)).all()):
            ledger[keys[L.report_key]] = L
    parks: dict[int, tuple] = {}
    for ch in _chunks(uids):
        for a in (db.query(ActionLog.target_id, ActionLog.reason, ActionLog.created_at)
                  .filter(ActionLog.action == "report.parked",
                          ActionLog.target_id.in_(ch))
                  .order_by(ActionLog.created_at).all()):
            parks.setdefault(uids[a.target_id], (a.reason, a.created_at))

    # The leader's bell: notify_profile writes one row per report it sends, so
    # this says what the LEADER was actually handed, independent of the ledger.
    bell: dict[tuple, list] = defaultdict(list)
    start_utc = datetime.fromisoformat(date_from).replace(tzinfo=TZ)
    for n in (db.query(Notification.recipient_profile, Notification.nkey,
                       Notification.params, Notification.created_at)
              .filter(Notification.nkey.in_(REPORT_NKEYS),
                      Notification.created_at >= start_utc,
                      Notification.recipient_profile.like("leader:%")).all()):
        try:
            lid = int(str(n.recipient_profile).split(":")[1])
        except (IndexError, ValueError):
            continue
        dstr = str((n.params or {}).get("date") or "")
        bell[(lid, dstr[:10])].append((n.nkey, dstr, n.created_at))

    lids = {d.leader_id for d in days}
    names = {p.id: p.name for p in db.query(RoleProfile.id, RoleProfile.name)
             .filter(RoleProfile.id.in_(lids or [0])).all()}
    units = {m.id: (m.name, m.shift) for m in db.query(Manager.id, Manager.name, Manager.shift).all()}
    cells = {c.id: c.verifix_code for c in db.query(Cell.id, Cell.verifix_code).all()}
    per_task = leader_tasks.per_task_units(db)
    holders: dict[int, int] = {}
    for lid in lids:
        try:
            holders[lid] = len(profile_holders(db, f"leader:{lid}"))
        except Exception:
            db.rollback()
            holders[lid] = -1
    window = leader_ai.auto_window_start()

    rows = []
    for d in days:
        unit, shift = units.get(d.manager_id, (f"#{d.manager_id}", None))
        es = entries.get(d.id, [])
        rs = reviews.get(d.id, [])
        closed = [e for e in es if e.closed_at is not None]
        last = max(closed, key=lambda e: _aware(e.closed_at)) if closed else None
        unfinished = [r for r in rs if r.status in ("pending", "error")
                      and (r.attempts or 0) < leader_ai.MAX_ATTEMPTS]
        flagged = [r for r in rs if r.status == "flagged" and r.resolution != "approved"]
        rev_times = [_aware(r.reviewed_at) for r in rs if r.reviewed_at]
        last_rev = max(rev_times) if rev_times else None
        L = ledger.get(d.id)
        park_reason, park_at = parks.get(d.id, (None, None))
        cell = cells.get(d.cell_id) if d.cell_id else None
        bells = [b for b in bell.get((d.leader_id, d.date), [])
                 if not cell or cell in b[1]]
        auto = leader_ai.in_auto_regime(d.date, shift)

        if L is not None and (L.sends or 0) > 0:
            verdict = SENT
        elif d.closed_at is None:
            verdict = OPEN
        elif L is not None:
            if park_reason == PARK_OPEN:
                verdict = MISSED_PARKED
            elif park_reason:
                verdict = PARKED_OK
            else:
                verdict = PARKED_UNKNOWN
        elif not auto:
            verdict = NOT_AUTO
        elif unfinished:
            verdict = WAITING
        elif not rs:
            verdict = MISSED_NO_REVIEW
        else:
            verdict = MISSED_FINISHED

        rows.append({
            "date": d.date, "shift": shift, "unit": unit, "manager_id": d.manager_id,
            "leader": names.get(d.leader_id, f"#{d.leader_id}"), "leader_id": d.leader_id,
            "cell": cell, "day_id": d.id, "uid": f"bot-{d.id}",
            "per_task_unit": d.manager_id in per_task, "auto_regime": auto,
            "verdict": verdict,
            "day_closed_at": _t(d.closed_at),
            "score": float(d.completion) if d.completion is not None else None,
            "tasks_entered": len(es), "tasks_closed": len(closed),
            "last_task": last.task_id if last else None,
            "last_task_closed_at": _t(last.closed_at) if last else None,
            "last_task_closed_by": _closer(last.reason) if last else None,
            "last_task_has_ai_review": bool(last and any(r.ref == f"bot:{last.id}"
                                                         for r in rs)),
            "ai_reviews": len(rs), "ai_unfinished": len(unfinished),
            "ai_rejected": len(flagged),
            "ai_rejected_tasks": ", ".join(f"#{r.task_id}" for r in sorted(flagged, key=lambda r: r.task_id)),
            "last_review_at": _t(last_rev),
            "reviews_done_before_day_closed": bool(
                last_rev and d.closed_at and last_rev <= _aware(d.closed_at)),
            "objections": len(disputes.get(d.id, [])),
            "ledger": ("none" if L is None else
                       "parked" if (L.sends or 0) == 0 else f"sent x{L.sends}"),
            "report_sent_at": _t(L.first_sent_at) if L is not None and (L.sends or 0) > 0 else None,
            "park_reason": park_reason, "parked_at": _t(park_at),
            "parked_before_day_closed": bool(park_at and d.closed_at
                                             and _aware(park_at) < _aware(d.closed_at)),
            "leader_bell_rows": len(bells),
            "leader_bell_first_at": _t(min((b[2] for b in bells), default=None)),
            "leader_telegram_accounts": holders.get(d.leader_id),
            "in_sweep_window": d.date >= window,
        })

    return {"generated_at": _t(now), "date_from": date_from, "date_to": today,
            "sweep_window_from": window, "summary": _summary(rows), "rows": rows}


def _summary(rows: list[dict]) -> dict:
    closed = [r for r in rows if r["verdict"] != OPEN]
    by = Counter(r["verdict"] for r in rows)
    missed = [r for r in rows if r["verdict"] in MISSED]
    pre = [r for r in closed if r["date"] < AUTO_CHECKS_FROM]
    post = [r for r in closed if r["date"] >= AUTO_CHECKS_FROM]

    def share(sub):
        m = sum(1 for r in sub if r["verdict"] in MISSED)
        return {"closed": len(sub), "missed": m,
                "missed_pct": round(100 * m / len(sub), 1) if sub else None}

    per_date = []
    for (dt, sh) in sorted({(r["date"], r["shift"]) for r in rows},
                           key=lambda k: (k[0], k[1] or 0)):
        sub = [r for r in rows if r["date"] == dt and r["shift"] == sh]
        c = Counter(r["verdict"] for r in sub)
        per_date.append({
            "date": dt, "shift": sh, "days": len(sub), "open": c[OPEN],
            "sent": c[SENT], "missed": sum(c[k] for k in MISSED),
            "missed_parked": c[MISSED_PARKED], "missed_no_review": c[MISSED_NO_REVIEW],
            "missed_finished": c[MISSED_FINISHED], "waiting": c[WAITING],
            "parked_on_purpose": c[PARKED_OK], "parked_unknown": c[PARKED_UNKNOWN],
        })
    mp = [r for r in missed if r["verdict"] == MISSED_PARKED]
    leaders = Counter((r["leader"], r["unit"]) for r in missed)
    return {
        "by_verdict": dict(by.most_common()),
        "before_auto_checks": share(pre),
        "since_auto_checks": share(post),
        "missed_parked_last_closed_by": dict(Counter(
            r["last_task_closed_by"] or "—" for r in mp).most_common()),
        "missed_parked_last_task": dict(Counter(
            f"#{r['last_task']}" for r in mp).most_common()),
        "missed_parked_before_close": sum(1 for r in mp if r["parked_before_day_closed"]),
        "missed_with_ai_rejections": sum(1 for r in missed if r["ai_rejected"]),
        "ai_rejected_tasks_on_missed_days": sum(r["ai_rejected"] for r in missed),
        "objections_on_missed_days": sum(r["objections"] for r in missed),
        "objections_on_sent_days": sum(r["objections"] for r in rows if r["verdict"] == SENT),
        "missed_but_leader_bell_has_report": sum(1 for r in missed if r["leader_bell_rows"]),
        "sent_but_no_leader_bell": sum(1 for r in rows if r["verdict"] == SENT
                                       and not r["leader_bell_rows"]),
        "sent_to_leader_without_telegram": sum(1 for r in rows if r["verdict"] == SENT
                                               and r["leader_telegram_accounts"] == 0),
        "parked_on_purpose_reasons": dict(Counter(
            r["park_reason"] for r in rows if r["verdict"] == PARKED_OK).most_common()),
        "top_leaders_missed": [{"leader": l, "unit": u, "missed": n}
                               for (l, u), n in leaders.most_common(15)],
        "per_date": per_date,
    }


# ── the workbook ─────────────────────────────────────────────────────────────

COLS = [
    ("date", 11), ("shift", 6), ("unit", 24), ("leader", 26), ("cell", 7),
    ("verdict", 44), ("day_closed_at", 16), ("score", 7),
    ("tasks_entered", 8), ("tasks_closed", 8), ("last_task", 7),
    ("last_task_closed_at", 16), ("last_task_closed_by", 20),
    ("last_task_has_ai_review", 9), ("ai_reviews", 8), ("ai_unfinished", 8),
    ("ai_rejected", 8), ("ai_rejected_tasks", 16), ("last_review_at", 16),
    ("reviews_done_before_day_closed", 10), ("objections", 8), ("ledger", 10),
    ("report_sent_at", 16), ("park_reason", 26), ("parked_at", 16),
    ("parked_before_day_closed", 10), ("leader_bell_rows", 8),
    ("leader_bell_first_at", 16), ("leader_telegram_accounts", 9),
    ("per_task_unit", 8), ("auto_regime", 8), ("in_sweep_window", 8),
    ("uid", 12), ("day_id", 8), ("leader_id", 8), ("manager_id", 8),
]


def build_workbook(rep: dict) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    bold = Font(bold=True)
    head_fill = PatternFill("solid", fgColor="E8E3D6")
    miss_fill = PatternFill("solid", fgColor="FBD5D5")
    s = rep["summary"]

    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    ws.append([f"Leader day reports {rep['date_from']} → {rep['date_to']}"
               f" (generated {rep['generated_at']} Tashkent)"])
    ws["A1"].font = Font(bold=True, size=13)
    ws.append([f"The catch-up sweep only looks back to {rep['sweep_window_from']}."])
    ws.append([])
    ws.append(["Verdict", "Leader-days"])
    for c in ws[ws.max_row]:
        c.font = bold
    for k, v in s["by_verdict"].items():
        ws.append([k, v])
    ws.append([])
    for label, part in (("Before 20 Sep (no auto checks)", s["before_auto_checks"]),
                        ("From 20 Sep (auto checks)", s["since_auto_checks"])):
        ws.append([label, f"{part['missed']} of {part['closed']} closed days missed"
                          f" ({part['missed_pct']}%)" if part["closed"] else "—"])
    ws.append([])
    ws.append(["«Tried while open» days — what closed the day last", ""])
    ws[ws.max_row][0].font = bold
    for k, v in s["missed_parked_last_closed_by"].items():
        ws.append([k, v])
    ws.append(["…parked BEFORE the day closed", s["missed_parked_before_close"]])
    ws.append([])
    for label, key in (
            ("Missed days carrying AI rejections", "missed_with_ai_rejections"),
            ("AI-rejected tasks on missed days", "ai_rejected_tasks_on_missed_days"),
            ("Objections filed on missed days", "objections_on_missed_days"),
            ("Objections filed on sent days", "objections_on_sent_days"),
            ("Missed, but the leader's bell has a report row", "missed_but_leader_bell_has_report"),
            ("Sent, but no report row in the leader's bell", "sent_but_no_leader_bell"),
            ("Sent to a leader with no Telegram account", "sent_to_leader_without_telegram")):
        ws.append([label, s[key]])
    ws.column_dimensions["A"].width = 58
    ws.column_dimensions["B"].width = 44

    def table(name, header, data, widths, fill_if=None):
        t = wb.create_sheet(name)
        t.append(header)
        for c in t[1]:
            c.font = bold
            c.fill = head_fill
            c.alignment = Alignment(wrap_text=True, vertical="top")
        for vals in data:
            t.append(vals)
            if fill_if and fill_if(vals):
                for c in t[t.max_row]:
                    c.fill = miss_fill
        for i, w in enumerate(widths, start=1):
            t.column_dimensions[t.cell(row=1, column=i).column_letter].width = w
        t.freeze_panes = "A2"
        if t.max_row > 1:
            t.auto_filter.ref = t.dimensions
        return t

    pd_cols = ["date", "shift", "days", "open", "sent", "missed", "missed_parked",
               "missed_no_review", "missed_finished", "waiting",
               "parked_on_purpose", "parked_unknown"]
    table("By date", pd_cols, [[r[c] for c in pd_cols] for r in s["per_date"]],
          [11, 6, 7, 7, 7, 8, 9, 9, 9, 8, 9, 9], lambda v: v[5])

    keys = [c for c, _ in COLS]
    widths = [w for _, w in COLS]
    missed = [r for r in rep["rows"] if r["verdict"] in MISSED]
    table("Missed", keys, [[r[k] for k in keys] for r in missed], widths)
    table("All days", keys, [[r[k] for k in keys] for r in rep["rows"]], widths,
          lambda v: v[5] in MISSED)
    table("Leaders", ["leader", "unit", "missed days"],
          [[x["leader"], x["unit"], x["missed"]] for x in s["top_leaders_missed"]],
          [30, 30, 12])

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ── delivery ─────────────────────────────────────────────────────────────────

def _clean(text: str) -> str:
    token = settings.telegram_bot_token or ""
    return text.replace(token, "***") if token else text


def _post(method: str, data: dict, files: dict | None = None, timeout: int = 60) -> None:
    last = ""
    for attempt in range(1, SEND_RETRIES + 1):
        wait = 0
        try:
            r = requests.post(f"{_API}/bot{settings.telegram_bot_token}/{method}",
                              data=data, files=files, timeout=timeout)
            body = r.json()
            if body.get("ok"):
                return
            last = body.get("description") or f"HTTP {r.status_code}"
            wait = int(((body.get("parameters") or {}).get("retry_after")) or 0)
        except Exception as exc:
            last = _clean(f"{type(exc).__name__}: {exc}")
        if attempt < SEND_RETRIES:
            time.sleep(max(wait, 3 * attempt))
    raise RuntimeError(_clean(last or f"{method} failed")[:300])


def _pct(part: dict) -> str:
    return "—" if part["missed_pct"] is None else f"{part['missed_pct']}%"


def summary_text(rep: dict) -> str:
    s = rep["summary"]
    b, a = s["before_auto_checks"], s["since_auto_checks"]
    by = s["by_verdict"]
    lines = [
        f"Leader day reports {rep['date_from']} → {rep['date_to']}: were they sent?",
        "",
        f"Sent: {by.get(SENT, 0)}",
        f"MISSED, tried while the day was still open: {by.get(MISSED_PARKED, 0)}",
        f"MISSED, no AI-reviewed task: {by.get(MISSED_NO_REVIEW, 0)}",
        f"MISSED, finished but never sent: {by.get(MISSED_FINISHED, 0)}",
        f"Not reported on purpose (excluded, voided…): {by.get(PARKED_OK, 0)}",
        f"Outside the automatic regime: {by.get(NOT_AUTO, 0)}",
        f"Parked, reason not logged: {by.get(PARKED_UNKNOWN, 0)}",
        f"Waiting for AI / day still open: {by.get(WAITING, 0)} / {by.get(OPEN, 0)}",
        "",
        f"Before 20 Sep: {b['missed']} of {b['closed']} closed days missed ({_pct(b)}).",
        f"From 20 Sep: {a['missed']} of {a['closed']} closed days missed ({_pct(a)}).",
        "",
        "What closed the day last, on the «tried while open» days: " + ", ".join(
            f"{k} {v}" for k, v in s["missed_parked_last_closed_by"].items()),
        f"AI-rejected tasks the leader was never sent a report for: "
        f"{s['ai_rejected_tasks_on_missed_days']} "
        f"(on {s['missed_with_ai_rejections']} days; objections filed on them: "
        f"{s['objections_on_missed_days']}).",
        "",
        "Excel: «By date», then «Missed» — one row per leader-day with its timeline.",
    ]
    return "\n".join(lines)


def send(db: Session, chat_id: int, *_window) -> int:
    if not settings.telegram_bot_token:
        raise RuntimeError("telegram bot token not configured")
    rep = collect(db)
    db.rollback()
    xlsx = build_workbook(rep)
    blob = json.dumps(rep, ensure_ascii=False, indent=1, default=str).encode()
    _post("sendMessage", {"chat_id": chat_id, "text": summary_text(rep)[:4000]})
    _post("sendDocument", {"chat_id": chat_id,
                           "caption": "Leader day reports — sent or missed, per leader-day."},
          files={"document": ("missed-day-reports.xlsx", xlsx,
                              "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
          timeout=300)
    _post("sendDocument", {"chat_id": chat_id, "caption": "Same report, full data."},
          files={"document": ("missed-day-reports.json", blob, "application/json")},
          timeout=300)
    return 3
