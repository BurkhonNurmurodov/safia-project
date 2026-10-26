"""The retroactive one-shot of the catalog follow (2026-10-26, WP6 of
docs/plan-catalog-follows-cells.md): `catalog_follow_retro.run` against a
scratch Postgres (`SAFIA_TEST_PG`, see `pg_scratch`) with
`app.database.SessionLocal` pointed at it — the engine's world (`_World` of
`test_catalog_follow_engine`), plus the evidence the pass reads: attendance
rows, closed live-day rosters and the action register.

What it proves (spec §4): T10 through the retro path (a single hop found from
the evidence, h = 8 Oct, the old unit importing its catalog inside the window),
T11 (every named kind from its own fixture writes nothing), a T12-style dry run
leaving every table as it was, T16 (a second run is a no-op — by the flag, and
by the move records when the flag is gone), the flag → RECORD → DM flow with
the Telegram sender mocked, and the pass that finds nothing.

    SAFIA_TEST_PG=postgresql://me@localhost:5432/postgres \\
        ./venv/bin/python -m unittest tests.test_catalog_follow_retro
"""
from __future__ import annotations

import io
import json
import unittest
from datetime import date, datetime, time, timedelta, timezone
from unittest import mock

try:
    from tests.pg_scratch import SERVER
    from tests.test_catalog_follow_engine import C, TZ, X, Y, Z, _days, _World
except ImportError:  # run from inside tests/
    from pg_scratch import SERVER
    from test_catalog_follow_engine import C, TZ, X, Y, Z, _days, _World

from app.services import catalog_follow as cf
from app.services import catalog_follow_retro as retro
from app.services import pp_catalog

TODAY = date(2026, 10, 22)
NOW = datetime(2026, 10, 22, 6, 0, tzinfo=timezone.utc)
H = date(2026, 10, 14)                 # a landed brigadir change's day (links era)


class _Retro(_World):
    """The engine's world, plus the evidence tables, emptied between tests."""

    def setUp(self):
        from app import models as m
        s = self.pg.Session()
        for t in (m.Attendance, m.LiveProjection, m.ActionLog):
            s.query(t).delete()
        s.commit()
        s.close()
        super().setUp()

    # ── the evidence ─────────────────────────────────────────────────────────

    def att(self, cell, unit, a, b, n=3):
        """`n` named workers with hours at the cell's code, filed under `unit`
        every day of [a, b]."""
        from app.models import Attendance
        for d in _days(a, b):
            for i in range(n):
                self.db.add(Attendance(manager_id=unit, date=d, worker_name=f"W{cell.verifix_code}-{i}",
                                       hours_worked=8, verifix_code=cell.verifix_code))

    def roster(self, unit, d, *cells):
        from app.models import LiveProjection
        self.db.add(LiveProjection(manager_id=unit, day=d, digest=f"{unit}{d}",
                                   roster=[[f"W{c.verifix_code}", c.verifix_code, "", True]
                                           for c in cells]))

    def applied(self, cell, to, eff, letter=None):
        """A brigadir change that has LANDED: the event applied, the column
        moved — what `cell_links.apply_due` leaves."""
        from app.models import CellLinkEvent
        ev = CellLinkEvent(cell_id=cell.id, field="manager_id", value=to, prev=cell.manager_id,
                           effective_from=eff,
                           effective_at=datetime.combine(eff, time(8), tzinfo=TZ),
                           applied_at=datetime.combine(eff, time(8, 1), tzinfo=TZ), src="manual",
                           vfx={"letter_before": letter if letter is not None else cell.wc_group})
        self.db.add(ev)
        cell.manager_id = to
        self.db.flush()
        return ev

    # ── running it ───────────────────────────────────────────────────────────

    def run_retro(self, **kw):
        self.db.commit()
        self.forget()
        rec = retro.run(self.db, today=kw.pop("today", TODAY), now=NOW)
        self.forget()
        return rec

    def flag(self):
        from app.models import AppSetting
        row = self.db.query(AppSetting).filter_by(key=retro.FLAG).first()
        return row.value if row else None

    def stored(self):
        from app.models import AppSetting
        row = self.db.query(AppSetting).filter_by(key=retro.RECORD).first()
        return json.loads(row.value) if row else None

    def catalog_fp(self):
        """Every catalog and quantity row and every cell — what a skipped
        move must leave exactly as it was."""
        return [r for r in self.fingerprint()
                if not r[0].startswith(("pp_catalog_moves", "cell_link_events"))]

    def skipped(self, rec, vcode):
        return [e for e in rec["skipped"] if e.get("vcode") == vcode]

    def assert_named(self, kind, vcode, before, rec, *, row=True):
        got = self.skipped(rec, vcode)
        self.assertTrue(got, rec["skipped"])
        self.assertEqual({e["kind"] for e in got}, {kind}, got)
        self.assertEqual(self.catalog_fp(), before)
        self.assertEqual(rec["moved"], [])
        if row:
            from app.models import Cell, PPCatalogMove
            cid = self.db.query(Cell.id).filter(Cell.verifix_code == vcode).scalar()
            rows = self.db.query(PPCatalogMove).filter(PPCatalogMove.cell_id == cid).all()
            self.assertTrue(rows)
            for r in rows:
                self.assertEqual((r.state, r.kind), ("named", "retro"))
                self.assertIn(kind, {n.get("kind") for n in r.named or []})
        self.assertTrue(self.flag().startswith("done:0/"))
        self.assertTrue(cf.enabled(self.db))

    def simple_event_world(self, vcode="3001"):
        """X: a lone cell at the code with one line and a day of quantities;
        its brigadir change to Y landed on H."""
        self.units(X, Y)
        c = self.cell(vcode, X)
        p = self.line(X, "F001", "Tort", labor=60)
        self.daily(X, H, "F001", plan=10, actual=8)
        self.daily(X, H - timedelta(days=1), "F001", plan=10, actual=8)
        ev = self.applied(c, Y, H)
        return c, p, ev


@unittest.skipUnless(SERVER, "SAFIA_TEST_PG is not set")
class RetroPass(_Retro):
    """The moves it applies."""

    # ── T10: a single hop found from the evidence, an import in the window ───

    def test_t10_single_hop_with_an_import_in_the_window(self):
        from app.models import PPProduct
        h = date(2026, 10, 8)
        imp = date(2026, 10, 12)
        last = date(2026, 10, 14)
        self.units(X, Y)
        c = self.cell("0701", Y)                       # moved before link history: no event
        old1 = self.line(X, "F001", "Tort", labor=60, sort=0)
        old2 = self.line(X, "F002", "Keks", labor=30, sort=1)
        other = self.line(X, "F003", "Non", wc="B100", labor=12, sort=2)
        self.cell("0702", X, sap="B100")
        for day in _days(date(2026, 10, 5), last):
            self.daily(X, day, "F001", plan=10, actual=10)
            self.daily(X, day, "F002", plan=20, actual=15)
            self.daily(X, day, "F003", wc="B100", plan=5, actual=5)
        self.att(c, X, date(2026, 9, 29), h - timedelta(days=1))
        self.att(c, Y, h, date(2026, 10, 11))
        self.db.flush()
        # An ABC import on 12 Oct: frozen, every line re-created with a new id.
        pp_catalog.freeze(self.db, X, "import", day=imp)
        for q in (old1, old2, other):
            self.db.delete(q)
        self.db.flush()
        new1 = self.line(X, "F001", "Tort", labor=60, sort=0)
        new2 = self.line(X, "F002", "Keks", labor=30, sort=1)
        new3 = self.line(X, "F003", "Non", wc="B100", labor=12, sort=2)
        self.db.commit()
        lo = date(2026, 10, 5)
        hist0 = self.history([X, Y], lo, h - timedelta(days=1))
        mins0 = self.minutes([X, Y], h, last)

        rec = self.run_retro()

        self.assertEqual(rec["skipped"], [], rec["skipped"])
        [m] = rec["moved"]
        self.assertEqual((m["vcode"], m["from"], m["to"], m["h"], m["S"], m["source"]),
                         ("0701", X, Y, h.isoformat(), imp.isoformat(), "attendance"))
        self.assertEqual(m["lines"], 4)
        self.assertEqual(self.history([X, Y], lo, h - timedelta(days=1)), hist0)
        self.assertEqual(self.minutes([X, Y], h, last), mins0)
        self.assert_unique([X, Y], lo, last)
        for day in _days(h, imp - timedelta(days=1)):
            self.assertEqual(self.ids_on(Y, day), {old1.id, old2.id})
            self.assertEqual(self.ids_on(X, day), {other.id})
        for day in _days(imp, last):
            self.assertEqual(self.ids_on(Y, day), {new1.id, new2.id})
            self.assertEqual(self.ids_on(X, day), {new3.id})
        [row] = self.move_rows()
        self.assertEqual((row.kind, row.state, row.event_id, row.home_from),
                         ("retro", "moved", None, h))
        self.assertEqual(set(row.line_ids), {old1.id, old2.id, new1.id, new2.id})
        self.assertEqual(self.db.query(PPProduct).filter_by(manager_id=Y).count(), 2)
        # Impact recorded on the pass, the cell named as moved before 12 Oct.
        imp_x = [i for i in rec["impact"] if i["unit"] == X and i["day"] == h.isoformat()]
        self.assertEqual(imp_x[0]["plan"], [21.0, 1.0])
        self.assertEqual(rec["pre12"], ["0701"])
        self.assertEqual(self.flag(), "done:1/0")
        self.assertTrue(cf.enabled(self.db))
        self.assertEqual(self.stored()["counts"]["moved"], 1)

    def test_moved_before_the_floor(self):
        """Standing in Y since before 6 Oct, last seen in X before the window:
        followed from 6 Oct."""
        self.units(X, Y)
        c = self.cell("0801", Y)
        p = self.line(X, "F001", "Tort")
        self.daily(X, cf.FOLLOW_FROM, "F001", plan=6, actual=6)
        self.att(c, X, date(2026, 8, 20), date(2026, 8, 25))
        self.att(c, Y, date(2026, 9, 29), date(2026, 10, 10))
        rec = self.run_retro()
        [m] = rec["moved"]
        self.assertEqual((m["h"], m["source"]), (cf.FOLLOW_FROM.isoformat(), "floor"))
        self.assertEqual(self.ids_on(Y, cf.FOLLOW_FROM), {p.id})
        self.assertEqual(self.ids_on(X, cf.FOLLOW_FROM - timedelta(days=1)), {p.id})

    def test_roster_evidence(self):
        """The closed live day's roster names the unit a cell stood in."""
        self.units(X, Y)
        c = self.cell("0901", Y)
        p = self.line(X, "F001", "Tort")
        for d in _days(date(2026, 10, 6), date(2026, 10, 8)):
            self.roster(X, d, c)
        for d in _days(date(2026, 10, 9), date(2026, 10, 11)):
            self.roster(Y, d, c)
        rec = self.run_retro()
        [m] = rec["moved"]
        self.assertEqual((m["h"], m["source"]), ("2026-10-09", "roster"))
        self.assertEqual(self.ids_on(Y, date(2026, 10, 9)), {p.id})
        self.assertEqual(self.ids_on(X, date(2026, 10, 8)), {p.id})

    def test_landed_event_moves(self):
        c, p, ev = self.simple_event_world()
        rec = self.run_retro()
        [m] = rec["moved"]
        self.assertEqual((m["event"], m["h"], m["source"]), (ev.id, H.isoformat(), "event"))
        self.assertEqual(self.ids_on(Y, H), {p.id})
        self.assertEqual(self.ids_on(X, H - timedelta(days=1)), {p.id})
        [row] = self.move_rows()
        self.assertEqual((row.state, row.kind, row.event_id), ("moved", "retro", ev.id))
        self.assertEqual(self.flag(), "done:1/0")


@unittest.skipUnless(SERVER, "SAFIA_TEST_PG is not set")
class RetroGuards(_Retro):
    """T11: every named kind, each from its own fixture, writes nothing."""

    def test_chain(self):
        self.units(X, Y, Z)
        c = self.cell("1101", X)
        self.line(X, "F001", "Tort")
        self.applied(c, Y, date(2026, 10, 13))
        self.applied(c, Z, date(2026, 10, 15))
        self.db.commit()
        before = self.catalog_fp()
        rec = self.run_retro()
        self.assert_named("chain", "1101", before, rec)
        self.assertEqual(len(self.move_rows()), 2)          # one per event

    def test_from_unknown(self):
        """In Y since before the window, never seen in X: the orphans in X
        are not shown to be this cell's."""
        self.units(X, Y)
        c = self.cell("1201", Y)
        self.line(X, "F001", "Tort")
        self.att(c, Y, date(2026, 9, 29), date(2026, 10, 11))
        self.db.commit()
        before = self.catalog_fp()
        rec = self.run_retro()
        self.assert_named("from_unknown", "1201", before, rec, row=False)

    def test_day_unknown(self):
        self.units(X, Y)
        c = self.cell("1301", Y)
        self.line(X, "F001", "Tort")
        self.att(c, X, date(2026, 9, 29), date(2026, 9, 30))
        self.att(c, Y, date(2026, 10, 9), date(2026, 10, 11))
        self.db.commit()
        before = self.catalog_fp()
        rec = self.run_retro()
        self.assert_named("day_unknown", "1301", before, rec, row=False)

    def test_letter_ambiguous(self):
        """X keeps a cell at the code, and the orphans carry two letters."""
        self.units(X, Y)
        self.cell("1402", X, letter="A")
        c = self.cell("1401", Y)
        self.line(X, "F001", "Tort", letter="A")
        self.line(X, "F002", "Keks", letter="B", sort=1)
        self.line(X, "F003", "Non", letter="C", sort=2)
        self.att(c, X, date(2026, 9, 29), date(2026, 10, 7))
        self.att(c, Y, date(2026, 10, 8), date(2026, 10, 11))
        self.db.commit()
        before = self.catalog_fp()
        rec = self.run_retro()
        self.assert_named("letter_ambiguous", "1401", before, rec, row=False)

    def test_shared_lines(self):
        """The cell's letter is clear, but X holds an unlettered line it shared
        with the cell that stayed."""
        self.units(X, Y)
        self.cell("1502", X, letter="A")
        c = self.cell("1501", Y)
        self.line(X, "F001", "Tort", letter="A")
        self.line(X, "F002", "Keks", letter="B", sort=1)
        self.line(X, "F003", "Non", sort=2)
        self.att(c, X, date(2026, 9, 29), date(2026, 10, 7))
        self.att(c, Y, date(2026, 10, 8), date(2026, 10, 11))
        self.db.commit()
        before = self.catalog_fp()
        rec = self.run_retro()
        self.assert_named("shared_lines", "1501", before, rec, row=False)

    def test_y_has_lines(self):
        c, _p, _ev = self.simple_event_world("1601")
        self.line(Y, "F009", "Rulet", labor=30)             # nobody of Y owns it
        self.db.commit()
        before = self.catalog_fp()
        rec = self.run_retro()
        self.assert_named("y_has_lines", "1601", before, rec)

    def test_window_cells_moved(self):
        """Two cells left X at one code on one day: followed one at a time
        they would take each other's positions — neither is."""
        self.units(X, Y, Z)
        c1 = self.cell("1701", X, letter="A")
        c2 = self.cell("1702", X, letter="B")
        self.line(X, "F001", "Tort", letter="A")
        self.line(X, "F002", "Keks", letter="B", sort=1)
        self.applied(c1, Y, H)
        self.applied(c2, Z, H)
        self.db.commit()
        before = self.catalog_fp()
        rec = self.run_retro()
        self.assert_named("window_cells_moved", "1701", before, rec)
        self.assert_named("window_cells_moved", "1702", before, rec)

    def test_letters_after(self):
        from app.models import CellLinkEvent
        c, _p, _ev = self.simple_event_world("1801")
        later = date(2026, 10, 18)
        self.db.add(CellLinkEvent(cell_id=c.id, field="manager_id", value=Y, prev=Y,
                                  effective_from=later,
                                  effective_at=datetime.combine(later, time(0), tzinfo=TZ),
                                  applied_at=NOW, src="catalog_follow",
                                  vfx={"catalog_mark": True, "wc_group_was": None}))
        self.db.commit()
        before = self.catalog_fp()
        rec = self.run_retro()
        self.assert_named("letters_after", "1801", before, rec)

    def test_archived(self):
        c, _p, _ev = self.simple_event_world("1901")
        c.archived_at = datetime(2026, 10, 18, 9, tzinfo=TZ)
        self.db.commit()
        before = self.catalog_fp()
        rec = self.run_retro()
        self.assert_named("archived", "1901", before, rec)

    def test_no_code(self):
        self.units(X, Y)
        c = self.cell("2001", X, sap=None)
        self.line(X, "F001", "Tort")
        self.applied(c, Y, H)
        self.db.commit()
        before = self.catalog_fp()
        rec = self.run_retro()
        self.assert_named("no_code", "2001", before, rec)

    def test_conservation(self):
        """Y's own line read a 0/0 snapshot the arrival's figures would fill:
        minutes from nowhere — the engine's strict check stops it."""
        self.units(X, Y)
        c = self.cell("2101", X)
        self.cell("2102", Y, letter="A")
        self.line(X, "F001", "Tort", labor=60)
        self.line(Y, "F001", "Keks", labor=30, letter="A")
        self.daily(X, H, "F001", plan=100, actual=90)
        self.daily(Y, H, "F001", plan=0, actual=0)
        self.applied(c, Y, H)
        self.db.commit()
        before = self.catalog_fp()
        rec = self.run_retro()
        self.assert_named("conservation", "2101", before, rec)

    def test_volume(self):
        """Too much at once: nothing is applied. A change that moves nothing
        (a new cell given its first brigadir) is no part of the volume and is
        still settled."""
        from app.models import Cell, PPCatalogMove
        self.simple_event_world("2201")
        fresh = self.cell("2202", None, sap="B300")
        self.applied(fresh, Y, H)
        self.db.commit()
        before = self.catalog_fp()
        with mock.patch.object(retro, "MAX_CELLS", 0):
            rec = self.run_retro()
        self.assert_named("volume", "2201", before, rec)
        [row] = self.db.query(PPCatalogMove).filter(PPCatalogMove.cell_id == fresh.id).all()
        self.assertEqual(row.state, "nothing")
        self.assertIn("from_no_unit", {n["kind"] for n in row.named})
        self.assertIsNotNone(self.db.get(Cell, fresh.id))


@unittest.skipUnless(SERVER, "SAFIA_TEST_PG is not set")
class RetroFlow(_Retro):
    """The rehearsal, idempotency, the flag → record → DM flow."""

    def test_t12_dry_run_leaves_everything(self):
        self.simple_event_world("2301")
        self.db.commit()
        before = self.fingerprint()
        with mock.patch.object(retro, "RETRO_APPLY", False):
            rec = self.run_retro()
            self.assertEqual(len(rec["moved"]), 1)
            self.assertEqual(rec["result"], "dry")
            self.assertEqual(self.fingerprint(), before)
            self.assertEqual(self.flag(), "dry:1")
            self.assertFalse(cf.enabled(self.db))
            self.assertFalse(retro.pending(self.db))
            self.assertEqual(retro.dm_mode(self.db), "dry")
            stored = self.stored()
            self.assertIn("SINOV", retro.text_of(stored))
            from openpyxl import load_workbook
            wb = load_workbook(io.BytesIO(retro.workbook(stored)))
            self.assertEqual(wb.sheetnames, ["Rejada", "Qoldirildi", "Dalil", "Ta'sir"])
            self.assertEqual(wb["Rejada"]["A2"].value, "2301")
        # Applying is due once the moves are to be applied.
        self.assertTrue(retro.pending(self.db))
        rec = self.run_retro()
        self.assertEqual(self.flag(), "done:1/0")
        self.assertEqual([r.state for r in self.move_rows()], ["moved"])

    def test_t16_second_run_is_a_noop(self):
        from app.models import AppSetting
        self.simple_event_world("2401")
        self.units(Z)
        e = self.cell("2402", Z, sap="B100")
        self.line(Y, "F005", "Pirog", wc="B100")           # an orphan of nobody's move
        self.att(e, Z, date(2026, 9, 29), date(2026, 10, 11))
        rec1 = self.run_retro()
        self.assertEqual(len(rec1["moved"]), 1)
        fp = self.fingerprint()
        flag = self.flag()
        # By the flag.
        self.assertEqual(retro.run(self.db, today=TODAY, now=NOW), {"state": "skipped"})
        self.assertEqual(self.fingerprint(), fp)
        # By the move records, with the flag gone.
        self.db.query(AppSetting).filter_by(key=retro.FLAG).delete()
        self.db.commit()
        rec2 = self.run_retro()
        self.assertEqual(rec2["moved"], [])
        self.assertEqual([e["kind"] for e in rec2["skipped"]],
                         [e["kind"] for e in rec1["skipped"]])
        self.assertEqual(self.fingerprint(), fp)
        self.assertEqual((flag, self.flag()), ("done:1/1", "done:0/1"))
        # …and the sweep finds nothing left either.
        self.assertEqual(self.groups(), [])

    def test_busy_is_left_to_the_sweep(self):
        """An edit holds the old unit's catalog: the transfer is recorded busy
        (a retro one), the pass still finishes, and the sweep — on once the
        flag reads done — follows it strict on its retry."""
        from app.models import PPCatalogMove
        _c, p, ev = self.simple_event_world("2701")
        self.db.commit()
        other = self.pg.Session()
        self.others.append(other)
        self.assertTrue(pp_catalog.lock(other, [X]))
        rec = self.run_retro()
        self.assertEqual([b["vcode"] for b in rec["busy"]], ["2701"])
        self.assertEqual(self.flag(), "done:0/0")
        [row] = self.move_rows()
        self.assertEqual((row.state, row.kind, row.event_id), ("busy", "retro", ev.id))
        other.rollback()
        self.db.query(PPCatalogMove).update(
            {"next_try_at": datetime.now(timezone.utc) - timedelta(minutes=1)})
        self.db.commit()
        with mock.patch("app.services.action_log.record_system"), mock.patch.object(cf, "_dm"):
            cf.run_sweep()
        self.forget()
        [row] = self.move_rows()
        self.assertEqual((row.state, row.kind), ("moved", "retro"))
        self.assertEqual(self.ids_on(Y, H), {p.id})

    def test_a_follow_that_raises_is_retried(self):
        from app.models import PPCatalogMove
        self.simple_event_world("2801")
        with mock.patch.object(cf, "follow_group", side_effect=RuntimeError("boom")):
            rec = self.run_retro()
        self.assertEqual([(b["vcode"], b["kind"]) for b in rec["busy"]], [("2801", "failed")])
        [row] = self.db.query(PPCatalogMove).all()
        self.assertEqual((row.state, row.kind), ("failed", "retro"))
        self.assertIn("boom", row.error)
        self.assertTrue(cf.enabled(self.db))

    def test_a_pass_that_keeps_failing_stops(self):
        self.simple_event_world("2901")
        before = self.fingerprint()
        with mock.patch.object(retro, "_orphans", side_effect=RuntimeError("bad data")):
            for n in range(1, retro.TRIES + 1):
                self.assertTrue(retro.pending(self.db))
                self.assertEqual(self.run_retro(), {"state": "failed", "attempt": n})
                self.assertEqual(self.flag(), f"failed:{n}")
        self.assertFalse(retro.pending(self.db))
        self.assertEqual(retro.dm_mode(self.db), "failed")
        self.assertFalse(cf.enabled(self.db))
        self.assertEqual(self.fingerprint(), before)
        self.assertIn("bad data", retro.text_of(self.stored()))

    def test_nothing_to_move(self):
        from app import startup
        self.units(X, Y)
        self.cell("2501", X)
        self.line(X, "F001", "Tort")
        self.db.commit()
        with mock.patch("app.startup.SessionLocal", self.pg.Session), \
                mock.patch("app.scheduler.schedule_at") as arm:
            startup.schedule_catalog_follow_retro()           # the pass is due: armed
        self.assertEqual(arm.call_args.args[0], "catalog-follow-retro")
        sent = []
        with mock.patch.object(retro, "_post", lambda m, d, files=None: sent.append((m, d, files))):
            rec = self.run_retro()
            self.assertEqual((rec["moved"], rec["skipped"], rec["busy"]), ([], [], []))
            self.assertEqual(self.flag(), "done:0/0")
            self.assertTrue(cf.enabled(self.db))
            n = retro.sender(self.db, 42)
        self.assertEqual(n, 1)
        [(method, data, files)] = sent
        self.assertEqual((method, files), ("sendMessage", None))
        self.assertIn("topilmadi", data["text"])

    def test_flag_record_dm_flow(self):
        """The startup job: the pass, then the DM from the stored record —
        a DM that fails is sent again at the next boot, the pass never re-run."""
        from app import startup
        from app.models import AppSetting
        self.simple_event_world("2601")
        self.units(Z)
        self.line(Y, "F009", "Rulet", wc="B200")            # an orphan nobody's move explains
        self.att(self.cell("2604", Z, sap="B200"), Z, date(2026, 9, 29), date(2026, 10, 11))
        self.db.commit()
        self.forget()
        sent = []

        def broken(method, data, files=None):
            raise RuntimeError("telegram down")

        session = mock.patch("app.startup.SessionLocal", self.pg.Session)
        with session, mock.patch.object(retro, "_post", broken), \
                mock.patch.object(retro, "_today", lambda: TODAY):
            startup._catalog_follow_retro_job()
        self.db.expire_all()
        self.assertEqual(self.flag(), "done:1/1")
        dm = self.db.query(AppSetting).filter_by(key=retro.dm_flag("done")).first()
        self.assertEqual(dm.value, "failed:1")
        rec = self.stored()
        self.assertEqual(len(rec["moved"]), 1)
        self.assertEqual([e["vcode"] for e in rec["skipped"]], ["2604"])

        def ok(method, data, files=None):
            sent.append((method, data, files))

        with session, mock.patch.object(retro, "_post", ok), \
                mock.patch.object(retro, "run", side_effect=AssertionError("re-run")):
            # The boot arms the job for the owed DM only.
            with mock.patch("app.scheduler.schedule_at") as arm:
                startup.schedule_catalog_follow_retro()
                self.assertTrue(arm.called)
            startup._catalog_follow_retro_job()
        self.db.expire_all()
        self.assertTrue(self.db.query(AppSetting).filter_by(
            key=retro.dm_flag("done")).first().value.startswith("sent:"))
        methods = [m for m, _d, _f in sent]
        self.assertEqual(methods[0], "sendDocument")
        self.assertIn("sendMessage", methods)
        text_ = "".join(d["text"] for m, d, _f in sent if m == "sendMessage")
        self.assertIn("2601", text_)
        self.assertIn("Avtomatik kuzatuv yoqildi", text_)
        from openpyxl import load_workbook
        _name, data, _mime = sent[0][2]["document"]
        wb = load_workbook(io.BytesIO(data))
        self.assertEqual(wb.sheetnames, ["Ko'chirildi", "Qoldirildi", "Dalil", "Ta'sir"])
        self.assertEqual(wb["Ko'chirildi"]["A2"].value, "2601")
        # Nothing left to arm.
        with session, mock.patch("app.scheduler.schedule_at") as arm:
            startup.schedule_catalog_follow_retro()
            self.assertFalse(arm.called)


if __name__ == "__main__":
    unittest.main()
